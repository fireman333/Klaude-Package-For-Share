#!/usr/bin/env python3
"""
Generate a formatted Word document (.docx) from manuscript text.
Takes a markdown-like text file and produces a properly formatted .docx
suitable for journal submission.

Usage:
    python generate_docx.py input.md output.docx [--language english|chinese]
"""

import sys
import os
import re

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def create_document(language='english'):
    """Create a new document with standard formatting."""
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)

    # Set East Asian font for Chinese support
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '新細明體' if language == 'chinese' else 'Times New Roman')

    # Set paragraph formatting - double spacing
    pf = style.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = 2.0

    # Set margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    return doc


def parse_markdown(text):
    """Parse markdown-like text into structured elements."""
    elements = []
    lines = text.split('\n')
    current_paragraph = []

    for line in lines:
        stripped = line.strip()

        if not stripped:
            # Empty line = paragraph break
            if current_paragraph:
                elements.append(('paragraph', '\n'.join(current_paragraph)))
                current_paragraph = []
            continue

        # Heading detection
        if stripped.startswith('# ') and not stripped.startswith('## '):
            if current_paragraph:
                elements.append(('paragraph', '\n'.join(current_paragraph)))
                current_paragraph = []
            elements.append(('heading1', stripped[2:].strip()))
        elif stripped.startswith('## '):
            if current_paragraph:
                elements.append(('paragraph', '\n'.join(current_paragraph)))
                current_paragraph = []
            elements.append(('heading2', stripped[3:].strip()))
        elif stripped.startswith('### '):
            if current_paragraph:
                elements.append(('paragraph', '\n'.join(current_paragraph)))
                current_paragraph = []
            elements.append(('heading3', stripped[4:].strip()))
        else:
            current_paragraph.append(stripped)

    if current_paragraph:
        elements.append(('paragraph', '\n'.join(current_paragraph)))

    return elements


def add_formatted_text(paragraph, text):
    """Add text to paragraph with inline formatting (bold, italic)."""
    # Split by bold and italic markers
    parts = re.split(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*)', text)

    for part in parts:
        if part.startswith('***') and part.endswith('***'):
            # Bold italic
            run = paragraph.add_run(part[3:-3])
            run.bold = True
            run.italic = True
        elif part.startswith('**') and part.endswith('**'):
            # Bold
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            # Italic
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def generate_docx(input_text, output_path, language='english'):
    """Generate the formatted .docx document."""
    doc = create_document(language)
    elements = parse_markdown(input_text)

    for elem_type, content in elements:
        if elem_type == 'heading1':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(content)
            run.bold = True
            run.font.size = Pt(14)
            p.paragraph_format.space_before = Pt(24)
            p.paragraph_format.space_after = Pt(12)

        elif elem_type == 'heading2':
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.bold = True
            run.font.size = Pt(12)
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)

        elif elem_type == 'heading3':
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.bold = True
            run.italic = True
            run.font.size = Pt(12)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)

        elif elem_type == 'paragraph':
            p = doc.add_paragraph()
            # First line indent
            p.paragraph_format.first_line_indent = Cm(1.27)
            add_formatted_text(p, content.replace('\n', ' '))

    doc.save(output_path)
    print(f"  Document saved: {output_path}")

    # Count words
    word_count = sum(len(content.split()) for _, content in elements if _ == 'paragraph')
    print(f"  Word count: ~{word_count}")

    return word_count


def main():
    if len(sys.argv) < 3:
        print("Usage: python generate_docx.py input.md output.docx [--language english|chinese]")
        sys.exit(1)

    input_path = sys.argv[1]
    output_path = sys.argv[2]

    language = 'english'
    if '--language' in sys.argv:
        idx = sys.argv.index('--language')
        if idx + 1 < len(sys.argv):
            language = sys.argv[idx + 1]

    with open(input_path, 'r', encoding='utf-8') as f:
        text = f.read()

    generate_docx(text, output_path, language)


if __name__ == '__main__':
    main()
