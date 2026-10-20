#!/usr/bin/env python3
"""
Extract structured data from analysis output Excel files.
Reads Excel files produced by table1, univariate-forest, multivariate-interaction,
kaplan-meier, competing-risk, psm-matching, and auc-cstat skills.

Outputs a structured JSON summary for manuscript writing.

Usage:
    python extract_data.py config.json
"""

import sys
import json
import os
import re
import pandas as pd
import numpy as np
from openpyxl import load_workbook


def load_config(path):
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def safe_str(val):
    """Convert cell value to string safely."""
    if val is None:
        return ''
    return str(val).strip()


# ═══════════════════════════════════════════════════════════════════════════════
# Shared header-scanning helpers (advanced-analyses extractors, 2026-08-11)
#
# extract_weighting's docstring records the bug that motivates every one of
# these: a sheet carries a subtitle row ONLY when the research plan set
# study.title, and the production pipeline always sets one while a hand-built
# test fixture can easily omit it. Hard-coding "the header is on row 3" reads
# the subtitle row as data on any workbook missing it. Every extractor added
# below locates its header row by the text in it, never by position.
# ═══════════════════════════════════════════════════════════════════════════════

def _to_float(s):
    """Best-effort float parse of a formatted stat cell ('<0.001*', '0.625')."""
    if s is None:
        return None
    s = str(s).replace('<', '').replace('*', '').strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def find_header_row(ws, header_text, col=1, max_scan=15):
    """Row index whose given column holds exactly `header_text`."""
    for row in range(1, min(ws.max_row, max_scan) + 1):
        if safe_str(ws.cell(row, col).value) == header_text:
            return row
    return None


def read_table(ws, header_row, max_col=None):
    """(headers, rows, next_row). `rows` are dicts keyed by the header text;
    reading stops at the first row whose first column is blank."""
    max_col = max_col or ws.max_column
    headers = [safe_str(ws.cell(header_row, c).value) for c in range(1, max_col + 1)]
    rows = []
    r = header_row + 1
    while r <= ws.max_row:
        first = safe_str(ws.cell(r, 1).value)
        if not first:
            break
        row = {headers[c - 1] if c - 1 < len(headers) else f'col_{c}': safe_str(ws.cell(r, c).value)
               for c in range(1, max_col + 1)}
        rows.append(row)
        r += 1
    return headers, rows, r


def read_footnote(ws, from_row):
    """First non-blank column-1 text at or after from_row (the methodological
    footnote every one of these sheets writes as a single wide paragraph)."""
    r = from_row
    while r <= ws.max_row:
        val = safe_str(ws.cell(r, 1).value)
        if val:
            return val
        r += 1
    return ''


def read_kv_block(ws, start_row):
    """(dict, next_row) of Item/Value pairs starting at start_row, stopping at
    the first row whose column 1 is blank."""
    kv = {}
    r = start_row
    while r <= ws.max_row:
        item = safe_str(ws.cell(r, 1).value)
        if not item:
            break
        kv[item] = safe_str(ws.cell(r, 2).value)
        r += 1
    return kv, r


def read_kv_sections(ws, header_row):
    """Flat {Item: Value} dict for a sheet that groups many Item/Value rows
    under several section-marker rows (column 1 only, column 2 blank) with
    blank-row gaps between sections — immortal_time.py's Estimand sheet is
    built exactly this way ('What is being estimated' / 'Exposure definition'
    / 'Landmark times (pre-specified)' / ... each its own little block). Unlike
    read_table(), this does NOT stop at the first blank row — a plain
    read_table on this sheet only ever returns the first section, silently
    dropping everything after the first blank (that is how this bug was
    found: exposure_time_variable and landmark_times came back empty on a
    real production workbook). Rows with a blank value (section markers,
    the trailing footnote) are skipped rather than stored.
    """
    kv = {}
    for r in range(header_row + 1, ws.max_row + 1):
        item = safe_str(ws.cell(r, 1).value)
        val = safe_str(ws.cell(r, 2).value)
        if item and val:
            kv[item] = val
    return kv


def read_sectioned_table(ws, header_row, section_prefixes, max_col=None):
    """(headers, {section_label: [row_dict, ...]}, footnote) for a sheet that
    groups several small tables under one header row, each introduced by a
    section-marker row (column 1 only) — immortal_time.py's Naive_comparison is
    the sheet this exists for: 'Naive model...' / 'Corrected models' / 'Size of
    the bias', each with its own row count, separated by blank rows. A
    column-1-only row that does NOT start with one of `section_prefixes` is
    read as the sheet's trailing footnote instead of a fourth section, so the
    footnote paragraph never gets mistaken for an empty section header.
    """
    max_col = max_col or ws.max_column
    headers = [safe_str(ws.cell(header_row, c).value) for c in range(1, max_col + 1)]
    sections = {}
    current = None
    footnote = ''
    for r in range(header_row + 1, ws.max_row + 1):
        vals = [safe_str(ws.cell(r, c).value) for c in range(1, max_col + 1)]
        if not any(vals):
            continue
        if vals[0] and not any(vals[1:]):
            if vals[0].startswith(section_prefixes):
                current = vals[0]
                sections[current] = []
            else:
                footnote = vals[0]
            continue
        if current is not None:
            sections[current].append(
                {headers[c] if c < len(headers) else f'col_{c + 1}': vals[c]
                 for c in range(len(vals))})
    return headers, sections, footnote


def extract_table1(path):
    """Extract data from Table 1 (baseline characteristics) Excel file."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active

    result = {
        'type': 'table1',
        'title': safe_str(ws.cell(1, 1).value),
        'subtitle': safe_str(ws.cell(2, 1).value),
        'groups': [],
        'significant_vars': [],
        'non_significant_vars': [],
        'all_rows': []
    }

    # Extract group names and sizes from subtitle or header row
    subtitle = result['subtitle']
    # Try to extract N from subtitle like "... (N = 1234)"
    n_match = re.findall(r'[Nn]\s*=\s*(\d+)', subtitle)

    # Read header row to get group names
    header_row = 3
    headers = []
    for col in range(1, ws.max_column + 1):
        val = safe_str(ws.cell(header_row, col).value)
        if val:
            headers.append(val)

    # Read data rows
    for row in range(4, ws.max_row + 1):
        row_data = {}
        for col in range(1, min(ws.max_column + 1, len(headers) + 1)):
            row_data[headers[col - 1] if col - 1 < len(headers) else f'col_{col}'] = safe_str(ws.cell(row, col).value)

        if not any(row_data.values()):
            continue

        # Check if this is a section header (typically has value in col 1 only)
        first_val = safe_str(ws.cell(row, 1).value)
        if not first_val:
            continue

        # Look for p-value column
        p_val_str = ''
        for key, val in row_data.items():
            if 'p' in key.lower() and val:
                p_val_str = val
                break

        row_entry = {
            'variable': first_val,
            'values': row_data,
            'p_value': p_val_str
        }

        result['all_rows'].append(row_entry)

        # Classify as significant or not
        if p_val_str:
            try:
                p_clean = p_val_str.replace('<', '').replace('*', '').strip()
                p_num = float(p_clean)
                if p_num < 0.05 or '<' in p_val_str:
                    result['significant_vars'].append(row_entry)
                else:
                    result['non_significant_vars'].append(row_entry)
            except ValueError:
                pass

    wb.close()
    return result


def extract_univariate(path):
    """Extract data from univariate regression Excel file."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active

    result = {
        'type': 'univariate',
        'title': safe_str(ws.cell(1, 1).value),
        'subtitle': safe_str(ws.cell(2, 1).value),
        'significant_vars': [],
        'non_significant_vars': [],
        'all_rows': []
    }

    # Read headers
    headers = []
    for col in range(1, ws.max_column + 1):
        val = safe_str(ws.cell(3, col).value)
        if val:
            headers.append(val)
        # Also check row 4 for merged headers
        val4 = safe_str(ws.cell(4, col).value)
        if val4 and not val:
            headers.append(val4)

    # Read data rows
    for row in range(5, ws.max_row + 1):
        first_val = safe_str(ws.cell(row, 1).value)
        if not first_val:
            continue

        row_data = {}
        for col in range(1, ws.max_column + 1):
            col_header = headers[col - 1] if col - 1 < len(headers) else f'col_{col}'
            row_data[col_header] = safe_str(ws.cell(row, col).value)

        # Look for HR/OR and p-value in Overall columns
        hr_or = ''
        p_val = ''
        for key, val in row_data.items():
            key_lower = key.lower()
            if ('hr' in key_lower or 'or' in key_lower) and 'ci' in key_lower and val and val != 'N/A':
                hr_or = val
            if 'p' in key_lower and 'value' in key_lower and val:
                p_val = val

        row_entry = {
            'variable': first_val,
            'hr_or': hr_or,
            'p_value': p_val,
            'all_values': row_data
        }

        result['all_rows'].append(row_entry)

        if p_val:
            try:
                p_clean = p_val.replace('<', '').replace('*', '').strip()
                p_num = float(p_clean)
                if p_num < 0.05 or '<' in p_val:
                    result['significant_vars'].append(row_entry)
                else:
                    result['non_significant_vars'].append(row_entry)
            except ValueError:
                pass

    wb.close()
    return result


def extract_multivariate(path):
    """Extract data from multivariate regression + interaction Excel file."""
    wb = load_workbook(path, data_only=True)

    result = {
        'type': 'multivariate',
        'title': '',
        'subtitle': '',
        'adjusted_results': [],
        'interactions': [],
        'covariates_list': ''
    }

    # Try to read the main sheet
    ws = wb.active
    result['title'] = safe_str(ws.cell(1, 1).value)
    result['subtitle'] = safe_str(ws.cell(2, 1).value)

    # Read all rows for adjusted results
    for row in range(5, ws.max_row + 1):
        first_val = safe_str(ws.cell(row, 1).value)
        if not first_val:
            continue

        row_data = {}
        for col in range(1, ws.max_column + 1):
            row_data[f'col_{col}'] = safe_str(ws.cell(row, col).value)

        # Look for interaction p-value
        p_interaction = ''
        for col in range(1, ws.max_column + 1):
            val = safe_str(ws.cell(row, col).value)
            if val and ('interaction' in safe_str(ws.cell(3, col).value).lower() or
                        'interaction' in safe_str(ws.cell(4, col).value).lower()):
                p_interaction = val

        entry = {
            'variable': first_val,
            'all_values': row_data,
            'p_interaction': p_interaction
        }

        result['adjusted_results'].append(entry)

        if p_interaction:
            try:
                p_clean = p_interaction.replace('<', '').replace('*', '').strip()
                p_num = float(p_clean)
                if p_num < 0.05 or '<' in p_interaction:
                    result['interactions'].append(entry)
            except ValueError:
                pass

    wb.close()
    return result


def extract_kaplan_meier(path):
    """Extract data from KM survival analysis Excel file."""
    wb = load_workbook(path, data_only=True)

    result = {
        'type': 'kaplan_meier',
        'title': '',
        'subtitle': '',
        'groups': [],
        'logrank_p': '',
        'logrank_stat': ''
    }

    # Sheet 1: Survival Summary
    ws1 = wb['Survival Summary'] if 'Survival Summary' in wb.sheetnames else wb.active
    result['title'] = safe_str(ws1.cell(1, 1).value)
    result['subtitle'] = safe_str(ws1.cell(2, 1).value)

    # Read headers
    headers = []
    for col in range(1, ws1.max_column + 1):
        headers.append(safe_str(ws1.cell(3, col).value))

    # Read group data
    for row in range(4, ws1.max_row + 1):
        first_val = safe_str(ws1.cell(row, 1).value)
        if not first_val or first_val.startswith('Log-rank') or first_val.startswith('KM') or first_val.startswith('Survival') or first_val.startswith('Abbrevi'):
            # Check if this is a footnote with log-rank info
            cell_val = safe_str(ws1.cell(row, 1).value)
            if 'log-rank' in cell_val.lower() or 'Log-rank' in cell_val:
                p_match = re.search(r'p\s*[=<]\s*([0-9.<]+)', cell_val, re.IGNORECASE)
                if p_match:
                    result['logrank_p'] = p_match.group(0).replace('p', 'P').strip()
                stat_match = re.search(r'χ²\s*=\s*([0-9.]+)', cell_val)
                if stat_match:
                    result['logrank_stat'] = stat_match.group(1)
            continue

        group_data = {'label': first_val}
        for col in range(2, ws1.max_column + 1):
            header = headers[col - 1] if col - 1 < len(headers) else f'col_{col}'
            group_data[header] = safe_str(ws1.cell(row, col).value)

        result['groups'].append(group_data)

    # Sheet 2: Log-Rank Test
    if 'Log-Rank Test' in wb.sheetnames:
        ws2 = wb['Log-Rank Test']
        for row in range(3, ws2.max_row + 1):
            metric = safe_str(ws2.cell(row, 1).value)
            value = safe_str(ws2.cell(row, 2).value)
            if 'p-value' in metric.lower():
                result['logrank_p'] = value
            elif 'statistic' in metric.lower():
                result['logrank_stat'] = value

    wb.close()
    return result


def extract_competing_risk(path):
    """Extract data from competing risk CIF Excel file."""
    wb = load_workbook(path, data_only=True)

    result = {
        'type': 'competing_risk',
        'title': '',
        'subtitle': '',
        'groups': [],
        'grays_tests': [],
        'footnotes': []
    }

    ws1 = wb['CIF Summary'] if 'CIF Summary' in wb.sheetnames else wb.active
    result['title'] = safe_str(ws1.cell(1, 1).value)
    result['subtitle'] = safe_str(ws1.cell(2, 1).value)

    # Read headers
    headers = []
    for col in range(1, ws1.max_column + 1):
        headers.append(safe_str(ws1.cell(3, col).value))

    # Read group data
    for row in range(4, ws1.max_row + 1):
        first_val = safe_str(ws1.cell(row, 1).value)
        if not first_val:
            continue
        if 'gray' in first_val.lower() or 'cif' in first_val.lower() or 'abbrevi' in first_val.lower():
            result['footnotes'].append(first_val)
            continue

        group_data = {'label': first_val}
        for col in range(2, ws1.max_column + 1):
            header = headers[col - 1] if col - 1 < len(headers) else f'col_{col}'
            group_data[header] = safe_str(ws1.cell(row, col).value)
        result['groups'].append(group_data)

    # Sheet 2: Cause-Specific Log-Rank (new name) or legacy "Gray's Test" (old)
    ws2 = None
    for sheet_name in ["Cause-Specific Log-Rank", "Gray's Test"]:
        if sheet_name in wb.sheetnames:
            ws2 = wb[sheet_name]
            break
    if ws2 is not None:
        for row in range(3, ws2.max_row + 1):
            event_type = safe_str(ws2.cell(row, 1).value)
            if not event_type or 'gray' in event_type.lower() or 'cause-specific' in event_type.lower():
                continue
            stat = safe_str(ws2.cell(row, 2).value)
            df = safe_str(ws2.cell(row, 3).value)
            p_val = safe_str(ws2.cell(row, 4).value)
            result['grays_tests'].append({
                'event_type': event_type,
                'statistic': stat,
                'df': df,
                'p_value': p_val
            })

    wb.close()
    return result


def extract_psm(path):
    """Extract data from PSM matching Excel file."""
    wb = load_workbook(path, data_only=True)

    result = {
        'type': 'psm',
        'title': '',
        'subtitle': '',
        'outcomes': [],
        'balance_summary': '',
        'matched_n': ''
    }

    ws = wb.active
    result['title'] = safe_str(ws.cell(1, 1).value)
    result['subtitle'] = safe_str(ws.cell(2, 1).value)

    # Extract matched N from subtitle
    n_match = re.findall(r'[Nn]\s*=\s*(\d+)', result['subtitle'])
    if n_match:
        result['matched_n'] = n_match

    # Read all data rows
    headers = []
    for col in range(1, ws.max_column + 1):
        val = safe_str(ws.cell(3, col).value)
        if not val:
            val = safe_str(ws.cell(4, col).value)
        headers.append(val)

    for row in range(5, ws.max_row + 1):
        first_val = safe_str(ws.cell(row, 1).value)
        if not first_val:
            continue

        row_data = {'outcome': first_val}
        for col in range(2, ws.max_column + 1):
            header = headers[col - 1] if col - 1 < len(headers) else f'col_{col}'
            row_data[header] = safe_str(ws.cell(row, col).value)
        result['outcomes'].append(row_data)

    # Check for balance sheet
    if 'Balance' in wb.sheetnames:
        ws_bal = wb['Balance']
        max_smd = 0
        for row in range(2, ws_bal.max_row + 1):
            for col in range(1, ws_bal.max_column + 1):
                val = safe_str(ws_bal.cell(row, col).value)
                if 'smd' in safe_str(ws_bal.cell(1, col).value).lower():
                    try:
                        smd_val = abs(float(val))
                        max_smd = max(max_smd, smd_val)
                    except (ValueError, TypeError):
                        pass
        result['balance_summary'] = f"Maximum SMD after matching: {max_smd:.3f}"

    wb.close()
    return result


def extract_auc(path):
    """Extract data from AUC/C-statistics Excel file."""
    wb = load_workbook(path, data_only=True)

    result = {
        'type': 'auc',
        'title': '',
        'subtitle': '',
        'variables': [],
        'clinical_model_stat': ''
    }

    ws = wb.active
    result['title'] = safe_str(ws.cell(1, 1).value)
    result['subtitle'] = safe_str(ws.cell(2, 1).value)

    # Read headers
    headers = []
    for col in range(1, ws.max_column + 1):
        val = safe_str(ws.cell(3, col).value)
        if not val:
            val = safe_str(ws.cell(4, col).value)
        headers.append(val)

    # Read data rows
    for row in range(5, ws.max_row + 1):
        first_val = safe_str(ws.cell(row, 1).value)
        if not first_val or first_val.startswith('Abbrevi') or first_val.startswith('AUC') or first_val.startswith('C-'):
            continue

        row_data = {'variable': first_val}
        for col in range(2, ws.max_column + 1):
            header = headers[col - 1] if col - 1 < len(headers) else f'col_{col}'
            row_data[header] = safe_str(ws.cell(row, col).value)
        result['variables'].append(row_data)

    wb.close()
    return result


def extract_weighting(path):
    """Extract data from a propensity-score weighting (IPTW / overlap) file.

    The estimates come out as a set, never one at a time. Under marginal
    positivity the weighting step adds an overlap-weighted (ATO) companion beside
    the IPTW estimate, and the two answer different questions — on real data they
    can differ by more than twofold and point opposite ways. Pulling only the
    headline row here would rebuild, inside the drafting material, exactly the
    selective reporting the companion exists to prevent. So every weighted row is
    returned together with its estimand, and `must_report_together` is set when
    more than one is present.
    """
    wb = load_workbook(path, data_only=True)

    result = {
        'type': 'weighting',
        'title': '',
        'subtitle': '',
        'estimates': [],
        'must_report_together': False,
        'has_ato_companion': False,
        'positivity': {},
        'ess': '',
        'max_smd_after': '',
        'n_smd_over_threshold': '',
    }

    ws = wb['Outcome_model'] if 'Outcome_model' in wb.sheetnames else wb.active
    result['title'] = safe_str(ws.cell(1, 1).value)

    # Find the header row rather than assuming where it is. The sheet carries a
    # subtitle only when the plan has a study title — which it always does in the
    # real pipeline and never did in the first test fixtures — and that one line
    # shifts every row below it. Assuming a position meant the header itself was
    # read as an estimate, `must_report_together` was true even for a single
    # estimand, and balance came back empty, on every production workbook.
    header_row = None
    for row in range(1, min(ws.max_row, 12) + 1):
        if safe_str(ws.cell(row, 1).value) == 'Analysis':
            header_row = row
            break
    if header_row is None:
        wb.close()
        raise ValueError(
            f"{os.path.basename(path)}: no 'Analysis' header found in Outcome_model; "
            f"the sheet layout is not what this extractor understands")

    # Rows run "Analysis | Estimand | Effect (95% CI) | p" until the first blank.
    for row in range(header_row + 1, ws.max_row + 1):
        analysis = safe_str(ws.cell(row, 1).value)
        estimand = safe_str(ws.cell(row, 2).value)
        effect = safe_str(ws.cell(row, 3).value)
        if not analysis:
            continue
        if not effect or not estimand:
            # Past the table body — the footnote paragraph starts here.
            break
        result['estimates'].append({
            'analysis': analysis,
            'estimand': estimand,
            'effect': effect,
            'p': safe_str(ws.cell(row, 4).value),
            'is_companion': 'ATO companion' in analysis,
        })

    weighted = [e for e in result['estimates']
                if e['estimand'] and e['estimand'].lower() not in ('none — confounded',
                                                                   'none - confounded')]
    result['has_ato_companion'] = any(e['is_companion'] for e in result['estimates'])
    result['must_report_together'] = len(weighted) > 1

    # Positivity and precision context. A weighted effect without the overlap
    # diagnostics and the effective sample size cannot be read: the confidence
    # interval reflects the ESS, not the row count.
    if 'PS_distribution' in wb.sheetnames:
        wsp = wb['PS_distribution']
        for row in range(1, wsp.max_row + 1):
            key = safe_str(wsp.cell(row, 1).value)
            val = safe_str(wsp.cell(row, 2).value)
            if key in ('Region of common support', 'Subjects outside common support',
                       'Overlap coefficient'):
                result['positivity'][key] = val

    if 'Weights' in wb.sheetnames:
        wsw = wb['Weights']
        for row in range(1, wsw.max_row + 1):
            if safe_str(wsw.cell(row, 1).value).startswith('Effective sample size'):
                result['ess'] = safe_str(wsw.cell(row, 4).value)
                break

    if 'Balance' in wb.sheetnames:
        wsb = wb['Balance']
        # Locate "SMD after" by its header. The last column is "Balanced?" and is
        # blank on the summary rows, and column 4 is "SMD before" — reading either
        # by position would report balance that the weighting has not achieved.
        # Same reason as above: locate the header row first, then the column
        # inside it. Row 2 is the header only when there is no subtitle.
        smd_after_col = None
        for hrow in range(1, min(wsb.max_row, 12) + 1):
            if safe_str(wsb.cell(hrow, 1).value) != 'Variable':
                continue
            for col in range(1, wsb.max_column + 1):
                if safe_str(wsb.cell(hrow, col).value) == 'SMD after':
                    smd_after_col = col
                    break
            break
        if smd_after_col:
            for row in range(1, wsb.max_row + 1):
                key = safe_str(wsb.cell(row, 1).value)
                if key == 'Maximum SMD':
                    result['max_smd_after'] = safe_str(wsb.cell(row, smd_after_col).value)
                elif key.startswith('Variables with SMD'):
                    result['n_smd_over_threshold'] = safe_str(
                        wsb.cell(row, smd_after_col).value)

    wb.close()
    return result


def extract_nonlinearity(path):
    """Extract data from a restricted-cubic-spline (RCS) non-linearity file.

    advanced_analyses.md section 3: p-nonlinear >= 0.05 is NOT evidence that a
    linear specification is correct — it only means this sample size did not
    detect a departure from it — so that reading rule is pulled verbatim from
    the sheet's own footnote rather than re-typed here, where it could drift
    out of sync with the script. Every variable also carries its reference
    value (every reported HR/OR is relative to it) and the linear-term result
    from the same adjusted model, so a reader can see what the linear
    specification would have claimed.
    """
    wb = load_workbook(path, data_only=True)
    result = {
        'type': 'nonlinearity',
        'title': '', 'subtitle': '',
        'variables': [],
        'reading_rule': '',
    }
    if 'Estimand' in wb.sheetnames:
        wse = wb['Estimand']
        result['title'] = safe_str(wse.cell(1, 1).value)
        result['subtitle'] = safe_str(wse.cell(2, 1).value)
    ws = wb['Nonlinearity_tests'] if 'Nonlinearity_tests' in wb.sheetnames else wb.active
    hdr = find_header_row(ws, 'Variable')
    if hdr is None:
        wb.close()
        raise ValueError(
            f"{os.path.basename(path)}: no 'Variable' header found in Nonlinearity_tests; "
            f"the sheet layout is not what this extractor understands")
    _, rows, after = read_table(ws, hdr)
    for row in rows:
        linear_key = next((k for k in row if k.startswith('Linear-term') and 'CI' in k), None)
        p_nonlin = row.get('p-nonlinear', '')
        p_val = _to_float(p_nonlin)
        result['variables'].append({
            'variable': row.get('Variable', ''),
            'column': row.get('Column', ''),
            'n': row.get('N', ''),
            'events': row.get('Events', ''),
            'knots': row.get('Knots (k)', ''),
            'knot_percentiles': row.get('Knot percentiles', ''),
            'knot_values': row.get('Knot values', ''),
            'reference_value': row.get('Reference value', ''),
            'p_nonlinear': p_nonlin,
            'events_per_spline_param': row.get('Events per spline parameter', ''),
            'linear_effect': row.get(linear_key, '') if linear_key else '',
            'linear_p': row.get('Linear-term p', ''),
            'p_overall': row.get('p-overall', ''),
            'flags': row.get('Flags', ''),
            # advanced_analyses.md section 3's reading rule as a boolean so the
            # manuscript drafter cannot quietly claim linearity holds.
            'nonlinearity_not_established': p_val is not None and p_val >= 0.05,
        })
    result['reading_rule'] = read_footnote(ws, after)
    wb.close()
    return result


def extract_robust_variance(path):
    """Extract data from a cluster-robust / GEE / mixed variance file.

    advanced_analyses.md section 4: "與未修正版本並排比較，讓人看得到修正前後 CI
    差多少 — 這個對照本身就是說服 reviewer 的材料." The primary payload is
    therefore the Naive_vs_Robust table (both the naive and the corrected
    estimate for every term), not the corrected-only Adjusted_robust sheet —
    pulling only the corrected numbers would delete the exact comparison this
    step exists to produce. Cluster count and the dominant-cluster share come
    along because a cluster-robust SE below ~30 clusters is itself biased
    downward (same section) — the correction can still understate uncertainty.
    """
    wb = load_workbook(path, data_only=True)
    result = {
        'type': 'robust-variance',
        'title': '', 'subtitle': '',
        'method': '', 'correlation_structure': '',
        'terms': [],
        'must_report_together': True,   # naive vs corrected IS the point of this file
        'n_clusters': None,
        'clusters_below_min': None,     # < 30 (advanced_analyses.md section 4)
        'largest_cluster_share_pct': None,
        'largest_cluster_dominant': None,   # > 20% of the sample in one cluster
    }
    MIN_CLUSTERS = 30
    DOMINANT_SHARE_PCT = 20.0

    if 'Estimand' in wb.sheetnames:
        wse = wb['Estimand']
        result['title'] = safe_str(wse.cell(1, 1).value)
        result['subtitle'] = safe_str(wse.cell(2, 1).value)
        hdr = find_header_row(wse, 'Item')
        if hdr is not None:
            _, erows, _ = read_table(wse, hdr)
            for r in erows:
                if r.get('Item') == 'Method used':
                    result['method'] = r.get('Statement', '')
                elif r.get('Item') == 'Correlation structure':
                    result['correlation_structure'] = r.get('Statement', '')

    ws = wb['Naive_vs_Robust'] if 'Naive_vs_Robust' in wb.sheetnames else wb.active
    hdr = find_header_row(ws, 'Variable')
    if hdr is None:
        wb.close()
        raise ValueError(
            f"{os.path.basename(path)}: no 'Variable' header found in Naive_vs_Robust; "
            f"the sheet layout is not what this extractor understands")
    _, rows, _ = read_table(ws, hdr)
    for row in rows:
        naive = {k[len('Naive '):].strip(): v for k, v in row.items() if k.startswith('Naive ')}
        skip = {'Variable', 'Model term', 'SE ratio (corrected / naive)',
                'CI width ratio (corrected / naive)', 'Crosses the null differently?'}
        corrected = {k: v for k, v in row.items() if k not in skip and not k.startswith('Naive ')}
        result['terms'].append({
            'variable': row.get('Variable', ''),
            'term': row.get('Model term', ''),
            'naive': naive,
            'corrected': corrected,
            'se_ratio': row.get('SE ratio (corrected / naive)', ''),
            'ci_width_ratio': row.get('CI width ratio (corrected / naive)', ''),
            'significance_changed': row.get('Crosses the null differently?', ''),
        })

    if 'Cluster_summary' in wb.sheetnames:
        wsc = wb['Cluster_summary']
        hdr = find_header_row(wsc, 'Item')
        if hdr is not None:
            kv, _ = read_kv_block(wsc, hdr + 1)
            n_clusters = _to_float(kv.get('Number of clusters'))
            share = kv.get('Largest cluster as % of the sample', '')
            share_val = _to_float(share.replace('%', '')) if share else None
            result['n_clusters'] = kv.get('Number of clusters', '')
            result['clusters_below_min'] = (n_clusters is not None and n_clusters < MIN_CLUSTERS)
            result['largest_cluster_share_pct'] = share
            result['largest_cluster_dominant'] = (share_val is not None
                                                   and share_val > DOMINANT_SHARE_PCT)
    wb.close()
    return result


def extract_calibration(path):
    """Extract data from a calibration / discrimination / internal-validation file.

    advanced_analyses.md section 5: this whole file is gated on
    `study_purpose` — "若研究問題是病因學...不要把 [C-index] 說成模型效能，也不需要
    calibration" — so that flag is extracted first and used to decide whether
    Calibration/Brier/Internal_validation even exist to read (calibration.py
    only writes them when study_purpose == "prediction"; an etiologic run has
    only Discrimination + Method, and treating its absence as a parsing error
    would be wrong). Where internal validation IS reported, the
    optimism-corrected C-index and the optimism-corrected calibration slope are
    returned together — the protocol requires both numbers, not the C-index
    alone, before this can be called a validated prediction model.
    """
    wb = load_workbook(path, data_only=True)
    result = {
        'type': 'calibration',
        'title': '', 'subtitle': '',
        'study_purpose': '',
        'is_prediction_model': False,
        'n_analyzed': '', 'n_total': '', 'n_dropped_listwise': '', 'n_events': '',
        'bootstrap_n': '', 'fast_mode': False,
        'discrimination': {},
        'calibration_slope': {},        # prediction only: apparent + REPORTED, paired
        'calibration_per_time': [],     # prediction only, Cox
        'calibration_bins': [],         # prediction only: predicted vs observed by decile
        'brier': [],                    # prediction only
        'internal_validation': {},      # prediction only: c-index + slope, both optimism-corrected
        'must_report_together': False,  # set True below when internal_validation is present
    }

    if 'Method' in wb.sheetnames:
        wsm = wb['Method']
        # No sheet in this workbook carries the outcome-specific config['title']
        # ("Calibration & internal validation: {outcome}") — every sheet_header()
        # call in calibration.py uses a fixed per-sheet title instead. Use a
        # constant label rather than borrowing a generic sheet title such as
        # "Methodological choices and defaults", which would misrepresent what
        # the file is. The study title (subtitle) IS shared across sheets.
        result['title'] = 'Calibration & internal validation'
        result['subtitle'] = safe_str(wsm.cell(2, 1).value)
        hdr = find_header_row(wsm, 'Setting')
        if hdr is not None:
            _, rows, _ = read_table(wsm, hdr)
            by_setting = {r.get('Setting', ''): r for r in rows}
            result['study_purpose'] = by_setting.get('Study purpose', {}).get('Value', '')
            boot = by_setting.get('Bootstrap resamples', {}).get('Value', '')
            result['bootstrap_n'] = boot
            result['fast_mode'] = 'FAST MODE' in boot

    result['is_prediction_model'] = (result['study_purpose'] == 'prediction')

    if 'Estimand' in wb.sheetnames:
        wse = wb['Estimand']
        for row in range(1, wse.max_row + 1):
            # safe_str() strips leading whitespace, so the sheet's own
            # 4-space "    Rows in analysis: ..." indentation is gone by the
            # time it gets here — match on the stripped text, not the raw cell.
            text = safe_str(wse.cell(row, 1).value)
            if text.startswith('Rows in analysis:'):
                m = re.search(r'(\d+)\s*\(of\s*(\d+);\s*(\d+)\s*dropped', text)
                if m:
                    result['n_analyzed'], result['n_total'], result['n_dropped_listwise'] = m.groups()
            elif text.startswith('Events:'):
                result['n_events'] = text.split(':', 1)[1].strip()

    if 'Discrimination' in wb.sheetnames:
        wsd = wb['Discrimination']
        hdr = find_header_row(wsd, 'Metric')
        if hdr is not None:
            _, rows, _ = read_table(wsd, hdr)
            if rows:
                result['discrimination'] = rows[0]

    if 'Calibration' in wb.sheetnames:
        wsc = wb['Calibration']
        hdr = find_header_row(wsc, 'Item')
        if hdr is not None:
            _, rows, after = read_table(wsc, hdr)
            result['calibration_slope'] = {r.get('Item', ''): r for r in rows}
            hdr2 = find_header_row(wsc, 'Time', max_scan=wsc.max_row)
            if hdr2 is not None and hdr2 >= after:
                _, per_time, after2 = read_table(wsc, hdr2)
                result['calibration_per_time'] = per_time
                hdr3 = None
                for row in range(after2, wsc.max_row + 1):
                    if safe_str(wsc.cell(row, 1).value) == 'Bin (time)':
                        hdr3 = row
                        break
                if hdr3 is not None:
                    _, bins, _ = read_table(wsc, hdr3)
                    result['calibration_bins'] = bins

    if 'Brier' in wb.sheetnames:
        wsb = wb['Brier']
        hdr = find_header_row(wsb, 'Time')
        if hdr is not None:
            _, rows, _ = read_table(wsb, hdr)
            result['brier'] = rows

    if 'Internal_validation' in wb.sheetnames:
        wsi = wb['Internal_validation']
        hdr = find_header_row(wsi, 'Quantity')
        if hdr is not None:
            _, rows, _ = read_table(wsi, hdr)
            result['internal_validation'] = {r.get('Quantity', ''): r for r in rows}
            if any(k.startswith('Harrell') or 'C-index' in k or 'index' in k.lower()
                   for k in result['internal_validation']) and any(
                    'slope' in k.lower() for k in result['internal_validation']):
                result['must_report_together'] = True

    wb.close()
    return result


def extract_immortal_time(path):
    """Extract data from a landmark / extended-Cox immortal-time file.

    advanced_analyses.md section 7: "兩種都做時要說明為什麼結果不同，不可只挑一個
    報" and the naive baseline-covariate model is written into every one of
    these workbooks specifically so the size of the immortal-time bias is
    visible — the gap between the naive row and the corrected row(s) IS the
    finding. The naive and corrected rows are therefore always returned
    together (`must_report_together`), and for landmark mode the excluded-population
    table comes along too, because "被排除的人不是隨機的" (same section) —
    a landmark HR without knowing who it dropped is not fully interpretable.
    """
    wb = load_workbook(path, data_only=True)
    result = {
        'type': 'immortal-time',
        'title': '', 'subtitle': '',
        'mode': '',
        'estimand': '', 'target_population': '',
        'exposure_time_variable': '', 'landmark_times': '', 'landmark_rationale': '',
        'naive': [],
        'corrected': [],
        'bias_size': [],
        'landmark_results': [],       # mode == landmark
        'excluded_at_landmark': [],   # mode == landmark
        'extended_cox_results': [],   # mode == extended_cox
        'must_report_together': True,  # naive vs corrected is the whole point of this file
    }
    result['mode'] = ('landmark' if 'Landmark_results' in wb.sheetnames
                      else 'extended_cox' if 'Extended_cox_results' in wb.sheetnames
                      else '')

    if 'Estimand' in wb.sheetnames:
        wse = wb['Estimand']
        result['title'] = safe_str(wse.cell(1, 1).value)
        result['subtitle'] = safe_str(wse.cell(2, 1).value)
        hdr = find_header_row(wse, 'Item')
        if hdr is not None:
            # This sheet groups its Item/Specification rows under several
            # section markers with blank-row gaps between them (see
            # read_kv_sections' docstring) — a plain read_table() here only
            # ever returns the first section and silently drops the rest.
            by_item = read_kv_sections(wse, hdr)
            result['estimand'] = by_item.get('Estimand', '')
            result['target_population'] = by_item.get('Target population', '')
            result['exposure_time_variable'] = by_item.get('Exposure time variable', '')
            result['landmark_times'] = by_item.get('Landmark times', '')
            result['landmark_rationale'] = by_item.get('Rationale for these times', '')

    if 'Landmark_results' in wb.sheetnames:
        wsl = wb['Landmark_results']
        # The header text embeds config['time_unit'] ("Landmark (days)" /
        # "(months)" / "(years)"), so it cannot be matched as one literal string.
        hdr = None
        for unit in ('days', 'months', 'years'):
            hdr = find_header_row(wsl, f'Landmark ({unit})')
            if hdr is not None:
                break
        if hdr is not None:
            _, rows, _ = read_table(wsl, hdr)
            result['landmark_results'] = rows

    if 'Excluded_at_landmark' in wb.sheetnames:
        wsx = wb['Excluded_at_landmark']
        hdr = None
        for row in range(1, min(wsx.max_row, 10) + 1):
            if safe_str(wsx.cell(row, 2).value) == 'n excluded':
                hdr = row
                break
        if hdr is not None:
            _, rows, _ = read_table(wsx, hdr)
            result['excluded_at_landmark'] = rows

    if 'Extended_cox_results' in wb.sheetnames:
        wse2 = wb['Extended_cox_results']
        hdr = find_header_row(wse2, 'Model')
        if hdr is not None:
            _, rows, _ = read_table(wse2, hdr)
            result['extended_cox_results'] = rows

    if 'Naive_comparison' in wb.sheetnames:
        wsn = wb['Naive_comparison']
        hdr = find_header_row(wsn, 'Model')
        if hdr is None:
            wb.close()
            raise ValueError(
                f"{os.path.basename(path)}: no 'Model' header found in Naive_comparison; "
                f"the sheet layout is not what this extractor understands")
        _, sections, _ = read_sectioned_table(
            wsn, hdr, section_prefixes=('Naive model', 'Corrected models', 'Size of the bias'))
        for label, rows in sections.items():
            if label.startswith('Naive model'):
                result['naive'] = rows
            elif label == 'Corrected models':
                result['corrected'] = rows
            elif label == 'Size of the bias':
                result['bias_size'] = rows

    wb.close()
    return result


def extract_evalue(path):
    """Extract data from an E-value (unmeasured-confounding sensitivity) file.

    advanced_analyses.md section 1: "只給一個數字是沒有意義的" — the point E-value
    alone says nothing about how fragile a finding is; the CI E-value is what
    shows how far the evidence can be pushed toward the null, so both travel
    together in every row (`must_report_together`). The plain-language
    comparison against the strongest measured confounder — required by the
    same section so the number has a sense of scale — is already composed into
    each row's interpretation text by evalue.py, so it is kept verbatim rather
    than reconstructed.
    """
    wb = load_workbook(path, data_only=True)
    ws = wb['E-values'] if 'E-values' in wb.sheetnames else wb.active
    result = {
        'type': 'evalue',
        'title': safe_str(ws.cell(1, 1).value),
        'subtitle': safe_str(ws.cell(2, 1).value),
        'estimates': [],
        'must_report_together': True,  # point E-value and CI E-value, always
        'effect_measure': '', 'outcome_prevalence': '', 'scale_conversion': '',
        'formula': '', 'caution': '',
    }
    hdr = find_header_row(ws, 'Variable')
    if hdr is None:
        wb.close()
        raise ValueError(
            f"{os.path.basename(path)}: no 'Variable' header found in E-values; "
            f"the sheet layout is not what this extractor understands")
    _, rows, after = read_table(ws, hdr)
    result['estimates'] = rows

    r = after
    while r <= ws.max_row and safe_str(ws.cell(r, 1).value) != 'Method':
        r += 1
    if r <= ws.max_row:
        kv, _ = read_kv_block(ws, r + 1)
        result['effect_measure'] = kv.get('Effect measure', '')
        result['outcome_prevalence'] = kv.get('Outcome prevalence', '')
        result['scale_conversion'] = kv.get('Scale conversion', '')
        result['formula'] = kv.get('Formula', '')
        result['caution'] = kv.get('Caution', '')

    wb.close()
    return result


def extract_strobe(path):
    """Extract data from the STROBE reporting-completeness checklist.

    advanced_analyses.md section 2: "打勾但指不出東西的 checklist 比沒有還糟" —
    every row's evidence text is kept, not just its status, because `covered`
    is only meaningful with a pointer to the artifact behind it. The full list
    of `missing` items is surfaced directly: those are exactly the ones the
    protocol says the pipeline cannot fill in and the author must write by
    hand (study setting, funding, ethics approval, etc.) — dropping them from
    the extraction would silently turn "known gap" back into "nobody checked".
    """
    wb = load_workbook(path, data_only=True)
    ws = wb['STROBE'] if 'STROBE' in wb.sheetnames else wb.active
    subtitle = safe_str(ws.cell(2, 1).value)
    result = {
        'type': 'strobe',
        'title': safe_str(ws.cell(1, 1).value),
        'subtitle': subtitle,
        'design_registered': '', 'strobe_version_used': '',
        'design_unmapped_warning': '',
        'counts': {'covered': 0, 'partial': 0, 'missing': 0, 'n/a': 0},
        'items': [],
        'missing_items': [],
    }
    m = re.search(r'計畫登記的設計\s*([^\s　]+)', subtitle)
    if m:
        result['design_registered'] = m.group(1)
    m = re.search(r'採用的\s*STROBE\s*版本\s*([^\s　]+)', subtitle)
    if m:
        result['strobe_version_used'] = m.group(1)
    r3 = safe_str(ws.cell(3, 1).value)
    if r3.startswith('⚠'):
        result['design_unmapped_warning'] = r3[1:].strip()

    hdr = find_header_row(ws, '#')
    if hdr is None:
        wb.close()
        raise ValueError(
            f"{os.path.basename(path)}: no '#' header found in STROBE; "
            f"the sheet layout is not what this extractor understands")
    _, rows, _ = read_table(ws, hdr)
    for row in rows:
        status_cell = row.get('狀態', '')
        status = status_cell.split(' ', 1)[1] if ' ' in status_cell else status_cell
        item = {
            'item': row.get('#', ''),
            'section': row.get('區段', ''),
            'requirement': row.get('要求', ''),
            'status': status,
            'evidence': row.get('證據 / 待辦', ''),
        }
        result['items'].append(item)
        if status in result['counts']:
            result['counts'][status] += 1
        if status == 'missing':
            result['missing_items'].append(item)
    wb.close()
    return result


def extract_missing_data(path):
    """Extract data from a multiple-imputation (MI) file.

    missing_data_protocol.md section 5: "同一個模型的 complete-case 與 MI 結果
    並排...這張表本身就是 reviewer 想看的 sensitivity analysis" — CC_vs_MI is
    therefore the primary payload, not the Pooled sheet alone; reporting only
    the pooled MI estimate would delete the exact comparison the protocol
    exists to produce. m, seed and the imputation model's variable list travel
    with it because a pooled estimate without them is not reproducible
    (protocol section 5.1 items 2-3).

    Note: this extractor is for the per-outcome `mi_{outcome}` step
    (MI_{outcome}.xlsx). The unconditional `missingness` assessment step
    (Missingness.xlsx) is a different artifact — see generate_configs.py's
    gen_manuscript() docstring for why it is deliberately not wired in here.
    """
    wb = load_workbook(path, data_only=True)
    result = {
        'type': 'missing-data',
        'title': '', 'subtitle': '',
        'pooled': [],
        'cc_vs_mi': [],
        'must_report_together': True,   # CC and MI must not be reported alone
        'flagged_terms': [],            # CC_vs_MI rows where Flag != '-'
        'imputation_model': {'m': '', 'seed': '', 'maxit': '', 'engine': '',
                              'variables': []},
        'convergence': [],
    }

    if 'Pooled' in wb.sheetnames:
        wsp = wb['Pooled']
        result['title'] = safe_str(wsp.cell(1, 1).value)
        result['subtitle'] = safe_str(wsp.cell(2, 1).value)
        hdr = find_header_row(wsp, 'Variable')
        if hdr is None:
            wb.close()
            raise ValueError(
                f"{os.path.basename(path)}: no 'Variable' header found in Pooled; "
                f"the sheet layout is not what this extractor understands")
        _, rows, _ = read_table(wsp, hdr)
        result['pooled'] = rows

    if 'CC_vs_MI' in wb.sheetnames:
        wsc = wb['CC_vs_MI']
        hdr = find_header_row(wsc, 'Variable')
        if hdr is not None:
            _, rows, _ = read_table(wsc, hdr)
            result['cc_vs_mi'] = rows
            for row in rows:
                flag = row.get('Flag', '')
                if flag and flag != '-':
                    result['flagged_terms'].append(row)

    if 'Imputation_model' in wb.sheetnames:
        wsi = wb['Imputation_model']
        hdr = find_header_row(wsi, 'Column')
        if hdr is not None:
            _, rows, after = read_table(wsi, hdr)
            result['imputation_model']['variables'] = rows
            for r in range(after, wsi.max_row + 1):
                item = safe_str(wsi.cell(r, 1).value)
                val = safe_str(wsi.cell(r, 2).value)
                if item == 'Number of imputations (m)':
                    result['imputation_model']['m'] = val
                elif item == 'Random seed':
                    result['imputation_model']['seed'] = val
                elif item == 'Iterations (maxit)':
                    result['imputation_model']['maxit'] = val
                elif item == 'Engine':
                    result['imputation_model']['engine'] = val

    if 'Convergence_notes' in wb.sheetnames:
        wsv = wb['Convergence_notes']
        hdr = find_header_row(wsv, 'Variable')
        if hdr is not None:
            _, rows, _ = read_table(wsv, hdr)
            result['convergence'] = rows

    wb.close()
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════════════════════

EXTRACTORS = {
    'table1': extract_table1,
    'univariate': extract_univariate,
    'multivariate': extract_multivariate,
    'kaplan_meier': extract_kaplan_meier,
    'competing_risk': extract_competing_risk,
    'psm': extract_psm,
    'auc': extract_auc,
    'weighting': extract_weighting,
    'nonlinearity': extract_nonlinearity,
    'robust-variance': extract_robust_variance,
    'calibration': extract_calibration,
    'immortal-time': extract_immortal_time,
    'evalue': extract_evalue,
    'strobe': extract_strobe,
    'missing-data': extract_missing_data,
}


def main():
    if len(sys.argv) < 2:
        print("Usage: python extract_data.py config.json")
        sys.exit(1)

    config = load_config(sys.argv[1])
    extracted = {}
    # TP-1 audit fix (2026-05-08, mr-read-01 + mr-read-02): track per-record
    # failure types so script exits non-zero on real errors. Previously
    # extraction exceptions were stored as {'error': ...} dict and the script
    # always exited 0 — pipeline saw "success" with hidden extraction failures.
    missing_files = 0
    unknown_type_count = 0
    extraction_failures = 0
    total_processed = 0

    for file_info in config.get('analysis_files', []):
        file_type = file_info['type']
        file_path = file_info['path']

        if not os.path.exists(file_path):
            print(f"  [QC_BLOCKING] File not found: {file_path}", file=sys.stderr)
            missing_files += 1
            continue

        extractor = EXTRACTORS.get(file_type)
        if not extractor:
            print(f"  [QC_BLOCKING] Unknown file type: {file_type}", file=sys.stderr)
            unknown_type_count += 1
            continue

        print(f"  Extracting {file_type} from {os.path.basename(file_path)}...")
        total_processed += 1
        try:
            data = extractor(file_path)
            extracted[file_type] = data
            print(f"    OK")
        except Exception as e:
            print(f"    [QC_BLOCKING] Extraction failed: {type(e).__name__}: {e}",
                  file=sys.stderr)
            extracted[file_type] = {'type': file_type, 'error': str(e)}
            extraction_failures += 1

    # Add config metadata
    output = {
        'config': {
            'study_design': config.get('study_design', 'cohort'),
            'primary_outcome': config.get('primary_outcome', ''),
            'secondary_outcomes': config.get('secondary_outcomes', []),
            'group_description': config.get('group_description', {}),
            'table_figure_map': config.get('table_figure_map', {}),
            'follow_up_description': config.get('follow_up_description', ''),
            'language': config.get('language', 'english'),
        },
        'extracted': extracted,
        'qc_summary': {
            'total_processed': total_processed,
            'missing_files': missing_files,
            'unknown_types': unknown_type_count,
            'extraction_failures': extraction_failures,
        },
    }

    # Save to JSON
    output_path = '/tmp/manuscript_extracted.json'
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(output, f, ensure_ascii=False, indent=2, default=str)

    print(f"\n  Extracted data saved to: {output_path}")
    print(f"  Analysis types found: {list(extracted.keys())}")
    print(f"  QC summary: total_processed={total_processed}, "
          f"missing_files={missing_files}, unknown_types={unknown_type_count}, "
          f"extraction_failures={extraction_failures}")

    # TP-1 audit fix (mr-read-01): exit non-zero if any blocking failure surfaced
    blocking_count = missing_files + unknown_type_count + extraction_failures
    if blocking_count > 0:
        print(f"\n  [QC_BLOCKING] {blocking_count} blocking issue(s) — exiting non-zero",
              file=sys.stderr)
        sys.exit(2)


if __name__ == '__main__':
    main()
