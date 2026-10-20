#!/usr/bin/env python3
"""
Missing Data Assessment + Multiple Imputation (research-plan pipeline)

Implements references/missing_data_protocol.md. One script, two modes selected by
the config's `mode` field:

    mode = "assessment"   -> Missingness.xlsx + Missingness_pattern.png
    mode = "imputation"   -> MI_{outcome}.xlsx + MI_{outcome}_diagnostics.png

Usage:
    python missing_data.py config.json                  # full run
    python missing_data.py config.json --check-missing  # preview only, no MICE

Design notes
------------
* The MICE engine lives in mice_driver.R (R `mice`). This script owns config
  parsing, filters/derived variables, Excel, plots and QC messaging.
* Complete-case is never silently substituted for MI. If R or mice is missing,
  the script emits [QC_BLOCKING] and exits non-zero.
* Pooled estimates come back from R on the log scale and are exponentiated here,
  so every HR/OR written to Excel is on the ratio scale.
"""

import sys
import os
import json
import math
import shutil
import operator
import tempfile
import subprocess
from collections import Counter

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D

from scipy import stats as sps

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
R_DRIVER = os.path.join(SCRIPT_DIR, 'mice_driver.R')

DEFAULT_RSCRIPT = '/opt/anaconda3/bin/Rscript'

# QC messages are collected so the summary block at the end mirrors what the
# runner scraped off stderr.
_QC_MESSAGES = []


def qc(kind, message):
    """Emit a QC line on stderr with the prefix run_pipeline.py greps for."""
    line = f"[QC_{kind}] {message}"
    _QC_MESSAGES.append(line)
    print(line, file=sys.stderr)


def qc_warning(message):
    qc('WARNING', message)


def qc_blocking(message):
    qc('BLOCKING', message)


# ═══════════════════════════════════════════════════════════════════════════════
# Configuration
# ═══════════════════════════════════════════════════════════════════════════════

def load_config(path):
    with open(path, 'r', encoding='utf-8') as f:
        config = json.load(f)

    config.setdefault('mode', 'assessment')
    config.setdefault('filters', [])
    config.setdefault('derived_variables', [])
    config.setdefault('title', '')
    config.setdefault('subtitle', '')
    config.setdefault('abbreviations', '')

    if config['mode'] == 'assessment':
        config.setdefault('analysis_variables', [])
        config.setdefault('outcomes', [])
        config.setdefault('candidate_auxiliary', [])
        th = config.setdefault('thresholds', {})
        th.setdefault('low', 5)
        th.setdefault('high', 40)
    else:
        config.setdefault('sections', [])
        config.setdefault('auxiliary_variables', [])
        config.setdefault('cox_settings', None)
        config.setdefault('logistic_settings', None)
        mi = config.setdefault('mi_settings', {})
        mi.setdefault('m', None)
        mi.setdefault('seed', 42)
        mi.setdefault('maxit', 20)
        mi.setdefault('method_overrides', {})
        th = config.setdefault('thresholds', {})
        th.setdefault('low', 5)
        th.setdefault('high', 40)

    if config['mode'] not in ('assessment', 'imputation'):
        sys.exit(f"[FATAL] Unknown mode: {config['mode']!r} (expected 'assessment' or 'imputation')")

    return config


# ═══════════════════════════════════════════════════════════════════════════════
# Data loading (same semantics as univariate_forest.py / multivariate_interaction.py)
# ═══════════════════════════════════════════════════════════════════════════════

OPS = {
    '==': operator.eq, '!=': operator.ne,
    '>': operator.gt, '>=': operator.ge,
    '<': operator.lt, '<=': operator.le,
}


def load_data(config):
    path = config['data_file']
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xls'):
        df = pd.read_excel(path)
    elif ext in ('.csv',):
        df = pd.read_csv(path)
    elif ext in ('.tsv',):
        df = pd.read_csv(path, sep='\t')
    elif ext in ('.sas7bdat',):
        df = pd.read_sas(path)
    elif ext in ('.dta',):
        df = pd.read_stata(path)
    else:
        df = pd.read_csv(path)

    for f in config['filters']:
        col, op, val = f['column'], f['operator'], f['value']
        if op in OPS:
            df = df[OPS[op](pd.to_numeric(df[col], errors='coerce'), val)]
        elif op == 'in':
            df = df[df[col].isin(val)]
        elif op == 'not_in':
            df = df[~df[col].isin(val)]

    for dv in config['derived_variables']:
        if dv['type'] == 'coalesce':
            result = pd.to_numeric(df[dv['columns'][0]], errors='coerce')
            for c in dv['columns'][1:]:
                result = result.combine_first(pd.to_numeric(df[c], errors='coerce'))
            df[dv['name']] = result
        elif dv['type'] == 'expression':
            local_ns = {}
            for dep in dv.get('dependencies', []):
                local_ns[dep] = pd.to_numeric(df[dep], errors='coerce')
            local_ns['np'] = np
            df[dv['name']] = eval(dv['expression'], {"__builtins__": {}}, local_ns)

    return df.reset_index(drop=True)


def normalize_series(df, col, var_type):
    """Blank-aware NA handling: categoricals keep their labels, everything else
    is coerced to numeric (mirrors univariate_forest.apply_listwise_deletion)."""
    s = df[col]
    if var_type == 'categorical':
        return s.where(s.notna() & (s.astype(str).str.strip() != ''), np.nan)
    return pd.to_numeric(s, errors='coerce')


# ═══════════════════════════════════════════════════════════════════════════════
# Variable collection
# ═══════════════════════════════════════════════════════════════════════════════

def assessment_variables(config):
    """[{column, type, label}] for mode=assessment."""
    out = []
    for v in config['analysis_variables']:
        out.append({
            'column': v['column'],
            'type': v.get('type', 'continuous'),
            'label': v.get('label', v['column']),
        })
    return out


def imputation_variables(config):
    """[{column, type, label, scale, reference_level, section}] for mode=imputation."""
    out = []
    for section in config['sections']:
        for v in section['variables']:
            out.append({
                'column': v['column'],
                'type': v.get('type', 'continuous'),
                'label': v.get('label', v['column']),
                'scale': v.get('scale'),
                'reference_level': v.get('reference_level'),
                'section': section.get('header', ''),
            })
    return out


def outcome_columns(config):
    """(list_of_outcome_columns, description_rows)."""
    cols, desc = [], []
    if config['mode'] == 'assessment':
        for oc in config['outcomes']:
            if oc.get('type') == 'time_to_event':
                cols += [oc['time_column'], oc['event_column']]
                desc.append({'id': oc.get('id', 'outcome'), 'type': 'time_to_event',
                             'time_column': oc['time_column'], 'event_column': oc['event_column']})
            else:
                c = oc.get('outcome_column') or oc.get('column')
                cols.append(c)
                desc.append({'id': oc.get('id', 'outcome'), 'type': 'binary', 'outcome_column': c})
    else:
        if config.get('analysis_type') == 'cox':
            cs = config['cox_settings']
            cols += [cs['time_column'], cs['event_column']]
            desc.append({'id': config.get('outcome_id', 'primary'), 'type': 'time_to_event',
                         'time_column': cs['time_column'], 'event_column': cs['event_column']})
        else:
            c = config['logistic_settings']['outcome_column']
            cols.append(c)
            desc.append({'id': config.get('outcome_id', 'primary'), 'type': 'binary',
                         'outcome_column': c})
    return [c for c in cols if c], desc


def grouping_info(config):
    gc = config.get('grouping')
    if not gc:
        return None, None, None
    labels = [g.get('label', str(g.get('value'))) for g in gc.get('groups', [])]
    values = [g.get('value') for g in gc.get('groups', [])]
    return gc['column'], values, labels


def split_groups(df, config):
    """dtype-aware group split; abort on an empty group (same guard as the
    other sub-skills, which silently produced empty groups before)."""
    col, values, labels = grouping_info(config)
    if col is None:
        return None
    if col not in df.columns:
        sys.exit(f"[FATAL] Grouping column '{col}' not found in data")
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in values):
        col_vals = pd.to_numeric(df[col], errors='coerce')
    else:
        col_vals = df[col]
    groups = []
    for v, lab in zip(values, labels):
        sub = df[col_vals == v]
        if len(sub) == 0:
            sys.exit(f"[FATAL] split_groups produced an empty group: {lab} "
                     f"(column '{col}', value {v!r})")
        groups.append((lab, v, sub))
    return groups


# ═══════════════════════════════════════════════════════════════════════════════
# Statistical helpers for the mechanism tests
# ═══════════════════════════════════════════════════════════════════════════════

SMALL_N = 30


def test_categorical(indicator, other, other_name):
    """Missingness indicator (0/1) vs a categorical variable.

    chi-square by default; Fisher's exact when the table is 2x2 and any expected
    count < 5 (or n is small). Returns dict or None when the test is impossible.
    """
    mask = indicator.notna() & other.notna()
    a = indicator[mask].astype(int)
    b = other[mask]
    if a.nunique() < 2 or b.nunique() < 2:
        return None
    table = pd.crosstab(a, b)
    if table.shape[0] < 2 or table.shape[1] < 2:
        return None
    try:
        chi2, p_chi, dof, expected = sps.chi2_contingency(table.values)
    except Exception:
        return None
    small_expected = bool((expected < 5).any())
    if table.shape == (2, 2) and (small_expected or table.values.sum() < SMALL_N * 2):
        try:
            odds, p = sps.fisher_exact(table.values)
            return {'test': "Fisher's exact", 'statistic': float(odds),
                    'statistic_name': 'OR', 'p': float(p),
                    'note': 'expected count < 5' if small_expected else 'small n'}
        except Exception:
            pass
    note = ''
    if small_expected:
        note = ('expected count < 5 and table is larger than 2x2; '
                'chi-square p is approximate')
    return {'test': 'Chi-square', 'statistic': float(chi2), 'statistic_name': 'chi2',
            'p': float(p_chi), 'note': note}


def test_continuous(indicator, values, values_name):
    """Missingness indicator (0/1) vs a continuous variable.

    Welch t-test by default; Mann-Whitney U when either arm has n < 30.
    """
    mask = indicator.notna() & values.notna()
    a = indicator[mask].astype(int)
    v = pd.to_numeric(values[mask], errors='coerce')
    ok = v.notna()
    a, v = a[ok], v[ok]
    g0 = v[a == 0].values
    g1 = v[a == 1].values
    if len(g0) < 2 or len(g1) < 2:
        return None
    if min(len(g0), len(g1)) < SMALL_N:
        try:
            u, p = sps.mannwhitneyu(g0, g1, alternative='two-sided')
        except Exception:
            return None
        return {'test': 'Mann-Whitney U', 'statistic': float(u), 'statistic_name': 'U',
                'p': float(p), 'note': f'n < {SMALL_N} in one arm',
                'mean_observed': float(np.mean(g0)), 'mean_missing': float(np.mean(g1))}
    t, p = sps.ttest_ind(g0, g1, equal_var=False)
    return {'test': "Welch's t-test", 'statistic': float(t), 'statistic_name': 't',
            'p': float(p), 'note': '',
            'mean_observed': float(np.mean(g0)), 'mean_missing': float(np.mean(g1))}


def is_numeric_like(series, min_frac=0.8):
    s = series.dropna()
    if len(s) == 0:
        return False
    conv = pd.to_numeric(s, errors='coerce')
    return conv.notna().mean() >= min_frac


# ═══════════════════════════════════════════════════════════════════════════════
# Excel style constants (shared with the other pipeline sub-skills)
# ═══════════════════════════════════════════════════════════════════════════════

DARK_BLUE = "1F3864"
HEADER_BG = "2E5FA3"
SECTION_BG = "EBF2FA"
ALT_ROW = "F7FAFD"
SIG_COLOR = "C00000"
BORDER_COLOR = "BFBFBF"
FOOTNOTE_COLOR = "595959"
FLAG_BG = "FDEBD0"
GOOD_COLOR = "1E8449"

THIN_BORDER = Border(
    left=Side(style='thin', color=BORDER_COLOR),
    right=Side(style='thin', color=BORDER_COLOR),
    top=Side(style='thin', color=BORDER_COLOR),
    bottom=Side(style='thin', color=BORDER_COLOR),
)

TITLE_FONT = Font(name='Arial', size=12, bold=True, color="FFFFFF")
SUBTITLE_FONT = Font(name='Arial', size=10, italic=True, color="FFFFFF")
HEADER_FONT = Font(name='Arial', size=10, bold=True, color="FFFFFF")
SECTION_FONT = Font(name='Arial', size=10, bold=True, color=DARK_BLUE)
DATA_FONT = Font(name='Arial', size=10)
SIG_FONT = Font(name='Arial', size=10, bold=True, color=SIG_COLOR)
FOOTNOTE_FONT = Font(name='Arial', size=8, italic=True, color=FOOTNOTE_COLOR)

TITLE_FILL = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type='solid')
HEADER_FILL = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type='solid')
SECTION_FILL = PatternFill(start_color=SECTION_BG, end_color=SECTION_BG, fill_type='solid')
ALT_FILL = PatternFill(start_color=ALT_ROW, end_color=ALT_ROW, fill_type='solid')
FLAG_FILL = PatternFill(start_color=FLAG_BG, end_color=FLAG_BG, fill_type='solid')

CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', indent=1, wrap_text=True)


def sheet_header(ws, title, subtitle, headers, widths):
    """Write the title / subtitle / header rows. Returns the first data row."""
    ncol = len(headers)
    for c, w in zip(range(1, ncol + 1), widths):
        ws.column_dimensions[get_column_letter(c)].width = w

    r = 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    cell = ws.cell(row=r, column=1, value=title)
    cell.font = TITLE_FONT; cell.alignment = CENTER
    for c in range(1, ncol + 1):
        ws.cell(row=r, column=c).fill = TITLE_FILL
    r += 1

    if subtitle:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        cell = ws.cell(row=r, column=1, value=subtitle)
        cell.font = SUBTITLE_FONT; cell.alignment = CENTER
        for c in range(1, ncol + 1):
            ws.cell(row=r, column=c).fill = TITLE_FILL
        r += 1

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL
        cell.alignment = CENTER; cell.border = THIN_BORDER
    r += 1
    ws.freeze_panes = ws.cell(row=r, column=1)
    return r


def write_row(ws, r, values, alt=False, flag=False, bold_cols=(), sig_cols=()):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = SIG_FONT if c in sig_cols else DATA_FONT
        if c in bold_cols:
            cell.font = Font(name='Arial', size=10, bold=True)
        cell.border = THIN_BORDER
        cell.alignment = LEFT if c == 1 else CENTER
        if flag:
            cell.fill = FLAG_FILL
        elif alt:
            cell.fill = ALT_FILL
    return r + 1


def write_footnote(ws, r, ncol, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    cell = ws.cell(row=r, column=1, value=text)
    cell.font = FOOTNOTE_FONT
    cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    ws.row_dimensions[r].height = max(14, 12 * (1 + text.count('\n')))
    return r + 1


def fmt_p(p):
    if p is None or (isinstance(p, float) and math.isnan(p)):
        return "N/A"
    if p < 0.001:
        return "<0.001*"
    if p < 0.05:
        return f"{p:.3f}*"
    return f"{p:.3f}"


def fmt_effect(effect, lo, hi):
    if effect is None or any(x is None or (isinstance(x, float) and math.isnan(x))
                             for x in (effect, lo, hi)):
        return "N/A"
    return f"{effect:.2f} ({lo:.2f}–{hi:.2f})"


def fmt_num(x, nd=3):
    if x is None or (isinstance(x, float) and (math.isnan(x) or math.isinf(x))):
        return "N/A"
    return f"{x:.{nd}f}"


# ═══════════════════════════════════════════════════════════════════════════════
# MODE 1 — Missingness assessment
# ═══════════════════════════════════════════════════════════════════════════════

def compute_missing_summary(df, variables, groups, extra_rows):
    """Per-variable missing n / % overall and by group."""
    rows = []
    for item in list(variables) + list(extra_rows):
        col = item['column']
        if col not in df.columns:
            rows.append({'label': item.get('label', col), 'column': col,
                         'role': item.get('role', 'analysis'), 'type': item.get('type', ''),
                         'n': 0, 'n_missing': None, 'pct': None, 'groups': [],
                         'note': 'COLUMN NOT FOUND'})
            continue
        s = normalize_series(df, col, item.get('type', 'continuous'))
        n = len(s)
        nm = int(s.isna().sum())
        grp_stats = []
        if groups:
            for lab, val, sub in groups:
                ss = normalize_series(sub, col, item.get('type', 'continuous'))
                gm = int(ss.isna().sum())
                gn = len(ss)
                grp_stats.append({'label': lab, 'n_missing': gm, 'n': gn,
                                  'pct': 100.0 * gm / gn if gn else 0.0})
        rows.append({'label': item.get('label', col), 'column': col,
                     'role': item.get('role', 'analysis'), 'type': item.get('type', ''),
                     'n': n, 'n_missing': nm, 'pct': 100.0 * nm / n if n else 0.0,
                     'groups': grp_stats, 'note': ''})
    return rows


def compute_patterns(df, variables):
    """Missingness pattern matrix + monotone check + top-10 patterns."""
    cols = [v['column'] for v in variables if v['column'] in df.columns]
    if not cols:
        return {'columns': [], 'n_complete': 0, 'pct_complete': 0.0,
                'monotone': None, 'top_patterns': [], 'matrix': None}
    mat = pd.DataFrame({c: normalize_series(df, c, next(
        (v.get('type', 'continuous') for v in variables if v['column'] == c), 'continuous')).isna()
        for c in cols})
    n = len(mat)
    n_complete = int((~mat.any(axis=1)).sum())

    # monotone: order columns by number of missing (ascending); once a variable
    # is missing in a row, every later variable must be missing too.
    order = sorted(cols, key=lambda c: mat[c].sum())
    ordered = mat[order].values
    monotone = True
    for row in ordered:
        seen_missing = False
        for flag in row:
            if flag:
                seen_missing = True
            elif seen_missing:
                monotone = False
                break
        if not monotone:
            break

    keys = mat.apply(lambda r: tuple(bool(x) for x in r), axis=1)
    counter = Counter(keys)
    top = []
    for pat, cnt in counter.most_common(10):
        missing_names = [c for c, flag in zip(cols, pat) if flag]
        top.append({
            'pattern': ''.join('0' if flag else '1' for flag in pat),
            'missing_vars': missing_names,
            'n_missing_vars': len(missing_names),
            'n': int(cnt),
            'pct': 100.0 * cnt / n if n else 0.0,
        })
    return {'columns': cols, 'n_complete': n_complete,
            'pct_complete': 100.0 * n_complete / n if n else 0.0,
            'monotone': monotone, 'top_patterns': top, 'matrix': mat,
            'n_patterns': len(counter), 'n_total': n}


def compute_listwise_impact(df, variables, out_cols, group_col, groups, outcome_desc):
    """What complete-case analysis actually costs, overall and per group."""
    cols = [v['column'] for v in variables if v['column'] in df.columns]
    cols += [c for c in out_cols if c in df.columns and c not in cols]
    if group_col and group_col in df.columns and group_col not in cols:
        cols.append(group_col)

    work = df.copy()
    types = {v['column']: v.get('type', 'continuous') for v in variables}
    for c in cols:
        work[c] = normalize_series(work, c, types.get(c, 'continuous'))
    complete = work.dropna(subset=cols)

    def events_of(sub):
        ev = {}
        for od in outcome_desc:
            if od['type'] == 'time_to_event':
                col = od['event_column']
            else:
                col = od['outcome_column']
            if col in sub.columns:
                ev[od['id']] = int(pd.to_numeric(sub[col], errors='coerce').fillna(0).sum())
        return ev

    rows = [{'population': 'Overall', 'n_before': len(df), 'n_after': len(complete),
             'events_before': events_of(df), 'events_after': events_of(complete)}]
    if groups:
        for lab, val, sub in groups:
            if group_col in complete.columns:
                gvals = complete[group_col]
                if isinstance(val, (int, float)) and not isinstance(val, bool):
                    gvals = pd.to_numeric(gvals, errors='coerce')
                sub_after = complete[gvals == val]
            else:
                sub_after = complete.iloc[0:0]
            rows.append({'population': lab, 'n_before': len(sub), 'n_after': len(sub_after),
                         'events_before': events_of(sub), 'events_after': events_of(sub_after)})
    for r in rows:
        r['n_dropped'] = r['n_before'] - r['n_after']
        r['pct_dropped'] = 100.0 * r['n_dropped'] / r['n_before'] if r['n_before'] else 0.0
    return rows, cols


def compute_mechanism_tests(df, variables, group_col, outcome_desc):
    """For every variable with missing values, test the missingness indicator
    against exposure/group, outcome event and follow-up time."""
    results = []
    targets = []
    if group_col and group_col in df.columns:
        targets.append(('Group / exposure', group_col, 'categorical'))
    for od in outcome_desc:
        if od['type'] == 'time_to_event':
            if od['event_column'] in df.columns:
                targets.append((f"Outcome event [{od['id']}]", od['event_column'], 'categorical'))
            if od['time_column'] in df.columns:
                targets.append((f"Follow-up time [{od['id']}]", od['time_column'], 'continuous'))
        else:
            if od['outcome_column'] in df.columns:
                targets.append((f"Outcome [{od['id']}]", od['outcome_column'], 'categorical'))

    for v in variables:
        col = v['column']
        if col not in df.columns:
            continue
        s = normalize_series(df, col, v.get('type', 'continuous'))
        nm = int(s.isna().sum())
        if nm == 0:
            continue
        ind = s.isna().astype(int)
        for target_label, target_col, target_kind in targets:
            if target_col == col:
                continue
            if target_kind == 'categorical':
                other = df[target_col]
                if is_numeric_like(other):
                    other = pd.to_numeric(other, errors='coerce')
                res = test_categorical(ind, other, target_col)
            else:
                res = test_continuous(ind, pd.to_numeric(df[target_col], errors='coerce'), target_col)
            results.append({
                'variable': v.get('label', col), 'column': col,
                'n_missing': nm, 'pct': 100.0 * nm / len(df) if len(df) else 0.0,
                'target_label': target_label, 'target_column': target_col,
                'result': res,
            })
    return results


def compute_auxiliary_candidates(df, variables, config, group_col, out_cols):
    """A useful auxiliary is correlated with the incomplete variable AND with its
    missingness indicator. Every evaluated pair is reported; the `recommended`
    column is the flag, nothing is silently filtered out."""
    declared = list(config.get('candidate_auxiliary') or [])
    analysis_cols = {v['column'] for v in variables}
    excluded = analysis_cols | set(out_cols) | ({group_col} if group_col else set())

    if declared:
        candidates = [c for c in declared if c in df.columns]
        auto = False
    else:
        candidates = []
        for c in df.columns:
            if c in excluded:
                continue
            s = df[c]
            if s.isna().mean() > 0.5:
                continue
            if s.nunique(dropna=True) < 2 or s.nunique(dropna=True) > 0.9 * len(s):
                continue
            if not is_numeric_like(s):
                continue
            candidates.append(c)
        auto = True

    rows = []
    for v in variables:
        col = v['column']
        if col not in df.columns:
            continue
        s = normalize_series(df, col, v.get('type', 'continuous'))
        if s.isna().sum() == 0:
            continue
        ind = s.isna().astype(int)
        s_num = pd.to_numeric(s, errors='coerce')
        for cand in candidates:
            if cand == col:
                continue
            cs = df[cand]
            cand_numeric = is_numeric_like(cs)
            cs_num = pd.to_numeric(cs, errors='coerce') if cand_numeric else cs

            # (a) association with the incomplete variable, among observed rows
            rho = p_rho = None
            corr_test = 'Spearman rho'
            if cand_numeric and s_num.notna().sum() > 3:
                mask = s_num.notna() & cs_num.notna()
                if mask.sum() > 3 and s_num[mask].nunique() > 1 and cs_num[mask].nunique() > 1:
                    rho, p_rho = sps.spearmanr(s_num[mask], cs_num[mask])
                    rho, p_rho = float(rho), float(p_rho)
            elif not cand_numeric and s_num.notna().sum() > 3:
                corr_test = 'Kruskal-Wallis'
                mask = s_num.notna() & cs.notna()
                grps = [g.values for _, g in s_num[mask].groupby(cs[mask])]
                grps = [g for g in grps if len(g) > 1]
                if len(grps) >= 2:
                    try:
                        h, p_rho = sps.kruskal(*grps)
                        rho, p_rho = float(h), float(p_rho)
                    except Exception:
                        rho = p_rho = None

            # (b) association with the missingness indicator
            if cand_numeric:
                miss_res = test_continuous(ind, cs_num, cand)
            else:
                miss_res = test_categorical(ind, cs, cand)

            p_miss = miss_res['p'] if miss_res else None
            recommended = (p_rho is not None and p_rho < 0.05 and
                           p_miss is not None and p_miss < 0.05)
            rows.append({
                'variable': v.get('label', col), 'column': col, 'candidate': cand,
                'corr_test': corr_test, 'corr_stat': rho, 'corr_p': p_rho,
                'miss_test': miss_res['test'] if miss_res else 'not testable',
                'miss_stat': miss_res['statistic'] if miss_res else None,
                'miss_p': p_miss,
                'recommended': recommended,
            })
    return rows, candidates, auto


def compute_verdict(summary_rows, mech_results, out_cols, thresholds):
    """Protocol section 2 decision table."""
    low, high = thresholds['low'], thresholds['high']
    drivers = []

    analysis_rows = [r for r in summary_rows if r['role'] == 'analysis' and r['n_missing'] is not None]
    outcome_rows = [r for r in summary_rows if r['role'] == 'outcome' and r['n_missing'] is not None]

    max_pct = max([r['pct'] for r in analysis_rows], default=0.0)
    over_high = [r for r in analysis_rows if r['pct'] > high]
    over_low = [r for r in analysis_rows if r['pct'] >= low]

    outcome_missing = [r for r in outcome_rows if r['n_missing'] > 0]

    outcome_assoc = []
    for m in mech_results:
        res = m['result']
        if res and res['p'] is not None and res['p'] < 0.05 and 'Outcome' in m['target_label']:
            outcome_assoc.append(m)
    group_assoc = [m for m in mech_results
                   if m['result'] and m['result']['p'] is not None and m['result']['p'] < 0.05
                   and 'Group' in m['target_label']]
    time_assoc = [m for m in mech_results
                  if m['result'] and m['result']['p'] is not None and m['result']['p'] < 0.05
                  and 'Follow-up' in m['target_label']]

    if outcome_missing:
        verdict = 'outcome_missing'
        drivers.append("Outcome variable(s) themselves contain missing values: "
                       + ", ".join(f"{r['column']} ({r['pct']:.1f}%)" for r in outcome_missing))
        rec = ("MI-then-delete (von Hippel 2007) is required: the outcome enters the imputation "
               "model to help impute covariates, but cases with a missing outcome are excluded "
               "from the analysis. This pipeline does not do that automatically -- handle it "
               "manually before running the imputation step.")
    elif outcome_assoc:
        verdict = 'mcar_refuted'
        drivers.append("Missingness is significantly associated with the outcome for: "
                       + ", ".join(sorted({m['variable'] for m in outcome_assoc})))
        rec = ("MCAR is refuted: missingness depends on the OBSERVED outcome. That rules "
               "out MCAR, but it does not diagnose MNAR -- missingness that depends on "
               "observed data is exactly the MAR case, and MAR cannot be tested from the "
               "data either way. The right response is to make MAR plausible: put the "
               "outcome in the imputation model (this pipeline does), and add auxiliary "
               "variables that predict both the missing values and their missingness. "
               "Complete-case is expected to be biased here. Reach for an MNAR "
               "sensitivity analysis (delta-adjustment / tipping point), which this pipeline "
               "does not generate.")
    elif over_high:
        verdict = 'mi_with_caution'
        drivers.append(f"Variable(s) above the {high}% threshold: "
                       + ", ".join(f"{r['label']} ({r['pct']:.1f}%)" for r in over_high))
        rec = (f"MI remains usable, but imputed values for these variables lean heavily on the "
               f"imputation model. Consider dropping the variable or sourcing it elsewhere; if it "
               f"is kept, run a with/without sensitivity analysis. Note that >50% missing does NOT "
               f"mandate exclusion (Madley-Dowd 2019) -- dropping it is a decision that belongs in "
               f"decision_log.md with a permanent ID.")
    elif over_low:
        verdict = 'mi_recommended'
        drivers.append(f"Variable(s) at or above the {low}% threshold: "
                       + ", ".join(f"{r['label']} ({r['pct']:.1f}%)" for r in over_low))
        rec = ("Run multiple imputation. Which of MI or complete-case is the primary analysis must "
               "be pre-specified and registered in decision_log.md -- the pipeline computes both "
               "and does not choose for you.")
    else:
        verdict = 'complete_case_ok'
        drivers.append(f"All analysis variables are below {low}% missing "
                       f"(maximum {max_pct:.1f}%)")
        rec = ("Complete-case analysis is defensible. Still report every variable's missing "
               "proportion in the manuscript.")

    if group_assoc and verdict not in ('outcome_missing', 'mcar_refuted'):
        drivers.append("Missingness differs by group/exposure for: "
                       + ", ".join(sorted({m['variable'] for m in group_assoc}))
                       + " -- MCAR is not tenable, but MAR conditional on the group may hold.")
    if time_assoc and verdict not in ('outcome_missing', 'mcar_refuted'):
        drivers.append("Missingness is associated with follow-up time for: "
                       + ", ".join(sorted({m['variable'] for m in time_assoc})))

    return {'verdict': verdict, 'drivers': drivers, 'recommendation': rec,
            'max_pct': max_pct, 'thresholds': thresholds}


# ── assessment Excel ──────────────────────────────────────────────────────────

def write_assessment_excel(config, summary_rows, patterns, listwise_rows, listwise_cols,
                           mech_results, aux_rows, aux_candidates, aux_auto, verdict,
                           groups, n_total):
    wb = Workbook()
    title = config.get('title') or 'Missing Data Assessment'
    subtitle = config.get('subtitle') or f'N = {n_total}'
    group_labels = [g[0] for g in groups] if groups else []

    # ── Summary ────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Summary'
    headers = ['Variable', 'Column', 'Role', 'Type', 'N', 'Missing n', 'Missing %']
    widths = [30, 20, 12, 13, 8, 11, 11]
    for lab in group_labels:
        headers += [f'{lab}: missing n (%)']
        widths += [22]
    if len(group_labels) == 2:
        headers += ['Difference (pp)']
        widths += [16]
    r = sheet_header(ws, title, subtitle, headers, widths)
    for i, row in enumerate(summary_rows):
        vals = [row['label'], row['column'], row['role'], row['type'],
                row['n'], row['n_missing'] if row['n_missing'] is not None else row['note'],
                fmt_num(row['pct'], 1) + '%' if row['pct'] is not None else 'N/A']
        for g in row['groups']:
            vals.append(f"{g['n_missing']}/{g['n']} ({g['pct']:.1f}%)")
        if len(group_labels) == 2 and len(row['groups']) == 2:
            vals.append(f"{row['groups'][1]['pct'] - row['groups'][0]['pct']:+.1f}")
        flag = row['pct'] is not None and row['pct'] >= config['thresholds']['low']
        r = write_row(ws, r, vals, alt=(i % 2 == 1), flag=flag)
    r += 1
    r = write_footnote(ws, r, len(headers),
                       "Shaded rows are at or above the low threshold "
                       f"({config['thresholds']['low']}% missing). Group-stratified missing rates "
                       "are shown because a difference between groups is itself evidence against MCAR. "
                       "Percentages are of the filtered analysis population, not the raw file.")

    # ── Pattern ────────────────────────────────────────────────────────────
    ws = wb.create_sheet('Pattern')
    headers = ['Rank', 'Pattern (1 = observed, 0 = missing)', 'Variables missing',
               '# variables missing', 'n', '% of cohort']
    r = sheet_header(ws, 'Missingness patterns', subtitle, headers, [8, 34, 46, 18, 10, 12])
    r = write_row(ws, r, ['Complete cases', '-', '-', 0,
                          patterns['n_complete'], f"{patterns['pct_complete']:.1f}%"],
                  bold_cols=(1,))
    r = write_row(ws, r, ['Distinct patterns', '-', '-', '-', patterns.get('n_patterns', 0), '-'],
                  alt=True, bold_cols=(1,))
    mono = patterns['monotone']
    r = write_row(ws, r, ['Monotone pattern?',
                          'Yes' if mono else ('No' if mono is not None else 'N/A'),
                          'Monotone patterns can be imputed sequentially; non-monotone patterns '
                          'require the full MICE (fully conditional specification) algorithm.',
                          '-', '-', '-'], bold_cols=(1,))
    r += 1
    r = write_row(ws, r, ['Top 10 patterns', '', '', '', '', ''], bold_cols=(1,))
    for i, p in enumerate(patterns['top_patterns'], 1):
        r = write_row(ws, r, [i, p['pattern'],
                              ', '.join(p['missing_vars']) if p['missing_vars'] else '(none missing)',
                              p['n_missing_vars'], p['n'], f"{p['pct']:.1f}%"],
                      alt=(i % 2 == 1))
    r += 1
    r = write_footnote(ws, r, len(headers),
                       "Pattern string reads left to right in this variable order: "
                       + ", ".join(patterns['columns']))

    # ── Listwise_impact ────────────────────────────────────────────────────
    ws = wb.create_sheet('Listwise_impact')
    headers = ['Population', 'N before', 'N after listwise deletion', 'N dropped', '% dropped',
               'Events before', 'Events after']
    r = sheet_header(ws, 'Cost of doing nothing (complete-case / listwise deletion)',
                     subtitle, headers, [24, 12, 26, 12, 12, 18, 18])
    for i, row in enumerate(listwise_rows):
        ev_b = ', '.join(f"{k}: {v}" for k, v in row['events_before'].items()) or '-'
        ev_a = ', '.join(f"{k}: {v}" for k, v in row['events_after'].items()) or '-'
        r = write_row(ws, r, [row['population'], row['n_before'], row['n_after'],
                              row['n_dropped'], f"{row['pct_dropped']:.1f}%", ev_b, ev_a],
                      alt=(i % 2 == 1), flag=(row['pct_dropped'] >= 10))
    r += 1
    r = write_footnote(ws, r, len(headers),
                       "Listwise deletion here drops a case if ANY of these columns is missing: "
                       + ", ".join(listwise_cols)
                       + ". Rows are shaded when 10% or more of the population is lost. This is the "
                         "sample size the regression models actually run on today.")

    # ── Mechanism_tests ────────────────────────────────────────────────────
    ws = wb.create_sheet('Mechanism_tests')
    headers = ['Variable', 'Column', 'Missing n', 'Missing %', 'Tested against',
               'Test', 'Statistic', 'p-value', 'MCAR violated?', 'Note']
    r = sheet_header(ws, 'Missingness mechanism tests (missingness indicator vs observed data)',
                     subtitle, headers, [26, 18, 11, 11, 26, 18, 14, 12, 15, 30])
    if not mech_results:
        r = write_row(ws, r, ['(no variable has missing values -- nothing to test)',
                              '', '', '', '', '', '', '', '', ''])
    for i, m in enumerate(mech_results):
        res = m['result']
        if res is None:
            vals = [m['variable'], m['column'], m['n_missing'], f"{m['pct']:.1f}%",
                    m['target_label'], 'not testable', 'N/A', 'N/A', 'N/A',
                    'insufficient variation or sample size']
            sig_cols = ()
        else:
            violated = 'YES' if res['p'] < 0.05 else 'no'
            vals = [m['variable'], m['column'], m['n_missing'], f"{m['pct']:.1f}%",
                    m['target_label'], res['test'],
                    f"{res['statistic_name']} = {res['statistic']:.3f}",
                    fmt_p(res['p']), violated, res.get('note', '')]
            sig_cols = (8, 9) if res['p'] < 0.05 else ()
        r = write_row(ws, r, vals, alt=(i % 2 == 1), sig_cols=sig_cols,
                      flag=(res is not None and res['p'] < 0.05))
    r += 1
    r = write_footnote(ws, r, len(headers),
                       "A missingness indicator (1 = missing) was built for every incomplete variable and "
                       "tested against the group/exposure, the outcome event and the follow-up time. "
                       "Categorical comparisons use chi-square, switching to Fisher's exact for 2x2 tables "
                       f"with an expected count below 5; continuous comparisons use Welch's t-test, switching "
                       f"to Mann-Whitney U when either arm has fewer than {SMALL_N} observations. "
                       "A significant result refutes MCAR. A non-significant result does NOT establish MCAR -- "
                       "it only means this dataset did not detect a departure from it, and MAR itself is not "
                       "testable from observed data.")

    # ── Auxiliary_candidates ───────────────────────────────────────────────
    ws = wb.create_sheet('Auxiliary_candidates')
    headers = ['Incomplete variable', 'Column', 'Candidate auxiliary',
               'Assoc. with variable (test)', 'Statistic', 'p-value',
               'Assoc. with missingness (test)', 'Statistic', 'p-value', 'Recommended?']
    r = sheet_header(ws, 'Auxiliary variable candidates for the imputation model',
                     subtitle, headers, [24, 16, 22, 22, 12, 11, 24, 12, 11, 14])
    if not aux_rows:
        r = write_row(ws, r, ['(no incomplete variable, or no candidate column available)',
                              '', '', '', '', '', '', '', '', ''])
    for i, a in enumerate(aux_rows):
        vals = [a['variable'], a['column'], a['candidate'],
                a['corr_test'], fmt_num(a['corr_stat']), fmt_p(a['corr_p']),
                a['miss_test'], fmt_num(a['miss_stat']), fmt_p(a['miss_p']),
                'YES' if a['recommended'] else 'no']
        r = write_row(ws, r, vals, alt=(i % 2 == 1), flag=a['recommended'])
    r += 1
    src = ('auto-detected numeric columns not already in the analysis model'
           if aux_auto else 'the candidate_auxiliary list in the config')
    r = write_footnote(ws, r, len(headers),
                       f"Candidates came from {src}: "
                       + (", ".join(aux_candidates) if aux_candidates else "(none)")
                       + ". EVERY evaluated pair is listed -- nothing was filtered out. "
                         "'Recommended' marks pairs where the candidate is associated both with the "
                         "incomplete variable itself (p < 0.05) and with its missingness indicator "
                         "(p < 0.05); those belong in the imputation model even when they are not in "
                         "the analysis model. p-values here are unadjusted and exploratory.")

    # ── Verdict ────────────────────────────────────────────────────────────
    ws = wb.create_sheet('Verdict')
    headers = ['Item', 'Value']
    r = sheet_header(ws, 'Verdict', subtitle, headers, [34, 110])
    r = write_row(ws, r, ['Verdict code', verdict['verdict']], bold_cols=(1, 2))
    labels = {
        'complete_case_ok': 'Complete-case analysis is acceptable',
        'mi_recommended': 'Multiple imputation recommended',
        'mi_with_caution': 'Multiple imputation possible, but with caution',
        'mcar_refuted': 'MCAR refuted (missingness depends on observed data) -- complete-case expected to be biased; MAR plausibility needs review',
        'outcome_missing': 'Outcome itself is missing -- MI-then-delete required',
    }
    r = write_row(ws, r, ['Verdict', labels.get(verdict['verdict'], verdict['verdict'])], alt=True)
    r = write_row(ws, r, ['Maximum missing % (analysis variables)',
                          f"{verdict['max_pct']:.1f}%"])
    r = write_row(ws, r, ['Thresholds used',
                          f"low = {verdict['thresholds']['low']}%, high = {verdict['thresholds']['high']}%"],
                  alt=True)
    r += 1
    r = write_row(ws, r, ['Why', ''], bold_cols=(1,))
    for i, d in enumerate(verdict['drivers']):
        r = write_row(ws, r, [f'  driver {i + 1}', d], alt=(i % 2 == 1))
    r += 1
    r = write_row(ws, r, ['Recommendation', verdict['recommendation']], bold_cols=(1,))
    r += 1
    r = write_footnote(ws, r, len(headers),
                       "This verdict is advisory. It does not change any other pipeline step: Table 1, "
                       "Kaplan-Meier, PSM and AUC keep their existing complete-case behaviour, so their N "
                       "will not match an MI model's N. Whether MI or complete-case is the primary analysis "
                       "is a pre-specified human decision that belongs in decision_log.md with a permanent ID.")

    out = config['output_file']
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    return [ws.title for ws in wb.worksheets]


def plot_assessment(config, patterns, summary_rows):
    """Pattern matrix heat map + per-variable missing-rate bars."""
    path = config.get('plot_file')
    if not path:
        return None
    mat = patterns['matrix']
    if mat is None or mat.shape[1] == 0:
        return None

    cols = list(mat.columns)
    order = mat.sum(axis=1).sort_values(kind='stable').index
    arr = mat.loc[order].values.astype(float)

    fig, axes = plt.subplots(1, 2, figsize=(max(10, 0.6 * len(cols) + 8), 6),
                             gridspec_kw={'width_ratios': [2.0, 1.2]})

    ax = axes[0]
    ax.imshow(arr, aspect='auto', interpolation='nearest',
              cmap=matplotlib.colors.ListedColormap(['#DCE6F1', '#C0392B']),
              vmin=0, vmax=1)
    ax.set_xticks(range(len(cols)))
    ax.set_xticklabels(cols, rotation=45, ha='right', fontsize=8)
    ax.set_ylabel('Cases (sorted by number of missing values)', fontsize=9)
    ax.set_title('Missingness pattern matrix', fontsize=11, fontweight='bold')
    ax.set_yticks([])
    ax.legend(handles=[Line2D([0], [0], marker='s', color='w', markerfacecolor='#DCE6F1',
                              markersize=10, label='observed'),
                       Line2D([0], [0], marker='s', color='w', markerfacecolor='#C0392B',
                              markersize=10, label='missing')],
              loc='upper right', fontsize=8, framealpha=0.9)

    ax = axes[1]
    rows = [r for r in summary_rows if r['pct'] is not None]
    rows = sorted(rows, key=lambda r: r['pct'])
    y = np.arange(len(rows))
    ax.barh(y, [r['pct'] for r in rows], color='#2E5FA3')
    ax.set_yticks(y)
    ax.set_yticklabels([r['label'] for r in rows], fontsize=8)
    ax.set_xlabel('% missing', fontsize=9)
    ax.set_title('Missing rate per variable', fontsize=11, fontweight='bold')
    low = config['thresholds']['low']; high = config['thresholds']['high']
    xmax = max(10, max([r['pct'] for r in rows], default=0) * 1.25)
    if low <= xmax:
        ax.axvline(low, color='#E67E22', ls='--', lw=1, label=f'low threshold {low}%')
    if high <= xmax:
        ax.axvline(high, color='#C0392B', ls='--', lw=1, label=f'high threshold {high}%')
    for i, r_ in enumerate(rows):
        ax.text(r_['pct'] + 0.4, i, f"{r_['pct']:.1f}%", va='center', fontsize=7)
    handles, labels = ax.get_legend_handles_labels()
    if handles:
        ax.legend(fontsize=8)
    if high > xmax:
        ax.set_xlabel(f'% missing (high threshold {high}% is off scale)', fontsize=9)
    ax.set_xlim(0, xmax)

    fig.tight_layout()
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    fig.savefig(path, dpi=300, bbox_inches='tight')
    plt.close(fig)
    return path


# ═══════════════════════════════════════════════════════════════════════════════
# MODE 2 — Multiple imputation
# ═══════════════════════════════════════════════════════════════════════════════

def resolve_rscript():
    """$RSCRIPT wins outright when set -- an explicit override that does not
    exist is a failure, not a reason to quietly use a different interpreter."""
    override = os.environ.get('RSCRIPT')
    if override:
        return override if os.path.exists(override) else None
    for cand in (DEFAULT_RSCRIPT, shutil.which('Rscript')):
        if cand and os.path.exists(cand):
            return cand
    return None


def check_r_available():
    """(ok, rscript_path, message)."""
    rscript = resolve_rscript()
    if not rscript:
        override = os.environ.get('RSCRIPT')
        if override:
            return False, None, f"$RSCRIPT is set to {override!r} but that path does not exist"
        return False, None, ("Rscript executable not found (looked at $RSCRIPT, "
                             f"{DEFAULT_RSCRIPT}, and $PATH)")
    env = dict(os.environ, LANG='en_US.UTF-8', LC_ALL='en_US.UTF-8')
    probe = ('cat(all(sapply(c("mice","survival","jsonlite"), '
             'function(p) requireNamespace(p, quietly=TRUE))))')
    try:
        res = subprocess.run([rscript, '-e', probe], capture_output=True, text=True,
                             timeout=120, env=env)
    except Exception as e:
        return False, rscript, f"Rscript probe failed: {e}"
    if res.returncode != 0 or 'TRUE' not in res.stdout:
        return False, rscript, ("R is present but one of mice / survival / jsonlite is not "
                                f"installed (probe output: {res.stdout.strip()!r})")
    return True, rscript, "ok"


def method_for(var_type, n_levels):
    if var_type == 'continuous':
        return 'pmm'
    if var_type == 'binary':
        return 'logreg'
    if var_type == 'ordinal':
        return 'polr' if (n_levels is None or n_levels > 2) else 'logreg'
    if var_type == 'categorical':
        return 'polyreg' if (n_levels is None or n_levels > 2) else 'logreg'
    return 'pmm'


def compute_m(max_pct):
    """protocol 4.4: m = max(20, ceil(max missing %)), capped at 100.

    `max_pct` is the maximum over the *analysis* variables only; auxiliaries are
    excluded on purpose (see the call sites).
    """
    return int(min(100, max(20, math.ceil(max_pct))))


def detect_competing_risk(config, df):
    """Return a reason string when the request looks like competing risks."""
    if config.get('analysis_type') in ('competing_risk', 'fine_gray', 'cif'):
        return f"analysis_type = {config['analysis_type']!r}"
    for key in ('competing_risk_settings', 'cif_settings', 'fine_gray_settings'):
        if config.get(key):
            return f"config contains {key!r}"
    cs = config.get('cox_settings') or {}
    for key in ('competing_event_column', 'competing_risk_column', 'cause_column'):
        if cs.get(key):
            return f"cox_settings.{key} = {cs[key]!r}"
    if config.get('analysis_type') == 'cox' and cs.get('event_column') in df.columns:
        vals = pd.to_numeric(df[cs['event_column']], errors='coerce').dropna().unique()
        vals = sorted(v for v in vals)
        if len(vals) > 2:
            return (f"event column {cs['event_column']!r} has {len(vals)} distinct codes "
                    f"({', '.join(str(v) for v in vals[:6])}), which is multi-state / "
                    f"competing-risk coding rather than a 0/1 indicator")
    return None


CLUSTER_HINT = ('cluster', 'center', 'centre', 'site', 'hospital', 'ward',
                'institut', 'facility', 'physician', 'surgeon', 'clinic')
# Bare 'weight' is deliberately NOT here: body weight is a routine covariate in
# cardiology cohorts and would raise a PSM warning on every study.
PSM_HINT = ('psm', 'matched', 'ps_weight', 'ps_score', 'propensity', 'iptw')


def detect_cluster(config, df, model_cols):
    """Explicit config key, else a name-based scan of the data columns."""
    for key in ('cluster_column', 'cluster_variable', 'cluster_variables', 'id_cluster'):
        if config.get(key):
            return f"config key {key!r} = {config[key]!r}"
    mi = config.get('mi_settings') or {}
    if mi.get('cluster_column'):
        return f"mi_settings.cluster_column = {mi['cluster_column']!r}"
    n = len(df)
    for col in df.columns:
        lc = str(col).lower()
        if not any(h in lc for h in CLUSTER_HINT):
            continue
        nu = df[col].nunique(dropna=True)
        if 1 < nu <= max(2, n / 5):
            return (f"column {col!r} looks like a cluster identifier "
                    f"({nu} distinct values over {n} rows)")
    return None


def detect_psm(config, df):
    for key in ('psm_settings', 'psm', 'matching_settings'):
        if config.get(key):
            return f"config contains {key!r}"
    for f in config.get('filters', []):
        lc = str(f.get('column', '')).lower()
        if any(h in lc for h in PSM_HINT):
            return f"filter on column {f['column']!r}"
    for col in df.columns:
        lc = str(col).lower()
        if lc.startswith(('psm', 'ps_', 'match')) or 'propensity' in lc or 'iptw' in lc:
            return f"data contains column {col!r}"
    return None


def run_mice(config, df, variables, aux_cols, out_cols, out_desc, m, seed, maxit,
             rscript, workdir):
    """Write the prepared frame + spec, call mice_driver.R, return the parsed JSON."""
    keep = []
    for c in [v['column'] for v in variables] + list(aux_cols) + list(out_cols):
        if c not in keep and c in df.columns:
            keep.append(c)
    sub = df[keep].copy()
    types = {v['column']: v.get('type', 'continuous') for v in variables}
    for c in keep:
        sub[c] = normalize_series(sub, c, types.get(c, 'continuous'))

    data_csv = os.path.join(workdir, 'mi_input.csv')
    sub.to_csv(data_csv, index=False, na_rep='NA')

    spec = {
        'data_csv': data_csv,
        'analysis_type': config['analysis_type'],
        'model_variables': [
            {'column': v['column'], 'type': v['type'],
             'scale': v.get('scale'), 'reference_level': v.get('reference_level')}
            for v in variables
        ],
        'auxiliary_variables': list(aux_cols),
        'm': m, 'seed': seed, 'maxit': maxit,
        'method_overrides': config['mi_settings'].get('method_overrides') or {},
        'max_dist_values': 4000,
    }
    if config['analysis_type'] == 'cox':
        spec['time_column'] = config['cox_settings']['time_column']
        spec['event_column'] = config['cox_settings']['event_column']
    else:
        spec['outcome_column'] = config['logistic_settings']['outcome_column']

    spec_path = os.path.join(workdir, 'mi_spec.json')
    res_path = os.path.join(workdir, 'mi_result.json')
    with open(spec_path, 'w', encoding='utf-8') as f:
        json.dump(spec, f, ensure_ascii=False)

    env = dict(os.environ, LANG='en_US.UTF-8', LC_ALL='en_US.UTF-8')
    proc = subprocess.run([rscript, R_DRIVER, spec_path, res_path],
                          capture_output=True, text=True, timeout=3600, env=env)
    if proc.stdout.strip():
        for line in proc.stdout.strip().splitlines():
            print(f"  [R] {line}")
    if not os.path.exists(res_path):
        return {'ok': False,
                'error': (proc.stderr.strip() or proc.stdout.strip()
                          or f'mice_driver.R exited with {proc.returncode} and wrote no result')}
    with open(res_path, 'r', encoding='utf-8') as f:
        result = json.load(f)
    if not result.get('ok') and proc.stderr.strip():
        result.setdefault('stderr', proc.stderr.strip())
    return result


def match_term(term, variables):
    """Map an R coefficient name back to (label, level_suffix)."""
    t = str(term).strip('`')
    for v in variables:
        if t == v['column']:
            return v['label'], ''
    best = None
    for v in variables:
        c = v['column']
        if t.startswith(c) and (best is None or len(c) > len(best['column'])):
            best = v
    if best is not None:
        return best['label'], t[len(best['column']):]
    return t, ''


# Convergence screening thresholds. Both statistics are scaled by the typical
# within-chain fluctuation, so they are unit-free and comparable across
# variables. Calibration note: for well-mixed chains the drift statistic sits
# around 0.0-0.2 and the mixing ratio around 0.3, because a chain-tail mean
# averages away roughly sqrt(maxit/2) of the within-chain noise.
DRIFT_THRESHOLD = 0.5
MIXING_THRESHOLD = 1.0


def convergence_stats(chain_means):
    """(drift, mixing) for one variable's chain-mean traces.

    drift  = |mean of the last half - mean of the first half| of the across-chain
             ensemble mean, divided by the typical within-chain SD. Picks up a
             sampler that is still moving.
    mixing = SD across chains of the last-half chain means, divided by the typical
             within-chain SD. Picks up chains sitting in different places.
    """
    chains = [np.asarray(c, dtype=float) for c in (chain_means or []) if c is not None]
    chains = [c[np.isfinite(c)] for c in chains]
    chains = [c for c in chains if len(c) > 0]
    if not chains:
        return None, None
    L = min(len(c) for c in chains)
    arr = np.vstack([c[:L] for c in chains])          # m x maxit
    within = float(np.mean(np.std(arr, axis=1))) if L > 1 else 0.0
    # Below ~8 iterations each half-mean averages only a handful of points, so
    # both statistics are dominated by noise and would flag healthy runs.
    if L < 8 or within <= 0:
        return None, None
    half = L // 2
    ens = arr.mean(axis=0)
    drift = float(abs(ens[half:].mean() - ens[:half].mean()) / within)
    tail_means = arr[:, half:].mean(axis=1)
    mixing = float(np.std(tail_means) / within) if arr.shape[0] > 1 else 0.0
    return drift, mixing


def assess_convergence(conv):
    """Machine-readable convergence read-out per imputed variable."""
    notes = []
    for v in (conv.get('variables') or []):
        drift, mixing = convergence_stats(conv.get('chain_mean', {}).get(v))
        sds = conv.get('chain_sd', {}).get(v) or []
        if drift is None:
            notes.append({'variable': v, 'verdict': 'not assessed', 'trend': None,
                          'between_within': None,
                          'note': 'too few iterations or a degenerate trace; '
                                  'read the trace panel in the diagnostics PNG directly'})
            continue

        problems = []
        if drift > DRIFT_THRESHOLD:
            problems.append(f'chain mean shifted {drift:.2f} within-chain SD between the '
                            f'first and second half of the iterations '
                            f'(threshold {DRIFT_THRESHOLD})')
        if mixing > MIXING_THRESHOLD:
            problems.append(f'chains separated: between-chain spread is {mixing:.2f}x the '
                            f'within-chain SD (threshold {MIXING_THRESHOLD})')
        sd_arr = [np.asarray(s, dtype=float) for s in sds] if sds else []
        sd_flat = np.concatenate(sd_arr) if sd_arr else np.array([])
        if sd_flat.size and not np.all(np.isfinite(sd_flat)):
            problems.append('non-finite chain SD')

        notes.append({
            'variable': v,
            'verdict': 'inspect' if problems else 'ok',
            'trend': drift, 'between_within': mixing,
            'note': '; '.join(problems) if problems else
                    'no systematic shift between halves and chains overlap',
        })
    return notes


# ── imputation Excel ──────────────────────────────────────────────────────────

def write_imputation_excel(config, result, variables, m, seed, maxit, conv_notes,
                           n_total, aux_cols):
    analysis_type = config['analysis_type']
    effect_label = 'HR' if analysis_type == 'cox' else 'OR'
    title = config.get('title') or f'Multiple Imputation — pooled {effect_label}'
    cc = result.get('cc') or {}
    subtitle = (config.get('subtitle') or '')
    meta = (f"m = {m}, seed = {seed}, maxit = {maxit}  |  "
            f"MI N = {result.get('n_total', n_total)}, complete-case N = {cc.get('n', 'N/A')}")
    subtitle = f"{subtitle}  |  {meta}" if subtitle else meta

    wb = Workbook()

    # ── Pooled ─────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Pooled'
    headers = ['Variable', 'Level', 'Term', f'Pooled {effect_label} (95% CI)',
               'log(estimate)', 'SE', 'p-value', 'df', 'FMI', 'RIV']
    r = sheet_header(ws, title, subtitle, headers, [28, 14, 24, 24, 14, 11, 12, 12, 10, 10])
    pooled = result.get('pooled') or []
    for i, p in enumerate(pooled):
        label, level = match_term(p['term'], variables)
        est = math.exp(p['estimate']) if p.get('estimate') is not None else None
        lo = math.exp(p['conf_low']) if p.get('conf_low') is not None else None
        hi = math.exp(p['conf_high']) if p.get('conf_high') is not None else None
        pv = p.get('p_value')
        fmi = p.get('fmi')
        flag = fmi is not None and fmi > 0.5
        r = write_row(ws, r, [label, level or '-', str(p['term']),
                              fmt_effect(est, lo, hi), fmt_num(p['estimate'], 4),
                              fmt_num(p.get('std_error'), 4), fmt_p(pv),
                              fmt_num(p.get('df'), 1), fmt_num(fmi), fmt_num(p.get('riv'))],
                      alt=(i % 2 == 1), flag=flag,
                      sig_cols=(7,) if (pv is not None and pv < 0.05) else ())
    r += 1
    r = write_footnote(ws, r, len(headers),
                       f"Pooled with mice::pool() (Rubin's rules, Barnard-Rubin degrees of freedom). "
                       f"{effect_label} and its 95% CI are on the ratio scale (exponentiated); "
                       f"log(estimate) and SE are the pooled log-scale quantities they came from. "
                       "FMI = fraction of missing information, the amount of information the missing "
                       "data actually cost for that coefficient -- it is not the same quantity as the "
                       "variable's missing percentage. RIV = relative increase in variance. Rows with "
                       "FMI > 0.5 are shaded: the estimate leans heavily on the imputation model, and "
                       "a larger m is worth considering.")

    # ── CC_vs_MI ───────────────────────────────────────────────────────────
    ws = wb.create_sheet('CC_vs_MI')
    headers = ['Variable', 'Term',
               f'Complete-case {effect_label} (95% CI)', 'CC p-value', 'CC N (events)',
               f'MI {effect_label} (95% CI)', 'MI p-value', 'MI N (events)',
               'Direction flip', 'Crosses p = 0.05', f'|Δ {effect_label}| (%)', 'Flag']
    r = sheet_header(ws, f'Complete-case vs multiple imputation — {effect_label}',
                     subtitle, headers,
                     [26, 22, 24, 12, 16, 24, 12, 16, 14, 16, 14, 26])
    cc_terms = {str(t['term']).strip('`'): t for t in (cc.get('terms') or [])}
    cc_n = cc.get('n')
    cc_ev = cc.get('events')
    mi_n = result.get('n_total', n_total)
    mi_ev = result.get('mi_events')
    n_flagged = 0
    for i, p in enumerate(pooled):
        term = str(p['term']).strip('`')
        label, level = match_term(p['term'], variables)
        disp = f"{label} [{level}]" if level else label
        mi_est = math.exp(p['estimate'])
        mi_lo = math.exp(p['conf_low']); mi_hi = math.exp(p['conf_high'])
        mi_p = p.get('p_value')

        c = cc_terms.get(term)
        if c is None:
            r = write_row(ws, r, [disp, term, 'N/A', 'N/A',
                                  f"{cc_n} ({cc_ev})" if cc_n is not None else 'N/A',
                                  fmt_effect(mi_est, mi_lo, mi_hi), fmt_p(mi_p),
                                  f"{mi_n} ({mi_ev})", 'N/A', 'N/A', 'N/A',
                                  'term absent from the complete-case model'],
                          alt=(i % 2 == 1), flag=True)
            n_flagged += 1
            continue
        cc_est = math.exp(c['estimate'])
        cc_lo = math.exp(c['conf_low']); cc_hi = math.exp(c['conf_high'])
        cc_p = c.get('p_value')

        flip = (cc_est - 1.0) * (mi_est - 1.0) < 0
        crosses = (cc_p is not None and mi_p is not None and
                   ((cc_p < 0.05) != (mi_p < 0.05)))
        delta = abs(mi_est - cc_est) / abs(cc_est) * 100.0 if cc_est not in (0, None) else float('nan')
        big = (not math.isnan(delta)) and delta > 20.0

        flags = []
        if flip:
            flags.append('direction flip')
        if crosses:
            flags.append('significance crosses 0.05')
        if big:
            flags.append(f'effect moved {delta:.0f}%')
        if flags:
            n_flagged += 1

        r = write_row(ws, r, [disp, term,
                              fmt_effect(cc_est, cc_lo, cc_hi), fmt_p(cc_p),
                              f"{cc_n} ({cc_ev})",
                              fmt_effect(mi_est, mi_lo, mi_hi), fmt_p(mi_p),
                              f"{mi_n} ({mi_ev})",
                              'YES' if flip else 'no',
                              'YES' if crosses else 'no',
                              fmt_num(delta, 1),
                              '; '.join(flags) if flags else '-'],
                      alt=(i % 2 == 1), flag=bool(flags))
    r += 1
    r = write_footnote(ws, r, len(headers),
                       "Both columns are the SAME model specification; only the handling of missing data "
                       "differs. The complete-case fit uses the rows with no missing value among the model "
                       f"variables (N = {cc_n} of {mi_n}); the MI fit uses all {mi_n} rows across m = {m} "
                       "imputed datasets. Flagged rows are where the two disagree in a way a reviewer will "
                       "ask about: the effect changes direction relative to the null, significance crosses "
                       "0.05, or the point estimate moves by more than 20%. Agreement between the two is "
                       "itself the sensitivity analysis; disagreement is a finding to explain, not to hide. "
                       "Which one is the primary analysis is a pre-specified decision, not an output of this "
                       "sheet. Crossing p = 0.05 means the two analyses differ in whether they reached "
                       "statistical significance -- it does not mean one of them showed 'no difference'.")

    # ── Imputation_model ───────────────────────────────────────────────────
    ws = wb.create_sheet('Imputation_model')
    headers = ['Column', 'Role', 'Type', 'Imputation method', 'Missing n', 'Missing %',
               'Used as predictor']
    r = sheet_header(ws, 'Imputation model specification', subtitle, headers,
                     [26, 22, 14, 20, 12, 12, 18])
    role_text = {
        'analysis': 'analysis covariate',
        'auxiliary': 'auxiliary (imputation only)',
        'outcome': 'outcome (in imputation model)',
        'outcome_event': 'outcome event indicator (in imputation model)',
        'outcome_cumhaz': 'Nelson-Aalen cumulative hazard (in imputation model)',
        'outcome_time_excluded': 'raw survival time (EXCLUDED as predictor)',
    }
    for i, row in enumerate(result.get('imputation_model') or []):
        r = write_row(ws, r, [row['column'], role_text.get(row['role'], row['role']),
                              row['type'], row['method'] or '(none - complete)',
                              int(row['n_missing']), f"{row['pct_missing']:.1f}%",
                              'yes' if row['used_as_predictor'] else 'no'],
                      alt=(i % 2 == 1))
    r += 1
    r = write_row(ws, r, ['Reproducibility', '', '', '', '', '', ''], bold_cols=(1,))
    versions = result.get('versions') or {}
    meta_rows = [
        ('Engine', 'R mice, fully conditional specification (MICE)'),
        ('Number of imputations (m)', str(m)),
        ('m rule', 'max(20, ceil(max missing % among analysis variables)), '
                   'capped at 100 (White 2011)'
                   if config['mi_settings'].get('m') in (None, 0) else 'set explicitly in config'),
        ('Random seed', str(seed)),
        ('Iterations (maxit)', str(maxit)),
        ('R version', versions.get('r', 'N/A')),
        ('mice version', versions.get('mice', 'N/A')),
        ('survival version', versions.get('survival', 'N/A')),
        ('Analysis model', 'Cox proportional hazards' if analysis_type == 'cox'
                           else 'Logistic regression'),
        ('Auxiliary variables', ', '.join(aux_cols) if aux_cols else '(none)'),
    ]
    for i, (k, v) in enumerate(meta_rows):
        r = write_row(ws, r, [k, v, '', '', '', '', ''], alt=(i % 2 == 1))
    r += 1
    logged = result.get('logged_events') or []
    if logged:
        r = write_row(ws, r, ['mice logged events', '', '', '', '', '', ''], bold_cols=(1,))
        for i, le in enumerate(logged):
            r = write_row(ws, r, ['', le, '', '', '', '', ''], alt=(i % 2 == 1))
        r += 1
    r = write_footnote(ws, r, len(headers),
                       "The outcome is deliberately inside the imputation model: leaving it out biases "
                       "covariate-outcome associations toward the null (Moons 2006). For a time-to-event "
                       "outcome the model carries the event indicator plus the Nelson-Aalen cumulative "
                       "hazard, NOT the raw survival time (White & Royston 2009) -- the raw time column is "
                       "kept in the frame only so the Cox model can be fitted and is switched off in the "
                       "predictor matrix. Methods follow the protocol: continuous = pmm, binary = logreg, "
                       "unordered categorical = polyreg, ordered categorical = polr. Complete variables get "
                       "no method. Everything in this sheet is what the manuscript's missing-data paragraph "
                       "needs; the MAR justification is the one item you still have to write yourself.")

    # ── Convergence_notes ──────────────────────────────────────────────────
    ws = wb.create_sheet('Convergence_notes')
    headers = ['Variable', 'Verdict', 'Drift (half-to-half shift / within-chain SD)',
               'Mixing (between/within chain spread)', 'Note']
    r = sheet_header(ws, 'Convergence diagnostics (machine-readable summary)',
                     subtitle, headers, [26, 14, 24, 26, 60])
    if not conv_notes:
        r = write_row(ws, r, ['(no imputed variable produced a trace)', '', '', '', ''])
    for i, n in enumerate(conv_notes):
        r = write_row(ws, r, [n['variable'], n['verdict'],
                              fmt_num(n['trend']), fmt_num(n['between_within']),
                              n['note']],
                      alt=(i % 2 == 1), flag=(n['verdict'] != 'ok'))
    r += 1
    warns = result.get('warnings') or []
    if warns:
        r = write_row(ws, r, ['R warnings', '', '', '', ''], bold_cols=(1,))
        for i, w in enumerate(warns[:30]):
            r = write_row(ws, r, ['', '', '', '', w], alt=(i % 2 == 1))
        r += 1
    r = write_footnote(ws, r, len(headers),
                       f"'Drift' is how far the across-chain mean moved between the first and second half "
                       f"of the iterations, expressed in within-chain SD units; above {DRIFT_THRESHOLD} "
                       f"suggests the sampler had not settled. 'Mixing' is the spread of the chain means "
                       f"across chains divided by the typical spread inside a chain; above "
                       f"{MIXING_THRESHOLD} suggests the chains sat in different places instead of "
                       f"exploring the same distribution. For a well-behaved run both sit well below "
                       f"their thresholds (drift near 0, mixing around 0.3). These are screening "
                       f"heuristics only -- the trace panel in the diagnostics PNG is the thing to "
                       f"actually look at, and the protocol requires that it be looked at, not just "
                       f"produced. For categorical and binary variables the trace is on mice's "
                       f"internal factor-level codes (levels numbered from 1), not on the original "
                       f"0/1 coding, so read the level relative to itself rather than as a "
                       f"prevalence.")

    out = config['output_file']
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    return [w.title for w in wb.worksheets]


def plot_diagnostics(config, result):
    """Convergence traces (chain mean + sd) and observed vs imputed distributions."""
    path = config.get('diagnostics_plot')
    if not path:
        return None
    conv = result.get('convergence') or {}
    conv_vars = conv.get('variables') or []
    dists = result.get('distributions') or {}
    dist_vars = list(dists.keys())

    n_rows = max(len(conv_vars), len(dist_vars), 1)
    fig, axes = plt.subplots(n_rows, 3, figsize=(14, max(3.0, 2.6 * n_rows)),
                             squeeze=False)

    for i in range(n_rows):
        # column 0 + 1: convergence traces
        if i < len(conv_vars):
            v = conv_vars[i]
            for j, (key, name) in enumerate((('chain_mean', 'mean'), ('chain_sd', 'sd'))):
                ax = axes[i][j]
                chains = conv.get(key, {}).get(v) or []
                for ch in chains:
                    arr = np.asarray(ch, dtype=float)
                    ax.plot(np.arange(1, len(arr) + 1), arr, lw=0.9, alpha=0.75)
                ax.set_title(f'{v}: chain {name}', fontsize=9, fontweight='bold')
                ax.set_xlabel('iteration', fontsize=8)
                ax.tick_params(labelsize=7)
        else:
            axes[i][0].axis('off'); axes[i][1].axis('off')

        # column 2: observed vs imputed
        ax = axes[i][2]
        if i < len(dist_vars):
            v = dist_vars[i]
            d = dists[v]
            if d.get('kind') == 'numeric':
                obs = np.asarray(d.get('observed') or [], dtype=float)
                imp = np.asarray(d.get('imputed') or [], dtype=float)
                bins = 25
                if obs.size:
                    ax.hist(obs, bins=bins, density=True, alpha=0.55,
                            color='#2E5FA3', label='observed')
                if imp.size:
                    ax.hist(imp, bins=bins, density=True, alpha=0.55,
                            color='#C0392B', label='imputed')
                ax.legend(fontsize=7)
            else:
                levels = d.get('levels') or []
                obs = np.asarray(d.get('observed_counts') or [], dtype=float)
                imp = np.asarray(d.get('imputed_counts') or [], dtype=float)
                obs_p = obs / obs.sum() * 100 if obs.sum() else obs
                imp_p = imp / imp.sum() * 100 if imp.sum() else imp
                x = np.arange(len(levels))
                ax.bar(x - 0.2, obs_p, width=0.4, color='#2E5FA3', label='observed')
                ax.bar(x + 0.2, imp_p, width=0.4, color='#C0392B', label='imputed')
                ax.set_xticks(x)
                ax.set_xticklabels([str(l) for l in levels], fontsize=7, rotation=30, ha='right')
                ax.set_ylabel('%', fontsize=8)
                ax.legend(fontsize=7)
            ax.set_title(f'{v}: observed vs imputed', fontsize=9, fontweight='bold')
            ax.tick_params(labelsize=7)
        else:
            ax.axis('off')

    fig.suptitle(f"MICE diagnostics — m = {result.get('m')}, maxit = {result.get('maxit')}, "
                 f"seed = {result.get('seed')}", fontsize=11, fontweight='bold')
    fig.tight_layout(rect=[0, 0, 1, 0.97])
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    fig.savefig(path, dpi=300, bbox_inches='tight')
    plt.close(fig)
    return path


# ═══════════════════════════════════════════════════════════════════════════════
# Preview (--check-missing): never runs MICE
# ═══════════════════════════════════════════════════════════════════════════════

def preview(config, df, variables, aux_cols, out_cols, out_desc, groups):
    mode = config['mode']
    print("=" * 96)
    print(f"MISSING DATA PREVIEW (--check-missing)  |  mode = {mode}")
    print("=" * 96)
    print(f"Data file : {config['data_file']}")
    print(f"Rows after filters/derived variables : {len(df)}")
    if groups:
        for lab, val, sub in groups:
            print(f"  {lab}: n = {len(sub)}")
    print("-" * 96)

    pct_map = {}
    header = f"{'Variable':<28} {'Column':<20} {'Type':<12} {'Missing':>14}"
    if mode == 'imputation':
        header += f"  {'Method':<10}"
    print(header)
    print("-" * 96)

    rows = list(variables)
    time_cols = {od['time_column'] for od in out_desc if od['type'] == 'time_to_event'}
    for c in out_cols:
        rows.append({'column': c,
                     'type': 'continuous' if c in time_cols else 'binary',
                     'label': f'[outcome] {c}'})
    for c in aux_cols:
        rows.append({'column': c, 'type': 'continuous', 'label': f'[auxiliary] {c}'})

    for v in rows:
        col = v['column']
        if col not in df.columns:
            print(f"{v['label']:<28} {col:<20} {v.get('type',''):<12} {'COLUMN NOT FOUND':>14}")
            continue
        s = normalize_series(df, col, v.get('type', 'continuous'))
        nm = int(s.isna().sum())
        pct = 100.0 * nm / len(df) if len(df) else 0.0
        if not str(v['label']).startswith('['):
            pct_map[col] = pct
        line = (f"{v['label']:<28} {col:<20} {v.get('type',''):<12} "
                f"{nm:>5}/{len(df):<5} ({pct:>4.1f}%)")
        if mode == 'imputation':
            if nm == 0:
                meth = '(none)'
            else:
                nl = s.nunique(dropna=True)
                meth = method_for(v.get('type', 'continuous'), nl)
                ov = (config['mi_settings'].get('method_overrides') or {}).get(col)
                if ov:
                    meth = f"{ov} (override)"
            line += f"  {meth:<10}"
        print(line)

    print("-" * 96)
    if mode == 'imputation':
        aux_pct = {}
        for c in aux_cols:
            if c in df.columns:
                aux_pct[c] = 100.0 * df[c].isna().mean()
        # Analysis variables only — see the note beside the same calculation in
        # the imputation run. An auxiliary's own missingness does not enter any
        # reported estimate, so it must not drive m.
        max_pct = max(list(pct_map.values()) + [0.0])
        max_aux_pct = max(list(aux_pct.values()) + [0.0])
        cfg_m = config['mi_settings'].get('m')
        m = int(cfg_m) if cfg_m else compute_m(max_pct)
        print(f"Maximum missing % across the analysis variables : {max_pct:.1f}%")
        if aux_pct:
            print(f"Maximum missing % across the auxiliary variables : "
                  f"{max_aux_pct:.1f}%  (does not affect m)")
        print(f"Planned m : {m}"
              + ("  (from config)" if cfg_m else "  (auto: max(20, ceil(max analysis-variable missing %)), cap 100)"))
        print(f"Planned seed : {config['mi_settings'].get('seed')}   "
              f"maxit : {config['mi_settings'].get('maxit')}")
        ok, rscript, msg = check_r_available()
        print(f"R engine : {'AVAILABLE' if ok else 'NOT AVAILABLE'} "
              f"({rscript or 'no Rscript'}) -- {msg}")
        if not ok:
            qc_blocking(f"R/mice unavailable, multiple imputation cannot run: {msg}")
    else:
        th = config['thresholds']
        over_low = [c for c, p in pct_map.items() if p >= th['low']]
        over_high = [c for c, p in pct_map.items() if p > th['high']]
        print(f"Thresholds : low = {th['low']}%, high = {th['high']}%")
        print(f"At/above low threshold  : {', '.join(over_low) if over_low else '(none)'}")
        print(f"Above high threshold    : {', '.join(over_high) if over_high else '(none)'}")

    print("=" * 96)
    print("[preview] MICE was NOT executed (--check-missing). No Excel or PNG was written.")
    print("=" * 96)


# ═══════════════════════════════════════════════════════════════════════════════
# Mode drivers
# ═══════════════════════════════════════════════════════════════════════════════

def run_assessment(config, df, groups):
    variables = assessment_variables(config)
    out_cols, out_desc = outcome_columns(config)
    group_col, _, _ = grouping_info(config)

    time_cols = {od['time_column'] for od in out_desc if od['type'] == 'time_to_event'}
    extra = [{'column': c,
              'type': 'continuous' if c in time_cols else 'binary',
              'label': f'[outcome] {c}', 'role': 'outcome'}
             for c in out_cols]
    if group_col:
        extra.append({'column': group_col, 'type': 'binary',
                      'label': f'[group] {group_col}', 'role': 'grouping'})
    for v in variables:
        v.setdefault('role', 'analysis')

    print("\n=== Missing data summary ===")
    summary_rows = compute_missing_summary(df, variables, groups, extra)
    for row in summary_rows:
        if row['n_missing']:
            print(f"  {row['label']:<30} {row['n_missing']:>5}/{row['n']:<5} ({row['pct']:.1f}%)")
    n_incomplete = sum(1 for r in summary_rows if r['n_missing'])
    print(f"  Variables with at least one missing value: {n_incomplete}/{len(summary_rows)}")

    print("\n=== Missingness pattern ===")
    patterns = compute_patterns(df, variables)
    print(f"  Complete cases: {patterns['n_complete']}/{len(df)} "
          f"({patterns['pct_complete']:.1f}%), distinct patterns: {patterns.get('n_patterns', 0)}, "
          f"monotone: {patterns['monotone']}")

    print("\n=== Listwise deletion impact ===")
    listwise_rows, listwise_cols = compute_listwise_impact(
        df, variables, out_cols, group_col, groups, out_desc)
    for row in listwise_rows:
        print(f"  {row['population']:<20} {row['n_before']:>6} -> {row['n_after']:<6} "
              f"(dropped {row['n_dropped']}, {row['pct_dropped']:.1f}%)")

    print("\n=== Mechanism tests ===")
    mech_results = compute_mechanism_tests(df, variables, group_col, out_desc)
    n_sig = sum(1 for m in mech_results if m['result'] and m['result']['p'] < 0.05)
    print(f"  {len(mech_results)} tests run, {n_sig} significant at p < 0.05")

    print("\n=== Auxiliary candidates ===")
    aux_rows, aux_candidates, aux_auto = compute_auxiliary_candidates(
        df, variables, config, group_col, out_cols)
    n_rec = sum(1 for a in aux_rows if a['recommended'])
    print(f"  {len(aux_rows)} pairs evaluated over {len(aux_candidates)} candidate columns "
          f"({'auto-detected' if aux_auto else 'from config'}), {n_rec} recommended")

    verdict = compute_verdict(summary_rows, mech_results, out_cols, config['thresholds'])
    print(f"\n=== Verdict: {verdict['verdict']} ===")
    for d in verdict['drivers']:
        print(f"  - {d}")

    print("\n=== Writing outputs ===")
    sheets = write_assessment_excel(config, summary_rows, patterns, listwise_rows, listwise_cols,
                                    mech_results, aux_rows, aux_candidates, aux_auto, verdict,
                                    groups, len(df))
    print(f"  Excel : {config['output_file']}")
    print(f"  Sheets: {', '.join(sheets)}")
    plot_path = plot_assessment(config, patterns, summary_rows)
    if plot_path:
        print(f"  Plot  : {plot_path}")

    # ── QC ─────────────────────────────────────────────────────────────────
    high = config['thresholds']['high']
    for row in summary_rows:
        if row['pct'] is not None and row['pct'] > high and row['role'] == 'analysis':
            qc_warning(f"Variable '{row['column']}' is {row['pct']:.1f}% missing "
                       f"(above the {high}% high threshold): imputed values will lean heavily on "
                       f"the imputation model; consider dropping it or running a with/without "
                       f"sensitivity analysis")
        if row['note'] == 'COLUMN NOT FOUND':
            qc_warning(f"Column '{row['column']}' declared in the config is not present in the data")

    if verdict['verdict'] == 'outcome_missing':
        qc_blocking("Outcome variable itself has missing values: MI-then-delete (von Hippel 2007) "
                    "is required and this implementation does not do it automatically. The "
                    "assessment output was still written; do not run the imputation step until "
                    "the outcome missingness is handled.")
    if verdict['verdict'] == 'mcar_refuted':
        qc_warning("Missingness is significantly associated with the observed outcome, which "
                   "refutes MCAR. This does NOT establish MNAR -- dependence on observed data "
                   "is the MAR case, and neither MAR nor MNAR is testable from the data. "
                   "Complete-case is expected to be biased; MI under MAR is the appropriate "
                   "response, with the outcome in the imputation model and auxiliary variables "
                   "supporting the MAR assumption. An MNAR sensitivity analysis "
                   "(delta-adjustment / tipping point) is warranted only when external "
                   "information suggests missingness depends on the unobserved values; this "
                   "pipeline does not generate one.")

    group_col_, _, _ = grouping_info(config)
    cl = detect_cluster(config, df, [v['column'] for v in variables])
    if cl:
        qc_warning(f"Possible clustering ({cl}): multiple imputation here would be single-level "
                   f"and would understate the variance; multilevel imputation (2l.pmm) is the "
                   f"correct method and the clustering must be disclosed in the limitations. This "
                   f"is raised as a warning, not a blocker, because the detection is heuristic -- "
                   f"once clustering is confirmed the protocol treats it as blocking, so confirm "
                   f"or dismiss it before running the imputation step.")
    ps = detect_psm(config, df)
    if ps:
        qc_warning(f"Propensity-score matching present ({ps}) but MI is NOT applied to the PSM "
                   f"step: the matched analysis stays complete-case, so its results are not "
                   f"directly comparable with the MI results")
    return 0


def run_imputation(config, df, groups, workdir):
    variables = imputation_variables(config)
    out_cols, out_desc = outcome_columns(config)
    group_col, _, _ = grouping_info(config)

    # Grouping column is the exposure main effect; keep it in the model like
    # multivariate_interaction.run_multivariate does.
    if group_col and group_col not in [v['column'] for v in variables]:
        variables = [{'column': group_col, 'type': 'binary',
                      'label': config['grouping'].get('label', group_col),
                      'scale': None, 'reference_level': None,
                      'section': 'Exposure'}] + variables

    missing_cols = [v['column'] for v in variables if v['column'] not in df.columns]
    if missing_cols:
        qc_blocking(f"Model columns not found in the data: {', '.join(missing_cols)}")
        return 3

    aux_cols = [c for c in (config.get('auxiliary_variables') or [])
                if c not in [v['column'] for v in variables]]
    missing_aux = [c for c in aux_cols if c not in df.columns]
    for c in missing_aux:
        qc_warning(f"Auxiliary variable '{c}' is not in the data and was skipped")
    aux_cols = [c for c in aux_cols if c in df.columns]

    # ── blocking check 0: clustering confirmed by Step 3.5 diagnostics ─────
    # Severity depends on where the finding came from. A column-name guess is a
    # warning (see the post-run check); a clustering flag that the survival
    # diagnostics actually raised is blocking, because then the variance
    # understatement of single-level imputation is known, not suspected.
    # Checked here, before MICE runs, so no misleading output file is produced.
    if config.get('cluster_confirmed'):
        cl_desc = detect_cluster(config, df, [v['column'] for v in variables]) or \
            "flagged by research_plan.json$.diagnostics"
        qc_blocking(
            f"Step 3.5 diagnostics confirmed clustering ({cl_desc}), but this imputation is "
            "single-level. Ignoring clusters understates the variance, so the pooled CIs would "
            "be too narrow. Use multilevel imputation (2l.pmm) or drop this MI step — do not "
            "report single-level pooled estimates as if the clustering were absent "
            "(protocol section 6).")
        return 3

    # ── blocking check 1: competing risk ───────────────────────────────────
    cr = detect_competing_risk(config, df)
    if cr:
        qc_blocking("Multiple imputation was requested for a competing-risk analysis "
                    f"({cr}). Each event type needs its own Nelson-Aalen cumulative hazard in "
                    "the imputation model; a single na_cumhaz is not enough (protocol section 6). "
                    "Use a cause-specific single endpoint, or handle the imputation manually.")
        return 3

    # ── blocking check 2: outcome itself missing ───────────────────────────
    outcome_missing = []
    for c in out_cols:
        if c not in df.columns:
            qc_blocking(f"Outcome column '{c}' not found in the data")
            return 3
        nm = int(pd.to_numeric(df[c], errors='coerce').isna().sum())
        if nm > 0:
            outcome_missing.append((c, nm, 100.0 * nm / len(df)))
    if outcome_missing:
        detail = ", ".join(f"{c}: {n} ({p:.1f}%)" for c, n, p in outcome_missing)
        qc_blocking("Outcome variable itself has missing values "
                    f"({detail}). Imputing the outcome and then analysing it as observed inflates "
                    "precision; the correct approach is MI-then-delete (von Hippel 2007) -- the "
                    "outcome stays in the imputation model but cases with a missing outcome are "
                    "dropped before analysis. This implementation does not do that automatically.")
        return 3

    # ── blocking check 3: R / mice availability ────────────────────────────
    r_ok, rscript, r_msg = check_r_available()
    if not r_ok:
        qc_blocking(f"R engine unavailable for multiple imputation: {r_msg}. "
                    "Complete-case is NOT an acceptable silent fallback (protocol section 4.1) -- "
                    "install R with the mice, survival and jsonlite packages and rerun.")
        return 3

    # ── m ──────────────────────────────────────────────────────────────────
    # m follows the missingness of the *analysis* variables (White 2011 ties m to
    # the fraction of missing information in the quantities being estimated).
    # Auxiliaries are in the imputation model to borrow strength, and their own
    # missingness does not enter any reported estimate — including them, as this
    # did until 2026-08-11, let one badly-observed auxiliary inflate m and the
    # runtime with it. Their maximum is still printed, because it is worth
    # knowing when an auxiliary is too sparse to be helping.
    pcts = []
    for v in variables:
        s = normalize_series(df, v['column'], v.get('type', 'continuous'))
        pcts.append(100.0 * s.isna().mean())
    aux_pcts = [100.0 * df[c].isna().mean() for c in aux_cols]
    max_pct = max(pcts + [0.0])
    max_aux_pct = max(aux_pcts + [0.0])
    cfg_m = config['mi_settings'].get('m')
    m = int(cfg_m) if cfg_m else compute_m(max_pct)
    # The protocol's cap of 100 bounds the *rule*, not the analyst. An explicit m
    # is a deliberate override and is honoured as given — but it is worth saying
    # out loud, because the cap exists to keep the run finite and nothing else
    # will mention that a hand-set m is above it.
    if cfg_m and int(cfg_m) > 100:
        qc_warning(f"mi_settings.m is set to {int(cfg_m)}, above the protocol's cap of 100 "
                   f"for the automatic rule. Honoured as an explicit choice; note that "
                   f"runtime scales with m, and record the reason in decision_log.md")
    seed = int(config['mi_settings'].get('seed', 42))
    maxit = int(config['mi_settings'].get('maxit', 20))

    print(f"\n=== Multiple imputation ===")
    print(f"  Model variables : {len(variables)}  "
          f"({', '.join(v['column'] for v in variables)})")
    print(f"  Auxiliary       : {', '.join(aux_cols) if aux_cols else '(none)'}")
    print(f"  Max missing %   : {max_pct:.1f}%  (analysis variables; m is set from this)")
    if aux_cols:
        print(f"  Max missing % (auxiliary) : {max_aux_pct:.1f}%  (does not affect m)")
    print(f"  m = {m} ({'config' if cfg_m else 'auto: max(20, ceil(max analysis-variable missing %)), cap 100'}), "
          f"seed = {seed}, maxit = {maxit}")
    print(f"  Engine          : {rscript}")

    if max_pct == 0:
        qc_blocking("No variable in the imputation model has any missing value: multiple "
                    "imputation is not applicable. Report the complete-case model instead.")
        return 3

    result = run_mice(config, df, variables, aux_cols, out_cols, out_desc,
                      m, seed, maxit, rscript, workdir)
    if not result.get('ok'):
        err = result.get('error') or 'unknown error'
        if result.get('stderr'):
            err = f"{err} | R stderr: {result['stderr'][:800]}"
        qc_blocking(f"mice_driver.R failed: {err}")
        return 3

    pooled = result.get('pooled') or []
    cc = result.get('cc') or {}
    print(f"  MICE completed: {len(pooled)} pooled coefficients, "
          f"MI N = {result.get('n_total')}, complete-case N = {cc.get('n')}")

    conv_notes = assess_convergence(result.get('convergence') or {})
    n_inspect = sum(1 for n in conv_notes if n['verdict'] != 'ok')
    print(f"  Convergence: {len(conv_notes)} imputed variables, {n_inspect} flagged for inspection")

    print("\n=== Writing outputs ===")
    sheets = write_imputation_excel(config, result, variables, m, seed, maxit,
                                    conv_notes, len(df), aux_cols)
    print(f"  Excel : {config['output_file']}")
    print(f"  Sheets: {', '.join(sheets)}")
    plot_path = plot_diagnostics(config, result)
    if plot_path:
        print(f"  Plot  : {plot_path}")

    # ── QC ─────────────────────────────────────────────────────────────────
    high = config['thresholds']['high']
    for row in result.get('imputation_model') or []:
        if row['pct_missing'] > high and row['role'] in ('analysis', 'auxiliary'):
            qc_warning(f"Variable '{row['column']}' is {row['pct_missing']:.1f}% missing "
                       f"(above the {high}% high threshold): its imputed values depend heavily on "
                       f"the imputation model; a with/without sensitivity analysis is expected")
    for p in pooled:
        fmi = p.get('fmi')
        if fmi is not None and fmi > 0.5:
            qc_warning(f"Coefficient '{p['term']}' has FMI = {fmi:.2f} (> 0.5): more than half the "
                       f"information for this estimate comes from the imputation model; consider "
                       f"increasing m and report the FMI in the manuscript")
    n_unassessed = sum(1 for n in conv_notes if n['verdict'] == 'not assessed')
    if n_unassessed:
        qc_warning(f"Convergence could not be screened for {n_unassessed} variable(s) because "
                   f"maxit = {maxit} is too short for the drift/mixing statistics (8 or more "
                   f"iterations are needed; the protocol's working default is 20). The trace panel "
                   f"in the diagnostics PNG still has to be read by eye.")
    for n in conv_notes:
        if n['verdict'] == 'inspect':
            qc_warning(f"Convergence for '{n['variable']}' needs inspection: {n['note']}")

    cl = detect_cluster(config, df, [v['column'] for v in variables])
    if cl:
        # Severity depends on the *source* of the finding, which is what the
        # protocol actually conditions on. A name-based guess is a warning; a
        # clustering flag that Step 3.5's survival diagnostics genuinely raised
        # (passed through as `cluster_confirmed`) is blocking, because then the
        # single-level variance understatement is known rather than suspected.
        # `cluster_confirmed` already blocked before MICE ran, so anything
        # reaching here is the column-name heuristic: warn, do not block.
        qc_warning(f"Possible clustering ({cl}) but this imputation is single-level: ignoring "
                   f"clusters understates the variance. Multilevel imputation (2l.pmm) is the "
                   f"correct method. Raised as a warning rather than a blocker because the "
                   f"detection is a column-name heuristic. If the clustering is real, the "
                   f"protocol treats it as blocking — these results should not be reported "
                   f"without either multilevel imputation or an explicit limitation.")
    ps = detect_psm(config, df)
    if ps:
        qc_warning(f"Propensity-score matching present ({ps}) but MI was NOT applied to the PSM "
                   f"step: matching stays complete-case, so the matched results and these MI "
                   f"results are not directly comparable and must not be discussed as one analysis")

    for w in result.get('warnings') or []:
        qc_warning(f"R: {w}")

    return 0


# ═══════════════════════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════════════════════

def _apply_academic_style():
    """academic-style integration (Excel-only). Consumes --style=<journal> from
    sys.argv. No-op when the skill is not installed or the flag is absent."""
    from pathlib import Path as _P
    _as = _P(__file__).resolve().parents[2] / "academic-style"
    if not _as.exists():
        return None
    if str(_as) not in sys.path:
        sys.path.insert(0, str(_as))
    try:
        from core.biostats_integration import apply_style_from_argv
    except ImportError:
        return None
    try:
        return apply_style_from_argv(sys.argv)
    except Exception:
        return None


def main():
    _AS_STYLE = _apply_academic_style()
    if _AS_STYLE:
        print(f"  Applied {_AS_STYLE} journal style (Excel formatting hints)")

    if len(sys.argv) < 2:
        print("Usage: python missing_data.py config.json [--check-missing] [--style=<journal>]")
        sys.exit(1)

    config_path = sys.argv[1]
    check_only = '--check-missing' in sys.argv

    config = load_config(config_path)
    df = load_data(config)
    groups = split_groups(df, config)

    print(f"Mode: {config['mode']}")
    print(f"Data loaded: n={len(df)}")
    if groups:
        for lab, val, sub in groups:
            print(f"  {lab}: n={len(sub)}")

    if config['mode'] == 'assessment':
        variables = assessment_variables(config)
        aux = [c for c in (config.get('candidate_auxiliary') or []) if c in df.columns]
    else:
        variables = imputation_variables(config)
        group_col, _, _ = grouping_info(config)
        if group_col and group_col not in [v['column'] for v in variables]:
            variables = [{'column': group_col, 'type': 'binary',
                          'label': group_col, 'scale': None,
                          'reference_level': None}] + variables
        aux = [c for c in (config.get('auxiliary_variables') or []) if c in df.columns]
    out_cols, out_desc = outcome_columns(config)

    if check_only:
        preview(config, df, variables, aux, out_cols, out_desc, groups)
        sys.exit(0)

    workdir = tempfile.mkdtemp(prefix='missing_data_')
    try:
        if config['mode'] == 'assessment':
            rc = run_assessment(config, df, groups)
        else:
            rc = run_imputation(config, df, groups, workdir)
    finally:
        shutil.rmtree(workdir, ignore_errors=True)

    print("\nDone!")
    if _QC_MESSAGES:
        n_block = sum(1 for m in _QC_MESSAGES if m.startswith('[QC_BLOCKING]'))
        n_warn = len(_QC_MESSAGES) - n_block
        print(f"  QC: {n_block} blocking, {n_warn} warning(s) written to stderr")
    sys.exit(rc)


if __name__ == '__main__':
    main()
