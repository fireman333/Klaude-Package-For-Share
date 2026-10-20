#!/usr/bin/env python3
"""
Stratified Cox — the executable half of `ph_violation_remedy = "stratify"`
(research-plan pipeline; 2026-08-12 adjudication dispute 5, remedy generators)

Why this exists
---------------
`multivariate-interaction` detects a proportional-hazards violation and then
downgrades it from blocking to a warning if the plan pre-declared a remedy. For
`stratify` that downgrade used to buy nothing: no step generator shipped, so the
"remedy" was a word in a JSON file. The artifact gate could only ask whether
*some* file existed, which any step named `stratified_*` satisfied by writing
anything at all.

This script is the analysis that word was promising. It refits the adjusted model
with the offending covariate(s) moved from the design matrix into `strata=`, so
each stratum gets its own baseline hazard and the remaining coefficients — the
exposure among them — are estimated within strata, where proportionality is not
being asked of the offender any more.

It writes a THIRD table. It does not touch the multivariate step's Sheet 1 or
Sheet 2
-------------------------------------------------------------------------------
The obvious implementation — pass `strata=` inside `multivariate_interaction.py`
— was considered and rejected in the adjudication, for a reason that survives
here: stratifying inside that script would leave Sheet 1 without an HR for the
stratified variable while Sheet 2's interaction models remained unstratified, so
one workbook would carry two tables fitted under different assumptions with
nothing in the layout saying so. Two tables that disagree silently is worse than
two tables that disagree loudly.

So the main analysis is left exactly as it was — complete, internally consistent,
and openly unremediated — and the remedy is a separate artifact with its own
cross-reference. The cost is honest and stated in the footnote: the reader has to
look at a second workbook. The alternative cost was a workbook that looks
coherent and is not.

Evidence sidecar
----------------
Alongside the workbook this writes `<stem>_remedy.json` (schema
`ph_remedy_evidence/1`): what was stratified on, why those variables, how many
strata and events, which coefficients survived, and a re-test of the PH
assumption on the stratified fit. `biostats-guardian/scripts/ph_remedy_gate.py`
reads it and cross-checks it against the `*_ph_test.json` that triggered the
remedy in the first place, so "declared stratify" now has to mean "stratified on
the variable that actually violated".

The sidecar is bound to the workbook by name and sha256, the same binding the PSM
contrast sidecar uses. Same reason: an unbound sidecar might describe a different
run, and a sidecar that might describe something else is not evidence.

Usage:
    python ph_stratified.py config.json                  # full run
    python ph_stratified.py config.json --check-missing  # preview only, no models

Config (mode = "ph_stratified"):
    {"mode":"ph_stratified","remedy":"stratify","outcome_id":"primary",
     "data_file":"...","filters":[],"derived_variables":[],
     "grouping":{"column":"AF","groups":[{"value":0,"label":"No AF"},
                                          {"value":1,"label":"AF"}]},
     "analysis_type":"cox",
     "cox_settings":{"time_column":"t","event_column":"e"},
     "logistic_settings":null,
     "sections":[{"header":"...","variables":[...]}],
     "strata_variables":"auto",          # or an explicit ["ckd_stage"]
     "ph_test_file":".../Multi_primary_ph_test.json",
     "source_table":".../Multi_primary.xlsx",
     "title":"...","subtitle":"...","abbreviations":"...",
     "output_file":".../Stratified_primary.xlsx",
     "evidence_file":".../Stratified_primary_remedy.json"}

Design notes
------------
* `strata_variables: "auto"` reads the violating variables out of the PH test
  artifact. Which covariate violated is a property of the data, not a decision —
  the pre-specified decision is *"if PH fails, we stratify"*, and stratifying on
  something other than the offender would not be a remedy. An explicit list in
  the plan overrides, and the sidecar records which of the two happened.
* Stratifying on the exposure is refused (exit 2), not silently done. `strata=`
  removes the variable's coefficient; if that variable is the exposure there is
  no estimate left and the "remedy" would have deleted the result. `time_varying`
  or `rmst` are the remedies for a non-proportional exposure.
* Continuous violators are refused, not auto-binned. Binning is an analytic
  decision with a defensible range of answers; making it silently, inside a
  remedy, would put an unrecorded choice underneath a published estimate.
* Every failure to compute is `status: "not_computed"` in the sidecar with a
  reason, never a missing file and never a silent pass — "could not" and "did
  not need to" are different states and the gate treats them differently.
"""

import hashlib
import json
import operator
import os
import sys
from datetime import datetime, timezone

import numpy as np
import pandas as pd

from lifelines import CoxPHFitter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EVIDENCE_SCHEMA = "ph_remedy_evidence/1"
REMEDY = "stratify"

# Same gate alpha as multivariate_interaction.PH_ALPHA. Duplicated as a literal
# rather than imported for the same reason generate_configs duplicates
# PH_REMEDIES: this script must not drag that module's import graph in, and the
# value is a published threshold, not an implementation detail.
PH_ALPHA = 0.01

# A stratum with almost no events contributes almost nothing but still costs a
# baseline hazard. Not an error — a warning, because the right answer (collapse
# levels, or choose a different remedy) is the author's call.
MIN_EVENTS_PER_STRATUM = 10
# Above this, "stratified" has become "matched on", and the within-stratum
# comparison is running out of data. Also a warning, for the same reason.
MAX_STRATA_LEVELS = 10
# Distinct values above which a variable is treated as continuous and refused.
CONTINUOUS_LEVEL_THRESHOLD = 20

_QC_MESSAGES = []


def qc(kind, message):
    line = f"[QC_{kind}] {message}"
    _QC_MESSAGES.append(line)
    print(line, file=sys.stderr)


def qc_warning(message):
    qc('WARNING', message)


def qc_blocking(message):
    qc('BLOCKING', message)


# ═══════════════════════════════════════════════════════════════════════════════
# Configuration
# ═══════════════════════════════════════════════════════════════════════════════

def load_config(path):
    with open(path, 'r', encoding='utf-8') as f:
        config = json.load(f)

    config.setdefault('mode', 'ph_stratified')
    config.setdefault('remedy', REMEDY)
    config.setdefault('filters', [])
    config.setdefault('derived_variables', [])
    config.setdefault('sections', [])
    config.setdefault('title', '')
    config.setdefault('subtitle', '')
    config.setdefault('abbreviations', '')
    config.setdefault('analysis_type', 'cox')
    config.setdefault('cox_settings', None)
    config.setdefault('logistic_settings', None)
    config.setdefault('strata_variables', 'auto')
    config.setdefault('ph_test_file', None)
    config.setdefault('source_table', None)
    config.setdefault('missing_data_handling', 'listwise')

    if config['mode'] != 'ph_stratified':
        sys.exit(f"[FATAL] Unknown mode: {config['mode']!r} (expected 'ph_stratified')")
    if config['remedy'] != REMEDY:
        sys.exit(f"[FATAL] This script implements ph_violation_remedy = {REMEDY!r}, "
                 f"but the config says {config['remedy']!r}.")
    if config['analysis_type'] != 'cox':
        # Stratification of a baseline hazard is a survival-model idea. A
        # logistic step has no proportional-hazards assumption to violate, so a
        # config that asks for this is wired wrong rather than merely unusual.
        sys.exit(f"[FATAL] analysis_type={config['analysis_type']!r}: a stratified "
                 f"baseline hazard only exists for a Cox model. A logistic step has "
                 f"no proportional-hazards assumption and needs no remedy.")
    if not config.get('output_file'):
        sys.exit("[FATAL] config has no output_file; the gate matches the evidence "
                 "sidecar to the workbook by name, so the workbook must be named.")
    config.setdefault('evidence_file',
                      f"{os.path.splitext(config['output_file'])[0]}_remedy.json")
    return config


# ═══════════════════════════════════════════════════════════════════════════════
# Data loading (identical semantics to robust_variance.py / univariate_forest.py)
# ═══════════════════════════════════════════════════════════════════════════════

OPS = {
    '==': operator.eq, '!=': operator.ne,
    '>': operator.gt, '>=': operator.ge,
    '<': operator.lt, '<=': operator.le,
}


def load_data(config):
    path = config['data_file']
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xls'):
        df = pd.read_excel(path)
    elif ext == '.tsv':
        df = pd.read_csv(path, sep='\t')
    elif ext == '.sas7bdat':
        df = pd.read_sas(path)
    elif ext == '.dta':
        df = pd.read_stata(path)
    else:
        df = pd.read_csv(path)

    for f in config['filters']:
        col, op, val = f['column'], f['operator'], f['value']
        if op in OPS:
            df = df[OPS[op](pd.to_numeric(df[col], errors='coerce'), val)]
        elif op == 'in':
            df = df[df[col].isin(val)]
        elif op == 'not_in':
            df = df[~df[col].isin(val)]

    for dv in config['derived_variables']:
        if dv['type'] == 'coalesce':
            result = pd.to_numeric(df[dv['columns'][0]], errors='coerce')
            for c in dv['columns'][1:]:
                result = result.combine_first(pd.to_numeric(df[c], errors='coerce'))
            df[dv['name']] = result
        elif dv['type'] == 'expression':
            local_ns = {}
            for dep in dv.get('dependencies', []):
                local_ns[dep] = pd.to_numeric(df[dep], errors='coerce')
            local_ns['np'] = np
            df[dv['name']] = eval(dv['expression'], {"__builtins__": {}}, local_ns)

    return df.reset_index(drop=True)


def normalize_series(df, col, var_type):
    s = df[col]
    if var_type == 'categorical':
        return s.where(s.notna() & (s.astype(str).str.strip() != ''), np.nan)
    return pd.to_numeric(s, errors='coerce')


def collect_variables(config):
    out = []
    for section in config['sections']:
        for v in section['variables']:
            out.append({
                'column': v['column'],
                'type': v.get('type', 'continuous'),
                'label': v.get('label', v['column']),
                'scale': v.get('scale'),
                'reference_level': v.get('reference_level'),
                'section': section.get('header', ''),
            })
    return out


def outcome_columns(config):
    cs = config['cox_settings'] or {}
    return [cs.get('time_column'), cs.get('event_column')]


def grouping_column(config):
    gc = config.get('grouping')
    return gc.get('column') if gc else None


# ═══════════════════════════════════════════════════════════════════════════════
# Which variables get stratified on
# ═══════════════════════════════════════════════════════════════════════════════

def read_ph_test(config):
    """(payload, path, error). The artifact that triggered this remedy."""
    path = config.get('ph_test_file')
    if not path:
        return None, None, "config has no ph_test_file"
    if not os.path.exists(path):
        return None, path, f"not found: {path}"
    try:
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f), path, None
    except Exception as exc:
        return None, path, f"{type(exc).__name__}: {exc}"


def resolve_strata(config, ph_art):
    """(strata_variables, source, error).

    `"auto"` takes the violating variables from the PH test. That is not a
    loosening of pre-specification: the plan pre-specified *stratify*, and which
    covariate broke proportionality is an observation about the data. Stratifying
    on anything other than the offender would leave the violation in the model
    while claiming to have addressed it, so "auto" is the only reading of
    "stratify" that is a remedy at all.

    An explicit list in the plan wins, and the sidecar records which happened, so
    a reader can tell a pre-named stratifier from a data-named one.
    """
    declared = config.get('strata_variables')
    if isinstance(declared, str) and declared.lower() == 'auto':
        if not ph_art:
            return [], 'auto_from_ph_test', (
                "strata_variables='auto' needs the PH test artifact to know which "
                "covariate violated, and it could not be read")
        violators = list(ph_art.get('violating_variables') or [])
        if not violators:
            return [], 'auto_from_ph_test', (
                "strata_variables='auto' but the PH test lists no violating "
                "covariate. A global violation with no single offender cannot be "
                "stratified away — name the variables explicitly, or use a remedy "
                "that does not need one (rmst / time_varying)")
        return violators, 'auto_from_ph_test', None
    if isinstance(declared, list) and declared:
        return list(declared), 'plan', None
    return [], 'plan', (f"strata_variables={declared!r} is neither 'auto' nor a "
                        f"non-empty list")


def validate_strata(df, strata_vars, config):
    """Refuse the two cases where stratifying would destroy the thing it fixes.

    Returns an error string, or None. Both refusals are exit-2 rather than a
    warning: a stratified model that dropped the exposure, or one stratified on a
    silently binned continuous variable, would still write a workbook and would
    still satisfy any gate that only asks whether a file exists.
    """
    grp = grouping_column(config)
    if grp and grp in strata_vars:
        return (f"the violating variable is the exposure ({grp!r}). Stratifying on it "
                f"removes its coefficient, so the remedy would delete the estimate the "
                f"analysis exists to produce. A non-proportional exposure needs "
                f"ph_violation_remedy='time_varying' (an HR that is allowed to change "
                f"with time) or 'rmst' (a summary that does not assume a constant "
                f"ratio) — change the plan, do not stratify.")
    missing = [c for c in strata_vars if c not in df.columns]
    if missing:
        return f"strata variable(s) not in the data: {missing}"
    for c in strata_vars:
        levels = df[c].dropna().unique()
        if len(levels) < 2:
            return (f"strata variable {c!r} has {len(levels)} distinct value(s) in the "
                    f"analysis rows; a single stratum is not a stratification")
        if len(levels) > CONTINUOUS_LEVEL_THRESHOLD:
            return (f"strata variable {c!r} takes {len(levels)} distinct values and is "
                    f"being treated as continuous. Stratifying needs categories, and "
                    f"choosing the cut points is an analytic decision with several "
                    f"defensible answers — this script will not make it silently "
                    f"underneath a published estimate. Add the binned variable to the "
                    f"plan (with its cut points in the decision log) and name it in "
                    f"outcomes[].ph_strata_variables, or use rmst / time_varying.")
    return None


# ═══════════════════════════════════════════════════════════════════════════════
# Model matrix
# ═══════════════════════════════════════════════════════════════════════════════

def _prepare_categorical(d, col, var_type, reference_level, context_label=""):
    """Reference level + dummy encoding, lifted from robust_variance.py so the
    encoding matches the unremediated model this table is compared against."""
    if var_type != 'categorical':
        return d, [col]

    non_na = d[col].dropna()
    unique_vals = sorted(non_na.unique().tolist(), key=lambda x: str(x))
    if len(unique_vals) < 2:
        return d, [col]

    if reference_level is None:
        first_cat = unique_vals[0]
        qc_warning(f"{col}: no reference_level specified, using alphabetical first "
                   f"({first_cat}). This may not match research_plan.")
        ref = first_cat
    else:
        ref = reference_level
        if ref not in unique_vals:
            qc_warning(f"{col}: reference_level '{ref}' not present in data "
                       f"{context_label}; falling back to alphabetical first "
                       f"({unique_vals[0]}).")
            ref = unique_vals[0]

    ordered = [ref] + [c for c in unique_vals if c != ref]
    d = d.copy()
    d[col] = pd.Categorical(d[col], categories=ordered)
    d = pd.get_dummies(d, columns=[col], prefix=col, drop_first=True, dtype=float)
    dummy_cols = [f"{col}_{c}" for c in ordered[1:]]
    dummy_cols = [c for c in dummy_cols if c in d.columns]
    return d, dummy_cols


def build_design(df, config, variables, strata_vars):
    """(design_df, model_cols, strata_cols, label_map, missing_cols).

    The strata variables are carried through the same complete-case filter as
    everything else and then held out of the design matrix — that holding out is
    the whole mechanism, and it is why the stratified variable has no row in the
    output table.
    """
    grp = grouping_column(config)
    out_cols = [c for c in outcome_columns(config) if c]

    model_vars = list(variables)
    if grp and grp not in [v['column'] for v in model_vars]:
        model_vars = [{'column': grp, 'type': 'binary', 'label': grp,
                       'scale': None, 'reference_level': None,
                       'section': 'Exposure'}] + model_vars
    # A variable that is being stratified on must not also be a covariate: it
    # would be collinear with its own baseline hazard.
    model_vars = [v for v in model_vars if v['column'] not in strata_vars]

    needed = ([v['column'] for v in model_vars] + out_cols + list(strata_vars))
    needed = list(dict.fromkeys(needed))
    missing_cols = [c for c in needed if c not in df.columns]
    if missing_cols:
        return None, None, None, None, missing_cols

    d = df[needed].copy()
    for v in model_vars:
        d[v['column']] = normalize_series(d, v['column'], v['type'])
    for c in out_cols:
        d[c] = pd.to_numeric(d[c], errors='coerce')
    for c in strata_vars:
        d[c] = d[c].where(d[c].notna() & (d[c].astype(str).str.strip() != ''), np.nan)

    d = d.dropna(subset=needed).reset_index(drop=True)

    model_cols, label_map = [], {}
    for v in model_vars:
        col, vtype = v['column'], v['type']
        if vtype == 'categorical':
            d, dummies = _prepare_categorical(d, col, vtype, v['reference_level'],
                                              context_label='(stratified Cox model)')
            for dc in dummies:
                label_map[dc] = f"{v['label']}: {dc[len(col) + 1:]}"
            model_cols += dummies
        else:
            if v['scale']:
                d[col] = d[col] / v['scale']
                label_map[col] = f"{v['label']} (per {v['scale']})"
            else:
                label_map[col] = v['label']
            model_cols.append(col)

    keep = []
    for c in model_cols:
        if d[c].nunique(dropna=True) < 2:
            qc_warning(f"dropped covariate '{c}' (zero variance in the complete-case rows)")
        else:
            keep.append(c)
    # lifelines wants the strata columns as plain categories.
    for c in strata_vars:
        d[c] = d[c].astype(str)
    return d, keep, list(strata_vars), label_map, []


# ═══════════════════════════════════════════════════════════════════════════════
# Fitting
# ═══════════════════════════════════════════════════════════════════════════════

def fit_stratified(d, model_cols, strata_cols, config):
    """(results, cph, fit_frame) or (None, None, reason)."""
    cs = config['cox_settings']
    time_col, event_col = cs['time_column'], cs['event_column']
    frame = d[model_cols + strata_cols + [time_col, event_col]].copy()
    n_events = int(frame[event_col].sum())
    if len(frame) < 10 or n_events < 5:
        return None, None, (f"n={len(frame)}, events={n_events}: too few complete cases "
                            f"to fit a stratified model")
    try:
        cph = CoxPHFitter()
        cph.fit(frame, duration_col=time_col, event_col=event_col, strata=strata_cols)
    except Exception as exc:
        return None, None, f"{type(exc).__name__}: {exc}"

    s = cph.summary
    results = {}
    for col in model_cols:
        if col in s.index:
            results[col] = {
                'coef': float(s.loc[col, 'coef']),
                'effect': float(np.exp(s.loc[col, 'coef'])),
                'ci_lower': float(np.exp(s.loc[col, 'coef lower 95%'])),
                'ci_upper': float(np.exp(s.loc[col, 'coef upper 95%'])),
                'p': float(s.loc[col, 'p']),
            }
    return results, cph, frame


def ph_recheck(cph, frame):
    """Did stratifying actually fix it?

    A stratified model can still violate proportional hazards — on a *different*
    covariate, or on the same one if the strata do not capture the way it moves.
    Re-testing is the difference between "we applied a remedy" and "the remedy
    worked", and only the second is worth putting in a paper. Reported as a
    three-state status rather than a bare boolean: "could not re-test" must not
    read as "re-tested and clean".
    """
    try:
        from lifelines.statistics import proportional_hazard_test
        res = proportional_hazard_test(cph, frame, time_transform='rank')
        summ = res.summary
        per_var = {}
        for idx, row in summ.iterrows():
            name = idx[0] if isinstance(idx, tuple) else idx
            per_var[str(name)] = float(row['p'])
        if not per_var:
            return {'status': 'not_computed',
                    'reason': 'the PH re-test returned no per-covariate rows'}
        min_p = min(per_var.values())
        violators = sorted([k for k, v in per_var.items() if v < PH_ALPHA])
        return {
            'status': 'computed',
            'alpha': PH_ALPHA,
            'time_transform': 'rank',
            'min_p': min_p,
            'violated': bool(violators),
            'violating_variables': violators,
            'per_variable_p': per_var,
        }
    except Exception as exc:
        return {'status': 'not_computed', 'reason': f"{type(exc).__name__}: {exc}"}


def strata_description(frame, strata_cols, event_col):
    """Per-stratum n and events, keyed by the joined level label."""
    if not strata_cols:
        return {}, {}
    grouped = frame.groupby(strata_cols, dropna=False)
    sizes, events = {}, {}
    for key, sub in grouped:
        label = " | ".join(str(k) for k in (key if isinstance(key, tuple) else (key,)))
        sizes[label] = int(len(sub))
        events[label] = int(sub[event_col].sum())
    return sizes, events


# ═══════════════════════════════════════════════════════════════════════════════
# Excel
# ═══════════════════════════════════════════════════════════════════════════════

HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
TITLE_FONT = Font(bold=True, size=14)
SUB_FONT = Font(italic=True, size=10, color="555555")
DATA_FONT = Font(size=10)
SIG_FONT = Font(size=10, bold=True)
SECTION_FONT = Font(bold=True, size=10)
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
ALT_FILL = PatternFill("solid", fgColor="F2F2F2")
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def fmt_p(p):
    if p is None:
        return "—"
    return "<0.001" if p < 0.001 else f"{p:.3f}"


def fmt_effect(effect, lo, hi):
    if effect is None:
        return "—"
    return f"{effect:.2f} ({lo:.2f}–{hi:.2f})"


def write_excel(config, results, label_map, model_cols, strata_vars, strata_source,
                sizes, events, n, n_events, recheck, ph_art):
    wb = Workbook()

    # ── Sheet 1: the stratified model ──────────────────────────────────────
    ws = wb.active
    ws.title = "Stratified Cox"
    for c, w in zip(range(1, 5), [34, 14, 26, 12]):
        ws.column_dimensions[get_column_letter(c)].width = w

    r = 1
    ws.cell(row=r, column=1, value=config.get('title') or "Stratified Cox model").font = TITLE_FONT
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1
    if config.get('subtitle'):
        ws.cell(row=r, column=1, value=config['subtitle']).font = SUB_FONT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1
    ws.cell(row=r, column=1,
            value=f"n = {n}, events = {n_events}; stratified by "
                  f"{', '.join(strata_vars)} ({len(sizes)} strata)").font = SUB_FONT
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 2

    for c, h in enumerate(["Variable", "Coefficient",
                           "Adjusted HR (95% CI)", "p-value"], 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, CENTER, THIN
    r += 1

    alt = False
    for col in model_cols:
        fill = ALT_FILL if alt else None
        alt = not alt
        cell = ws.cell(row=r, column=1, value=label_map.get(col, col))
        cell.font, cell.alignment, cell.border = DATA_FONT, LEFT, THIN
        if fill:
            cell.fill = fill
        res = results.get(col)
        if res:
            font = SIG_FONT if res['p'] < 0.05 else DATA_FONT
            vals = [f"{res['coef']:.3f}",
                    fmt_effect(res['effect'], res['ci_lower'], res['ci_upper']),
                    fmt_p(res['p'])]
        else:
            font, vals = DATA_FONT, ["N/A", "N/A", "N/A"]
        for i, v in enumerate(vals, 2):
            cell = ws.cell(row=r, column=i, value=v)
            cell.font, cell.alignment, cell.border = font, CENTER, THIN
            if fill:
                cell.fill = fill
        r += 1

    # The stratified variables, shown as present-but-without-an-estimate. Leaving
    # them off the table entirely is how a reader ends up thinking they were not
    # adjusted for.
    for sv in strata_vars:
        cell = ws.cell(row=r, column=1, value=f"{sv} (stratum)")
        cell.font, cell.alignment, cell.border = SECTION_FONT, LEFT, THIN
        cell.fill = SECTION_FILL
        for i in range(2, 5):
            cell = ws.cell(row=r, column=i, value="— (stratified)")
            cell.font, cell.alignment, cell.border = DATA_FONT, CENTER, THIN
            cell.fill = SECTION_FILL
        r += 1

    r += 1
    footnotes = [
        f"Stratified Cox proportional-hazards model: {', '.join(strata_vars)} "
        f"entered as strata (each stratum has its own baseline hazard) rather than "
        f"as covariates, so no hazard ratio is estimable for them by construction. "
        f"All other covariates are estimated within strata.",
        f"Stratification variables were chosen "
        + ("from the covariates that violated the proportional-hazards assumption in "
           f"the unremediated model (scaled Schoenfeld residuals, alpha = {PH_ALPHA})."
           if strata_source == 'auto_from_ph_test'
           else "explicitly in research_plan.json (outcomes[].ph_strata_variables)."),
    ]
    if config.get('source_table'):
        footnotes.append(
            f"This table is the pre-specified remedy for the proportional-hazards "
            f"violation reported in {os.path.basename(config['source_table'])}. That "
            f"workbook is left unremediated on purpose and its Sheet 1 hazard ratios "
            f"must not be presented as the primary estimate for the affected "
            f"covariates; this table is.")
    if recheck.get('status') == 'computed':
        footnotes.append(
            f"Proportional hazards re-tested on the stratified fit: min p = "
            f"{recheck['min_p']:.4g} (gate alpha = {PH_ALPHA}); "
            + ("no covariate violates after stratification."
               if not recheck['violated'] else
               f"still violated by {', '.join(recheck['violating_variables'])} — "
               f"stratification did not resolve the violation."))
    else:
        footnotes.append(f"Proportional hazards could not be re-tested on the "
                         f"stratified fit ({recheck.get('reason', 'no reason given')}). "
                         f"This is not evidence that the assumption now holds.")
    if config.get('abbreviations'):
        footnotes.append(f"Abbreviations: {config['abbreviations']}")
    for note in footnotes:
        cell = ws.cell(row=r, column=1, value=note)
        cell.font, cell.alignment = SUB_FONT, LEFT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1

    # ── Sheet 2: strata composition + PH re-test ───────────────────────────
    ws2 = wb.create_sheet("Strata & PH recheck")
    for c, w in zip(range(1, 4), [40, 14, 14]):
        ws2.column_dimensions[get_column_letter(c)].width = w
    r = 1
    ws2.cell(row=r, column=1, value="Strata composition").font = TITLE_FONT
    r += 2
    for c, h in enumerate(["Stratum", "n", "Events"], 1):
        cell = ws2.cell(row=r, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, CENTER, THIN
    r += 1
    for label in sorted(sizes):
        ws2.cell(row=r, column=1, value=label).font = DATA_FONT
        ws2.cell(row=r, column=2, value=sizes[label]).font = DATA_FONT
        ev = events.get(label, 0)
        cell = ws2.cell(row=r, column=3, value=ev)
        cell.font = SIG_FONT if ev < MIN_EVENTS_PER_STRATUM else DATA_FONT
        for i in range(1, 4):
            ws2.cell(row=r, column=i).border = THIN
        r += 1

    r += 1
    ws2.cell(row=r, column=1, value="Proportional hazards re-test (stratified fit)").font = TITLE_FONT
    r += 2
    if recheck.get('status') == 'computed':
        for c, h in enumerate(["Covariate", "p"], 1):
            cell = ws2.cell(row=r, column=c, value=h)
            cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, CENTER, THIN
        r += 1
        for name, p in sorted(recheck['per_variable_p'].items()):
            ws2.cell(row=r, column=1, value=name).font = DATA_FONT
            cell = ws2.cell(row=r, column=2, value=fmt_p(p))
            cell.font = SIG_FONT if p < PH_ALPHA else DATA_FONT
            cell.alignment = CENTER
            r += 1
    else:
        ws2.cell(row=r, column=1,
                 value=f"not computed: {recheck.get('reason', '')}").font = SUB_FONT
        r += 1

    if ph_art:
        r += 1
        ws2.cell(row=r, column=1, value="Violation being remedied").font = TITLE_FONT
        r += 1
        g = (ph_art.get('global') or {})
        ws2.cell(row=r, column=1,
                 value=f"Unremediated model: global chi2 = {g.get('chi2')}, "
                       f"p = {g.get('p')}; violating covariates: "
                       f"{', '.join(ph_art.get('violating_variables') or []) or 'none individually'}"
                 ).font = SUB_FONT
        ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)

    out = config['output_file']
    os.makedirs(os.path.dirname(os.path.abspath(out)) or ".", exist_ok=True)
    wb.save(out)
    return [ws.title, ws2.title]


# ═══════════════════════════════════════════════════════════════════════════════
# Evidence sidecar
# ═══════════════════════════════════════════════════════════════════════════════

def file_sha256(path):
    h = hashlib.sha256()
    try:
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 20), b""):
                h.update(chunk)
    except OSError:
        return None
    return h.hexdigest()


def write_evidence(config, config_path, payload):
    """Write `<stem>_remedy.json` atomically, bound to the workbook it describes.

    Written even when the analysis could not run (`status: "not_computed"`), and
    that is the point: an absent sidecar and a sidecar that says "I could not
    compute this" are different claims, and the gate is allowed to treat them
    differently only if the second one exists.
    """
    out = config['output_file']
    payload = dict(payload)
    payload.update({
        'schema': EVIDENCE_SCHEMA,
        'remedy': REMEDY,
        'generated_by': 'research-plan/scripts/ph_stratified.py',
        'generated_at': datetime.now(timezone.utc).isoformat(),
        'outcome_id': config.get('outcome_id'),
        'table_basename': os.path.basename(out),
        'table_sha256': file_sha256(out) if os.path.exists(out) else None,
        'config_basename': os.path.basename(config_path),
        'config_sha256': file_sha256(config_path),
        'source_table': (os.path.basename(config['source_table'])
                         if config.get('source_table') else None),
    })
    path = config['evidence_file']
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    tmp = f"{path}.{os.getpid()}.tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)
        f.write("\n")
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)
    return path


def _not_computed(config, config_path, reason, **extra):
    """Record why, on disk, then exit 2. Never leaves the gate guessing."""
    payload = {'status': 'not_computed', 'reason': reason}
    payload.update(extra)
    path = write_evidence(config, config_path, payload)
    qc_blocking(f"stratified Cox not computed: {reason}")
    print(f"Evidence written: {path} (status = not_computed)")
    return 2


def _not_applicable(config, config_path, reason, **extra):
    """The remedy was pre-declared and turned out not to be needed. Exit 0.

    Three states, not two. The remedy is declared before the Schoenfeld test is
    run — that is the entire point of pre-specification — so a step that finds no
    violation has not failed, and a step that could not fit has not succeeded.
    Collapsing either into the other is how "we could not compute the remedy"
    ends up reading as "the remedy was not required".
    """
    payload = {'status': 'not_applicable', 'reason': reason}
    payload.update(extra)
    path = write_evidence(config, config_path, payload)
    print(f"No proportional-hazards violation to remedy: {reason}")
    print(f"Evidence written: {path} (status = not_applicable)")
    return 0


# ═══════════════════════════════════════════════════════════════════════════════
# Preview / run
# ═══════════════════════════════════════════════════════════════════════════════

def preview(config, df, variables):
    print("STRATIFIED COX PREVIEW (--check-missing)  |  mode = ph_stratified")
    ph_art, ph_path, ph_err = read_ph_test(config)
    if ph_err:
        qc_warning(f"PH test artifact unreadable ({ph_err}). With "
                   f"strata_variables='auto' the run will not be able to choose "
                   f"stratification variables.")
    strata_vars, source, err = resolve_strata(config, ph_art)
    print(f"  strata_variables : {strata_vars or '(unresolved)'}  [{source}]")
    if err:
        qc_warning(err)
    else:
        verr = validate_strata(df, strata_vars, config)
        if verr:
            qc_warning(f"strata would be refused at run time: {verr}")
    cols = [v['column'] for v in variables] + [c for c in outcome_columns(config) if c]
    missing = [c for c in dict.fromkeys(cols + strata_vars) if c not in df.columns]
    print(f"  rows loaded      : {len(df)}")
    print(f"  columns missing  : {missing or 'none'}")
    print("[preview] No model was fitted (--check-missing). No Excel was written.")


def run(config, config_path, df, variables):
    ph_art, ph_path, ph_err = read_ph_test(config)
    if ph_err:
        qc_warning(f"PH test artifact: {ph_err}")

    # The remedy is declared before the test result is known, so this step is
    # routinely scheduled for a model that turns out not to violate anything.
    # That is `not_applicable` and a clean exit — not a pass to be explained away,
    # and not a failure either.
    ph_status = (ph_art or {}).get('status')
    if ph_art and ph_status == 'computed' and not ph_art.get('violated'):
        return _not_applicable(
            config, config_path,
            f"the proportional-hazards assumption was not violated in "
            f"{os.path.basename(ph_path)} (gate alpha = {ph_art.get('alpha', PH_ALPHA)}), "
            f"so there is nothing for the declared 'stratify' remedy to remedy",
            source_ph_test=os.path.basename(ph_path),
            source_ph_test_sha256=file_sha256(ph_path),
            ph_violating_variables=list(ph_art.get('violating_variables') or []),
            ph_borderline=bool(ph_art.get('borderline')))
    if ph_art and ph_status not in (None, 'computed'):
        # NOT_COMPUTED / NOT_APPLICABLE upstream: the violation is unknown, not
        # absent. Stratifying on nothing and reporting a clean table would be the
        # worst of the three outcomes.
        return _not_computed(
            config, config_path,
            f"the upstream PH test is status={ph_status!r} "
            f"({(ph_art.get('reason') or 'no reason given')}), so whether the "
            f"assumption is violated is unknown. An unknown violation is not an "
            f"absent one.",
            source_ph_test=os.path.basename(ph_path) if ph_path else None)

    strata_vars, source, err = resolve_strata(config, ph_art)
    if err:
        return _not_computed(config, config_path, err,
                             strata_variables_source=source,
                             ph_test=os.path.basename(ph_path) if ph_path else None)

    verr = validate_strata(df, strata_vars, config)
    if verr:
        return _not_computed(config, config_path, verr,
                             strata_variables=strata_vars,
                             strata_variables_source=source)

    d, model_cols, strata_cols, label_map, missing = build_design(
        df, config, variables, strata_vars)
    if missing:
        return _not_computed(config, config_path,
                             f"columns not in the data: {missing}",
                             strata_variables=strata_vars,
                             strata_variables_source=source)
    if not model_cols:
        return _not_computed(config, config_path,
                             "no covariates left in the design matrix after removing "
                             "the strata variables and zero-variance columns",
                             strata_variables=strata_vars,
                             strata_variables_source=source)

    results, cph, frame_or_reason = fit_stratified(d, model_cols, strata_cols, config)
    if results is None:
        return _not_computed(config, config_path,
                             f"stratified Cox did not fit: {frame_or_reason}",
                             strata_variables=strata_vars,
                             strata_variables_source=source)
    frame = frame_or_reason

    event_col = config['cox_settings']['event_column']
    n, n_events = int(len(frame)), int(frame[event_col].sum())
    sizes, events = strata_description(frame, strata_cols, event_col)
    recheck = ph_recheck(cph, frame)

    thin = sorted([k for k, v in events.items() if v < MIN_EVENTS_PER_STRATUM])
    if thin:
        qc_warning(f"{len(thin)} stratum/strata with fewer than "
                   f"{MIN_EVENTS_PER_STRATUM} events ({', '.join(thin[:5])}"
                   f"{'…' if len(thin) > 5 else ''}). They cost a baseline hazard each "
                   f"and contribute almost nothing; consider collapsing levels.")
    if len(sizes) > MAX_STRATA_LEVELS:
        qc_warning(f"{len(sizes)} strata: at this many, 'stratified' is closer to "
                   f"'matched on' and the within-stratum comparison is running out of "
                   f"data. Consider collapsing levels or a different remedy.")
    if recheck.get('status') == 'computed' and recheck['violated']:
        qc_warning(f"the proportional-hazards assumption is STILL violated after "
                   f"stratification ({', '.join(recheck['violating_variables'])}, "
                   f"min p = {recheck['min_p']:.4g}). The remedy was applied but did "
                   f"not work; this table is not a fix and must not be presented as "
                   f"one.")

    grp = grouping_column(config)
    exposure = None
    if grp and grp in results:
        exposure = {'variable': grp, 'hr': results[grp]['effect'],
                    'ci': [results[grp]['ci_lower'], results[grp]['ci_upper']],
                    'p': results[grp]['p']}

    sheets = write_excel(config, results, label_map, model_cols, strata_vars, source,
                         sizes, events, n, n_events, recheck, ph_art)

    evidence_path = write_evidence(config, config_path, {
        'status': 'computed',
        'method': 'stratified Cox proportional hazards '
                  '(lifelines CoxPHFitter, strata=)',
        'strata_variables': strata_vars,
        'strata_variables_source': source,
        'ph_violating_variables': list((ph_art or {}).get('violating_variables') or []),
        'source_ph_test': os.path.basename(ph_path) if ph_path else None,
        'source_ph_test_sha256': file_sha256(ph_path) if ph_path else None,
        'n': n,
        'n_events': n_events,
        'n_strata': len(sizes),
        'strata_sizes': sizes,
        'events_per_stratum': events,
        'covariates_in_model': list(model_cols),
        # By construction, and the gate cross-checks it against strata_variables:
        # a stratified variable with a hazard ratio would mean the strata never
        # reached the fit.
        'variables_without_hr': list(strata_vars),
        'exposure': exposure,
        'ph_recheck': recheck,
        'time_column': config['cox_settings']['time_column'],
        'event_column': event_col,
    })

    print(f"\n{'Term':<34} {'HR (95% CI)':>26} {'p':>10}")
    print("-" * 72)
    for col in model_cols:
        res = results.get(col)
        if res:
            print(f"{label_map.get(col, col):<34} "
                  f"{fmt_effect(res['effect'], res['ci_lower'], res['ci_upper']):>26} "
                  f"{fmt_p(res['p']):>10}")
    for sv in strata_vars:
        print(f"{sv + ' (stratum)':<34} {'— (stratified)':>26} {'—':>10}")

    print(f"\nExcel written    : {config['output_file']}")
    print(f"  Sheets         : {', '.join(sheets)}")
    print(f"Evidence written : {evidence_path}")
    print(f"  strata         : {', '.join(strata_vars)} [{source}], "
          f"{len(sizes)} strata, n = {n}, events = {n_events}")
    return 0


def main():
    if len(sys.argv) < 2:
        print("Usage: python ph_stratified.py config.json [--check-missing]")
        sys.exit(1)

    config_path = sys.argv[1]
    check_only = '--check-missing' in sys.argv

    config = load_config(config_path)
    df = load_data(config)
    variables = collect_variables(config)

    print(f"Mode: {config['mode']}  (remedy = {config['remedy']})")
    print(f"Data loaded: n={len(df)}")

    if check_only:
        preview(config, df, variables)
        sys.exit(0)

    rc = run(config, config_path, df, variables)

    print("\nDone!")
    if _QC_MESSAGES:
        n_block = sum(1 for m in _QC_MESSAGES if m.startswith('[QC_BLOCKING]'))
        n_warn = len(_QC_MESSAGES) - n_block
        print(f"  QC: {n_block} blocking, {n_warn} warning(s) written to stderr")
    sys.exit(rc)


if __name__ == '__main__':
    main()
