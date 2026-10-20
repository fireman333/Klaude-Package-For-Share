#!/usr/bin/env python3
"""Route tests on synthetic fixtures with known ground truth.

Why this exists
---------------
`test_pipeline_contract.py` has 516 passing assertions and mutation testing
showed it does not test wiring: deleting the `reviewer_gate_satisfied()` call
left it ALL PASS. Contract tests prove the code does what the code says. They
cannot prove the code answers the research question, because they never run an
analysis on data whose answer is known in advance.

These tests do. Each fixture is generated from a specified data-generating
process, so the correct answer is known before the pipeline runs, and each test
asserts that the route actually taken produces it:

  psm            an exposure that is genuinely harmful must come back OR > 1,
                 not its reciprocal
  psm direction  a PSM output whose group labels are inverted must be caught by
                 precheck_configs.py, not by five LLM reviewers in series
  left trunc.    delayed entry, where the naive fit is biased and the
                 entry-aware fit is not
  competing ev.  a competing event, where 1 - KM overstates incidence and the
                 CIF does not
  non-PH         a time-varying hazard ratio, which the Schoenfeld column must
                 flag
  MAR missing    missingness that depends on an observed covariate, which the
                 missing-data route must report rather than silently drop

Plus the deterministic units that precheck_configs.py depends on (the
success_criteria evaluator and the EPV parameter budget) and two wiring
assertions that a mutation to generate_configs.py would break.

Run:  cd <research-plan>/scripts && python3 -m pytest test_routes_synthetic.py -v

The route tests shell out to the real sub-skill scripts, which is the point:
a route test that stubs the analysis tests nothing. Set RP_SKIP_SLOW_ROUTES=1
to run only the fast deterministic ones.
"""
from __future__ import annotations

import hashlib
import json
import math
import os
import subprocess
import sys
from datetime import datetime

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from precheck_configs import (  # noqa: E402
    EVAL_FAIL,
    EVAL_PASS,
    EVAL_UNKNOWN,
    EVAL_UNPARSABLE,
    _LOG_BOUND,
    _RATIO_HIGH,
    _RATIO_LOW,
    _implied_ratio,
    estimate_scale,
    evaluate_criterion,
    parameter_budget,
    parse_labelled_effect,
    scan_output_numbers,
)

SCRIPTS = os.path.dirname(os.path.abspath(__file__))
GENERATE_CONFIGS = os.path.join(SCRIPTS, "generate_configs.py")
PRECHECK = os.path.join(SCRIPTS, "precheck_configs.py")

SKILL_ROOT = str(__import__("pathlib").Path(__file__).resolve().parents[2])
PSM_SCRIPT = f"{SKILL_ROOT}/psm-matching/scripts/psm_matching.py"
UNI_SCRIPT = f"{SKILL_ROOT}/univariate-forest/scripts/univariate_forest.py"
CR_SCRIPT = f"{SKILL_ROOT}/competing-risk/scripts/competing_risk.py"
KM_SCRIPT = f"{SKILL_ROOT}/kaplan-meier/scripts/kaplan_meier.py"
MISSING_SCRIPT = os.path.join(SCRIPTS, "missing_data.py")

SUBPROCESS_TIMEOUT = 300

# The route tests shell out to the real sub-skill scripts, which is the point of
# them, and each subprocess pays ~20s of lifelines/matplotlib import. Set
# RP_SKIP_SLOW_ROUTES=1 to run only the deterministic units. A skipif is used
# rather than a custom marker because registering a marker needs a conftest.py
# or pytest.ini, and this file is meant to stand alone.
slow = pytest.mark.skipif(
    os.environ.get("RP_SKIP_SLOW_ROUTES") == "1",
    reason="RP_SKIP_SLOW_ROUTES=1: skipping tests that run the real sub-skill scripts",
)


# ─────────────────────────────────────────────────────────────────────────────
# Fixture plumbing
# ─────────────────────────────────────────────────────────────────────────────

def _run(argv, cwd=None):
    """Run a script and return CompletedProcess. Output is never swallowed."""
    proc = subprocess.run([sys.executable, *argv], cwd=cwd, capture_output=True,
                          text=True, timeout=SUBPROCESS_TIMEOUT)
    if proc.returncode != 0:
        sys.stderr.write(f"\n--- {argv[0]} exited {proc.returncode} ---\n")
        sys.stderr.write(proc.stdout[-4000:])
        sys.stderr.write(proc.stderr[-4000:])
    return proc


# These fixtures state that Step 3.5 ran and found nothing, which for a dataset
# whose data-generating process is written a few lines below is simply true.
#
# `mode: "off"` is deliberately NOT used: generate_configs.py routes that through
# _diagnostic_override(plan, "diagnostics_off", ...), and "diagnostics_off" is not
# in its _CANONICAL_FLAGS table, so no override key can ever match it and the
# mode is currently unsatisfiable. Reported to the owner of that file rather than
# worked around there.
SYNTHETIC_DIAGNOSTICS = {"mode": "completed", "flags_detected": {}}

# Written next to every fixture plan. Only the left-truncation override test
# needs it, but an override that names a record has to find that record, and
# generate_configs.py looks for it in <output_dir>/decision_log.md.
DECISION_LOG = """# Decision log (synthetic fixtures)

## SYN-002 left truncation
Used only by the left-truncation override test in test_routes_synthetic.py:
delayed entry is present, and the cohort is restricted to a common time origin
before analysis rather than handled in-pipeline. Applies to outcomes: primary.
"""


# Most fixtures below write their output workbooks by hand, because what they
# test is how precheck reads a table, not whether the numbers in it are right.
# RES-014 refits the model in R and compares — which on a fabricated workbook
# correctly reports a huge disagreement, since the numbers were never computed
# from the data. results-audit.md provides exactly this escape hatch
# ("適用 toy run、教學"), so the toy fixtures declare themselves as such and
# TestCrossLanguageGrounding, which runs the real analysis, turns it back on.
TOY_GUARDIAN = {"cross_lang_check": False}


def _base_plan(tmpdir, data_path, outcome, covariates, pipeline, grouping=None,
               diagnostics=None, guardian=None):
    return {
        "version": 1,
        "study": {"title": "Synthetic fixture", "design": "cohort",
                  "data_file": str(data_path)},
        "population": {"filters": [], "exclusion_steps": []},
        "grouping": grouping or {
            "column": "exposed",
            "groups": [{"value": 1.0, "label": "Exposed"},
                       {"value": 0.0, "label": "Unexposed"}],
        },
        "outcomes": {"primary": outcome, "secondary": []},
        "covariates": covariates,
        "analyses": {},
        "missing_data": {"strategy": "assess_only", "primary_analysis": "complete_case"},
        "diagnostics": SYNTHETIC_DIAGNOSTICS if diagnostics is None else diagnostics,
        "guardian": dict(TOY_GUARDIAN) if guardian is None else guardian,
        "pipeline": pipeline,
        "output_dir": str(tmpdir),
    }


def _write_plan(tmpdir, plan):
    """Write the plan, plus the two records generate_configs.py insists exist.

    A `completed` diagnostics block has to point at a report that is on disk,
    carry a timestamp no earlier than the dataset, and (for an override) name a
    decision_log.md entry. Those are reasonable demands on a real project and
    they are cheap to satisfy honestly here: the fixture really does know what
    is in its own data.
    """
    out_dir = str(plan.get("output_dir") or tmpdir)
    os.makedirs(out_dir, exist_ok=True)

    with open(os.path.join(out_dir, "decision_log.md"), "w", encoding="utf-8") as f:
        f.write(DECISION_LOG)

    diag = plan.get("diagnostics") or {}
    if str(diag.get("mode", "")).lower() in ("completed", "on"):
        diag = dict(diag)
        report_dir = os.path.join(out_dir, "diagnostics")
        os.makedirs(report_dir, exist_ok=True)
        report_path = os.path.join(report_dir, "diagnostic_report.json")
        with open(report_path, "w", encoding="utf-8") as f:
            json.dump({"source": "test_routes_synthetic.py",
                       "flags_detected": diag.get("flags_detected", {}),
                       "note": "synthetic fixture; the data-generating process is "
                               "specified in the test module"}, f, indent=2)
        diag.setdefault("report_path", report_path)
        diag.setdefault("completed_at", datetime.now().isoformat())
        data_file = (plan.get("study") or {}).get("data_file")
        if data_file and os.path.exists(data_file):
            h = hashlib.sha256()
            with open(data_file, "rb") as f:
                for chunk in iter(lambda: f.read(1 << 20), b""):
                    h.update(chunk)
            diag.setdefault("data_fingerprint", h.hexdigest()[:16])
        plan = dict(plan, diagnostics=diag)

    path = os.path.join(str(tmpdir), "research_plan.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(plan, f, indent=2)
    return path


def _generate(plan_path):
    proc = _run([GENERATE_CONFIGS, plan_path])
    assert proc.returncode == 0, "generate_configs.py failed on a valid synthetic plan"
    return proc


def _config(tmpdir, step_id):
    with open(os.path.join(str(tmpdir), "configs", f"{step_id}_config.json"),
              encoding="utf-8") as f:
        return json.load(f)


def _xlsx_rows(path):
    """Every row of every sheet as a tuple, for assertion against output tables."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        return [(ws.title, row) for ws in wb.worksheets
                for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


# ─────────────────────────────────────────────────────────────────────────────
# Data-generating processes — the ground truth lives here
# ─────────────────────────────────────────────────────────────────────────────

def make_harmful_exposure(n=500, seed=11, log_or=1.0):
    """Confounded cohort where exposure genuinely RAISES the odds of the event.

    Truth: log OR = +log_or, so any correct analysis must report OR > 1.
    Confounding is by age and sex, both of which push exposure and outcome the
    same way, so a crude analysis over-states the effect but never flips it.
    """
    rng = np.random.default_rng(seed)
    age = rng.normal(65, 10, n)
    sex = rng.binomial(1, 0.5, n)
    ps = 1 / (1 + np.exp(-(-0.5 + 0.03 * (age - 65) + 0.4 * sex)))
    exposed = rng.binomial(1, ps)
    lin = -1.5 + log_or * exposed + 0.02 * (age - 65) + 0.3 * sex
    event = rng.binomial(1, 1 / (1 + np.exp(-lin)))
    return pd.DataFrame({"age": age, "sex": sex,
                         "exposed": exposed.astype(float), "dead": event})


def make_left_truncated(n=800, seed=13, true_log_hr=0.0):
    """Cohort with delayed entry.

    Subjects only enter the risk set at `entry_time`, and entry is later for
    the exposed. Truth: log HR = true_log_hr (0.0 = no effect). A fit that
    ignores entry counts pre-entry person-time as at risk and therefore has to
    invent a survival advantage for the group that entered later.
    """
    rng = np.random.default_rng(seed)
    exposed = rng.binomial(1, 0.5, n)
    baseline_rate = 0.10
    t_event = rng.exponential(1.0 / (baseline_rate * np.exp(true_log_hr * exposed)))
    entry = np.where(exposed == 1,
                     rng.uniform(0.5, 3.0, n),
                     rng.uniform(0.0, 0.3, n))
    admin_censor = 10.0
    observed = np.minimum(t_event, admin_censor)
    event = (t_event <= admin_censor).astype(int)
    keep = observed > entry  # left truncation: unobserved unless they survive to entry
    return pd.DataFrame({
        "exposed": exposed.astype(float)[keep],
        "entry_time": entry[keep],
        "time": observed[keep],
        "event": event[keep],
    }).reset_index(drop=True)


def make_competing_event(n=700, seed=17):
    """Cohort with a frequent competing event.

    Truth: the cumulative incidence of the event of interest is strictly below
    the 1 - KM estimate that treats the competing event as censoring, because
    censoring assumes those subjects could still have the event and they cannot.
    """
    rng = np.random.default_rng(seed)
    exposed = rng.binomial(1, 0.5, n)
    t_primary = rng.exponential(1 / (0.12 * np.exp(0.4 * exposed)))
    t_compete = rng.exponential(1 / 0.35, n)       # competing event is the common one
    admin = 8.0
    t = np.minimum(np.minimum(t_primary, t_compete), admin)
    status = np.where((t_primary <= t_compete) & (t_primary <= admin), 1,
                      np.where(t_compete < admin, 2, 0))
    return pd.DataFrame({"exposed": exposed.astype(float), "time": t,
                         "status": status.astype(float),
                         "event": (status == 1).astype(float)})


def make_non_ph(n=900, seed=19):
    """Cohort with a hazard ratio that crosses 1 during follow-up.

    Truth: the exposed have a HIGHER hazard early and a LOWER hazard late, so
    proportional hazards is false by construction and the Schoenfeld test must
    say so. A single averaged HR is meaningless here whatever its p-value.
    """
    rng = np.random.default_rng(seed)
    exposed = rng.binomial(1, 0.5, n)
    u = rng.uniform(size=n)
    # Piecewise-constant hazard: exposed 3x baseline before t=1, 1/3 after.
    early_rate = np.where(exposed == 1, 0.9, 0.3)
    late_rate = np.where(exposed == 1, 0.1, 0.3)
    cut = 1.0
    s_at_cut = np.exp(-early_rate * cut)
    t = np.where(u > s_at_cut,
                 -np.log(u) / early_rate,
                 cut + (-np.log(u) - early_rate * cut) / late_rate)
    admin = 6.0
    event = (t <= admin).astype(float)
    return pd.DataFrame({"exposed": exposed.astype(float),
                         "time": np.minimum(t, admin), "event": event})


def make_unbalanced_events(n_exposed=100, e_exposed=8, n_control=900, e_control=300,
                           seed=31):
    """A small exposed arm and a large control arm, with plenty of events overall.

    Truth: 308 events against a 2-parameter model — EPV 154, which is ample.
    Only the *exposure* coefficient is fragile, because one arm carries 8
    events. This is the shape that a min-group EPV numerator mis-reports as
    EPV 4.0 and blocks.
    """
    rng = np.random.default_rng(seed)
    n = n_exposed + n_control
    return pd.DataFrame({
        "exposed": np.r_[np.ones(n_exposed), np.zeros(n_control)],
        "age": rng.normal(65, 10, n),
        "dead": np.r_[np.r_[np.ones(e_exposed), np.zeros(n_exposed - e_exposed)],
                      np.r_[np.ones(e_control), np.zeros(n_control - e_control)]],
    })


def make_sparse_total_events(n=500, total_events=12, n_covariates=14, seed=37):
    """A cohort that really is too small for its model: 12 events, 15 parameters."""
    rng = np.random.default_rng(seed)
    data = pd.DataFrame({"exposed": rng.integers(0, 2, n).astype(float)})
    for i in range(n_covariates):
        data[f"c{i}"] = rng.integers(0, 2, n).astype(float)
    dead = np.zeros(n)
    dead[:total_events] = 1.0
    data["dead"] = dead
    half = total_events // 2
    data.loc[data.index[:half], "exposed"] = 1.0
    data.loc[data.index[half:total_events], "exposed"] = 0.0
    return data


def make_crude_or(target_or, n_per_arm=400, seed=41, p_control=0.30):
    """A cohort whose *crude* odds ratio is approximately `target_or`.

    Used to set up Simpson's-paradox scenarios: the crude association points one
    way and the correctly adjusted estimate points the other, which is what
    adjustment exists to do and must not be reported as a labelling error.
    """
    odds_control = p_control / (1 - p_control)
    p_exposed = (odds_control * target_or) / (1 + odds_control * target_or)
    rng = np.random.default_rng(seed)
    return pd.DataFrame({
        "exposed": np.r_[np.ones(n_per_arm), np.zeros(n_per_arm)],
        "age": rng.normal(65, 10, 2 * n_per_arm),
        "dead": np.r_[rng.binomial(1, p_exposed, n_per_arm),
                      rng.binomial(1, p_control, n_per_arm)].astype(float),
    })


def make_mar_missing(n=600, seed=23, missing_frac=0.60):
    """Cohort where a covariate is missing at random conditional on age.

    Truth: `biomarker` is missing for the older half at ~`missing_frac` and not
    at all for the younger half, so overall missingness clears the 20% COL-005
    threshold. Complete-case therefore drops a non-random, identifiable subset —
    which is exactly the claim the missing-data step exists to substantiate.
    """
    rng = np.random.default_rng(seed)
    age = rng.normal(65, 10, n)
    exposed = rng.binomial(1, 0.5, n)
    biomarker = rng.normal(5, 1.5, n)
    p_missing = np.where(age > np.median(age), missing_frac, 0.0)
    biomarker = np.where(rng.uniform(size=n) < p_missing, np.nan, biomarker)
    event = rng.binomial(1, 1 / (1 + np.exp(-(-1.0 + 0.6 * exposed + 0.02 * (age - 65)))))
    return pd.DataFrame({"age": age, "exposed": exposed.astype(float),
                         "biomarker": biomarker, "dead": event.astype(float)})


# ─────────────────────────────────────────────────────────────────────────────
# Fast deterministic units — the pieces precheck_configs.py stands on
# ─────────────────────────────────────────────────────────────────────────────

class TestSuccessCriteriaEvaluator:
    """The mini-DSL in plan_template.md:236-265, which had no parser until now."""

    @pytest.mark.parametrize("expr,metrics,expected", [
        ("events_per_group >= 1", {"events_per_group": 4}, EVAL_PASS),
        ("events_per_group >= 5", {"events_per_group": 4}, EVAL_FAIL),
        ("model_converged == true", {"model_converged": True}, EVAL_PASS),
        ("model_converged == true", {"model_converged": False}, EVAL_FAIL),
        ("hr_is_finite == true", {"hr_is_finite": True}, EVAL_PASS),
        ("ci_width < 100", {"ci_width": 12.5}, EVAL_PASS),
        ("ci_width < 100", {"ci_width": 1e9}, EVAL_FAIL),
        ("sample_size >= 30", {"sample_size": 30}, EVAL_PASS),
        ("smd_max < 0.1", {"smd_max": 0.035}, EVAL_PASS),
        ("groups_have_data", {"groups_have_data": True}, EVAL_PASS),
        ("input_files_exist", {"input_files_exist": False}, EVAL_FAIL),
        ("final_n > 0", {"final_n": 0}, EVAL_FAIL),
    ])
    def test_evaluates_the_documented_criteria(self, expr, metrics, expected):
        outcome, _detail = evaluate_criterion(expr, metrics)
        assert outcome == expected

    def test_missing_metric_is_unknown_not_pass(self):
        """A criterion that cannot be evaluated must never be recorded as passing."""
        outcome, detail = evaluate_criterion("smd_max < 0.1", {})
        assert outcome == EVAL_UNKNOWN
        assert "smd_max" in detail

    def test_none_metric_is_unknown_not_pass(self):
        outcome, _ = evaluate_criterion("ci_width < 100", {"ci_width": None})
        assert outcome == EVAL_UNKNOWN

    @pytest.mark.parametrize("expr", [
        "__import__('os').system('touch /tmp/pwned') > 0",
        "events > 1 and __import__('sys').exit()",
        "open('/etc/passwd').read() != ''",
        "events_per_group >= 1; print('x')",
        "lambda: 1",
        "",
    ])
    def test_rejects_anything_that_is_not_a_single_comparison(self, expr):
        """No eval(), no exec(): a criterion is one comparison or it is refused."""
        outcome, _detail = evaluate_criterion(expr, {"events": 5, "events_per_group": 5})
        assert outcome == EVAL_UNPARSABLE

    def test_no_side_effect_reached_the_filesystem(self):
        assert not os.path.exists("/tmp/pwned"), "the evaluator executed a criterion"

    def test_metric_to_metric_comparison(self):
        outcome, _ = evaluate_criterion("events >= events_per_group",
                                        {"events": 10, "events_per_group": 4})
        assert outcome == EVAL_PASS


class TestEstimateScaleTagging:
    """Which scale a column is on, taken from its header. The bound RES-001
    applies is chosen from this, so a wrong tag is a wrong verdict."""

    def _book(self, tmp_path, header, row):
        import openpyxl
        path = os.path.join(str(tmp_path), "t.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(header)
        ws.append(row)
        wb.save(path)
        return path

    @pytest.mark.parametrize("header,expected", [
        ("HR", "ratio"), ("OR", "ratio"), ("RR", "ratio"), ("sHR", "ratio"),
        ("Coefficient", "log"), ("Coef", "log"), ("Beta", "log"),
        ("Estimate", "unknown"),
        ("AUC", "unit_interval"), ("C-index", "unit_interval"),
    ])
    def test_header_fixes_the_scale(self, tmp_path, header, expected):
        path = self._book(tmp_path, ["Variable", header, "p-value"],
                          ["Age", 0.5, 0.04])
        effects, _p, _n, _u = scan_output_numbers(path)
        assert [e["scale"] for e in effects] == [expected]

    def test_formatted_cell_under_an_unrecognised_header_defaults_to_ratio(self, tmp_path):
        """`OR (95% CI)` is a header, just not one the resolver knows — which is
        the univariate-forest layout. Defaulting to "unknown" there would switch
        off the check that catches HR=1.2e6.
        """
        path = self._book(tmp_path, ["Variable", "N (events)", "OR (95% CI)"],
                          ["Age", "300 (95)", "2.22 (1.29-3.82)"])
        effects, _p, _n, _u = scan_output_numbers(path)
        assert [e["scale"] for e in effects] == ["ratio"]

    @pytest.mark.parametrize("header,expected", [
        # Exactly what auc_cstat.py:780 composes, newlines and all.
        ("C-index\n(variable alone)\n[95% CI]", "unit_interval"),
        ("C-index\n(clinical model\n+ variable)\n[95% CI]", "unit_interval"),
        ("AUC [95% CI]", "unit_interval"),
        ("Delta C-index\n[95% CI]", "signed_unit_interval"),
        ("Delta AUC\n[95% CI]", "signed_unit_interval"),
    ])
    def test_the_headers_auc_cstat_actually_writes_resolve(self, header, expected):
        """The previous version keyed on bare 'c-index' / 'auc', which no producer
        emits — so every discrimination column fell through to the ratio default.
        """
        assert estimate_scale(header) == expected

    @pytest.mark.parametrize("header", [
        "log HR", "SE(log HR)", "HR vs reference",     # nonlinearity.py:813
        "Coefficient (log scale)", "Naive coef",       # robust_variance.py:772/800
        "Observed RR (95% CI)", "E-value (CI)",        # evalue.py:267
        "Adjusted HR (95% CI)", "Risk-ratio scale",
    ])
    def test_the_resolver_stays_narrow(self, header):
        """Widening past the discrimination family is not free. nonlinearity's
        Spline_estimates sheet has recognised `CI lower` / `CI upper` columns that
        belong to the RATIO column; a newly recognised `log HR` beside them would
        pair with the wrong interval and block on 'outside its own CI'.
        """
        assert estimate_scale(header) is None

    @pytest.mark.parametrize("coef,expect", [
        (-1000.0, "0"),           # not "inf" — the opposite end of the scale
        (27.4, "7.93723e+11"),
        (700.0, "1.01423e+304"),  # large, but finite and representable
    ])
    def test_the_implied_ratio_keeps_its_sign(self, coef, expect):
        """Guarding on abs(coef) reported a large negative coefficient as implying
        an infinite ratio, in the one sentence a reader diagnoses the fit from.
        """
        assert _implied_ratio(coef) == expect

    def test_an_unrepresentable_ratio_says_so_rather_than_raising(self):
        assert "too large" in _implied_ratio(710.0)

    @pytest.mark.parametrize("cell,expected", [
        # weighting.py writes f"{scale}: " + fmt_effect(...), so every one of its
        # estimates was invisible and the step blocked on "no effect estimate".
        ("Hazard ratio: 1.036 (0.827–1.297)", (1.036, 0.827, 1.297)),
        ("Odds ratio: 2.10 (1.40–3.15)", (2.10, 1.40, 3.15)),
        # Prose that merely contains a bracketed range is not an estimate.
        ("Note: see figure 3 (panel 2–4)", None),
        ("Method: complete case (n 2–3)", None),
        # ph_rmst's contrast sentence. Signed, and must stay unreadable here.
        ("Unexposed minus Exposed: +0.518 years (95% CI 0.370 to 0.665); p = <0.001",
         None),
        # No label at all: that is parse_effect's job, not this one's.
        ("2.22 (1.29–3.82)", None),
    ])
    def test_a_labelled_effect_is_read_without_swallowing_prose(self, cell, expected):
        got = parse_labelled_effect(cell)
        assert (got[1] if got else None) == expected

    def test_the_label_supplies_the_scale(self):
        assert estimate_scale(parse_labelled_effect(
            "Hazard ratio: 1.04 (0.83–1.30)")[0]) == "ratio"

    def test_an_unrecognised_label_is_unknown_rather_than_ratio(self):
        """A producer that named its scale must not have a different one assumed.
        Reading `Risk difference: -0.043 (...)` as a ratio would block a correct
        negative difference — the Delta C-index bug, one layer down.
        """
        label, eff = parse_labelled_effect("Risk difference: -0.043 (-0.081–-0.005)")
        assert eff[0] < 0
        assert estimate_scale(label) is None  # -> the caller stores "unknown"

    def test_formatted_cell_inherits_its_columns_scale(self, tmp_path):
        """A coefficient printed with its interval is still on the log scale."""
        path = self._book(tmp_path, ["Variable", "Coefficient", "p-value"],
                          ["Male sex", "-0.36 (-0.60--0.10)", 0.004])
        effects, _p, _n, _u = scan_output_numbers(path)
        assert [(e["scale"], e["estimate"]) for e in effects] == [("log", -0.36)]

    def test_the_log_bound_is_the_ratio_bound_transported(self):
        """One policy, two scales. Stated twice, they drift."""
        assert math.isclose(math.exp(_LOG_BOUND), _RATIO_HIGH)
        assert math.isclose(math.exp(-_LOG_BOUND), _RATIO_LOW)


class TestParameterBudget:
    """EPV denominators. Counting variables instead of parameters is how a
    15-parameter model gets recorded as 8 and passes a rule-of-10 gate."""

    def _plan(self, covariates, nonlinearity=None):
        plan = {
            "grouping": {"column": "exposed",
                         "groups": [{"value": 1}, {"value": 0}]},
            "covariates": covariates,
        }
        if nonlinearity:
            plan["advanced_analyses"] = {"nonlinearity": nonlinearity}
        return plan

    def test_continuous_and_binary_cost_one_each(self):
        total, _breakdown, _prov = parameter_budget(self._plan([
            {"column": "age", "type": "continuous"},
            {"column": "sex", "type": "binary"},
        ]))
        assert total == 3  # 1 exposure contrast + age + sex

    def test_categorical_costs_levels_minus_one(self):
        total, breakdown, _prov = parameter_budget(self._plan([
            {"column": "nyha", "type": "ordinal",
             "label_map": {"1": "I", "2": "II", "3": "III", "4": "IV"}},
        ]))
        assert total == 4  # 1 exposure + 3 for a four-level factor
        assert any(b["params"] == 3 for b in breakdown)

    def test_spline_costs_knots_minus_one_not_one(self):
        """An RCS on 5 knots is 4 parameters. This is the count an external reviewer insisted on."""
        plan = self._plan(
            [{"column": "age", "type": "continuous"},
             {"column": "sex", "type": "binary"}],
            nonlinearity={"enabled": True, "variables": ["age"], "knots": 5},
        )
        total, breakdown, _prov = parameter_budget(plan)
        assert total == 6  # 1 exposure + 4 spline + 1 binary
        age = next(b for b in breakdown if b["term"] == "age")
        assert age["params"] == 4

    def test_spline_not_counted_when_nonlinearity_disabled(self):
        plan = self._plan(
            [{"column": "age", "type": "continuous"}],
            nonlinearity={"enabled": False, "variables": ["age"], "knots": 5},
        )
        total, _breakdown, _prov = parameter_budget(plan)
        assert total == 2

    def test_knot_percentiles_override_the_knot_count(self):
        plan = self._plan(
            [{"column": "age", "type": "continuous"}],
            nonlinearity={"enabled": True, "variables": ["age"], "knots": 4,
                          "knot_percentiles": [5, 27.5, 50, 72.5, 95]},
        )
        total, _breakdown, _prov = parameter_budget(plan)
        assert total == 5  # 1 exposure + 4 spline parameters from 5 knots

    def test_low_epv_scenario_is_computed_not_assumed(self):
        """n=500, 12 events, 15 parameters — EPV 0.8.

        The Stage Reflector rule this replaces triggered on `Any group n < 30`
        while its formula used events (checkpoint_protocol.md:135), so this
        cohort never fired it.
        """
        covs = [{"column": f"c{i}", "type": "binary"} for i in range(14)]
        total, _breakdown, _prov = parameter_budget(self._plan(covs))
        assert total == 15
        assert 12 / total < 5


# ─────────────────────────────────────────────────────────────────────────────
# Wiring assertions — a mutation to generate_configs.py must break these
# ─────────────────────────────────────────────────────────────────────────────

class TestGeneratorWiring:

    def test_manuscript_without_reviewer_gate_is_refused(self, tmp_path):
        """SKILL.md:579 requires manuscript to depend on reviewer_challenge.

        Deleting `_validate_manuscript_is_gated` from validate_plan() makes this
        test fail, which is what the 516 contract assertions did not do.
        """
        data = make_harmful_exposure(n=120)
        data_path = tmp_path / "data.csv"
        data.to_csv(data_path, index=False)
        plan = _base_plan(
            tmp_path, data_path,
            {"id": "primary", "name": "Death", "type": "binary", "event_column": "dead"},
            [{"name": "Age", "column": "age", "type": "continuous"}],
            [{"id": "missingness", "skill": "missing-data", "status": "pending",
              "depends_on": []},
             {"id": "uni_primary", "skill": "univariate-forest", "status": "pending",
              "depends_on": [], "outcome": "primary", "output_file": "Uni_primary"},
             {"id": "manuscript", "skill": "manuscript-results", "status": "pending",
              "depends_on": ["uni_primary"]}],
        )
        proc = _run([GENERATE_CONFIGS, _write_plan(tmp_path, plan)])
        assert proc.returncode != 0, "an ungated manuscript step was accepted"
        assert "reviewer" in (proc.stdout + proc.stderr).lower()

    def test_models_without_a_missingness_step_are_refused(self, tmp_path):
        """missing_data_protocol.md §1: the assessment is not optional."""
        data = make_harmful_exposure(n=120)
        data_path = tmp_path / "data.csv"
        data.to_csv(data_path, index=False)
        plan = _base_plan(
            tmp_path, data_path,
            {"id": "primary", "name": "Death", "type": "binary", "event_column": "dead"},
            [{"name": "Age", "column": "age", "type": "continuous"}],
            [{"id": "uni_primary", "skill": "univariate-forest", "status": "pending",
              "depends_on": [], "outcome": "primary", "output_file": "Uni_primary"}],
        )
        proc = _run([GENERATE_CONFIGS, _write_plan(tmp_path, plan)])
        assert proc.returncode != 0
        assert "missing" in (proc.stdout + proc.stderr).lower()

    def test_competing_risk_primary_routes_generic_skills_to_cause_specific_cox(self, tmp_path):
        """SKILL.md:311 — competing-risk primaries run cause-specific Cox in the
        generic skills, and the substitution has to be announced, not silent."""
        data = make_competing_event(n=200)
        data_path = tmp_path / "data.csv"
        data.to_csv(data_path, index=False)
        plan = _base_plan(
            tmp_path, data_path,
            {"id": "primary", "name": "Event", "type": "competing_risk",
             "time_column": "time", "event_column": "event",
             "competing_event_column": "status", "time_unit": "years", "time_scale": 1},
            [{"name": "Exposure", "column": "exposed", "type": "binary", "reference": 0}],
            [{"id": "missingness", "skill": "missing-data", "status": "pending",
              "depends_on": []},
             {"id": "uni_primary", "skill": "univariate-forest", "status": "pending",
              "depends_on": [], "outcome": "primary", "output_file": "Uni_primary"}],
        )
        proc = _generate(_write_plan(tmp_path, plan))
        cfg = _config(tmp_path, "uni_primary")
        assert cfg["analysis_type"] == "cox"
        assert cfg["cox_settings"]["event_column"] == "event"
        assert "METHOD_WARNING" in proc.stderr, \
            "cause-specific substitution was applied without announcing it"


# ─────────────────────────────────────────────────────────────────────────────
# Route tests — known truth in, assert the answer out
# ─────────────────────────────────────────────────────────────────────────────

@slow
class TestPsmDirection:

    @pytest.fixture(scope="class")
    def psm_run(self, tmp_path_factory):
        tmp = tmp_path_factory.mktemp("psm")
        data = make_harmful_exposure(n=500, log_or=1.0)
        data_path = tmp / "data.csv"
        data.to_csv(data_path, index=False)
        plan = _base_plan(
            tmp, data_path,
            {"id": "primary", "name": "Death", "type": "binary", "event_column": "dead"},
            [{"name": "Age", "column": "age", "type": "continuous"},
             {"name": "Male", "column": "sex", "type": "binary", "reference": 0}],
            [{"id": "missingness", "skill": "missing-data", "status": "pending",
              "depends_on": []},
             {"id": "psm", "skill": "psm-matching", "status": "pending",
              "depends_on": [], "output_file": "PSM"}],
        )
        plan_path = _write_plan(tmp, plan)
        _generate(plan_path)
        proc = _run([PSM_SCRIPT, os.path.join(str(tmp), "configs", "psm_config.json")],
                    cwd=str(tmp))
        assert proc.returncode == 0, "psm-matching route failed on a clean fixture"
        return tmp, plan_path

    def test_harmful_exposure_comes_back_as_or_above_one(self, psm_run):
        """The fixture's true log OR is +1.0. An OR below 1 is the reciprocal bug."""
        tmp, _ = psm_run
        psm = pd.read_csv(os.path.join(str(tmp), "PSM.csv"))
        assert len(psm) == 1
        row = psm.iloc[0]
        assert row["OR"] > 1.0, f"harmful exposure reported as protective (OR {row['OR']})"
        assert row["95% CI Lower"] > 1.0, "CI includes the null for a strong true effect"
        assert 1.5 < row["OR"] < 8.0, f"OR {row['OR']} implausible for a true log OR of 1.0"

    def test_exposed_column_comes_first(self, psm_run):
        """grouping.groups[0] is the arm psm_matching.py treats as exposed."""
        tmp, _ = psm_run
        cols = list(pd.read_csv(os.path.join(str(tmp), "PSM.csv")).columns)
        assert cols.index("Exposed N") < cols.index("Unexposed N")

    def test_precheck_accepts_the_correct_direction(self, psm_run):
        tmp, plan_path = psm_run
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "[QC_BLOCKING] PSM-DIR" not in proc.stderr, \
            "the direction anchor fired on a correctly oriented output"

    def test_precheck_catches_a_reciprocal_psm_output(self, psm_run, tmp_path):
        """The failure that five separate LLM reviewers each found once.

        The output is rewritten so the reported OR is the reciprocal of the
        crude effect — exactly what an inverted group label produces — and the
        deterministic anchor has to catch it without anyone reading the table.
        """
        tmp, _ = psm_run
        work = tmp_path / "inverted"
        work.mkdir()
        for name in ("data.csv", "research_plan.json"):
            (work / name).write_text((tmp / name).read_text(encoding="utf-8"),
                                     encoding="utf-8")
        plan = json.loads((work / "research_plan.json").read_text(encoding="utf-8"))
        plan["output_dir"] = str(work)
        plan["study"]["data_file"] = str(work / "data.csv")
        (work / "research_plan.json").write_text(json.dumps(plan, indent=2), encoding="utf-8")
        _generate(str(work / "research_plan.json"))

        original = pd.read_csv(os.path.join(str(tmp), "PSM.csv"))
        inverted = original.copy()
        inverted["OR"] = 1.0 / original["OR"]
        inverted["95% CI Lower"] = 1.0 / original["95% CI Upper"]
        inverted["95% CI Upper"] = 1.0 / original["95% CI Lower"]
        inverted.to_csv(work / "PSM.csv", index=False)

        proc = _run([PRECHECK, str(work / "research_plan.json"), "--stage", "post"])
        assert "PSM-DIR" in proc.stderr
        assert "[QC_BLOCKING] PSM-DIR" in proc.stderr, \
            "an inverted PSM effect passed the direction anchor"
        assert proc.returncode == 1


@slow
class TestLeftTruncation:
    """Delayed entry. The pipeline has no entry-aware route, so the assertion
    here is the one that justifies blocking at plan time: the naive fit the
    pipeline *does* offer is biased on data whose truth is known."""

    @pytest.fixture(scope="class")
    def truncated(self):
        return make_left_truncated(n=1200, true_log_hr=0.0)

    def test_fixture_really_is_left_truncated(self, truncated):
        assert (truncated["entry_time"] > 0).mean() > 0.5
        assert (truncated["time"] > truncated["entry_time"]).all()
        late = truncated.loc[truncated["exposed"] == 1, "entry_time"].mean()
        early = truncated.loc[truncated["exposed"] == 0, "entry_time"].mean()
        assert late > early + 0.5, "the fixture has no differential delayed entry"

    def test_naive_fit_is_biased_and_entry_aware_fit_is_not(self, truncated):
        """True log HR is 0. Ignoring entry manufactures a protective effect."""
        from lifelines import CoxPHFitter

        naive = CoxPHFitter().fit(truncated[["time", "event", "exposed"]],
                                  duration_col="time", event_col="event")
        naive_log_hr = float(naive.params_["exposed"])

        aware = CoxPHFitter().fit(truncated[["entry_time", "time", "event", "exposed"]],
                                  duration_col="time", event_col="event",
                                  entry_col="entry_time")
        aware_log_hr = float(aware.params_["exposed"])

        assert abs(aware_log_hr) < 0.20, \
            f"entry-aware fit did not recover the null (log HR {aware_log_hr:.3f})"
        assert naive_log_hr < -0.30, \
            f"naive fit was expected to be biased protective, got {naive_log_hr:.3f}"
        assert abs(naive_log_hr) > abs(aware_log_hr) + 0.20, \
            "the naive and entry-aware fits are not distinguishable on this fixture"

    def _truncated_plan(self, tmp_path, truncated, diagnostics):
        data_path = tmp_path / "data.csv"
        truncated.to_csv(data_path, index=False)
        plan = _base_plan(
            tmp_path, data_path,
            {"id": "primary", "name": "Death", "type": "time_to_event",
             "time_column": "time", "event_column": "event",
             "time_unit": "years", "time_scale": 1},
            [{"name": "Exposure", "column": "exposed", "type": "binary", "reference": 0}],
            [{"id": "missingness", "skill": "missing-data", "status": "pending",
              "depends_on": []},
             {"id": "uni_primary", "skill": "univariate-forest", "status": "pending",
              "depends_on": [], "outcome": "primary", "output_file": "Uni_primary"}],
            diagnostics=diagnostics,
        )
        return _write_plan(tmp_path, plan)

    def test_flagged_delayed_entry_is_refused_at_plan_time(self, tmp_path, truncated):
        """No pipeline skill handles delayed entry, so the only correct output is
        a refusal. The bias measured above is what a "limitation" would be
        papering over."""
        plan_path = self._truncated_plan(
            tmp_path, truncated,
            {"mode": "completed", "flags_detected": ["left_truncation"]})
        proc = _run([GENERATE_CONFIGS, plan_path])
        assert proc.returncode != 0, "a left-truncated cohort was accepted"
        assert "left_truncation" in (proc.stdout + proc.stderr)

    def test_the_refusal_is_overridable_only_on_the_record(self, tmp_path, truncated):
        plan_path = self._truncated_plan(
            tmp_path, truncated,
            {"mode": "completed", "flags_detected": ["left_truncation"],
             "overrides": {"left_truncation": {
                 "decision_id": "SYN-002",
                 "reason": "cohort restricted to a common time origin",
                 "outcomes": ["primary"]}}})
        proc = _run([GENERATE_CONFIGS, plan_path])
        assert proc.returncode == 0
        assert "METHOD_WARNING" in proc.stderr, \
            "the override was accepted without recording a methodological warning"

    def test_the_naive_route_runs_when_nothing_flagged_it(self, tmp_path, truncated):
        """With diagnostics off, the same data produces a number and no column in
        the config carries the entry time — which is why the gate above has to
        exist upstream rather than downstream."""
        plan_path = self._truncated_plan(tmp_path, truncated, None)
        _generate(plan_path)
        cfg = _config(tmp_path, "uni_primary")
        assert "entry_time" not in json.dumps(cfg), \
            "a delayed-entry column reached the config; update this test and the gate"
        proc = _run([UNI_SCRIPT, os.path.join(str(tmp_path), "configs",
                                              "uni_primary_config.json")],
                    cwd=str(tmp_path))
        assert proc.returncode == 0
        assert os.path.exists(os.path.join(str(tmp_path), "Uni_primary.xlsx"))


@slow
class TestCompetingEvent:

    @pytest.fixture(scope="class")
    def cr_run(self, tmp_path_factory):
        tmp = tmp_path_factory.mktemp("cr")
        data = make_competing_event(n=700)
        data_path = tmp / "data.csv"
        data.to_csv(data_path, index=False)
        plan = _base_plan(
            tmp, data_path,
            {"id": "primary", "name": "Event of interest", "type": "competing_risk",
             "time_column": "time", "event_column": "event",
             "competing_event_column": "status",
             "competing_event_labels": {"1": "Event of interest", "2": "Competing death"},
             "time_unit": "years", "time_scale": 1, "time_points": [1, 3, 5]},
            [{"name": "Exposure", "column": "exposed", "type": "binary", "reference": 0}],
            [{"id": "missingness", "skill": "missing-data", "status": "pending",
              "depends_on": []},
             {"id": "km_primary", "skill": "kaplan-meier", "status": "pending",
              "depends_on": [], "outcome": "primary", "output_file": "KM_primary"},
             {"id": "cr_primary", "skill": "competing-risk", "status": "pending",
              "depends_on": [], "outcome": "primary", "output_file": "CR_primary"}],
        )
        plan_path = _write_plan(tmp, plan)
        _generate(plan_path)
        for step in ("km_primary", "cr_primary"):
            script = KM_SCRIPT if step.startswith("km") else CR_SCRIPT
            proc = _run([script, os.path.join(str(tmp), "configs", f"{step}_config.json")],
                        cwd=str(tmp))
            assert proc.returncode == 0, f"{step} route failed"
        return tmp, data

    def test_competing_risk_route_reads_the_multistate_column(self, cr_run):
        tmp, _ = cr_run
        cfg = _config(tmp, "cr_primary")
        assert cfg["event_column"] == "status", \
            "the competing-risk route was pointed at the binary indicator, not the status column"
        assert cfg["primary_event"] == 1

    def test_km_overstates_incidence_relative_to_the_cif(self, cr_run):
        """1 - KM treats the competing event as censoring and must come out higher."""
        from lifelines import AalenJohansenFitter, KaplanMeierFitter

        _tmp, data = cr_run
        t, status = data["time"].values, data["status"].values
        km = KaplanMeierFitter().fit(t, (status == 1).astype(int))
        aj = AalenJohansenFitter().fit(t, status, event_of_interest=1)

        horizon = 3.0
        km_incidence = 1.0 - float(km.predict(horizon))
        cif = float(aj.predict(horizon).iloc[0]) if hasattr(aj.predict(horizon), "iloc") \
            else float(aj.predict(horizon))
        assert km_incidence > cif + 0.02, \
            f"1-KM {km_incidence:.3f} did not exceed the CIF {cif:.3f}; fixture is wrong"
        assert cif > 0.05, "the event of interest is too rare to test anything"

    def test_both_outputs_exist_so_the_pair_can_be_compared(self, cr_run):
        tmp, _ = cr_run
        assert os.path.exists(os.path.join(str(tmp), "KM_primary.xlsx"))
        assert os.path.exists(os.path.join(str(tmp), "CR_primary.xlsx"))


@slow
class TestNonProportionalHazards:

    @pytest.fixture(scope="class")
    def ph_run(self, tmp_path_factory):
        tmp = tmp_path_factory.mktemp("nonph")
        data = make_non_ph(n=900)
        data_path = tmp / "data.csv"
        data.to_csv(data_path, index=False)
        plan = _base_plan(
            tmp, data_path,
            {"id": "primary", "name": "Death", "type": "time_to_event",
             "time_column": "time", "event_column": "event",
             "time_unit": "years", "time_scale": 1},
            [{"name": "Exposure", "column": "exposed", "type": "binary", "reference": 0}],
            [{"id": "missingness", "skill": "missing-data", "status": "pending",
              "depends_on": []},
             {"id": "uni_primary", "skill": "univariate-forest", "status": "pending",
              "depends_on": [], "outcome": "primary", "output_file": "Uni_primary"}],
        )
        _generate(_write_plan(tmp, plan))
        proc = _run([UNI_SCRIPT, os.path.join(str(tmp), "configs",
                                              "uni_primary_config.json")], cwd=str(tmp))
        assert proc.returncode == 0
        return tmp, data

    def test_fixture_hazard_ratio_really_does_cross_one(self, ph_run):
        from lifelines import CoxPHFitter

        _tmp, data = ph_run
        early = data.copy()
        early["event"] = np.where(data["time"] <= 1.0, data["event"], 0.0)
        early["time"] = np.minimum(data["time"], 1.0)
        late = data[data["time"] > 1.0].copy()
        late["time"] = late["time"] - 1.0

        hr_early = math.exp(float(CoxPHFitter().fit(
            early[["time", "event", "exposed"]], "time", "event").params_["exposed"]))
        hr_late = math.exp(float(CoxPHFitter().fit(
            late[["time", "event", "exposed"]], "time", "event").params_["exposed"]))
        assert hr_early > 1.5, f"early HR {hr_early:.2f} is not the harmful phase"
        assert hr_late < 0.8, f"late HR {hr_late:.2f} is not the protective phase"

    def test_route_emits_a_ph_column_that_flags_the_violation(self, ph_run):
        """The Schoenfeld p is the only thing standing between a crossing hazard
        and a single averaged HR in Table 2."""
        tmp, _ = ph_run
        rows = _xlsx_rows(os.path.join(str(tmp), "Uni_primary.xlsx"))
        headers = [r for _sheet, r in rows if r and "PH p" in [str(c) for c in r if c]]
        assert headers, "the Cox route produced no PH p column"

        ph_values = []
        for _sheet, row in rows:
            cells = [str(c) for c in row if c is not None]
            if not any("Exposure" in c for c in cells):
                continue
            for c in cells:
                try:
                    ph_values.append(float(c.lstrip("<").rstrip("*")))
                except ValueError:
                    continue
        assert ph_values, "no numeric cells on the exposure row"
        assert min(ph_values) < 0.05, \
            "nothing on the exposure row was significant; the PH violation went unflagged"


@slow
class TestMarMissingness:

    @pytest.fixture(scope="class")
    def mar_run(self, tmp_path_factory):
        tmp = tmp_path_factory.mktemp("mar")
        data = make_mar_missing(n=600, missing_frac=0.60)
        data_path = tmp / "data.csv"
        data.to_csv(data_path, index=False)
        plan = _base_plan(
            tmp, data_path,
            {"id": "primary", "name": "Death", "type": "binary", "event_column": "dead"},
            [{"name": "Age", "column": "age", "type": "continuous"},
             {"name": "Biomarker", "column": "biomarker", "type": "continuous"}],
            [{"id": "missingness", "skill": "missing-data", "status": "pending",
              "depends_on": [], "output_file": "Missingness"}],
        )
        plan_path = _write_plan(tmp, plan)
        _generate(plan_path)
        proc = _run([MISSING_SCRIPT, os.path.join(str(tmp), "configs",
                                                  "missingness_config.json")],
                    cwd=str(tmp))
        return tmp, plan_path, data, proc

    def test_fixture_missingness_depends_on_an_observed_covariate(self, mar_run):
        _tmp, _plan, data, _proc = mar_run
        miss = data["biomarker"].isna()
        old = data["age"] > data["age"].median()
        assert miss.mean() > 0.10
        assert miss[old].mean() > miss[~old].mean() + 0.10, \
            "the fixture's missingness is not related to age; it is not MAR by design"

    def test_missing_data_route_runs_and_reports(self, mar_run):
        tmp, _plan, _data, proc = mar_run
        combined = proc.stdout + proc.stderr
        assert proc.returncode == 0, "missing-data route failed on a MAR fixture"
        assert "biomarker" in combined.lower() or \
            any(f.startswith("Missingness") for f in os.listdir(str(tmp))), \
            "the missing-data step neither named the incomplete column nor wrote an output"

    def test_precheck_warns_about_the_incomplete_covariate(self, mar_run):
        """COL-005: a covariate above the 20% threshold has to be surfaced before
        every downstream model silently drops those rows."""
        _tmp, plan_path, _data, _proc = mar_run
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert "COL-005" in proc.stderr
        assert "biomarker" in proc.stderr
        assert "[QC_WARNING] COL-005" in proc.stderr


# ─────────────────────────────────────────────────────────────────────────────
# precheck end-to-end on a fixture with planted defects
# ─────────────────────────────────────────────────────────────────────────────

class TestPrecheckCatchesPlantedDefects:

    def _plan_with(self, tmp_path, data, covariates, outcome_overrides=None):
        data_path = tmp_path / "data.csv"
        data.to_csv(data_path, index=False)
        outcome = {"id": "primary", "name": "Death", "type": "binary",
                   "event_column": "dead"}
        outcome.update(outcome_overrides or {})
        plan = _base_plan(
            tmp_path, data_path, outcome, covariates,
            [{"id": "missingness", "skill": "missing-data", "status": "pending",
              "depends_on": []},
             {"id": "uni_primary", "skill": "univariate-forest", "status": "pending",
              "depends_on": [], "outcome": "primary", "output_file": "Uni_primary"}],
        )
        return _write_plan(tmp_path, plan)

    def test_missing_column_is_blocking(self, tmp_path):
        plan_path = self._plan_with(
            tmp_path, make_harmful_exposure(n=200),
            [{"name": "Ghost", "column": "not_a_column", "type": "continuous"}])
        _generate(plan_path)
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert proc.returncode == 1
        assert "[QC_BLOCKING] COL-001" in proc.stderr
        assert "not_a_column" in proc.stderr

    def test_absent_reference_level_is_blocking(self, tmp_path):
        plan_path = self._plan_with(
            tmp_path, make_harmful_exposure(n=200),
            [{"name": "Male", "column": "sex", "type": "binary", "reference": 7}])
        _generate(plan_path)
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert proc.returncode == 1
        assert "[QC_BLOCKING] COL-003" in proc.stderr

    def test_non_binary_event_column_is_blocking(self, tmp_path):
        data = make_harmful_exposure(n=200)
        data.loc[data.index[:20], "dead"] = 9
        plan_path = self._plan_with(
            tmp_path, data,
            [{"name": "Age", "column": "age", "type": "continuous"}])
        _generate(plan_path)
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert proc.returncode == 1
        assert "[QC_BLOCKING] COL-002" in proc.stderr

    def test_low_epv_is_blocking(self, tmp_path):
        """12 events against 15 fitted parameters — EPV 0.8, never previously computed."""
        rng = np.random.default_rng(3)
        n = 500
        data = pd.DataFrame({"exposed": rng.binomial(1, 0.5, n).astype(float)})
        for i in range(14):
            data[f"c{i}"] = rng.binomial(1, 0.4, n).astype(float)
        dead = np.zeros(n)
        dead[:12] = 1.0
        data["dead"] = dead
        data.loc[data.index[:6], "exposed"] = 1.0
        data.loc[data.index[6:12], "exposed"] = 0.0
        plan_path = self._plan_with(
            tmp_path, data,
            [{"name": f"C{i}", "column": f"c{i}", "type": "binary", "reference": 0}
             for i in range(14)])
        _generate(plan_path)
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert "[QC_BLOCKING] RES-006" in proc.stderr
        assert proc.returncode == 1

    def test_clean_fixture_passes(self, tmp_path):
        plan_path = self._plan_with(
            tmp_path, make_harmful_exposure(n=400),
            [{"name": "Age", "column": "age", "type": "continuous"},
             {"name": "Male", "column": "sex", "type": "binary", "reference": 0}])
        _generate(plan_path)
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert proc.returncode == 0, f"clean fixture reported blocking issues:\n{proc.stderr}"

    def test_report_json_is_written_and_machine_readable(self, tmp_path):
        plan_path = self._plan_with(
            tmp_path, make_harmful_exposure(n=300),
            [{"name": "Age", "column": "age", "type": "continuous"}])
        _generate(plan_path)
        _run([PRECHECK, plan_path, "--stage", "pre"])
        report_path = os.path.join(str(tmp_path), "precheck_report.json")
        assert os.path.exists(report_path)
        with open(report_path, encoding="utf-8") as f:
            payload = json.load(f)
        for key in ("stage", "counters", "epv", "issues", "skipped_checks", "cohort"):
            assert key in payload
        assert payload["counters"]["checks_run"] > 0


# ─────────────────────────────────────────────────────────────────────────────
# estimate_inventory / check_numbers_sourced round trip
# ─────────────────────────────────────────────────────────────────────────────

class TestCoverageAndSourcing:

    @pytest.fixture
    def executed_outputs(self, tmp_path):
        """A workbook holding three estimates, of which the report will quote one."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Univariate Analysis"
        ws.append(["Variable", "N (events)", "OR (95% CI)", "p-value"])
        ws.append(["Exposure", "400 (120)", "2.22 (1.29–3.82)", "0.004*"])
        ws.append(["Age", "400 (120)", "1.05 (1.01–1.09)", "0.012*"])
        ws.append(["Sex", "400 (120)", "0.88 (0.55–1.41)", "0.601"])
        wb.save(os.path.join(str(tmp_path), "Uni_primary.xlsx"))
        with open(os.path.join(str(tmp_path), "sensitivity_stats.json"), "w",
                  encoding="utf-8") as f:
            json.dump({"stratum_a": {"or": 3.11}, "stratum_b": {"or": 0.74}}, f)
        return tmp_path

    def test_omitted_estimates_are_listed(self, executed_outputs, tmp_path):
        report = tmp_path / "report.md"
        report.write_text(
            "Exposure was associated with the outcome (OR 2.22, 95% CI 1.29-3.82).\n",
            encoding="utf-8")
        proc = _run([os.path.join(SCRIPTS, "estimate_inventory.py"), str(tmp_path),
                     "--reported", str(report)])
        assert proc.returncode == 1, "an incomplete report was reported as fully covered"
        with open(os.path.join(str(tmp_path), "coverage_diff.json"), encoding="utf-8") as f:
            payload = json.load(f)
        omitted = {round(o["estimate"], 2) for o in payload["omitted"]}
        assert 1.05 in omitted and 0.88 in omitted
        assert 3.11 in omitted and 0.74 in omitted, \
            "estimates that only exist in the stats JSON were not counted as executed"
        assert all(round(r["estimate"], 2) != 1.05 for r in payload["reported"])

    def test_full_coverage_exits_zero(self, executed_outputs, tmp_path):
        report = tmp_path / "report.md"
        report.write_text(
            "ORs were 2.22, 1.05, 0.88 in the primary model; sensitivity 3.11 and 0.74.\n",
            encoding="utf-8")
        proc = _run([os.path.join(SCRIPTS, "estimate_inventory.py"), str(tmp_path),
                     "--reported", str(report)])
        assert proc.returncode == 0, proc.stderr[-2000:]

    def test_invented_number_is_flagged_as_unsourced(self, executed_outputs, tmp_path):
        report = tmp_path / "report.md"
        report.write_text(
            "The odds ratio was 2.22 (95% CI 1.29-3.82).\n"
            "A second analysis gave 9.87, which no output contains.\n",
            encoding="utf-8")
        proc = _run([os.path.join(SCRIPTS, "check_numbers_sourced.py"), str(report),
                     "--output-dir", str(tmp_path)])
        assert proc.returncode == 1
        with open(os.path.join(str(tmp_path), "numbers_sourced_report.json"),
                  encoding="utf-8") as f:
            payload = json.load(f)
        unsourced = {e["text"] for e in payload["unsourced"]}
        assert "9.87" in unsourced
        assert "2.22" not in unsourced

    def test_structural_numbers_are_excluded_with_a_recorded_reason(self, executed_outputs,
                                                                    tmp_path):
        report = tmp_path / "report.md"
        report.write_text(
            "## 2. Results\n"
            "See Table 3 and Figure 2. The OR was 2.22 (95% CI 1.29-3.82).\n",
            encoding="utf-8")
        proc = _run([os.path.join(SCRIPTS, "check_numbers_sourced.py"), str(report),
                     "--output-dir", str(tmp_path)])
        assert proc.returncode == 0, proc.stderr[-2000:]
        with open(os.path.join(str(tmp_path), "numbers_sourced_report.json"),
                  encoding="utf-8") as f:
            payload = json.load(f)
        reasons = {e["reason"] for e in payload["structural_ignored"]}
        assert reasons, "no exclusion reasons were recorded"
        assert all(isinstance(r, str) and r for r in reasons)


# ─────────────────────────────────────────────────────────────────────────────
# Fail-closed behaviour: missing evidence must exit 1, not 0
# ─────────────────────────────────────────────────────────────────────────────

class PrecheckScenario:
    """Fixture builders shared by the precheck test classes.

    Deliberately not named `Test*`: pytest collects test classes by name, and a
    test class that inherits from another re-runs every one of its parent's
    tests. Keeping the helpers here means each class below runs its own
    assertions once.
    """

    def _scenario(self, tmp_path, criteria=None, groups=None, extra_steps=None,
                  data=None, covariates=None, analyses=None):
        """A post-stage scenario with fabricated but well-formed outputs."""
        (data if data is not None else make_harmful_exposure(n=300)).to_csv(
            tmp_path / "data.csv", index=False)
        pipeline = [
            {"id": "flowchart", "skill": "flowchart", "status": "done",
             "depends_on": [], "output_file": "Flowchart"},
            {"id": "missingness", "skill": "missing-data", "status": "done",
             "depends_on": []},
            {"id": "uni_primary", "skill": "univariate-forest", "status": "done",
             "depends_on": [], "outcome": "primary", "output_file": "Uni_primary",
             "success_criteria": criteria if criteria is not None else
             [{"check": "events_per_group >= 1", "severity": "blocking"}]},
        ]
        pipeline.extend(extra_steps or [])
        plan = _base_plan(
            tmp_path, tmp_path / "data.csv",
            {"id": "primary", "name": "Death", "type": "binary", "event_column": "dead"},
            covariates if covariates is not None else
            [{"name": "Age", "column": "age", "type": "continuous"},
             {"name": "Male", "column": "sex", "type": "binary", "reference": 0}],
            pipeline,
            grouping=groups,
        )
        if analyses:
            plan["analyses"] = analyses
        plan_path = _write_plan(tmp_path, plan)
        _generate(plan_path)
        cohort_n = len(data) if data is not None else 300
        events = int(data["dead"].sum()) if data is not None else 95
        self._write_flowchart(os.path.join(str(tmp_path), "Flowchart.xlsx"), cohort_n)
        self._write_uni(os.path.join(str(tmp_path), "Uni_primary.xlsx"),
                        n=cohort_n, events=events)
        self._write_uni(os.path.join(str(tmp_path), "Missingness.xlsx"),
                        n=cohort_n, events=events)
        return plan_path

    @staticmethod
    def _write_flowchart(path, final_n):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Patient Selection"
        ws.append(["Patient Selection Flowchart Summary"])
        ws.append([f"Initial N = {final_n} → Final N = {final_n}"])
        ws.append(["Step", "Description", "N Excluded", "N Remaining"])
        ws.append(["Initial", "Repro", "—", str(final_n)])
        ws.append(["Final", "Final study cohort", "—", str(final_n)])
        wb.save(path)

    @staticmethod
    def _write_uni(path, effect="2.22 (1.29–3.82)", n=300, events=95):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Univariate Analysis"
        ws.append(["Variable", "N (events)", "OR (95% CI)", "p-value"])
        ws.append(["Age", f"{n} ({events})", effect, "0.004*"])
        wb.save(path)

    @staticmethod
    def _write_multivariable(path, rows, n=300, events=95):
        """The layout multivariate_interaction.py / ph_stratified.py emit: a
        log-scale coefficient beside the formatted ratio it exponentiates to.
        """
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Multivariable"
        ws.append(["Multivariable model"])
        ws.append([f"n = {n}, events = {events}"])
        ws.append([])
        ws.append(["Variable", "Coefficient", "Adjusted HR (95% CI)", "p-value"])
        for variable, coef, effect, p in rows:
            ws.append([variable, coef, effect, p])
        ws.append(["Total", f"{n} ({events})", "", ""])
        wb.save(path)

    @staticmethod
    def _write_discrimination(path, rows, metric="C-index"):
        """The discrimination table auc_cstat.py:780 writes, headers verbatim."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{metric} Analysis"
        ws.append([f"{metric} / discrimination"])
        ws.append([])
        ws.append(["Variable", "N", "Events",
                   f"{metric}\n(variable alone)\n[95% CI]",
                   f"{metric}\n(clinical model)\n[95% CI]",
                   f"{metric}\n(clinical model\n+ variable)\n[95% CI]",
                   f"Delta {metric}\n[95% CI]",
                   "p-value"])
        for row in rows:
            ws.append(list(row))
        wb.save(path)

    @staticmethod
    def _write_psm(path, exposed_label, control_label, effect="2.24 (1.30–3.85)",
                   exposed_counts=(120, 60), control_counts=(120, 35)):
        """A PSM outcome table in the layout psm_matching.py emits."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Outcome Analysis"
        ws.append(["Propensity Score Matched Analysis"])
        ws.append(["Outcome", f"{exposed_label}\nN (events)",
                   f"{control_label}\nN (events)", "OR (95% CI)", "p-value"])
        ws.append(["Death",
                   f"{exposed_counts[0]} ({exposed_counts[1]})",
                   f"{control_counts[0]} ({control_counts[1]})", effect, "<0.001*"])
        wb.save(path)


class TestFailsClosedWhenEvidenceIsMissing(PrecheckScenario):
    """Every case here is a way the check could end up asserting nothing — no
    manifest, no artifact, an unreadable flowchart, a criterion that cannot be
    evaluated — and every one of them used to exit 0. "I could not tell" and
    "it is fine" sharing an exit code is the one failure mode that makes a
    checker worse than no checker, because it is the one nobody notices.
    """

    def test_clean_scenario_is_the_control(self, tmp_path):
        proc = _run([PRECHECK, self._scenario(tmp_path), "--stage", "post"])
        assert proc.returncode == 0, proc.stderr[-3000:]

    def test_missing_config_manifest_blocks(self, tmp_path):
        plan_path = self._scenario(tmp_path)
        os.remove(os.path.join(str(tmp_path), "configs", "manifest.json"))
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert proc.returncode == 1
        assert "[QC_BLOCKING] COL-001" in proc.stderr

    def test_empty_config_manifest_blocks(self, tmp_path):
        plan_path = self._scenario(tmp_path)
        with open(os.path.join(str(tmp_path), "configs", "manifest.json"), "w",
                  encoding="utf-8") as f:
            json.dump({"configs": []}, f)
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert proc.returncode == 1
        assert "zero configs" in proc.stderr

    def test_manifest_that_misses_an_enabled_step_blocks(self, tmp_path):
        plan_path = self._scenario(tmp_path)
        man_path = os.path.join(str(tmp_path), "configs", "manifest.json")
        with open(man_path, encoding="utf-8") as f:
            man = json.load(f)
        man["configs"] = [e for e in man["configs"] if e["step_id"] != "uni_primary"]
        with open(man_path, "w", encoding="utf-8") as f:
            json.dump(man, f)
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert proc.returncode == 1
        assert "uni_primary" in proc.stderr and "not audited" in proc.stderr

    def test_unapplicable_exclusion_blocks_and_stops_the_cohort_checks(self, tmp_path):
        """A cohort that could not be built must not then be measured."""
        plan_path = self._scenario(tmp_path)
        with open(plan_path, encoding="utf-8") as f:
            plan = json.load(f)
        plan["population"]["exclusion_steps"] = [
            {"description": "Age under 20", "condition": "agee < 20"}]
        with open(plan_path, "w", encoding="utf-8") as f:
            json.dump(plan, f)
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert proc.returncode == 1
        assert "[QC_BLOCKING] COL-005" in proc.stderr
        with open(os.path.join(str(tmp_path), "precheck_report.json"),
                  encoding="utf-8") as f:
            payload = json.load(f)
        skipped = {s["check_id"] for s in payload["skipped_checks"]}
        assert {"RES-006", "REFLECT", "COL-005"} <= skipped, \
            "cohort-dependent checks ran on a cohort already known to be wrong"

    def test_missing_declared_output_blocks(self, tmp_path):
        plan_path = self._scenario(tmp_path)
        os.remove(os.path.join(str(tmp_path), "Uni_primary.xlsx"))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "[QC_BLOCKING] RES-000 uni_primary" in proc.stderr

    def test_blocking_criterion_that_cannot_be_evaluated_blocks_at_post(self, tmp_path):
        plan_path = self._scenario(tmp_path, criteria=[
            {"check": "smd_max < 0.1", "severity": "blocking"}])
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "[QC_BLOCKING] RES-004" in proc.stderr

    def test_the_same_criterion_is_only_informational_before_the_run(self, tmp_path):
        plan_path = self._scenario(tmp_path, criteria=[
            {"check": "smd_max < 0.1", "severity": "blocking"}])
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert proc.returncode == 0, "a not-yet-computable metric blocked the pre stage"

    def test_compound_criterion_is_unparsable_even_when_metrics_exist(self):
        """Validity is decided before availability, otherwise a malformed
        criterion is excused as "metric unavailable" — which sounds temporary
        and is permanent."""
        outcome, _ = evaluate_criterion("smd_max < 0.1 and matched_pairs >= 50", {})
        assert outcome == EVAL_UNPARSABLE
        outcome, _ = evaluate_criterion("smd_max < 0.1 and matched_pairs >= 50",
                                        {"smd_max": 0.02, "matched_pairs": 100})
        assert outcome == EVAL_UNPARSABLE

    def test_invalid_dsl_blocks(self, tmp_path):
        plan_path = self._scenario(tmp_path, criteria=[
            {"check": "smd_max < 0.1 and matched_pairs >= 50", "severity": "blocking"}])
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "not valid DSL" in proc.stderr

    def test_empty_criteria_warns(self, tmp_path):
        plan_path = self._scenario(tmp_path, criteria=[])
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "[QC_WARNING] RES-004 uni_primary" in proc.stderr
        assert "no success_criteria" in proc.stderr

    def test_unreadable_flowchart_blocks(self, tmp_path):
        """Losing the flowchart must not be a way to skip the N-alignment check."""
        import openpyxl
        plan_path = self._scenario(tmp_path)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Patient Selection Flowchart Summary"])
        ws.append(["Final cohort (n=438)"])
        wb.save(os.path.join(str(tmp_path), "Flowchart.xlsx"))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "[QC_BLOCKING] RES-002" in proc.stderr

    def test_infinite_estimate_is_parsed_and_blocks(self, tmp_path):
        plan_path = self._scenario(tmp_path)
        self._write_uni(os.path.join(str(tmp_path), "Uni_primary.xlsx"),
                        effect="Inf (0.00–Inf)")
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "non-finite estimate" in proc.stderr

    def test_split_estimate_columns_are_read(self, tmp_path):
        """An OR/CI split across columns is still an OR that has to be plausible."""
        import openpyxl
        plan_path = self._scenario(tmp_path)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Univariate Analysis"
        ws.append(["Variable", "N (events)", "OR", "95% CI Lower", "95% CI Upper",
                   "p-value"])
        ws.append(["Exposure", "300 (95)", 1945429.0, 1.2e-05, 3.66e58, 0.97])
        wb.save(os.path.join(str(tmp_path), "Uni_primary.xlsx"))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "outside [0.01, 100]" in proc.stderr

    def test_log_scale_coefficient_column_does_not_block(self, tmp_path):
        """The layout every multivariable table uses: a log-scale `Coefficient`
        column beside the formatted HR. On that scale a protective covariate is
        negative and a near-null one is ~0.002, so the ratio bound turned a
        correct model into three blocking findings.
        """
        plan_path = self._scenario(tmp_path)
        self._write_multivariable(
            os.path.join(str(tmp_path), "Uni_primary.xlsx"),
            [("Age (per year)", 0.0021, "1.00 (1.00–1.01)", "0.31"),
             ("Male sex", -0.3567, "0.70 (0.55–0.89)", "0.004*"),
             ("Stage III/IV", 1.2528, "3.50 (2.40–5.10)", "<0.001*")])
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "[QC_BLOCKING] RES-001" not in proc.stderr, proc.stderr
        assert proc.returncode == 0, proc.stderr

    def test_log_scale_coefficient_still_blocks_when_it_blows_up(self, tmp_path):
        """The bound moves to the log scale; it does not go away. A coefficient of
        27.4 is a separation artefact (HR ~ 8e11) and has to block, exactly as the
        same fit blocks when printed as a ratio.
        """
        plan_path = self._scenario(tmp_path)
        self._write_multivariable(
            os.path.join(str(tmp_path), "Uni_primary.xlsx"),
            [("Rare marker", 27.4, "1.50 (1.10–2.00)", "0.99")])
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "log-scale coefficient" in proc.stderr
        assert "outside [0.01, 100]" in proc.stderr

    def test_log_scale_estimate_outside_its_own_ci_still_blocks(self, tmp_path):
        """Containment is scale-free, so dropping the ratio bound must not drop it."""
        import openpyxl
        plan_path = self._scenario(tmp_path)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Univariate Analysis"
        ws.append(["Variable", "N (events)", "Coefficient", "95% CI Lower",
                   "95% CI Upper", "p-value"])
        ws.append(["Male sex", "300 (95)", -0.36, 0.10, 0.60, 0.004])
        wb.save(os.path.join(str(tmp_path), "Uni_primary.xlsx"))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "outside its own CI" in proc.stderr

    def test_a_negative_delta_c_index_does_not_block(self, tmp_path):
        """Adding a variable that COSTS discrimination gives a negative delta.
        That is a result, not a defect. It used to block as a non-positive ratio,
        because `Delta C-index [95% CI]` matched no header and the formatted cell
        fell through to the ratio default.
        """
        plan_path = self._scenario(tmp_path)
        self._write_discrimination(
            os.path.join(str(tmp_path), "Uni_primary.xlsx"),
            [("Biomarker", 300, 95, "0.681 (0.63–0.73)", "0.742 (0.70–0.79)",
              "0.721 (0.67–0.77)", "-0.021 (-0.045–0.003)", "0.31"),
             ("Stage", 300, 95, "0.700 (0.65–0.75)", "0.742 (0.70–0.79)",
              "0.788 (0.74–0.83)", "+0.046 (0.012–0.080)", "0.008")])
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "[QC_BLOCKING] RES-001" not in proc.stderr, proc.stderr
        assert proc.returncode == 0, proc.stderr

    def test_an_impossible_c_index_still_blocks(self, tmp_path):
        """The bound moves onto the unit interval; it does not go away."""
        plan_path = self._scenario(tmp_path)
        self._write_discrimination(
            os.path.join(str(tmp_path), "Uni_primary.xlsx"),
            [("Broken", 300, 95, "1.400 (1.20–1.60)", "0.742 (0.70–0.79)",
              "0.721 (0.67–0.77)", "+0.010 (0.001–0.020)", "0.44")])
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "discrimination metric" in proc.stderr
        assert "[0, 1]" in proc.stderr

    def test_a_discrimination_header_is_not_reported_as_an_unparsable_estimate(self, tmp_path):
        """`C-index (variable alone) [95% CI]` carries digits and is not an exact
        key. Asking "is this a header?" against a literal tuple rather than the
        resolver reports the header of a column as an unreadable estimate.
        """
        plan_path = self._scenario(tmp_path)
        self._write_discrimination(
            os.path.join(str(tmp_path), "Uni_primary.xlsx"),
            [("Biomarker", 300, 95, "0.681 (0.63–0.73)", "0.742 (0.70–0.79)",
              "0.721 (0.67–0.77)", "+0.010 (0.001–0.020)", "0.31")])
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "could not parse" not in proc.stderr, proc.stderr

    def test_an_effect_labelled_inside_the_cell_is_not_a_missing_effect(self, tmp_path):
        """weighting.py writes `Hazard ratio: 1.04 (0.83–1.30)` into one cell, and
        its header row has a single recognised name so no column roles are
        assigned either. Every weighting run therefore parsed zero estimates —
        and `weighting` is in _ESTIMATE_PRODUCING_SKILLS, so zero is a blocker.
        """
        import openpyxl
        plan_path = self._scenario(tmp_path)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Outcome_model"
        ws.append(["Weighted outcome model"])
        ws.append(["Analysis", "Estimand", "Effect (95% CI)", "p", "SE type", "Note"])
        ws.append(["Crude", "—", "Hazard ratio: 1.036 (0.827–1.297)", "0.76",
                   "model", ""])
        ws.append(["IPTW", "ATE", "Hazard ratio: 1.077 (0.807–1.438)", "0.61",
                   "robust", ""])
        wb.save(os.path.join(str(tmp_path), "Uni_primary.xlsx"))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "no effect estimate could be read" not in proc.stderr, proc.stderr
        assert "[QC_BLOCKING] RES-001" not in proc.stderr, proc.stderr

    def test_unparsable_estimate_shape_blocks_rather_than_disappearing(self, tmp_path):
        plan_path = self._scenario(tmp_path)
        self._write_uni(os.path.join(str(tmp_path), "Uni_primary.xlsx"),
                        effect="2.22 (1.29 to 3.82]")
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert ("could not parse" in proc.stderr
                or "no effect estimate could be read" in proc.stderr)

    def test_psm_direction_is_checked_from_the_workbook_when_no_csv_exists(self, tmp_path):
        """The pipeline writes PSM.xlsx. Reading only PSM.csv meant this check —
        the one guarding the reciprocal-OR defect — never ran on a normal run."""
        psm_step = {"id": "psm", "skill": "psm-matching", "status": "done",
                    "depends_on": [], "output_file": "PSM",
                    "success_criteria": [{"check": "matched_pairs > 0",
                                          "severity": "blocking"}]}
        plan_path = self._scenario(
            tmp_path, extra_steps=[psm_step],
            groups={"column": "exposed",
                    "groups": [{"value": 0.0, "label": "Unexposed"},
                               {"value": 1.0, "label": "Exposed"}]})
        self._write_psm(os.path.join(str(tmp_path), "PSM.xlsx"), "Exposed", "Unexposed")
        assert not os.path.exists(os.path.join(str(tmp_path), "PSM.csv"))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "[QC_BLOCKING] PSM-DIR" in proc.stderr, \
            "the workbook-only PSM output was not checked for direction"
        assert proc.returncode == 1

    def test_psm_step_without_any_artifact_blocks(self, tmp_path):
        psm_step = {"id": "psm", "skill": "psm-matching", "status": "done",
                    "depends_on": [], "output_file": "PSM"}
        plan_path = self._scenario(tmp_path, extra_steps=[psm_step])
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "PSM-DIR" in proc.stderr or "RES-000 psm" in proc.stderr


class TestIncrementalValueIsReachableFromAPlan(PrecheckScenario):
    """The whole incremental-value path, from plan declaration to graded output.

    gen_auc hard-coded `clinical_model.enabled = False` with no plan field to
    turn it on, so a project could not ask "does this marker add anything to the
    standard model" without hand-editing a generated config — and the
    signed_unit_interval rule that grades the Delta this question produces was
    only reachable from a workbook nobody's pipeline wrote.

    The two ends are asserted together on purpose. Unit tests of the generator
    say the config is shaped right; only a real auc_cstat run says the shape is
    the one that script reads, and only precheck reading the resulting workbook
    says the numbers land inside the bounds their scale actually has.
    """

    AUC_STEP = {"id": "auc", "skill": "auc-cstat", "status": "done",
                "depends_on": [], "outcome": "primary", "output_file": "AUC"}

    COVARIATES = [{"name": "Age", "column": "age", "type": "continuous"},
                  {"name": "Male", "column": "sex", "type": "binary", "reference": 0},
                  {"name": "Exposed", "column": "exposed", "type": "binary",
                   "reference": 0}]

    def _run_auc(self, tmp_path, analyses):
        plan_path = self._scenario(
            tmp_path, extra_steps=[self.AUC_STEP], covariates=self.COVARIATES,
            data=make_harmful_exposure(n=400, seed=3), analyses=analyses)
        cfg = os.path.join(str(tmp_path), "configs", "auc_config.json")
        assert os.path.exists(cfg), os.listdir(os.path.join(str(tmp_path), "configs"))
        # Resolved from the runner's own registry rather than spelled again here,
        # so a script that moves takes this test with it instead of leaving it
        # green against a path nothing runs.
        import run_pipeline
        script = os.path.expanduser(run_pipeline.SKILL_SCRIPTS["auc-cstat"][0])
        proc = _run([script, cfg], cwd=str(tmp_path))
        return plan_path, cfg, proc

    def test_a_declared_base_model_produces_a_graded_delta(self, tmp_path):
        import openpyxl
        plan_path, cfg, proc = self._run_auc(
            tmp_path, {"auc": {"clinical_model_columns": ["age", "sex"]}})
        assert proc.returncode == 0, proc.stderr
        with open(cfg, encoding="utf-8") as f:
            written = json.load(f)
        assert written["clinical_model"]["enabled"] is True
        assert [v["column"] for v in written["clinical_model"]["variables"]] == ["age", "sex"]

        wb = openpyxl.load_workbook(os.path.join(str(tmp_path), "AUC.xlsx"))
        headers = [str(c.value) for ws in wb.worksheets
                   for row in ws.iter_rows() for c in row if c.value]
        assert any("Delta" in h for h in headers), \
            f"the incremental-value column never reached the workbook: {headers[:12]}"
        # The base model's own variables are not candidates: auc_cstat fits
        # `base + candidate`, so age against a base containing age is age twice.
        assert not any(h.startswith("Age") for h in headers), headers[:12]

        post = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "[QC_BLOCKING] RES-001 auc" not in post.stderr, post.stderr
        assert "no effect estimate could be read" not in post.stderr, post.stderr

    def test_every_estimate_in_the_table_gets_its_own_id(self, tmp_path):
        """One row, several numbers, one row label.

        The incremental-value table is the first output in this pipeline to put
        more than one estimate on a row — the candidate alone, the clinical
        model, the clinical model plus the candidate. estimate_id keyed on
        `stem:sheet:row:label` until it did, so all of them answered to the same
        handle: the omission listing printed one string beside three different
        values, and a join on that key could not tell which model each belonged
        to. Asserted on the real workbook, and on the ledger too, because the
        existing equality contract uses a fixture with one estimate per row and
        so cannot see this.
        """
        import estimate_inventory
        import run_pipeline
        _plan, _cfg, proc = self._run_auc(
            tmp_path, {"auc": {"clinical_model_columns": ["age", "sex"]}})
        assert proc.returncode == 0, proc.stderr

        path = os.path.join(str(tmp_path), "AUC.xlsx")
        records = estimate_inventory.scan_workbook(path)
        ids = [estimate_inventory.estimate_id(r, path, i)
               for i, r in enumerate(records)]
        rows = {(r["sheet"], r["row"]) for r in records}
        assert len(records) > len(rows), \
            "this table is supposed to hold several estimates per row"
        assert len(set(ids)) == len(ids), ids

        ledger, error = run_pipeline.step_estimate_inventory(
            {"id": "auc", "output_file": "AUC"}, str(tmp_path))
        assert error is None, error
        assert ledger["ids"] == ids

    def test_without_the_declaration_the_table_has_no_delta_column(self, tmp_path):
        """The default is unchanged: every covariate is a candidate, on its own."""
        import openpyxl
        _plan, cfg, proc = self._run_auc(tmp_path, None)
        assert proc.returncode == 0, proc.stderr
        with open(cfg, encoding="utf-8") as f:
            assert json.load(f)["clinical_model"] == {"enabled": False}
        wb = openpyxl.load_workbook(os.path.join(str(tmp_path), "AUC.xlsx"))
        headers = [str(c.value) for ws in wb.worksheets
                   for row in ws.iter_rows() for c in row if c.value]
        assert not any("Delta" in h for h in headers), headers[:12]


class TestAdjudicatedEpvAndPsmDirection(PrecheckScenario):
    """The two designs a third-party adjudication ruled on (03-fable-adjudication.md).

    EPV: the numerator is the whole analysis set's events, not the smallest
    exposure arm's. The min-arm version blocks a model with EPV 154 whenever the
    arms are unbalanced, which is most of clinical research. The sparse-arm
    concern is real but is a different quantity and is reported as such.

    PSM direction: the crude-OR reciprocal anchor cannot distinguish a labelling
    error from Simpson's paradox, so it is a warning. The blocking guard moved to
    a check with no confounding failure mode — whether the reported effect agrees
    with the counts printed in its own row.
    """

    def _psm_scenario(self, tmp_path, data, effect, exposed_counts, control_counts,
                      groups=None):
        psm_step = {"id": "psm", "skill": "psm-matching", "status": "done",
                    "depends_on": [], "output_file": "PSM"}
        plan_path = self._scenario(
            tmp_path, extra_steps=[psm_step], data=data, groups=groups,
            covariates=[{"name": "Age", "column": "age", "type": "continuous"}])
        self._write_psm(os.path.join(str(tmp_path), "PSM.xlsx"), "Exposed", "Unexposed",
                        effect=effect, exposed_counts=exposed_counts,
                        control_counts=control_counts)
        return plan_path

    # ── EPV ─────────────────────────────────────────────────────────────────

    def test_unbalanced_arms_with_ample_total_events_do_not_block(self, tmp_path):
        """8 events vs 300, two fitted parameters.

        Total EPV is 154. Blocking here would refuse an entirely estimable model
        for being unbalanced.
        """
        plan_path = self._scenario(
            tmp_path,
            data=make_unbalanced_events(),
            covariates=[{"name": "Age", "column": "age", "type": "continuous"}])
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert "[QC_BLOCKING] RES-006" not in proc.stderr, \
            "an unbalanced but well-powered cohort was blocked on EPV"
        with open(os.path.join(str(tmp_path), "precheck_report.json"),
                  encoding="utf-8") as f:
            payload = json.load(f)
        epv = payload["epv"]["primary"]
        assert epv["total_events"] == 308
        assert epv["min_group_events"] == 8
        assert epv["n_fitted_params"] == 2
        assert epv["epv"] > 100, f"EPV computed as {epv['epv']}, expected ~154"

    def test_the_sparse_arm_is_still_reported_under_its_own_name(self, tmp_path):
        """The concern does not disappear — it stops impersonating a fatal gate."""
        plan_path = self._scenario(
            tmp_path,
            data=make_unbalanced_events(),
            covariates=[{"name": "Age", "column": "age", "type": "continuous"}])
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert "[QC_WARNING] RES-006b" in proc.stderr
        assert "sparse-data bias" in proc.stderr
        assert "8 event(s)" in proc.stderr

    def test_genuinely_low_total_epv_still_blocks(self, tmp_path):
        """12 events against 15 parameters — EPV 0.8, and still fatal."""
        plan_path = self._scenario(
            tmp_path,
            data=make_sparse_total_events(),
            covariates=[{"name": f"C{i}", "column": f"c{i}", "type": "binary",
                         "reference": 0} for i in range(14)])
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert proc.returncode == 1
        assert "[QC_BLOCKING] RES-006" in proc.stderr
        with open(os.path.join(str(tmp_path), "precheck_report.json"),
                  encoding="utf-8") as f:
            epv = json.load(f)["epv"]["primary"]
        assert epv["effective_events"] == 12
        assert epv["n_fitted_params"] == 15
        assert epv["epv"] < 1

    def test_logistic_numerator_uses_the_smaller_of_events_and_non_events(self):
        """Peduzzi's rule for a logistic model, not the raw event count."""
        data = make_unbalanced_events(n_exposed=100, e_exposed=90,
                                      n_control=900, e_control=880)
        assert data["dead"].sum() == 970
        assert len(data) - data["dead"].sum() == 30

    # ── unpriced terms: total is a lower bound, EPV an upper bound ──────────

    @staticmethod
    def _unpriced_setup(n_events=100, n_rows=400, n_unknown=2):
        """A plan whose categorical covariates have no determinable level count.

        No `label_map`, and the columns are absent from the dataframe, so each
        costs L-1 parameters for an unknown L >= 2. `parameter_budget` charges
        the floor of 1 and says so rather than guessing in either direction.
        """
        covs = [{"name": "Age", "column": "age", "type": "continuous"}]
        covs += [{"name": f"Cat{i}", "column": f"cat{i}", "type": "ordinal"}
                 for i in range(n_unknown)]
        plan = {"grouping": {"column": "exposed",
                             "groups": [{"value": 1.0}, {"value": 0.0}]},
                "covariates": covs,
                "outcomes": {"primary": {"id": "primary", "name": "D", "type": "binary",
                                         "event_column": "dead"}}}
        rng = np.random.default_rng(2)
        half = n_rows // 2
        dead = np.zeros(n_rows)
        dead[:n_events // 2] = 1.0
        dead[half:half + n_events // 2] = 1.0
        df = pd.DataFrame({"exposed": np.r_[np.ones(half), np.zeros(half)],
                           "age": rng.normal(65, 10, n_rows), "dead": dead})
        configs = {"multi_primary": {
            "grouping": plan["grouping"],
            "sections": [{"header": "x", "variables": [
                {"column": c["column"], "type": c["type"]} for c in covs]}]}}
        return plan, df, configs

    def test_unpriced_terms_are_listed_in_provenance(self):
        plan, df, configs = self._unpriced_setup()
        total, breakdown, prov = parameter_budget(plan, df, configs, "primary")
        assert total == 4, "each unpriceable term should be charged its floor of 1"
        unpriced = prov["unpriced_terms"]
        assert {u["column"] for u in unpriced} == {"cat0", "cat1"}
        assert all(u["reason"] == "categorical with undeclared levels" for u in unpriced)
        assert all(u["floor"] == 1 for u in unpriced)
        assert all(b.get("is_floor") for b in breakdown if b["term"] in ("cat0", "cat1"))

    def test_a_fully_priced_plan_reports_no_unpriced_terms(self):
        plan, df, configs = self._unpriced_setup(n_unknown=0)
        _total, _breakdown, prov = parameter_budget(plan, df, configs, "primary")
        assert prov["unpriced_terms"] == []

    def test_declaring_levels_prices_the_term_exactly(self):
        plan, df, configs = self._unpriced_setup(n_unknown=1)
        label_map = {"1": "I", "2": "II", "3": "III", "4": "IV"}
        plan["covariates"][1]["label_map"] = label_map
        configs["multi_primary"]["sections"][0]["variables"][1]["label_map"] = label_map
        total, _breakdown, prov = parameter_budget(plan, df, configs, "primary")
        assert prov["unpriced_terms"] == []
        assert total == 5, "1 exposure + 1 continuous + 3 for a four-level factor"

    def test_upper_bound_below_the_threshold_still_blocks(self):
        """If even the upper bound is under 5, the true EPV is under 5 too.

        Unpriced terms can only make the denominator larger, so this direction
        is safe and stays blocking.
        """
        import precheck_configs as PC

        plan, df, configs = self._unpriced_setup(n_events=12)
        report = PC.Report(quiet=True)
        results = PC.check_epv(report, plan, df, "pre", configs)
        assert results["primary"]["epv"] == 3.0
        assert results["primary"]["epv_is_upper_bound"] is True
        blocking = [i for i in report.issues
                    if i["check_id"] == "RES-006" and i["severity"] == "block"]
        assert len(blocking) == 1
        assert "upper bound on EPV" in blocking[0]["message"]

    def test_upper_bound_above_the_threshold_warns_instead_of_passing(self):
        """Not knowing whether a gate is passed is not the same as failing it —
        but it is not the same as passing it either, which is what a silent
        floor of 1 used to make it look like."""
        import precheck_configs as PC

        plan, df, configs = self._unpriced_setup(n_events=100)
        report = PC.Report(quiet=True)
        results = PC.check_epv(report, plan, df, "pre", configs)
        assert results["primary"]["epv"] == 25.0
        assert results["primary"]["n_fitted_params_is_lower_bound"] is True
        res006 = [i for i in report.issues if i["check_id"] == "RES-006"]
        assert [i["severity"] for i in res006] == ["warn"]
        assert "may be lower" in res006[0]["message"]
        assert res006[0]["evidence"]["unpriced_terms"]

    def test_a_fully_priced_plan_above_the_threshold_stays_silent(self):
        """The warning is about the unknown, not about the EPV value itself."""
        import precheck_configs as PC

        plan, df, configs = self._unpriced_setup(n_events=100, n_unknown=0)
        report = PC.Report(quiet=True)
        results = PC.check_epv(report, plan, df, "pre", configs)
        assert results["primary"]["epv_is_upper_bound"] is False
        assert not [i for i in report.issues
                    if i["check_id"] == "RES-006" and i["severity"] in ("block", "warn")]

    def test_parameter_budget_prefers_the_config_over_the_plan(self, tmp_path):
        """A hand-edited or interaction-bearing config changes the real df."""
        plan_path = self._scenario(tmp_path)
        configs_dir = os.path.join(str(tmp_path), "configs")
        with open(os.path.join(configs_dir, "manifest.json"), encoding="utf-8") as f:
            manifest = json.load(f)
        with open(plan_path, encoding="utf-8") as f:
            plan = json.load(f)

        multi_config = {
            "grouping": plan["grouping"],
            "sections": [{"header": "x", "variables": [
                {"column": "age", "type": "continuous"},
                {"column": "sex", "type": "binary"},
                {"column": "extra", "type": "continuous"},
            ]}],
        }
        path = os.path.join(configs_dir, "multi_primary_config.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(multi_config, f)
        manifest["configs"].append({"step_id": "multi_primary",
                                    "skill": "multivariate-interaction",
                                    "config_path": path})
        with open(os.path.join(configs_dir, "manifest.json"), "w", encoding="utf-8") as f:
            json.dump(manifest, f)

        configs = {"multi_primary": multi_config}
        total, _breakdown, prov = parameter_budget(plan, None, configs, "primary")
        assert prov["source"] == "config"
        assert total == 4, "config-declared terms were not counted"
        assert prov["plan_total"] == 3
        assert prov["drift"] == 1

    def test_parameter_budget_falls_back_to_the_plan_without_blocking(self, tmp_path):
        """Not being able to read the fitted model is a caveat, not a refusal."""
        plan_path = self._scenario(tmp_path)
        with open(plan_path, encoding="utf-8") as f:
            plan = json.load(f)
        total, _breakdown, prov = parameter_budget(plan, None, {}, "primary")
        assert prov["source"] == "plan"
        assert total == 3
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert proc.returncode == 0
        assert "approximated from plan.covariates" in proc.stderr

    # ── PSM direction ───────────────────────────────────────────────────────

    def test_simpson_reversal_is_a_warning_not_a_block(self, tmp_path):
        """Crude OR 2.0, correct matched OR 0.50.

        Adjustment reversing the sign of a confounded association is the whole
        point of adjusting. Blocking it teaches people to write overrides.
        """
        plan_path = self._psm_scenario(
            tmp_path, data=make_crude_or(2.0), effect="0.50 (0.35–0.72)",
            exposed_counts=(200, 40), control_counts=(200, 67))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "[QC_BLOCKING] PSM-DIR" not in proc.stderr, \
            "a legitimate confounding reversal was blocked as a labelling error"
        assert "[QC_WARNING] PSM-DIR" in proc.stderr
        assert "points the opposite way to the crude effect" in proc.stderr
        assert proc.returncode == 0, proc.stderr[-2500:]

    def test_row_that_contradicts_its_own_counts_blocks(self, tmp_path):
        """The counts say 2.41, the row reports 0.41 — its own reciprocal.

        Same matched cohort, same row, so no amount of confounding explains it.
        This is the blocking guard that replaces the crude anchor.
        """
        plan_path = self._psm_scenario(
            tmp_path, data=make_crude_or(2.4), effect="0.41 (0.28–0.60)",
            exposed_counts=(120, 60), control_counts=(120, 35))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "[QC_BLOCKING] PSM-DIR" in proc.stderr
        assert "reciprocal of the one implied by its own row's counts" in proc.stderr

    def test_a_self_consistent_row_is_accepted(self, tmp_path):
        """Counts imply 2.41 and the row reports 2.24 — non-collapsibility, not inversion."""
        plan_path = self._psm_scenario(
            tmp_path, data=make_crude_or(2.4), effect="2.24 (1.30–3.85)",
            exposed_counts=(120, 60), control_counts=(120, 35))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "[QC_BLOCKING] PSM-DIR" not in proc.stderr
        assert "row is self-consistent" in proc.stderr

    def test_missing_per_arm_counts_warn_that_the_strong_check_did_not_run(self, tmp_path):
        import openpyxl
        plan_path = self._psm_scenario(
            tmp_path, data=make_crude_or(2.4), effect="2.24 (1.30–3.85)",
            exposed_counts=(120, 60), control_counts=(120, 35))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Outcome Analysis"
        ws.append(["Outcome", "Exposed\nsummary", "Unexposed\nsummary", "OR (95% CI)"])
        ws.append(["Death", "n/a", "n/a", "2.24 (1.30–3.85)"])
        wb.save(os.path.join(str(tmp_path), "PSM.xlsx"))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "could not be checked against itself" in proc.stderr

    def test_column_order_assert_is_still_blocking(self, tmp_path):
        """Kept blocking on adjudication: psm_matching.py hardcodes groups[0]-first
        and carries no separate contrast metadata, so the column order *is* the
        contrast declaration in this pipeline."""
        plan_path = self._psm_scenario(
            tmp_path, data=make_crude_or(2.4), effect="2.24 (1.30–3.85)",
            exposed_counts=(120, 60), control_counts=(120, 35),
            groups={"column": "exposed",
                    "groups": [{"value": 0.0, "label": "Unexposed"},
                               {"value": 1.0, "label": "Exposed"}]})
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "puts 'Exposed' before 'Unexposed'" in proc.stderr

    # ── producer sidecar ────────────────────────────────────────────────────

    def _write_sidecar(self, tmp_path, exposed_value, exposed=(200, 40),
                       control=(200, 67), table="PSM.xlsx"):
        """A sidecar carrying the binding fields, as the producer now writes them."""
        import hashlib
        artifact = os.path.join(str(tmp_path), table)
        digest = hashlib.sha256()
        with open(artifact, "rb") as f:
            digest.update(f.read())
        payload = {"sidecar_version": 1,
                   "table_basename": table,
                   "table_sha256": digest.hexdigest(),
                   "config_sha256": None,
                   "exposed_label": "Exposed", "exposed_value": exposed_value,
                   "exposed": {"n": exposed[0], "events": exposed[1]},
                   "control": {"n": control[0], "events": control[1]}}
        stem = os.path.splitext(table)[0]
        with open(os.path.join(str(tmp_path), f"{stem}_contrast.json"), "w",
                  encoding="utf-8") as f:
            json.dump(payload, f)

    def test_absent_sidecar_warns_but_does_not_block(self, tmp_path):
        plan_path = self._psm_scenario(
            tmp_path, data=make_crude_or(2.0), effect="0.50 (0.35–0.72)",
            exposed_counts=(200, 40), control_counts=(200, 67))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "no PSM_contrast.json" in proc.stderr
        assert proc.returncode == 0

    def test_sidecar_naming_the_wrong_arm_blocks(self, tmp_path):
        plan_path = self._psm_scenario(
            tmp_path, data=make_crude_or(2.0), effect="0.50 (0.35–0.72)",
            exposed_counts=(200, 40), control_counts=(200, 67))
        self._write_sidecar(tmp_path, exposed_value=0.0)
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "fitted a different contrast" in proc.stderr

    def test_sidecar_counts_that_disagree_with_the_table_block(self, tmp_path):
        plan_path = self._psm_scenario(
            tmp_path, data=make_crude_or(2.0), effect="0.50 (0.35–0.72)",
            exposed_counts=(200, 40), control_counts=(200, 67))
        self._write_sidecar(tmp_path, exposed_value=1.0, exposed=(200, 67),
                            control=(200, 40))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "is not the arm the table displays" in proc.stderr

    def test_agreeing_sidecar_is_accepted(self, tmp_path):
        plan_path = self._psm_scenario(
            tmp_path, data=make_crude_or(2.0), effect="0.50 (0.35–0.72)",
            exposed_counts=(200, 40), control_counts=(200, 67))
        self._write_sidecar(tmp_path, exposed_value=1.0)
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "producer sidecar agrees with the table" in proc.stderr
        assert proc.returncode == 0


@slow
class TestPsmContrastSidecarProducer(PrecheckScenario):
    """The producer half of the sidecar contract, run for real.

    `precheck_configs.py` blocks when PSM_contrast.json disagrees with the output
    table. That check is only worth something if the sidecar is derived
    independently of the table — one built from the same dict that fills the
    table agrees with it unconditionally and can never catch anything. The last
    test here is what proves it: the table-building path is corrupted while the
    fitting path is left alone, and the sidecar has to keep telling the truth.
    """

    def _psm_plan(self, tmp_path):
        make_harmful_exposure(n=600, seed=5).to_csv(tmp_path / "data.csv", index=False)
        plan = _base_plan(
            tmp_path, tmp_path / "data.csv",
            {"id": "primary", "name": "Death", "type": "binary", "event_column": "dead"},
            [{"name": "Age", "column": "age", "type": "continuous"},
             {"name": "Male", "column": "sex", "type": "binary", "reference": 0}],
            [{"id": "missingness", "skill": "missing-data", "status": "done",
              "depends_on": []},
             {"id": "psm", "skill": "psm-matching", "status": "done",
              "depends_on": [], "output_file": "PSM"}],
        )
        plan_path = _write_plan(tmp_path, plan)
        _generate(plan_path)
        # Stand-in for the missing-data artifact so RES-000 does not add noise.
        self._write_uni(os.path.join(str(tmp_path), "Missingness.xlsx"), n=600, events=200)
        return plan_path

    @pytest.fixture(scope="class")
    def produced(self, tmp_path_factory):
        tmp = tmp_path_factory.mktemp("sidecar")
        plan_path = self._psm_plan(tmp)
        proc = _run([PSM_SCRIPT, os.path.join(str(tmp), "configs", "psm_config.json")],
                    cwd=str(tmp))
        assert proc.returncode == 0, proc.stdout[-2000:] + proc.stderr[-2000:]
        return tmp, plan_path

    @staticmethod
    def _repoint(src, dst):
        """Copy a finished run and rewrite the absolute paths embedded in it."""
        import shutil
        shutil.copytree(str(src), str(dst))
        targets = [os.path.join(str(dst), "research_plan.json")]
        cfg_dir = os.path.join(str(dst), "configs")
        targets += [os.path.join(cfg_dir, f) for f in os.listdir(cfg_dir)]
        for path in targets:
            with open(path, encoding="utf-8") as f:
                text = f.read()           # read fully before opening for write
            with open(path, "w", encoding="utf-8") as f:
                f.write(text.replace(str(src), str(dst)))
        return dst

    def test_sidecar_is_written_beside_the_workbook(self, produced):
        tmp, _ = produced
        assert os.path.exists(os.path.join(str(tmp), "PSM_contrast.json")), \
            "the producer did not write the sidecar where the consumer looks for it"

    def test_sidecar_matches_the_cohort_the_model_was_fitted_on(self, produced):
        tmp, _ = produced
        with open(os.path.join(str(tmp), "PSM_contrast.json"), encoding="utf-8") as f:
            side = json.load(f)
        table = pd.read_csv(os.path.join(str(tmp), "PSM.csv")).iloc[0]
        assert side["exposed"]["n"] == int(table["Exposed N"])
        assert side["exposed"]["events"] == int(table["Exposed Events"])
        assert side["control"]["n"] == int(table["Unexposed N"])
        assert side["control"]["events"] == int(table["Unexposed Events"])

    def test_exposed_value_is_the_raw_data_value_not_the_label(self, produced):
        """The consumer blocks on `exposed_value != groups[0].value`, so the
        sidecar has to carry the coded value, read back out of the data."""
        tmp, _ = produced
        with open(os.path.join(str(tmp), "PSM_contrast.json"), encoding="utf-8") as f:
            side = json.load(f)
        assert side["exposed_value"] == 1.0
        assert side["exposed_label"] == "Exposed"
        assert side["control_value"] == 0.0

    def test_consumer_goes_from_warning_to_pass(self, produced):
        tmp, plan_path = produced
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "no PSM_contrast.json" not in proc.stderr
        assert "producer sidecar agrees with the table" in proc.stderr
        assert "[QC_BLOCKING] PSM-DIR" not in proc.stderr
        assert proc.returncode == 0, proc.stderr[-2500:]

    def test_sidecar_naming_the_other_arm_blocks(self, produced, tmp_path):
        tmp, _ = produced
        work = self._repoint(tmp, tmp_path / "wrong_value")
        path = os.path.join(str(work), "PSM_contrast.json")
        with open(path, encoding="utf-8") as f:
            side = json.load(f)
        side["exposed_value"] = 0.0
        with open(path, "w", encoding="utf-8") as f:
            json.dump(side, f)
        proc = _run([PRECHECK, os.path.join(str(work), "research_plan.json"),
                     "--stage", "post"])
        assert proc.returncode == 1
        assert "fitted a different contrast" in proc.stderr

    def test_sidecar_counts_that_disagree_with_the_table_block(self, produced, tmp_path):
        tmp, _ = produced
        work = self._repoint(tmp, tmp_path / "wrong_counts")
        path = os.path.join(str(work), "PSM_contrast.json")
        with open(path, encoding="utf-8") as f:
            side = json.load(f)
        side["exposed"], side["control"] = side["control"], side["exposed"]
        with open(path, "w", encoding="utf-8") as f:
            json.dump(side, f)
        proc = _run([PRECHECK, os.path.join(str(work), "research_plan.json"),
                     "--stage", "post"])
        assert proc.returncode == 1
        assert "is not the arm the table displays" in proc.stderr

    def test_sidecar_is_independent_of_the_output_table(self, tmp_path):
        """Corrupt the table-building path only; the sidecar must not follow.

        `analyze_outcome_logistic` is patched to return the two arms swapped,
        which is the shape of the historical defect. The fitting path and
        `write_contrast_sidecar` are untouched. A sidecar derived from the
        outcome-result dict would swap too, agree with the table, and let this
        through — which is the whole thing this test exists to prevent.
        """
        sys.path.insert(0, os.path.join(SKILL_ROOT, "psm-matching", "scripts"))
        import psm_matching as P

        plan_path = self._psm_plan(tmp_path)
        config = P.load_config(os.path.join(str(tmp_path), "configs", "psm_config.json"))
        df_clean = P.prepare_analysis_df(P.load_data(config), config)
        grp1, grp2 = P.split_groups(df_clean, config)
        matched_df, match_info = P.perform_matching(
            P.estimate_propensity_scores(df_clean, config), config)
        balance = P.compute_balance_table(
            grp1, grp2, matched_df[matched_df["_treatment"] == 1],
            matched_df[matched_df["_treatment"] == 0], config)

        real = P.analyze_outcome_logistic

        def swapped(mdf, outcome, cfg):
            res = real(mdf, outcome, cfg)
            res["n_treatment"], res["n_control"] = res["n_control"], res["n_treatment"]
            res["events_treatment"], res["events_control"] = (
                res["events_control"], res["events_treatment"])
            return res

        P.analyze_outcome_logistic = swapped
        try:
            results = P.run_outcome_analyses(matched_df, config)
        finally:
            P.analyze_outcome_logistic = real
        P.write_excel(results, balance, match_info, config)
        P.write_contrast_sidecar(matched_df, config)

        with open(os.path.join(str(tmp_path), "PSM_contrast.json"), encoding="utf-8") as f:
            side = json.load(f)
        table = pd.read_csv(os.path.join(str(tmp_path), "PSM.csv")).iloc[0]
        assert side["exposed"]["events"] != int(table["Exposed Events"]), \
            "the sidecar tracked the corrupted table — it is not an independent source"

        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "is not the arm the table displays" in proc.stderr, \
            "the consumer did not catch a table that disagrees with the fitted cohort"


# ─────────────────────────────────────────────────────────────────────────────
# RES-015-det — the delegation seam to biostats-guardian's PH remedy gate
# ─────────────────────────────────────────────────────────────────────────────

PH_VIOLATED = {"status": "computed", "violated": True, "borderline": False,
               "min_p": 0.001, "global": {"p": 0.001},
               "effect_size": {"early": 1.9, "late": 0.7}}
PH_CLEAN = {"status": "computed", "violated": False, "borderline": False,
            "min_p": 0.42, "global": {"p": 0.42}}


class TestPhRemedyGateDelegation:
    """RES-015-det is delegated to biostats-guardian, not reimplemented here.

    These tests are the seam: the post stage must reproduce the gate's verdicts,
    and the two scope guards must hold — the gate runs at the manuscript, once,
    and only where a proportional-hazards assumption is actually being made.
    """

    def _cox_scenario(self, tmp_path, ph_artifact, outcome_extra=None,
                      decision_log=None, with_manuscript=True, cox=True):
        rng = np.random.default_rng(3)
        n = 400
        pd.DataFrame({
            "exposed": rng.integers(0, 2, n).astype(float),
            "age": rng.normal(65, 10, n),
            "time": rng.exponential(3.0, n).round(3),
            "event": rng.binomial(1, 0.45, n).astype(float),
            "dead": rng.binomial(1, 0.4, n).astype(float),
        }).to_csv(tmp_path / "data.csv", index=False)

        if cox:
            outcome = {"id": "primary", "name": "Death", "type": "time_to_event",
                       "time_column": "time", "event_column": "event",
                       "time_unit": "years", "time_scale": 1}
        else:
            outcome = {"id": "primary", "name": "Death", "type": "binary",
                       "event_column": "dead"}
        outcome.update(outcome_extra or {})

        pipeline = [
            {"id": "flowchart", "skill": "flowchart", "status": "done",
             "depends_on": [], "output_file": "Flowchart"},
            {"id": "missingness", "skill": "missing-data", "status": "done",
             "depends_on": []},
            {"id": "uni_primary", "skill": "univariate-forest", "status": "done",
             "depends_on": [], "outcome": "primary", "output_file": "Uni_primary"},
            {"id": "multi_primary", "skill": "multivariate-interaction", "status": "done",
             "depends_on": ["uni_primary"], "outcome": "primary",
             "output_file": "Multi_primary"},
        ]
        if with_manuscript:
            pipeline += [
                {"id": "reviewer_challenge", "skill": "reviewer-challenge",
                 "status": "done", "depends_on": ["multi_primary"]},
                {"id": "manuscript", "skill": "manuscript-results", "status": "done",
                 "depends_on": ["reviewer_challenge"]},
            ]
        plan = _base_plan(
            tmp_path, tmp_path / "data.csv", outcome,
            [{"name": "Age", "column": "age", "type": "continuous"}], pipeline)
        plan_path = _write_plan(tmp_path, plan)
        _generate(plan_path)

        builder = TestFailsClosedWhenEvidenceIsMissing
        builder._write_flowchart(os.path.join(str(tmp_path), "Flowchart.xlsx"), n)
        for name in ("Missingness.xlsx", "Uni_primary.xlsx", "Multi_primary.xlsx"):
            builder._write_uni(os.path.join(str(tmp_path), name),
                               effect="1.42 (1.10–1.83)", n=n, events=180)
        if ph_artifact is not None:
            with open(os.path.join(str(tmp_path), "Multi_primary_ph_test.json"), "w",
                      encoding="utf-8") as f:
                json.dump(ph_artifact, f)
        if decision_log is not None:
            with open(os.path.join(str(tmp_path), "decision_log.md"), "w",
                      encoding="utf-8") as f:
                f.write(decision_log)
        return plan_path

    def _res015(self, plan_path):
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        lines = [x for x in proc.stderr.splitlines() if "RES-015-det" in x]
        blocked = any(x.startswith("[QC_BLOCKING]") for x in lines)
        return blocked, lines, proc

    # ── the gate's own six cases, reproduced through the post stage ──────────

    def test_a_declared_stratify_with_no_artifact_blocks(self, tmp_path):
        plan_path = self._cox_scenario(
            tmp_path, dict(PH_VIOLATED, declared_remedy="stratify"))
        blocked, lines, _ = self._res015(plan_path)
        assert blocked, lines
        assert any("has no artifact" in x for x in lines)

    def test_b_disclose_without_a_decision_id_blocks(self, tmp_path):
        plan_path = self._cox_scenario(
            tmp_path, dict(PH_VIOLATED, declared_remedy="disclose"))
        blocked, lines, _ = self._res015(plan_path)
        assert blocked, lines
        assert any("ph_disclose_decision_id" in x for x in lines)

    def test_c_disclose_with_an_id_present_in_the_decision_log_passes(self, tmp_path):
        plan_path = self._cox_scenario(
            tmp_path, dict(PH_VIOLATED, declared_remedy="disclose"),
            outcome_extra={"ph_disclose_decision_id": "D-017"},
            decision_log="# Decision log\n\n## D-017 PH disclosure\n"
                         "PH violated for the primary outcome; disclosed in the "
                         "limitations rather than remediated.\n")
        blocked, lines, _ = self._res015(plan_path)
        assert not blocked, lines

    def test_d_decision_id_absent_from_the_log_blocks(self, tmp_path):
        plan_path = self._cox_scenario(
            tmp_path, dict(PH_VIOLATED, declared_remedy="disclose"),
            outcome_extra={"ph_disclose_decision_id": "D-999"},
            decision_log="# Decision log\n\n## D-017 something else\nUnrelated.\n")
        blocked, lines, _ = self._res015(plan_path)
        assert blocked, lines
        assert any("D-999" in x for x in lines)

    def test_e_violation_with_no_declared_remedy_blocks(self, tmp_path):
        plan_path = self._cox_scenario(tmp_path, dict(PH_VIOLATED))
        blocked, lines, _ = self._res015(plan_path)
        assert blocked, lines

    def test_f_no_violation_passes(self, tmp_path):
        plan_path = self._cox_scenario(tmp_path, PH_CLEAN)
        blocked, lines, _ = self._res015(plan_path)
        assert not blocked, lines

    # ── scope guards ────────────────────────────────────────────────────────

    def test_the_gate_is_a_manuscript_gate_not_a_per_step_one(self, tmp_path):
        """Same violated artifact, no manuscript step: nothing to guard, no block.

        The gate is explicitly designed not to block the multivariate step
        itself — that table may legitimately exist as an intermediate. What may
        not happen is it reaching a manuscript unaccompanied.
        """
        plan_path = self._cox_scenario(
            tmp_path, dict(PH_VIOLATED, declared_remedy="stratify"),
            with_manuscript=False)
        blocked, lines, proc = self._res015(plan_path)
        assert not blocked, lines
        assert "no manuscript step" in proc.stderr
        with open(os.path.join(str(tmp_path), "precheck_report.json"),
                  encoding="utf-8") as f:
            skipped = {s["check_id"] for s in json.load(f)["skipped_checks"]}
        assert "RES-015-det" in skipped, "the skip was not recorded"

    def test_a_logistic_plan_is_not_asked_for_ph_artifacts(self, tmp_path):
        """results-audit.md's producer table: only multivariate-interaction on a
        Cox model writes *_ph_test.json. Running the gate on a logistic plan
        would block it for missing a file nothing ever produces."""
        plan_path = self._cox_scenario(tmp_path, ph_artifact=None, cox=False)
        blocked, lines, proc = self._res015(plan_path)
        assert not blocked, lines
        assert "no proportional-hazards assumption" in proc.stderr

    def test_an_uninstalled_gate_blocks_rather_than_skipping(self, tmp_path):
        """Not auditing the remedies is not the same as them passing."""
        import precheck_configs as PC

        plan_path = self._cox_scenario(tmp_path, dict(PH_VIOLATED))
        with open(plan_path, encoding="utf-8") as f:
            plan = json.load(f)
        report = PC.Report(quiet=True)
        original = PC.PH_REMEDY_GATE
        PC.PH_REMEDY_GATE = os.path.join(str(tmp_path), "not_installed.py")
        try:
            PC.check_ph_remedy_gate(report, plan, plan_path, plan["pipeline"],
                                    plan["output_dir"])
        finally:
            PC.PH_REMEDY_GATE = original
        assert report.n_blocking == 1
        assert "not installed" in report.issues[0]["message"]

    def test_findings_reach_the_machine_readable_report(self, tmp_path):
        plan_path = self._cox_scenario(
            tmp_path, dict(PH_VIOLATED, declared_remedy="stratify"))
        self._res015(plan_path)
        gate_json = os.path.join(str(tmp_path), "guardian", "ph_remedy_gate.json")
        assert os.path.exists(gate_json), "the delegated gate's own JSON was not kept"
        with open(os.path.join(str(tmp_path), "precheck_report.json"),
                  encoding="utf-8") as f:
            issues = json.load(f)["issues"]
        assert any(i["check_id"] == "RES-015-det" and i["severity"] == "block"
                   for i in issues), "the delegated findings did not reach the report"


def make_non_ph_covariate(n=1200, seed=23):
    """Exposure is proportional; a three-level covariate is not.

    Truth: `stage` multiplies the hazard strongly before t=1 and not at all
    after, so it violates proportional hazards by construction, while `exposed`
    carries a constant 1.6x throughout and does not. That asymmetry is the whole
    point — it is the shape where stratifying on the offender is the right remedy
    and leaves the exposure estimate intact. `make_non_ph` is the opposite shape
    (the exposure itself is non-proportional) and stratifying there would delete
    the result, which is why ph_stratified.py refuses it.
    """
    rng = np.random.default_rng(seed)
    exposed = rng.binomial(1, 0.5, n)
    stage = rng.integers(0, 3, n)
    u = rng.uniform(size=n)
    hr_exp = np.where(exposed == 1, 1.6, 1.0)
    early = 0.25 * hr_exp * (1.0 + 1.2 * stage)
    late = 0.25 * hr_exp
    cut = 1.0
    s_at_cut = np.exp(-early * cut)
    t = np.where(u > s_at_cut,
                 -np.log(u) / early,
                 cut + (-np.log(u) - early * cut) / late)
    admin = 6.0
    return pd.DataFrame({"exposed": exposed.astype(float),
                         "stage": stage.astype(float),
                         "age": rng.normal(65, 10, n),
                         "time": np.minimum(t, admin),
                         "event": (t <= admin).astype(float)})


class TestPhRemedyEvidenceSidecar:
    """A declared `stratify` remedy now RUNS, and the gate checks what it did.

    Before this, `ph_violation_remedy: "stratify"` was a word in a JSON file:
    no generator shipped, so the analysis had to be hand-built, and the gate
    could only ask whether *a file* existed — name a step `stratified_os`, write
    anything at all, green light. Every test below that ends in a block is a case
    the old gate passed; they are written that way on purpose, because a check
    that cannot be shown to turn red is not evidence that it works.

    The pair that matters most is `test_the_gate_verifies_the_real_thing` and
    `test_stratifying_on_the_wrong_variable_is_caught`: same step, same workbook
    on disk, same file-existence answer — and opposite verdicts, because one of
    them stratified on the covariate that actually violated and the other did not.
    """

    STRATA_VIOLATION = {"status": "computed", "violated": True, "borderline": False,
                        "min_p": 0.0004, "alpha": 0.01,
                        "global": {"chi2": 18.2, "df": 2, "p": 0.0004},
                        "violating_variables": ["stage"],
                        "declared_remedy": "stratify"}

    def _gate(self, tmp_path, json_out="ph_gate.json"):
        import precheck_configs as PC
        out = os.path.join(str(tmp_path), json_out)
        proc = _run([PC.PH_REMEDY_GATE, "--plan",
                     os.path.join(str(tmp_path), "research_plan.json"),
                     "--output-dir", str(tmp_path), "--json", out])
        report = None
        if os.path.exists(out):
            with open(out, encoding="utf-8") as f:
                report = json.load(f)
        return proc, report

    def _scenario(self, tmp_path, remedy="stratify", ph_artifact=None,
                  strata_override=None, run_remedy=True):
        data = make_non_ph_covariate()
        data_path = tmp_path / "data.csv"
        data.to_csv(data_path, index=False)

        outcome = {"id": "primary", "name": "Death", "type": "time_to_event",
                   "time_column": "time", "event_column": "event",
                   "time_unit": "years", "time_scale": 1}
        if remedy:
            outcome["ph_violation_remedy"] = remedy
        if strata_override is not None:
            outcome["ph_strata_variables"] = strata_override

        pipeline = [
            {"id": "missingness", "skill": "missing-data", "status": "done",
             "depends_on": []},
            {"id": "multi_primary", "skill": "multivariate-interaction",
             "status": "done", "depends_on": [], "outcome": "primary",
             "output_file": "Multi_primary"},
        ]
        plan = _base_plan(
            tmp_path, data_path, outcome,
            [{"name": "Stage", "column": "stage", "type": "categorical",
              "reference_level": 0.0},
             {"name": "Age", "column": "age", "type": "continuous"}],
            pipeline)
        # Restructuring the pipeline is opt-in for plans written before schema 2 —
        # an older plan gets the change *reported* and has to run --migrate. A
        # current plan says so, and the remedy step is added on the spot.
        plan["plan_schema_version"] = 2
        plan_path = _write_plan(tmp_path, plan)
        _generate(plan_path)

        builder = TestFailsClosedWhenEvidenceIsMissing
        for name in ("Missingness.xlsx", "Multi_primary.xlsx"):
            builder._write_uni(os.path.join(str(tmp_path), name),
                               effect="1.42 (1.10–1.83)", n=1200, events=600)
        art = self.STRATA_VIOLATION if ph_artifact is None else ph_artifact
        with open(os.path.join(str(tmp_path), "Multi_primary_ph_test.json"), "w",
                  encoding="utf-8") as f:
            json.dump(art, f)

        cfg = os.path.join(str(tmp_path), "configs", "ph_stratified_primary_config.json")
        proc = None
        if run_remedy and os.path.exists(cfg):
            proc = _run([os.path.join(SCRIPTS, "ph_stratified.py"), cfg],
                        cwd=str(tmp_path))
        return plan_path, cfg, proc

    def _sidecar_path(self, tmp_path):
        return os.path.join(str(tmp_path), "Stratified_primary_remedy.json")

    def _tamper(self, tmp_path, **changes):
        path = self._sidecar_path(tmp_path)
        with open(path, encoding="utf-8") as f:
            side = json.load(f)
        side.update(changes)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(side, f)
        return side

    def _res015(self, report):
        return [i for i in (report or {}).get("findings", [])
                if i["check_id"] == "RES-015-det"]

    # ── the declaration now schedules the analysis ───────────────────────────

    def test_a_declared_remedy_adds_the_step_that_performs_it(self, tmp_path):
        plan_path, cfg, proc = self._scenario(tmp_path, run_remedy=False)
        with open(plan_path, encoding="utf-8") as f:
            pipeline = json.load(f)["pipeline"]
        added = [s for s in pipeline if s.get("skill") == "ph-stratified"]
        assert added, ("a plan declaring ph_violation_remedy='stratify' still has no "
                       "step that performs it — the declaration is a word again")
        assert added[0]["depends_on"] == ["multi_primary"], \
            "the remedy must run after the model whose PH test it reads"
        assert os.path.exists(cfg), "the added step got no config"

    def test_no_declared_remedy_adds_no_step(self, tmp_path):
        """The control for the check above: a clean plan is left alone."""
        plan_path, _cfg, _ = self._scenario(tmp_path, remedy=None, run_remedy=False)
        with open(plan_path, encoding="utf-8") as f:
            pipeline = json.load(f)["pipeline"]
        assert not [s for s in pipeline if s.get("skill") == "ph-stratified"], \
            "a remedy step was added to a plan that never asked for one"

    def test_the_step_produces_a_workbook_and_a_bound_sidecar(self, tmp_path):
        _plan, _cfg, proc = self._scenario(tmp_path)
        assert proc.returncode == 0, proc.stderr
        table = os.path.join(str(tmp_path), "Stratified_primary.xlsx")
        assert os.path.exists(table) and os.path.getsize(table) > 0
        with open(self._sidecar_path(tmp_path), encoding="utf-8") as f:
            side = json.load(f)
        assert side["status"] == "computed", side.get("reason")
        assert side["strata_variables"] == ["stage"]
        assert side["strata_variables_source"] == "auto_from_ph_test"
        assert side["variables_without_hr"] == ["stage"]
        assert "exposed" in side["covariates_in_model"], \
            "stratifying on the confounder must leave the exposure estimable"
        assert side["exposure"]["hr"] > 0
        assert side["table_basename"] == "Stratified_primary.xlsx"
        h = hashlib.sha256()
        with open(table, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 20), b""):
                h.update(chunk)
        assert side["table_sha256"] == h.hexdigest(), \
            "the sidecar is not bound to the workbook it sits beside"
        assert side["ph_recheck"]["status"] == "computed", \
            "'we applied a remedy' and 'the remedy worked' are different claims"

    def test_the_gate_verifies_the_real_thing(self, tmp_path):
        self._scenario(tmp_path)
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 0, proc.stderr
        assert len(report["verified_remedies"]) == 1, report["findings"]
        v = report["verified_remedies"][0]
        assert v["remedy"] == "stratify"
        assert "stage" in v["detail"], v["detail"]
        assert v["source_table"] == "Multi_primary.xlsx", \
            "the verified entry is also the manuscript's cross-reference back to the " \
            "unremediated table"
        assert report["attestations_unverified"] == [], \
            "a verified remedy must not also demand a signature — that is the whole point"
        with open(os.path.join(str(tmp_path), "configs", "manifest.json"),
                  encoding="utf-8") as f:
            manifest = json.load(f)
        assert not any(w.get("kind") == "ph_remedy_declared_not_computed"
                       for w in manifest.get("methodological_warnings", [])), \
            "a verified remedy was still recorded as unverified"

    # ── the content checks, each shown turning red ───────────────────────────

    def test_stratifying_on_the_wrong_variable_is_caught(self, tmp_path):
        """The case the old gate passed: right kind of step, wrong analysis."""
        self._scenario(tmp_path)
        table = os.path.join(str(tmp_path), "Stratified_primary.xlsx")
        self._tamper(tmp_path, strata_variables=["age"], variables_without_hr=["age"])
        assert os.path.getsize(table) > 0, \
            "the file-existence test the old gate ran is still satisfied here"
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 1
        msgs = " ".join(i["message"] for i in self._res015(report))
        assert "沒有被分層" in msgs, msgs
        assert report["verified_remedies"] == []

    def test_a_sidecar_that_read_a_different_ph_test_is_refused(self, tmp_path):
        self._scenario(tmp_path)
        self._tamper(tmp_path, ph_violating_variables=["some_other_covariate"])
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 1
        assert any("不一致" in i["message"] for i in self._res015(report))

    def test_a_sidecar_naming_another_table_is_refused(self, tmp_path):
        self._scenario(tmp_path)
        self._tamper(tmp_path, table_basename="Stratified_secondary.xlsx")
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 1
        assert any("自稱描述" in i["message"] for i in self._res015(report))

    def test_a_sidecar_written_for_an_older_version_of_the_table_is_refused(self, tmp_path):
        self._scenario(tmp_path)
        with open(os.path.join(str(tmp_path), "Stratified_primary.xlsx"), "ab") as f:
            f.write(b"\0")                      # the table was regenerated since
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 1
        assert any("另一個版本" in i["message"] for i in self._res015(report))

    def test_internally_inconsistent_evidence_is_refused(self, tmp_path):
        self._scenario(tmp_path)
        self._tamper(tmp_path, variables_without_hr=[])
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 1
        assert any("自相矛盾" in i["message"] for i in self._res015(report))

    def test_not_computed_evidence_is_not_a_pass(self, tmp_path):
        self._scenario(tmp_path)
        self._tamper(tmp_path, status="not_computed", reason="model did not converge")
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 1
        assert any("算不出來不是通過" in i["message"] for i in self._res015(report))

    def test_not_applicable_evidence_contradicting_a_violation_is_refused(self, tmp_path):
        """Three states, and the third one has to disagree with the artifact."""
        self._scenario(tmp_path)
        self._tamper(tmp_path, status="not_applicable", reason="no violation seen")
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 1
        assert any("矛盾" in i["message"] for i in self._res015(report))

    def test_a_wrong_remedy_in_the_sidecar_is_refused(self, tmp_path):
        self._scenario(tmp_path)
        self._tamper(tmp_path, remedy="rmst")
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 1
        assert any("必須是同一個 remedy" in i["message"] for i in self._res015(report))

    # ── quality signals warn, they do not block ─────────────────────────────

    def test_a_remedy_that_did_not_work_warns_without_blocking(self, tmp_path):
        """Stratifying and still violating is a real finding — and an author's
        call between time_varying, rmst and a limitation. Blocking on a judgement
        is how a check gets routed around."""
        self._scenario(tmp_path)
        self._tamper(tmp_path, ph_recheck={"status": "computed", "violated": True,
                                           "min_p": 0.002,
                                           "violating_variables": ["age"]})
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 0, "a quality signal must not block"
        assert any(i["severity"] == "warn" and "仍然違反" in i["message"]
                   for i in self._res015(report))
        assert len(report["verified_remedies"]) == 1

    # ── the legacy path, unchanged ──────────────────────────────────────────

    def test_a_step_without_a_sidecar_still_takes_the_signature_path(self, tmp_path):
        """Zero migration cost: an output built before sidecars existed still
        passes, on the old terms — file existence plus a named signature."""
        self._scenario(tmp_path)
        os.remove(self._sidecar_path(tmp_path))
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 0, proc.stderr
        assert report["verified_remedies"] == []
        assert len(report["attestations_unverified"]) == 1
        with open(os.path.join(str(tmp_path), "configs", "manifest.json"),
                  encoding="utf-8") as f:
            manifest = json.load(f)
        assert any(w.get("kind") == "ph_remedy_declared_not_computed"
                   for w in manifest.get("methodological_warnings", [])), \
            "the unverified path stopped asking anybody to sign for it"

    # ── the third state, and the two statistical refusals ───────────────────

    def test_no_violation_makes_the_remedy_not_applicable(self, tmp_path):
        """The step is scheduled before the test result is known, so finding
        nothing to fix is a clean exit — not a failure, and not a pass either."""
        clean = {"status": "computed", "violated": False, "borderline": False,
                 "min_p": 0.42, "alpha": 0.01, "global": {"p": 0.42},
                 "violating_variables": []}
        _plan, _cfg, proc = self._scenario(tmp_path, ph_artifact=clean)
        assert proc.returncode == 0, proc.stderr
        with open(self._sidecar_path(tmp_path), encoding="utf-8") as f:
            side = json.load(f)
        assert side["status"] == "not_applicable"
        gate_proc, report = self._gate(tmp_path)
        assert gate_proc.returncode == 0
        assert self._res015(report) == [], "a clean PH test raised remedy findings"

    def test_an_unknown_upstream_ph_status_is_not_treated_as_absent(self, tmp_path):
        unknown = {"status": "not_computed", "reason": "all events share one time",
                   "violating_variables": []}
        _plan, _cfg, proc = self._scenario(tmp_path, ph_artifact=unknown)
        assert proc.returncode == 2
        with open(self._sidecar_path(tmp_path), encoding="utf-8") as f:
            side = json.load(f)
        assert side["status"] == "not_computed"
        assert "unknown violation is not an absent one" in side["reason"]

    def test_stratifying_on_the_exposure_is_refused(self, tmp_path):
        """`strata=` deletes the coefficient. If that is the exposure, the
        'remedy' deletes the result the study exists to report."""
        art = dict(self.STRATA_VIOLATION, violating_variables=["exposed"])
        _plan, _cfg, proc = self._scenario(tmp_path, ph_artifact=art)
        assert proc.returncode == 2
        with open(self._sidecar_path(tmp_path), encoding="utf-8") as f:
            side = json.load(f)
        assert side["status"] == "not_computed"
        assert "time_varying" in side["reason"] and "rmst" in side["reason"], \
            "a refusal has to name the remedies that would work"

    def test_a_continuous_violator_is_refused_rather_than_binned(self, tmp_path):
        art = dict(self.STRATA_VIOLATION, violating_variables=["age"])
        _plan, _cfg, proc = self._scenario(tmp_path, ph_artifact=art)
        assert proc.returncode == 2
        with open(self._sidecar_path(tmp_path), encoding="utf-8") as f:
            side = json.load(f)
        assert side["status"] == "not_computed"
        assert "cut points" in side["reason"], \
            "silently binning would put an unrecorded analytic decision under an estimate"

    # ── who reads the numbers in the remedy table ───────────────────────────
    # Nobody but RES-001. The gate verifies the sidecar and deliberately never
    # parses the workbook, and the sidecar is written from the same fit rather
    # than from the sheet — so a producer regression that emptied the table
    # would leave the sidecar saying `computed` and the gate saying green.
    # `ph-stratified` and `ph-time-varying` shipped without being named in
    # _ESTIMATE_PRODUCING_SKILLS, the tuple that makes "no readable effect in
    # this workbook" a blocker, so that silence was the whole answer.

    @staticmethod
    def _empty_the_effects(path):
        """Strip every readable estimate, leaving the table otherwise intact."""
        import openpyxl
        wb = openpyxl.load_workbook(path)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    text = str(value).strip()
                    # Both shapes: the formatted cell, and the bare number in a
                    # role-typed column. ph_stratified writes its coefficient as
                    # TEXT, so an isinstance check on float misses it entirely.
                    if "(" in text and any(d in text for d in "–—-"):
                        cell.value = "—"
                        continue
                    try:
                        float(text)
                    except ValueError:
                        continue
                    cell.value = "—"
        wb.save(path)

    def test_a_stratified_table_that_reports_hazard_ratios_does_not_block(self, tmp_path):
        """The control, on the workbook ph_stratified.py really wrote."""
        plan_path, _cfg, proc = self._scenario(tmp_path)
        assert proc is not None and proc.returncode == 0, "the remedy did not run"
        post = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "no effect estimate could be read" not in post.stderr, post.stderr

    def test_an_emptied_stratified_table_blocks(self, tmp_path):
        """A workbook only exists on the path where the model fitted.

        _not_computed and _not_applicable both write the sidecar and return
        before `wb.save`, so "the file is there and holds no effect" is never
        the honest degraded state — it is a table that lost its numbers.
        """
        plan_path, _cfg, proc = self._scenario(tmp_path)
        assert proc is not None and proc.returncode == 0, "the remedy did not run"
        self._empty_the_effects(
            os.path.join(str(tmp_path), "Stratified_primary.xlsx"))
        post = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "no effect estimate could be read" in post.stderr, post.stderr
        assert post.returncode == 1

    def test_the_rmst_table_is_deliberately_left_out(self):
        """ph_rmst states a difference in years as prose, and
        parse_labelled_effect refuses that shape on purpose. Adding ph-rmst here
        for symmetry would block every rmst run that ever succeeds.
        """
        import precheck_configs
        assert "ph-rmst" not in precheck_configs._ESTIMATE_PRODUCING_SKILLS
        assert {"ph-stratified", "ph-time-varying"} <= set(
            precheck_configs._ESTIMATE_PRODUCING_SKILLS)


class TestPhRemedyRmstAndTimeVarying:
    """The other two remedies, and the two things only they have to defend.

    `stratify` is checked by whether it stratified on the covariate that actually
    violated. These two have a second failure mode the coverage check cannot see:
    both are parameterised by a TIME the author chooses — tau for RMST, the
    window boundaries for the piecewise model — and both answers move with that
    choice. A tau picked after seeing which value flatters the comparison is the
    entire p-hacking surface of RMST; windows drawn where the effect happens to
    turn are a subgroup search along the time axis. So the gate demands that the
    number came from `outcomes[]`, which is what puts it inside the plan
    fingerprint.

    The analysis still runs when it did not — the author has to be able to see
    the numbers and the exact fix — so "runs but is blocked at the manuscript"
    is asserted as a pair, not just the block.
    """

    VIOLATION = {"status": "computed", "violated": True, "borderline": False,
                 "min_p": 0.0004, "alpha": 0.01,
                 "global": {"chi2": 18.2, "df": 2, "p": 0.0004},
                 "violating_variables": ["stage"]}

    STEPS = {"rmst": ("ph_rmst_primary", "RMST_primary"),
             "time_varying": ("ph_time_varying_primary", "TimeVarying_primary")}

    def _gate(self, tmp_path):
        import precheck_configs as PC
        out = os.path.join(str(tmp_path), "ph_gate.json")
        proc = _run([PC.PH_REMEDY_GATE, "--plan",
                     os.path.join(str(tmp_path), "research_plan.json"),
                     "--output-dir", str(tmp_path), "--json", out])
        report = None
        if os.path.exists(out):
            with open(out, encoding="utf-8") as f:
                report = json.load(f)
        return proc, report

    def _scenario(self, tmp_path, remedy, outcome_extra=None, ph_artifact=None):
        data = make_non_ph_covariate()
        data_path = tmp_path / "data.csv"
        data.to_csv(data_path, index=False)

        outcome = {"id": "primary", "name": "Death", "type": "time_to_event",
                   "time_column": "time", "event_column": "event",
                   "time_unit": "years", "time_scale": 1,
                   "ph_violation_remedy": remedy}
        outcome.update(outcome_extra or {})
        pipeline = [
            {"id": "missingness", "skill": "missing-data", "status": "done",
             "depends_on": []},
            {"id": "multi_primary", "skill": "multivariate-interaction",
             "status": "done", "depends_on": [], "outcome": "primary",
             "output_file": "Multi_primary"},
        ]
        plan = _base_plan(
            tmp_path, data_path, outcome,
            [{"name": "Stage", "column": "stage", "type": "categorical",
              "reference_level": 0.0},
             {"name": "Age", "column": "age", "type": "continuous"}],
            pipeline)
        plan["plan_schema_version"] = 2
        plan_path = _write_plan(tmp_path, plan)
        _generate(plan_path)

        builder = TestFailsClosedWhenEvidenceIsMissing
        for name in ("Missingness.xlsx", "Multi_primary.xlsx"):
            builder._write_uni(os.path.join(str(tmp_path), name),
                               effect="1.42 (1.10–1.83)", n=1200, events=600)
        with open(os.path.join(str(tmp_path), "Multi_primary_ph_test.json"), "w",
                  encoding="utf-8") as f:
            json.dump(self.VIOLATION if ph_artifact is None else ph_artifact, f)

        step_id, stem = self.STEPS[remedy]
        cfg = os.path.join(str(tmp_path), "configs", f"{step_id}_config.json")
        proc = None
        if os.path.exists(cfg):
            script = {"rmst": "ph_rmst.py",
                      "time_varying": "ph_time_varying.py"}[remedy]
            proc = _run([os.path.join(SCRIPTS, script), cfg], cwd=str(tmp_path))
        return plan_path, cfg, proc, os.path.join(str(tmp_path), f"{stem}_remedy.json")

    def _res015(self, report):
        return [i for i in (report or {}).get("findings", [])
                if i["check_id"] == "RES-015-det"]

    def _tamper(self, path, **changes):
        with open(path, encoding="utf-8") as f:
            side = json.load(f)
        side.update(changes)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(side, f)

    # ── RMST ────────────────────────────────────────────────────────────────

    def test_rmst_runs_and_is_verified_when_tau_is_prespecified(self, tmp_path):
        _p, cfg, proc, side_path = self._scenario(
            tmp_path, "rmst", {"rmst_tau": 4.0})
        assert os.path.exists(cfg), "declaring rmst added no step"
        assert proc.returncode == 0, proc.stderr
        with open(side_path, encoding="utf-8") as f:
            side = json.load(f)
        assert side["status"] == "computed", side.get("reason")
        assert side["tau"] == 4.0 and side["tau_source"] == "plan_value"
        assert side["adjusted"] is False, \
            "an unadjusted comparison must say so in the evidence, not only in prose"
        proc2, report = self._gate(tmp_path)
        assert proc2.returncode == 0, proc2.stderr
        assert len(report["verified_remedies"]) == 1
        assert report["verified_remedies"][0]["remedy"] == "rmst"

    def test_rmst_standard_error_is_the_estimators_not_the_populations(self, tmp_path):
        """Regression: lifelines' `return_variance=True` gives the spread of
        restricted survival time ACROSS INDIVIDUALS, not the sampling variance of
        the estimate. Using its square root as a standard error inflated it by
        roughly sqrt(n) — measured here as SE 2.06 against an RMST of 1.45 — and
        made a genuine half-year difference read as p = 0.80."""
        _p, _cfg, proc, side_path = self._scenario(
            tmp_path, "rmst", {"rmst_tau": 4.0})
        assert proc.returncode == 0, proc.stderr
        with open(side_path, encoding="utf-8") as f:
            side = json.load(f)
        for arm in side["arms"]:
            assert arm["se"] is not None
            assert arm["se"] < arm["rmst"] / 5, (
                f"{arm['label']}: SE {arm['se']:.3f} against RMST {arm['rmst']:.3f} — "
                f"that is a population standard deviation, not a standard error")
        c = side["contrast"]
        assert c["p"] < 0.05, (
            "the fixture's arms differ by construction; a non-significant result "
            "here means the variance is being computed on the wrong quantity again")
        assert c["ci"][0] < c["difference"] < c["ci"][1]

    def test_rmst_without_a_declared_tau_runs_but_is_blocked(self, tmp_path):
        """A door with a sign on it, not a locked one: the numbers are produced
        so the author can see them, and the manuscript is what stops."""
        _p, _cfg, proc, side_path = self._scenario(tmp_path, "rmst")
        assert proc.returncode == 0, "the analysis should still run"
        assert os.path.exists(os.path.join(str(tmp_path), "RMST_primary.xlsx")), \
            "the table should exist — the author has to be able to see the numbers"
        with open(side_path, encoding="utf-8") as f:
            side = json.load(f)
        assert side["tau_source"] == "default_rule"
        gate_proc, report = self._gate(tmp_path)
        assert gate_proc.returncode == 1
        msgs = " ".join(i["message"] for i in self._res015(report))
        assert "無法證明它早於結果" in msgs, msgs
        assert report["verified_remedies"] == []

    def test_a_named_tau_rule_counts_as_prespecification(self, tmp_path):
        _p, _cfg, proc, side_path = self._scenario(
            tmp_path, "rmst", {"rmst_tau_rule": "min_of_arm_max_followup"})
        assert proc.returncode == 0, proc.stderr
        with open(side_path, encoding="utf-8") as f:
            side = json.load(f)
        assert side["tau_source"] == "plan_rule"
        assert side["tau_rule"] == "min_of_arm_max_followup"
        assert side["tau"] > 0
        gate_proc, report = self._gate(tmp_path)
        assert gate_proc.returncode == 0, proc.stderr
        assert len(report["verified_remedies"]) == 1

    def test_a_tau_beyond_follow_up_warns_without_blocking(self, tmp_path):
        _p, _cfg, proc, side_path = self._scenario(
            tmp_path, "rmst", {"rmst_tau": 99.0})
        assert proc.returncode == 0, proc.stderr
        with open(side_path, encoding="utf-8") as f:
            assert json.load(f)["tau_within_followup"] is False
        gate_proc, report = self._gate(tmp_path)
        assert gate_proc.returncode == 0, "extrapolation is a judgement, not a defect"
        assert any(i["severity"] == "warn" and "外推" in i["message"]
                   for i in self._res015(report))

    # ── time_varying ────────────────────────────────────────────────────────

    def test_time_varying_runs_and_recovers_the_moving_effect(self, tmp_path):
        """The fixture's `stage` multiplies the hazard only before t=1. A model
        that reports one averaged HR hides that; this one has to find it."""
        _p, cfg, proc, side_path = self._scenario(
            tmp_path, "time_varying", {"tv_cut_points": [1.0]})
        assert os.path.exists(cfg), "declaring time_varying added no step"
        assert proc.returncode == 0, proc.stderr
        with open(side_path, encoding="utf-8") as f:
            side = json.load(f)
        assert side["status"] == "computed", side.get("reason")
        assert side["cut_points"] == [1.0] and side["cut_points_source"] == "plan"
        assert side["tv_method"] == "piecewise_constant"
        early, late = side["period_effects"]["stage"]
        assert early["hr"] > 1.4 > late["hr"], (
            f"early HR {early['hr']:.2f}, late HR {late['hr']:.2f} — the fixture's "
            f"effect is present only in the first window and this did not find it")
        assert side["heterogeneity"]["stage"]["status"] == "computed", \
            side["heterogeneity"]["stage"].get("reason")
        assert side["heterogeneity"]["stage"]["differs"] is True

    def test_time_varying_is_verified_by_the_gate(self, tmp_path):
        self._scenario(tmp_path, "time_varying", {"tv_cut_points": [1.0]})
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 0, proc.stderr
        assert len(report["verified_remedies"]) == 1
        assert report["verified_remedies"][0]["remedy"] == "time_varying"
        assert report["attestations_unverified"] == []

    def test_undeclared_windows_run_but_are_blocked(self, tmp_path):
        _p, _cfg, proc, side_path = self._scenario(tmp_path, "time_varying")
        assert proc.returncode == 0, "the analysis should still run"
        with open(side_path, encoding="utf-8") as f:
            assert json.load(f)["cut_points_source"] == "default_rule"
        gate_proc, report = self._gate(tmp_path)
        assert gate_proc.returncode == 1
        msgs = " ".join(i["message"] for i in self._res015(report))
        assert "次族群搜尋" in msgs, msgs

    def test_making_the_wrong_variable_time_varying_is_caught(self, tmp_path):
        """The same failure stratify has: right kind of analysis, wrong target."""
        _p, _cfg, _proc, side_path = self._scenario(
            tmp_path, "time_varying", {"tv_cut_points": [1.0]})
        self._tamper(side_path, time_varying_variables=["age"])
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 1
        assert any("係數仍是固定的" in i["message"] for i in self._res015(report))

    def test_indistinguishable_windows_warn_without_blocking(self, tmp_path):
        _p, _cfg, _proc, side_path = self._scenario(
            tmp_path, "time_varying", {"tv_cut_points": [1.0]})
        self._tamper(side_path, heterogeneity={
            "stage": {"status": "computed", "differs": False, "p": 0.41}})
        proc, report = self._gate(tmp_path)
        assert proc.returncode == 0, "whether the effect moved is a finding, not a defect"
        assert any(i["severity"] == "warn" and "兩個不同的宣稱" in i["message"]
                   for i in self._res015(report))

    # ── controls ────────────────────────────────────────────────────────────

    @pytest.mark.parametrize("remedy,extra", [("rmst", {"rmst_tau": 4.0}),
                                              ("time_varying", {"tv_cut_points": [1.0]})])
    def test_no_violation_makes_either_remedy_not_applicable(self, tmp_path, remedy, extra):
        clean = {"status": "computed", "violated": False, "borderline": False,
                 "min_p": 0.42, "alpha": 0.01, "global": {"p": 0.42},
                 "violating_variables": []}
        _p, _cfg, proc, side_path = self._scenario(
            tmp_path, remedy, extra, ph_artifact=clean)
        assert proc.returncode == 0, proc.stderr
        with open(side_path, encoding="utf-8") as f:
            assert json.load(f)["status"] == "not_applicable"
        gate_proc, report = self._gate(tmp_path)
        assert gate_proc.returncode == 0
        assert self._res015(report) == []

    def test_disclose_still_adds_no_remedy_step(self, tmp_path):
        """The control for the auto-add table: `disclose` is not in it."""
        data = make_non_ph_covariate()
        data.to_csv(tmp_path / "data.csv", index=False)
        outcome = {"id": "primary", "name": "Death", "type": "time_to_event",
                   "time_column": "time", "event_column": "event",
                   "time_unit": "years", "time_scale": 1,
                   "ph_violation_remedy": "disclose",
                   "ph_disclose_decision_id": "D-017"}
        plan = _base_plan(
            tmp_path, tmp_path / "data.csv", outcome,
            [{"name": "Age", "column": "age", "type": "continuous"}],
            [{"id": "missingness", "skill": "missing-data", "status": "done",
              "depends_on": []},
             {"id": "multi_primary", "skill": "multivariate-interaction",
              "status": "done", "depends_on": [], "outcome": "primary",
              "output_file": "Multi_primary"}])
        plan["plan_schema_version"] = 2
        plan_path = _write_plan(tmp_path, plan)
        _generate(plan_path)
        with open(plan_path, encoding="utf-8") as f:
            skills = {s.get("skill") for s in json.load(f)["pipeline"]}
        assert not any(s and s.startswith("ph-") for s in skills), \
            f"disclose gained a remedy step: {skills}"

    # ── the two halves of "this table must hold an effect" ──────────────────

    def test_a_time_varying_run_that_succeeded_is_not_read_as_holding_nothing(
            self, tmp_path):
        """ph-time-varying is in _ESTIMATE_PRODUCING_SKILLS, so a successful run
        whose table were unreadable would block every time. The control has to
        run the real producer: the whole point of the tuple is a claim about
        what this script writes.
        """
        plan_path, cfg, proc, _side = self._scenario(
            tmp_path, "time_varying", {"tv_cut_points": [1.0]})
        assert os.path.exists(cfg) and proc.returncode == 0, proc.stderr
        post = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "no effect estimate could be read" not in post.stderr, post.stderr

    def test_a_successful_rmst_run_is_not_blocked_for_holding_no_ratio(self, tmp_path):
        """The other half, and the reason ph-rmst is left out of that tuple.

        A difference in years is not an effect parse_labelled_effect will read,
        by design — so if ph-rmst is ever added for symmetry, this turns red on
        a run that did everything right.
        """
        plan_path, cfg, proc, _side = self._scenario(
            tmp_path, "rmst", {"rmst_tau": 4.0})
        assert os.path.exists(cfg) and proc.returncode == 0, proc.stderr
        post = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "no effect estimate could be read" not in post.stderr, post.stderr


class TestCannotRunExitsTwo:
    """Every "could not run" path must exit 2 AND print a scannable line.

    `sys.exit("some string")` prints the string and exits **1**. A caller that
    branches on the exit code reads 1 as "ran, found blocking problems", then
    reports the findings by counting `[QC_*]` lines — and a bare `[FATAL]` line
    is not one of those. The one line explaining the failure got dropped and the
    run was logged as `exit 1 — 0 blocking / 0 warning`. Both halves are asserted
    here, because fixing either alone leaves the other kind of caller broken.
    """

    def _plan(self, tmp_path, name, payload):
        path = os.path.join(str(tmp_path), name)
        with open(path, "w", encoding="utf-8") as f:
            if isinstance(payload, str):
                f.write(payload)
            else:
                json.dump(payload, f)
        return path

    def test_plan_file_not_found(self, tmp_path):
        proc = _run([PRECHECK, os.path.join(str(tmp_path), "nope.json")])
        assert proc.returncode == 2
        assert "[QC_BLOCKING]" in proc.stderr and "not found" in proc.stderr

    def test_plan_is_not_valid_json(self, tmp_path):
        path = self._plan(tmp_path, "bad.json", "{not json")
        proc = _run([PRECHECK, path, "--stage", "pre"])
        assert proc.returncode == 2
        assert "[QC_BLOCKING]" in proc.stderr and "not valid JSON" in proc.stderr

    def test_plan_missing_a_required_key(self, tmp_path):
        path = self._plan(tmp_path, "nokey.json", {"study": {}, "pipeline": []})
        proc = _run([PRECHECK, path, "--stage", "pre"])
        assert proc.returncode == 2
        assert "[QC_BLOCKING]" in proc.stderr and "output_dir" in proc.stderr

    def test_dataset_path_does_not_exist(self, tmp_path):
        path = self._plan(tmp_path, "nodata.json", {
            "output_dir": str(tmp_path),
            "study": {"data_file": os.path.join(str(tmp_path), "missing.csv")},
            "pipeline": []})
        proc = _run([PRECHECK, path, "--stage", "pre"])
        assert proc.returncode == 2
        assert "[QC_BLOCKING]" in proc.stderr and "dataset not found" in proc.stderr

    def test_dataset_cannot_be_parsed(self, tmp_path):
        corrupt = os.path.join(str(tmp_path), "corrupt.xlsx")
        with open(corrupt, "w", encoding="utf-8") as f:
            f.write("not really a workbook")
        path = self._plan(tmp_path, "unreadable.json", {
            "output_dir": str(tmp_path), "study": {"data_file": corrupt},
            "pipeline": []})
        proc = _run([PRECHECK, path, "--stage", "pre"])
        assert proc.returncode == 2
        assert "[QC_BLOCKING]" in proc.stderr and "cannot read dataset" in proc.stderr

    def test_no_plan_argument_at_all(self):
        proc = _run([PRECHECK, "--stage", "pre"])
        assert proc.returncode == 2
        assert "[QC_BLOCKING]" in proc.stderr and "no plan given" in proc.stderr

    def test_configs_is_not_a_directory(self, tmp_path):
        path = self._plan(tmp_path, "nokey.json", {"study": {}, "pipeline": []})
        proc = _run([PRECHECK, path, "--configs", os.path.join(str(tmp_path), "nope")])
        assert proc.returncode == 2
        assert "[QC_BLOCKING]" in proc.stderr and "not a directory" in proc.stderr

    def test_no_fatal_prefix_survives_anywhere(self):
        """`[FATAL]` is invisible to a caller scanning for `[QC_*]`."""
        for name in ("precheck_configs.py", "estimate_inventory.py",
                     "check_numbers_sourced.py"):
            with open(os.path.join(SCRIPTS, name), encoding="utf-8") as f:
                body = f.read()
            emitted = [ln for ln in body.splitlines()
                       if "[FATAL]" in ln and not ln.lstrip().startswith("#")
                       and "`[FATAL]`" not in ln]
            assert not emitted, f"{name} still emits [FATAL]: {emitted}"


class TestCoverageInventoryIsProduced:
    """C7 reads coverage_diff.json. Nothing used to write one."""

    def _scenario(self, tmp_path, with_manuscript=True, quote=None):
        rng = np.random.default_rng(4)
        n = 400
        pd.DataFrame({"exposed": rng.integers(0, 2, n).astype(float),
                      "age": rng.normal(65, 10, n),
                      "dead": rng.binomial(1, 0.4, n).astype(float)}).to_csv(
            tmp_path / "data.csv", index=False)
        pipeline = [
            {"id": "flowchart", "skill": "flowchart", "status": "done",
             "depends_on": [], "output_file": "Flowchart"},
            {"id": "missingness", "skill": "missing-data", "status": "done",
             "depends_on": []},
            {"id": "uni_primary", "skill": "univariate-forest", "status": "done",
             "depends_on": [], "outcome": "primary", "output_file": "Uni_primary"},
        ]
        if with_manuscript:
            pipeline += [
                {"id": "reviewer_challenge", "skill": "reviewer-challenge",
                 "status": "done", "depends_on": ["uni_primary"]},
                {"id": "manuscript", "skill": "manuscript-results", "status": "done",
                 "depends_on": ["reviewer_challenge"]},
            ]
        plan = _base_plan(
            tmp_path, tmp_path / "data.csv",
            {"id": "primary", "name": "Death", "type": "binary", "event_column": "dead"},
            [{"name": "Age", "column": "age", "type": "continuous"}], pipeline)
        plan_path = _write_plan(tmp_path, plan)
        _generate(plan_path)

        builder = TestFailsClosedWhenEvidenceIsMissing
        builder._write_flowchart(os.path.join(str(tmp_path), "Flowchart.xlsx"), n)
        builder._write_uni(os.path.join(str(tmp_path), "Missingness.xlsx"),
                           effect="1.10 (1.02–1.19)", n=n, events=160)
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Univariate Analysis"
        ws.append(["Variable", "N (events)", "OR (95% CI)", "p-value"])
        ws.append(["Age", f"{n} (160)", "2.22 (1.29–3.82)", "0.004*"])
        ws.append(["Exposure", f"{n} (160)", "1.55 (1.01–2.38)", "0.045*"])
        ws.append(["Sex", f"{n} (160)", "0.88 (0.55–1.41)", "0.601"])
        wb.save(os.path.join(str(tmp_path), "Uni_primary.xlsx"))

        if with_manuscript and quote is not None:
            import docx
            document = docx.Document()
            document.add_paragraph("Results")
            document.add_paragraph(quote)
            document.save(os.path.join(str(tmp_path), "Results_Section.docx"))
        return plan_path

    def test_post_stage_writes_coverage_diff(self, tmp_path):
        plan_path = self._scenario(
            tmp_path, quote="Age was associated with mortality "
                            "(OR 2.22, 95% CI 1.29-3.82, p=0.004).")
        assert not os.path.exists(os.path.join(str(tmp_path), "coverage_diff.json"))
        _run([PRECHECK, plan_path, "--stage", "post"])
        assert os.path.exists(os.path.join(str(tmp_path), "coverage_diff.json")), \
            "C7 still has nothing to read"

    def test_the_docx_manuscript_is_actually_read(self, tmp_path):
        """manuscript-results writes .docx. Reading only markdown made the
        pipeline's own deliverable invisible, and every estimate in it would
        have been reported as omitted."""
        plan_path = self._scenario(
            tmp_path, quote="Age was associated with mortality "
                            "(OR 2.22, 95% CI 1.29-3.82, p=0.004).")
        _run([PRECHECK, plan_path, "--stage", "post"])
        with open(os.path.join(str(tmp_path), "coverage_diff.json"),
                  encoding="utf-8") as f:
            payload = json.load(f)
        assert payload["status"] == "compared"
        assert payload["reported_source_files"] == ["Results_Section.docx"]
        assert payload["counters"]["reported_estimates"] == 1
        omitted = {round(o["estimate"], 2) for o in payload["omitted"]}
        assert omitted == {1.10, 1.55, 0.88}

    def test_omissions_are_surfaced_as_a_set_difference_not_a_verdict(self, tmp_path):
        plan_path = self._scenario(
            tmp_path, quote="Age was associated with mortality (OR 2.22).")
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "[QC_WARNING] C7" in proc.stderr
        assert "do not appear in the deliverable" in proc.stderr
        assert "[QC_BLOCKING] C7" not in proc.stderr, \
            "coverage is a question for a reader, not a blocking verdict"

    def test_a_complete_report_produces_no_coverage_warning(self, tmp_path):
        plan_path = self._scenario(
            tmp_path, quote="ORs were 2.22, 1.55, 0.88 in the primary model, "
                            "and 1.10 for the missingness model.")
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        with open(os.path.join(str(tmp_path), "coverage_diff.json"),
                  encoding="utf-8") as f:
            payload = json.load(f)
        assert payload["counters"]["omitted_estimates"] == 0
        assert "[QC_WARNING] C7" not in proc.stderr

    def test_no_deliverable_is_a_distinct_state_not_an_empty_omitted_list(self, tmp_path):
        """"The manuscript does not exist yet" and "the manuscript omits nothing"
        must not both look like `omitted: []`."""
        plan_path = self._scenario(tmp_path, with_manuscript=False)
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        with open(os.path.join(str(tmp_path), "coverage_diff.json"),
                  encoding="utf-8") as f:
            payload = json.load(f)
        assert payload["status"] == "no_deliverable_found"
        assert payload["counters"]["executed_estimates"] > 0, \
            "the executed inventory should still be built"
        assert payload["omitted"] == []
        assert "no deliverable could be read" in proc.stderr


class TestMissingnessTiersMatchTheirMessages:
    """A message that names a threshold the check does not have is worse than no
    message: a reader infers from its absence that the threshold was tested and
    cleared. COL-005's top band said "(>50%)" while testing `miss >= 0.60`, so a
    covariate at 55% produced only a 40%-band line, and nothing anywhere
    mentioned 50%. Tiers per missing_data_protocol.md — COL-005 is 20/40/60, the
    Stage Reflector is 20/50, and the difference is deliberate.
    """

    def _scenario(self, tmp_path, fractions, allow_high_missing=()):
        n = 1000
        rng = np.random.default_rng(9)
        data = {"exposed": rng.integers(0, 2, n).astype(float),
                "dead": rng.binomial(1, 0.4, n).astype(float)}
        covariates = []
        for name, frac in fractions.items():
            values = rng.normal(0, 1, n)
            values[:int(round(frac * n))] = np.nan
            data[name] = values
            cov = {"name": name.upper(), "column": name, "type": "continuous"}
            if name in allow_high_missing:
                cov["allow_high_missing"] = True
            covariates.append(cov)
        pd.DataFrame(data).to_csv(tmp_path / "data.csv", index=False)
        plan = _base_plan(
            tmp_path, tmp_path / "data.csv",
            {"id": "primary", "name": "Death", "type": "binary", "event_column": "dead"},
            covariates,
            [{"id": "missingness", "skill": "missing-data", "status": "done",
              "depends_on": []}])
        plan_path = _write_plan(tmp_path, plan)
        _generate(plan_path)
        return plan_path

    @staticmethod
    def _lines(proc, check_id, column):
        return [ln for ln in proc.stderr.splitlines()
                if f"] {check_id}" in ln and f"'{column}'" in ln]

    @pytest.fixture(scope="class")
    def three_tiers(self, tmp_path_factory):
        tmp = tmp_path_factory.mktemp("missingness")
        plan_path = self._scenario(tmp, {"m45": 0.45, "m55": 0.55, "m65": 0.65})
        return _run([PRECHECK, plan_path, "--stage", "pre"])

    def test_45_percent_lands_in_the_40_to_60_band(self, three_tiers):
        lines = self._lines(three_tiers, "COL-005", "m45")
        assert len(lines) == 1
        assert "45.0% missing (40-60%)" in lines[0]
        assert "50%" not in lines[0], "the message must not imply a 50% tier here"

    def test_55_percent_lands_in_the_same_band_and_says_so(self, three_tiers):
        """The case that used to be silent about 50%."""
        lines = self._lines(three_tiers, "COL-005", "m55")
        assert len(lines) == 1
        assert "55.0% missing (40-60%)" in lines[0]

    def test_55_percent_also_trips_the_reflector_fifty_percent_rule(self, three_tiers):
        """So nobody can infer from COL-005's silence that 50% was cleared."""
        lines = self._lines(three_tiers, "REFLECT-missing", "m55")
        assert len(lines) == 1
        assert "55.0% missing (>50%)" in lines[0]
        assert "sensitivity analysis" in lines[0]

    def test_65_percent_lands_in_the_top_band(self, three_tiers):
        lines = self._lines(three_tiers, "COL-005", "m65")
        assert len(lines) == 1
        assert "65.0% missing (>=60%)" in lines[0]
        assert "sensitivity analysis" in lines[0]

    def test_45_percent_does_not_trip_the_fifty_percent_rule(self, three_tiers):
        assert self._lines(three_tiers, "REFLECT-missing", "m45") == []

    def test_every_stated_band_matches_the_value_it_reports(self, three_tiers):
        """Parse each message back and check the number really is in the band."""
        import re as _re

        for line in three_tiers.stderr.splitlines():
            m = _re.search(r"is (\d+\.\d)% missing \((>=60%|40-60%|20-40%|>50%)\)", line)
            if not m:
                continue
            value, band = float(m.group(1)), m.group(2)
            if band == ">=60%":
                assert value >= 60
            elif band == "40-60%":
                assert 40 <= value < 60
            elif band == "20-40%":
                assert 20 <= value < 40
            else:
                assert value > 50

    def test_20_to_40_band_is_reported_as_such(self, tmp_path):
        plan_path = self._scenario(tmp_path, {"m30": 0.30})
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        lines = self._lines(proc, "COL-005", "m30")
        assert len(lines) == 1
        assert "30.0% missing (20-40%)" in lines[0]

    def test_below_twenty_percent_says_nothing(self, tmp_path):
        plan_path = self._scenario(tmp_path, {"m10": 0.10})
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        assert self._lines(proc, "COL-005", "m10") == []
        assert self._lines(proc, "REFLECT-missing", "m10") == []

    def test_allow_high_missing_does_not_downgrade_to_the_twenty_band(self, tmp_path):
        """A 65%-missing column with the flag set used to be announced as ">20%"."""
        plan_path = self._scenario(tmp_path, {"m65": 0.65}, allow_high_missing=("m65",))
        proc = _run([PRECHECK, plan_path, "--stage", "pre"])
        warnings = [ln for ln in self._lines(proc, "COL-005", "m65")
                    if ln.startswith("[QC_WARNING]")]
        assert warnings == []
        acknowledged = [ln for ln in proc.stderr.splitlines()
                        if "allow_high_missing is set" in ln and "m65" in ln]
        assert len(acknowledged) == 1


class TestRoleThreeChecksRestored:
    """RES-005 / 008 / 012 / 013 / 014 — coverage the script layer dropped.

    Each has a trigger case and a clean case. The clean cases matter as much as
    the triggers: on this module a false block has twice cost more than the miss
    it prevented, so "the honest output raises nothing" is asserted every time.
    """

    @pytest.fixture(scope="class")
    def survival_project(self, tmp_path_factory):
        """A real KM + CIF + Table 1 run — the outputs these checks parse."""
        tmp = tmp_path_factory.mktemp("roleiii")
        rng = np.random.default_rng(11)
        n = 500
        exposed = rng.integers(0, 2, n)
        t_primary = rng.exponential(1 / (0.12 * np.exp(0.4 * exposed)))
        t_compete = rng.exponential(1 / 0.35, n)
        t = np.minimum(np.minimum(t_primary, t_compete), 8.0)
        status = np.where((t_primary <= t_compete) & (t_primary <= 8.0), 1,
                          np.where(t_compete < 8.0, 2, 0))
        pd.DataFrame({"exposed": exposed.astype(float),
                      "age": rng.normal(65, 10, n),
                      "sex": rng.integers(0, 2, n).astype(float),
                      "time": t.round(3), "event": (status == 1).astype(float),
                      "status": status.astype(float)}).to_csv(tmp / "data.csv", index=False)
        outcome = {"id": "primary", "name": "Event", "type": "competing_risk",
                   "time_column": "time", "event_column": "event",
                   "competing_event_column": "status",
                   "competing_event_labels": {"1": "Event", "2": "Death"},
                   "time_unit": "years", "time_scale": 1, "time_points": [1, 3, 5]}
        plan = _base_plan(
            tmp, tmp / "data.csv", outcome,
            [{"name": "Age", "column": "age", "type": "continuous"},
             {"name": "Male", "column": "sex", "type": "binary", "reference": 0}],
            [{"id": "table1", "skill": "table1", "status": "done", "depends_on": [],
              "output_file": "Table1"},
             {"id": "missingness", "skill": "missing-data", "status": "done",
              "depends_on": []},
             {"id": "km_primary", "skill": "kaplan-meier", "status": "done",
              "depends_on": [], "outcome": "primary", "output_file": "KM_primary"},
             {"id": "cr_primary", "skill": "competing-risk", "status": "done",
              "depends_on": [], "outcome": "primary", "output_file": "CR_primary"}])
        plan_path = _write_plan(tmp, plan)
        _generate(plan_path)
        jobs = [(f"{SKILL_ROOT}/table1/scripts/make_table1.py", "table1"),
                (KM_SCRIPT, "km_primary"), (CR_SCRIPT, "cr_primary")]
        for script, step in jobs:
            proc = _run([script, os.path.join(str(tmp), "configs", f"{step}_config.json")],
                        cwd=str(tmp))
            assert proc.returncode == 0, f"{step}: {proc.stdout[-1500:]}{proc.stderr[-1500:]}"
        TestFailsClosedWhenEvidenceIsMissing._write_uni(
            os.path.join(str(tmp), "Missingness.xlsx"), n=n, events=157)
        return tmp, plan_path

    @staticmethod
    def _clone(src, dst):
        import shutil
        shutil.copytree(str(src), str(dst))
        targets = [os.path.join(str(dst), "research_plan.json")]
        cfg = os.path.join(str(dst), "configs")
        targets += [os.path.join(cfg, f) for f in os.listdir(cfg)]
        for path in targets:
            with open(path, encoding="utf-8") as f:
                text = f.read()
            with open(path, "w", encoding="utf-8") as f:
                f.write(text.replace(str(src), str(dst)))
        return os.path.join(str(dst), "research_plan.json")

    def _verdict(self, plan_path, check):
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        lines = [x for x in proc.stderr.splitlines()
                 if f"] {check}" in x or f"[SKIP] {check}" in x]
        if any(x.lstrip().startswith("[QC_BLOCKING]") for x in lines):
            return "block", lines, proc
        if any(x.lstrip().startswith("[QC_WARNING]") for x in lines):
            return "warn", lines, proc
        if any(x.lstrip().startswith("[SKIP]") for x in lines):
            return "skip", lines, proc
        return "quiet", lines, proc

    # ── the clean run: nothing fires ────────────────────────────────────────

    @pytest.mark.parametrize("check", ["RES-008", "RES-012", "RES-013"])
    def test_honest_output_raises_nothing(self, survival_project, check):
        _tmp, plan_path = survival_project
        verdict, lines, _ = self._verdict(plan_path, check)
        assert verdict == "quiet", lines

    # ── RES-005 ─────────────────────────────────────────────────────────────

    def test_res005_repeated_subject_id_blocks(self, survival_project, tmp_path):
        src, _ = survival_project
        plan_path = self._clone(src, tmp_path / "dup")
        pd.DataFrame({"subject_id": [1, 2, 2, 3], "grp": [1, 0, 0, 1]}).to_csv(
            os.path.join(str(tmp_path / "dup"), "matched_cohort.csv"), index=False)
        verdict, lines, _ = self._verdict(plan_path, "RES-005")
        assert verdict == "block", lines
        assert any("subject_id" in x for x in lines)

    def test_res005_unique_ids_are_quiet(self, survival_project, tmp_path):
        src, _ = survival_project
        plan_path = self._clone(src, tmp_path / "uniq")
        pd.DataFrame({"subject_id": [1, 2, 3, 4], "grp": [1, 0, 0, 1]}).to_csv(
            os.path.join(str(tmp_path / "uniq"), "matched_cohort.csv"), index=False)
        verdict, lines, _ = self._verdict(plan_path, "RES-005")
        assert verdict != "block", lines

    def test_res005_records_the_unmet_precondition(self, survival_project):
        """No shipped skill writes a subject-level CSV. That is a skip with a
        reason, not a pass."""
        _tmp, plan_path = survival_project
        verdict, lines, _ = self._verdict(plan_path, "RES-005")
        assert verdict == "skip"
        assert any("no subject-level CSV" in x for x in lines)

    # ── RES-008 ─────────────────────────────────────────────────────────────

    def _edit(self, tmp_path, src, name, workbook, sheet, row, col, value):
        plan_path = self._clone(src, tmp_path / name)
        import openpyxl
        path = os.path.join(str(tmp_path / name), workbook)
        wb = openpyxl.load_workbook(path)
        wb[sheet].cell(row=row, column=col).value = value
        wb.save(path)
        return plan_path

    def test_res008_survival_cannot_increase(self, survival_project, tmp_path):
        src, _ = survival_project
        plan_path = self._edit(tmp_path, src, "rise", "KM_primary.xlsx",
                               "Survival Summary", 4, 6, "95.0% (90.0–98.0)")
        verdict, lines, _ = self._verdict(plan_path, "RES-008")
        assert verdict == "block", lines
        assert "cannot increase" in " ".join(lines)

    def test_res008_cif_plus_survival_cannot_exceed_one(self, survival_project, tmp_path):
        """Raised at the last time point so monotonicity still holds and this
        rule is the one under test."""
        src, _ = survival_project
        plan_path = self._edit(tmp_path, src, "sum", "CR_primary.xlsx",
                               "CIF Summary", 4, 7, "70.0% (64.0–76.0)")
        verdict, lines, _ = self._verdict(plan_path, "RES-008")
        assert verdict == "block", lines
        assert "exceeds 100%" in " ".join(lines)

    def test_res008_median_not_reached_needs_under_half_the_events(self, survival_project,
                                                                   tmp_path):
        src, _ = survival_project
        plan_path = self._clone(src, tmp_path / "median")
        import openpyxl
        path = os.path.join(str(tmp_path / "median"), "KM_primary.xlsx")
        wb = openpyxl.load_workbook(path)
        ws = wb["Survival Summary"]
        ws.cell(row=4, column=3).value = "160 (64.3%)"
        ws.cell(row=4, column=4).value = "NR"
        wb.save(path)
        verdict, lines, _ = self._verdict(plan_path, "RES-008")
        assert verdict == "block", lines
        assert "64.3% had the event" in " ".join(lines)

    # ── RES-012 ─────────────────────────────────────────────────────────────

    def test_res012_blank_figure_blocks(self, survival_project, tmp_path):
        src, _ = survival_project
        plan_path = self._clone(src, tmp_path / "blank")
        with open(os.path.join(str(tmp_path / "blank"), "KM_primary.png"), "wb") as f:
            f.write(b"\x89PNG\r\n\x1a\n" + b"\0" * 200)
        verdict, lines, _ = self._verdict(plan_path, "RES-012")
        assert verdict == "block", lines
        assert any("blank canvas" in x for x in lines)

    def test_res012_says_it_is_not_a_visual_check(self, survival_project, tmp_path):
        """A reader who thinks this verified the picture has been misled."""
        src, _ = survival_project
        plan_path = self._clone(src, tmp_path / "blank2")
        with open(os.path.join(str(tmp_path / "blank2"), "KM_primary.png"), "wb") as f:
            f.write(b"\x89PNG\r\n\x1a\n" + b"\0" * 200)
        _verdict, lines, _ = self._verdict(plan_path, "RES-012")
        assert any("does not check what the figure shows" in x for x in lines)

    # ── RES-013 ─────────────────────────────────────────────────────────────

    def test_res013_group_columns_must_sum_to_the_total(self, survival_project, tmp_path):
        src, _ = survival_project
        plan_path = self._edit(tmp_path, src, "sumn", "Table1.xlsx", "Table 1",
                               3, 3, "Exposed\n(n=149)")
        verdict, lines, _ = self._verdict(plan_path, "RES-013")
        assert verdict == "block", lines
        assert "does not account for" in " ".join(lines)

    def test_res013_percentage_must_match_its_own_denominator(self, survival_project,
                                                              tmp_path):
        src, _ = survival_project
        plan_path = self._edit(tmp_path, src, "pct", "Table1.xlsx", "Table 1",
                               7, 2, "261 (87.0%)")
        verdict, lines, _ = self._verdict(plan_path, "RES-013")
        assert verdict == "warn", lines
        assert "implies 52.2%" in " ".join(lines)

    # ── message / condition consistency ─────────────────────────────────────

    def test_every_reported_number_matches_its_stated_condition(self, survival_project,
                                                                tmp_path):
        """Parse the numbers back out of the messages and re-check them.

        Same guard as the missingness bands: a message that states a figure the
        condition does not actually test is how a reader concludes something was
        checked when it was not.
        """
        import re as _re

        src, _ = survival_project
        plan_path = self._edit(tmp_path, src, "parseback", "CR_primary.xlsx",
                               "CIF Summary", 4, 7, "70.0% (64.0–76.0)")
        _verdict, _lines, proc = self._verdict(plan_path, "RES-008")
        checked = 0
        for line in proc.stderr.splitlines():
            m = _re.search(r"CIF (\d+(?:\.\d+)?)% \+ survival (\d+(?:\.\d+)?)% = "
                           r"(\d+(?:\.\d+)?)%, which exceeds 100%", line)
            if m:
                cif, surv, total = (float(m.group(i)) for i in (1, 2, 3))
                assert abs(cif + surv - total) < 0.05, line
                assert total > 100.0, line
                checked += 1
            m = _re.search(r"is (\d+) bytes \(<10 KB\)", line)
            if m:
                assert int(m.group(1)) < 10 * 1024, line
                checked += 1
            m = _re.search(r"hold (\d+) subjects but the total column says (\d+); "
                           r"Table 1 does not account for (\d+)", line)
            if m:
                held, total, gap = (int(m.group(i)) for i in (1, 2, 3))
                assert abs(total - held) == gap, line
                checked += 1
        assert checked >= 1, "no message with an arithmetic claim was produced"


@slow
class TestCrossLanguageGrounding:
    """RES-014 — the only external grounding on the default path.

    Everything else compares the pipeline against its own inputs. This re-fits
    in R and compares, which is the one check that can catch an error the
    Python side makes consistently.
    """

    @pytest.fixture(scope="class")
    def cox_project(self, tmp_path_factory):
        tmp = tmp_path_factory.mktemp("xlang")
        rng = np.random.default_rng(21)
        n = 600
        age = rng.normal(65, 10, n)
        sex = rng.integers(0, 2, n)
        t = rng.exponential(1 / (0.15 * np.exp(0.03 * (age - 65) + 0.5 * sex)))
        event = (t <= 8).astype(float)
        pd.DataFrame({"exposed": rng.integers(0, 2, n).astype(float), "age": age,
                      "sex": sex.astype(float), "time": np.minimum(t, 8.0).round(4),
                      "event": event}).to_csv(tmp / "data.csv", index=False)
        plan = _base_plan(
            tmp, tmp / "data.csv",
            {"id": "primary", "name": "Death", "type": "time_to_event",
             "time_column": "time", "event_column": "event",
             "time_unit": "years", "time_scale": 1},
            [{"name": "Male", "column": "sex", "type": "binary", "reference": 0},
             {"name": "Age", "column": "age", "type": "continuous"}],
            [{"id": "missingness", "skill": "missing-data", "status": "done",
              "depends_on": []},
             {"id": "uni_primary", "skill": "univariate-forest", "status": "done",
              "depends_on": [], "outcome": "primary", "output_file": "Uni_primary"}],
            guardian={"cross_lang_check": True})
        plan_path = _write_plan(tmp, plan)
        _generate(plan_path)
        proc = _run([UNI_SCRIPT, os.path.join(str(tmp), "configs",
                                              "uni_primary_config.json")], cwd=str(tmp))
        assert proc.returncode == 0
        TestFailsClosedWhenEvidenceIsMissing._write_uni(
            os.path.join(str(tmp), "Missingness.xlsx"), n=600, events=466)
        return tmp, plan_path

    @staticmethod
    def _res014(plan_path, env=None):
        environ = dict(os.environ)
        environ.update(env or {})
        proc = subprocess.run([sys.executable, PRECHECK, plan_path, "--stage", "post"],
                              capture_output=True, text=True, env=environ,
                              timeout=SUBPROCESS_TIMEOUT)
        return [x for x in proc.stderr.splitlines()
                if "RES-014" in x and not x.strip().startswith("RES-014:")]

    def test_r_and_python_agree_on_an_untouched_run(self, cox_project):
        _tmp, plan_path = cox_project
        lines = self._res014(plan_path)
        assert lines, "RES-014 produced no line at all"
        assert any(x.startswith("[QC_INFO]") and "agree" in x for x in lines), lines

    def test_missing_r_is_an_explicit_skip_not_a_pass(self, cox_project):
        """The hard rule: a check that could not run never reports success."""
        _tmp, plan_path = cox_project
        lines = self._res014(plan_path, env={"PATH": "/nonexistent-bin"})
        assert any(x.strip().startswith("[SKIP] RES-014") for x in lines), lines
        assert any("Rscript is not on PATH" in x for x in lines)
        assert not any(x.startswith("[QC_INFO]") and "agree" in x for x in lines), \
            "a run without R claimed the two languages agreed"

    def _tamper(self, tmp_path, src, name, replacement):
        plan_path = TestRoleThreeChecksRestored._clone(src, tmp_path / name)
        import openpyxl
        path = os.path.join(str(tmp_path / name), "Uni_primary.xlsx")
        wb = openpyxl.load_workbook(path)
        for row in wb["Univariate Analysis"].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("1.45 ("):
                    cell.value = replacement
        wb.save(path)
        return plan_path

    def test_a_large_disagreement_blocks_and_asks_for_adjudication(self, cox_project,
                                                                    tmp_path):
        src, _ = cox_project
        plan_path = self._tamper(tmp_path, src, "big", "1.75 (1.45–2.11)")
        lines = self._res014(plan_path)
        blocking = [x for x in lines if x.startswith("[QC_BLOCKING]")]
        assert blocking, lines
        assert "coding error or a variable-definition mismatch" in blocking[0]

    def test_a_small_disagreement_warns_and_names_the_two_usual_causes(self, cox_project,
                                                                        tmp_path):
        """1-5% is the tie-handling / reference-level band, per results-audit.md."""
        src, _ = cox_project
        plan_path = self._tamper(tmp_path, src, "small", "1.49 (1.24–1.79)")
        lines = self._res014(plan_path)
        warnings = [x for x in lines if x.startswith("[QC_WARNING]")]
        assert warnings, lines
        assert "tied-event handling" in warnings[0]
        assert "reference level" in warnings[0]
        assert not [x for x in lines if x.startswith("[QC_BLOCKING]")]

    def test_the_escalation_evidence_is_machine_readable(self, cox_project, tmp_path):
        src, _ = cox_project
        plan_path = self._tamper(tmp_path, src, "evidence", "1.75 (1.45–2.11)")
        self._res014(plan_path)
        with open(os.path.join(str(tmp_path / "evidence"), "precheck_report.json"),
                  encoding="utf-8") as f:
            issues = json.load(f)["issues"]
        found = [i for i in issues if i["check_id"] == "RES-014"
                 and i["evidence"].get("type") == "cross_lang_repro"]
        assert found, "no cross_lang_repro evidence was recorded for an agent to read"
        evidence = found[0]["evidence"]
        for key in ("main_HR_python", "alt_HR_R", "rel_diff_pct",
                    "python_n_events", "R_n_events", "term"):
            assert key in evidence, key


class TestCrossFileJoints:
    """Assumptions that hold between files, not inside one.

    Each of these was true by coincidence rather than by contract: a fixed
    output path that nothing proved was written this run, a sidecar lookup that
    would settle for any sidecar, and an id built independently on both sides of
    a join.
    """

    # ── the PH gate must never read a report it did not just produce ────────

    def _cox_plan(self, tmp_path, gate_path=None):
        return TestPhRemedyGateDelegation()._cox_scenario(
            tmp_path, dict(PH_VIOLATED, declared_remedy="stratify"))

    def test_a_stale_report_is_not_mistaken_for_this_run(self, tmp_path):
        """A previous "passed" report plus a killed child must block.

        The gate is replaced with one that writes nothing and dies on a signal,
        while a passing report sits at the path the old code read from.
        """
        import precheck_configs as PC

        plan_path = self._cox_plan(tmp_path)
        with open(plan_path, encoding="utf-8") as f:
            plan = json.load(f)
        guardian_dir = os.path.join(str(tmp_path), "guardian")
        os.makedirs(guardian_dir, exist_ok=True)
        stale = os.path.join(guardian_dir, "ph_remedy_gate.json")
        with open(stale, "w", encoding="utf-8") as f:
            json.dump({"plan": plan_path, "output_dir": str(tmp_path),
                       "n_ph_artifacts": 1, "n_blocking": 0, "n_warning": 0,
                       "findings": []}, f)

        killer = os.path.join(str(tmp_path), "killed_gate.py")
        with open(killer, "w", encoding="utf-8") as f:
            f.write("import os, signal\nos.kill(os.getpid(), signal.SIGKILL)\n")

        report = PC.Report(quiet=True)
        original = PC.PH_REMEDY_GATE
        PC.PH_REMEDY_GATE = killer
        try:
            PC.check_ph_remedy_gate(report, plan, plan_path, plan["pipeline"],
                                    plan["output_dir"])
        finally:
            PC.PH_REMEDY_GATE = original

        assert report.n_blocking == 1, report.issues
        assert "did not produce a usable result" in report.issues[0]["message"]
        assert not os.path.exists(stale), \
            "the stale report survived and can still be mistaken for this run"

    def test_a_report_about_another_plan_is_refused(self, tmp_path):
        import precheck_configs as PC

        plan_path = self._cox_plan(tmp_path)
        with open(plan_path, encoding="utf-8") as f:
            plan = json.load(f)
        impostor = os.path.join(str(tmp_path), "impostor_gate.py")
        with open(impostor, "w", encoding="utf-8") as f:
            f.write(
                "import json, sys\n"
                "out = sys.argv[sys.argv.index('--json') + 1]\n"
                "json.dump({'plan': '/somewhere/else/research_plan.json',\n"
                "           'output_dir': '/somewhere/else', 'n_ph_artifacts': 1,\n"
                "           'n_blocking': 0, 'n_warning': 0, 'findings': []},\n"
                "          open(out, 'w'))\n")
        report = PC.Report(quiet=True)
        original = PC.PH_REMEDY_GATE
        PC.PH_REMEDY_GATE = impostor
        try:
            PC.check_ph_remedy_gate(report, plan, plan_path, plan["pipeline"],
                                    plan["output_dir"])
        finally:
            PC.PH_REMEDY_GATE = original
        assert report.n_blocking == 1, report.issues
        assert "not the plan being checked" in report.issues[0]["message"]

    def test_an_unknown_schema_is_refused(self, tmp_path):
        import precheck_configs as PC

        plan_path = self._cox_plan(tmp_path)
        with open(plan_path, encoding="utf-8") as f:
            plan = json.load(f)
        old_version = os.path.join(str(tmp_path), "old_gate.py")
        with open(old_version, "w", encoding="utf-8") as f:
            f.write("import json, sys\n"
                    "out = sys.argv[sys.argv.index('--json') + 1]\n"
                    "json.dump({'issues': []}, open(out, 'w'))\n")
        report = PC.Report(quiet=True)
        original = PC.PH_REMEDY_GATE
        PC.PH_REMEDY_GATE = old_version
        try:
            PC.check_ph_remedy_gate(report, plan, plan_path, plan["pipeline"],
                                    plan["output_dir"])
        finally:
            PC.PH_REMEDY_GATE = original
        assert report.n_blocking == 1, report.issues
        assert "not the schema this check knows how to read" in report.issues[0]["message"]

    def test_exit_code_and_report_must_agree(self, tmp_path):
        import precheck_configs as PC

        plan_path = self._cox_plan(tmp_path)
        with open(plan_path, encoding="utf-8") as f:
            plan = json.load(f)
        liar = os.path.join(str(tmp_path), "liar_gate.py")
        with open(liar, "w", encoding="utf-8") as f:
            f.write("import json, sys\n"
                    "out = sys.argv[sys.argv.index('--json') + 1]\n"
                    "json.dump({'plan': sys.argv[sys.argv.index('--plan') + 1],\n"
                    "           'output_dir': sys.argv[sys.argv.index('--output-dir') + 1],\n"
                    "           'n_ph_artifacts': 1, 'n_blocking': 0, 'n_warning': 0,\n"
                    "           'findings': []}, open(out, 'w'))\n"
                    "sys.exit(1)\n")
        report = PC.Report(quiet=True)
        original = PC.PH_REMEDY_GATE
        PC.PH_REMEDY_GATE = liar
        try:
            PC.check_ph_remedy_gate(report, plan, plan_path, plan["pipeline"],
                                    plan["output_dir"])
        finally:
            PC.PH_REMEDY_GATE = original
        assert report.n_blocking == 1, report.issues
        assert "exit code and report disagree" in report.issues[0]["message"]

    def test_the_real_gate_still_passes_a_clean_plan(self, tmp_path):
        """The hardening must not break the working path."""
        plan_path = TestPhRemedyGateDelegation()._cox_scenario(tmp_path, PH_CLEAN)
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        lines = [x for x in proc.stderr.splitlines() if "RES-015-det" in x]
        assert not any(x.lstrip().startswith("[QC_BLOCKING]") for x in lines), lines
        assert os.path.exists(os.path.join(str(tmp_path), "guardian",
                                           "ph_remedy_gate.json"))

    # ── a sidecar must be bound to the table it describes ───────────────────

    def _two_psm_scenario(self, tmp_path):
        """A primary and a sensitivity PSM step; only the primary gets a sidecar."""
        helper = TestAdjudicatedEpvAndPsmDirection()
        steps = [{"id": "psm", "skill": "psm-matching", "status": "done",
                  "depends_on": [], "output_file": "PSM"},
                 {"id": "psm_sens", "skill": "psm-matching", "status": "done",
                  "depends_on": [], "output_file": "PSM_sensitivity"}]
        plan_path = helper._scenario(
            tmp_path, extra_steps=steps, data=make_crude_or(2.4),
            covariates=[{"name": "Age", "column": "age", "type": "continuous"}])
        # gen_psm hardcodes output_file to "PSM.xlsx", so two psm-matching steps
        # are handed the same path and would overwrite each other. Point the
        # sensitivity step at its own workbook, which is what a plan running two
        # matched analyses has to do anyway.
        sens_config = os.path.join(str(tmp_path), "configs", "psm_sens_config.json")
        with open(sens_config, encoding="utf-8") as f:
            cfg = json.load(f)
        cfg["output_file"] = os.path.join(str(tmp_path), "PSM_sensitivity.xlsx")
        with open(sens_config, "w", encoding="utf-8") as f:
            json.dump(cfg, f)
        for table in ("PSM.xlsx", "PSM_sensitivity.xlsx"):
            helper._write_psm(os.path.join(str(tmp_path), table), "Exposed", "Unexposed",
                              effect="2.24 (1.30–3.85)",
                              exposed_counts=(120, 60), control_counts=(120, 35))
        helper._write_sidecar(tmp_path, exposed_value=1.0, exposed=(120, 60),
                              control=(120, 35), table="PSM.xlsx")
        return plan_path

    def test_a_generic_sidecar_is_refused_for_another_table(self, tmp_path):
        """Two PSM steps, one sidecar. The unbound one must not stand in."""
        plan_path = self._two_psm_scenario(tmp_path)
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        blocking = [x for x in proc.stderr.splitlines()
                    if x.startswith("[QC_BLOCKING] PSM-DIR psm_sens")]
        assert blocking, proc.stderr[-2500:]
        assert "not known to describe this table" in blocking[0]
        assert proc.returncode == 1

    def test_the_bound_table_in_the_same_run_still_passes(self, tmp_path):
        """The same run's correctly-bound step must not be caught by the above."""
        plan_path = self._two_psm_scenario(tmp_path)
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "producer sidecar agrees with the table" in proc.stderr
        assert not [x for x in proc.stderr.splitlines()
                    if x.startswith("[QC_BLOCKING] PSM-DIR psm:")], proc.stderr[-2000:]

    def test_a_sidecar_for_a_different_table_blocks(self, tmp_path):
        helper = TestAdjudicatedEpvAndPsmDirection()
        plan_path = helper._psm_scenario(
            tmp_path, data=make_crude_or(2.4), effect="2.24 (1.30–3.85)",
            exposed_counts=(120, 60), control_counts=(120, 35))
        helper._write_sidecar(tmp_path, exposed_value=1.0, exposed=(120, 60),
                              control=(120, 35))
        path = os.path.join(str(tmp_path), "PSM_contrast.json")
        with open(path, encoding="utf-8") as f:
            side = json.load(f)
        side["table_basename"] = "SomeOtherAnalysis.xlsx"
        with open(path, "w", encoding="utf-8") as f:
            json.dump(side, f)
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "says it describes 'SomeOtherAnalysis.xlsx'" in proc.stderr

    def test_a_sidecar_for_a_stale_version_of_the_table_blocks(self, tmp_path):
        """The table was regenerated after the sidecar was written."""
        helper = TestAdjudicatedEpvAndPsmDirection()
        plan_path = helper._psm_scenario(
            tmp_path, data=make_crude_or(2.4), effect="2.24 (1.30–3.85)",
            exposed_counts=(120, 60), control_counts=(120, 35))
        helper._write_sidecar(tmp_path, exposed_value=1.0, exposed=(120, 60),
                              control=(120, 35))
        helper._write_psm(os.path.join(str(tmp_path), "PSM.xlsx"), "Exposed", "Unexposed",
                          effect="2.30 (1.35–3.92)",
                          exposed_counts=(120, 60), control_counts=(120, 35))
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert proc.returncode == 1
        assert "written for a different version of PSM.xlsx" in proc.stderr

    def test_an_unbound_legacy_sidecar_is_not_used_to_verify(self, tmp_path):
        """Backward compatible, but it must not silently vouch for anything."""
        helper = TestAdjudicatedEpvAndPsmDirection()
        plan_path = helper._psm_scenario(
            tmp_path, data=make_crude_or(2.4), effect="2.24 (1.30–3.85)",
            exposed_counts=(120, 60), control_counts=(120, 35))
        with open(os.path.join(str(tmp_path), "PSM_contrast.json"), "w",
                  encoding="utf-8") as f:
            json.dump({"exposed_label": "Exposed", "exposed_value": 1.0,
                       "exposed": {"n": 120, "events": 60},
                       "control": {"n": 120, "events": 35}}, f)
        proc = _run([PRECHECK, plan_path, "--stage", "post"])
        assert "predates the binding fields" in proc.stderr
        assert "was NOT used to verify the contrast" in proc.stderr
        assert "producer sidecar agrees with the table" not in proc.stderr

    # ── one estimate_id builder, shared across the join ─────────────────────

    def test_run_pipeline_and_the_inventory_produce_identical_ids(self, tmp_path):
        """The equality contract, on real ids rather than two format strings.

        Turns red the moment either side changes shape — which is the point,
        because the two used to be independent and agreed only by luck.
        """
        import estimate_inventory
        import run_pipeline

        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Univariate Analysis"
        ws.append(["Variable", "N (events)", "OR (95% CI)", "p-value"])
        ws.append(["Age (per 10)", "400 (120)", "2.22 (1.29–3.82)", "0.004*"])
        ws.append(["Male, n (%)", "400 (120)", "1.05 (1.01–1.09)", "0.012*"])
        wb.save(os.path.join(str(tmp_path), "Uni_primary.xlsx"))

        step = {"id": "uni_primary", "output_file": "Uni_primary"}
        ledger, error = run_pipeline.step_estimate_inventory(step, str(tmp_path))
        assert error is None, error
        assert ledger["ids"], "the ledger recorded no estimates to join on"

        records, _counters = estimate_inventory.collect_executed(str(tmp_path), quiet=True)
        inventory_ids = [r["estimate_id"] for r in records
                         if r["source"] == "Uni_primary.xlsx"]

        assert ledger["ids"] == inventory_ids, (
            "run_pipeline's ledger ids and coverage_diff.json's ids have diverged; "
            "a join between the two would silently return nothing\n"
            f"  ledger:    {ledger['ids']}\n  inventory: {inventory_ids}")

    # ── coverage_diff.json provenance ───────────────────────────────────────

    def test_coverage_diff_carries_provenance(self, tmp_path):
        """C7 can only tell a stale audit from a current one if the audit says
        which bytes it covered."""
        helper = TestCoverageInventoryIsProduced()
        plan_path = helper._scenario(
            tmp_path, quote="Age was associated with mortality (OR 2.22).")
        _run([PRECHECK, plan_path, "--stage", "post"])
        with open(os.path.join(str(tmp_path), "coverage_diff.json"),
                  encoding="utf-8") as f:
            payload = json.load(f)
        prov = payload["provenance"]
        assert prov["hash_algorithm"] == "sha256"
        assert prov["generated_at"]
        assert prov["plan_fingerprint"], prov.get("plan_fingerprint_error")
        assert prov["plan_fingerprint_error"] is None
        sources = {r["source"] for r in payload["executed"]}
        assert sources, "nothing was scanned"
        assert sources <= set(prov["source_hashes"]), \
            "an executed source has no hash, so C7 cannot check it"
        assert all(v and len(v) == 64 for v in prov["source_hashes"].values())
        assert "Results_Section.docx" in prov["deliverable_hashes"]

    def test_the_fingerprint_is_reviewer_challenges_not_a_local_copy(self, tmp_path):
        import estimate_inventory
        import reviewer_challenge

        helper = TestCoverageInventoryIsProduced()
        plan_path = helper._scenario(tmp_path, quote="OR 2.22.")
        with open(plan_path, encoding="utf-8") as f:
            plan = json.load(f)
        computed, error = estimate_inventory.plan_fingerprint_for(plan)
        assert error is None, error
        assert computed == reviewer_challenge.plan_fingerprint(plan)

    def test_editing_the_plan_changes_the_recorded_fingerprint(self, tmp_path):
        import estimate_inventory

        helper = TestCoverageInventoryIsProduced()
        plan_path = helper._scenario(tmp_path, quote="OR 2.22.")
        with open(plan_path, encoding="utf-8") as f:
            plan = json.load(f)
        before, _ = estimate_inventory.plan_fingerprint_for(plan)
        plan["covariates"].append({"name": "Extra", "column": "age",
                                   "type": "continuous"})
        after, _ = estimate_inventory.plan_fingerprint_for(plan)
        assert before != after, "a changed analysis produced the same fingerprint"

    def test_a_rebuilt_workbook_changes_its_hash_even_with_an_older_mtime(self,
                                                                          tmp_path):
        """The reason this is a hash and not an mtime.

        The workbook is rewritten and its mtime back-dated, which is exactly the
        state an mtime comparison reports as "not stale".
        """
        import openpyxl

        helper = TestCoverageInventoryIsProduced()
        plan_path = helper._scenario(tmp_path, quote="OR 2.22.")
        _run([PRECHECK, plan_path, "--stage", "post"])
        cov = os.path.join(str(tmp_path), "coverage_diff.json")
        with open(cov, encoding="utf-8") as f:
            before = json.load(f)["provenance"]["source_hashes"]["Uni_primary.xlsx"]

        workbook = os.path.join(str(tmp_path), "Uni_primary.xlsx")
        wb = openpyxl.load_workbook(workbook)
        ws = wb[wb.sheetnames[0]]
        ws.cell(row=2, column=3).value = "2.99 (1.70–5.26)"
        wb.save(workbook)
        stat = os.stat(workbook)
        os.utime(workbook, (stat.st_atime - 99999, stat.st_mtime - 99999))

        _run([PRECHECK, plan_path, "--stage", "post"])
        with open(cov, encoding="utf-8") as f:
            after = json.load(f)["provenance"]["source_hashes"]["Uni_primary.xlsx"]
        assert before != after, \
            "the rebuilt workbook hashed the same; C7 would call a stale audit current"

    def test_a_missing_plan_records_why_rather_than_omitting_the_field(self, tmp_path):
        import estimate_inventory

        fingerprint, error = estimate_inventory.plan_fingerprint_for(None)
        assert fingerprint is None
        assert error and "no plan" in error

    def test_the_shared_builder_is_insensitive_to_the_workbook_spelling(self):
        """`output_file` may be declared with or without the extension."""
        from estimate_inventory import estimate_id

        record = {"sheet": "Univariate Analysis", "row": 4, "label": "Age (per 10)"}
        assert (estimate_id(record, "Uni_primary.xlsx")
                == estimate_id(record, "Uni_primary")
                == estimate_id(record, "/tmp/out/Uni_primary.xlsx"))

    def test_a_changed_id_shape_breaks_the_contract_test(self, monkeypatch):
        """The equality test is only worth having if it can fail."""
        import estimate_inventory

        record = {"sheet": "S", "row": 2, "label": "x"}
        baseline = estimate_inventory.estimate_id(record, "W.xlsx")
        monkeypatch.setattr(estimate_inventory, "estimate_id",
                            lambda r, w, i=0: baseline + "|v2")
        assert estimate_inventory.estimate_id(record, "W.xlsx") != baseline


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v"]))
