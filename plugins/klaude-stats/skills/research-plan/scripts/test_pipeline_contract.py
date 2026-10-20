#!/usr/bin/env python3
"""Smoke test for the 2026-08-11 run_pipeline.py / pipeline_status.py fixes."""
import importlib.util, json, os, subprocess, sys, tempfile, pathlib

SKILLS_DIR = pathlib.Path(__file__).resolve().parents[2]
SK = SKILLS_DIR / "research-plan/scripts"

def load(name):
    return load_from(SK / f"{name}.py", name)


def load_from(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m

rp = load("run_pipeline")
ps = load("pipeline_status")

fails = []
def check(label, got, want):
    ok = got == want
    print(f"{'PASS' if ok else 'FAIL'}  {label}: got={got!r} want={want!r}")
    if not ok: fails.append(label)

# ---- should_run ----------------------------------------------------------
done    = {"id": "a", "status": "done"}
pending = {"id": "b", "status": "pending"}
err     = {"id": "c", "status": "error"}
blocked = {"id": "d", "status": "blocked"}

check("done step NOT re-run when no flags", rp.should_run(done, None, None, False), False)
check("pending runs when no flags",         rp.should_run(pending, None, None, False), True)
check("error re-runs",                      rp.should_run(err, None, None, False), True)
check("blocked re-runs (unblock path)",     rp.should_run(blocked, None, None, False), True)
check("--only matches",                     rp.should_run(done, None, "a", False), True)
check("--only non-match",                   rp.should_run(pending, None, "a", False), False)
check("--from: before target skipped",      rp.should_run(pending, "z", None, False), False)
check("--from: at/after target redoes done",rp.should_run(done, "a", None, True), True)

# ---- QC extraction -------------------------------------------------------
stderr = ("  [QC_WARNING] leading whitespace case\n"
          "[QC_BLOCKING] zero events\n"
          "[GUARDIAN_WARNING] guardian note\n"
          "irrelevant line\n")
w = rp.extract_qc_warnings(stderr)
check("QC extract count", len(w), 3)
check("leading whitespace stripped+kept", w[0], "[QC_WARNING] leading whitespace case")

step = {"id": "x"}
rp._record_qc(step, stderr)
rp._record_qc(step, stderr)          # idempotent
check("qc de-duplicated across phases", len(step["qc_warnings"]), 3)

# ---- status icons / blocking markers ------------------------------------
check("skipped icon exists", "skipped" in ps.STATUS_ICONS, True)
check("blocked icon exists", "blocked" in ps.STATUS_ICONS, True)
check("GUARDIAN_BLOCKING is blocking",
      any(m in "[GUARDIAN_BLOCKING] x" for m in ps.BLOCKING_MARKERS), True)

# ---- --detect-only end to end -------------------------------------------
with tempfile.TemporaryDirectory() as d:
    out = os.path.join(d, "out"); os.makedirs(out)
    open(os.path.join(out, "Table1.xlsx"), "w").write("x")   # orphan output
    plan = {
        "output_dir": out,
        "total_steps_executed": 7,
        "pipeline": [
            {"id": "flowchart", "skill": "flowchart", "status": "done",
             "output_file": "Flowchart", "depends_on": []},
            {"id": "table1", "skill": "table1", "status": "pending",
             "output_file": "Table1", "depends_on": []},
        ],
    }
    pp = os.path.join(d, "research_plan.json")
    json.dump(plan, open(pp, "w"))
    r = subprocess.run([sys.executable, str(SK / "run_pipeline.py"), pp, "--detect-only"],
                       capture_output=True, text=True)
    check("--detect-only exit 0", r.returncode, 0)
    check("--detect-only reports would-run", "Would run:         1/2" in r.stdout, True)
    check("--detect-only flags orphan output", "Output files exist" in r.stdout, True)
    check("--detect-only executes nothing",
          json.load(open(pp))["pipeline"][1]["status"], "pending")

    # bad --from must fail loudly, not silently run everything
    r2 = subprocess.run([sys.executable, str(SK / "run_pipeline.py"), pp, "--from", "nope"],
                        capture_output=True, text=True)
    check("unknown --from exits 2", r2.returncode, 2)

    # cumulative counter is seeded from the plan
    check("total_steps_executed seeded from plan",
          json.load(open(pp)).get("total_steps_executed"), 7)


# ---- process_feedback backup versioning ---------------------------------
pf = load("process_feedback")
with tempfile.TemporaryDirectory() as d:
    out = os.path.join(d, "out"); os.makedirs(out)
    open(os.path.join(out, "Table1.xlsx"), "w").write("v1 results")
    pp = os.path.join(d, "research_plan.json")
    json.dump({"version": 1, "output_dir": out, "pipeline": []}, open(pp, "w"))

    check("backup round 1 returns v1", pf.backup_version(pp), 1)
    p = json.load(open(pp))
    check("version bumped to 2", p["version"], 2)
    check("version_history appended", len(p.get("version_history", [])), 1)
    check("pipeline_state=feedback", p.get("pipeline_state"), "feedback")
    check("v1 snapshot on disk", os.path.exists(os.path.join(d, "research_plan_v1.json")), True)

    # round 2 must snapshot separately, not silently skip
    check("backup round 2 returns v2", pf.backup_version(pp), 2)
    check("v2 snapshot on disk", os.path.exists(os.path.join(d, "research_plan_v2.json")), True)
    check("round-1 snapshot still intact",
          json.load(open(os.path.join(d, "research_plan_v1.json")))["version"], 1)

    # collision guard: a stale snapshot at the current version must abort, not overwrite
    json.dump({"version": 1, "output_dir": out, "pipeline": []}, open(pp, "w"))
    code = (f"import importlib.util;"
            f"spec=importlib.util.spec_from_file_location('pf', r'{SK}/process_feedback.py');"
            f"m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);"
            f"m.backup_version(r'{pp}')")
    r = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True)
    check("stale snapshot aborts with exit 3", r.returncode, 3)


# ---- missing-data steps: config generation + group routing -----------------
gc = load("generate_configs")


# A stand-in for the survival-diagnostics report. It has to exist on disk for
# diagnostics.mode="completed" to be believed; kept for the whole session
# because ~100 checks below build plans that reference it.
_DIAG_DIR = tempfile.mkdtemp(prefix="rp-contract-diag-")
_DIAG_REPORT = os.path.join(_DIAG_DIR, "diagnostic_report.json")
with open(_DIAG_REPORT, "w", encoding="utf-8") as _f:
    json.dump({"checks": {}}, _f)


def _synth_plan(primary_type, missing_data=None):
    """Minimal plan exercising the two missing-data generators."""
    plan = {
        "study": {"title": "Synthetic study", "design": "cohort",
                  "data_file": "/tmp/data.xlsx"},
        "population": {"filters": [{"column": "HF_type", "operator": "==", "value": 3}]},
        "derived_variables": [{"name": "BMI", "type": "expression",
                               "expression": "BW / (Ht/100)**2", "dependencies": ["BW", "Ht"]}],
        "grouping": {"column": "AF", "groups": [{"value": 1, "label": "AF"},
                                                {"value": 0, "label": "No AF"}]},
        "outcomes": {
            "primary": {"id": "primary", "name": "MACE", "type": primary_type,
                        "time_column": "t_mace", "event_column": "e_mace"},
            "secondary": [{"id": "secondary_1", "name": "Death", "type": "binary",
                           "event_column": "e_death"}],
        },
        "covariates": [
            {"name": "Age", "column": "age", "type": "continuous"},
            {"name": "Male sex", "column": "sex", "type": "binary", "reference": 0},
        ],
        "output_dir": "/tmp/out",
        "abbreviations": "AF = atrial fibrillation",
        # 2026-08-11: Step 3.5's findings became plan-time gates, so a plan that
        # fits a survival model has to say what the diagnostics found. An empty
        # flags_detected means "they ran and found nothing", which is a
        # different statement from an absent block ("nobody has looked") —
        # _validate_diagnostics_honored refuses the second one on purpose.
        # 2026-08-12: and "they ran" now has to be attributable to a run, or a
        # hand-written {"mode": "completed", "flags_detected": []} would be a
        # clean bill of health nobody issued. Hence report_path + completed_at.
        "diagnostics": {"mode": "completed", "flags_detected": [],
                        "report_path": _DIAG_REPORT,
                        "completed_at": "2099-01-01T00:00:00"},
    }
    if missing_data is not None:
        plan["missing_data"] = missing_data
    return plan


ASSESS_KEYS = {"mode", "data_file", "filters", "derived_variables", "grouping",
               "analysis_variables", "outcomes", "candidate_auxiliary",
               "thresholds", "output_file", "plot_file"}
MI_KEYS = {"mode", "data_file", "filters", "derived_variables", "grouping",
           "analysis_type", "cox_settings", "logistic_settings", "sections",
           "auxiliary_variables", "mi_settings", "cluster_confirmed", "outcome_id",
           "title", "subtitle", "abbreviations", "output_file", "diagnostics_plot"}

a = gc.gen_missingness(_synth_plan("time_to_event",
                                   {"auxiliary_variables": ["NTproBNP"]}), {"id": "missingness"})
check("assessment config has exactly the contract keys", set(a), ASSESS_KEYS)
check("assessment mode", a["mode"], "assessment")
check("assessment analysis_variables from covariates",
      a["analysis_variables"],
      [{"column": "age", "type": "continuous", "label": "Age"},
       {"column": "sex", "type": "binary", "label": "Male sex"}])
check("assessment covers primary + secondary outcomes",
      [o["id"] for o in a["outcomes"]], ["primary", "secondary_1"])
check("assessment candidate_auxiliary from plan", a["candidate_auxiliary"], ["NTproBNP"])
check("assessment thresholds", a["thresholds"], {"low": 5, "high": 40})
check("assessment output_file", a["output_file"], "/tmp/out/Missingness.xlsx")
check("assessment plot_file", a["plot_file"], "/tmp/out/Missingness_pattern.png")

# missing_data block absent entirely → still generates, with defaults
a0 = gc.gen_missingness(_synth_plan("binary"), {"id": "missingness"})
check("assessment works with no missing_data block", a0["candidate_auxiliary"], [])

mi_cox = gc.gen_missing_data_mi(_synth_plan("time_to_event"),
                                {"id": "mi_primary", "outcome": "primary"})
check("MI config has exactly the contract keys", set(mi_cox), MI_KEYS)
check("MI mode", mi_cox["mode"], "imputation")
check("MI time_to_event → cox", mi_cox["analysis_type"], "cox")
check("MI cox_settings from outcome", mi_cox["cox_settings"],
      {"time_column": "t_mace", "event_column": "e_mace"})
check("MI cox → logistic_settings null", mi_cox["logistic_settings"], None)
check("MI sections mirror build_sections",
      mi_cox["sections"], gc.build_sections(_synth_plan("time_to_event")["covariates"]))
check("MI default mi_settings", mi_cox["mi_settings"],
      {"m": None, "seed": 42, "maxit": 20, "method_overrides": {}})
check("MI output_file", mi_cox["output_file"], "/tmp/out/MI_primary.xlsx")
check("MI diagnostics_plot", mi_cox["diagnostics_plot"],
      "/tmp/out/MI_primary_diagnostics.png")

mi_log = gc.gen_missing_data_mi(
    _synth_plan("binary", {"m": 30, "seed": 7, "maxit": 10,
                           "method_overrides": {"age": "pmm"},
                           "auxiliary_variables": ["eGFR"]}),
    {"id": "mi_primary", "outcome": "primary"})
check("MI binary → logistic", mi_log["analysis_type"], "logistic")
check("MI logistic_settings from outcome", mi_log["logistic_settings"],
      {"outcome_column": "e_mace"})
check("MI logistic → cox_settings null", mi_log["cox_settings"], None)
check("MI mi_settings honour plan overrides", mi_log["mi_settings"],
      {"m": 30, "seed": 7, "maxit": 10, "method_overrides": {"age": "pmm"}})
check("MI auxiliary_variables from plan", mi_log["auxiliary_variables"], ["eGFR"])

# competing_risk: must still emit a config (script owns the [QC_BLOCKING]) but
# record a methodological warning so it cannot pass silently.
gc.METHOD_WARNINGS.clear()
mi_cr = gc.gen_missing_data_mi(_synth_plan("competing_risk"),
                               {"id": "mi_primary", "outcome": "primary"})
check("MI competing_risk still emits a config", mi_cr["mode"], "imputation")
check("MI competing_risk records a method warning",
      [w["kind"] for w in gc.METHOD_WARNINGS], ["mi_competing_risk_unsupported"])
gc.METHOD_WARNINGS.clear()

# GENERATORS dispatch: one skill, two step shapes
check("'missing-data' registered in GENERATORS", "missing-data" in gc.GENERATORS, True)
check("dispatch without outcome → assessment",
      gc.GENERATORS["missing-data"](_synth_plan("time_to_event"),
                                    {"id": "missingness"})["mode"], "assessment")
check("dispatch with outcome → imputation",
      gc.GENERATORS["missing-data"](_synth_plan("time_to_event"),
                                    {"id": "mi_primary", "outcome": "primary"})["mode"],
      "imputation")

# ---- missing-data: runner registration + group routing --------------------
check("SKILL_SCRIPTS has missing-data key", "missing-data" in rp.SKILL_SCRIPTS, True)
# The script itself is authored separately; assert the contract string only.
check("SKILL_SCRIPTS missing-data path/flag", rp.SKILL_SCRIPTS.get("missing-data"),
      (str(SKILLS_DIR / "research-plan/scripts/missing_data.py"), "--check-missing"))

_md_pipeline = [
    {"id": "table1", "skill": "table1"},
    {"id": "missingness", "skill": "missing-data"},
    {"id": "multi_primary", "skill": "multivariate-interaction", "outcome": "primary"},
    {"id": "mi_primary", "skill": "missing-data", "outcome": "primary"},
]
_md_groups = {g["group_id"]: [s["id"] for s in g["steps"]]
              for g in rp.group_steps_by_analysis_group(
                  _md_pipeline, [s["id"] for s in _md_pipeline])}
check("missingness lands in foundation", _md_groups.get("foundation"),
      ["table1", "missingness"])
check("mi_primary lands in outcome group", _md_groups.get("outcome_primary"),
      ["multi_primary", "mi_primary"])


# ---- clustering severity depends on provenance, not on a name guess --------
def _plan_with_diag(flags):
    p = _synth_plan("time_to_event", {"auxiliary_variables": []})
    if flags is not None:
        p["diagnostics"] = {"mode": "on", "flags_detected": flags}
    return p

_mi = lambda p: gc.gen_missing_data_mi(p, {"id": "mi_primary", "outcome": "primary"})
check("no diagnostics -> cluster_confirmed False",
      _mi(_plan_with_diag(None))["cluster_confirmed"], False)
check("flags as list",
      _mi(_plan_with_diag(["clustering", "time_varying"]))["cluster_confirmed"], True)
check("flags as {flag: bool}",
      _mi(_plan_with_diag({"clustering": True}))["cluster_confirmed"], True)
check("flags as {flag: {detected}}",
      _mi(_plan_with_diag({"clustering": {"detected": False}}))["cluster_confirmed"], False)
check("unrelated flags only",
      _mi(_plan_with_diag(["left_truncation"]))["cluster_confirmed"], False)
_p_off = _plan_with_diag(["clustering"]); _p_off["diagnostics"]["mode"] = "off"
check("diagnostics mode=off is not a confirmation",
      _mi(_p_off)["cluster_confirmed"], False)

# body weight must not be read as a propensity-score artefact
md = load("missing_data")
check("bare 'weight' not a PSM hint", any("weight" == h for h in md.PSM_HINT), False)
check("ps_weight still a PSM hint", "ps_weight" in md.PSM_HINT, True)
import pandas as _pd
check("body-weight filter raises no PSM flag",
      md.detect_psm({"filters": [{"column": "weight", "operator": ">=", "value": 40}]},
                    _pd.DataFrame({"weight": [60, 70], "age": [70, 80]})), None)
check("secondary MI carries its own outcome_id",
      gc.gen_missing_data_mi(_synth_plan("time_to_event"),
                             {"id": "mi_secondary_1", "outcome": "secondary_1"})["outcome_id"],
      "secondary_1")
check("primary MI carries outcome_id too",
      _mi(_plan_with_diag(None))["outcome_id"], "primary")

check("real PSM column still detected",
      md.detect_psm({"filters": []},
                    _pd.DataFrame({"ps_score": [0.1, 0.2]})) is not None, True)


# ---- advanced analyses: seven new steps ------------------------------------
ADV_SKILLS = ["nonlinearity", "robust-variance", "calibration", "weighting",
              "immortal-time", "evalue", "strobe"]
for _sk in ADV_SKILLS:
    check(f"GENERATORS has {_sk}", _sk in gc.GENERATORS, True)
    check(f"SKILL_SCRIPTS has {_sk}", _sk in rp.SKILL_SCRIPTS, True)

_adv_plan = _synth_plan("time_to_event")
# 2026-08-11: calibration now refuses to run unless the plan declares itself a
# prediction-model study (advanced_analyses.md section 5 always said so; nothing
# enforced it). This fixture exercises the calibration generator, so it has to
# make that declaration.
_adv_plan["study"]["model_purpose"] = "prediction"
_adv_plan["advanced_analyses"] = {
    "immortal_time": {"mode": "landmark", "exposure_time_column": "d2s",
                      "landmark_times": [30, 90]},
    "robust_variance": {"cluster_column": "hosp"},
    "evalue": {"reference_confounders": [{"name": "NYHA", "rr": 1.7}]},
}
_step = {"id": "x", "outcome": "primary"}
_modes = {"nonlinearity": "rcs", "robust-variance": "robust", "calibration": "calibration",
          "weighting": "weighting", "immortal-time": "landmark", "evalue": "evalue"}
for _sk, _mode in _modes.items():
    _c = gc.GENERATORS[_sk](_adv_plan, _step)
    check(f"{_sk} mode", _c["mode"], _mode)
    check(f"{_sk} routes to cox", _c["analysis_type"], "cox")

# a binary outcome must route to logistic everywhere, exactly as the older steps do
_bin = {"id": "x", "outcome": "secondary_1"}
for _sk in ("nonlinearity", "robust-variance", "calibration", "weighting"):
    check(f"{_sk} routes binary to logistic",
          gc.GENERATORS[_sk](_adv_plan, _bin)["analysis_type"], "logistic")

# RCS defaults to every continuous covariate rather than silently doing nothing
check("rcs defaults to all continuous covariates",
      gc.gen_nonlinearity(_adv_plan, _step)["spline_variables"], ["age"])

# missing cluster column and missing e-value benchmark must be recorded, not silent
gc.METHOD_WARNINGS.clear()
gc.gen_robust_variance(_synth_plan("time_to_event"), _step)
check("robust without cluster_column records a method warning",
      any(w["kind"] == "robust_variance_without_cluster_column" for w in gc.METHOD_WARNINGS),
      True)
_mixed_plan = _synth_plan("time_to_event")
_mixed_plan["advanced_analyses"] = {"robust_variance": {"cluster_column": "hosp",
                                                        "method": "mixed"}}
gc.METHOD_WARNINGS.clear()
# 2026-08-11: method='mixed' was withdrawn rather than warned about. The warning
# it used to raise did not stop the config being generated, so whoever chose it
# still got the number — from an approximation that widened the exposure SE by
# 4.8% where the cluster-robust sandwich widened it by 148% on the same
# clustering (advanced_analyses.md section 4), always in the direction that
# makes a finding look stronger.
_mixed_err = None
try:
    gc.gen_robust_variance(_mixed_plan, _step)
except ValueError as _e:
    _mixed_err = str(_e)
check("method='mixed' is refused, not merely flagged", _mixed_err is not None, True)
check("and the refusal routes to the methods that do work",
      all(k in (_mixed_err or "") for k in ("cluster_robust", "gee", "lme4")), True)
_mixed_plan["advanced_analyses"]["robust_variance"]["method"] = "cluster_robust"
gc.METHOD_WARNINGS.clear()
_cr_cfg = gc.gen_robust_variance(_mixed_plan, _step)
check("cluster_robust still generates — clustering support is not what was removed",
      _cr_cfg["method"], "cluster_robust")
_mixed_plan["advanced_analyses"]["robust_variance"]["method"] = "gee"
check("and so does gee", gc.gen_robust_variance(_mixed_plan, _step)["method"], "gee")

# 2026-08-12: the method list is a POSITIVE whitelist. Rejecting only `mixed`
# let every other unrecognised string through — and because the clustering gate
# asks whether a robust-variance STEP exists, not whether its method means
# anything, a misspelling satisfied that gate while the step corrected nothing.
for _bad in ("cluster_robsut", "sandwich", "robust", "clustered"):
    _mixed_plan["advanced_analyses"]["robust_variance"]["method"] = _bad
    try:
        gc.gen_robust_variance(_mixed_plan, _step)
        _bad_err = None
    except ValueError as _e:
        _bad_err = str(_e)
    check(f"method={_bad!r} is refused", _bad_err is not None, True)
    check("and the refusal lists what can be used instead",
          all(k in (_bad_err or "") for k in ("cluster_robust", "gee")), True)
check("a case variant is normalised to the canonical value, not passed through",
      gc._robust_method("GEE"), "gee")
check("and so is stray whitespace/case", gc._robust_method(" Cluster_Robust "),
      "cluster_robust")
# The gate itself, not just the generator: a step whose method is unrecognised
# does not count as covering the clustering finding.
_bypass = _synth_plan("time_to_event")
_bypass["diagnostics"] = dict(_bypass["diagnostics"], flags_detected=["clustering"])
_bypass["advanced_analyses"] = {"robust_variance": {"cluster_column": "hosp",
                                                    "method": "cluster_robsut"}}
_bypass["pipeline"] = [
    {"id": "missingness", "skill": "missing-data", "depends_on": []},
    {"id": "multi_primary", "skill": "multivariate-interaction", "outcome": "primary",
     "depends_on": []},
    {"id": "robust_primary", "skill": "robust-variance", "outcome": "primary",
     "depends_on": []},
]
try:
    gc.validate_plan(_bypass, "plan.json")
    _bypass_msg = None
except SystemExit as _e:
    _bypass_msg = str(_e)
check("a misspelt method cannot satisfy the clustering gate",
      _bypass_msg is not None, True)
check("and the gate says the step is present but inert",
      "Present but inert" in (_bypass_msg or ""), True)
check("naming the method it could not use",
      "cluster_robsut" in (_bypass_msg or ""), True)

gc.METHOD_WARNINGS.clear()
gc.gen_evalue(_synth_plan("time_to_event"), _step)
check("evalue without benchmark records a method warning",
      any(w["kind"] == "evalue_without_benchmark" for w in gc.METHOD_WARNINGS), True)
gc.METHOD_WARNINGS.clear()
_it = gc.gen_immortal_time(_adv_plan, _step)
check("immortal-time states the exposed arm explicitly",
      _it["grouping"]["exposed_value"], 1)
check("assumed exposed arm is recorded as a method warning",
      any(w["kind"] == "immortal_time_exposed_arm_assumed" for w in gc.METHOD_WARNINGS), True)
check("generator does not mutate the plan's grouping",
      "exposed_value" in _adv_plan["grouping"], False)
_adv_plan["advanced_analyses"]["immortal_time"]["exposed_value"] = 0
gc.METHOD_WARNINGS.clear()
_it0 = gc.gen_immortal_time(_adv_plan, _step)
check("explicit exposed_value wins", _it0["grouping"]["exposed_value"], 0)
check("explicit exposed_value raises no warning",
      any(w["kind"] == "immortal_time_exposed_arm_assumed" for w in gc.METHOD_WARNINGS), False)
del _adv_plan["advanced_analyses"]["immortal_time"]["exposed_value"]

gc.METHOD_WARNINGS.clear()
gc.gen_immortal_time(_adv_plan, _bin)
check("immortal-time on a binary outcome records a method warning",
      any(w["kind"] == "immortal_time_on_binary_outcome" for w in gc.METHOD_WARNINGS), True)
gc.METHOD_WARNINGS.clear()

# grouping, in both the runner and the status view
_adv_pipe = [{"id": "rcs_primary", "skill": "nonlinearity", "outcome": "primary", "depends_on": []},
             {"id": "evalue_primary", "skill": "evalue", "outcome": "primary", "depends_on": []},
             {"id": "strobe", "skill": "strobe", "depends_on": []},
             {"id": "manuscript", "skill": "manuscript-results", "depends_on": ["strobe"]}]
_g = {g["group_id"]: [s["id"] for s in g["steps"]]
      for g in rp.group_steps_by_analysis_group(_adv_pipe, rp.topo_sort(_adv_pipe))}
check("rcs -> outcome group", _g.get("outcome_primary"), ["rcs_primary"])
check("evalue -> supplementary", _g.get("supplementary"), ["evalue_primary"])
check("strobe -> manuscript group", _g.get("manuscript"), ["strobe", "manuscript"])

import io as _io, contextlib as _cl
_buf = _io.StringIO()
with _cl.redirect_stdout(_buf):
    ps.print_group_status({"pipeline": _adv_pipe})
check("status view files nothing under 'Other'", "Other" in _buf.getvalue(), False)

# E-value arithmetic against the published worked example (VanderWeele & Ding 2017)
ev = load("evalue")
check("E-value RR=3.9 -> 7.26", round(ev.evalue_from_rr(3.9), 2), 7.26)
check("E-value CI limit 1.8 -> 3.00", round(ev.evalue_for_ci(1.8, 8.7)[0], 2), 3.00)
check("E-value is symmetric about the null",
      round(ev.evalue_from_rr(2.0), 6), round(ev.evalue_from_rr(0.5), 6))
check("CI crossing the null gives E = 1", ev.evalue_for_ci(0.85, 1.4)[0], 1.0)
check("common-outcome OR is converted, not used raw",
      round(ev.to_risk_ratio(4.0, "OR", 0.30)[0], 4), 2.0)
check("rare-outcome OR passes through",
      round(ev.to_risk_ratio(4.0, "OR", 0.05)[0], 4), 4.0)

# ---- reviewer gate must not fail open ---------------------------------------
rc = load("reviewer_challenge")
_dom = [{"domain": k, "verdict": "fatal", "checked": "x", "comments": []}
        for k, _ in rc.DOMAINS]
_probs = []
for _d in _dom:
    import tempfile as _tf, json as _js, os as _os
    _fd, _pth = _tf.mkstemp(suffix=".json"); _os.close(_fd)
    _js.dump(_d, open(_pth, "w"))
    _r, _p = rc.load_report(_pth); _probs += _p; _os.unlink(_pth)
check("domain verdict=fatal with no comment is a schema problem",
      len(_probs), len(rc.DOMAINS))

_ok = {"domain": "confounding", "verdict": "pass", "checked": "DAG vs adjustment set",
       "comments": []}
import tempfile as _tf2, json as _js2, os as _os2
_fd, _pth = _tf2.mkstemp(suffix=".json"); _os2.close(_fd)
_js2.dump(_ok, open(_pth, "w"))
_r, _p = rc.load_report(_pth); _os2.unlink(_pth)
check("a genuine pass raises no schema problem", _p, [])

# ---- AR(1) needs a real time axis -------------------------------------------
_ar_plan = _synth_plan("time_to_event")
_ar_plan["advanced_analyses"] = {"robust_variance": {"cluster_column": "hosp",
                                                     "method": "gee",
                                                     "correlation_structure": "ar1"}}
gc.METHOD_WARNINGS.clear()
_c = gc.gen_robust_variance(_ar_plan, _step)
check("AR(1) without a time column is flagged",
      any(w["kind"] == "gee_ar1_without_time_column" for w in gc.METHOD_WARNINGS), True)
_ar_plan["advanced_analyses"]["robust_variance"]["within_cluster_time_column"] = "visit"
gc.METHOD_WARNINGS.clear()
_c2 = gc.gen_robust_variance(_ar_plan, _step)
check("AR(1) with a time column passes it through",
      _c2["within_cluster_time_column"], "visit")
check("AR(1) with a time column raises no warning",
      any(w["kind"] == "gee_ar1_without_time_column" for w in gc.METHOD_WARNINGS), False)
gc.METHOD_WARNINGS.clear()


# ---- STROBE is built for the study's design ---------------------------------
sb = load("strobe")
for _raw, _want in [("cohort", "cohort"),
                    ("retrospective cohort study", "cohort"),
                    ("target trial emulation", "cohort"),
                    ("case-control", "case-control"),
                    ("nested case-control within a cohort", "case-control"),
                    ("cross-sectional", "cross-sectional"),
                    ("prevalence survey", "cross-sectional"),
                    ("rct", "unknown"),
                    ("randomized controlled trial", "unknown"),
                    ("我亂打的", "unknown"),
                    (None, "unknown"),
                    ("", "unknown")]:
    check(f"design {_raw!r} -> {_want}", sb.normalize_design(_raw)[0], _want)
check("case-cohort is not filed under cohort",
      sb.normalize_design("case-cohort study")[0], "unknown")
check("and it says why", "case-cohort" in (sb.normalize_design("case-cohort")[1] or ""), True)
check("an unmappable design says why", bool(sb.normalize_design("rct")[1]), True)
check("a mapped design needs no note", sb.normalize_design("cohort")[1], None)

for _d in ("cohort", "case-control", "cross-sectional", "unknown"):
    _items = sb.items_for_design(_d)
    check(f"{_d}: 34 checklist rows", len(_items), 34)
    check(f"{_d}: no unresolved DESIGN marker",
          any(r is sb.DESIGN for _, _, _, r, _k in _items), False)

_keys = {d: {f"{i}{s}": k for i, s, _sec, _r, k in sb.items_for_design(d)}
         for d in ("cohort", "case-control", "cross-sectional", "unknown")}
check("cohort 12d asks about loss to follow-up", _keys["cohort"]["12d"], "stats_ltfu")
check("case-control 12d asks about matching",
      _keys["case-control"]["12d"], "stats_matching_analysis")
check("cross-sectional 12d asks about sampling",
      _keys["cross-sectional"]["12d"], "stats_sampling")
check("cohort 14c asks for follow-up time", _keys["cohort"]["14c"], "followup")
check("case-control 14c is not applicable",
      _keys["case-control"]["14c"], "not_applicable")
check("cross-sectional 14c is not applicable",
      _keys["cross-sectional"]["14c"], "not_applicable")
check("cross-sectional 6b is not applicable",
      _keys["cross-sectional"]["6b"], "not_applicable")
check("cohort 6b still asks about matching", _keys["cohort"]["6b"], "matching")
check("case-control 15 reports exposure, not outcome events",
      _keys["case-control"]["15"], "exposure_data")
check("cohort 15 reports outcome events", _keys["cohort"]["15"], "outcome_data")
_psm_dir = tempfile.mkdtemp()
open(os.path.join(_psm_dir, "PSM.xlsx"), "w").write("x")
_plan_stub = {"output_dir": "/nonexistent", "study": {"design": "cohort", "title": "t"},
              "outcomes": {"primary": {"id": "p", "name": "P", "type": "survival"}},
              "covariates": [], "population": {}}
_plan_stub_nc = dict(_plan_stub)
_plan_stub_psm = dict(_plan_stub, output_dir=_psm_dir)

# An unrecognised design must never mark a row not applicable *because of the
# design*: guessing which rows do not apply is worse than listing one that does.
# Checking the resolver key is not enough — `matching` used to resolve to n/a
# whenever no PSM artifact existed, so the unknown variant reached n/a by a side
# door. Assert on the rendered status.
check("unknown design marks nothing n/a on resolver keys",
      "not_applicable" in _keys["unknown"].values(), False)
_unknown_rows = {f"{i}{sub}": sb.resolve(k, _plan_stub_nc, {"decision_log": None})[0]
                 for i, sub, _sec, _r, k in sb.items_for_design("unknown")}
check("unknown design 6b is not rendered n/a", _unknown_rows["6b"], "partial")
check("unknown design 14c is not rendered n/a", _unknown_rows["14c"] != "n/a", True)
check("unknown design 12d is not rendered n/a", _unknown_rows["12d"] != "n/a", True)
# Case-control 6b asks how cases and controls were sampled. A PSM artifact is
# analysis-stage matching and answers a different question, so it must not tick
# this box even when PSM.xlsx exists.
_cc_key = _keys["case-control"]["6b"]
check("case-control 6b does not reuse the cohort matching resolver",
      _cc_key, "matching_case_control")
check("case-control 6b stays missing even with a PSM artifact",
      sb.resolve(_cc_key, _plan_stub_psm, {"decision_log": None})[0], "missing")
check("cohort 6b is still covered by a PSM artifact",
      sb.resolve("matching", _plan_stub_psm, {"decision_log": None})[0], "covered")

# Every resolver key the table can produce has a branch in resolve().
for _d in ("cohort", "case-control", "cross-sectional", "unknown"):
    for _i, _s, _sec, _req, _k in sb.items_for_design(_d):
        _st, _ev = sb.resolve(_k, _plan_stub, {"decision_log": None})
        check(f"{_d} {_i}{_s} resolves to a known status", _st in sb.ICON, True)
        # The 13b ternary once bound to the wrong operand and returned a tuple
        # as the evidence string, which only surfaced as a crash inside openpyxl.
        check(f"{_d} {_i}{_s} evidence is a string", isinstance(_ev, str), True)

# ---- the psm config this skill emits must be readable by psm-matching --------
# gen_psm emits a mixed outcomes list whenever the primary is time-to-event and a
# secondary is binary — a common shape. psm_matching read a single analysis_type
# for every outcome and raised KeyError on the odd one out, so the two skills
# disagreed about the contract between them. Pin it here, on this side of the
# boundary, because this is where the config is built.
_psm_mod = load_from(SKILLS_DIR / "psm-matching/scripts/psm_matching.py",
                     "psm_matching")
_psm_plan = _synth_plan("time_to_event")          # primary survival, secondary_1 binary
_psm_cfg = gc.gen_psm(_psm_plan, {"id": "psm"})
_shapes = [_psm_mod.outcome_is_survival(o, _psm_cfg) for o in _psm_cfg["outcomes"]]
check("gen_psm really does emit a mixed outcomes list", sorted(set(_shapes)), [False, True])
check("psm-matching can resolve every outcome's columns",
      len(_psm_mod.get_outcome_columns(_psm_cfg)) > 0, True)

# The crash was only half of it: with mixed outcomes the workbook used to label
# every row "HR" and describe the whole table as Cox, including the row a
# conditional logistic produced. A mislabelled effect measure in a manuscript
# table is worse than a crash, because nothing announces it.
with tempfile.TemporaryDirectory() as _pd_dir:
    import numpy as _n6, pandas as _p6
    _r6 = _n6.random.default_rng(5)
    n = 400
    _age6 = _r6.normal(70, 8, n)
    _grp6 = (_r6.random(n) < 1 / (1 + _n6.exp(-(_age6 - 70) / 8))).astype(int)
    _csv6 = os.path.join(_pd_dir, "d.csv")
    _p6.DataFrame({"grp": _grp6, "age": _age6,
                   "t_mace": _r6.exponential(400, n),
                   "e_mace": _r6.integers(0, 2, n),
                   "e_death": _r6.integers(0, 2, n)}).to_csv(_csv6, index=False)
    # Regenerate from a plan rooted in the temp dir, so every derived path
    # (love plot, balance csv) lands somewhere that exists.
    _plan6 = _synth_plan("time_to_event")
    _plan6["output_dir"] = _pd_dir
    _plan6["study"]["data_file"] = _csv6
    _plan6["population"]["filters"] = []
    _plan6["derived_variables"] = []
    _plan6["grouping"] = {"column": "grp", "groups": [{"value": 1, "label": "A"},
                                                      {"value": 0, "label": "B"}]}
    _cfg6 = gc.gen_psm(_plan6, {"id": "psm"})
    _cfg6["matching_variables"] = [{"header": "C", "variables": [
        {"column": "age", "label": "Age", "type": "continuous"}]}]
    _cp6 = os.path.join(_pd_dir, "c.json"); json.dump(_cfg6, open(_cp6, "w"))
    _r6run = subprocess.run(
        [sys.executable, str(SKILLS_DIR / "psm-matching/scripts/psm_matching.py"), _cp6],
        capture_output=True, text=True)
    _psm_xlsx = _cfg6["output_file"]
    _psm_made = os.path.exists(_psm_xlsx)
    _labels6, _note6 = [], ""
    if _psm_made:
        from openpyxl import load_workbook as _lw6
        _ws6 = _lw6(_psm_xlsx, data_only=True).worksheets[0]
        for _row in _ws6.iter_rows(values_only=True):
            for _c in _row:
                if isinstance(_c, str):
                    if _c.endswith("(HR)") or _c.endswith("(OR)"):
                        _labels6.append(_c[-3:-1])
                    if "regression" in _c and "outcomes:" in _c:
                        _note6 = _c
check("a mixed-outcome PSM run produces a workbook", _psm_made, True)
check("each row carries its own effect measure", sorted(set(_labels6)), ["HR", "OR"])
check("and the footnote names both methods, not just one",
      ("Cox" in _note6 and "logistic" in _note6), True)

# ---- the weighting workbook must be extracted as a set, never one row --------
# Under marginal positivity the weighting step writes an ATE and an ATO side by
# side, and on real data they can differ more than twofold and point opposite
# ways. An extractor that took only the headline row would rebuild the selective
# reporting the companion exists to prevent, inside the drafting material.
_ed = load_from(SKILLS_DIR / "manuscript-results/scripts/extract_data.py",
                "extract_data")
check("manuscript-results knows how to read a weighting workbook",
      "weighting" in _ed.EXTRACTORS, True)
_wx = os.path.join(tempfile.mkdtemp(), "W.xlsx")
_wt_mod = load("weighting")


def _weighting_workbook(sep, subtitle=None):
    """Run the weighting step for real and return the extracted dict.

    `subtitle` matters more than it looks: gen_weighting always fills it from
    study.title, and it shifts every row on the sheet down by one. The first
    version of these tests never set it, so the extractor was validated against a
    layout no real run produces.
    """
    import numpy as _n, pandas as _p2
    rng = _n.random.default_rng(11)
    n = 300
    a = _n.r_[_n.zeros(n, dtype=int), _n.ones(n, dtype=int)]
    age = _n.r_[rng.normal(60, 5, n), rng.normal(60 + sep * 5, 5, n)]
    d = tempfile.mkdtemp()
    csv = os.path.join(d, "d.csv")
    _p2.DataFrame({"grp": a, "age": age, "time": rng.exponential(300, 2 * n),
                   "event": rng.integers(0, 2, 2 * n)}).to_csv(csv, index=False)
    out = os.path.join(d, "W.xlsx")
    cfg = {"mode": "weighting", "data_file": csv, "method": "iptw", "analysis_type": "cox",
           "cox_settings": {"time_column": "time", "event_column": "event"},
           "grouping": {"column": "grp", "groups": [{"value": 0, "label": "C"},
                                                    {"value": 1, "label": "T"}]},
           "sections": [{"header": "C", "variables": [
               {"column": "age", "label": "Age", "type": "continuous"}]}],
           "output_file": out, "plot_file": os.path.join(d, "W.png"),
           "love_plot_file": os.path.join(d, "L.png")}
    if subtitle:
        cfg["subtitle"] = subtitle
    cp = os.path.join(d, "c.json"); json.dump(cfg, open(cp, "w"))
    subprocess.run([sys.executable, str(SK / "weighting.py"), cp],
                   capture_output=True, text=True)
    extracted = _ed.extract_weighting(out)
    extracted["_path"] = out
    return extracted


_ex_marg = _weighting_workbook(4.0)
_estimands = [e["estimand"] for e in _ex_marg["estimates"]]
check("both weighted estimands are extracted",
      sorted(e for e in _estimands if e in ("ATE", "ATO")), ["ATE", "ATO"])
check("the companion is flagged as one", _ex_marg["has_ato_companion"], True)
check("and the pair is marked as inseparable", _ex_marg["must_report_together"], True)
check("the effect estimates come through, not just labels",
      all(e["effect"] for e in _ex_marg["estimates"]), True)
# A weighted effect cannot be read without the ESS and the overlap diagnostics:
# the interval reflects the effective sample size, not the row count.
check("the effective sample size is carried along", bool(_ex_marg["ess"]), True)
check("so are the positivity diagnostics",
      "Overlap coefficient" in _ex_marg["positivity"], True)
# Balance must be the post-weighting figure. "SMD after" is neither the last
# column nor column 4 (that is SMD before), so reading it by position reports
# balance the weighting never achieved.
# Pin the post-weighting figure against the pre-weighting one read from the same
# sheet. The previous version compared against a dict key that does not exist, so
# it degenerated into "is not empty" and passed on the before-value too.
def _balance_smds(path):
    """(SMD before, SMD after) from the Maximum SMD row, read independently."""
    from openpyxl import load_workbook as _lw
    wsb = _lw(path, data_only=True)["Balance"]
    hrow = next(r for r in range(1, 12)
                if str(wsb.cell(r, 1).value or "") == "Variable")
    cols = {str(wsb.cell(hrow, c).value or ""): c for c in range(1, wsb.max_column + 1)}
    mrow = next(r for r in range(1, wsb.max_row + 1)
                if str(wsb.cell(r, 1).value or "") == "Maximum SMD")
    return (str(wsb.cell(mrow, cols["SMD before"]).value or ""),
            str(wsb.cell(mrow, cols["SMD after"]).value or ""))

_smd_before, _smd_after = _balance_smds(_ex_marg["_path"])
check("the fixture actually distinguishes before from after",
      _smd_before != _smd_after, True)
check("balance is reported after weighting", _ex_marg["max_smd_after"], _smd_after)
check("and is not the pre-weighting figure",
      _ex_marg["max_smd_after"] != _smd_before, True)

_ex_ok = _weighting_workbook(0.2)
check("a run with good overlap has no companion", _ex_ok["has_ato_companion"], False)
check("and is not marked as a pair", _ex_ok["must_report_together"], False)

# Every real workbook carries a subtitle, which pushes the table down one row.
# Reading by fixed position swallowed the header as an estimate, made a single
# estimand look like a pair, and blanked the balance figure — on every production
# file, while these tests stayed green against a shape the pipeline never emits.
_ex_sub = _weighting_workbook(4.0, subtitle="A study title, as every real plan has")
check("a workbook with a subtitle extracts the same estimands",
      [e["estimand"] for e in _ex_sub["estimates"]],
      [e["estimand"] for e in _ex_marg["estimates"]])
check("the header row is not read as an estimate",
      any(e["estimand"] == "Estimand" for e in _ex_sub["estimates"]), False)
check("balance survives the shift", _ex_sub["max_smd_after"], _ex_marg["max_smd_after"])
_ex_sub_ok = _weighting_workbook(0.2, subtitle="A study title")
check("and a single-estimand run is still not called a pair",
      _ex_sub_ok["must_report_together"], False)

# ---- the reviewer-challenge gate is wired into the runner --------------------
check("GENERATORS has reviewer-challenge", "reviewer-challenge" in gc.GENERATORS, True)
check("SKILL_SCRIPTS has reviewer-challenge",
      "reviewer-challenge" in rp.SKILL_SCRIPTS, True)

_gate_pipe = [{"id": "missingness", "skill": "missing-data", "depends_on": []},
              {"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary",
               "depends_on": []},
              {"id": "reviewer_challenge", "skill": "reviewer-challenge",
               "depends_on": ["km_primary"]},
              {"id": "manuscript", "skill": "manuscript-results",
               "depends_on": ["reviewer_challenge"]}]
_gg = {g["group_id"]: [s["id"] for s in g["steps"]]
       for g in rp.group_steps_by_analysis_group(_gate_pipe, rp.topo_sort(_gate_pipe))}
check("reviewer-challenge -> manuscript group",
      _gg.get("manuscript"), ["reviewer_challenge", "manuscript"])
_buf2 = _io.StringIO()
with _cl.redirect_stdout(_buf2):
    ps.print_group_status({"pipeline": _gate_pipe})
check("the status view does not file the gate under 'Other'",
      "Other" in _buf2.getvalue(), False)

def _validates(pipeline):
    """True when validate_plan accepts the pipeline, False when it exits."""
    p = _synth_plan("time_to_event"); p["pipeline"] = pipeline
    try:
        gc.validate_plan(p, "plan.json")
        return True
    except SystemExit:
        return False

check("wired manuscript validates", _validates(_gate_pipe), True)
check("manuscript with no gate step is rejected",
      _validates([s for s in _gate_pipe if s["skill"] != "reviewer-challenge"]), False)
check("manuscript that skips the gate is rejected",
      _validates(_gate_pipe[:3] + [{"id": "manuscript", "skill": "manuscript-results",
                                    "depends_on": ["km_primary"]}]), False)
check("the gate reached transitively counts",
      _validates(_gate_pipe[:3]
                 + [{"id": "strobe", "skill": "strobe", "depends_on": ["reviewer_challenge"]},
                    {"id": "manuscript", "skill": "manuscript-results",
                     "depends_on": ["strobe"]}]), True)
# Duplicate ids collapse in every id-keyed dict downstream, so a gate can be
# shadowed by a later step of the same name: it validates, and never runs.
check("duplicate step ids are rejected",
      _validates(_gate_pipe + [{"id": "reviewer_challenge", "skill": "kaplan-meier",
                                "outcome": "primary", "depends_on": []}]), False)

# depends_on is not the only way into the manuscript step: `--only manuscript`
# skips dependency checking entirely. The manuscript path therefore re-runs the
# gate itself, and that check has to fail when there is no gate at all.
_ok_no_gate, _why_no_gate = rp.reviewer_gate_satisfied(
    {"pipeline": [{"id": "manuscript", "skill": "manuscript-results"}]}, "/nonexistent")
check("the manuscript path refuses to run with no gate step", _ok_no_gate, False)
check("and says why", "reviewer-challenge" in _why_no_gate, True)
_ok_no_cfg, _why_no_cfg = rp.reviewer_gate_satisfied(
    {"pipeline": [{"id": "rc", "skill": "reviewer-challenge"}]}, "/nonexistent")
check("the manuscript path refuses when the gate config is missing", _ok_no_cfg, False)

# Calling reviewer_gate_satisfied() directly proves the function works, not that
# run_step uses it — deleting the call site left every test green. Drive run_step
# itself, with a config on disk so it gets past the config-exists check and
# reaches the gate.
with tempfile.TemporaryDirectory() as _d:
    _out = os.path.join(_d, "out")
    os.makedirs(os.path.join(_out, "configs")); os.makedirs(os.path.join(_out, "logs"))
    json.dump({"mode": "manuscript"},
              open(os.path.join(_out, "configs", "manuscript_config.json"), "w"))
    _ms_plan = {"output_dir": _out,
                "pipeline": [{"id": "manuscript", "skill": "manuscript-results",
                              "depends_on": []}]}
    _pp = os.path.join(_d, "research_plan.json")
    json.dump(_ms_plan, open(_pp, "w"))
    _ok, _msg = rp.run_step(_ms_plan["pipeline"][0], _ms_plan, _pp, sys.executable)
    check("run_step refuses to write the manuscript with no gate", _ok, False)
    check("and names the reviewer challenge as the reason",
          "reviewer-challenge" in _msg, True)

check("a pipeline with no manuscript needs no gate",
      _validates([{"id": "missingness", "skill": "missing-data", "depends_on": []},
                  {"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary",
                   "depends_on": []}]), True)

# The missing-data assessment is not optional once anything is modelled: every
# model drops incomplete rows, and how many is not optional information.
check("a plan that models without assessing missingness is rejected",
      _validates([{"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary",
                   "depends_on": []}]), False)
check("the mi step does not count as the assessment",
      _validates([{"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary",
                   "depends_on": []},
                  {"id": "mi_primary", "skill": "missing-data", "outcome": "primary",
                   "depends_on": []}]), False)
check("a plan that models nothing needs no assessment",
      _validates([{"id": "flowchart", "skill": "flowchart", "depends_on": []},
                  {"id": "table1", "skill": "table1", "depends_on": []}]), True)

# The gate reads the reports, not the summary the reports were meant to produce.
_GATE_PLAN = {"version": 2, "output_dir": "/tmp/x", "pipeline": [],
              "study": {"title": "t", "design": "cohort"}, "covariates": []}

def _gate(agg, plan=None, decision_log=None):
    plan = plan or _GATE_PLAN
    with tempfile.TemporaryDirectory() as d:
        pp = os.path.join(d, "research_plan.json")
        json.dump(plan, open(pp, "w"))
        if decision_log is not None:
            open(os.path.join(d, "decision_log.md"), "w").write(decision_log)
        ap = os.path.join(d, "agg.json")
        json.dump(agg, open(ap, "w"))
        cp = os.path.join(d, "cfg.json")
        json.dump({"mode": "reviewer-challenge", "plan_path": pp, "aggregate_file": ap},
                  open(cp, "w"))
        buf = _io.StringIO()
        with _cl.redirect_stdout(buf):
            return rc.cmd_gate(cp)

def _agg(reports, version=2, fingerprint=None, **summary):
    fo = sum(1 for r in reports for c in r.get("comments", [])
             if c["verdict"] == "fatal" and c["status"] == "open")
    base = {"version": version,
            "plan_fingerprint": (fingerprint if fingerprint is not None
                                 else rc.plan_fingerprint(_GATE_PLAN)),
            "domains_reported": sorted({r["domain"] for r in reports}),
            "domains_missing": [k for k in rc.DOMAIN_KEYS
                                if k not in {r["domain"] for r in reports}],
            "domains_duplicated": [], "schema_problems": [],
            "counts": {"fatal_open": fo, "concern_open": 0, "comments_total": 0},
            "reports": reports}
    base.update(summary)
    return base

_full = [{"domain": k, "verdict": "pass", "checked": "x", "comments": []}
         for k, _ in rc.DOMAINS]
_fatal = {"id": "RC-01", "verdict": "fatal", "status": "open", "severity": "major",
          "comment": "unmeasured confounding", "evidence": "Multi.xlsx!B7", "remedy": "E-value"}

check("ten clean domains pass the gate", _gate(_agg(_full)), 0)
check("a missing domain blocks the gate", _gate(_agg(_full[:8])), 3)
check("an open fatal blocks the gate",
      _gate(_agg(_full[:9] + [dict(_full[9], comments=[_fatal])])), 3)
_wontfix_agg = _agg(_full[:9] + [dict(_full[9], comments=[
    dict(_fatal, status="wontfix", decision_log_id="D-014")])])
check("a fatal closed as wontfix passes when the decision-log row exists",
      _gate(_wontfix_agg,
            decision_log="# Decisions\n\n| D-014 | override RC-01 | ... |\n"), 0)
# "leaves a trace" is only true if something reads the log. Any non-empty string
# used to unlock the override, and the row it named did not have to exist.
check("a wontfix pointing at a decision-log row that does not exist blocks",
      _gate(_wontfix_agg, decision_log="# Decisions\n\n| D-001 | something else |\n"), 3)
check("a wontfix with no decision_log.md at all blocks",
      _gate(_wontfix_agg), 3)
# Omitting `evidence` used to be an undocumented third exit for a fatal: the
# comment was demoted to concern and the gate passed in silence.
_no_evidence = dict(_fatal, evidence="")
check("a fatal with no evidence pointer blocks instead of being demoted quietly",
      _gate(_agg(_full[:9] + [dict(_full[9], comments=[_no_evidence])])), 3)
# "wontfix" with nothing behind it is deleting the finding by another name.
check("a fatal closed as wontfix with no decision-log id blocks",
      _gate(_agg(_full[:9] + [dict(_full[9],
                                   comments=[dict(_fatal, status="wontfix")])])), 3)
# Ten reports that all say pass with an empty `checked` are a skipped challenge
# wearing a pass. The gate re-validates the reports instead of believing the
# summary's schema_problems list.
check("domains claiming pass without saying what they checked block",
      _gate(_agg([{"domain": k, "verdict": "pass", "checked": "", "comments": []}
                  for k, _ in rc.DOMAINS])), 3)
# The plan can change without anyone bumping `version`; the fingerprint is what
# notices.
_edited = dict(_GATE_PLAN, covariates=[{"name": "Age", "column": "age",
                                        "type": "continuous"}])
check("an edited plan at the same version still blocks",
      _gate(_agg(_full), plan=_edited), 3)
check("an aggregate with no fingerprint blocks",
      _gate(_agg(_full, fingerprint=""), plan=_GATE_PLAN), 3)
check("a challenge for an older plan version blocks",
      _gate(_agg(_full, version=1)), 3)
check("schema problems block the gate",
      _gate(_agg(_full, schema_problems=["r3: unreadable"])), 3)
# Editing the summary must not launder an open fatal into a pass.
check("a summary that disagrees with its reports blocks",
      _gate(_agg(_full[:9] + [dict(_full[9], comments=[_fatal])],
                 counts={"fatal_open": 0, "concern_open": 0, "comments_total": 1})), 3)
check("a summary hiding a missing domain blocks",
      _gate(_agg(_full[:8], domains_missing=[])), 3)

# ---- positivity: only structural failure blocks ------------------------------
# Marginal overlap is a limitation to report, not a reason to withhold the
# estimate; no-overlap-at-all is arithmetic and still stops the run.
_wt = load("weighting")
check("positivity thresholds are still surfaced in DEFAULTS",
      ("overlap_coefficient_threshold" in _wt.DEFAULTS
       and "nonoverlap_fraction_threshold" in _wt.DEFAULTS), True)

def _weighting_run(sep, method="iptw"):
    """Fit the weighting step on two groups separated by `sep` SDs of age."""
    import numpy as _np, pandas as _pd2
    rng = _np.random.default_rng(11)
    n = 300
    a = _np.r_[_np.zeros(n, dtype=int), _np.ones(n, dtype=int)]
    age = _np.r_[rng.normal(60, 5, n), rng.normal(60 + sep * 5, 5, n)]
    t = rng.exponential(300, 2 * n)
    with tempfile.TemporaryDirectory() as d:
        csv = os.path.join(d, "data.csv")
        _pd2.DataFrame({"grp": a, "age": age, "time": t,
                        "event": rng.integers(0, 2, 2 * n)}).to_csv(csv, index=False)
        cfg = {"mode": "weighting", "data_file": csv, "method": method,
               "analysis_type": "cox",
               "cox_settings": {"time_column": "time", "event_column": "event"},
               "grouping": {"column": "grp",
                            "groups": [{"value": 0, "label": "Control"},
                                       {"value": 1, "label": "Treated"}]},
               "sections": [{"header": "Covariates", "variables": [
                   {"column": "age", "label": "Age", "type": "continuous"}]}],
               "output_file": os.path.join(d, "W.xlsx"),
               "plot_file": os.path.join(d, "W.png"),
               "love_plot_file": os.path.join(d, "W_love.png")}
        cp = os.path.join(d, "cfg.json"); json.dump(cfg, open(cp, "w"))
        r = subprocess.run([sys.executable, str(SK / "weighting.py"), cp],
                           capture_output=True, text=True)
        return r.returncode, r.stderr

_rc_marginal, _err_marginal = _weighting_run(4.0)
check("severe-but-real overlap no longer blocks", _rc_marginal, 0)
check("severe overlap is reported as a warning",
      any(l.startswith("[QC_WARNING]") and "positivity" in l.lower()
          for l in _err_marginal.splitlines()), True)
check("severe overlap raises no blocking QC",
      any(l.startswith("[QC_BLOCKING]") for l in _err_marginal.splitlines()), False)

# Both graded checks have to be evaluated independently. While they were an
# elif chain, a run breaching both reported only the first, and a test that
# breaches one cannot tell the difference.
_pos_lines = [l for l in _err_marginal.splitlines()
              if l.startswith("[QC_WARNING]") and "Marginal positivity" in l]
check("a run breaching both positivity thresholds reports both",
      len(_pos_lines) >= 2, True)
check("one of them is the overlap coefficient",
      any("overlap coefficient" in l for l in _pos_lines), True)
check("the other is the fraction outside common support",
      any("outside the region of common support" in l for l in _pos_lines), True)

# Marginal overlap must not produce a lone IPTW estimate. The companion is
# computed in the same step from the same propensity model, so no plan wiring can
# route around it.
check("marginal positivity adds an ATO companion",
      any("ATO) companion analysis has been computed" in l
          for l in _err_marginal.splitlines()), True)

_rc_ok, _err_ok = _weighting_run(0.2)
check("good overlap runs clean", _rc_ok, 0)
check("good overlap raises no positivity warning",
      any(l.startswith("[QC_WARNING]") and "positivity" in l.lower()
          for l in _err_ok.splitlines()), False)
check("and adds no ATO companion when it is not needed",
      any("companion analysis has been computed" in l for l in _err_ok.splitlines()), False)

# method='overlap' already estimates the ATO, so a companion would duplicate it.
_rc_ato, _err_ato = _weighting_run(4.0, method="overlap")
check("an overlap-weighted run needs no companion of its own",
      any("companion analysis has been computed" in l for l in _err_ato.splitlines()), False)
check("but it still reports the marginal positivity",
      any("Marginal positivity" in l for l in _err_ato.splitlines()), True)


# ---- an explicit m above the protocol cap is honoured, but said out loud ------
# The cap of 100 bounds the automatic rule, not the analyst. Runtime scales with
# m and nothing else would mention that a hand-set value is above it. Analysis
# variable complete + auxiliary missing, so the run stops before MICE starts.
with tempfile.TemporaryDirectory() as d:
    import numpy as _n5, pandas as _p5
    _r5 = _n5.random.default_rng(3)
    n = 100
    _a5 = _n5.arange(n, dtype=float); _a5[:60] = _n5.nan
    csv = os.path.join(d, "d.csv")
    _p5.DataFrame({"grp": _r5.integers(0, 2, n), "age": _n5.arange(n, dtype=float),
                   "aux": _a5, "time": _r5.exponential(300, n),
                   "event": _r5.integers(0, 2, n)}).to_csv(csv, index=False)
    _base = {"mode": "imputation", "data_file": csv, "analysis_type": "cox",
             "cox_settings": {"time_column": "time", "event_column": "event"},
             "grouping": {"column": "grp", "groups": [{"value": 0, "label": "A"},
                                                      {"value": 1, "label": "B"}]},
             "sections": [{"header": "C", "variables": [
                 {"column": "age", "label": "Age", "type": "continuous"}]}],
             "auxiliary_variables": ["aux"],
             "output_file": os.path.join(d, "MI.xlsx")}

    def _mi_out(m):
        cfg = dict(_base, mi_settings={"m": m, "seed": 42, "maxit": 20})
        cp = os.path.join(d, f"c{m}.json"); json.dump(cfg, open(cp, "w"))
        r = subprocess.run([sys.executable, str(SK / "missing_data.py"), cp],
                           capture_output=True, text=True)
        return r.stdout + r.stderr

    # inside the with-block: the temp dir is gone by the time it exits
    _over = _mi_out(150)
    _under = _mi_out(40)


check("an explicit m above 100 is honoured", "m = 150" in _over, True)
check("and warned about", "above the protocol's cap of 100" in _over, True)
check("an explicit m within the cap raises no such warning",
      "above the protocol's cap" in _under, False)

# ---- guardian findings must survive a re-run ---------------------------------
# run_step clears the step's QC lines so a fixed problem stops counting. But
# biostats-guardian writes its verdict back into the same list after a review
# that re-running the analysis does not repeat, so clearing those too would let
# `--only <step>` wash a guardian block away with nobody looking at it again.
with tempfile.TemporaryDirectory() as _d:
    _out = os.path.join(_d, "out")
    os.makedirs(os.path.join(_out, "configs")); os.makedirs(os.path.join(_out, "logs"))
    json.dump({"mode": "x"}, open(os.path.join(_out, "configs", "s_config.json"), "w"))
    _gplan = {"output_dir": _out, "pipeline": [{"id": "s", "skill": "manuscript-results",
                                                "depends_on": []}]}
    _gstep = _gplan["pipeline"][0]
    _gstep["qc_warnings"] = ["[QC_BLOCKING] stale, from a run whose cause is fixed",
                             "[QC_WARNING] also stale",
                             "[GUARDIAN_BLOCKING] role 3 rejected the effect estimate"]
    _gpp = os.path.join(_d, "research_plan.json"); json.dump(_gplan, open(_gpp, "w"))
    rp.run_step(_gstep, _gplan, _gpp, sys.executable)
    _kept = _gstep["qc_warnings"]
check("this script's own stale QC lines are cleared on re-run",
      any(w.startswith("[QC_") for w in _kept), False)
check("but the guardian's verdict is not washed away",
      any(w.startswith("[GUARDIAN_BLOCKING]") for w in _kept), True)

# ---- MI as primary while PSM runs is not a silent pairing --------------------
_mi_pipe = [{"id": "mi_primary", "skill": "missing-data", "outcome": "primary",
             "depends_on": []},
            {"id": "uni_primary", "skill": "univariate-forest", "outcome": "primary",
             "depends_on": []},
            {"id": "psm", "skill": "psm-matching", "depends_on": []}]
def _checkpoint(plan):
    gs = rp.group_steps_by_analysis_group(_mi_pipe, rp.topo_sort(_mi_pipe), plan)
    return {g["group_id"]: g["checkpoint_type"] for g in gs}["outcome_primary"]

check("MI outcome group is FULL by default",
      _checkpoint({"missing_data": {"primary_analysis": "complete_case"},
                   "pipeline": _mi_pipe}), "FULL")
check("MI primary + PSM forces a MANDATORY checkpoint",
      _checkpoint({"missing_data": {"primary_analysis": "mi"}, "pipeline": _mi_pipe}),
      "MANDATORY")
def _group(plan):
    gs = rp.group_steps_by_analysis_group(_mi_pipe, rp.topo_sort(_mi_pipe), plan)
    return {g["group_id"]: g for g in gs}["outcome_primary"]

# MANDATORY on its own only pauses under --hitl, so marking the group and
# stopping there would have been a label with no behaviour behind it.
check("the MI/PSM group also halts in auto mode",
      _group({"missing_data": {"primary_analysis": "mi"},
              "pipeline": _mi_pipe}).get("halt_in_auto"), True)
check("and carries the reason for the halt",
      "primary" in _group({"missing_data": {"primary_analysis": "mi"},
                           "pipeline": _mi_pipe}).get("halt_reason", ""), True)
check("an ordinary outcome group does not halt in auto mode",
      _group({"missing_data": {"primary_analysis": "complete_case"},
              "pipeline": _mi_pipe}).get("halt_in_auto"), None)

# The flag on the group dict is inert unless run() reads it — ignoring it left
# every test green. Drive the runner itself and look for the halt banner.
def _halt_run(primary_analysis):
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "out")
        os.makedirs(os.path.join(out, "configs")); os.makedirs(os.path.join(out, "logs"))
        plan = {"output_dir": out,
                "missing_data": {"primary_analysis": primary_analysis},
                "analyses": {"psm_enabled": True},
                "pipeline": [
                    {"id": "mi_primary", "skill": "missing-data", "outcome": "primary",
                     "status": "pending", "depends_on": []},
                    {"id": "psm", "skill": "psm-matching", "status": "pending",
                     "depends_on": []}]}
        pp = os.path.join(d, "research_plan.json")
        json.dump(plan, open(pp, "w"))
        # --skip-precheck: this stub plan carries no dataset and no configs, so the
        # deterministic pre-flight cannot run and (correctly) stops the pipeline
        # before any group executes. What is under test here is the halt banner,
        # not the pre-flight; leaving the gate in the path would make this
        # assertion fail whenever precheck's entry conditions change.
        r = subprocess.run([sys.executable, str(SK / "run_pipeline.py"), pp,
                            "--skip-precheck"],
                           capture_output=True, text=True)
        return r.stdout, json.load(open(pp))

_halt_out, _halt_plan = _halt_run("mi")
check("the runner prints the halt banner in auto mode",
      "HALT [MANDATORY]" in _halt_out, True)
check("and stops before the supplementary group",
      _halt_plan["pipeline"][1]["status"], "pending")
_nohalt_out, _ = _halt_run("complete_case")
check("no halt banner when MI is not the primary analysis",
      "HALT [MANDATORY]" in _nohalt_out, False)
check("MI primary without PSM stays FULL",
      _checkpoint({"missing_data": {"primary_analysis": "mi"},
                   "pipeline": [s for s in _mi_pipe if s["skill"] != "psm-matching"]}),
      "FULL")
check("psm_enabled counts even without a PSM step",
      rp._mi_primary_with_psm({"missing_data": {"primary_analysis": "mi"},
                               "analyses": {"psm_enabled": True}, "pipeline": []}), True)
check("no plan means no upgrade", rp._mi_primary_with_psm(None), False)

gc.METHOD_WARNINGS.clear()
gc._warn_mi_primary_with_psm({"missing_data": {"primary_analysis": "mi"},
                              "pipeline": _mi_pipe})
check("MI primary + PSM is recorded as a methodological warning",
      any(w["kind"] == "psm_complete_case_under_mi_primary" for w in gc.METHOD_WARNINGS),
      True)
gc.METHOD_WARNINGS.clear()
gc._warn_mi_primary_with_psm({"missing_data": {"primary_analysis": "complete_case"},
                              "pipeline": _mi_pipe})
check("complete-case primary + PSM raises nothing", list(gc.METHOD_WARNINGS), [])
gc.METHOD_WARNINGS.clear()

# ---- bootstrap default is 500, and below it is labelled ----------------------
cal = load("calibration")
check("calibration bootstrap default is 500", cal.DEFAULTS["bootstrap_n"], 500)
_cal_plan = _synth_plan("time_to_event")
_cal_plan["advanced_analyses"] = {"calibration": {"study_purpose": "prediction"}}
check("generator default matches the script default",
      gc.gen_calibration(_cal_plan, _step)["bootstrap_n"], 500)
_cal_plan["advanced_analyses"]["calibration"]["bootstrap_n"] = 200
check("an explicit 200 is still honoured",
      gc.gen_calibration(_cal_plan, _step)["bootstrap_n"], 200)

# AUC's resample count was hard-coded and could time the step out with no lever.
_auc_plan = _synth_plan("time_to_event")
check("auc bootstrap defaults to 2000",
      gc.gen_auc(_auc_plan, {"id": "auc"})["bootstrap"]["n_iterations"], 2000)
_auc_plan["analyses"] = {"auc": {"bootstrap_n": 500, "seed": 7}}
_auc_cfg = gc.gen_auc(_auc_plan, {"id": "auc"})
check("and can be turned down from the plan", _auc_cfg["bootstrap"]["n_iterations"], 500)
check("along with its seed", _auc_cfg["bootstrap"]["random_seed"], 7)

# ---- incremental value: the base model is a declaration ----------------------
# gen_auc hard-coded clinical_model.enabled = False with no plan lever, so the
# pipeline could not ask "does this marker add to the standard model" at all —
# and precheck's signed_unit_interval rule, which exists to grade the negative
# Delta C-index that question produces, was unreachable from a generated config.
# The lever lives in `analyses`, which is already a FINGERPRINT_FIELD: WHICH
# variables count as already-known decides the delta, so choosing them after
# seeing the answer is the p-hacking surface, and the fingerprint is what makes
# that choice a pre-specification rather than a preference.
_inc = _synth_plan("time_to_event")
check("no declaration leaves incremental value off",
      gc.gen_auc(_inc, {"id": "auc"})["clinical_model"], {"enabled": False})
check("and evaluates every covariate",
      [v["column"] for v in gc.gen_auc(_inc, {"id": "auc"})["sections"][0]["variables"]],
      ["age", "sex"])
_inc["analyses"] = {"auc": {"clinical_model_columns": ["age"]}}
_inc_cfg = gc.gen_auc(_inc, {"id": "auc"})
check("a declared base model turns it on", _inc_cfg["clinical_model"]["enabled"], True)
check("with the declared columns as the base",
      [v["column"] for v in _inc_cfg["clinical_model"]["variables"]], ["age"])
check("and the OTHER covariates as the candidates — auc_cstat fits base + candidate, "
      "so a candidate inside its own base scores a delta of zero by construction",
      [v["column"] for v in _inc_cfg["sections"][0]["variables"]], ["sex"])
check("the base model's label comes from covariates, not a second copy in analyses",
      [v["label"] for v in _inc_cfg["clinical_model"]["variables"]], ["Age"])


def _auc_refuses(columns):
    gc.METHOD_WARNINGS.clear()
    plan = _synth_plan("time_to_event")
    plan["analyses"] = {"auc": {"clinical_model_columns": columns}}
    try:
        gc.gen_auc(plan, {"id": "auc"})
        return None
    except SystemExit as exc:
        return str(exc)


check("a base column the plan does not model is refused, not silently dropped",
      "plan.covariates does not model" in (_auc_refuses(["egfr"]) or ""), True)
check("naming every covariate is refused — nothing is left to measure",
      "no variable is left" in (_auc_refuses(["age", "sex"]) or ""), True)
check("a repeated column is refused rather than fitted twice",
      "repeats" in (_auc_refuses(["age", "age"]) or ""), True)
check("and a non-list is refused",
      "must be a list" in (_auc_refuses("age") or ""), True)
_rc_inc = load("reviewer_challenge")
check("declaring the base model moves the plan fingerprint, so a challenge that ran "
      "against the old declaration is stale rather than silently reused",
      _rc_inc.plan_fingerprint(_synth_plan("time_to_event"))
      != _rc_inc.plan_fingerprint(_inc), True)
gc.METHOD_WARNINGS.clear()

# 2026-08-11: extract_data.py grew extractors for all seven advanced-analysis
# skills plus the per-outcome MI step, so this block moved from "none of these
# reach the Results draft" to "all but the unconditional missingness assessment
# do" (see gen_manuscript's docstring for why that one stays out on purpose).
# Silence for whatever remains unextracted is still the "ran it, did not
# report it" failure mode, so the omission still has to be recorded — the
# still-unextracted `missingness` step is what exercises that path now.
#
# 2026-08-11 (flow-review gap #16): `flowchart` and the missingness assessment
# moved OUT of that warning into gen_manuscript._NOT_EXTRACTED. Neither has an
# extractor by design, so both fired on every run of every project, and a
# warning whose contents never change is a warning nobody reads — which is
# precisely how a genuinely unextracted analysis would have gone unnoticed. They
# are now printed as an informational line instead, and the warning is reserved
# for steps whose absence is a real oversight.
gc.METHOD_WARNINGS.clear()
_man_plan = _synth_plan("time_to_event")
_man_plan["pipeline"] = [
    {"id": "flowchart", "skill": "flowchart", "depends_on": []},
    {"id": "table1", "skill": "table1", "depends_on": []},
    {"id": "weighting_primary", "skill": "weighting", "outcome": "primary", "depends_on": []},
    {"id": "rcs_primary", "skill": "nonlinearity", "outcome": "primary", "depends_on": []},
    # missing-data is one skill with two step shapes: an "outcome" key means
    # the per-outcome MI + pooled-analysis run (extractable); no "outcome" key
    # means the unconditional missingness assessment (deliberately not).
    {"id": "mi_primary", "skill": "missing-data", "outcome": "primary", "depends_on": []},
    {"id": "missingness", "skill": "missing-data", "depends_on": []},
    {"id": "reviewer_challenge", "skill": "reviewer-challenge", "depends_on": []},
    {"id": "manuscript", "skill": "manuscript-results", "depends_on": ["reviewer_challenge"]},
]
_man_cfg = gc.gen_manuscript(_man_plan, {"id": "manuscript"})
check("weighting, RCS and MI outputs now reach the manuscript step",
      sorted(f["type"] for f in _man_cfg["analysis_files"]),
      ["missing-data", "nonlinearity", "table1", "weighting"])
_w = [w for w in gc.METHOD_WARNINGS
      if w["kind"] == "steps_missing_from_manuscript_extraction"]
check("the by-design exclusions no longer raise a warning every single run",
      len(_w), 0)
# …but a step that really has no extractor still must not pass in silence.
gc.METHOD_WARNINGS.clear()
_man_plan["pipeline"].insert(
    -1, {"id": "custom_probe", "skill": "some-external-analysis", "depends_on": []})
gc.gen_manuscript(_man_plan, {"id": "manuscript"})
_w = [w for w in gc.METHOD_WARNINGS
      if w["kind"] == "steps_missing_from_manuscript_extraction"]
check("a genuinely unextracted step is still recorded", len(_w), 1)
check("and the record names it", "custom_probe" in _w[0]["message"], True)
check("without the permanent noise entries",
      ("missingness" in _w[0]["message"] or "flowchart" in _w[0]["message"]), False)
check("the gate step is not counted as a missing extraction",
      "reviewer_challenge" in _w[0]["message"], False)
# Paths must be the workbook each generator actually writes. The old fallback
# was the step *id* (km_primary.xlsx), a file nothing ever creates.
_path_plan = _synth_plan("time_to_event")
_path_plan["pipeline"] = [
    {"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary", "depends_on": []},
    {"id": "multi_primary", "skill": "multivariate-interaction", "outcome": "primary",
     "depends_on": []},
    {"id": "manuscript", "skill": "manuscript-results", "depends_on": []},
]
_paths = {os.path.basename(f["path"])
          for f in gc.gen_manuscript(_path_plan, {"id": "manuscript"})["analysis_files"]}
check("a step with no output_file resolves to the real workbook name",
      sorted(_paths), ["KM_primary.xlsx", "Multi_primary.xlsx"])
gc.METHOD_WARNINGS.clear()

# ---- m follows the analysis variables, not the auxiliaries -------------------
md_mod = load("missing_data")
check("compute_m floors at 20", md_mod.compute_m(3.0), 20)
check("compute_m follows the missing %", md_mod.compute_m(41.2), 42)
check("compute_m caps at 100", md_mod.compute_m(140.0), 100)

with tempfile.TemporaryDirectory() as d:
    import numpy as _np3, pandas as _pd3
    _rng = _np3.random.default_rng(7)
    n = 200
    # 35% on the analysis variable, deliberately above the m floor of 20, so the
    # assertion below can tell "m came from the analysis variable" apart from
    # "m hit the floor". 70% on the auxiliary: if auxiliaries still fed into m,
    # this run would produce m = 70.
    _age = _np3.arange(n, dtype=float)
    _age[:70] = _np3.nan                      # 35% missing, analysis variable
    _aux = _np3.arange(n, dtype=float)
    _aux[:140] = _np3.nan                     # 70% missing, auxiliary only
    csv = os.path.join(d, "d.csv")
    _pd3.DataFrame({"grp": _rng.integers(0, 2, n), "age": _age, "aux": _aux,
                    "time": _rng.exponential(300, n),
                    "event": _rng.integers(0, 2, n)}).to_csv(csv, index=False)
    cfg = {"mode": "imputation", "data_file": csv,
           "analysis_type": "cox",
           "cox_settings": {"time_column": "time", "event_column": "event"},
           "grouping": {"column": "grp",
                        "groups": [{"value": 0, "label": "A"}, {"value": 1, "label": "B"}]},
           "sections": [{"header": "C", "variables": [
               {"column": "age", "label": "Age", "type": "continuous"}]}],
           "auxiliary_variables": ["aux"],
           "mi_settings": {"m": None, "seed": 42, "maxit": 20},
           "output_file": os.path.join(d, "MI.xlsx")}
    cp = os.path.join(d, "cfg.json"); json.dump(cfg, open(cp, "w"))
    r = subprocess.run([sys.executable, str(SK / "missing_data.py"), cp, "--check-missing"],
                       capture_output=True, text=True)
    _out = r.stdout
check("m is set from the analysis variable's 35%, not the auxiliary's 70%",
      "Planned m : 35" in _out, True)
check("the auxiliary's missingness is still reported",
      "auxiliary variables : 70.0%" in _out, True)
check("the analysis-variable maximum is labelled as such",
      "analysis variables : 35.0%" in _out, True)

# The check above runs the preview path only, so putting auxiliaries back into m
# inside run_imputation left it green. Drive the imputation path: an analysis
# variable with no missingness and an auxiliary with plenty. Correct behaviour is
# max_pct = 0 -> the run blocks before MICE starts ("multiple imputation is not
# applicable"). If auxiliaries fed into m, max_pct would be 60 and the run would
# go on to fit the imputations instead.
with tempfile.TemporaryDirectory() as d:
    import numpy as _np4, pandas as _pd4
    _r4 = _np4.random.default_rng(9)
    n = 120
    _aux2 = _np4.arange(n, dtype=float); _aux2[:72] = _np4.nan   # 60% missing
    csv = os.path.join(d, "d.csv")
    _pd4.DataFrame({"grp": _r4.integers(0, 2, n),
                    "age": _np4.arange(n, dtype=float),          # 0% missing
                    "aux": _aux2,
                    "time": _r4.exponential(300, n),
                    "event": _r4.integers(0, 2, n)}).to_csv(csv, index=False)
    cfg = {"mode": "imputation", "data_file": csv, "analysis_type": "cox",
           "cox_settings": {"time_column": "time", "event_column": "event"},
           "grouping": {"column": "grp",
                        "groups": [{"value": 0, "label": "A"}, {"value": 1, "label": "B"}]},
           "sections": [{"header": "C", "variables": [
               {"column": "age", "label": "Age", "type": "continuous"}]}],
           "auxiliary_variables": ["aux"],
           "mi_settings": {"m": None, "seed": 42, "maxit": 20},
           "output_file": os.path.join(d, "MI.xlsx")}
    cp = os.path.join(d, "cfg.json"); json.dump(cfg, open(cp, "w"))
    _r = subprocess.run([sys.executable, str(SK / "missing_data.py"), cp],
                        capture_output=True, text=True)
_imp = _r.stdout + _r.stderr
check("the imputation run also takes m from the analysis variables only",
      "multiple imputation is not applicable" in _imp, True)
check("and reports the analysis-variable maximum as 0",
      "Max missing %   : 0.0%" in _imp, True)
check("while still reporting the auxiliary's 60%",
      "auxiliary) : 60.0%" in _imp, True)


# ---- numeric contracts: Brier, IPCW Brier, counting-process reshape ----------
# Hand-checkable cases, not "it ran". Each of these kernels is a place where a
# sign or an index slip produces a plausible number rather than an error.
_y = [0, 1, 0, 1]
_p = [0.1, 0.9, 0.2, 0.8]
_bb = cal.brier_binary(_y, _p)
check("Brier = mean squared error", round(_bb["brier"], 6), 0.025)
check("null Brier uses the observed rate", round(_bb["brier_null"], 6), 0.25)
check("IPA = 1 - B/B_null", round(_bb["ipa"], 6), 0.9)

# With no censoring the IPCW weights are all 1, so the IPCW Brier must collapse
# onto the plain Brier of "had the event by t". If it does not, the weighting or
# the case split is wrong.
_t = [1.0, 2.0, 3.0, 4.0, 5.0, 6.0]
_e = [1, 1, 1, 1, 1, 1]
_risk = [0.9, 0.8, 0.6, 0.3, 0.2, 0.1]
_at = 3.5
_ipcw = cal.brier_ipcw(_t, _e, _risk, _at)
_plain = cal.brier_binary([1 if x <= _at else 0 for x in _t], _risk)
check("uncensored IPCW Brier equals the plain Brier",
      round(_ipcw["brier"], 10), round(_plain["brier"], 10))

# With censoring the weights stop being 1, which is the half of the estimator the
# uncensored case cannot exercise. Hand-computed: subjects at times 1(cens),
# 2(event), 3(event), 4(cens), predicted risks .5/.8/.4/.2, evaluated at t=2.5.
# The censoring KM drops to 0.75 at time 1, so G(2.5) = G(2-) = 0.75 and every
# contributing subject is weighted 4/3. The subject censored at 1 is before t and
# contributes nothing. B = (4/3)[(1-.8)^2 + .4^2 + .2^2] / 4 = 0.08.
_ct = [1.0, 2.0, 3.0, 4.0]
_ce = [0, 1, 1, 0]
_cprisk = [0.5, 0.8, 0.4, 0.2]
check("IPCW Brier with censoring matches the hand calculation",
      round(cal.brier_ipcw(_ct, _ce, _cprisk, 2.5)["brier"], 10), 0.08)
# The subject censored before t must drop out entirely, so its prediction cannot
# move the score. If it leaked into either branch, these two would differ.
# The case above still cannot see the difference between G(T-) and G(T): the
# censoring distribution is flat between the censoring time and the event time,
# so both readings give the same weight. A tie — one subject censored and another
# having the event at the same time — is what separates them. Times 2(event),
# 2(censored), 3(event), 4(censored); risks .6/.5/.4/.2; evaluated at t=2.5.
# Correct: the event at t=2 is weighted by G(2-) = 1, giving B = 0.10666...
# Reading G(2) = 0.75 instead would give 0.12.
check("IPCW Brier uses G(T-), not G(T), when a censoring time ties an event time",
      round(cal.brier_ipcw([2.0, 2.0, 3.0, 4.0], [1, 0, 1, 0],
                           [0.6, 0.5, 0.4, 0.2], 2.5)["brier"], 10),
      round((1 * (1 - 0.6) ** 2 + (4 / 3) * 0.4 ** 2 + (4 / 3) * 0.2 ** 2) / 4, 10))

check("a subject censored before t contributes nothing",
      round(cal.brier_ipcw(_ct, _ce, [0.99, 0.8, 0.4, 0.2], 2.5)["brier"], 10),
      round(cal.brier_ipcw(_ct, _ce, [0.01, 0.8, 0.4, 0.2], 2.5)["brier"], 10))

it_mod = load("immortal_time")
import numpy as _np
_cp_work = _pd.DataFrame({
    it_mod.PID:     [1, 2, 3, 4],
    it_mod.TIME:    [100.0, 100.0, 100.0, 100.0],
    it_mod.EVENT:   [1, 1, 0, 1],
    it_mod.EXPOSED: [0, 1, 1, 1],
    it_mod.EXPTIME: [_np.nan, 40.0, 150.0, -5.0],
    "x":            [0.0, 0.0, 0.0, 0.0],
})
it_mod._BLOCKING.clear()
_cp, _meta = it_mod.build_counting_process(_cp_work, [("x", "x")])
check("counting-process reshape raises no blocking issue", it_mod._BLOCKING, [])
check("one extra row only for exposure inside follow-up", _meta["n_rows"], 5)
check("exposure after follow-up contributes no exposed time",
      _meta["n_exposure_after_fu"], 1)
check("exposure at or before time zero is exposed throughout",
      _meta["n_exposed_at_zero"], 1)
check("total person-time is preserved",
      round(_meta["exposed_person_time"] + _meta["unexposed_person_time"], 6),
      float(_cp_work[it_mod.TIME].sum()))
# Subject 2 is exposed at 40 of 100, subject 4 from time zero: 60 + 100 = 160.
check("exposed person-time counts only time after exposure",
      round(_meta["exposed_person_time"], 6), 160.0)
# The 40 days subject 2 survived before exposure is exactly the immortal time a
# naive ever/never analysis credits to the exposed arm.
check("immortal time is the pre-exposure window", round(_meta["immortal_person_time"], 6), 40.0)
check("events are preserved by the reshape",
      int(_cp["event"].sum()), int(_cp_work[it_mod.EVENT].sum()))
# Preserving the event *count* is not enough — the count survives even when the
# event is attached to the wrong interval, which is exactly the reshape bug that
# would hand the exposed arm its immortal time. Subject 2 must carry its event on
# the second (exposed) row and nothing on the first.
_s2 = _cp[_cp[it_mod.PID] == 2].sort_values("tstart")
check("subject 2 is split into two intervals", len(_s2), 2)
check("the pre-exposure interval carries no event", int(_s2["event"].values[0]), 0)
check("the event lands on the exposed interval", int(_s2["event"].values[1]), 1)
check("the exposed interval starts at the exposure time",
      float(_s2["tstart"].values[1]), 40.0)
check("and that interval is flagged exposed", float(_s2["exposed"].values[1]), 1.0)
it_mod._BLOCKING.clear()


# ---- GEE AR(1) reads the time column, not the row order ----------------------
rv = load("robust_variance")
_rng2 = _np.random.default_rng(3)
_nc, _nv = 60, 3
_gee = _pd.DataFrame({
    "cid": _np.repeat(_np.arange(_nc), _nv),
    "visit": _np.tile(_np.arange(1, _nv + 1), _nc),
    "x": _rng2.normal(0, 1, _nc * _nv),
})
_lin = -0.3 + 0.8 * _gee["x"] + _np.repeat(_rng2.normal(0, 0.8, _nc), _nv)
_gee["y"] = (_rng2.random(_nc * _nv) < 1 / (1 + _np.exp(-_lin))).astype(float)
_gcfg = {"correlation_structure": "ar1", "within_cluster_time_column": "visit",
         "logistic_settings": {"outcome_column": "y"}}

_ordered = rv.fit_gee(_gee, ["x"], _gcfg, _gee["cid"].to_numpy())
# Shuffling rows within each cluster must not move the estimate: AR(1) is defined
# by the time gaps, and `visit` still carries them. If the fit reads row order
# instead, these two disagree.
_shuf = (_gee.groupby("cid", group_keys=False)
         .apply(lambda g: g.sample(frac=1, random_state=5)).reset_index(drop=True))
_shuffled = rv.fit_gee(_shuf, ["x"], _gcfg, _shuf["cid"].to_numpy())
check("AR(1) is invariant to row order within a cluster",
      round(_ordered["terms"]["x"]["coef"], 6), round(_shuffled["terms"]["x"]["coef"], 6))
check("AR(1) SE is invariant to row order too",
      round(_ordered["terms"]["x"]["se"], 6), round(_shuffled["terms"]["x"]["se"], 6))

# Singleton clusters have no within-cluster pairs, so there is no correlation to
# estimate and the fit must collapse onto the independence model.
_single = _gee.copy()
_single["cid"] = _np.arange(len(_single))
_indep = rv.fit_gee(_single, ["x"], dict(_gcfg, correlation_structure="independence"),
                    _single["cid"].to_numpy())
_ar_single = rv.fit_gee(_single, ["x"], _gcfg, _single["cid"].to_numpy())
check("with singleton clusters AR(1) equals independence",
      round(_ar_single["terms"]["x"]["coef"], 6), round(_indep["terms"]["x"]["coef"], 6))
check("and the estimated within-cluster correlation is zero",
      round(_ar_single["extra"]["estimated_within_cluster_correlation"] or 0.0, 6), 0.0)

# The AR(1)-without-a-time-column guard is a hard failure, not a silent fallback.
try:
    rv.fit_gee(_gee, ["x"], {"correlation_structure": "ar1",
                             "logistic_settings": {"outcome_column": "y"}},
               _gee["cid"].to_numpy())
    _guarded = False
except rv.FitFailure:
    _guarded = True
check("AR(1) without a time column refuses to fit", _guarded, True)

# The three checks above all still hold if AR(1) silently degenerated to
# independence, so one of them has to see a correlation actually being estimated.
# Strongly autocorrelated repeated measures: rho must come back clearly positive.
_ar = _pd.DataFrame({"cid": _np.repeat(_np.arange(80), 4),
                     "visit": _np.tile(_np.arange(1, 5), 80)})
_u = _np.repeat(_rng2.normal(0, 2.0, 80), 4)      # strong cluster-level shift
_ar["x"] = _rng2.normal(0, 1, len(_ar))
_arlin = -0.2 + 0.5 * _ar["x"] + _u
_ar["y"] = (_rng2.random(len(_ar)) < 1 / (1 + _np.exp(-_arlin))).astype(float)
_ar_fit = rv.fit_gee(_ar, ["x"], _gcfg, _ar["cid"].to_numpy())
_rho = _ar_fit["extra"]["estimated_within_cluster_correlation"]
check("AR(1) actually estimates a within-cluster correlation", _rho > 0.15, True)
_indep_fit = rv.fit_gee(_ar, ["x"], dict(_gcfg, correlation_structure="independence"),
                        _ar["cid"].to_numpy())
check("and the AR(1) fit differs from the independence fit",
      round(_ar_fit["terms"]["x"]["coef"], 6)
      != round(_indep_fit["terms"]["x"]["coef"], 6), True)


# ---- the seven advanced-analysis extractors, 2026-08-11 ----------------------
# Each one is validated against a workbook the real analysis script actually
# wrote, not a hand-typed dict — the weighting extractor's own history (see
# _weighting_workbook's docstring above) is why: a fixture missing the
# subtitle row every production plan carries reads the header as data. Every
# config below sets `subtitle`, the same lesson applied here.
import numpy as _np12
import pandas as _pd12

def _run_script(script, cfg, dirpath=None):
    d = dirpath or tempfile.mkdtemp()
    cp = os.path.join(d, "c.json")
    json.dump(cfg, open(cp, "w"))
    r = subprocess.run([sys.executable, str(SK / f"{script}.py"), cp],
                       capture_output=True, text=True)
    return d, r

_SUBTITLE = "A study title, as every real plan has"

# -- nonlinearity (RCS) ---------------------------------------------------------
def _rcs_workbook():
    rng = _np12.random.default_rng(41)
    n = 200
    df = _pd12.DataFrame({
        "grp": rng.integers(0, 2, n), "age": rng.uniform(30, 90, n),
        "time": rng.exponential(300, n) + 1, "event": rng.integers(0, 2, n)})
    d = tempfile.mkdtemp()
    csv = os.path.join(d, "d.csv"); df.to_csv(csv, index=False)
    out = os.path.join(d, "RCS.xlsx")
    cfg = {"mode": "rcs", "data_file": csv, "analysis_type": "cox",
           "cox_settings": {"time_column": "time", "event_column": "event"},
           "grouping": {"column": "grp", "groups": [{"value": 0, "label": "A"},
                                                    {"value": 1, "label": "B"}]},
           "sections": [{"header": "C", "variables": [
               {"column": "age", "label": "Age", "type": "continuous"}]}],
           "spline_variables": ["age"],
           "spline_settings": {"knots": 4, "knot_percentiles": None, "reference": None},
           "title": "RCS test", "subtitle": _SUBTITLE,
           "output_file": out, "plot_file": os.path.join(d, "RCS.png")}
    _run_script("nonlinearity", cfg, d)
    ex = _ed.extract_nonlinearity(out); ex["_path"] = out
    return ex

_rcs = _rcs_workbook()
check("nonlinearity extractor is registered", "nonlinearity" in _ed.EXTRACTORS, True)
check("title/subtitle come from the Estimand sheet, not Nonlinearity_tests",
      (bool(_rcs["title"]) and _rcs["subtitle"] == _SUBTITLE), True)
check("one spline variable extracted", len(_rcs["variables"]), 1)
check("reference value is non-empty (every reported HR is relative to it)",
      bool(_rcs["variables"][0]["reference_value"]), True)
check("the reading-rule footnote is captured verbatim",
      "does NOT mean linearity is correct" in _rcs["reading_rule"], True)
check("the not-established flag agrees with p-nonlinear itself",
      _rcs["variables"][0]["nonlinearity_not_established"],
      (_ed._to_float(_rcs["variables"][0]["p_nonlinear"]) >= 0.05))

# -- robust-variance --------------------------------------------------------------
def _robust_workbook():
    rng = _np12.random.default_rng(43)
    n = 200
    df = _pd12.DataFrame({
        "grp": rng.integers(0, 2, n), "age": rng.normal(60, 10, n),
        "site": rng.integers(0, 10, n),
        "time": rng.exponential(300, n) + 1, "event": rng.integers(0, 2, n)})
    d = tempfile.mkdtemp()
    csv = os.path.join(d, "d.csv"); df.to_csv(csv, index=False)
    out = os.path.join(d, "Robust.xlsx")
    cfg = {"mode": "robust", "data_file": csv, "analysis_type": "cox",
           "cox_settings": {"time_column": "time", "event_column": "event"},
           "grouping": {"column": "grp", "groups": [{"value": 0, "label": "A"},
                                                    {"value": 1, "label": "B"}]},
           "sections": [{"header": "C", "variables": [
               {"column": "age", "label": "Age", "type": "continuous"}]}],
           "cluster_column": "site", "method": "cluster_robust",
           "correlation_structure": "exchangeable", "compare_to_naive": True,
           "title": "Robust test", "subtitle": _SUBTITLE, "output_file": out}
    _run_script("robust_variance", cfg, d)
    ex = _ed.extract_robust_variance(out); ex["_path"] = out
    return ex

_rob = _robust_workbook()
check("robust-variance extractor is registered", "robust-variance" in _ed.EXTRACTORS, True)
check("naive and corrected are both present for every term (the comparison IS the file)",
      all(t["naive"] and t["corrected"] for t in _rob["terms"]), True)
check("must_report_together is always true for this file",
      _rob["must_report_together"], True)
check("cluster count is extracted from Cluster_summary", _rob["n_clusters"], "10")
check("10 clusters is flagged below the 30-cluster minimum",
      _rob["clusters_below_min"], True)

# -- calibration: prediction mode and etiologic mode are different workbooks -----
def _calibration_workbook(purpose):
    rng = _np12.random.default_rng(47)
    n = 200
    df = _pd12.DataFrame({
        "grp": rng.integers(0, 2, n), "age": rng.normal(60, 10, n),
        "time": rng.exponential(300, n) + 1, "event": rng.integers(0, 2, n)})
    d = tempfile.mkdtemp()
    csv = os.path.join(d, "d.csv"); df.to_csv(csv, index=False)
    out = os.path.join(d, "Cal.xlsx")
    cfg = {"mode": "calibration", "data_file": csv, "analysis_type": "cox",
           "cox_settings": {"time_column": "time", "event_column": "event"},
           "sections": [{"header": "C", "variables": [
               {"column": "age", "label": "Age", "type": "continuous"},
               {"column": "grp", "label": "Group", "type": "binary"}]}],
           "eval_times": [100], "bootstrap_n": 200, "seed": 1, "n_bins": 5,
           "min_bootstrap_success_frac": 0.5, "predicted_column": None,
           "study_purpose": purpose, "title": "Cal test", "subtitle": _SUBTITLE,
           "abbreviations": "", "output_file": out,
           "plot_file": os.path.join(d, "Cal.png")}
    _run_script("calibration", cfg, d)
    ex = _ed.extract_calibration(out); ex["_path"] = out
    return ex

_cal_pred = _calibration_workbook("prediction")
check("calibration extractor is registered", "calibration" in _ed.EXTRACTORS, True)
check("prediction-purpose run is flagged as a prediction model",
      _cal_pred["is_prediction_model"], True)
check("optimism-corrected C-index and calibration slope are reported together",
      _cal_pred["must_report_together"], True)
check("both optimism-corrected quantities are actually present, not just the flag",
      (bool(_cal_pred["internal_validation"].get("Harrell C-index", {})
            .get("Optimism-corrected"))
       and bool(_cal_pred["internal_validation"].get("Calibration slope", {})
                .get("Optimism-corrected"))),
      True)

_cal_etio = _calibration_workbook("etiologic")
check("etiologic-purpose run is NOT flagged as a prediction model",
      _cal_etio["is_prediction_model"], False)
check("etiologic mode has no calibration/Brier/internal-validation sheets to read",
      (_cal_etio["calibration_slope"] == {} and _cal_etio["brier"] == []
       and _cal_etio["internal_validation"] == {}),
      True)
check("discrimination is still reported for an etiologic run",
      bool(_cal_etio["discrimination"]), True)

# -- immortal-time: landmark and extended_cox are different workbooks ------------
def _immortal_time_workbook(mode):
    rng = _np12.random.default_rng(53)
    n = 200
    grp = rng.integers(0, 2, n)
    time = rng.exponential(300, n) + 1
    ext = _np12.full(n, _np12.nan)
    mask = grp == 1
    ext[mask] = rng.uniform(1, _np12.minimum(time[mask], 100))
    df = _pd12.DataFrame({"grp": grp, "age": rng.normal(60, 10, n), "extime": ext,
                          "time": time, "event": rng.integers(0, 2, n)})
    d = tempfile.mkdtemp()
    csv = os.path.join(d, "d.csv"); df.to_csv(csv, index=False)
    out = os.path.join(d, "IT.xlsx")
    cfg = {"mode": mode, "data_file": csv, "analysis_type": "cox",
           "cox_settings": {"time_column": "time", "event_column": "event"},
           "grouping": {"column": "grp",
                       "groups": [{"value": 1, "label": "Exposed"},
                                 {"value": 0, "label": "Control"}],
                       "exposed_value": 1},
           "sections": [{"header": "C", "variables": [
               {"column": "age", "label": "Age", "type": "continuous"}]}],
           "exposure_time_column": "extime",
           "landmark_times": [30] if mode == "landmark" else [],
           "landmark_rationale": "pre-specified for this test fixture",
           "time_unit": "days", "title": "IT test", "subtitle": _SUBTITLE,
           "output_file": out, "plot_file": os.path.join(d, "IT.png")}
    _run_script("immortal_time", cfg, d)
    ex = _ed.extract_immortal_time(out); ex["_path"] = out
    return ex

_land = _immortal_time_workbook("landmark")
check("immortal-time extractor is registered", "immortal-time" in _ed.EXTRACTORS, True)
check("mode is correctly detected as landmark", _land["mode"], "landmark")
check("naive and corrected rows are both present (the gap between them IS the finding)",
      bool(_land["naive"]) and bool(_land["corrected"]), True)
check("must_report_together is always true for this file",
      _land["must_report_together"], True)
check("landmark_results has one row per pre-specified landmark time",
      len(_land["landmark_results"]), 1)
check("excluded_at_landmark is populated (who got dropped is not optional)",
      bool(_land["excluded_at_landmark"]), True)

_ext = _immortal_time_workbook("extended_cox")
check("mode is correctly detected as extended_cox", _ext["mode"], "extended_cox")
check("extended_cox_results has the crude and adjusted rows",
      len(_ext["extended_cox_results"]), 2)
check("landmark-only fields stay empty in extended_cox mode",
      (_ext["landmark_results"] == [] and _ext["excluded_at_landmark"] == []), True)

# -- evalue -----------------------------------------------------------------------
def _evalue_workbook():
    rng = _np12.random.default_rng(59)
    n = 60
    df = _pd12.DataFrame({"time": rng.exponential(300, n) + 1,
                          "event": rng.integers(0, 2, n)})
    d = tempfile.mkdtemp()
    csv = os.path.join(d, "d.csv"); df.to_csv(csv, index=False)
    out = os.path.join(d, "EV.xlsx")
    cfg = {"data_file": csv, "analysis_type": "cox",
           "cox_settings": {"time_column": "time", "event_column": "event"},
           "estimates": [{"label": "Exposed vs Control", "effect": 1.5,
                         "ci_lower": 1.1, "ci_upper": 2.0}],
           "reference_confounders": [{"name": "Age", "rr": 1.3}],
           "title": "E-value test", "subtitle": _SUBTITLE, "output_file": out}
    _run_script("evalue", cfg, d)
    ex = _ed.extract_evalue(out); ex["_path"] = out
    return ex

_ev = _evalue_workbook()
check("evalue extractor is registered", "evalue" in _ed.EXTRACTORS, True)
check("one estimate extracted", len(_ev["estimates"]), 1)
check("point E-value and CI E-value are both present, never just one",
      (bool(_ev["estimates"][0]["E-value (point)"])
       and bool(_ev["estimates"][0]["E-value (CI)"])),
      True)
check("must_report_together is always true for this file",
      _ev["must_report_together"], True)
check("the interpretation names the reference confounder, giving the number a scale",
      "Age" in _ev["estimates"][0]["Interpretation"], True)
check("outcome prevalence is captured (decides which RR approximation applies)",
      bool(_ev["outcome_prevalence"]), True)

# -- strobe -------------------------------------------------------------------------
def _strobe_workbook(design):
    d = tempfile.mkdtemp()
    plan = {"study": {"title": "STROBE fixture", "design": design},
           "outcomes": {"primary": {"name": "Death"}},
           "covariates": [{"column": "age", "type": "continuous", "name": "Age"}],
           "population": {"filters": [], "exclusion_steps": []},
           "output_dir": d}
    plan_path = os.path.join(d, "research_plan.json")
    json.dump(plan, open(plan_path, "w"))
    out = os.path.join(d, "STROBE.xlsx")
    cfg = {"mode": "strobe", "plan_path": plan_path, "output_dir": d,
           "output_file": out, "md_file": os.path.join(d, "STROBE.md")}
    _run_script("strobe", cfg, d)
    ex = _ed.extract_strobe(out); ex["_path"] = out
    return ex

_st = _strobe_workbook("cohort")
check("strobe extractor is registered", "strobe" in _ed.EXTRACTORS, True)
check("design registered as cohort", _st["design_registered"], "cohort")
check("strobe version used matches the registered design", _st["strobe_version_used"], "cohort")
check("34 checklist rows extracted", len(_st["items"]), 34)
check("status counts sum to the item count",
      sum(_st["counts"].values()), len(_st["items"]))
check("missing_items only ever contains status == missing",
      all(i["status"] == "missing" for i in _st["missing_items"]), True)

_st_unk = _strobe_workbook("some_weird_design")
check("an unmapped design raises the warning", bool(_st_unk["design_unmapped_warning"]), True)
# The rule is that an unmapped design must not mark a row n/a *because of the
# design* — advanced_analyses.md §2 says so and then says plainly that a data
# fact can still produce one (16b is n/a whenever no continuous variable was
# categorised, whatever the design). Asserting zero n/a overall tests a rule
# nobody wrote, and the only way to make it pass would be to break the correct
# behaviour. Assert the rule that exists: none of the five design-specific rows.
_st_unk_designed = {i["item"]: i["status"] for i in _st_unk["items"]
                    if i["item"] in ("6a", "6b", "12d", "14c", "15")}
check("an unmapped design marks none of the design-specific rows n/a",
      [k for k, v in _st_unk_designed.items() if v == "n/a"], [])
check("and it really did evaluate all five of them", len(_st_unk_designed), 5)
# Any n/a that remains has to be a data fact, and 16b is the one the doc names.
check("a remaining n/a is data-driven, not a design guess",
      {i["item"] for i in _st_unk["items"] if i["status"] == "n/a"} <= {"6b", "16b"}, True)

# -- missing-data (per-outcome MI step; NOT the missingness assessment) ----------
def _mi_workbook():
    rng = _np12.random.default_rng(61)
    n = 200
    bmi = rng.normal(25, 4, n)
    bmi[rng.random(n) < 0.15] = _np12.nan
    df = _pd12.DataFrame({
        "grp": rng.integers(0, 2, n), "age": rng.normal(60, 10, n), "bmi": bmi,
        "time": rng.exponential(300, n) + 1, "event": rng.integers(0, 2, n)})
    d = tempfile.mkdtemp()
    csv = os.path.join(d, "d.csv"); df.to_csv(csv, index=False)
    out = os.path.join(d, "MI.xlsx")
    cfg = {"mode": "imputation", "data_file": csv, "analysis_type": "cox",
           "cox_settings": {"time_column": "time", "event_column": "event"},
           "grouping": {"column": "grp", "groups": [{"value": 0, "label": "A"},
                                                    {"value": 1, "label": "B"}]},
           "sections": [{"header": "C", "variables": [
               {"column": "age", "label": "Age", "type": "continuous"},
               {"column": "bmi", "label": "BMI", "type": "continuous"}]}],
           "auxiliary_variables": [], "mi_settings": {"m": None, "seed": 42, "maxit": 10},
           "outcome_id": "primary", "title": "MI test", "subtitle": _SUBTITLE,
           "abbreviations": "", "output_file": out,
           "diagnostics_plot": os.path.join(d, "MI.png")}
    _, r = _run_script("missing_data", cfg, d)
    ex = _ed.extract_missing_data(out); ex["_path"] = out; ex["_stderr"] = r.stderr
    return ex

_mi = _mi_workbook()
check("missing-data extractor is registered", "missing-data" in _ed.EXTRACTORS, True)
check("pooled has one row per analysis variable (grp, age, bmi)", len(_mi["pooled"]), 3)
check("CC_vs_MI has the same rows as Pooled — CC and MI are the same specification",
      len(_mi["cc_vs_mi"]), len(_mi["pooled"]))
check("must_report_together is always true for this file",
      _mi["must_report_together"], True)
check("m and seed are captured (a pooled estimate without them is not reproducible)",
      (bool(_mi["imputation_model"]["m"]) and _mi["imputation_model"]["seed"] == "42"),
      True)


# ---- 2026-08-11: Step 3.5's findings are plan-time gates --------------------
# Detected-and-unhandled used to produce a "⚠ suggest adding" line, with the
# corresponding fatal arriving 60 days later from a reviewer-challenge agent.
# The cost of that trade is one round of ten agents plus a full pipeline re-run,
# to learn something a set-membership test knew on day one.
_DL_DIR = tempfile.mkdtemp(prefix="rp-contract-log-")
with open(os.path.join(_DL_DIR, "decision_log.md"), "w", encoding="utf-8") as _f:
    _f.write("# Decision log\n\n## D-017 clustering left unmodelled\n"
             "Single centre, so the cluster is the study.\n\n"
             "## D-021 survival diagnostics not run\n"
             "The registry carries no dates, so the survival diagnostics have nothing\n"
             "to test; recorded as a limitation.\n")


def _diag_plan(flags, pipeline=None, outcomes=None, **diag):
    p = _synth_plan("time_to_event")
    p["output_dir"] = _DL_DIR                    # so decision_log.md is findable
    if outcomes:
        p["outcomes"] = outcomes
    p["diagnostics"] = dict(p["diagnostics"], flags_detected=flags, **diag)
    p["pipeline"] = pipeline if pipeline is not None else [
        {"id": "missingness", "skill": "missing-data", "depends_on": []},
        {"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary",
         "depends_on": []},
    ]
    return p


def _diag_validates(flags, pipeline=None, outcomes=None, **diag):
    gc.METHOD_WARNINGS.clear()
    try:
        gc.validate_plan(_diag_plan(flags, pipeline, outcomes, **diag), "plan.json")
        return True
    except SystemExit:
        return False


def _diag_message(flags, pipeline=None, outcomes=None, **diag):
    gc.METHOD_WARNINGS.clear()
    try:
        gc.validate_plan(_diag_plan(flags, pipeline, outcomes, **diag), "plan.json")
        return None
    except SystemExit as exc:
        return str(exc)


check("competing risks detected with no competing-risk step is refused",
      _diag_validates({"competing_risk": {"detected": True}}), False)
check("left truncation is refused outright — the estimates are biased, not incomplete",
      _diag_validates({"left_truncation": {"detected": True}}), False)
check("time-varying exposure at high risk with no landmark/extended-Cox step is refused",
      _diag_validates({"time_varying": {"detected": True, "risk": "high"}}), False)
check("the same flag graded low risk is not refused",
      _diag_validates({"time_varying": {"detected": True, "risk": "low"}}), True)
check("clustering with no robust-variance step is refused",
      _diag_validates(["clustering"]), False)
_REMEDIED = [
    {"id": "missingness", "skill": "missing-data", "depends_on": []},
    {"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary",
     "depends_on": []},
    {"id": "robust_primary", "skill": "robust-variance", "outcome": "primary",
     "depends_on": []},
]

# --- 2026-08-12, from an adversarial review: each of these WAS a bypass --
# 1. the value of `detected` was never read; only the key name was matched.
check("detected:false is not read as detected",
      _diag_validates({"competing_risk": {"detected": False, "risk": "none"}}), True)
check("a flag name that merely contains a flag word is refused, not matched",
      _diag_validates(["no_time_varying_detected"]), False)
check("and the refusal names the flag it could not read",
      "unrecognised flag" in (_diag_message(["no_time_varying_detected"]) or ""), True)
check("a non-boolean detected value is refused rather than treated as true",
      _diag_validates({"clustering": {"detected": "no"}}), False)

# 2. any diagnostics block at all satisfied "the diagnostics must have run":
#    {"mode": "not_run", "flags_detected": []} passed, which is a plan stating
#    in so many words that nothing was run.
check("mode='not_run' cannot stand in for a run",
      _diag_validates([], mode="not_run"), False)
check("mode='completed' with no report on disk is refused",
      _diag_validates([], report_path="/nonexistent/report.json"), False)
check("mode='completed' with no completed_at is refused",
      _diag_validates([], completed_at=None), False)

# 3. one non-empty string released all four gates, with nothing checked.
check("a blanket override_reason releases nothing",
      _diag_validates(["clustering", "left_truncation"],
                      override_reason="N/A"), False)
check("and says the free-text field has been withdrawn",
      "no longer releases any" in (_diag_message(["clustering"],
                                                 override_reason="N/A") or ""), True)
_OV = {"clustering": {"decision_id": "D-017", "reason": "single centre",
                      "outcomes": ["primary"]}}
check("a structured override naming a real decision releases its own gate",
      _diag_validates(["clustering"], overrides=_OV), True)
check("and it is recorded as a methodological decision, not waved through",
      any(w["kind"].startswith("diagnostic_gate_overridden")
          for w in gc.METHOD_WARNINGS), True)
check("an override whose decision_id is not in decision_log.md is rejected",
      _diag_validates(["clustering"], overrides={
          "clustering": dict(_OV["clustering"], decision_id="D-999")}), False)
check("an override pointing at a decision about something else is rejected",
      _diag_validates(["left_truncation"], overrides={
          "left_truncation": {"decision_id": "D-017", "reason": "see log",
                              "outcomes": ["primary"]}}), False)
check("an override for one flag does not release a different flag",
      _diag_validates(["clustering", "left_truncation"], overrides=_OV), False)

# 4. "some remedy step exists" counted as "the finding is handled", even when
#    the step pointed at a different outcome than the model being reported.
_TWO_OUTCOMES = {
    "primary": {"id": "primary", "name": "MACE", "type": "time_to_event",
                "time_column": "t_mace", "event_column": "e_mace"},
    "secondary": [{"id": "sec", "name": "Death", "type": "time_to_event",
                   "time_column": "t_d", "event_column": "e_d"}],
}
_TWO_MODELS = [
    {"id": "missingness", "skill": "missing-data", "depends_on": []},
    {"id": "multi_primary", "skill": "multivariate-interaction", "outcome": "primary",
     "depends_on": []},
    {"id": "multi_sec", "skill": "multivariate-interaction", "outcome": "sec",
     "depends_on": []},
]
_WITH_CLUSTER_COL = {"robust_variance": {"cluster_column": "hosp"}}
_sec_only = _diag_plan(["clustering"], pipeline=_TWO_MODELS + [
    {"id": "robust_sec", "skill": "robust-variance", "outcome": "sec",
     "depends_on": []}], outcomes=_TWO_OUTCOMES)
_sec_only["advanced_analyses"] = _WITH_CLUSTER_COL
try:
    gc.validate_plan(_sec_only, "plan.json")
    _sec_msg = None
except SystemExit as _e:
    _sec_msg = str(_e)
check("a remedy for the secondary outcome does not clear the primary model",
      _sec_msg is not None, True)
check("and the message names the outcome left uncovered",
      "['primary']" in (_sec_msg or ""), True)
_both = _diag_plan(["clustering"], pipeline=_TWO_MODELS + [
    {"id": "robust_primary", "skill": "robust-variance", "outcome": "primary",
     "depends_on": []},
    {"id": "robust_sec", "skill": "robust-variance", "outcome": "sec",
     "depends_on": []}], outcomes=_TWO_OUTCOMES)
_both["advanced_analyses"] = _WITH_CLUSTER_COL
gc.METHOD_WARNINGS.clear()
try:
    gc.validate_plan(_both, "plan.json")
    _both_ok = True
except SystemExit:
    _both_ok = False
check("covering every modelled outcome does clear it", _both_ok, True)
# A remedy step whose own config would make it refuse to run covers nothing.
# _REMEDIED names no cluster_column anywhere, so robust_variance.py would block.
check("a remedy step that would block at run time is not coverage",
      _diag_validates(["clustering"], pipeline=_REMEDIED), False)
check("and it is reported as present-but-inert",
      "Present but inert" in (_diag_message(["clustering"], pipeline=_REMEDIED) or ""),
      True)
check("an override covering only some of the affected outcomes is rejected",
      _diag_validates(["clustering"], pipeline=_TWO_MODELS, outcomes=_TWO_OUTCOMES,
                      overrides=_OV), False)

check("a clean diagnostics run validates", _diag_validates([]), True)

# "Nobody has looked" must not read as "they looked and found nothing".
_nodiag = _synth_plan("time_to_event")
del _nodiag["diagnostics"]
_nodiag["pipeline"] = [{"id": "missingness", "skill": "missing-data", "depends_on": []},
                       {"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary",
                        "depends_on": []}]
try:
    gc.validate_plan(_nodiag, "plan.json")
    _nodiag_msg = None
except SystemExit as _e:
    _nodiag_msg = str(_e)
check("a survival plan with no diagnostics block is blocked", _nodiag_msg is not None, True)
check("and the message says how to fix it, including --migrate",
      all(k in (_nodiag_msg or "") for k in ("Step 3.5", "--migrate", "override_reason")),
      True)
# mode='off' has to be a path a plan can actually take. It was briefly a
# setting no plan could satisfy — the override lookup did not recognise
# `diagnostics_off` as a key, so every mode='off' plan was blocked and the
# "problem" line came out empty, leaving nothing to act on.
check("mode='off' with a substantive reason is accepted",
      _diag_validates([], mode="off",
                      reason="registry has no dates; the survival diagnostics have "
                             "nothing to test"), True)
check("and it is recorded as a limitation rather than a clean bill of health",
      any(w["kind"] == "diagnostics_turned_off" for w in gc.METHOD_WARNINGS), True)
check("a token reason does not count",
      _diag_validates([], mode="off", override_reason="N/A"), False)
check("and the refusal says why that string is not a reason",
      "too short to be a reason" in (_diag_message([], mode="off",
                                                   override_reason="N/A") or ""), True)
check("no reason at all is refused, naming what is missing",
      "gives no reason" in (_diag_message([], mode="off") or ""), True)
check("the stronger decision-log form is accepted too",
      _diag_validates([], mode="off", overrides={
          "diagnostics_off": {"decision_id": "D-021", "reason": "no dates in registry",
                              "outcomes": ["primary"]}}), True)
check("but only when the decision it names is about the diagnostics",
      _diag_validates([], mode="off", overrides={
          "diagnostics_off": {"decision_id": "D-017", "reason": "single centre",
                              "outcomes": ["primary"]}}), False)
# Every refusal in this area has to leave the reader something to do. A FATAL
# whose explanation is blank is indistinguishable from a crash.
for _case in ({"mode": "off"}, {"mode": "off", "reason": "N/A"},
              {"mode": "later", "flags_detected": []},
              {"mode": "completed", "flags_detected": ["whatever"]}):
    _p = _synth_plan("time_to_event")
    _p["output_dir"] = _DL_DIR
    _p["diagnostics"] = dict(_p["diagnostics"], **_case)
    _p["pipeline"] = [{"id": "missingness", "skill": "missing-data", "depends_on": []},
                      {"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary",
                       "depends_on": []}]
    try:
        gc.validate_plan(_p, "plan.json")
        _msg = ""
    except SystemExit as _e:
        _msg = str(_e)
    check(f"diagnostics {_case.get('mode')!r}/{list(_case)[-1]} fails with an "
          f"explanation", len(_msg) > 200 and "Problem: \n" not in _msg, True)
_nodiag["pipeline"] = [{"id": "table1", "skill": "table1", "depends_on": []}]
check("a plan that fits no survival model is not asked for diagnostics",
      _diag_validates([], pipeline=[{"id": "table1", "skill": "table1",
                                     "depends_on": []}]), True)

# ---- effect-measure compatibility: E-value after cause-specific Cox ---------
_ev_plan = _synth_plan("competing_risk")
_ev_err = None
try:
    gc.gen_evalue(_ev_plan, {"id": "evalue_primary", "outcome": "primary"})
except ValueError as _e:
    _ev_err = str(_e)
check("an E-value on a competing-risk outcome is refused", _ev_err is not None, True)
check("and the refusal explains the scale mismatch, not just 'unsupported'",
      ("RISK" in (_ev_err or "") and "cause-specific" in (_ev_err or "")), True)
gc.METHOD_WARNINGS.clear()
gc.gen_evalue(_synth_plan("time_to_event"), {"id": "ev", "outcome": "primary"})
check("an E-value without a declared causal estimand is recorded, not blocked",
      any(w["kind"] == "evalue_without_declared_causal_estimand"
          for w in gc.METHOD_WARNINGS), True)
_causal = _synth_plan("time_to_event")
_causal["study"]["estimand_type"] = "causal"
gc.METHOD_WARNINGS.clear()
gc.gen_evalue(_causal, {"id": "ev", "outcome": "primary"})
check("declaring the estimand causal settles it",
      any(w["kind"] == "evalue_without_declared_causal_estimand"
          for w in gc.METHOD_WARNINGS), False)
# The source workbook is overridable: a project whose main model moved used to
# have to hand-patch the generated config, which the next run overwrote.
_causal["advanced_analyses"] = {"evalue": {"source_file": "Multi_primary_T5.xlsx"}}
check("evalue reads an overridden source workbook",
      os.path.basename(gc.gen_evalue(_causal, {"id": "ev", "outcome": "primary"})
                       ["source_file"]), "Multi_primary_T5.xlsx")

# ---- calibration is for prediction models only ------------------------------
_cal_err = None
try:
    gc.gen_calibration(_synth_plan("time_to_event"), {"id": "cal", "outcome": "primary"})
except ValueError as _e:
    _cal_err = str(_e)
check("calibration refuses when model_purpose is not declared", _cal_err is not None, True)
check("and says which field to set", "model_purpose" in (_cal_err or ""), True)
_pred = _synth_plan("time_to_event")
_pred["study"]["model_purpose"] = "prediction"
check("a declared prediction model gets its calibration step",
      gc.gen_calibration(_pred, {"id": "cal", "outcome": "primary"})["mode"], "calibration")
# 2026-08-12: purpose belongs to the MODEL, not to the study. A study-level
# field alone cannot describe an etiologic paper carrying one prognostic
# sub-model — it would either block the legitimate calibration or license
# calibration for every association model in the plan.
_mixed = _synth_plan("time_to_event")
_mixed["study"]["model_purpose"] = "etiologic"
check("a per-step model_purpose overrides the study default",
      gc.gen_calibration(_mixed, {"id": "cal", "outcome": "primary",
                                  "model_purpose": "prediction"})["mode"], "calibration")
check("and the study default still refuses the steps that do not claim it",
      gc._model_purpose(_mixed, {"id": "other"}), "etiologic")

# ---- schema v2 fields: absent is fine, unreadable is not --------------------
check("an absent pipeline_profile means 'submit'",
      gc._pipeline_profile(_synth_plan("time_to_event")), "submit")
_typo = _synth_plan("time_to_event")
_typo["pipeline_profile"] = "exploratory-ish"
_typo["pipeline"] = [{"id": "table1", "skill": "table1", "depends_on": []}]
try:
    gc.validate_plan(_typo, "plan.json")
    _typo_ok = True
except SystemExit:
    _typo_ok = False
check("an unreadable profile value is rejected rather than silently defaulted",
      _typo_ok, False)
check("a prose estimand block is not read as an estimand TYPE",
      gc._estimand_type({"estimand": {"population": "…", "exposure": "…"}}), None)
check("covariates_selection_rule counts as a rationale for covariates without one",
      gc._covariate_rationale({"covariates_selection_rule": "pre-specified DAG"},
                              {"column": "age"}) is not None, True)

# ---- explore vs submit -------------------------------------------------------
_prof = _synth_plan("time_to_event")
_prof["plan_schema_version"] = gc.PLAN_SCHEMA_VERSION
_prof["pipeline_profile"] = "explore"
_prof["pipeline"] = [
    {"id": "missingness", "skill": "missing-data", "depends_on": []},
    {"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary", "depends_on": []},
    {"id": "reviewer_challenge", "skill": "reviewer-challenge", "depends_on": []},
    {"id": "manuscript", "skill": "manuscript-results", "depends_on": ["reviewer_challenge"]},
]
gc.METHOD_WARNINGS.clear()
gc.validate_plan(_prof, "plan.json")
_st = {s["id"]: s.get("status") for s in _prof["pipeline"]}
check("explore parks the submission steps rather than deleting them",
      [_st["manuscript"], _st["reviewer_challenge"]], ["skipped", "skipped"])
check("but leaves the analyses alone", _st["km_primary"], None)
# 2026-08-12: the park has to be reversible, and reversible WITHOUT guessing.
# It was one-way — a plan that had once run as explore kept manuscript and
# reviewer_challenge marked skipped, so switching back to submit produced a
# paper with no reviewer challenge behind it and no error anywhere.
check("the park is marked so it can be told from a step skipped by hand",
      [s.get("skip_reason") for s in _prof["pipeline"] if s["id"] == "manuscript"],
      ["profile"])
_prof["pipeline_profile"] = "submit"
gc.METHOD_WARNINGS.clear()
gc.validate_plan(_prof, "plan.json")
_st = {s["id"]: (s.get("status"), s.get("skip_reason")) for s in _prof["pipeline"]}
check("switching back to submit restores them",
      [_st["manuscript"], _st["reviewer_challenge"]],
      [("pending", None), ("pending", None)])
check("and the restoration is reported like any other pipeline change",
      sorted(c[1] for c in _prof["_pipeline_changes"] if c[0] == "unpark"),
      ["manuscript", "reviewer_challenge"])
_hand = _synth_plan("time_to_event")
_hand["plan_schema_version"] = gc.PLAN_SCHEMA_VERSION
_hand["pipeline"] = [
    {"id": "missingness", "skill": "missing-data", "depends_on": []},
    {"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary", "depends_on": []},
    {"id": "reviewer_challenge", "skill": "reviewer-challenge", "depends_on": []},
    {"id": "manuscript", "skill": "manuscript-results", "status": "skipped",
     "depends_on": ["reviewer_challenge"]},
]
gc.validate_plan(_hand, "plan.json")
check("a step skipped by hand is not resurrected by the profile",
      [s["status"] for s in _hand["pipeline"] if s["id"] == "manuscript"], ["skipped"])
# Explore must still run the deterministic core — exploration is exactly when
# you want to find out the shape is not linear.
_expl = _synth_plan("time_to_event")
_expl["plan_schema_version"] = gc.PLAN_SCHEMA_VERSION
_expl["pipeline_profile"] = "explore"
_expl["pipeline"] = [
    {"id": "missingness", "skill": "missing-data", "depends_on": []},
    {"id": "multi_primary", "skill": "multivariate-interaction", "outcome": "primary",
     "depends_on": []},
    {"id": "manuscript", "skill": "manuscript-results", "depends_on": []},
]
gc.METHOD_WARNINGS.clear()
gc.validate_plan(_expl, "plan.json")
_expl_ids = [s["id"] for s in _expl["pipeline"]]
check("explore still plans the non-linearity diagnostic",
      "rcs_primary" in _expl_ids, True)
check("but does not add the submission-only STROBE step",
      "strobe" in _expl_ids, False)

# The functional-form diagnostic is added for a submit-profile plan that fits a
# multivariable model over a continuous covariate — but it never replaces the
# linear term in the main model, which stays a pre-specified decision.
_auto = _synth_plan("time_to_event")
_auto["plan_schema_version"] = gc.PLAN_SCHEMA_VERSION
_auto["pipeline"] = [
    {"id": "missingness", "skill": "missing-data", "depends_on": []},
    {"id": "multi_primary", "skill": "multivariate-interaction", "outcome": "primary",
     "depends_on": []},
]
gc.METHOD_WARNINGS.clear()
gc.validate_plan(_auto, "plan.json")
_auto_ids = [s["id"] for s in _auto["pipeline"]]
check("a non-linearity diagnostic is added by default", "rcs_primary" in _auto_ids, True)
check("the main model keeps its linear term",
      "spline_variables" in gc.gen_multivariate_interaction(
          _auto, {"id": "multi_primary", "outcome": "primary"}), False)
check("and the addition is reported as a pipeline change",
      [c[1] for c in _auto["_pipeline_changes"]], ["rcs_primary"])
# …and NOT when the declared event count cannot support the parameters.
_low = _synth_plan("time_to_event")
_low["plan_schema_version"] = gc.PLAN_SCHEMA_VERSION
_low["outcomes"]["primary"]["n_events"] = 12
_low["pipeline"] = [
    {"id": "missingness", "skill": "missing-data", "depends_on": []},
    {"id": "multi_primary", "skill": "multivariate-interaction", "outcome": "primary",
     "depends_on": []},
]
gc.METHOD_WARNINGS.clear()
gc.validate_plan(_low, "plan.json")
check("low events per parameter suppresses the spline diagnostic",
      "rcs_primary" in [s["id"] for s in _low["pipeline"]], False)
check("and says so instead of doing it quietly",
      any(w["kind"] == "rcs_diagnostic_skipped_low_epv" for w in gc.METHOD_WARNINGS), True)

# ---- the EPV denominator (adjudicated 2026-08-12) ---------------------------
# There must be ONE definition of the parameter count. There were two: this file
# charged 1 for a categorical with no label_map where precheck_configs charged
# (L-1), and the looser count was the one feeding the "should we add a spline"
# decision — an overfitting guard with a back door that issued splines.
# generate_configs now calls precheck_configs.parameter_budget() rather than
# keeping a second copy, for the same reason reviewer_challenge imports
# _normalize_model_purpose from here instead of mirroring the enum table.
_pc = load("precheck_configs")


def _budget_probe(covariates, spline_cols=(), knots=4):
    p = _synth_plan("time_to_event")
    p["covariates"] = covariates
    total, unpriced, prov = gc._prospective_param_budget(p, "primary", spline_cols, knots)
    probe = json.loads(json.dumps(p))
    probe.setdefault("advanced_analyses", {}).setdefault("nonlinearity", {}).update(
        {"enabled": True, "variables": list(spline_cols), "knots": knots})
    ref, _bd, _pv = _pc.parameter_budget(probe, df=None, configs=None,
                                         outcome_id="primary")
    return total, unpriced, ref


# 1 parameter for the two-group exposure term is included by both.
_b = _budget_probe([{"column": "nyha", "type": "ordinal",
                     "label_map": {"1": "I", "2": "II", "3": "III", "4": "IV"}}])
check("a 4-level ordinal costs 3 parameters, not 1", _b[0], 1 + 3)
check("and the two counters agree", _b[0], _b[2])
_b = _budget_probe([{"column": "site", "type": "categorical"}])
check("a nominal categorical with NO label_map: both counters give the same number",
      (_b[0], _b[2]), (_b[2], _b[2]))
check("and parameter_budget itself names it as unpriced",
      [u["column"] for u in _b[1]], ["site"])
check("with a floor, so `total` is explicitly a lower bound",
      _b[1][0].get("floor"), 1)
_b = _budget_probe([{"column": "age", "type": "continuous"}], spline_cols=["age"])
check("a splined continuous costs knots-1", _b[0], 1 + 3)
check("still agreeing with precheck", _b[0], _b[2])
check("nothing is unpriced when every term is declared", _b[1], [])
check("the budget is the plan-level fallback at plan time (there is no dataframe)",
      gc._prospective_param_budget(_synth_plan("time_to_event"), "primary", [], 4)[2]
      ["source"], "plan")

# 2026-08-12 adjudication: the surcharge for an unpriced term is NOT applied
# here. `parameter_budget` owns the pricing; a caller free to re-price its items
# turns "the parameters this model fits" into a suggestion, which is the
# two-parameter-counts failure moved up a level rather than fixed. This side
# reads `provenance["unpriced_terms"]` and treats an unbounded one as
# undemonstrable: `total` is a lower bound, so the events-per-parameter figure
# is an UPPER bound, and clearing a floor on an upper bound proves nothing.
def _auto_rcs(covariates, n_events):
    p = _synth_plan("time_to_event")
    p["plan_schema_version"] = gc.PLAN_SCHEMA_VERSION
    p["covariates"] = covariates
    p["outcomes"]["primary"]["n_events"] = n_events
    p["pipeline"] = [
        {"id": "missingness", "skill": "missing-data", "depends_on": []},
        {"id": "multi_primary", "skill": "multivariate-interaction",
         "outcome": "primary", "depends_on": []},
    ]
    gc.METHOD_WARNINGS.clear()
    gc.validate_plan(p, "plan.json")
    return "rcs_primary" in [s["id"] for s in p["pipeline"]]

_UNDECLARED = [{"name": "Age", "column": "age", "type": "continuous"},
               {"name": "Sex", "column": "sex", "type": "binary"},
               {"name": "Site", "column": "site", "type": "categorical"}]
_DECLARED = [dict(c) for c in _UNDECLARED]
_DECLARED[2]["label_map"] = {"0": "a", "1": "b", "2": "c", "3": "d"}
check("120 events, one unpriced categorical: no spline is granted",
      _auto_rcs(_UNDECLARED, 120), False)
_unpriced_msg = next((w["message"] for w in gc.METHOD_WARNINGS
                      if w["kind"] == "rcs_diagnostic_skipped_unpriced_terms"), "")
check("and the refusal points at unpriced_terms", bool(_unpriced_msg), True)
check("naming the term and why it could not be priced",
      ("site" in _unpriced_msg and "undeclared levels" in _unpriced_msg), True)
check("and explaining that the ratio is an upper bound",
      "UPPER bound" in _unpriced_msg, True)
check("no invented level count appears in the reason",
      ("ten-level" in _unpriced_msg or "ten levels" in _unpriced_msg), False)
check("declaring the 4 levels shows the budget was in fact sufficient",
      _auto_rcs(_DECLARED, 120), True)
# Events cannot buy their way past an unpriced term: with no ceiling there is
# no finite worst case to clear, however many events there are.
check("even 5000 events do not grant a spline while a term is unpriced",
      _auto_rcs(_UNDECLARED, 5000), False)
check("while a declared budget that genuinely fails is refused on the EPV itself",
      _auto_rcs(_DECLARED, 40), False)
check("with the ordinary low-EPV reason, not the unpriced one",
      any(w["kind"] == "rcs_diagnostic_skipped_low_epv" for w in gc.METHOD_WARNINGS), True)
# A term that IS bounded from above can still clear, via the worst case — the
# clause exists so a future `ceiling` in unpriced_terms is honoured rather than
# ignored, without this file inventing the bound itself.
check("an unbounded unpriced term has no computable worst case",
      gc._worst_case_params(6, [{"column": "site", "floor": 1}]), None)
check("a ceiling makes the worst case computable",
      gc._worst_case_params(6, [{"column": "site", "floor": 1, "ceiling": 4}]), 8)

# The compatibility path, for a precheck_configs older than `unpriced_terms`.
# Its first version matched the human-readable type label against "? levels"
# while the producer writes "(levels unknown)", so it never fired: the fallback
# reported nothing unpriced, and the gate then treated a lower bound as an exact
# parameter count and granted a spline. It now reads `is_floor` — the same
# structured flag the producer uses to decide what goes into unpriced_terms.
# Formatting is not an interface; a field is.
_real_precheck = gc._precheck_module()
check("the producer marks an unpriced term with a structured flag",
      [b.get("is_floor") for b in _real_precheck.parameter_budget(
          {"covariates": [{"column": "site", "type": "categorical"}],
           "grouping": {"column": "g", "groups": [{"value": 1}, {"value": 0}]}},
          df=None, configs=None, outcome_id="primary")[1] if b.get("is_floor")], [True])


class _OlderPrecheck:
    """parameter_budget as it was before `unpriced_terms` existed."""

    @staticmethod
    def parameter_budget(*a, **k):
        total, breakdown, prov = _real_precheck.parameter_budget(*a, **k)
        return total, breakdown, {kk: vv for kk, vv in prov.items()
                                  if kk != "unpriced_terms"}


_skew = _synth_plan("time_to_event")
_skew["plan_schema_version"] = gc.PLAN_SCHEMA_VERSION
_skew["covariates"] = [{"name": "Age", "column": "age", "type": "continuous"},
                       {"name": "Site", "column": "site", "type": "categorical"}]
_skew["outcomes"]["primary"]["n_events"] = 400
_skew["pipeline"] = [
    {"id": "missingness", "skill": "missing-data", "depends_on": []},
    {"id": "multi_primary", "skill": "multivariate-interaction", "outcome": "primary",
     "depends_on": []},
]
_saved_precheck = gc._precheck_module
gc._precheck_module = lambda: _OlderPrecheck
try:
    gc.METHOD_WARNINGS.clear()
    gc.validate_plan(_skew, "plan.json")
    _skew_kinds = [w["kind"] for w in gc.METHOD_WARNINGS]
    check("under version skew the fallback detects the unpriced term",
          "parameter_budget_missing_unpriced_terms" in _skew_kinds, True)
    check("and still withholds the spline — the fallback must not fail open",
          "rcs_primary" in [s["id"] for s in _skew["pipeline"]], False)
    check("for the unpriced reason specifically",
          "rcs_diagnostic_skipped_unpriced_terms" in _skew_kinds, True)
finally:
    gc._precheck_module = _saved_precheck

_auto_rcs(_UNDECLARED, 400)
check("the retired 'assumed' warning, which named an invented level count, is gone",
      any(w["kind"] == "rcs_diagnostic_parameter_budget_assumed"
          for w in gc.METHOD_WARNINGS), False)
# The numerator is the TOTAL events in the fitted analysis set, never the
# smaller arm's: the min-group version blocks imbalanced cohorts that have
# plenty of events, and a gate that false-positives teaches people to override
# it — after which the true positives get overridden too.
check("the event count is read as a total",
      gc._declared_events({"total_events": 308}), 308)

# ---- --only must not empty the manifest of the other steps' warnings --------
# The configs list was merged with the previous manifest while the warnings were
# overwritten with just this call's, so a selective re-run silently dropped
# every methodological warning belonging to a step it did not touch — and the
# reviewer-challenge agents read that file.
with tempfile.TemporaryDirectory() as _d:
    _out = os.path.join(_d, "out")
    _p = _synth_plan("time_to_event")
    _p["output_dir"] = _out
    _p["outcomes"]["secondary"] = [{"id": "sec", "name": "Death", "type": "time_to_event",
                                    "time_column": "t2", "event_column": "e2"}]
    _p["pipeline"] = [
        {"id": "missingness", "skill": "missing-data", "depends_on": []},
        # neither states time_scale, so each raises its own time_scale_assumed
        {"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary",
         "depends_on": []},
        {"id": "km_sec", "skill": "kaplan-meier", "outcome": "sec", "depends_on": []},
    ]
    _pp = os.path.join(_d, "research_plan.json")
    json.dump(_p, open(_pp, "w"))
    _gcp = str(SK / "generate_configs.py")
    subprocess.run([sys.executable, _gcp, _pp], capture_output=True, text=True)
    _mf = json.load(open(os.path.join(_out, "configs", "manifest.json")))
    check("both steps' warnings are in the manifest",
          sorted({w["step_id"] for w in _mf["methodological_warnings"] if w.get("step_id")}),
          ["km_primary", "km_sec"])
    subprocess.run([sys.executable, _gcp, _pp, "--only", "km_primary"],
                   capture_output=True, text=True)
    _mf2 = json.load(open(os.path.join(_out, "configs", "manifest.json")))
    check("a selective re-run keeps the untouched step's warning",
          sorted({w["step_id"] for w in _mf2["methodological_warnings"] if w.get("step_id")}),
          ["km_primary", "km_sec"])
    check("and still lists every config",
          sorted(e["step_id"] for e in _mf2["configs"]),
          ["km_primary", "km_sec", "missingness"])

# ---- --migrate: a live plan must not lose its reviewer challenge ------------
# reviewer_challenge.plan_fingerprint() gates the manuscript. Adding a
# documentary field to a hashed block would invalidate a completed ten-domain
# challenge over text that changes no number, so migration writes only fields
# the fingerprint cannot see, and writes nothing at all where absence already
# means the default it would have written.
_rc = load("reviewer_challenge")
with tempfile.TemporaryDirectory() as _d:
    _legacy = _synth_plan("time_to_event")
    _legacy["output_dir"] = os.path.join(_d, "out")
    _legacy["pipeline"] = [
        {"id": "missingness", "skill": "missing-data", "depends_on": []},
        {"id": "km_primary", "skill": "kaplan-meier", "outcome": "primary",
         "depends_on": []},
        # a multivariable model over a continuous covariate is what the
        # functional-form diagnostic attaches to, so this plan has a pipeline
        # change to preview
        {"id": "multi_primary", "skill": "multivariate-interaction",
         "outcome": "primary", "depends_on": []},
    ]
    _pp = os.path.join(_d, "research_plan.json")
    json.dump(_legacy, open(_pp, "w"))
    _fp_before = _rc.plan_fingerprint(json.load(open(_pp)))
    _r = subprocess.run([sys.executable, str(SK / "generate_configs.py"), _pp, "--migrate"],
                        capture_output=True, text=True)
    _migrated = json.load(open(_pp))
    check("--migrate exits clean", _r.returncode, 0)
    check("it raises the schema version",
          _migrated.get("plan_schema_version"), gc.PLAN_SCHEMA_VERSION)
    check("it backs the plan up first",
          any(f.startswith("research_plan.json.bak-") for f in os.listdir(_d)), True)
    check("the reviewer-challenge fingerprint is untouched",
          _rc.plan_fingerprint(_migrated), _fp_before)
    check("so nothing has to be re-run for the migration itself",
          "fingerprint unchanged" in _r.stdout, True)
    check("it does not write pipeline_profile, which absence already answers",
          "pipeline_profile" in _migrated, False)
    check("nor invent an estimand", "estimand_type" in _migrated, False)
    # 2026-08-12: nor a model purpose. Defaulting to 'etiologic' would have
    # broken the plans that need the field most — an existing prognostic study
    # whose calibration step predates the field would start failing on the very
    # next run, and a mixed study cannot be described by one study-level value.
    check("nor guess a model purpose", "model_purpose" in _migrated, False)
    check("it does state what it declined to write and why",
          all(k in _r.stdout for k in ("pipeline_profile", "estimand_type",
                                       "model_purpose", "multiplicity")), True)
    # …and it does not file the outcomes by position either. A pre-specified
    # co-primary endpoint often sits in outcomes.secondary, and calling it
    # exploratory would demote it silently. Writing the placeholder "unknown" is
    # not free either: reviewer_challenge hashes analysis_intent as a structured
    # decision, so the placeholder alone would cost a ten-agent challenge.
    check("no multiplicity labels are invented", "multiplicity" in _migrated, False)
    check("and it warns about the pipeline change the next run will make",
          ("NEXT" in _r.stdout and "rcs_primary" in _r.stdout), True)
    _r2 = subprocess.run([sys.executable, str(SK / "generate_configs.py"), _pp, "--migrate"],
                         capture_output=True, text=True)
    check("running it twice is a no-op", "Nothing to migrate" in _r2.stdout, True)
    # A plan from a NEWER generator must not be silently downgraded.
    _future = json.load(open(_pp))
    _future["plan_schema_version"] = gc.PLAN_SCHEMA_VERSION + 1
    json.dump(_future, open(_pp, "w"))
    _r3 = subprocess.run([sys.executable, str(SK / "generate_configs.py"), _pp, "--migrate"],
                         capture_output=True, text=True)
    check("a newer schema version fails closed instead of being downgraded",
          _r3.returncode, 1)
    check("and the plan is left exactly as it was",
          json.load(open(_pp))["plan_schema_version"], gc.PLAN_SCHEMA_VERSION + 1)
    _r4 = subprocess.run([sys.executable, str(SK / "generate_configs.py"), _pp],
                         capture_output=True, text=True)
    check("an ordinary run refuses it too", _r4.returncode, 1)
    check("rather than treating it as an old plan and disabling the auto steps",
          "do NOT lower the number" in (_r4.stdout + _r4.stderr), True)

# ---- an undeclared time unit must be reported on every survival route --------
# resolve_time_scale() was only called by the KM and competing-risk generators,
# so the same undeclared field was warned about or passed over in silence
# depending on which analysis happened to be in the pipeline: KM warned, every
# Cox route said nothing.
_ts_plan = _synth_plan("time_to_event")
_ts_plan["outcomes"]["primary"].pop("time_scale", None)
_ts_plan["outcomes"]["primary"].pop("time_unit", None)


def _ts_warn(fn, *a):
    gc.METHOD_WARNINGS.clear()
    cfg = fn(*a)
    msgs = [w["message"] for w in gc.METHOD_WARNINGS if w["kind"] == "time_scale_assumed"]
    return cfg, (msgs[0] if msgs else None)


_uni_cfg, _uni_msg = _ts_warn(gc.gen_univariate_forest, _ts_plan,
                              {"id": "uni_primary", "outcome": "primary"})
_psm_cfg, _psm_msg = _ts_warn(gc.gen_psm, _ts_plan, {"id": "psm"})
_auc_cfg, _auc_msg = _ts_warn(gc.gen_auc, _ts_plan, {"id": "auc"})
for _label, _cfg, _msg in (("univariate-forest", _uni_cfg, _uni_msg),
                           ("psm-matching", _psm_cfg, _psm_msg),
                           ("auc-cstat", _auc_cfg, _auc_msg)):
    check(f"{_label} carries the time scale", "time_scale" in _cfg, True)
    check(f"and warns when nothing declared it ({_label})", _msg is not None, True)
# The wording is graded, because the statistical consequence differs. A shared
# sentence would overstate the Cox case (a constant rescaling cannot move a
# hazard ratio — the partial likelihood uses only the ordering of event times)
# and understate the AUC one (a time-dependent AUC at t consumes t in the
# declared unit, so the number itself is computed at the wrong time point).
check("the three messages are not the same text",
      len({_uni_msg, _psm_msg, _auc_msg}), 3)
check("the Cox routes say the hazard ratios are unaffected",
      all("unaffected" in m for m in (_uni_msg, _psm_msg)), True)
check("and point at the time-valued outputs instead",
      all("TIME-VALUED" in m for m in (_uni_msg, _psm_msg)), True)
check("the AUC route says the result itself changes",
      "does change the result" in _auc_msg, True)
check("naming the mechanism — discrimination computed at a different t",
      ("DIFFERENT time point" in _auc_msg and "C-index" in _auc_msg), True)
check("and it does not claim the hazard ratios are wrong",
      "unaffected" in _auc_msg, False)
# A logistic outcome has no time axis: the fields would be inert and the
# warning would be about nothing.
_ts_bin = _synth_plan("binary")
_ts_bin["outcomes"]["primary"].pop("time_scale", None)
_bin_cfg, _bin_msg = _ts_warn(gc.gen_univariate_forest, _ts_bin,
                              {"id": "uni", "outcome": "primary"})
check("a logistic route gets no time_scale field", "time_scale" in _bin_cfg, False)
check("and raises no time-scale warning", _bin_msg, None)
# Declared: no warning, on any route.
_ts_ok = _synth_plan("time_to_event")
_ts_ok["outcomes"]["primary"]["time_scale"] = 1
_ok_cfg, _ok_msg = _ts_warn(gc.gen_auc, _ts_ok, {"id": "auc"})
check("a declared time_scale is used verbatim", _ok_cfg["time_scale"], 1)
check("and raises no warning", _ok_msg, None)

# ---- two steps of one skill must not write one file -------------------------
# gen_psm wrote "PSM.xlsx" unconditionally. With a primary and a sensitivity PSM
# step the overwrite was the visible half; the worse half was that both steps
# RESOLVED to the same artifact, so every post-run check (outputs exist, numeric
# plausibility, group direction, sidecar match) read the second step off the
# first one's file and the second was never independently examined.
def _psm_plan(step_ids, declare=False):
    p = _synth_plan("time_to_event")
    p["pipeline"] = [{"id": "missingness", "skill": "missing-data", "depends_on": []},
                     {"id": "table1", "skill": "table1", "depends_on": []}]
    for sid in step_ids:
        step = {"id": sid, "skill": "psm-matching", "depends_on": []}
        if declare:
            step["output_file"] = f"custom_{sid}"
        p["pipeline"].append(step)
    return p


_one = _psm_plan(["psm"])
check("a single PSM step keeps the historical filename",
      os.path.basename(gc.gen_psm(_one, _one["pipeline"][2])["output_file"]), "PSM.xlsx")
check("and so do the other fixed-name generators",
      [os.path.basename(gc.gen_table1(_one, _one["pipeline"][1])["output_file"]),
       os.path.basename(gc.gen_missingness(_one, _one["pipeline"][0])["output_file"])],
      ["Table1.xlsx", "Missingness.xlsx"])
_two = _psm_plan(["psm_primary", "psm_sensitivity"])
_paths = [os.path.basename(gc.gen_psm(_two, s)["output_file"])
          for s in _two["pipeline"][2:]]
check("two PSM steps get distinct workbooks", len(set(_paths)), 2)
check("named after the steps that wrote them", sorted(_paths),
      ["PSM_psm_primary.xlsx", "PSM_psm_sensitivity.xlsx"])
check("and the manuscript resolver agrees with what the generators wrote",
      sorted(os.path.basename(gc._manuscript_source_path(_two, s))
             for s in _two["pipeline"][2:]), sorted(_paths))
gc.METHOD_WARNINGS.clear()
gc.validate_plan(_two, "plan.json")
check("the rename is reported, never silent",
      any(w["kind"] == "output_file_disambiguated" for w in gc.METHOD_WARNINGS), True)
gc.METHOD_WARNINGS.clear()
gc.validate_plan(_one, "plan.json")
check("and is not reported when nothing collided",
      any(w["kind"] == "output_file_disambiguated" for w in gc.METHOD_WARNINGS), False)
_declared = _psm_plan(["psm_a", "psm_b"], declare=True)
check("an explicit output_file on the step still wins",
      sorted(os.path.basename(gc.gen_psm(_declared, s)["output_file"])
             for s in _declared["pipeline"][2:]),
      ["custom_psm_a.xlsx", "custom_psm_b.xlsx"])
# The same rule covers every other fixed-name generator; none of them is
# guaranteed by construction to appear once, so none of them hard-codes a path.
_dup_t1 = _synth_plan("time_to_event")
_dup_t1["pipeline"] = [{"id": "table1", "skill": "table1", "depends_on": []},
                       {"id": "table1_subgroup", "skill": "table1", "depends_on": []}]
check("duplicate table1 steps are disambiguated too",
      sorted(os.path.basename(gc.gen_table1(_dup_t1, s)["output_file"])
             for s in _dup_t1["pipeline"]),
      ["Table1_table1.xlsx", "Table1_table1_subgroup.xlsx"])

# ---- the PH remedy has to reach the model config ----------------------------
# multivariate_interaction.py blocks (exit 2) when Schoenfeld finds a violation
# and no remedy was pre-specified. Nothing carried that declaration out of
# research_plan.json, so the only way to set it inside a pipeline was to edit
# the generated config after the run had already blocked — the after-the-fact
# route the gate exists to prevent. multivariate-interaction/SKILL.md described
# this wiring as existing; grep found nothing.
_ph = _synth_plan("time_to_event")
_ph["outcomes"]["primary"]["ph_violation_remedy"] = "rmst"
_ph_cfg = gc.gen_multivariate_interaction(_ph, {"id": "multi_primary",
                                                "outcome": "primary"})
check("a declared PH remedy reaches the model config",
      _ph_cfg.get("ph_violation_remedy"), "rmst")
check("an undeclared one is absent, so a violation stays blocking",
      "ph_violation_remedy" in gc.gen_multivariate_interaction(
          _synth_plan("time_to_event"), {"id": "m", "outcome": "primary"}), False)
_ph["outcomes"]["primary"]["ph_violation_remedy"] = "disclose"
_ph["outcomes"]["primary"]["ph_disclose_decision_id"] = "D-042"
_ph_cfg = gc.gen_multivariate_interaction(_ph, {"id": "m", "outcome": "primary"})
check("the disclosure decision ID travels with it",
      (_ph_cfg.get("ph_violation_remedy"), _ph_cfg.get("ph_disclose_decision_id")),
      ("disclose", "D-042"))
_ph_typo = _synth_plan("time_to_event")
_ph_typo["outcomes"]["primary"]["ph_violation_remedy"] = "startify"
_ph_err = None
try:
    gc.gen_multivariate_interaction(_ph_typo, {"id": "m", "outcome": "primary"})
except ValueError as _e:
    _ph_err = str(_e)
check("a typo is refused rather than passed through as a pre-specification",
      _ph_err is not None, True)
gc.METHOD_WARNINGS.clear()
_ph_orphan = _synth_plan("time_to_event")
_ph_orphan["outcomes"]["primary"]["ph_disclose_decision_id"] = "D-042"
gc.gen_multivariate_interaction(_ph_orphan, {"id": "m", "outcome": "primary"})
check("a decision ID with no remedy is reported as the no-op it is",
      any(w["kind"] == "ph_disclose_decision_id_without_remedy"
          for w in gc.METHOD_WARNINGS), True)
gc.METHOD_WARNINGS.clear()
# The field is read from outcomes[] and NOT from the plan's top level, and that
# placement is the enforcement rather than a style choice: reviewer_challenge
# hashes `outcomes`, so adding a remedy after seeing the Schoenfeld p-value
# invalidates the challenge that had already passed. At the top level the same
# edit is free, and a post-hoc declaration would then be indistinguishable from
# a pre-specification — which is the only thing this gate checks.
_base_plan = _synth_plan("time_to_event")
_in_outcome = json.loads(json.dumps(_base_plan))
_in_outcome["outcomes"]["primary"]["ph_violation_remedy"] = "rmst"
_at_top = json.loads(json.dumps(_base_plan))
_at_top["ph_violation_remedy"] = "rmst"
check("declaring the remedy in outcomes[] moves the fingerprint",
      _rc.plan_fingerprint(_in_outcome) != _rc.plan_fingerprint(_base_plan), True)
check("declaring it at the top level would not — which is why it is not read there",
      _rc.plan_fingerprint(_at_top), _rc.plan_fingerprint(_base_plan))
check("and a top-level value is ignored",
      "ph_violation_remedy" in gc.gen_multivariate_interaction(
          _at_top, {"id": "m", "outcome": "primary"}), False)

# generate_configs._PH_REMEDIES duplicates multivariate_interaction.PH_REMEDIES
# ON PURPOSE: importing the analysis script pulls in lifelines and statsmodels,
# and the generator has to run on a machine that only has the plan. The
# duplication is justified; leaving it to a comment is not. Twice in this
# codebase a definition kept in two places drifted — the two parameter counts
# for EPV (with the looser one feeding the auto-spline decision), and the enum
# table reviewer_challenge chose to import rather than mirror. Without this
# check the next divergence surfaces as a remedy value one side accepts and the
# other rejects, at run time, in someone's pipeline.
if importlib.util.find_spec("lifelines") and importlib.util.find_spec("statsmodels"):
    _mi_mod = load_from(SKILLS_DIR / "multivariate-interaction/scripts"
                        / "multivariate_interaction.py", "multivariate_interaction")
    check("the two copies of the PH remedy vocabulary are identical",
          tuple(gc._PH_REMEDIES), tuple(_mi_mod.PH_REMEDIES))
else:
    # Loud, not silent: a check that quietly disappears is the same failure mode
    # as the drift it exists to catch.
    print("SKIP  the two copies of the PH remedy vocabulary are identical: "
          "multivariate_interaction.py needs lifelines + statsmodels and one of them "
          "is not importable here, so the contract could not be enforced. Run this "
          "suite in the analysis environment to check it.")

# generate_configs._REMEDY_STEPS[*]["hints"] mirrors
# ph_remedy_gate.REMEDY_STEP_HINTS, and the two do different jobs: the gate's copy
# RECOGNISES a hand-built step as evidence for a remedy, the generator's copy
# stops a generated step being added next to one the user already built. They are
# still the same fact. If the generator's copy is the narrower of the two, a
# hand-built step it fails to recognise gets a generated twin, and the gate is
# handed two artifacts answering one violation with no way to tell which is the
# analysis. Same argument as the _PH_REMEDIES mirror above; a comment saying "keep
# these in sync" is what the last two drifts also had.
_gate_mod = load_from(SKILLS_DIR / "biostats-guardian/scripts"
                      / "ph_remedy_gate.py", "ph_remedy_gate")
check("the two copies of the remedy step hints are identical",
      {k: tuple(v["hints"]) for k, v in gc._REMEDY_STEPS.items()},
      {k: tuple(v) for k, v in _gate_mod.REMEDY_STEP_HINTS.items()})

# Third mirror of the same fact: which column header marks an estimate.
# precheck_configs grades numbers and estimate_inventory counts what the
# manuscript must account for, so these two are NOT identical by design — the
# inventory skips log-scale columns, because a coefficient beside the ratio it
# exponentiates to is one estimate written twice, and counting it moves the C7
# denominator. What has to hold is that the difference is EXACTLY the log scale,
# plus the two formatted-header variants only the inventory needs. They had
# already drifted once: the inventory carried "coef" and "beta" but not
# "coefficient", so the same column counted once or twice depending on how the
# producing script spelled its header — and only when the value was positive,
# because _P_CELL_RE rejects a leading minus.
_inv_mod = load("estimate_inventory")
_FORMATTED_ONLY = {"or (95% ci)", "hr (95% ci)"}
check("the inventory counts exactly the non-log headers precheck knows",
      sorted(set(_inv_mod._EFFECT_HEADERS) - _FORMATTED_ONLY),
      sorted(h for h, scale in _pc._ESTIMATE_HEADER_SCALES.items() if scale != "log"))
check("and no log-scale header leaks into the inventory",
      sorted(set(_inv_mod._EFFECT_HEADERS)
             & {h for h, s in _pc._ESTIMATE_HEADER_SCALES.items() if s == "log"}),
      [])
# Every remedy that needs its own step must have one, and every generator named in
# the table must be registered. The table is the single place that knows, so the
# thing worth checking is that nothing reads it and finds a hole.
check("every remedy step in the table has a registered generator",
      sorted(v["skill"] for v in gc._REMEDY_STEPS.values()
             if v["skill"] in gc.GENERATORS),
      sorted(v["skill"] for v in gc._REMEDY_STEPS.values()))
check("every remedy step in the table has an output stem",
      sorted(v["skill"] for v in gc._REMEDY_STEPS.values()
             if v["skill"] in gc._OUTPUT_STEM),
      sorted(v["skill"] for v in gc._REMEDY_STEPS.values()))
check("the table covers exactly the remedies that need their own step",
      sorted(gc._REMEDY_STEPS), sorted(_gate_mod.REMEDIES_NEEDING_STEP))
# The gate must have a content check for each of them, or a remedy silently
# degrades to the file-existence path while looking like the others.
check("every remedy step has a gate content check",
      sorted(gc._REMEDY_STEPS), sorted(_gate_mod.CONTENT_CHECKS))

# ---- the plan must not be reverted by a write that raced another writer -----
# The write held a lock, but the READ happened long before it. Anything the
# runner recorded in between (a step status) or a researcher edited (a covariate)
# was overwritten by the older in-memory snapshot — and because the backup
# captured the NEWER file, the live plan went backwards while the only evidence
# sat in a .bak nobody opens.
with tempfile.TemporaryDirectory() as _d:
    _pp = os.path.join(_d, "research_plan.json")
    json.dump({"version": 1, "pipeline": []}, open(_pp, "w"))
    _stale = gc._disk_hash(_pp)
    json.dump({"version": 2, "pipeline": [{"id": "x", "status": "done"}]},
              open(_pp, "w"))                      # the other writer lands here
    _raised = False
    try:
        gc._backup_and_write_plan({"version": 1, "pipeline": []}, _pp,
                                  expected_hash=_stale)
    except gc.PlanChangedOnDisk:
        _raised = True
    check("a write whose baseline moved is refused", _raised, True)
    check("and the other writer's version is still on disk",
          json.load(open(_pp))["version"], 2)
    check("with its work intact",
          json.load(open(_pp))["pipeline"], [{"id": "x", "status": "done"}])
    _ok_hash = gc._disk_hash(_pp)
    gc._backup_and_write_plan({"version": 3, "pipeline": []}, _pp,
                              expected_hash=_ok_hash)
    check("an unraced write still goes through", json.load(open(_pp))["version"], 3)

# ---- one abort, one vocabulary ----------------------------------------------
# `[FATAL]` and `[QC_BLOCKING]` both meant "this run stops", so grepping either
# one missed half the terminating paths. That is not tidiness: callers split
# into those that read exit codes and those that scan stderr for `[QC_*]`
# findings, and a `[FATAL]`-only line is invisible to the second kind —
# precheck_configs hit exactly that and logged a failed run as
# `exit 1 — 0 blocking / 0 warning`. Every abort now carries BOTH markers. Exit
# codes are deliberately NOT unified: 1 here means "the plan is unusable", 2
# means "a step names a skill with no generator", and precheck_configs uses the
# opposite convention — three downstream files depend on the current behaviour.
_gc_src = (SK / "generate_configs.py").read_text().splitlines()
_orphans = [ln.strip() for ln in _gc_src
            if "[FATAL]" in ln and "[QC_BLOCKING]" not in ln
            and not ln.lstrip().startswith(("#", "-"))
            and "prefixed only" not in ln and "line was dropped" not in ln]
check("no source line emits [FATAL] without [QC_BLOCKING]", _orphans, [])
check("the abort helper carries both markers",
      ("[QC_BLOCKING]" in gc._abort.__doc__ or True) and
      all(m in "".join(_gc_src) for m in ("[QC_BLOCKING] [FATAL]",)), True)
# …and the same, observed at runtime rather than in the source.
with tempfile.TemporaryDirectory() as _d:
    _p = _synth_plan("time_to_event")
    _p["output_dir"] = _d
    _p.pop("grouping")
    _pp = os.path.join(_d, "research_plan.json")
    json.dump(_p, open(_pp, "w"))
    _r = subprocess.run([sys.executable, str(SK / "generate_configs.py"), _pp],
                        capture_output=True, text=True)
    _emitted = (_r.stdout + _r.stderr).splitlines()
    check("a real abort exits non-zero", _r.returncode != 0, True)
    check("and every emitted [FATAL] line also carries [QC_BLOCKING]",
          [ln for ln in _emitted if "[FATAL]" in ln and "[QC_BLOCKING]" not in ln], [])
    check("so a caller scanning for QC findings sees the reason",
          any("[QC_BLOCKING]" in ln for ln in _emitted), True)
    check("exit-code semantics are unchanged: an unusable plan is 1, not 2",
          _r.returncode, 1)

print("\n" + ("ALL PASS" if not fails else f"{len(fails)} FAILED: {fails}"))
sys.exit(1 if fails else 0)
