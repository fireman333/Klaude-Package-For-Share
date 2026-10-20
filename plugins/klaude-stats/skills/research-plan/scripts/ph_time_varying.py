#!/usr/bin/env python3
"""
Time-varying coefficient — the executable half of `ph_violation_remedy = "time_varying"`
(research-plan pipeline; 2026-08-12 adjudication dispute 5, remedy generators)

Why this exists
---------------
A proportional-hazards violation means the hazard ratio moved during follow-up.
The other two remedies respond by removing the offender from the model
(`stratify`) or by changing the estimand to one that never assumed a constant
ratio (`rmst`). This one responds by **reporting the movement**: the violating
variable is given its own coefficient in each of several pre-specified time
windows, so instead of one averaged HR the table says "3.1 in the first year,
0.8 thereafter" — which is usually the finding, not a nuisance.

Like the other remedies this is its own step and its own table; it does not touch
the multivariate workbook. See ph_stratified.py's header for why.

HOW IT IS FITTED, AND WHAT THAT COSTS
-------------------------------------
Piecewise-constant coefficients, fitted with `CoxTimeVaryingFitter` on a
start/stop frame. There is no way around the reshaping: lifelines' `CoxPHFitter`
has no equivalent of R's `tt()`, so a coefficient that varies with time has to be
expressed as rows over intervals. What IS avoided is the unbounded version — each
subject contributes at most (number of cut points + 1) rows, not one row per
event time, and the intervals are declared in advance rather than derived from
where the effect happens to change.

That bound is what keeps the model interpretable and the pre-specification
checkable. A piecewise HR over windows chosen after seeing the Schoenfeld plot is
not a remedy, it is a subgroup search along the time axis.

WHAT HAS TO BE DEFENDED HERE IS THE CUT POINTS
----------------------------------------------
Same argument as tau in ph_rmst.py, same enforcement. Cut points must come from
`outcomes[<id>].tv_cut_points`, which is inside
`reviewer_challenge.plan_fingerprint()`. Undeclared is not fatal here — the
analysis runs so the author can see the numbers and the exact fix — but the
sidecar records `cut_points_source: "default_rule"` and the manuscript gate
blocks on it. A cut chosen at the median event time is a reasonable default and
a bad pre-specification, and only the second thing matters to a reviewer.

Usage:
    python ph_time_varying.py config.json                  # full run
    python ph_time_varying.py config.json --check-missing  # preview only

Config (mode = "ph_time_varying"):
    {"mode":"ph_time_varying","remedy":"time_varying","outcome_id":"primary",
     "data_file":"...","filters":[],"derived_variables":[],
     "grouping":{"column":"AF","groups":[...]},
     "analysis_type":"cox",
     "cox_settings":{"time_column":"t","event_column":"e"},
     "sections":[{"header":"...","variables":[...]}],
     "time_varying_variables":"auto",     # or an explicit ["stage"]
     "cut_points":null,                    # or [1.0, 3.0]; null -> default rule
     "cut_points_source":"default_rule",
     "ph_test_file":".../Multi_primary_ph_test.json",
     "source_table":".../Multi_primary.xlsx",
     "output_file":".../TimeVarying_primary.xlsx",
     "evidence_file":".../TimeVarying_primary_remedy.json"}
"""

import hashlib
import json
import operator
import os
import sys
from datetime import datetime, timezone

import numpy as np
import pandas as pd

from lifelines import CoxTimeVaryingFitter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EVIDENCE_SCHEMA = "ph_remedy_evidence/1"
REMEDY = "time_varying"

# Recorded in the sidecar so a later slice can add the unbounded form without the
# gate having to guess which one it is looking at.
TV_METHOD = "piecewise_constant"

# More windows than this and each one is estimated from a slice of the events;
# a warning, not an error, because the right number is the author's call.
MAX_PERIODS = 5
MIN_EVENTS_PER_PERIOD = 10

_QC_MESSAGES = []


def qc(kind, message):
    line = f"[QC_{kind}] {message}"
    _QC_MESSAGES.append(line)
    print(line, file=sys.stderr)


def qc_warning(message):
    qc('WARNING', message)


def qc_blocking(message):
    qc('BLOCKING', message)


# ═══════════════════════════════════════════════════════════════════════════════
# Configuration
# ═══════════════════════════════════════════════════════════════════════════════

def load_config(path):
    with open(path, 'r', encoding='utf-8') as f:
        config = json.load(f)

    config.setdefault('mode', 'ph_time_varying')
    config.setdefault('remedy', REMEDY)
    config.setdefault('filters', [])
    config.setdefault('derived_variables', [])
    config.setdefault('sections', [])
    config.setdefault('title', '')
    config.setdefault('subtitle', '')
    config.setdefault('abbreviations', '')
    config.setdefault('analysis_type', 'cox')
    config.setdefault('cox_settings', None)
    config.setdefault('time_varying_variables', 'auto')
    config.setdefault('cut_points', None)
    config.setdefault('cut_points_source', 'default_rule')
    config.setdefault('ph_test_file', None)
    config.setdefault('source_table', None)
    config.setdefault('missing_data_handling', 'listwise')

    if config['mode'] != 'ph_time_varying':
        sys.exit(f"[FATAL] Unknown mode: {config['mode']!r} (expected 'ph_time_varying')")
    if config['remedy'] != REMEDY:
        sys.exit(f"[FATAL] This script implements ph_violation_remedy = {REMEDY!r}, "
                 f"but the config says {config['remedy']!r}.")
    if config['analysis_type'] != 'cox':
        sys.exit(f"[FATAL] analysis_type={config['analysis_type']!r}: a coefficient that "
                 f"varies with follow-up time only exists in a survival model. A "
                 f"logistic step has no proportional-hazards assumption to violate.")
    if not config.get('output_file'):
        sys.exit("[FATAL] config has no output_file; the gate matches the evidence "
                 "sidecar to the workbook by name, so the workbook must be named.")
    config.setdefault('evidence_file',
                      f"{os.path.splitext(config['output_file'])[0]}_remedy.json")
    return config


# ═══════════════════════════════════════════════════════════════════════════════
# Data loading (identical semantics to ph_stratified.py / robust_variance.py)
# ═══════════════════════════════════════════════════════════════════════════════

OPS = {
    '==': operator.eq, '!=': operator.ne,
    '>': operator.gt, '>=': operator.ge,
    '<': operator.lt, '<=': operator.le,
}


def load_data(config):
    path = config['data_file']
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xls'):
        df = pd.read_excel(path)
    elif ext == '.tsv':
        df = pd.read_csv(path, sep='\t')
    elif ext == '.sas7bdat':
        df = pd.read_sas(path)
    elif ext == '.dta':
        df = pd.read_stata(path)
    else:
        df = pd.read_csv(path)

    for f in config['filters']:
        col, op, val = f['column'], f['operator'], f['value']
        if op in OPS:
            df = df[OPS[op](pd.to_numeric(df[col], errors='coerce'), val)]
        elif op == 'in':
            df = df[df[col].isin(val)]
        elif op == 'not_in':
            df = df[~df[col].isin(val)]

    for dv in config['derived_variables']:
        if dv['type'] == 'coalesce':
            result = pd.to_numeric(df[dv['columns'][0]], errors='coerce')
            for c in dv['columns'][1:]:
                result = result.combine_first(pd.to_numeric(df[c], errors='coerce'))
            df[dv['name']] = result
        elif dv['type'] == 'expression':
            local_ns = {}
            for dep in dv.get('dependencies', []):
                local_ns[dep] = pd.to_numeric(df[dep], errors='coerce')
            local_ns['np'] = np
            df[dv['name']] = eval(dv['expression'], {"__builtins__": {}}, local_ns)

    return df.reset_index(drop=True)


def collect_variables(config):
    out = []
    for section in config['sections']:
        for v in section['variables']:
            out.append({
                'column': v['column'],
                'type': v.get('type', 'continuous'),
                'label': v.get('label', v['column']),
                'scale': v.get('scale'),
                'section': section.get('header', ''),
            })
    return out


def read_ph_test(config):
    path = config.get('ph_test_file')
    if not path:
        return None, None, "config has no ph_test_file"
    if not os.path.exists(path):
        return None, path, f"not found: {path}"
    try:
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f), path, None
    except Exception as exc:
        return None, path, f"{type(exc).__name__}: {exc}"


# ═══════════════════════════════════════════════════════════════════════════════
# Which variables vary with time, and where the windows are
# ═══════════════════════════════════════════════════════════════════════════════

def resolve_tv_variables(config, ph_art):
    """(variables, source, error). Same argument as ph_stratified.resolve_strata:
    the pre-specified decision is 'if PH fails, model the effect as time-varying';
    WHICH covariate failed is an observation about the data."""
    declared = config.get('time_varying_variables')
    if isinstance(declared, str) and declared.lower() == 'auto':
        if not ph_art:
            return [], 'auto_from_ph_test', (
                "time_varying_variables='auto' needs the PH test artifact to know "
                "which covariate violated, and it could not be read")
        violators = list(ph_art.get('violating_variables') or [])
        if not violators:
            return [], 'auto_from_ph_test', (
                "time_varying_variables='auto' but the PH test lists no violating "
                "covariate. A global violation with no single offender gives this "
                "remedy nothing to make time-varying — name the variables "
                "explicitly, or use rmst, which needs no offender")
        return violators, 'auto_from_ph_test', None
    if isinstance(declared, list) and declared:
        return list(declared), 'plan', None
    return [], 'plan', (f"time_varying_variables={declared!r} is neither 'auto' nor "
                        f"a non-empty list")


CUT_POINT_SOURCES = ('plan', 'default_rule')


def cut_points_source_of(config):
    """Where the windows came from — read, never re-derived.

    Same rule, and the same reason, as ph_rmst.tau_source_of: only the generator
    can know whether the plan declared them, so it says so in the config and this
    script reports what it was told. Inferring provenance from the presence of a
    field the generator also fills in for its own default is how an undeclared
    choice gets reported as a pre-specified one.
    """
    declared = config.get('cut_points_source')
    if declared in CUT_POINT_SOURCES:
        return declared
    if declared is not None:
        qc_warning(f"cut_points_source={declared!r} is not one of "
                   f"{CUT_POINT_SOURCES}; treating it as 'default_rule', because an "
                   f"unrecognised provenance claim is not evidence of "
                   f"pre-specification.")
        return 'default_rule'
    return 'plan' if config.get('cut_points') else 'default_rule'


def resolve_cut_points(config, d, time_col, event_col):
    """(cuts, source, error). Declared beats derived; the source reaches the sidecar."""
    source = cut_points_source_of(config)
    declared = config.get('cut_points')
    if declared:
        try:
            cuts = sorted(float(c) for c in declared)
        except (TypeError, ValueError):
            return None, source, f"cut_points={declared!r} are not all numbers"
        if any(c <= 0 for c in cuts):
            return None, source, f"cut_points must all be positive, got {cuts}"
        if len(set(cuts)) != len(cuts):
            return None, source, f"cut_points contain duplicates: {cuts}"
        return cuts, source, None
    ev = d.loc[d[event_col] == 1, time_col]
    if len(ev) < 10:
        return None, source, (f"only {len(ev)} events; a data-derived cut "
                              f"point would be noise")
    return [float(ev.median())], source, None


# ═══════════════════════════════════════════════════════════════════════════════
# Design + episode splitting
# ═══════════════════════════════════════════════════════════════════════════════

def build_wide(df, config, variables, tv_vars):
    """(wide_frame, model_cols, label_map, error) — complete case, numeric."""
    cs = config['cox_settings']
    time_col, event_col = cs['time_column'], cs['event_column']
    grp = (config.get('grouping') or {}).get('column')

    model_vars = list(variables)
    if grp and grp not in [v['column'] for v in model_vars]:
        model_vars = [{'column': grp, 'type': 'binary', 'label': grp,
                       'scale': None, 'section': 'Exposure'}] + model_vars

    needed = [v['column'] for v in model_vars] + [time_col, event_col]
    needed = list(dict.fromkeys(needed))
    missing = [c for c in needed if c not in df.columns]
    if missing:
        return None, None, None, f"columns not in the data: {missing}"
    unknown_tv = [c for c in tv_vars if c not in [v['column'] for v in model_vars]]
    if unknown_tv:
        return None, None, None, (
            f"the PH test names {unknown_tv} as violating, but they are not covariates "
            f"in this model. The violating names come from the fitted model's own "
            f"terms — if a categorical was dummy-coded there, name the raw column in "
            f"outcomes[<id>].tv_cut_points' companion field "
            f"outcomes[<id>].time_varying_variables instead of relying on 'auto'.")

    d = df[needed].copy()
    for c in needed:
        d[c] = pd.to_numeric(d[c], errors='coerce')
    # The cut points are stated in `time_unit`, so the time column has to be in
    # `time_unit` too. Same divide as kaplan_meier.py:213. Skipping it would put
    # every window boundary in the wrong place while the table still read as if
    # the units matched.
    if config.get('time_scale') is not None:
        d[time_col] = d[time_col] / config['time_scale']
    d = d.dropna(subset=needed).reset_index(drop=True)
    d = d[d[time_col] > 0].reset_index(drop=True)

    model_cols, label_map = [], {}
    for v in model_vars:
        col = v['column']
        if v.get('scale'):
            d[col] = d[col] / v['scale']
            label_map[col] = f"{v['label']} (per {v['scale']})"
        else:
            label_map[col] = v['label']
        if d[col].nunique(dropna=True) < 2:
            qc_warning(f"dropped covariate '{col}' (zero variance in the "
                       f"complete-case rows)")
            continue
        model_cols.append(col)
    return d, model_cols, label_map, None


def split_episodes(d, config, model_cols, tv_vars, cuts):
    """Wide -> start/stop long frame with period-specific copies of `tv_vars`.

    Each subject contributes one row per time window they are still at risk in,
    and at most `len(cuts) + 1` rows in total. The event indicator is 1 only on
    the row whose stop is the subject's own event time, which is what keeps the
    partial likelihood the same as the unsplit model's when no coefficient
    actually varies.
    """
    cs = config['cox_settings']
    time_col, event_col = cs['time_column'], cs['event_column']
    edges = [0.0] + list(cuts) + [float('inf')]
    rows = []
    for i, rec in enumerate(d.itertuples(index=False)):
        rec = rec._asdict()
        t, e = float(rec[time_col]), int(rec[event_col])
        for p in range(len(edges) - 1):
            lo, hi = edges[p], edges[p + 1]
            if t <= lo:
                break
            stop = min(t, hi)
            row = {'_id': i, 'start': lo, 'stop': stop,
                   '_event': int(e == 1 and stop == t), '_period': p}
            for c in model_cols:
                row[c] = rec[c]
            for v in tv_vars:
                for q in range(len(edges) - 1):
                    row[f"{v}__p{q}"] = rec[v] if q == p else 0.0
            rows.append(row)
    long = pd.DataFrame(rows)
    if long.empty:
        return None, None, "episode splitting produced no rows"
    long = long[long['stop'] > long['start']].reset_index(drop=True)
    fit_cols = [c for c in model_cols if c not in tv_vars]
    fit_cols += [f"{v}__p{p}" for v in tv_vars for p in range(len(edges) - 1)]
    return long, fit_cols, None


def period_labels(cuts, unit):
    edges = [0.0] + list(cuts)
    out = []
    for i, lo in enumerate(edges):
        hi = cuts[i] if i < len(cuts) else None
        out.append(f"{lo:g}–{hi:g} {unit}" if hi is not None else f">{lo:g} {unit}")
    return out


# ═══════════════════════════════════════════════════════════════════════════════
# Fitting
# ═══════════════════════════════════════════════════════════════════════════════

def fit_time_varying(long, fit_cols):
    """(results, ctv) or (None, reason)."""
    frame = long[['_id', 'start', 'stop', '_event'] + fit_cols].copy()
    try:
        ctv = CoxTimeVaryingFitter()
        ctv.fit(frame, id_col='_id', event_col='_event',
                start_col='start', stop_col='stop', show_progress=False)
    except Exception as exc:
        return None, f"{type(exc).__name__}: {exc}"
    s = ctv.summary
    results = {}
    for col in fit_cols:
        if col in s.index:
            results[col] = {
                'coef': float(s.loc[col, 'coef']),
                'effect': float(np.exp(s.loc[col, 'coef'])),
                'ci_lower': float(np.exp(s.loc[col, 'coef lower 95%'])),
                'ci_upper': float(np.exp(s.loc[col, 'coef upper 95%'])),
                'p': float(s.loc[col, 'p']),
            }
    return results, ctv


def heterogeneity(ctv, var, n_periods):
    """Do the window-specific coefficients actually differ?

    "We fitted a time-varying model" and "the effect demonstrably moved" are
    different claims, and only the second one is a finding. Wald test on the
    difference between the first and last window's coefficients, taken from the
    fitted covariance matrix so the correlation between them is not ignored.
    """
    first, last = f"{var}__p0", f"{var}__p{n_periods - 1}"
    try:
        params = ctv.params_
        cov = ctv.variance_matrix_
        if first not in params.index or last not in params.index:
            return {'status': 'not_computed',
                    'reason': f'{first} or {last} is not in the fitted model'}
        # Positional, not `.loc[a, b]`: CoxTimeVaryingFitter's variance_matrix_
        # has NAMED ROWS and a plain RangeIndex for its COLUMNS, so a label pair
        # raises KeyError on the column side even though both names are in the
        # index. Reading the row names and indexing the array survives that
        # asymmetry whichever way a future lifelines labels it.
        names = list(cov.index)
        V = np.asarray(cov.values, dtype=float)
        if V.shape[0] != V.shape[1] or V.shape[0] != len(names):
            return {'status': 'not_computed',
                    'reason': f'variance matrix is {V.shape} against {len(names)} '
                              f'named terms; cannot line the two up'}
        i, j = names.index(first), names.index(last)
        diff = float(params[last] - params[first])
        var_d = float(V[i, i] + V[j, j] - 2 * V[i, j])
        if var_d <= 0:
            return {'status': 'not_computed',
                    'reason': f'non-positive variance for the contrast ({var_d})'}
        se = float(np.sqrt(var_d))
        z = diff / se
        from scipy import stats as sps
        p = float(2 * sps.norm.sf(abs(z)))
        return {'status': 'computed', 'variable': var,
                'log_hr_difference': diff, 'se': se, 'z': float(z), 'p': p,
                'differs': bool(p < 0.05),
                'hr_ratio_last_over_first': float(np.exp(diff))}
    except Exception as exc:
        return {'status': 'not_computed', 'reason': f"{type(exc).__name__}: {exc}"}


# ═══════════════════════════════════════════════════════════════════════════════
# Excel
# ═══════════════════════════════════════════════════════════════════════════════

HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
TITLE_FONT = Font(bold=True, size=14)
SUB_FONT = Font(italic=True, size=10, color="555555")
DATA_FONT = Font(size=10)
SIG_FONT = Font(size=10, bold=True)
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
SECTION_FONT = Font(bold=True, size=10)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def fmt_p(p):
    if p is None:
        return "—"
    return "<0.001" if p < 0.001 else f"{p:.3f}"


def fmt_effect(r):
    if not r:
        return "N/A"
    return f"{r['effect']:.2f} ({r['ci_lower']:.2f}–{r['ci_upper']:.2f})"


def write_excel(config, results, label_map, model_cols, tv_vars, tv_source,
                cuts, cuts_source, labels, het, n, n_events, events_by_period,
                n_long_rows, unit, ph_art):
    wb = Workbook()
    ws = wb.active
    ws.title = "Time-varying Cox"
    for c, w in zip(range(1, 5), [36, 14, 26, 12]):
        ws.column_dimensions[get_column_letter(c)].width = w

    r = 1
    ws.cell(row=r, column=1,
            value=config.get('title') or "Cox model with time-varying coefficients"
            ).font = TITLE_FONT
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1
    if config.get('subtitle'):
        ws.cell(row=r, column=1, value=config['subtitle']).font = SUB_FONT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1
    ws.cell(row=r, column=1,
            value=f"n = {n} subjects, {n_events} events, {n_long_rows} risk intervals; "
                  f"windows: {', '.join(labels)}").font = SUB_FONT
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 2

    for c, h in enumerate(["Variable", "Coefficient", "HR (95% CI)", "p-value"], 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, CENTER, THIN
    r += 1

    for v in tv_vars:
        cell = ws.cell(row=r, column=1, value=f"{label_map.get(v, v)} — by time window")
        cell.font, cell.fill, cell.alignment, cell.border = SECTION_FONT, SECTION_FILL, LEFT, THIN
        for i in range(2, 5):
            c2 = ws.cell(row=r, column=i)
            c2.fill, c2.border = SECTION_FILL, THIN
        r += 1
        for p, lab in enumerate(labels):
            res = results.get(f"{v}__p{p}")
            vals = [f"  {lab}",
                    f"{res['coef']:.3f}" if res else "N/A",
                    fmt_effect(res),
                    fmt_p(res['p']) if res else "N/A"]
            font = SIG_FONT if (res and res['p'] < 0.05) else DATA_FONT
            for i, val in enumerate(vals, 1):
                cell = ws.cell(row=r, column=i, value=val)
                cell.font = font if i > 1 else DATA_FONT
                cell.alignment = LEFT if i == 1 else CENTER
                cell.border = THIN
            r += 1

    for col in model_cols:
        if col in tv_vars:
            continue
        res = results.get(col)
        vals = [label_map.get(col, col),
                f"{res['coef']:.3f}" if res else "N/A",
                fmt_effect(res), fmt_p(res['p']) if res else "N/A"]
        font = SIG_FONT if (res and res['p'] < 0.05) else DATA_FONT
        for i, val in enumerate(vals, 1):
            cell = ws.cell(row=r, column=i, value=val)
            cell.font = font if i > 1 else DATA_FONT
            cell.alignment = LEFT if i == 1 else CENTER
            cell.border = THIN
        r += 1

    r += 1
    footnotes = [
        f"Cox model with piecewise-constant coefficients for "
        f"{', '.join(tv_vars)}: follow-up is split at {', '.join(f'{c:g}' for c in cuts)} "
        f"{unit} and each window gets its own hazard ratio for those variables, so no "
        f"single ratio is assumed to hold across follow-up. All other covariates are "
        f"constant over time and their hazard ratios are as usual. Fitted on a "
        f"start/stop (counting-process) frame; each subject contributes at most "
        f"{len(labels)} risk intervals.",
        f"Variables were made time-varying "
        + ("because they violated the proportional-hazards assumption in the "
           "unremediated model (scaled Schoenfeld residuals)."
           if tv_source == 'auto_from_ph_test'
           else "as named in research_plan.json (outcomes[<id>].time_varying_variables)."),
    ]
    if cuts_source == 'default_rule':
        footnotes.append(
            f"⚠ The time windows were NOT declared in research_plan.json. The cut "
            f"point came from this script's default (the median event time), so it is "
            f"not covered by the plan fingerprint and cannot be shown to have been "
            f"chosen before the results were seen. Windows picked after looking at "
            f"where the effect changes are a subgroup search along the time axis, not "
            f"a remedy. Set outcomes[<id>].tv_cut_points; the manuscript gate blocks "
            f"until you do.")
    else:
        footnotes.append("The time windows were pre-declared in research_plan.json "
                         "(outcomes[<id>].tv_cut_points).")
    for v, h in (het or {}).items():
        if h.get('status') == 'computed':
            footnotes.append(
                f"{v}: hazard ratio in the last window relative to the first = "
                f"{h['hr_ratio_last_over_first']:.2f} (p = {fmt_p(h['p'])}). "
                + ("The windows differ, so the effect demonstrably moved."
                   if h['differs'] else
                   "The windows are NOT distinguishable at p < 0.05 — the model was "
                   "fitted, but this table does not establish that the effect moved; "
                   "report it as such."))
        else:
            footnotes.append(f"{v}: the windows could not be compared "
                             f"({h.get('reason', 'no reason given')}).")
    if config.get('source_table'):
        footnotes.append(
            f"This table is the pre-specified remedy for the proportional-hazards "
            f"violation reported in {os.path.basename(config['source_table'])}. That "
            f"workbook is left unremediated on purpose and its hazard ratios for "
            f"{', '.join(tv_vars)} are averages over a ratio that was changing; this "
            f"table is what should be reported for them.")
    if config.get('abbreviations'):
        footnotes.append(f"Abbreviations: {config['abbreviations']}")
    for note in footnotes:
        cell = ws.cell(row=r, column=1, value=note)
        cell.font, cell.alignment = SUB_FONT, LEFT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1

    ws2 = wb.create_sheet("Windows")
    for c, w in zip(range(1, 4), [30, 16, 16]):
        ws2.column_dimensions[get_column_letter(c)].width = w
    r = 1
    ws2.cell(row=r, column=1, value="Events per time window").font = TITLE_FONT
    r += 2
    for c, h in enumerate(["Window", "Risk intervals", "Events"], 1):
        cell = ws2.cell(row=r, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, CENTER, THIN
    r += 1
    for p, lab in enumerate(labels):
        info = events_by_period.get(p, {'rows': 0, 'events': 0})
        ws2.cell(row=r, column=1, value=lab).font = DATA_FONT
        ws2.cell(row=r, column=2, value=info['rows']).font = DATA_FONT
        cell = ws2.cell(row=r, column=3, value=info['events'])
        cell.font = SIG_FONT if info['events'] < MIN_EVENTS_PER_PERIOD else DATA_FONT
        for i in range(1, 4):
            ws2.cell(row=r, column=i).border = THIN
        r += 1
    if ph_art:
        r += 1
        g = ph_art.get('global') or {}
        ws2.cell(row=r, column=1,
                 value=f"Violation being remedied — unremediated model: global chi2 = "
                       f"{g.get('chi2')}, p = {g.get('p')}; violating covariates: "
                       f"{', '.join(ph_art.get('violating_variables') or []) or 'global only'}"
                 ).font = SUB_FONT
        ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)

    out = config['output_file']
    os.makedirs(os.path.dirname(os.path.abspath(out)) or ".", exist_ok=True)
    wb.save(out)
    return [ws.title, ws2.title]


# ═══════════════════════════════════════════════════════════════════════════════
# Evidence sidecar
# ═══════════════════════════════════════════════════════════════════════════════

def file_sha256(path):
    h = hashlib.sha256()
    try:
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 20), b""):
                h.update(chunk)
    except OSError:
        return None
    return h.hexdigest()


def write_evidence(config, config_path, payload):
    out = config['output_file']
    payload = dict(payload)
    payload.update({
        'schema': EVIDENCE_SCHEMA,
        'remedy': REMEDY,
        'tv_method': TV_METHOD,
        'generated_by': 'research-plan/scripts/ph_time_varying.py',
        'generated_at': datetime.now(timezone.utc).isoformat(),
        'outcome_id': config.get('outcome_id'),
        'table_basename': os.path.basename(out),
        'table_sha256': file_sha256(out) if os.path.exists(out) else None,
        'config_basename': os.path.basename(config_path),
        'config_sha256': file_sha256(config_path),
        'source_table': (os.path.basename(config['source_table'])
                         if config.get('source_table') else None),
    })
    path = config['evidence_file']
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    tmp = f"{path}.{os.getpid()}.tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)
        f.write("\n")
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)
    return path


def _not_computed(config, config_path, reason, **extra):
    payload = {'status': 'not_computed', 'reason': reason}
    payload.update(extra)
    path = write_evidence(config, config_path, payload)
    qc_blocking(f"time-varying model not computed: {reason}")
    print(f"Evidence written: {path} (status = not_computed)")
    return 2


def _not_applicable(config, config_path, reason, **extra):
    payload = {'status': 'not_applicable', 'reason': reason}
    payload.update(extra)
    path = write_evidence(config, config_path, payload)
    print(f"No proportional-hazards violation to remedy: {reason}")
    print(f"Evidence written: {path} (status = not_applicable)")
    return 0


# ═══════════════════════════════════════════════════════════════════════════════
# Preview / run
# ═══════════════════════════════════════════════════════════════════════════════

def preview(config, df, variables):
    print("TIME-VARYING COX PREVIEW (--check-missing)  |  mode = ph_time_varying")
    ph_art, _p, ph_err = read_ph_test(config)
    if ph_err:
        qc_warning(f"PH test artifact unreadable ({ph_err}); with "
                   f"time_varying_variables='auto' the run cannot choose variables.")
    tv_vars, source, err = resolve_tv_variables(config, ph_art)
    print(f"  time-varying vars : {tv_vars or '(unresolved)'}  [{source}]")
    if err:
        qc_warning(err)
    cs = config['cox_settings'] or {}
    cols = [v['column'] for v in variables] + [cs.get('time_column'),
                                               cs.get('event_column')]
    missing = [c for c in dict.fromkeys([c for c in cols if c]) if c not in df.columns]
    print(f"  rows loaded       : {len(df)}")
    print(f"  columns missing   : {missing or 'none'}")
    print(f"  cut_points        : {config.get('cut_points') or '(default rule)'}")
    if not config.get('cut_points'):
        qc_warning("cut_points are not declared in research_plan.json; the manuscript "
                   "gate will block until outcomes[<id>].tv_cut_points is set.")
    print("[preview] No model was fitted (--check-missing). No Excel was written.")


def run(config, config_path, df, variables):
    ph_art, ph_path, ph_err = read_ph_test(config)
    if ph_err:
        qc_warning(f"PH test artifact: {ph_err}")

    ph_status = (ph_art or {}).get('status')
    if ph_art and ph_status == 'computed' and not ph_art.get('violated'):
        return _not_applicable(
            config, config_path,
            f"the proportional-hazards assumption was not violated in "
            f"{os.path.basename(ph_path)}, so there is nothing for the declared "
            f"'time_varying' remedy to remedy",
            source_ph_test=os.path.basename(ph_path),
            source_ph_test_sha256=file_sha256(ph_path),
            ph_violating_variables=list(ph_art.get('violating_variables') or []))
    if ph_art and ph_status not in (None, 'computed'):
        return _not_computed(
            config, config_path,
            f"the upstream PH test is status={ph_status!r} "
            f"({ph_art.get('reason') or 'no reason given'}), so whether the assumption "
            f"is violated is unknown. An unknown violation is not an absent one.",
            source_ph_test=os.path.basename(ph_path) if ph_path else None)

    tv_vars, tv_source, err = resolve_tv_variables(config, ph_art)
    if err:
        return _not_computed(config, config_path, err,
                             time_varying_variables_source=tv_source)

    d, model_cols, label_map, derr = build_wide(df, config, variables, tv_vars)
    if derr:
        return _not_computed(config, config_path, derr,
                             time_varying_variables=tv_vars,
                             time_varying_variables_source=tv_source)

    cs = config['cox_settings']
    time_col, event_col = cs['time_column'], cs['event_column']
    cuts, cuts_source, cerr = resolve_cut_points(config, d, time_col, event_col)
    if cerr:
        return _not_computed(config, config_path,
                             f"cut points could not be resolved: {cerr}",
                             time_varying_variables=tv_vars,
                             cut_points_source=cuts_source)
    if len(cuts) + 1 > MAX_PERIODS:
        qc_warning(f"{len(cuts) + 1} time windows: each one is estimated from a slice "
                   f"of the events, and the confidence intervals widen accordingly.")

    long, fit_cols, serr = split_episodes(d, config, model_cols, tv_vars, cuts)
    if serr:
        return _not_computed(config, config_path, serr,
                             time_varying_variables=tv_vars, cut_points=cuts,
                             cut_points_source=cuts_source)

    events_by_period = {}
    for p, sub in long.groupby('_period'):
        events_by_period[int(p)] = {'rows': int(len(sub)),
                                    'events': int(sub['_event'].sum())}
    thin = sorted([p for p, v in events_by_period.items()
                   if v['events'] < MIN_EVENTS_PER_PERIOD])
    if thin:
        qc_warning(f"time window(s) {thin} carry fewer than {MIN_EVENTS_PER_PERIOD} "
                   f"events; the hazard ratio there is estimated from very little and "
                   f"its interval will be wide. Consider fewer, wider windows.")

    results, ctv_or_reason = fit_time_varying(long, fit_cols)
    if results is None:
        return _not_computed(config, config_path,
                             f"the time-varying Cox model did not fit: {ctv_or_reason}",
                             time_varying_variables=tv_vars, cut_points=cuts,
                             cut_points_source=cuts_source)
    ctv = ctv_or_reason

    labels = period_labels(cuts, config.get('time_unit') or '')
    het = {v: heterogeneity(ctv, v, len(labels)) for v in tv_vars}
    for v, h in het.items():
        if h.get('status') == 'computed' and not h['differs']:
            qc_warning(f"{v}: the window-specific coefficients are not distinguishable "
                       f"(p = {h['p']:.3g}). The model was fitted, but it does not "
                       f"establish that the effect moved — 'we ran a time-varying "
                       f"model' is not the same finding as 'the effect changed'.")
    if cuts_source == 'default_rule':
        qc_warning(f"the time windows came from this script's default (median event "
                   f"time = {cuts[0]:.4g}), not from research_plan.json, so they are "
                   f"outside the plan fingerprint and cannot be shown to predate the "
                   f"result. Set outcomes[<id>].tv_cut_points — the manuscript gate "
                   f"blocks until you do.")

    n, n_events = int(len(d)), int(d[event_col].sum())
    sheets = write_excel(config, results, label_map, model_cols, tv_vars, tv_source,
                         cuts, cuts_source, labels, het, n, n_events,
                         events_by_period, int(len(long)),
                         config.get('time_unit') or '', ph_art)

    evidence_path = write_evidence(config, config_path, {
        'status': 'computed',
        'method': 'Cox model with piecewise-constant coefficients over pre-specified '
                  'time windows (lifelines CoxTimeVaryingFitter on a start/stop frame)',
        'time_varying_variables': tv_vars,
        'time_varying_variables_source': tv_source,
        'ph_violating_variables': list((ph_art or {}).get('violating_variables') or []),
        'cut_points': cuts,
        'cut_points_source': cuts_source,
        'period_labels': labels,
        'n': n,
        'n_events': n_events,
        'n_risk_intervals': int(len(long)),
        'events_by_period': {str(k): v for k, v in events_by_period.items()},
        'covariates_in_model': [c for c in model_cols if c not in tv_vars],
        'period_effects': {
            v: [{'window': labels[p],
                 'hr': (results.get(f"{v}__p{p}") or {}).get('effect'),
                 'ci': [(results.get(f"{v}__p{p}") or {}).get('ci_lower'),
                        (results.get(f"{v}__p{p}") or {}).get('ci_upper')],
                 'p': (results.get(f"{v}__p{p}") or {}).get('p')}
                for p in range(len(labels))]
            for v in tv_vars},
        'heterogeneity': het,
        'source_ph_test': os.path.basename(ph_path) if ph_path else None,
        'source_ph_test_sha256': file_sha256(ph_path) if ph_path else None,
        'time_column': time_col,
        'event_column': event_col,
    })

    print(f"\nwindows: {', '.join(labels)}   (cut points {cuts}, {cuts_source})")
    for v in tv_vars:
        for p, lab in enumerate(labels):
            res = results.get(f"{v}__p{p}")
            print(f"  {v} [{lab}]: {fmt_effect(res)}  p = "
                  f"{fmt_p(res['p']) if res else '—'}")
    print(f"\nExcel written    : {config['output_file']}")
    print(f"  Sheets         : {', '.join(sheets)}")
    print(f"Evidence written : {evidence_path}")
    return 0


def main():
    if len(sys.argv) < 2:
        print("Usage: python ph_time_varying.py config.json [--check-missing]")
        sys.exit(1)

    config_path = sys.argv[1]
    check_only = '--check-missing' in sys.argv

    config = load_config(config_path)
    df = load_data(config)
    variables = collect_variables(config)

    print(f"Mode: {config['mode']}  (remedy = {config['remedy']})")
    print(f"Data loaded: n={len(df)}")

    if check_only:
        preview(config, df, variables)
        sys.exit(0)

    rc = run(config, config_path, df, variables)

    print("\nDone!")
    if _QC_MESSAGES:
        n_block = sum(1 for m in _QC_MESSAGES if m.startswith('[QC_BLOCKING]'))
        n_warn = len(_QC_MESSAGES) - n_block
        print(f"  QC: {n_block} blocking, {n_warn} warning(s) written to stderr")
    sys.exit(rc)


if __name__ == '__main__':
    main()
