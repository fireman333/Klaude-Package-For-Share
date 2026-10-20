#!/usr/bin/env python3
"""STROBE checklist — can the pipeline's outputs actually fill the reporting form?

Spec: references/advanced_analyses.md section 2.

Usage:
    python strobe.py config.json [--check-missing]

Outputs (all three written together; the table is for people, the JSON is for the
downstream sign-off gate):
    <output_file>            .xlsx checklist
    <output_file>.md         markdown checklist  (or config.md_file)
    <output_file>_rows.json  machine-readable rows (or config.json_file) — schema
                             documented at write_rows_json(); consumers must read
                             this rather than re-parsing the markdown table.

Every `covered` item must name the artifact that covers it. A ticked box with
nothing behind it is worse than an empty one: it reads as done. Items the
pipeline cannot know — setting, recruitment period, funding, ethics — are
reported as `missing` on purpose rather than filled with something plausible.

The checklist is built for the study's design. STROBE's combined form words
items 6a, 6b, 12d, 14c and 15 differently for cohort, case-control and
cross-sectional studies, and leaves two of them blank outside cohort studies;
`study.design` in the plan selects the variant. A design string that cannot be
mapped raises a QC warning and falls back to listing every item with all three
wordings — never to marking rows not applicable on a guess.
"""

import json
import os
import sys
from datetime import datetime

_QC = []


def qc(kind, msg):
    line = f"[QC_{kind}] {msg}"
    _QC.append(line)
    print(line, file=sys.stderr)


# ── The checklist ─────────────────────────────────────────────────────────────
# (item, sub, section, paraphrased requirement, resolver key)
# Requirements are paraphrased, not reproduced from the STROBE statement.
#
# Five rows differ by study design in the official combined checklist — 6a, 6b,
# 12d, 14c and 15. They carry the DESIGN marker here and are filled in by
# items_for_design(); everything else is shared. Before 2026-08-11 this file had
# one fixed cohort-flavoured list, so a case-control study was asked about loss
# to follow-up and a cross-sectional study about matching criteria — items that
# do not apply to it, listed as if they did.
DESIGN = object()

ITEMS = [
    (1, 'a', 'Title and abstract', '標題或摘要要以常用詞彙標明研究設計', 'design_word'),
    (1, 'b', 'Title and abstract', '摘要要有平衡且資訊充分的總結', 'human'),
    (2, '', 'Introduction', '研究的科學背景與理由', 'human'),
    (3, '', 'Introduction', '明確的目標，含任何事前指定的假說', 'objectives'),
    (4, '', 'Methods', '研究設計的關鍵要素，在文章早期陳述', 'design'),
    (5, '', 'Methods', '場域、地點、招募與追蹤的相關日期', 'human'),
    (6, 'a', 'Methods', DESIGN, '6a'),
    (6, 'b', 'Methods', DESIGN, '6b'),
    (7, '', 'Methods', '結果、暴露、預測因子、干擾因子與效應修飾因子的定義', 'variables'),
    (8, '', 'Methods', '每個變項的資料來源與測量方法', 'human'),
    (9, '', 'Methods', '為處理潛在偏差所做的努力', 'bias'),
    (10, '', 'Methods', '樣本數是怎麼決定的', 'human'),
    (11, '', 'Methods', '連續變項怎麼處理；有分組時說明分組方式與理由', 'quantitative'),
    (12, 'a', 'Methods', '所有統計方法，含控制干擾的方法', 'stats_confounding'),
    (12, 'b', 'Methods', '檢查亞群與交互作用的方法', 'stats_subgroup'),
    (12, 'c', 'Methods', '缺失資料如何處理', 'stats_missing'),
    (12, 'd', 'Methods', DESIGN, '12d'),
    (12, 'e', 'Methods', '任何敏感度分析', 'stats_sensitivity'),
    (13, 'a', 'Results', '各階段人數：合格、納入、完成追蹤、納入分析', 'flow'),
    (13, 'b', 'Results', '各階段退出的理由', 'flow_reasons'),
    (13, 'c', 'Results', '建議畫流程圖', 'flow_diagram'),
    (14, 'a', 'Results', '參與者特徵與暴露、潛在干擾因子的資訊', 'table1'),
    (14, 'b', 'Results', '每個變項的缺失人數', 'missing_counts'),
    (14, 'c', 'Results', DESIGN, '14c'),
    (15, '', 'Results', DESIGN, '15'),
    (16, 'a', 'Results', '未校正與校正後的估計值及其精確度；校正了哪些干擾因子、為什麼', 'main_results'),
    (16, 'b', 'Results', '連續變項被分類時，報告分界點', 'category_bounds'),
    (16, 'c', 'Results', '有意義時，把相對風險換算成絕對風險', 'absolute_risk'),
    (17, '', 'Results', '其他分析：亞群、交互作用、敏感度分析', 'other_analyses'),
    (18, '', 'Discussion', '對照研究目標總結關鍵結果', 'human'),
    (19, '', 'Discussion', '研究限制，含潛在偏差與不精確的來源與方向', 'limitations'),
    (20, '', 'Discussion', '謹慎的整體詮釋，考量目標、限制、多重分析與其他證據', 'human'),
    (21, '', 'Discussion', '結果的外推性', 'human'),
    (22, '', 'Other', '本研究的資金來源與資助者的角色', 'human'),
]

DESIGNS = ('cohort', 'case-control', 'cross-sectional')

# marker -> design -> (requirement, resolver key).
# `unknown` is used when study.design cannot be mapped: it keeps every row,
# states all three variants, and never claims a row is not applicable — an
# unrecognised design is a reason to ask the author, not to drop items.
BY_DESIGN = {
    '6a': {
        'cohort': ('納入條件、參與者來源與選取方法，以及追蹤的方式', 'population'),
        'case-control': ('納入條件、病例確認與對照選取的來源與方法；並說明選擇這組病例與對照的理由',
                         'population'),
        'cross-sectional': ('納入條件、參與者來源與選取方法', 'population'),
        'unknown': ('納入條件與參與者來源；世代研究另述追蹤方式，病例對照另述病例確認與對照選取及其理由',
                    'population'),
    },
    '6b': {
        'cohort': ('配對研究：配對條件、暴露組與非暴露組的人數', 'matching'),
        'case-control': ('配對研究：配對條件、每個病例配到幾個對照', 'matching_case_control'),
        'cross-sectional': ('（橫斷面研究不適用：官方 checklist 此格留白）', 'not_applicable'),
        'unknown': ('配對研究：配對條件與各組人數（橫斷面研究不適用）', 'matching_unknown'),
    },
    '12d': {
        'cohort': ('若有追蹤流失，說明如何處理', 'stats_ltfu'),
        'case-control': ('若為配對設計，說明分析如何處理病例與對照的配對', 'stats_matching_analysis'),
        'cross-sectional': ('若有抽樣設計，說明分析方法如何反映抽樣策略', 'stats_sampling'),
        'unknown': ('依設計而定：追蹤流失（世代）／配對如何處理（病例對照）／抽樣策略（橫斷面）',
                    'stats_ltfu'),
    },
    '14c': {
        'cohort': ('追蹤時間總結（平均與總量）', 'followup'),
        'case-control': ('（病例對照研究不適用：官方 checklist 此格留白）', 'not_applicable'),
        'cross-sectional': ('（橫斷面研究不適用：官方 checklist 此格留白）', 'not_applicable'),
        'unknown': ('追蹤時間總結（僅世代研究適用）', 'followup'),
    },
    '15': {
        'cohort': ('結果事件數，或隨時間的摘要指標', 'outcome_data'),
        'case-control': ('各暴露類別的人數，或暴露的摘要指標', 'exposure_data'),
        'cross-sectional': ('結果事件數或摘要指標', 'outcome_data'),
        'unknown': ('結果事件數或摘要指標；病例對照研究改報各暴露類別的人數', 'outcome_data'),
    },
}


def normalize_design(raw):
    """Map study.design onto one of the three STROBE variants.

    Returns (design, note). `design` is one of DESIGNS, or 'unknown' when the
    string cannot be mapped — in which case `note` says why, and the caller
    raises it as a QC warning. Guessing a design silently is the failure this
    function exists to prevent: it decides which rows get marked not applicable.
    """
    s = str(raw or '').strip().lower().replace('_', ' ').replace('-', ' ')
    if not s:
        return 'unknown', ('research_plan.json$.study.design 未填，無法選擇 STROBE 的設計版本；'
                           '本表列出全部項目、不標任何「不適用」')
    # Ordered from most to least specific. `nested case-control` and `target
    # trial emulation` both contain a word that would otherwise misroute them.
    # `case cohort` before the cohort test: it contains the word but is neither
    # a plain cohort nor a case-control study, and STROBE's combined checklist
    # has no column for it. Silently filing it under cohort would apply the wrong
    # wording to 6a/12d/15 and mark 6b n/a on a guess.
    if 'case cohort' in s:
        return 'unknown', (f"study.design = {raw!r} is a case-cohort design, which the "
                           f"STROBE combined checklist does not cover as its own variant. "
                           f"本表列出全部項目、不依設計標任何「不適用」；請自行判斷 6a / 6b / "
                           f"12d / 14c / 15 應採哪一版措辭")
    if 'case control' in s or 'case referent' in s:
        return 'case-control', None
    if 'cross sectional' in s or 'crosssectional' in s or 'prevalence' in s:
        return 'cross-sectional', None
    if 'target trial' in s or 'cohort' in s or 'longitudinal' in s or 'follow up' in s:
        return 'cohort', None
    if 'rct' in s or 'randomi' in s or 'trial' in s:
        return 'unknown', (f"study.design = {raw!r} 看起來是隨機試驗。STROBE 是觀察性研究的"
                           f"報告準則，隨機試驗應改用 CONSORT；本表列出全部項目、不標「不適用」")
    return 'unknown', (f"study.design = {raw!r} 無法對應到 cohort / case-control / "
                       f"cross-sectional；本表列出全部項目、不標任何「不適用」")


def items_for_design(design):
    """Resolve the DESIGN-marked rows for one design, keeping row order."""
    variant = design if design in DESIGNS else 'unknown'
    out = []
    for item, sub, section, req, key in ITEMS:
        if req is DESIGN:
            req, key = BY_DESIGN[key][variant]
        out.append((item, sub, section, req, key))
    return out


def _artifact(output_dir, *names):
    """Return the first artifact that exists, as a display string."""
    for n in names:
        p = os.path.join(output_dir, n)
        if os.path.exists(p):
            return n
    return None


def _glob_like(output_dir, prefix, suffix='.xlsx'):
    if not os.path.isdir(output_dir):
        return []
    return sorted(f for f in os.listdir(output_dir)
                  if f.startswith(prefix) and f.endswith(suffix))


def resolve(key, plan, out):
    """Return (status, evidence). status in covered / partial / missing."""
    D = plan.get('output_dir', '')
    flow = _artifact(D, 'Flowchart.xlsx')
    t1 = _artifact(D, 'Table1.xlsx')
    miss = _artifact(D, 'Missingness.xlsx')
    mis = _glob_like(D, 'MI_')
    uni = _glob_like(D, 'Uni_')
    multi = _glob_like(D, 'Multi_')
    km = _glob_like(D, 'KM_') + _glob_like(D, 'CR_')
    psm = _artifact(D, 'PSM.xlsx')
    auc = _artifact(D, 'AUC.xlsx')
    evs = _glob_like(D, 'E_value')
    ev = evs[0] if evs else None
    rcs = _glob_like(D, 'RCS_')
    land = _glob_like(D, 'Landmark_')
    wt = _glob_like(D, 'IPTW_') + _glob_like(D, 'Weighting_')
    rob = _glob_like(D, 'Robust_')
    dlog = out.get('decision_log')

    if key == 'human':
        return 'missing', '本流程無從得知，必須由作者撰寫'
    if key == 'not_applicable':
        return 'n/a', '本研究設計不適用（STROBE 官方 checklist 於此設計留白）'

    if key == 'design_word':
        d = (plan.get('study') or {}).get('design')
        return ('covered', f"research_plan.json$.study.design = {d!r}") if d else \
               ('missing', '計畫未記錄 study.design')
    if key == 'design':
        d = (plan.get('study') or {}).get('design')
        return ('covered', f"study.design = {d!r}；設計要素見 decision_log.md C 節") if d else \
               ('missing', '計畫未記錄 study.design')
    if key == 'objectives':
        t = (plan.get('study') or {}).get('title')
        prim = ((plan.get('outcomes') or {}).get('primary') or {}).get('name')
        if t and prim:
            return 'partial', (f"目標與主要終點在 research_plan.json（title、outcomes.primary="
                               f"{prim!r}）；**事前假說的文字**仍須作者撰寫")
        return 'missing', '計畫缺 title 或 primary outcome'
    if key == 'population':
        n_f = len((plan.get('population') or {}).get('filters', []) or [])
        n_e = len((plan.get('population') or {}).get('exclusion_steps', []) or [])
        if flow:
            return 'covered', f"{flow}；納入排除條件 {n_f} 個 filter、{n_e} 個排除步驟"
        return 'partial', f"計畫有 {n_f} 個 filter、{n_e} 個排除步驟，但流程圖尚未產出"
    if key == 'matching':
        # Matched cohort: the PSM step is the matching, so the artifact answers
        # the question. This does not carry over to case-control — see below.
        if psm:
            return 'covered', f"{psm}（配對條件與配對後人數）"
        return 'n/a', ('**假設非配對設計**（本流程未產出任何配對產物）。若本研究其實有配對，'
                       '配對條件與各組人數必須自行補上——這格的「不適用」是推論，不是查核結果')
    if key == 'matching_case_control':
        # PSM.xlsx must not satisfy this row. Analysis-stage propensity-score
        # matching and design-stage case-control matching are different things:
        # the row asks how cases and controls were sampled, which no artifact
        # this pipeline produces can answer. Same distinction as item 12d.
        extra = (f"（注意：{psm} 是分析階段的傾向分數配對，回答的不是本項）" if psm else "")
        return 'missing', ('**配對條件與每個病例的對照數必須由作者填寫**——本流程產不出病例對照的'
                           '抽樣設計資訊。' + extra + '若為非配對設計，請在論文中明說本項不適用')
    if key == 'matching_unknown':
        # study.design could not be mapped, so whether this row applies is exactly
        # what is unknown. Marking it n/a would be the design guess this file
        # refuses to make elsewhere.
        return 'partial', ('**設計未知，無法判斷本項是否適用**（橫斷面研究此格留白，世代與病例對照'
                           '則需填寫）。先把 research_plan.json$.study.design 填成 cohort / '
                           'case-control / cross-sectional，再重跑本步驟')
    if key == 'variables':
        cov = plan.get('covariates') or []
        return 'covered', (f"research_plan.json：outcomes、grouping、covariates（{len(cov)} 個）；"
                           f"欄位對應表見 Step 2 的確認紀錄")
    if key == 'bias':
        parts = []
        if psm:
            parts.append(psm)
        if wt:
            parts.append(wt[0])
        if land:
            parts.append(land[0])
        if ev:
            parts.append(ev)
        if parts:
            return 'partial', (f"已量化的偏差處理：{', '.join(parts)}；**偏差來源的敘述**仍須作者撰寫")
        return 'missing', '沒有偏差處理的產物；敘述須作者撰寫'
    if key == 'quantitative':
        if rcs:
            return 'covered', f"{rcs[0]}（連續變項以 RCS 建模，含非線性檢定與參考值）"
        scaled = [c for c in (plan.get('covariates') or []) if c.get('scale')]
        return 'partial', (f"{len(scaled)} 個連續變項有 scale 設定（線性假設）；"
                           f"**未做非線性檢定**——跑 rcs 步驟可補足本項")
    if key == 'stats_confounding':
        if multi:
            return 'covered', f"{', '.join(multi)}（校正模型與共變項集合）"
        return 'missing', '尚無多變量模型產物'
    if key == 'stats_subgroup':
        if multi:
            return 'covered', f"{multi[0]} 的 Interaction Analysis 分頁"
        return 'missing', '尚無交互作用分析產物'
    if key == 'stats_missing':
        if mis:
            return 'covered', f"{miss or 'Missingness.xlsx'} + {', '.join(mis)}（MI 與 complete-case 對照）"
        if miss:
            return 'partial', f"{miss}（有評估，但未做多重插補；主分析仍為 complete-case）"
        return 'missing', '**沒有缺失資料的任何產物**——reviewer 必問。跑 missingness 步驟'
    if key == 'stats_ltfu':
        if km:
            return 'partial', (f"{km[0]} 有追蹤與 censoring 資訊；**流失是否與結果相關**"
                               f"須另行檢視（Missingness.xlsx 的 Mechanism_tests 可作參考）")
        return 'missing', '尚無存活分析產物'
    if key == 'stats_matching_analysis':
        # "If applicable" — the pipeline cannot tell a matched case-control study
        # from an unmatched one, and the models it fits are unconditional. Saying
        # so is the honest answer; silently ticking the box is not.
        if psm:
            return 'partial', (f"{psm} 是分析階段的傾向分數配對，**不等於**病例對照的設計配對。"
                               f"若病例與對照是配對抽樣，分析須用條件式邏輯迴歸或依配對層分層——"
                               f"本流程的模型是非條件式的")
        return 'missing', ('**配對的病例對照研究必須在分析中處理配對**（條件式邏輯迴歸／配對層分層），'
                           '本流程沒有這類產物。若本研究是非配對設計，請在論文中明說本項不適用')
    if key == 'stats_sampling':
        if rob:
            return 'partial', (f"{rob[0]}（群集穩健變異／GEE）可反映集群抽樣；**分層或不等機率抽樣的"
                               f"權重本流程未處理**，須確認研究是否為簡單抽樣")
        return 'missing', ('若採複雜抽樣（分層、集群、不等機率），分析須反映抽樣設計；'
                           '本流程未處理調查權重。簡單抽樣請在論文中明說本項不適用')
    if key == 'exposure_data':
        if t1:
            return 'covered', f"{t1}（依病例／對照分組，列出各暴露類別的人數與摘要指標）"
        return 'missing', '尚無 Table 1；病例對照研究要報各暴露類別的人數'
    if key == 'stats_sensitivity':
        parts = [x for x in [ev, psm, (wt[0] if wt else None), (land[0] if land else None),
                             (mis[0] if mis else None)] if x]
        if len(parts) >= 2:
            return 'covered', f"敏感度分析：{', '.join(parts)}"
        if parts:
            return 'partial', (f"只有一項敏感度分析（{parts[0]}）。單一 PSM 或單一 MI "
                               f"不構成 sensitivity framework")
        return 'missing', '**沒有任何敏感度分析**'
    if key == 'flow':
        return ('covered', flow) if flow else ('missing', '尚無流程圖產物')
    if key == 'flow_reasons':
        # The ternary here used to bind to the second element of the tuple, so a
        # run with no flowchart returned ('partial', ('missing', ...)) and blew up
        # in openpyxl. Written out longhand so it cannot happen again.
        n_e = len((plan.get('population') or {}).get('exclusion_steps', []) or [])
        if not flow:
            return 'missing', '尚無流程圖'
        if n_e:
            return 'covered', f"{flow}（{n_e} 個排除步驟各有標籤與人數）"
        return 'partial', f"{flow} 已產出，但計畫沒有逐步排除條件，退出理由無法逐步標註"
    if key == 'flow_diagram':
        p = _artifact((plan.get('output_dir') or ''), 'Flowchart.png')
        return ('covered', p) if p else ('partial', '有流程表但無圖檔')
    if key == 'table1':
        return ('covered', t1) if t1 else ('missing', '尚無 Table 1')
    if key == 'missing_counts':
        if miss:
            return 'covered', f"{miss} 的 Summary 分頁（逐變項、且依組別拆開）"
        if t1:
            return 'partial', f"{t1} 可能含 available-case 的 N，但沒有逐變項缺失表"
        return 'missing', '沒有缺失人數的產物'
    if key == 'followup':
        if km:
            return 'partial', f"{km[0]} 含追蹤時間；**中位追蹤與總人年**須確認是否已列出"
        return 'missing', '尚無存活分析產物'
    if key == 'outcome_data':
        if km or multi:
            return 'covered', f"{', '.join((km + multi)[:3])}"
        return 'missing', '尚無結果資料產物'
    if key == 'main_results':
        if uni and multi:
            return 'covered', f"未校正 {uni[0]}；校正後 {multi[0]}"
        if multi:
            return 'partial', f"只有校正後結果（{multi[0]}）；**STROBE 要求同時報未校正**'"
        return 'missing', '尚無主要結果產物'
    if key == 'category_bounds':
        cat = [c for c in (plan.get('covariates') or [])
               if c.get('type') in ('ordinal', 'categorical')]
        if cat:
            return 'partial', (f"{len(cat)} 個類別變項；**分界點與 reference level** 見 "
                               f"research_plan.json$.covariates，須確認已寫進論文")
        return 'n/a', '沒有連續變項被分類，本項不適用'
    if key == 'absolute_risk':
        if km:
            return 'covered', f"{km[0]}（KM/CIF 提供絕對風險）"
        return 'partial', '只有相對效應量；建議補時間點的絕對風險'
    if key == 'other_analyses':
        parts = [x for x in [psm, auc, ev, (wt[0] if wt else None), (rcs[0] if rcs else None),
                             (land[0] if land else None)] if x]
        if parts:
            return 'covered', ', '.join(parts)
        return 'missing', '尚無其他分析產物'
    if key == 'limitations':
        if dlog:
            return 'partial', f"{dlog} 的 F 節（已知限制）；**論文的限制段落**仍須作者撰寫"
        return 'missing', '沒有 decision_log.md；限制須作者自行撰寫'
    return 'missing', '未對應'


ICON = {'covered': '✓', 'partial': '◐', 'missing': '✗', 'n/a': '–'}


def build(plan, out, items):
    rows = []
    for item, sub, section, req, key in items:
        status, ev = resolve(key, plan, out)
        rows.append({'item': f"{item}{sub}", 'section': section, 'requirement': req,
                     'status': status, 'evidence': ev, 'kind': key})
    return rows


# ── strobe_rows.json ──────────────────────────────────────────────────────────
# Machine-readable twin of the .md / .xlsx checklist, written alongside them (not
# instead of: the table is what a human reads). It exists because the downstream
# sign-off gate — biostats-guardian/scripts/strobe_signoff.py — used to recover
# these fields by parsing the markdown table, and a column reorder or a wording
# change in write_md() would have broken that silently.
#
# schema strobe_rows/1:
#   schema        str    "strobe_rows/1"
#   generated_at  str    ISO-8601, seconds
#   study         obj    title, design_declared (plan's raw string),
#                        design_variant (the STROBE variant actually used),
#                        design_note (set when the design string could not be mapped)
#   counts        obj    covered / partial / missing / n-a / total
#   rows          list   one per checklist item, in checklist order:
#     item        str    "1a", "12d", ...
#     section     str    "Methods", "Results", ...
#     requirement str    paraphrased requirement (design-specific where it varies)
#     status      str    covered | partial | missing | n/a
#     evidence    str    the artifact that covers it, or what is still owed
#     kind        str    resolver key. "human" marks items no pipeline can produce
#                        (setting, funding, interpretation) — they are `missing` on
#                        purpose and are the sign-off's subject, not defects. This is
#                        the authoritative source for that distinction; consumers
#                        must not maintain their own copy of the list.
#   artifacts     obj    paths of the sibling outputs, for provenance
STROBE_ROWS_SCHEMA = 'strobe_rows/1'


def write_rows_json(rows, path, plan, design, note, artifacts):
    n = {k: sum(1 for r in rows if r['status'] == k) for k in ICON}
    payload = {
        'schema': STROBE_ROWS_SCHEMA,
        'generated_at': datetime.now().isoformat(timespec='seconds'),
        'study': {
            'title': (plan.get('study') or {}).get('title', ''),
            'design_declared': (plan.get('study') or {}).get('design'),
            'design_variant': design,
            'design_note': note or None,
        },
        'counts': {'covered': n['covered'], 'partial': n['partial'],
                   'missing': n['missing'], 'n/a': n['n/a'], 'total': len(rows)},
        'rows': rows,
        'artifacts': artifacts,
    }
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)
        f.write('\n')
    return path


def write_md(rows, path, plan, design, note):
    L = [f"# STROBE checklist — {(plan.get('study') or {}).get('title', '')}\n",
         f"_產生於 {datetime.now().isoformat(timespec='seconds')}；"
         f"計畫登記的設計：{(plan.get('study') or {}).get('design', '?')}　→　"
         f"採用的 STROBE 版本：**{design}**_\n",
         "`✓ covered` 有具體產物可指　`◐ partial` 有材料但要人補文字　"
         "`✗ missing` 本流程產不出來，必須作者撰寫　`– n/a` 本研究設計不適用\n"]
    if note:
        L.append(f"> ⚠ {note}\n")
    n = {k: sum(1 for r in rows if r['status'] == k) for k in ICON}
    L.append(f"**covered {n['covered']}　partial {n['partial']}　missing {n['missing']}"
             f"　n/a {n['n/a']}　（共 {len(rows)} 項）**\n")
    L.append("| # | 區段 | 要求 | 狀態 | 證據 / 待辦 |")
    L.append("|---|---|---|---|---|")
    for r in rows:
        L.append(f"| {r['item']} | {r['section']} | {r['requirement']} | "
                 f"{ICON[r['status']]} {r['status']} | {r['evidence']} |")
    L.append("\n> `covered` 一定附了指標。**打勾但指不出東西的 checklist 比沒有更糟**——"
             "它會讓人以為做過了。`missing` 的項目是刻意留白，不要用似是而非的文字填滿。")
    with open(path, 'w', encoding='utf-8') as f:
        f.write("\n".join(L))


def write_excel(rows, path, plan, design, note):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = 'STROBE'
    thin = Side(style='thin', color='BFBFBF')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws['A1'] = f"STROBE checklist — {(plan.get('study') or {}).get('title', '')}"
    ws['A1'].font = Font(size=13, bold=True, color='1F3864')
    ws['A2'] = (f"計畫登記的設計 {(plan.get('study') or {}).get('design', '?')}　→　"
                f"採用的 STROBE 版本 {design}　|　"
                f"covered 一定附產物指標；missing 是刻意留白，須作者撰寫")
    ws['A2'].font = Font(size=10, italic=True, color='595959')
    if note:
        ws['A3'] = f"⚠ {note}"
        ws['A3'].font = Font(size=10, bold=True, color='9C0006')

    heads = ['#', '區段', '要求', '狀態', '證據 / 待辦']
    for i, (h, w) in enumerate(zip(heads, [7, 14, 46, 12, 74]), start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='2E5FA3')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = box
        ws.column_dimensions[chr(64 + i)].width = w

    fills = {'covered': 'E8F5E9', 'partial': 'FFF8E1', 'missing': 'FDECEA',
             'n/a': 'F2F2F2'}
    r = 5
    for row in rows:
        for i, v in enumerate([row['item'], row['section'], row['requirement'],
                               f"{ICON[row['status']]} {row['status']}", row['evidence']],
                              start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = box
            c.alignment = Alignment(vertical='top', wrap_text=(i in (3, 5)),
                                    horizontal='center' if i in (1, 4) else 'left')
            c.fill = PatternFill('solid', fgColor=fills[row['status']])
        r += 1
    os.makedirs(os.path.dirname(os.path.abspath(path)) or '.', exist_ok=True)
    wb.save(path)


def main():
    if len(sys.argv) < 2:
        print('Usage: python strobe.py config.json [--check-missing]')
        sys.exit(1)
    with open(sys.argv[1], 'r', encoding='utf-8') as f:
        config = json.load(f)

    plan_path = config.get('plan_path')
    if not plan_path or not os.path.exists(plan_path):
        qc('BLOCKING', f"plan_path not found: {plan_path!r}. The checklist is built from "
                       f"research_plan.json plus the artifacts on disk.")
        sys.exit(3)
    with open(plan_path, 'r', encoding='utf-8') as f:
        plan = json.load(f)
    plan.setdefault('output_dir', config.get('output_dir', os.path.dirname(plan_path)))

    out = {'decision_log': None}
    for cand in ('report/decision_log.md', 'decision_log.md',
                 'analysis/report/decision_log.md'):
        p = os.path.join(plan['output_dir'], cand)
        if os.path.exists(p):
            out['decision_log'] = cand
            break

    design, design_note = normalize_design((plan.get('study') or {}).get('design'))
    if design_note:
        qc('WARNING', design_note)
    items = items_for_design(design)
    rows = build(plan, out, items)
    n = {k: sum(1 for r in rows if r['status'] == k) for k in ICON}

    if '--check-missing' in sys.argv:
        print('=' * 78)
        print('STROBE PREVIEW (--check-missing)')
        print('=' * 78)
        print(f"  Design      : {(plan.get('study') or {}).get('design')} "
              f"-> STROBE variant {design}")
        print(f"  Output dir  : {plan['output_dir']}")
        print(f"  covered {n['covered']}  partial {n['partial']}  "
              f"missing {n['missing']}  n/a {n['n/a']}")
        for r in rows:
            if r['status'] == 'missing':
                print(f"      ✗ {r['item']:<4s} {r['requirement'][:52]}")
        sys.exit(0)

    write_excel(rows, config['output_file'], plan, design, design_note)
    md = config.get('md_file') or os.path.splitext(config['output_file'])[0] + '.md'
    write_md(rows, md, plan, design, design_note)
    rows_json = (config.get('json_file')
                 or os.path.splitext(config['output_file'])[0] + '_rows.json')
    write_rows_json(rows, rows_json, plan, design, design_note,
                    {'excel': config['output_file'], 'md': md,
                     'plan': plan_path, 'output_dir': plan['output_dir']})

    # Items the pipeline could have covered but did not are the actionable ones;
    # the human-only items are expected to be missing and are not warned about.
    actionable = [r for r, (_, _, _, _, k) in zip(rows, items)
                  if r['status'] == 'missing' and k != 'human']
    for r in actionable:
        qc('WARNING', f"STROBE {r['item']} not covered by any artifact: {r['evidence']}")
    print(f"  Excel : {config['output_file']}")
    print(f"  MD    : {md}")
    print(f"  JSON  : {rows_json}")
    print(f"  covered {n['covered']}  partial {n['partial']}  missing {n['missing']}"
          f"  n/a {n['n/a']}  (of {len(rows)})")
    print('\nDone!')
    print(f"  QC: 0 blocking, {len(actionable)} warning(s) written to stderr")


if __name__ == '__main__':
    main()
