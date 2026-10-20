#!/usr/bin/env python3
"""
Calibration + internal validation (research-plan pipeline)

Implements references/advanced_analyses.md section 5 ("Calibration -- discrimination
is not the same as a usable model").

    python calibration.py config.json                  # full run
    python calibration.py config.json --check-missing  # preview only, no fitting

Why this script exists
----------------------
auc_cstat.py already reports discrimination (AUC / C-index). Discrimination only
says whether the model ranks patients correctly; it says nothing about whether the
absolute risks it produces are right. A model with C = 0.80 can systematically
double every patient's risk and still have C = 0.80. Clinical decisions are made on
absolute risk, so calibration is the other half and it is the half reviewers of
prediction-model papers reject papers over.

Design notes
------------
* study_purpose = "etiologic" hard-branches: only discrimination is reported and the
  first sheet states in plain words that this is NOT prediction-model performance and
  that calibration does not apply. Reporting a C-index for an etiologic model is fine;
  calling it model performance is not.
* Internal validation is optimism-corrected bootstrap (Harrell), never split-sample.
  Split-sample at typical clinical sample sizes is unstable and throws away data.
* Survival Brier scores are IPCW-weighted (Graf 1999); the unweighted version is
  biased whenever censoring is informative about follow-up length.
* Calibration slope / intercept for survival follow Crowson 2016 (Poisson framework
  with the model's own cumulative hazard as offset) plus the van Houwelingen
  calibration slope (Cox model refitted on the linear predictor alone).
* When the config supplies `predicted_column`, the predictions are taken as given
  (external / pre-specified model). Nothing is refitted, the apparent slope is then a
  genuine calibration slope, and bootstrap optimism correction is not applicable.
  Without it, the model is fitted from `sections` on the same data, the apparent
  slope is 1.0 by construction, and the honest number is the optimism-corrected one.
"""

import sys
import os
import json
import math
import operator

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


_QC_MESSAGES = []


def qc(kind, message):
    line = f"[QC_{kind}] {message}"
    _QC_MESSAGES.append(line)
    print(line, file=sys.stderr)


def qc_warning(message):
    qc('WARNING', message)


def qc_blocking(message):
    qc('BLOCKING', message)


def n_blocking():
    return sum(1 for m in _QC_MESSAGES if m.startswith('[QC_BLOCKING]'))


# ═══════════════════════════════════════════════════════════════════════════════
# Configuration
# ═══════════════════════════════════════════════════════════════════════════════

DEFAULTS = {
    # 500, not the protocol minimum of 200: percentile CIs and the optimism
    # correction are both still visibly noisy at 200, and the cost is only time.
    # 200 remains a legal fast mode and is labelled as one in the output.
    'bootstrap_n': 500,
    'n_bins': 10,
    'eval_times': None,
    'seed': 42,
    'study_purpose': 'prediction',
    'min_bootstrap_success_frac': 0.5,
}


def load_config(path):
    with open(path, 'r', encoding='utf-8') as f:
        config = json.load(f)

    config.setdefault('mode', 'calibration')
    if config['mode'] != 'calibration':
        sys.exit(f"[FATAL] Unknown mode: {config['mode']!r} (expected 'calibration')")

    config.setdefault('filters', [])
    config.setdefault('derived_variables', [])
    config.setdefault('title', 'Calibration and internal validation')
    config.setdefault('subtitle', '')
    config.setdefault('abbreviations', '')
    config.setdefault('sections', [])
    config.setdefault('predicted_column', None)

    # Track which methodological choices came from the config and which are this
    # script's defaults -- both go into the Method sheet because both are decisions
    # that belong in decision_log.md.
    provenance = {}
    for key, default in DEFAULTS.items():
        if key in ('min_bootstrap_success_frac',):
            continue
        provenance[key] = 'config' if key in config else 'default'
        config.setdefault(key, default)
    config.setdefault('min_bootstrap_success_frac', DEFAULTS['min_bootstrap_success_frac'])
    config['_provenance'] = provenance

    at = config.get('analysis_type')
    if at not in ('cox', 'logistic'):
        sys.exit("[FATAL] analysis_type must be 'cox' or 'logistic'")

    if at == 'cox':
        cs = config.get('cox_settings') or {}
        if not cs.get('time_column') or not cs.get('event_column'):
            sys.exit("[FATAL] cox_settings.time_column and cox_settings.event_column are required")
        if not config.get('eval_times'):
            sys.exit("[FATAL] eval_times is required for analysis_type='cox'")
    else:
        ls = config.get('logistic_settings') or {}
        if not ls.get('outcome_column'):
            sys.exit("[FATAL] logistic_settings.outcome_column is required")

    if config['study_purpose'] not in ('prediction', 'etiologic'):
        sys.exit("[FATAL] study_purpose must be 'prediction' or 'etiologic'")

    base = os.path.splitext(os.path.basename(path))[0]
    config.setdefault('output_file', f"{base}_calibration.xlsx")
    config.setdefault('plot_file', f"{base}_calibration.png")
    return config


# ═══════════════════════════════════════════════════════════════════════════════
# Data loading (same semantics as the other pipeline sub-skills)
# ═══════════════════════════════════════════════════════════════════════════════

EQ_OPS = {'==': operator.eq, '!=': operator.ne}
NUMERIC_OPS = {'>': operator.gt, '>=': operator.ge, '<': operator.lt, '<=': operator.le}


def load_data(config):
    path = config['data_file']
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xls'):
        df = pd.read_excel(path)
    elif ext == '.tsv':
        df = pd.read_csv(path, sep='\t')
    elif ext == '.sas7bdat':
        df = pd.read_sas(path)
    elif ext == '.dta':
        df = pd.read_stata(path)
    else:
        df = pd.read_csv(path)

    for f in config['filters']:
        col, op, val = f['column'], f['operator'], f['value']
        if op in EQ_OPS:
            df = df[EQ_OPS[op](df[col], val)]
        elif op in NUMERIC_OPS:
            df = df[NUMERIC_OPS[op](pd.to_numeric(df[col], errors='coerce'), val)]
        elif op == 'in':
            df = df[df[col].isin(val)]
        elif op == 'not_in':
            df = df[~df[col].isin(val)]

    for dv in config['derived_variables']:
        if dv['type'] == 'coalesce':
            result = pd.to_numeric(df[dv['columns'][0]], errors='coerce')
            for c in dv['columns'][1:]:
                result = result.combine_first(pd.to_numeric(df[c], errors='coerce'))
            df[dv['name']] = result
        elif dv['type'] == 'expression':
            ns = {dep: pd.to_numeric(df[dep], errors='coerce')
                  for dep in dv.get('dependencies', [])}
            ns['np'] = np
            df[dv['name']] = eval(dv['expression'], {"__builtins__": {}}, ns)

    return df.reset_index(drop=True)


def model_variables(config):
    out = []
    for section in config['sections']:
        for v in section['variables']:
            out.append({
                'column': v['column'],
                'type': v.get('type', 'continuous'),
                'label': v.get('label', v['column']),
                'scale': v.get('scale'),
                'reference_level': v.get('reference_level'),
                'section': section.get('header', ''),
            })
    return out


def build_design(df, model_vars):
    """Design matrix with sanitized column names (lifelines formulas choke on
    spaces / punctuation in real-world column names)."""
    cols, terms = {}, []
    for v in model_vars:
        col, vtype = v['column'], v['type']
        if col not in df.columns:
            sys.exit(f"[FATAL] Model variable column '{col}' not found in data")
        if vtype == 'categorical':
            s = df[col].astype(str)
            s = s.where(df[col].notna() & (s.str.strip() != ''), np.nan)
            levels = sorted(x for x in s.dropna().unique())
            ref = v.get('reference_level')
            ref = str(ref) if ref is not None else (levels[0] if levels else None)
            for lev in levels:
                if lev == ref:
                    continue
                name = f"x{len(terms)}"
                cols[name] = np.where(s.isna(), np.nan, (s == lev).astype(float))
                terms.append({'safe': name, 'label': f"{v['label']} = {lev}",
                              'column': col, 'type': 'categorical(dummy)',
                              'section': v['section']})
        else:
            s = pd.to_numeric(df[col], errors='coerce')
            if v.get('scale'):
                s = s / float(v['scale'])
            name = f"x{len(terms)}"
            cols[name] = s.values
            terms.append({'safe': name, 'label': v['label'], 'column': col,
                          'type': vtype, 'section': v['section']})
    X = pd.DataFrame(cols, index=df.index)
    return X, terms


# ═══════════════════════════════════════════════════════════════════════════════
# Excel style constants (shared with the other pipeline sub-skills)
# ═══════════════════════════════════════════════════════════════════════════════

DARK_BLUE = "1F3864"
HEADER_BG = "2E5FA3"
ALT_ROW = "F7FAFD"
SIG_COLOR = "C00000"
BORDER_COLOR = "BFBFBF"
FOOTNOTE_COLOR = "595959"
FLAG_BG = "FDEBD0"

THIN_BORDER = Border(
    left=Side(style='thin', color=BORDER_COLOR),
    right=Side(style='thin', color=BORDER_COLOR),
    top=Side(style='thin', color=BORDER_COLOR),
    bottom=Side(style='thin', color=BORDER_COLOR),
)

TITLE_FONT = Font(name='Arial', size=12, bold=True, color="FFFFFF")
SUBTITLE_FONT = Font(name='Arial', size=10, italic=True, color="FFFFFF")
HEADER_FONT = Font(name='Arial', size=10, bold=True, color="FFFFFF")
DATA_FONT = Font(name='Arial', size=10)
BOLD_FONT = Font(name='Arial', size=10, bold=True)
SIG_FONT = Font(name='Arial', size=10, bold=True, color=SIG_COLOR)
FOOTNOTE_FONT = Font(name='Arial', size=8, italic=True, color=FOOTNOTE_COLOR)

TITLE_FILL = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type='solid')
HEADER_FILL = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type='solid')
ALT_FILL = PatternFill(start_color=ALT_ROW, end_color=ALT_ROW, fill_type='solid')
FLAG_FILL = PatternFill(start_color=FLAG_BG, end_color=FLAG_BG, fill_type='solid')

CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', indent=1, wrap_text=True)


def sheet_header(ws, title, subtitle, headers, widths):
    ncol = len(headers)
    for c, w in zip(range(1, ncol + 1), widths):
        ws.column_dimensions[get_column_letter(c)].width = w
    r = 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    cell = ws.cell(row=r, column=1, value=title)
    cell.font = TITLE_FONT
    cell.alignment = CENTER
    for c in range(1, ncol + 1):
        ws.cell(row=r, column=c).fill = TITLE_FILL
    r += 1
    if subtitle:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        cell = ws.cell(row=r, column=1, value=subtitle)
        cell.font = SUBTITLE_FONT
        cell.alignment = CENTER
        for c in range(1, ncol + 1):
            ws.cell(row=r, column=c).fill = TITLE_FILL
        r += 1
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER
    r += 1
    ws.freeze_panes = ws.cell(row=r, column=1)
    return r


def write_row(ws, r, values, alt=False, flag=False, bold_cols=(), sig_cols=()):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = SIG_FONT if c in sig_cols else DATA_FONT
        if c in bold_cols:
            cell.font = BOLD_FONT
        cell.border = THIN_BORDER
        cell.alignment = LEFT if c == 1 else CENTER
        if flag:
            cell.fill = FLAG_FILL
        elif alt:
            cell.fill = ALT_FILL
    return r + 1


def write_para(ws, r, ncol, text, bold=False):
    """A full-width paragraph row (used by the Estimand sheet)."""
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    cell = ws.cell(row=r, column=1, value=text)
    cell.font = BOLD_FONT if bold else DATA_FONT
    cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True, indent=1)
    ws.row_dimensions[r].height = max(15, 13 * (1 + len(text) // 110 + text.count('\n')))
    return r + 1


def write_footnote(ws, r, ncol, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    cell = ws.cell(row=r, column=1, value=text)
    cell.font = FOOTNOTE_FONT
    cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    ws.row_dimensions[r].height = max(14, 11 * (1 + len(text) // 130 + text.count('\n')))
    return r + 1


def fmt_num(x, nd=3):
    if x is None or (isinstance(x, float) and (math.isnan(x) or math.isinf(x))):
        return "N/A"
    return f"{x:.{nd}f}"


def fmt_ci(est, lo, hi, nd=3):
    if est is None or (isinstance(est, float) and math.isnan(est)):
        return "N/A"
    if lo is None or hi is None or any(isinstance(v, float) and math.isnan(v) for v in (lo, hi)):
        return f"{est:.{nd}f}"
    return f"{est:.{nd}f} ({lo:.{nd}f}–{hi:.{nd}f})"


# ═══════════════════════════════════════════════════════════════════════════════
# Statistical helpers
# ═══════════════════════════════════════════════════════════════════════════════

def concordance_binary(y, p):
    """AUC == concordance for a binary outcome; rank-based, ties get 0.5."""
    y = np.asarray(y, dtype=float)
    p = np.asarray(p, dtype=float)
    pos, neg = p[y == 1], p[y == 0]
    if len(pos) == 0 or len(neg) == 0:
        return np.nan
    order = np.argsort(p)
    ranks = np.empty(len(p), dtype=float)
    sp = p[order]
    i = 0
    r = np.empty(len(p), dtype=float)
    while i < len(sp):
        j = i
        while j + 1 < len(sp) and sp[j + 1] == sp[i]:
            j += 1
        r[i:j + 1] = (i + j) / 2.0 + 1.0
        i = j + 1
    ranks[order] = r
    n1, n0 = len(pos), len(neg)
    return float((ranks[y == 1].sum() - n1 * (n1 + 1) / 2.0) / (n1 * n0))


def concordance_survival(time, event, risk_score):
    """Harrell's C for a risk score (higher = worse prognosis)."""
    from lifelines.utils import concordance_index
    return float(concordance_index(time, -np.asarray(risk_score, dtype=float), event))


def step_lookup(x_grid, y_vals, queries, strictly_before=False):
    """Right-continuous step function lookup (KM / Nelson-Aalen style)."""
    x_grid = np.asarray(x_grid, dtype=float)
    y_vals = np.asarray(y_vals, dtype=float)
    side = 'left' if strictly_before else 'right'
    idx = np.searchsorted(x_grid, np.asarray(queries, dtype=float), side=side) - 1
    out = np.where(idx < 0, 1.0, y_vals[np.clip(idx, 0, len(y_vals) - 1)])
    return out


def wilson_ci(k, n, z=1.96):
    if n == 0:
        return (np.nan, np.nan)
    p = k / n
    d = 1 + z ** 2 / n
    c = p + z ** 2 / (2 * n)
    h = z * math.sqrt(p * (1 - p) / n + z ** 2 / (4 * n ** 2))
    return ((c - h) / d, (c + h) / d)


# ═══════════════════════════════════════════════════════════════════════════════
# Model fitting
# ═══════════════════════════════════════════════════════════════════════════════

def fit_logistic(X, y):
    """statsmodels returns a fitted object even when the optimizer gave up (it only
    warns). Silently accepting those inside the bootstrap would count divergent fits
    as successes and shrink the optimism estimate, so non-convergence is raised."""
    import statsmodels.api as sm
    Xc = sm.add_constant(X.values.astype(float), has_constant='add')
    model = sm.Logit(y.astype(float), Xc).fit(disp=0, method='newton', maxiter=200)
    if not model.mle_retvals.get('converged', True):
        raise RuntimeError('logistic regression did not converge')
    if not np.all(np.isfinite(np.asarray(model.params, dtype=float))):
        raise RuntimeError('logistic regression produced non-finite coefficients')
    return model


def logistic_lp(model, X):
    import statsmodels.api as sm
    Xc = sm.add_constant(X.values.astype(float), has_constant='add')
    return np.asarray(Xc @ model.params, dtype=float)


def fit_cox(d, terms, time_col, event_col):
    from lifelines import CoxPHFitter
    cph = CoxPHFitter()
    cph.fit(d, duration_col=time_col, event_col=event_col,
            formula=' + '.join(t['safe'] for t in terms))
    return cph


def cox_lp(cph, Xdf):
    """Linear predictor on the model's own centering (constant offsets cancel in
    every calibration statistic we compute from it)."""
    beta = cph.params_
    vals = np.zeros(len(Xdf), dtype=float)
    for name in beta.index:
        vals += float(beta[name]) * Xdf[name].values.astype(float)
    return vals


def cox_cumhaz_at(cph, Xdf, target_times):
    """Subject-specific cumulative hazard H_i(t_i), evaluated on the fitted
    baseline. Used for the Poisson (Crowson) calibration framework."""
    target_times = np.asarray(target_times, dtype=float)
    uniq = np.unique(np.concatenate([[0.0], target_times]))
    if len(uniq) > 1500:
        uniq = np.unique(np.concatenate(
            [[0.0], np.quantile(target_times, np.linspace(0, 1, 1500))]))
    H = cph.predict_cumulative_hazard(Xdf, times=uniq)
    grid = np.asarray(H.index, dtype=float)
    Hv = H.values
    idx = np.clip(np.searchsorted(grid, target_times, side='right') - 1, 0, len(grid) - 1)
    return Hv[idx, np.arange(Hv.shape[1])]


def cox_risk_at(cph, Xdf, t):
    """Predicted absolute risk 1 - S(t|x)."""
    S = cph.predict_survival_function(Xdf, times=[float(t)])
    return 1.0 - np.asarray(S.values[0], dtype=float)


# ═══════════════════════════════════════════════════════════════════════════════
# Calibration statistics
# ═══════════════════════════════════════════════════════════════════════════════

def calib_slope_intercept_logistic(y, lp):
    """Slope: logit(y) ~ a + b*lp, b is the slope. Intercept
    (calibration-in-the-large): logit(y) ~ a with lp as a fixed offset."""
    import statsmodels.api as sm
    y = np.asarray(y, dtype=float)
    lp = np.asarray(lp, dtype=float)
    out = {'slope': np.nan, 'slope_se': np.nan, 'intercept': np.nan, 'intercept_se': np.nan}
    try:
        Xs = sm.add_constant(lp.reshape(-1, 1), has_constant='add')
        m = sm.Logit(y, Xs).fit(disp=0, method='newton', maxiter=200)
        out['slope'] = float(m.params[1])
        out['slope_se'] = float(m.bse[1])
    except Exception as e:
        print(f"  [calibration slope] {type(e).__name__}: {e}", file=sys.stderr)
    try:
        ones = np.ones((len(y), 1))
        m2 = sm.Logit(y, ones, offset=lp).fit(disp=0, method='newton', maxiter=200)
        out['intercept'] = float(m2.params[0])
        out['intercept_se'] = float(m2.bse[0])
    except Exception as e:
        print(f"  [calibration intercept] {type(e).__name__}: {e}", file=sys.stderr)
    return out


def calib_slope_cox(d, lp, time_col, event_col):
    """van Houwelingen calibration slope: Cox refitted on the linear predictor
    alone. 1.0 = the prognostic index is on the right scale; < 1 = predictions
    are too extreme (overfitting)."""
    from lifelines import CoxPHFitter
    dd = pd.DataFrame({'lp': np.asarray(lp, dtype=float),
                       time_col: d[time_col].values,
                       event_col: d[event_col].values})
    try:
        c = CoxPHFitter()
        c.fit(dd, duration_col=time_col, event_col=event_col, formula='lp')
        return float(c.params_['lp']), float(c.standard_errors_['lp'])
    except Exception as e:
        print(f"  [cox calibration slope] {type(e).__name__}: {e}", file=sys.stderr)
        return np.nan, np.nan


def poisson_calibration(observed, expected):
    """Crowson 2016: intercept = log(O/E) from a Poisson model with log(expected)
    as offset; slope from the same model with log(expected) as a covariate."""
    import statsmodels.api as sm
    observed = np.asarray(observed, dtype=float)
    expected = np.asarray(expected, dtype=float)
    ok = np.isfinite(expected) & (expected > 0)
    observed, expected = observed[ok], expected[ok]
    out = {'intercept': np.nan, 'slope': np.nan,
           'observed': float(observed.sum()), 'expected': float(expected.sum())}
    if len(observed) == 0 or observed.sum() == 0:
        return out
    out['oe_ratio'] = out['observed'] / out['expected'] if out['expected'] > 0 else np.nan
    try:
        m = sm.GLM(observed, np.ones((len(observed), 1)),
                   family=sm.families.Poisson(), offset=np.log(expected)).fit()
        out['intercept'] = float(m.params[0])
    except Exception:
        out['intercept'] = float(np.log(out['oe_ratio'])) if out.get('oe_ratio', 0) > 0 else np.nan
    try:
        Xp = sm.add_constant(np.log(expected).reshape(-1, 1), has_constant='add')
        m2 = sm.GLM(observed, Xp, family=sm.families.Poisson()).fit()
        out['slope'] = float(m2.params[1])
    except Exception:
        pass
    return out


def decile_table_binary(y, p, n_bins):
    y = np.asarray(y, dtype=float)
    p = np.asarray(p, dtype=float)
    ranks = pd.Series(p).rank(method='first')
    bins = pd.qcut(ranks, min(n_bins, len(np.unique(ranks))), labels=False, duplicates='drop')
    rows = []
    for b in sorted(pd.unique(bins[~pd.isna(bins)])):
        m = (bins == b).values
        k = int(y[m].sum())
        n = int(m.sum())
        lo, hi = wilson_ci(k, n)
        rows.append({'bin': int(b) + 1, 'n': n, 'events': k,
                     'pred_mean': float(p[m].mean()),
                     'pred_min': float(p[m].min()), 'pred_max': float(p[m].max()),
                     'obs': k / n if n else np.nan, 'obs_lo': lo, 'obs_hi': hi})
    return rows


def decile_table_survival(time, event, risk, t, n_bins):
    from lifelines import KaplanMeierFitter
    time = np.asarray(time, dtype=float)
    event = np.asarray(event, dtype=float)
    risk = np.asarray(risk, dtype=float)
    ranks = pd.Series(risk).rank(method='first')
    bins = pd.qcut(ranks, min(n_bins, len(np.unique(ranks))), labels=False, duplicates='drop')
    rows = []
    for b in sorted(pd.unique(bins[~pd.isna(bins)])):
        m = (bins == b).values
        kmf = KaplanMeierFitter()
        obs = obs_lo = obs_hi = np.nan
        try:
            kmf.fit(time[m], event_observed=event[m])
            sf = kmf.survival_function_
            obs = 1.0 - float(step_lookup(sf.index.values, sf.iloc[:, 0].values, [t])[0])
            ci = kmf.confidence_interval_
            hi_s = float(step_lookup(ci.index.values, ci.iloc[:, 0].values, [t])[0])
            lo_s = float(step_lookup(ci.index.values, ci.iloc[:, 1].values, [t])[0])
            obs_lo, obs_hi = 1.0 - lo_s, 1.0 - hi_s
        except Exception:
            pass
        rows.append({'bin': int(b) + 1, 'n': int(m.sum()),
                     'events': int(event[m & (time <= t)].sum()),
                     'pred_mean': float(risk[m].mean()),
                     'pred_min': float(risk[m].min()), 'pred_max': float(risk[m].max()),
                     'obs': obs, 'obs_lo': obs_lo, 'obs_hi': obs_hi})
    return rows


def brier_binary(y, p):
    y = np.asarray(y, dtype=float)
    p = np.asarray(p, dtype=float)
    b = float(np.mean((y - p) ** 2))
    pbar = float(y.mean())
    b_null = float(np.mean((y - pbar) ** 2))
    return {'brier': b, 'brier_null': b_null,
            'ipa': 1.0 - b / b_null if b_null > 0 else np.nan}


def brier_ipcw(time, event, pred_risk, t):
    """Graf 1999 IPCW Brier score at time t.

    Censoring distribution G is the reverse KM (KM of the censoring indicator).
    Subjects censored before t contribute nothing; the rest are up-weighted by
    1/G so the estimator stays unbiased under random censoring.
    """
    from lifelines import KaplanMeierFitter
    time = np.asarray(time, dtype=float)
    event = np.asarray(event, dtype=float)
    S_pred = 1.0 - np.asarray(pred_risk, dtype=float)

    kmc = KaplanMeierFitter()
    kmc.fit(time, event_observed=(1 - event))
    gx = kmc.survival_function_.index.values
    gy = kmc.survival_function_.iloc[:, 0].values

    G_ti = step_lookup(gx, gy, time, strictly_before=True)
    G_t = float(step_lookup(gx, gy, [t])[0])

    w = np.zeros(len(time))
    contrib = np.zeros(len(time))
    case1 = (time <= t) & (event == 1)
    case2 = time > t
    with np.errstate(divide='ignore', invalid='ignore'):
        w[case1] = np.where(G_ti[case1] > 0, 1.0 / G_ti[case1], 0.0)
        w[case2] = 1.0 / G_t if G_t > 0 else 0.0
    contrib[case1] = (0.0 - S_pred[case1]) ** 2
    contrib[case2] = (1.0 - S_pred[case2]) ** 2
    n = len(time)
    b = float(np.sum(w * contrib) / n)

    kmm = KaplanMeierFitter()
    kmm.fit(time, event_observed=event)
    sx = kmm.survival_function_.index.values
    sy = kmm.survival_function_.iloc[:, 0].values
    S_marg = float(step_lookup(sx, sy, [t])[0])
    contrib0 = np.zeros(len(time))
    contrib0[case1] = (0.0 - S_marg) ** 2
    contrib0[case2] = (1.0 - S_marg) ** 2
    b_null = float(np.sum(w * contrib0) / n)

    return {'brier': b, 'brier_null': b_null,
            'ipa': 1.0 - b / b_null if b_null > 0 else np.nan,
            'n_effective': int((w > 0).sum()),
            'n_censored_before_t': int(((time <= t) & (event == 0)).sum()),
            'G_at_t': G_t}


# ═══════════════════════════════════════════════════════════════════════════════
# Optimism-corrected bootstrap (Harrell). Never split-sample.
# ═══════════════════════════════════════════════════════════════════════════════

def bootstrap_optimism_logistic(X, y, n_boot, seed):
    rng = np.random.default_rng(seed)
    n = len(y)
    app_model = fit_logistic(X, y)
    app_lp = logistic_lp(app_model, X)
    app_p = 1.0 / (1.0 + np.exp(-app_lp))
    app_c = concordance_binary(y, app_p)
    app_stats = calib_slope_intercept_logistic(y, app_lp)

    opt_c, opt_slope, n_ok, n_fail = [], [], 0, 0
    for _ in range(n_boot):
        idx = rng.integers(0, n, size=n)
        Xb, yb = X.iloc[idx].reset_index(drop=True), np.asarray(y)[idx]
        if len(np.unique(yb)) < 2:
            n_fail += 1
            continue
        try:
            mb = fit_logistic(Xb, yb)
            lp_b = logistic_lp(mb, Xb)
            lp_o = logistic_lp(mb, X)
            c_b = concordance_binary(yb, lp_b)
            c_o = concordance_binary(y, lp_o)
            s_b = calib_slope_intercept_logistic(yb, lp_b)['slope']
            s_o = calib_slope_intercept_logistic(y, lp_o)['slope']
            if not (np.isfinite(c_b) and np.isfinite(c_o) and
                    np.isfinite(s_b) and np.isfinite(s_o)):
                n_fail += 1
                continue
            opt_c.append(c_b - c_o)
            opt_slope.append(s_b - s_o)
            n_ok += 1
        except Exception:
            n_fail += 1
    return _assemble_optimism(app_c, app_stats['slope'], opt_c, opt_slope,
                              n_ok, n_fail, n_boot)


def bootstrap_optimism_cox(d, X, terms, time_col, event_col, n_boot, seed):
    rng = np.random.default_rng(seed)
    n = len(d)
    cph = fit_cox(d, terms, time_col, event_col)
    app_lp = cox_lp(cph, X)
    app_c = concordance_survival(d[time_col].values, d[event_col].values, app_lp)
    app_slope, _ = calib_slope_cox(d, app_lp, time_col, event_col)

    opt_c, opt_slope, n_ok, n_fail = [], [], 0, 0
    for _ in range(n_boot):
        idx = rng.integers(0, n, size=n)
        db = d.iloc[idx].reset_index(drop=True)
        Xb = X.iloc[idx].reset_index(drop=True)
        if db[event_col].sum() < 2:
            n_fail += 1
            continue
        try:
            cb = fit_cox(db, terms, time_col, event_col)
            lp_b = cox_lp(cb, Xb)
            lp_o = cox_lp(cb, X)
            c_b = concordance_survival(db[time_col].values, db[event_col].values, lp_b)
            c_o = concordance_survival(d[time_col].values, d[event_col].values, lp_o)
            s_b, _ = calib_slope_cox(db, lp_b, time_col, event_col)
            s_o, _ = calib_slope_cox(d, lp_o, time_col, event_col)
            if not all(np.isfinite(v) for v in (c_b, c_o, s_b, s_o)):
                n_fail += 1
                continue
            opt_c.append(c_b - c_o)
            opt_slope.append(s_b - s_o)
            n_ok += 1
        except Exception:
            n_fail += 1
    return _assemble_optimism(app_c, app_slope, opt_c, opt_slope, n_ok, n_fail, n_boot)


def _assemble_optimism(app_c, app_slope, opt_c, opt_slope, n_ok, n_fail, n_boot):
    res = {'apparent_c': app_c, 'apparent_slope': app_slope,
           'n_success': n_ok, 'n_failed': n_fail, 'n_requested': n_boot,
           'optimism_c': np.nan, 'optimism_slope': np.nan,
           'corrected_c': np.nan, 'corrected_slope': np.nan,
           'optimism_c_sd': np.nan, 'optimism_slope_sd': np.nan}
    if n_ok > 0:
        res['optimism_c'] = float(np.mean(opt_c))
        res['optimism_slope'] = float(np.mean(opt_slope))
        res['optimism_c_sd'] = float(np.std(opt_c, ddof=1)) if n_ok > 1 else np.nan
        res['optimism_slope_sd'] = float(np.std(opt_slope, ddof=1)) if n_ok > 1 else np.nan
        res['corrected_c'] = app_c - res['optimism_c']
        res['corrected_slope'] = app_slope - res['optimism_slope']
    return res


def bootstrap_ci_fixed_predictions(stat_fn, n, n_boot, seed):
    """Percentile CI for a statistic computed on FIXED predictions (no refitting)."""
    rng = np.random.default_rng(seed + 1)
    vals = []
    for _ in range(n_boot):
        idx = rng.integers(0, n, size=n)
        try:
            v = stat_fn(idx)
            if np.isfinite(v):
                vals.append(v)
        except Exception:
            continue
    if len(vals) < 10:
        return (np.nan, np.nan)
    return (float(np.percentile(vals, 2.5)), float(np.percentile(vals, 97.5)))


# ═══════════════════════════════════════════════════════════════════════════════
# Analysis driver
# ═══════════════════════════════════════════════════════════════════════════════

def run_analysis(config, df):
    at = config['analysis_type']
    purpose = config['study_purpose']
    n_bins = int(config['n_bins'])
    n_boot = int(config['bootstrap_n'])
    seed = int(config['seed'])
    pred_col = config.get('predicted_column')

    model_vars = model_variables(config)
    result = {'analysis_type': at, 'study_purpose': purpose,
              'external_predictions': bool(pred_col), 'terms': [],
              'eval_times': [], 'calibration': {}, 'discrimination': {},
              'brier': [], 'internal': None, 'n_total': len(df)}

    if at == 'cox':
        cs = config['cox_settings']
        time_col, event_col = cs['time_column'], cs['event_column']
        keep = [time_col, event_col]
    else:
        out_col = config['logistic_settings']['outcome_column']
        keep = [out_col]

    if pred_col:
        if pred_col not in df.columns:
            sys.exit(f"[FATAL] predicted_column '{pred_col}' not found in data")
        # A single column holds the predicted risk at ONE horizon. Calibrating it
        # against several eval_times would reuse the same 1-year risk as if it were
        # also the 3-year risk — every statistic after the first horizon would be
        # wrong, and nothing in the output would say so.
        eval_times = config.get('eval_times') or []
        if len(eval_times) > 1:
            qc_blocking(
                f"predicted_column '{pred_col}' holds the predicted risk at a single "
                f"horizon, but {len(eval_times)} eval_times were requested "
                f"({eval_times}). The same column cannot be the risk at more than one "
                f"time point. Either give one eval_time matching the horizon these "
                f"predictions were made for, or drop predicted_column and let the "
                f"model be fitted here.")
            sys.exit(3)
        X = pd.DataFrame(index=df.index)
        terms = []
        work = df[keep + [pred_col]].copy()
    else:
        if not model_vars:
            sys.exit("[FATAL] sections must define model variables when predicted_column is absent")
        X, terms = build_design(df, model_vars)
        work = pd.concat([df[keep], X], axis=1)

    for c in keep:
        work[c] = pd.to_numeric(work[c], errors='coerce')
    before = len(work)
    work = work.dropna().reset_index(drop=True)
    n_dropped = before - len(work)
    result['n_analyzed'] = len(work)
    result['n_dropped'] = n_dropped
    result['terms'] = terms
    if n_dropped:
        print(f"  Listwise deletion removed {n_dropped} row(s) with missing model data")

    if len(work) < 20:
        qc_blocking(f"Only {len(work)} complete rows after listwise deletion; "
                    f"calibration cannot be estimated")
        return result

    X = work[[t['safe'] for t in terms]] if terms else pd.DataFrame(index=work.index)
    n_params = max(len(terms), 1)

    # ── Survival branch ──────────────────────────────────────────────────────
    if at == 'cox':
        time = work[time_col].values.astype(float)
        event = work[event_col].values.astype(float)
        n_events = int(event.sum())
        result['n_events'] = n_events
        result['max_followup'] = float(time.max())
        result['last_event_time'] = float(time[event == 1].max()) if n_events else np.nan

        eval_times = [float(t) for t in config['eval_times']]
        result['eval_times'] = eval_times
        for t in eval_times:
            if t > result['max_followup']:
                qc_blocking(f"eval_time {t:g} exceeds maximum observed follow-up "
                            f"({result['max_followup']:g}); no data support risk "
                            f"prediction at that horizon")
            elif n_events and t > result['last_event_time']:
                qc_warning(f"eval_time {t:g} is beyond the last observed event "
                           f"({result['last_event_time']:g}); the KM estimate is flat there")
        if n_blocking():
            return result

        if not pred_col:
            epp = n_events / n_params
            result['epp'] = epp
            if epp < 10:
                qc_warning(f"Events per parameter = {epp:.1f} (< 10): {n_events} events "
                           f"for {n_params} model parameter(s); the model is at high risk "
                           f"of overfitting")

        if pred_col:
            # External predictions: a single column of predicted RISK at the first
            # eval time. Discrimination uses it directly as the risk score.
            risk_ext = pd.to_numeric(work[pred_col], errors='coerce').values.astype(float)
            lp = np.log(np.clip(-np.log(np.clip(1 - risk_ext, 1e-12, 1)), 1e-12, None))
            cph = None
        else:
            try:
                cph = fit_cox(work, terms, time_col, event_col)
            except Exception as e:
                qc_blocking(f"The Cox model did not converge on the full dataset "
                            f"({type(e).__name__}: {e}); no calibration statistic can be "
                            f"computed from a model that could not be fitted")
                return result
            lp = cox_lp(cph, X)

        c_app = concordance_survival(time, event, lp)
        lo, hi = bootstrap_ci_fixed_predictions(
            lambda idx: concordance_survival(time[idx], event[idx], lp[idx]),
            len(work), n_boot, seed)
        result['discrimination'] = {'metric': 'Harrell C-index', 'apparent': c_app,
                                    'ci_lo': lo, 'ci_hi': hi, 'n': len(work),
                                    'n_events': n_events}

        if purpose == 'etiologic':
            return result

        slope, slope_se = calib_slope_cox(work, lp, time_col, event_col)
        result['calibration']['slope_apparent'] = slope
        result['calibration']['slope_se'] = slope_se
        result['calibration']['per_time'] = []

        for t in eval_times:
            if pred_col:
                risk = np.clip(risk_ext, 1e-12, 1 - 1e-12)
                expected = -np.log(1 - risk)  # cumulative hazard implied at t
                obs_ind = ((time <= t) & (event == 1)).astype(float)
                pois = poisson_calibration(obs_ind, expected)
            else:
                risk = cox_risk_at(cph, X, t)
                obs_ind = ((time <= t) & (event == 1)).astype(float)
                expected = cox_cumhaz_at(cph, X, np.minimum(time, t))
                pois = poisson_calibration(obs_ind, expected)
            deciles = decile_table_survival(time, event, risk, t, n_bins)
            result['calibration']['per_time'].append(
                {'t': t, 'intercept': pois['intercept'], 'poisson_slope': pois['slope'],
                 'observed': pois['observed'], 'expected': pois['expected'],
                 'oe_ratio': pois.get('oe_ratio', np.nan),
                 'deciles': deciles, 'pred_risk': risk})
            result['brier'].append({'t': t, **brier_ipcw(time, event, risk, t)})

        if pred_col:
            result['internal'] = None
        else:
            result['internal'] = bootstrap_optimism_cox(
                work, X, terms, time_col, event_col, n_boot, seed)

    # ── Binary branch ────────────────────────────────────────────────────────
    else:
        y = work[out_col].values.astype(float)
        n_events = int(y.sum())
        result['n_events'] = n_events
        if config['_provenance'].get('eval_times') == 'config' and config.get('eval_times'):
            qc_warning("eval_times is ignored for analysis_type='logistic' "
                       "(no time horizon in a binary outcome model)")
        if set(np.unique(y)) - {0.0, 1.0}:
            qc_blocking(f"Outcome '{out_col}' is not 0/1 coded")
            return result
        if n_events < 5 or n_events == len(y):
            qc_blocking(f"Outcome '{out_col}' has {n_events} events out of {len(y)}; "
                        f"calibration cannot be estimated")
            return result

        if not pred_col:
            epp = n_events / n_params
            result['epp'] = epp
            if epp < 10:
                qc_warning(f"Events per parameter = {epp:.1f} (< 10): {n_events} events "
                           f"for {n_params} model parameter(s); the model is at high risk "
                           f"of overfitting")

        if pred_col:
            p = np.clip(pd.to_numeric(work[pred_col], errors='coerce').values.astype(float),
                        1e-12, 1 - 1e-12)
            lp = np.log(p / (1 - p))
        else:
            try:
                model = fit_logistic(X, y)
            except Exception as e:
                qc_blocking(f"The logistic model did not converge on the full dataset "
                            f"({type(e).__name__}: {e}); this is usually complete or "
                            f"quasi-complete separation, i.e. some covariate pattern predicts "
                            f"the outcome perfectly. No calibration statistic can be computed "
                            f"from a model that could not be fitted")
                return result
            lp = logistic_lp(model, X)
            p = 1.0 / (1.0 + np.exp(-lp))

        auc = concordance_binary(y, p)
        lo, hi = bootstrap_ci_fixed_predictions(
            lambda idx: concordance_binary(y[idx], p[idx]), len(work), n_boot, seed)
        result['discrimination'] = {'metric': 'AUC (c-statistic)', 'apparent': auc,
                                    'ci_lo': lo, 'ci_hi': hi, 'n': len(work),
                                    'n_events': n_events}

        if purpose == 'etiologic':
            return result

        st = calib_slope_intercept_logistic(y, lp)
        result['calibration']['slope_apparent'] = st['slope']
        result['calibration']['slope_se'] = st['slope_se']
        result['calibration']['intercept'] = st['intercept']
        result['calibration']['intercept_se'] = st['intercept_se']
        result['calibration']['deciles'] = decile_table_binary(y, p, n_bins)
        result['calibration']['pred'] = p
        result['calibration']['y'] = y
        result['brier'].append({'t': None, **brier_binary(y, p)})
        result['internal'] = (None if pred_col
                              else bootstrap_optimism_logistic(X, y, n_boot, seed))

    # ── Reported slope + QC ──────────────────────────────────────────────────
    internal = result.get('internal')
    if internal and np.isfinite(internal.get('corrected_slope', np.nan)):
        result['calibration']['slope_reported'] = internal['corrected_slope']
        result['calibration']['slope_reported_kind'] = 'optimism-corrected'
    else:
        result['calibration']['slope_reported'] = result['calibration'].get('slope_apparent')
        result['calibration']['slope_reported_kind'] = (
            'apparent (external predictions; no refitting)' if pred_col
            else 'apparent (bootstrap correction unavailable)')

    if internal:
        need = max(1, int(math.ceil(config['min_bootstrap_success_frac'] * n_boot)))
        if internal['n_success'] < need:
            qc_blocking(f"Bootstrap internal validation did not converge: only "
                        f"{internal['n_success']}/{n_boot} resamples produced a usable "
                        f"fit (minimum {need}); the optimism correction is not trustworthy")

    sr = result['calibration'].get('slope_reported')
    if sr is not None and np.isfinite(sr) and sr < 0.8:
        qc_warning(f"Calibration slope = {sr:.3f} (< 0.8, {result['calibration']['slope_reported_kind']}): "
                   f"predictions are too extreme, i.e. the model is overfitted; "
                   f"shrinkage or a simpler model is indicated")

    return result


# ═══════════════════════════════════════════════════════════════════════════════
# Excel output
# ═══════════════════════════════════════════════════════════════════════════════

ESTIMAND_PREDICTION = (
    "This output evaluates a PREDICTION MODEL. The quantity being estimated is the "
    "absolute risk of the outcome for an individual patient with a given covariate "
    "pattern -- not a causal effect. No ATE, ATT or ATO is estimated here; the "
    "coefficients are not interpreted as effects of anything, and no causal claim "
    "follows from them."
)

ESTIMAND_ETIOLOGIC = (
    "This model was fitted for an ETIOLOGIC question (is X associated with Y after "
    "adjustment), not to predict risk. This is NOT prediction-model performance, and "
    "calibration does not apply. "
    "(這不是預測模型效能，calibration 不適用。) "
    "The c-statistic below is reported only as a descriptive summary of how well the "
    "fitted model ranks patients; it must not be described as model performance, and "
    "a high value is not evidence that the model is clinically usable."
)


def write_excel(config, result):
    wb = Workbook()
    subtitle = config.get('subtitle', '')
    purpose = result['study_purpose']
    at = result['analysis_type']
    cal = result.get('calibration') or {}

    # ── Estimand (always the first sheet) ────────────────────────────────────
    ws = wb.active
    ws.title = 'Estimand'
    headers = ['What this analysis estimates']
    r = sheet_header(ws, 'Estimand and interpretation', subtitle, headers, [130])
    r = write_para(ws, r, 1, f"Study purpose (from config): {purpose}", bold=True)
    r = write_para(ws, r, 1,
                   ESTIMAND_ETIOLOGIC if purpose == 'etiologic' else ESTIMAND_PREDICTION)
    r += 1
    r = write_para(ws, r, 1, "What is and is not reported in this file", bold=True)
    if purpose == 'etiologic':
        r = write_para(ws, r, 1,
                       "Reported: discrimination only (Discrimination sheet). "
                       "NOT reported: calibration plot, calibration slope / intercept, "
                       "Brier score, optimism-corrected internal validation -- all four are "
                       "prediction-model quantities and are deliberately omitted, not missing "
                       "by accident. To obtain them, re-run with study_purpose = \"prediction\", "
                       "which is a different research question and requires its own "
                       "pre-specification.")
    else:
        r = write_para(ws, r, 1,
                       "Reported: calibration slope and intercept, decile-level predicted vs "
                       "observed risk, Brier score (IPCW-weighted for survival outcomes), and "
                       "optimism-corrected internal validation by bootstrap. Discrimination is "
                       "reported alongside because on its own it is not sufficient: a model can "
                       "rank patients perfectly and still put every absolute risk in the wrong "
                       "place.")
    r += 1
    r = write_para(ws, r, 1, "Model and sample", bold=True)
    rows = [
        ('Outcome model', 'Cox proportional hazards' if at == 'cox' else 'Logistic regression'),
        ('Predictions', 'supplied externally via predicted_column (nothing refitted)'
         if result['external_predictions'] else 'fitted from the covariates in this dataset'),
        ('Rows in analysis', f"{result.get('n_analyzed', 0)} "
                             f"(of {result['n_total']}; {result.get('n_dropped', 0)} dropped listwise)"),
        ('Events', str(result.get('n_events', 'N/A'))),
        ('Model parameters', str(len(result['terms']))),
        ('Events per parameter', fmt_num(result.get('epp'), 1) if result.get('epp') else 'N/A (external predictions)'),
    ]
    if at == 'cox':
        rows.append(('Evaluation time points', ', '.join(f"{t:g}" for t in result['eval_times']) or 'N/A'))
        rows.append(('Maximum observed follow-up', fmt_num(result.get('max_followup'), 1)))
    for k, v in rows:
        r = write_para(ws, r, 1, f"    {k}: {v}")
    r += 1
    r = write_footnote(ws, r, 1,
                       "Wording convention: a calibration slope below 1 means the model's "
                       "predictions are too extreme in this dataset; it does not mean the "
                       "predictors are unrelated to the outcome. A confidence interval that "
                       "includes the null means no statistically significant difference was "
                       "detected, not that no difference exists.")

    # ── Calibration ──────────────────────────────────────────────────────────
    if purpose == 'prediction':
        ws = wb.create_sheet('Calibration')
        headers = ['Item', 'Value', 'SE', 'Interpretation']
        r = sheet_header(ws, 'Calibration', subtitle, headers, [46, 20, 14, 62])
        i = 0
        rows = []
        sa = cal.get('slope_apparent')
        rows.append(('Calibration slope (apparent)', fmt_num(sa), fmt_num(cal.get('slope_se')),
                     'Slope of the outcome on the linear predictor. Exactly 1.0 by '
                     'construction when the model was fitted on these same rows -- read the '
                     'optimism-corrected value instead.' if not result['external_predictions']
                     else 'Slope of the outcome on the supplied linear predictor. 1.0 = correct '
                          'scale; < 1 = predictions too extreme; > 1 = too compressed.'))
        sr = cal.get('slope_reported')
        rows.append(('Calibration slope (REPORTED)', fmt_num(sr), '',
                     f"{cal.get('slope_reported_kind','')}. This is the number the QC "
                     f"threshold (< 0.8) is applied to."))
        if at == 'logistic':
            rows.append(('Calibration intercept (calibration-in-the-large)',
                         fmt_num(cal.get('intercept')), fmt_num(cal.get('intercept_se')),
                         '0 = no systematic over- or under-prediction. Negative = the model '
                         'over-predicts risk overall; positive = under-predicts.'))
        for k, v, se, interp in rows:
            flag = (k.startswith('Calibration slope (REPORTED)') and sr is not None
                    and np.isfinite(sr) and sr < 0.8)
            r = write_row(ws, r, [k, v, se, interp], alt=(i % 2 == 1), flag=flag,
                          bold_cols=(1,) if 'REPORTED' in k else ())
            i += 1

        if at == 'cox':
            r += 1
            r = write_row(ws, r, ['Per evaluation time point', '', '', ''], bold_cols=(1,))
            r = write_row(ws, r, ['Time', 'Calibration intercept log(O/E)',
                                  'O/E ratio', 'Observed events / expected events'],
                          bold_cols=(1, 2, 3, 4))
            for i, pt in enumerate(cal.get('per_time', [])):
                r = write_row(ws, r, [f"t = {pt['t']:g}", fmt_num(pt['intercept']),
                                      fmt_num(pt.get('oe_ratio')),
                                      f"{pt['observed']:.0f} / {pt['expected']:.1f}"],
                              alt=(i % 2 == 1))

        r += 1
        r = write_row(ws, r, ['Predicted vs observed risk by bin', '', '', ''], bold_cols=(1,))
        bin_headers = ['Bin (time)', 'n', 'Events', 'Mean predicted risk',
                       'Observed risk (95% CI)', 'Predicted range', 'Difference (obs - pred)']
        for c, h in enumerate(bin_headers, 1):
            cell = ws.cell(row=r, column=c, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = CENTER
            cell.border = THIN_BORDER
        for c, w in zip(range(1, len(bin_headers) + 1), [22, 10, 10, 20, 26, 22, 22]):
            ws.column_dimensions[get_column_letter(c)].width = max(
                ws.column_dimensions[get_column_letter(c)].width or 0, w)
        r += 1
        if at == 'cox':
            groups = [(f"t = {pt['t']:g}", pt['deciles']) for pt in cal.get('per_time', [])]
        else:
            groups = [('', cal.get('deciles', []))]
        i = 0
        for tag, deciles in groups:
            for d in deciles:
                diff = (d['obs'] - d['pred_mean']) if np.isfinite(d['obs']) else np.nan
                r = write_row(ws, r, [
                    f"{d['bin']}" + (f" ({tag})" if tag else ''), d['n'], d['events'],
                    fmt_num(d['pred_mean'], 4),
                    (f"{d['obs']:.4f} ({d['obs_lo']:.4f}–{d['obs_hi']:.4f})"
                     if np.isfinite(d['obs']) and np.isfinite(d['obs_lo']) else fmt_num(d['obs'], 4)),
                    f"{d['pred_min']:.4f}–{d['pred_max']:.4f}",
                    fmt_num(diff, 4)], alt=(i % 2 == 1))
                i += 1
        r += 1
        r = write_footnote(ws, r, len(bin_headers),
                           "Bins are equal-sized groups of the predicted risk (deciles by "
                           "default). For a survival outcome the observed risk in each bin is "
                           "1 minus the Kaplan-Meier estimate at the evaluation time, so "
                           "censored patients are handled correctly; a simple event proportion "
                           "would under-count. Perfect calibration puts every bin on the "
                           "diagonal. Systematic deviation in one direction matters more than "
                           "scatter: scatter in the extreme bins is expected because those bins "
                           "hold the fewest events.")

        # ── Brier ────────────────────────────────────────────────────────────
        ws = wb.create_sheet('Brier')
        headers = ['Time', 'Brier score', 'Brier (null model)', 'IPA / scaled Brier',
                   'n contributing', 'Censored before t', 'G(t) censoring survival']
        r = sheet_header(ws, 'Brier score (overall accuracy)', subtitle, headers,
                         [16, 16, 20, 22, 16, 20, 24])
        for i, b in enumerate(result['brier']):
            r = write_row(ws, r, [
                f"t = {b['t']:g}" if b['t'] is not None else 'binary outcome',
                fmt_num(b['brier'], 4), fmt_num(b['brier_null'], 4), fmt_num(b['ipa'], 4),
                b.get('n_effective', result.get('n_analyzed')),
                b.get('n_censored_before_t', 'n/a'),
                fmt_num(b.get('G_at_t'), 4) if b.get('G_at_t') is not None else 'n/a'],
                alt=(i % 2 == 1))
        r += 1
        r = write_footnote(ws, r, len(headers),
                           "Brier score = mean squared difference between predicted risk and "
                           "what actually happened; lower is better, and it penalises both poor "
                           "discrimination and poor calibration, which is why it is reported "
                           "next to the c-statistic rather than instead of it. For survival "
                           "outcomes the score is inverse-probability-of-censoring weighted "
                           "(Graf 1999): patients censored before the evaluation time "
                           "contribute nothing and the remainder are up-weighted by 1/G, "
                           "otherwise the score is biased by however long people happened to be "
                           "followed. IPA (index of prediction accuracy) = 1 - Brier/Brier_null; "
                           "it is the proportion of the null model's error that the model "
                           "removes. IPA <= 0 means the model is no better than predicting the "
                           "overall event rate for everyone.")

    # ── Discrimination ───────────────────────────────────────────────────────
    ws = wb.create_sheet('Discrimination')
    headers = ['Metric', 'Estimate (95% CI)', 'Optimism', 'Optimism-corrected', 'n', 'Events']
    r = sheet_header(ws, 'Discrimination', subtitle, headers, [34, 26, 14, 22, 10, 10])
    d = result['discrimination']
    internal = result.get('internal')
    r = write_row(ws, r, [
        d.get('metric', ''), fmt_ci(d.get('apparent'), d.get('ci_lo'), d.get('ci_hi')),
        fmt_num(internal['optimism_c']) if internal else 'N/A',
        fmt_num(internal['corrected_c']) if internal else 'N/A',
        d.get('n'), d.get('n_events')], bold_cols=(1,))
    r += 1
    if purpose == 'etiologic':
        r = write_footnote(ws, r, len(headers),
                           "This model was fitted for an etiologic question. The c-statistic is "
                           "a descriptive summary of rank discrimination only; it is NOT "
                           "prediction-model performance and calibration is deliberately not "
                           "reported. See the Estimand sheet.")
    else:
        r = write_footnote(ws, r, len(headers),
                           "The confidence interval is a percentile bootstrap on the fixed "
                           "predictions (it quantifies sampling noise in the estimate). The "
                           "optimism correction is a separate quantity: it refits the model in "
                           "each bootstrap resample and measures how much the apparent value "
                           "flatters the model. Report the optimism-corrected value as the "
                           "model's internally validated discrimination.")

    # ── Internal_validation ──────────────────────────────────────────────────
    if purpose == 'prediction':
        ws = wb.create_sheet('Internal_validation')
        headers = ['Quantity', 'Apparent', 'Optimism (mean)', 'Optimism (SD)',
                   'Optimism-corrected', 'Note']
        r = sheet_header(ws, 'Internal validation — optimism-corrected bootstrap',
                         subtitle, headers, [30, 16, 18, 16, 22, 48])
        if internal is None:
            r = write_row(ws, r, ['(not applicable)', 'N/A', 'N/A', 'N/A', 'N/A',
                                  'Predictions were supplied externally (predicted_column); '
                                  'no model was refitted, so there is no optimism to correct. '
                                  'The apparent values are already out-of-sample for this '
                                  'dataset.'])
        else:
            r = write_row(ws, r, [
                d.get('metric', 'C-index'), fmt_num(internal['apparent_c']),
                fmt_num(internal['optimism_c']), fmt_num(internal['optimism_c_sd']),
                fmt_num(internal['corrected_c']),
                'Optimism = mean over resamples of (performance in the resample the model '
                'was fitted on) - (same model applied to the original data).'],
                bold_cols=(1,))
            flag_slope = (np.isfinite(internal['corrected_slope'])
                          and internal['corrected_slope'] < 0.8)
            r = write_row(ws, r, [
                'Calibration slope', fmt_num(internal['apparent_slope']),
                fmt_num(internal['optimism_slope']), fmt_num(internal['optimism_slope_sd']),
                fmt_num(internal['corrected_slope']),
                'The corrected slope doubles as the uniform shrinkage factor: multiply every '
                'coefficient by it if the model is to be used prospectively.'],
                alt=True, flag=flag_slope, bold_cols=(1,))
            r = write_row(ws, r, [
                'Bootstrap resamples', f"{internal['n_requested']} requested",
                f"{internal['n_success']} usable", f"{internal['n_failed']} failed", '',
                'Resamples that failed to converge or produced a degenerate outcome '
                'distribution are excluded and counted here.'])
        r += 1
        r = write_footnote(ws, r, len(headers),
                           "Method: Harrell's optimism-corrected bootstrap. Split-sample "
                           "validation is deliberately NOT used -- at typical clinical sample "
                           "sizes it discards data and the estimate depends heavily on which "
                           "split happened to be drawn. The whole modelling procedure is "
                           "repeated inside every resample, so the correction reflects the cost "
                           "of the model as actually built. If any variable selection was done "
                           "by hand before this script ran, that step is NOT inside the "
                           "bootstrap and the correction will be too small.")

    # ── Method ───────────────────────────────────────────────────────────────
    ws = wb.create_sheet('Method')
    headers = ['Setting', 'Value', 'Source', 'Why it matters']
    r = sheet_header(ws, 'Methodological choices and defaults', subtitle, headers,
                     [34, 30, 16, 64])
    prov = config.get('_provenance', {})
    meta = [
        ('Study purpose', purpose, prov.get('study_purpose', 'default'),
         'Determines the whole branch: "etiologic" reports discrimination only; '
         '"prediction" runs calibration, Brier and internal validation.'),
        ('Outcome model', 'Cox proportional hazards' if at == 'cox' else 'Logistic regression',
         'config', 'Fixed by analysis_type.'),
        ('Prediction source',
         'external (predicted_column)' if result['external_predictions'] else 'fitted here',
         'config', 'External predictions make the apparent slope a genuine calibration '
                   'slope; internally fitted predictions do not.'),
        ('Bootstrap resamples',
         str(config['bootstrap_n'])
         + (' — FAST MODE' if int(config['bootstrap_n']) < 500 else ''),
         prov.get('bootstrap_n', 'default'),
         'Default 500. The protocol minimum is 200, which is a fast mode: percentile '
         'CIs and the optimism correction are both still visibly noisy there, so a '
         'value below 500 is marked as such and should be stated in the manuscript. '
         'Below 200 is not acceptable.'),
        ('Random seed', str(config['seed']), prov.get('seed', 'default'),
         'Fixed so the bootstrap is reproducible. Changing it changes the third decimal '
         'place, not the conclusion.'),
        ('Calibration bins', str(config['n_bins']), prov.get('n_bins', 'default'),
         'Equal-sized groups of predicted risk. Deciles is convention; with few events, '
         'fewer bins are more stable.'),
        ('Evaluation times',
         ', '.join(f"{t:g}" for t in result['eval_times']) if at == 'cox' else 'n/a (binary outcome)',
         prov.get('eval_times', 'default'),
         'Must be pre-specified clinically meaningful horizons and must lie inside the '
         'observed follow-up. Choosing them after seeing the results is a form of '
         'selective reporting.'),
        ('Calibration slope method',
         'Cox refitted on the linear predictor (van Houwelingen)' if at == 'cox'
         else 'logistic regression of outcome on the linear predictor',
         'fixed', 'Standard definition; 1.0 = predictions on the correct scale.'),
        ('Calibration intercept method',
         'Poisson model with log(model cumulative hazard) as offset (Crowson 2016)'
         if at == 'cox' else 'logistic regression with the linear predictor as offset',
         'fixed', 'Cox has no intercept of its own, so calibration-in-the-large is '
                  'expressed as log(observed/expected events) at each evaluation time.'),
        ('Brier score', 'IPCW-weighted (Graf 1999)' if at == 'cox' else 'unweighted',
         'fixed', 'Censoring makes the unweighted survival Brier score biased.'),
        ('Internal validation', "Harrell optimism-corrected bootstrap", 'fixed',
         'Split-sample is explicitly not used.'),
        ('Missing data', 'listwise deletion on model variables', 'fixed',
         f"{result.get('n_dropped', 0)} row(s) dropped. If this is more than a few percent, "
         f"run the multiple-imputation step first and calibrate on the imputed data."),
        ('Blocking bootstrap threshold',
         f"{config['min_bootstrap_success_frac']:.0%} of resamples must converge",
         'default', 'Below this the optimism correction is reported as non-convergent and '
                    'the run exits with a blocking status.'),
    ]
    for i, (k, v, src, why) in enumerate(meta):
        r = write_row(ws, r, [k, v, src, why], alt=(i % 2 == 1))
    r += 1
    if _QC_MESSAGES:
        r = write_row(ws, r, ['QC messages raised by this run', '', '', ''], bold_cols=(1,))
        for i, m in enumerate(_QC_MESSAGES):
            r = write_row(ws, r, ['', m, '', ''], alt=(i % 2 == 1),
                          flag=m.startswith('[QC_BLOCKING]'))
    r += 1
    r = write_footnote(ws, r, len(headers),
                       "Every row above is a methodological choice that belongs in "
                       "decision_log.md with a permanent ID. \"default\" in the Source column "
                       "means this script chose it, not the analyst -- if a reviewer asks why, "
                       "that is the honest answer, and the value can be overridden in the "
                       "config.")

    out = config['output_file']
    os.makedirs(os.path.dirname(os.path.abspath(out)) or '.', exist_ok=True)
    wb.save(out)
    print(f"  Excel written: {out}")
    return out


# ═══════════════════════════════════════════════════════════════════════════════
# Plot
# ═══════════════════════════════════════════════════════════════════════════════

def write_plot(config, result):
    if result['study_purpose'] == 'etiologic':
        print("  Calibration plot skipped (study_purpose = 'etiologic')")
        return None
    cal = result.get('calibration') or {}
    at = result['analysis_type']
    if at == 'cox':
        panels = [(f"t = {pt['t']:g}   (O/E = {fmt_num(pt.get('oe_ratio'), 2)})",
                   pt['deciles'], pt['pred_risk']) for pt in cal.get('per_time', [])]
    else:
        panels = [(f"Binary outcome   (calibration intercept = "
                   f"{fmt_num(cal.get('intercept'), 2)})", cal.get('deciles', []), cal.get('pred'))]
    if not panels:
        return None

    ncol = len(panels)
    fig, axes = plt.subplots(1, ncol, figsize=(4.6 * ncol, 4.8), squeeze=False)
    for ax, (tag, deciles, pred) in zip(axes[0], panels):
        pr = [d['pred_mean'] for d in deciles]
        ob = [d['obs'] for d in deciles]
        lo = [max(0.0, d['obs'] - d['obs_lo']) if np.isfinite(d.get('obs_lo', np.nan)) else 0
              for d in deciles]
        hi = [max(0.0, d['obs_hi'] - d['obs']) if np.isfinite(d.get('obs_hi', np.nan)) else 0
              for d in deciles]
        top = max([v for v in pr + ob if np.isfinite(v)] + [0.05]) * 1.15
        ax.plot([0, top], [0, top], color='#888888', linestyle='--', linewidth=1,
                label='Perfect calibration', zorder=1)
        ax.errorbar(pr, ob, yerr=[lo, hi], fmt='o', color='#2E5FA3', markersize=6,
                    capsize=3, linewidth=1.2, label='Observed by decile', zorder=3)
        ax.plot(pr, ob, color='#2E5FA3', linewidth=1.0, alpha=0.6, zorder=2)
        if pred is not None and len(pred):
            # Spike histogram of predicted risk along the bottom axis: it shows where
            # the patients actually are, so a wild-looking extreme decile can be read
            # as "few patients" rather than "model broken".
            counts, edges = np.histogram(np.asarray(pred, dtype=float), bins=30,
                                         range=(0, top))
            if counts.max() > 0:
                heights = counts / counts.max() * top * 0.14
                ax.bar(edges[:-1], heights, width=np.diff(edges), align='edge',
                       color='#B7C9E2', alpha=0.75, zorder=0,
                       label='Distribution of predicted risk')
        ax.set_xlim(0, top)
        ax.set_ylim(0, top)
        ax.set_xlabel('Predicted risk')
        ax.set_ylabel('Observed risk' + (' (Kaplan-Meier)' if at == 'cox' else ''))
        ax.set_title(tag, fontsize=10, fontweight='bold')
        ax.grid(alpha=0.25, linewidth=0.5)
        ax.legend(fontsize=8, loc='upper left', frameon=False)

    sr = cal.get('slope_reported')
    kind = cal.get('slope_reported_kind', '')
    itc = cal.get('intercept')
    bits = [config.get('title', 'Calibration')]
    sub2 = f"Calibration slope = {fmt_num(sr)} ({kind})"
    if itc is not None and np.isfinite(itc):
        sub2 += f"   |   Calibration intercept = {fmt_num(itc)}"
    fig.suptitle(bits[0], fontsize=12, fontweight='bold', y=1.00)
    fig.text(0.5, 0.945, sub2, ha='center', fontsize=9, color='#444444')
    fig.tight_layout(rect=[0, 0, 1, 0.93])
    out = config['plot_file']
    os.makedirs(os.path.dirname(os.path.abspath(out)) or '.', exist_ok=True)
    fig.savefig(out, dpi=300, bbox_inches='tight')
    plt.close(fig)
    print(f"  Plot written: {out}")
    return out


# ═══════════════════════════════════════════════════════════════════════════════
# --check-missing preview (never fits anything)
# ═══════════════════════════════════════════════════════════════════════════════

def preview(config, df):
    at = config['analysis_type']
    print("=" * 96)
    print("CALIBRATION PREVIEW (--check-missing)")
    print("=" * 96)
    print(f"Data file : {config['data_file']}")
    print(f"Rows after filters/derived variables : {len(df)}")
    print(f"analysis_type : {at}    study_purpose : {config['study_purpose']}")
    print("-" * 96)

    rows = []
    for v in model_variables(config):
        rows.append((v['label'], v['column'], v['type']))
    if config.get('predicted_column'):
        rows.append(('[predictions]', config['predicted_column'], 'continuous'))
    if at == 'cox':
        cs = config['cox_settings']
        rows.append(('[outcome] time', cs['time_column'], 'continuous'))
        rows.append(('[outcome] event', cs['event_column'], 'binary'))
    else:
        rows.append(('[outcome]', config['logistic_settings']['outcome_column'], 'binary'))

    print(f"{'Variable':<30} {'Column':<24} {'Type':<14} {'Missing':>16}")
    print("-" * 96)
    present = []
    for label, col, vtype in rows:
        if col not in df.columns:
            print(f"{label:<30} {col:<24} {vtype:<14} {'COLUMN NOT FOUND':>16}")
            continue
        s = (df[col] if vtype == 'categorical'
             else pd.to_numeric(df[col], errors='coerce'))
        nm = int(s.isna().sum())
        pct = 100.0 * nm / len(df) if len(df) else 0.0
        present.append(col)
        print(f"{label:<30} {col:<24} {vtype:<14} {nm:>5}/{len(df):<5} ({pct:>4.1f}%)")

    if present:
        complete = df[present].apply(
            lambda c: pd.to_numeric(c, errors='coerce') if c.dtype != object else c)
        n_complete = int(complete.notna().all(axis=1).sum())
        print("-" * 96)
        print(f"Complete cases across all model variables : {n_complete}/{len(df)} "
              f"({100.0 * n_complete / len(df):.1f}%)")

    print("-" * 96)
    print(f"Planned bootstrap resamples : {config['bootstrap_n']} "
          f"({config['_provenance'].get('bootstrap_n', 'default')})"
          + ("  -- FAST MODE, below the 500 default"
             if int(config['bootstrap_n']) < 500 else ""))
    print(f"Planned calibration bins    : {config['n_bins']} "
          f"({config['_provenance'].get('n_bins', 'default')})")
    if at == 'cox':
        et = config['eval_times']
        print(f"Planned evaluation times    : {', '.join(f'{t:g}' for t in et)}")
        tc = config['cox_settings']['time_column']
        if tc in df.columns:
            mx = pd.to_numeric(df[tc], errors='coerce').max()
            print(f"Maximum observed follow-up  : {mx:g}")
            over = [t for t in et if t > mx]
            if over:
                print(f"  WARNING: eval_times beyond follow-up: "
                      f"{', '.join(f'{t:g}' for t in over)} -- the full run will BLOCK")
    print("=" * 96)
    print("[preview] No model was fitted (--check-missing). No Excel or PNG was written.")
    print("=" * 96)


# ═══════════════════════════════════════════════════════════════════════════════
# Entry point
# ═══════════════════════════════════════════════════════════════════════════════

def main():
    if len(sys.argv) < 2:
        print("Usage: python calibration.py config.json [--check-missing]")
        sys.exit(1)

    config = load_config(sys.argv[1])
    check_only = '--check-missing' in sys.argv

    df = load_data(config)
    print(f"Mode: calibration ({config['analysis_type']}, {config['study_purpose']})")
    print(f"Data loaded: n={len(df)}")

    if check_only:
        preview(config, df)
        sys.exit(0)

    result = run_analysis(config, df)

    if n_blocking() == 0 or result.get('discrimination'):
        write_excel(config, result)
        if n_blocking() == 0:
            write_plot(config, result)

    print("\nDone!")
    nb = n_blocking()
    nw = len(_QC_MESSAGES) - nb
    if _QC_MESSAGES:
        print(f"  QC: {nb} blocking, {nw} warning(s) written to stderr")
    sys.exit(3 if nb else 0)


if __name__ == '__main__':
    main()
