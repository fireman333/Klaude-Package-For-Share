#!/usr/bin/env python3
"""
Immortal Time Bias — Landmark analysis and extended (time-varying) Cox

Implements references/advanced_analyses.md section 7. One script, two modes
selected by the config's `mode` field:

    mode = "landmark"      -> Landmark_*.xlsx  + Landmark_*.png
    mode = "extended_cox"  -> ExtendedCox_*.xlsx + ExtendedCox_*.png

Usage:
    python immortal_time.py config.json                  # full run
    python immortal_time.py config.json --check-missing  # preview only, no models

Why this script exists
----------------------
When the exposure grouping uses information that only becomes available after
time zero (e.g. "did the patient get surgery during the admission?"), everyone
in the exposed arm had to survive long enough to be exposed. That guaranteed
event-free stretch is immortal time. Handing it to a Cox model as if it were
baseline exposure manufactures a protective effect out of nothing.

Both modes therefore always fit, and always report, the naive baseline-covariate
Cox model side by side with the corrected estimate — the size of the gap IS the
immortal time bias, and the reader is entitled to see it.

Design notes
------------
* Landmark mode runs EVERY landmark time in `landmark_times`. Reporting a single
  cherry-picked landmark is explicitly forbidden by the protocol, so the script
  has no option to do it.
* The people dropped at each landmark are not a random subset — `Excluded_at_landmark`
  reports how many they were and what their event rate was, per landmark.
* Extended Cox writes `Counting_process_sample`: the full tstart/tstop/event
  expansion for a handful of real patients so a human can verify the reshape.
* QC goes to stderr with the `[QC_WARNING]` / `[QC_BLOCKING]` prefixes the pipeline
  runner greps for. Any blocking finding exits 3.
"""

import os
import sys
import json
import math
import operator
import warnings

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ═══════════════════════════════════════════════════════════════════════════════
# QC messaging
# ═══════════════════════════════════════════════════════════════════════════════

_QC_MESSAGES = []
_BLOCKING = []


def qc(kind, message):
    line = f"[QC_{kind}] {message}"
    _QC_MESSAGES.append(line)
    print(line, file=sys.stderr)


def qc_warning(message):
    qc('WARNING', message)


def qc_blocking(message):
    _BLOCKING.append(message)
    qc('BLOCKING', message)


def die_blocking(message):
    """Blocking finding that makes any further computation meaningless."""
    qc_blocking(message)
    print("\n[ABORTED] Blocking QC finding — no output written.", file=sys.stderr)
    sys.exit(3)


# ═══════════════════════════════════════════════════════════════════════════════
# Configuration
# ═══════════════════════════════════════════════════════════════════════════════

def load_config(path):
    with open(path, 'r', encoding='utf-8') as f:
        config = json.load(f)

    config.setdefault('mode', 'landmark')
    config.setdefault('filters', [])
    config.setdefault('derived_variables', [])
    config.setdefault('sections', [])
    config.setdefault('title', '')
    config.setdefault('subtitle', '')
    config.setdefault('time_unit', 'days')
    config.setdefault('landmark_times', [])
    config.setdefault('landmark_rationale', '')
    config.setdefault('late_exposure_handling', 'unexposed_at_landmark')
    config.setdefault('tv_hr_cutpoints', None)
    config.setdefault('thresholds', {})
    config['thresholds'].setdefault('max_excluded_pct', 30.0)
    config['thresholds'].setdefault('min_events_per_group', 10)

    if config['mode'] not in ('landmark', 'extended_cox'):
        sys.exit(f"[FATAL] Unknown mode: {config['mode']!r} "
                 f"(expected 'landmark' or 'extended_cox')")

    for key in ('data_file', 'grouping', 'cox_settings', 'exposure_time_column'):
        if not config.get(key):
            sys.exit(f"[FATAL] config is missing required key: {key!r}")

    cs = config['cox_settings']
    for key in ('time_column', 'event_column'):
        if not cs.get(key):
            sys.exit(f"[FATAL] cox_settings is missing required key: {key!r}")

    if config['mode'] == 'landmark' and not config['landmark_times']:
        sys.exit("[FATAL] mode='landmark' requires a non-empty `landmark_times` list "
                 "(pre-specified, see advanced_analyses.md section 7)")

    config['landmark_times'] = sorted({float(t) for t in config['landmark_times']})

    if not config.get('output_file'):
        sys.exit("[FATAL] config is missing required key: 'output_file'")
    if not config.get('plot_file'):
        config['plot_file'] = os.path.splitext(config['output_file'])[0] + '.png'

    return config


# ═══════════════════════════════════════════════════════════════════════════════
# Data loading (same semantics as missing_data.py / multivariate_interaction.py)
# ═══════════════════════════════════════════════════════════════════════════════

OPS = {
    '==': operator.eq, '!=': operator.ne,
    '>': operator.gt, '>=': operator.ge,
    '<': operator.lt, '<=': operator.le,
}


def load_data(config):
    path = config['data_file']
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xls'):
        df = pd.read_excel(path)
    elif ext == '.tsv':
        df = pd.read_csv(path, sep='\t')
    elif ext == '.sas7bdat':
        df = pd.read_sas(path)
    elif ext == '.dta':
        df = pd.read_stata(path)
    else:
        df = pd.read_csv(path)

    for f in config['filters']:
        col, op, val = f['column'], f['operator'], f['value']
        if op in OPS:
            df = df[OPS[op](pd.to_numeric(df[col], errors='coerce'), val)]
        elif op == 'in':
            df = df[df[col].isin(val)]
        elif op == 'not_in':
            df = df[~df[col].isin(val)]

    for dv in config['derived_variables']:
        if dv['type'] == 'coalesce':
            result = pd.to_numeric(df[dv['columns'][0]], errors='coerce')
            for c in dv['columns'][1:]:
                result = result.combine_first(pd.to_numeric(df[c], errors='coerce'))
            df[dv['name']] = result
        elif dv['type'] == 'expression':
            local_ns = {}
            for dep in dv.get('dependencies', []):
                local_ns[dep] = pd.to_numeric(df[dep], errors='coerce')
            local_ns['np'] = np
            df[dv['name']] = eval(dv['expression'], {"__builtins__": {}}, local_ns)

    return df.reset_index(drop=True)


# ═══════════════════════════════════════════════════════════════════════════════
# Covariates / exposure definition
# ═══════════════════════════════════════════════════════════════════════════════

EXPOSED = '_exposed_ever'
EXPOSED_LM = '_exposed_at_landmark'
TIME = '_time'
EVENT = '_event'
EXPTIME = '_exposure_time'
PID = '_pid'


def covariate_variables(config):
    """Flatten config['sections'] into a covariate list."""
    out = []
    for section in config['sections']:
        for v in section.get('variables', []):
            out.append({
                'column': v['column'],
                'label': v.get('label', v['column']),
                'type': v.get('type', 'continuous'),
                'reference_level': v.get('reference_level'),
                'section': section.get('header', ''),
            })
    return out


def build_covariate_terms(df, variables):
    """Add model terms to df. Categoricals become dummies vs a reference level.

    Returns (df, [(term_column, term_label)]).
    """
    work = df.copy()
    terms = []
    for v in variables:
        col = v['column']
        if col not in work.columns:
            die_blocking(f"Covariate column '{col}' not found in the data file")
        if v['type'] == 'categorical':
            s = work[col].astype(str).where(work[col].notna(), np.nan)
            levels = sorted([x for x in s.dropna().unique()])
            if len(levels) < 2:
                qc_warning(f"Covariate '{col}' has < 2 observed levels; dropped from the model")
                continue
            ref = (str(v['reference_level'])
                   if v.get('reference_level') is not None else levels[0])
            if ref not in levels:
                qc_warning(f"Reference level {ref!r} of '{col}' not observed; "
                           f"using {levels[0]!r} instead")
                ref = levels[0]
            for lev in levels:
                if lev == ref:
                    continue
                tcol = f"{col}__{lev}"
                work[tcol] = np.where(s.isna(), np.nan, (s == lev).astype(float))
                terms.append((tcol, f"{v['label']}: {lev} vs {ref}"))
        else:
            work[col] = pd.to_numeric(work[col], errors='coerce')
            terms.append((col, v['label']))
    return work, terms


def exposure_labels(config):
    """(exposed_value, exposed_label, unexposed_label). groups[0] is the exposed arm."""
    gc = config['grouping']
    groups = gc.get('groups', [])
    if len(groups) < 2:
        sys.exit("[FATAL] grouping.groups must contain the exposed arm first and the "
                 "reference arm second")
    exposed_value = gc.get('exposed_value', groups[0].get('value'))
    exp_label = next((g.get('label', str(g.get('value'))) for g in groups
                      if g.get('value') == exposed_value), str(exposed_value))
    unexp_label = next((g.get('label', str(g.get('value'))) for g in groups
                        if g.get('value') != exposed_value), 'Unexposed')
    return exposed_value, exp_label, unexp_label


def prepare(config, df):
    """Build the analysis frame: _pid, _time, _event, _exposed_ever, _exposure_time,
    plus covariate terms. Returns (work, terms, prep_notes)."""
    cs = config['cox_settings']
    time_col, event_col = cs['time_column'], cs['event_column']
    grp_col = config['grouping']['column']
    exp_time_col = config['exposure_time_column']
    exposed_value, exp_label, _ = exposure_labels(config)

    for col, what in ((time_col, 'cox_settings.time_column'),
                      (event_col, 'cox_settings.event_column'),
                      (grp_col, 'grouping.column')):
        if col not in df.columns:
            die_blocking(f"{what} '{col}' not found in the data file")

    # ── Blocking: exposure_time_column must exist ────────────────────────────
    if exp_time_col not in df.columns:
        die_blocking(f"exposure_time_column '{exp_time_col}' not found in the data file "
                     f"— immortal time cannot be quantified without the exposure timing")

    work = df.copy()
    work[TIME] = pd.to_numeric(work[time_col], errors='coerce')
    work[EVENT] = pd.to_numeric(work[event_col], errors='coerce')
    grp_vals = pd.to_numeric(work[grp_col], errors='coerce') \
        if isinstance(exposed_value, (int, float)) and not isinstance(exposed_value, bool) \
        else work[grp_col]
    work[EXPOSED] = np.where(grp_vals.isna(), np.nan, (grp_vals == exposed_value).astype(float))
    work[EXPTIME] = pd.to_numeric(work[exp_time_col], errors='coerce')

    variables = covariate_variables(config)
    work, terms = build_covariate_terms(work, variables)

    n_raw = len(work)
    needed = [TIME, EVENT, EXPOSED] + [t for t, _ in terms]
    work = work.dropna(subset=needed).copy()
    n_listwise = n_raw - len(work)

    work = work[(work[TIME] > 0)].copy()
    n_nonpositive = n_raw - n_listwise - len(work)
    if n_nonpositive:
        qc_warning(f"{n_nonpositive} row(s) dropped for non-positive follow-up time")

    if len(work) == 0:
        die_blocking("No rows survive listwise deletion of the model variables")

    work[PID] = np.arange(1, len(work) + 1)
    work = work.reset_index(drop=True)

    # ── Blocking: exposed subjects must have a known exposure time ───────────
    exposed_mask = work[EXPOSED] == 1
    n_exposed = int(exposed_mask.sum())
    if n_exposed == 0:
        die_blocking(f"No subject matches the exposed arm ('{exp_label}', "
                     f"{grp_col} == {exposed_value!r})")
    n_missing_exptime = int(work.loc[exposed_mask, EXPTIME].isna().sum())
    if n_missing_exptime:
        die_blocking(f"{n_missing_exptime}/{n_exposed} exposed subject(s) have a missing "
                     f"'{exp_time_col}' — exposure timing is required to correct immortal time")
    n_negative = int((work.loc[exposed_mask, EXPTIME] < 0).sum())
    if n_negative:
        die_blocking(f"{n_negative} exposed subject(s) have a negative '{exp_time_col}'")

    # Unexposed subjects carrying an exposure time is a data-consistency smell.
    n_unexp_with_time = int(work.loc[~exposed_mask, EXPTIME].notna().sum())
    if n_unexp_with_time:
        qc_warning(f"{n_unexp_with_time} unexposed subject(s) carry a non-missing "
                   f"'{exp_time_col}'; their exposure time is ignored")
    work.loc[~exposed_mask, EXPTIME] = np.nan

    n_after_fu = int((work.loc[exposed_mask, EXPTIME] >= work.loc[exposed_mask, TIME]).sum())
    if n_after_fu:
        qc_warning(f"{n_after_fu} exposed subject(s) have exposure time >= follow-up time; "
                   f"they contribute no exposed person-time and are handled as unexposed "
                   f"in the counting-process reshape")

    prep = {
        'n_raw': n_raw,
        'n_listwise_dropped': n_listwise,
        'n_nonpositive_time': n_nonpositive,
        'n_analysis': len(work),
        'n_exposed': n_exposed,
        'n_unexposed': int(len(work) - n_exposed),
        'n_events': int(work[EVENT].sum()),
        'immortal_person_time': float(work.loc[exposed_mask, EXPTIME].clip(upper=None).sum()),
        'median_exposure_time': float(work.loc[exposed_mask, EXPTIME].median()),
        'n_exposure_after_fu': n_after_fu,
    }
    return work, terms, prep


# ═══════════════════════════════════════════════════════════════════════════════
# Cox fitting helpers
# ═══════════════════════════════════════════════════════════════════════════════

def _clean_terms(d, terms, target):
    """Drop zero-variance terms so lifelines does not choke."""
    keep = []
    for t in terms:
        if t == target:
            continue
        if d[t].nunique() < 2:
            qc_warning(f"Covariate term '{t}' has zero variance in this model; dropped")
            continue
        keep.append(t)
    return keep


MAX_ABS_COEF = 20.0     # |log HR| this large is separation, not an estimate
MAX_SE = 100.0          # a standard error this wide means the term is unidentified


def _scan_convergence(summary, context, caught_warnings=()):
    """Blocking check on the WHOLE fitted model, not only the exposure term.
    A covariate blowing up means the reported exposure HR cannot be trusted either."""
    ok = True
    for w in caught_warnings:
        if 'Convergence' in w.category.__name__:
            first = str(w.message).strip().splitlines()[0]
            qc_blocking(f"{context}: model did not converge — {first}")
            ok = False
    for term in summary.index:
        coef = float(summary.loc[term, 'coef'])
        se = float(summary.loc[term, 'se(coef)'])
        if not np.isfinite(coef) or not np.isfinite(se) or se == 0:
            qc_blocking(f"{context}: model did not converge (term '{term}': "
                        f"coef={coef}, se={se})")
            ok = False
        elif abs(coef) > MAX_ABS_COEF or se > MAX_SE:
            qc_blocking(f"{context}: model did not converge (term '{term}': "
                        f"coef={coef:.2f}, se={se:.2f}) — likely complete separation "
                        f"or collinearity")
            ok = False
    return ok


def fit_cox(data, term_cols, target, time_col=TIME, event_col=EVENT, context=''):
    """Fixed-covariate Cox. Returns dict with HR/CI/p for `target`, or None."""
    from lifelines import CoxPHFitter
    from lifelines.exceptions import ConvergenceError

    cols = [time_col, event_col, target] + [c for c in term_cols if c != target]
    d = data[cols].dropna().copy()
    if len(d) < 10 or d[event_col].sum() < 3:
        qc_warning(f"{context}: too few observations/events to fit a Cox model "
                   f"(n={len(d)}, events={int(d[event_col].sum())})")
        return None
    keep = _clean_terms(d, [c for c in cols if c not in (time_col, event_col)], target)
    if d[target].nunique() < 2:
        qc_warning(f"{context}: exposure has no variation; model not fitted")
        return None
    d = d[[time_col, event_col, target] + keep]

    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            cph = CoxPHFitter()
            cph.fit(d, duration_col=time_col, event_col=event_col)
    except (ConvergenceError, Exception) as e:  # noqa: B014 - lifelines raises broadly
        qc_blocking(f"{context}: Cox model did not converge ({type(e).__name__}: {e})")
        return None

    s = cph.summary
    if target not in s.index:
        qc_blocking(f"{context}: exposure term missing from the fitted model")
        return None
    if not _scan_convergence(s, context, caught):
        return None
    coef = float(s.loc[target, 'coef'])
    se = float(s.loc[target, 'se(coef)'])
    return {
        'coef': coef, 'se': se,
        'hr': float(np.exp(coef)),
        'ci_low': float(np.exp(s.loc[target, 'coef lower 95%'])),
        'ci_high': float(np.exp(s.loc[target, 'coef upper 95%'])),
        'p': float(s.loc[target, 'p']),
        'n': int(len(d)),
        'events': int(d[event_col].sum()),
        'n_covariates': len(keep),
    }


def fit_ctv(cp, term_cols, targets, context=''):
    """CoxTimeVaryingFitter on counting-process data. Returns {target: result} or None."""
    from lifelines import CoxTimeVaryingFitter
    from lifelines.exceptions import ConvergenceError

    cols = [PID, 'tstart', 'tstop', 'event'] + list(targets) + \
           [c for c in term_cols if c not in targets]
    d = cp[cols].dropna().copy()
    keep = []
    for c in [c for c in term_cols if c not in targets]:
        if d[c].nunique() < 2:
            qc_warning(f"{context}: covariate '{c}' has zero variance; dropped")
            continue
        keep.append(c)
    live_targets = [t for t in targets if d[t].nunique() >= 2]
    if not live_targets:
        qc_warning(f"{context}: exposure has no variation in the counting-process data")
        return None
    d = d[[PID, 'tstart', 'tstop', 'event'] + live_targets + keep]

    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            ctv = CoxTimeVaryingFitter()
            ctv.fit(d, id_col=PID, event_col='event', start_col='tstart', stop_col='tstop',
                    show_progress=False)
    except (ConvergenceError, Exception) as e:  # noqa: B014
        qc_blocking(f"{context}: extended Cox model did not converge "
                    f"({type(e).__name__}: {e})")
        return None

    s = ctv.summary
    if not _scan_convergence(s, context, caught):
        return None
    out = {}
    for t in targets:
        if t not in s.index:
            out[t] = None
            continue
        coef = float(s.loc[t, 'coef'])
        se = float(s.loc[t, 'se(coef)'])
        out[t] = {
            'coef': coef, 'se': se,
            'hr': float(np.exp(coef)),
            'ci_low': float(np.exp(s.loc[t, 'coef lower 95%'])),
            'ci_high': float(np.exp(s.loc[t, 'coef upper 95%'])),
            'p': float(s.loc[t, 'p']),
            'n_rows': int(len(d)),
            'n_subjects': int(d[PID].nunique()),
            'events': int(d['event'].sum()),
            'n_covariates': len(keep),
        }
    return out


# ═══════════════════════════════════════════════════════════════════════════════
# Naive (biased) model — always computed, in both modes
# ═══════════════════════════════════════════════════════════════════════════════

def run_naive(work, terms):
    """Exposure treated as a baseline covariate over the whole follow-up.
    This is the model that carries the immortal time bias."""
    term_cols = [t for t, _ in terms]
    crude = fit_cox(work, [EXPOSED], EXPOSED, context='Naive Cox (crude)')
    adj = fit_cox(work, [EXPOSED] + term_cols, EXPOSED, context='Naive Cox (adjusted)')
    return {'crude': crude, 'adjusted': adj}


# ═══════════════════════════════════════════════════════════════════════════════
# MODE 1 — Landmark analysis
# ═══════════════════════════════════════════════════════════════════════════════

def run_landmark(config, work, terms):
    """Run EVERY landmark time in the config. Never a single cherry-picked one."""
    term_cols = [t for t, _ in terms]
    thr = config['thresholds']
    results, excluded_rows = [], []

    for t in config['landmark_times']:
        n_total = len(work)
        keep_mask = work[TIME] > t
        kept = work[keep_mask].copy()
        dropped = work[~keep_mask].copy()

        n_excl = int(len(dropped))
        pct_excl = 100.0 * n_excl / n_total if n_total else float('nan')
        excl_events = int(dropped[EVENT].sum())
        kept_events = int(kept[EVENT].sum())

        # Exposure status AS OF the landmark.
        kept[EXPOSED_LM] = np.where(kept[EXPTIME].notna() & (kept[EXPTIME] <= t), 1.0, 0.0)
        n_late = int(((kept[EXPOSED] == 1) & (kept[EXPTIME] > t)).sum())
        kept['_lm_time'] = kept[TIME] - t

        n_exp = int((kept[EXPOSED_LM] == 1).sum())
        n_unexp = int((kept[EXPOSED_LM] == 0).sum())
        ev_exp = int(kept.loc[kept[EXPOSED_LM] == 1, EVENT].sum())
        ev_unexp = int(kept.loc[kept[EXPOSED_LM] == 0, EVENT].sum())

        ctx = f"Landmark t={t:g}"
        crude = adj = None
        if n_exp > 0 and n_unexp > 0 and len(kept) >= 10:
            crude = fit_cox(kept, [EXPOSED_LM], EXPOSED_LM,
                            time_col='_lm_time', context=f'{ctx} (crude)')
            adj = fit_cox(kept, [EXPOSED_LM] + term_cols, EXPOSED_LM,
                          time_col='_lm_time', context=f'{ctx} (adjusted)')
        else:
            qc_warning(f"{ctx}: one arm is empty at the landmark "
                       f"(exposed={n_exp}, unexposed={n_unexp}); model not fitted")

        # ── QC: exclusion fraction, event counts ─────────────────────────────
        if pct_excl > thr['max_excluded_pct']:
            qc_warning(f"{ctx}: excludes {n_excl}/{n_total} ({pct_excl:.1f}%) of the cohort, "
                       f"above the {thr['max_excluded_pct']:g}% threshold — the landmark "
                       f"population is far from the enrolled population")
        min_ev = thr['min_events_per_group']
        if ev_exp < min_ev or ev_unexp < min_ev:
            qc_warning(f"{ctx}: fewer than {min_ev} events in an arm after the landmark "
                       f"(exposed={ev_exp}, unexposed={ev_unexp}); the estimate is unstable")

        results.append({
            'landmark': t, 'n_total': n_total, 'n_analyzed': int(len(kept)),
            'n_excluded': n_excl, 'pct_excluded': pct_excl,
            'n_exposed': n_exp, 'n_unexposed': n_unexp,
            'events_exposed': ev_exp, 'events_unexposed': ev_unexp,
            'n_late_exposure': n_late,
            'crude': crude, 'adjusted': adj,
            'kept': kept,
        })
        excluded_rows.append({
            'landmark': t, 'n_total': n_total,
            'n_excluded': n_excl, 'pct_excluded': pct_excl,
            'excl_events': excl_events,
            'excl_event_rate': 100.0 * excl_events / n_excl if n_excl else float('nan'),
            'excl_censored': n_excl - excl_events,
            'excl_median_fu': float(dropped[TIME].median()) if n_excl else float('nan'),
            'excl_would_be_exposed': int((dropped[EXPOSED] == 1).sum()),
            'n_kept': int(len(kept)),
            'kept_events': kept_events,
            'kept_event_rate': 100.0 * kept_events / len(kept) if len(kept) else float('nan'),
        })

    # ── QC: direction consistency across landmarks ───────────────────────────
    signs = []
    for r in results:
        est = r['adjusted'] or r['crude']
        if est:
            signs.append((r['landmark'], est['coef'], est['hr']))
    if len(signs) >= 2 and len({1 if c > 0 else -1 for _, c, _ in signs}) > 1:
        detail = ', '.join(f"t={t:g}: HR {hr:.2f}" for t, _, hr in signs)
        qc_warning(f"Landmark results point in different directions across time points "
                   f"({detail}) — all of them must be reported together; selecting one "
                   f"is not permitted")

    return results, excluded_rows


# ═══════════════════════════════════════════════════════════════════════════════
# MODE 2 — Extended Cox (counting-process)
# ═══════════════════════════════════════════════════════════════════════════════

def check_counting_process(cp, work):
    """Blocking integrity checks on the reshaped data. A wrong reshape is the main
    failure mode of extended Cox, so this is checked, not assumed."""
    bad = cp[cp['tstart'] >= cp['tstop']]
    if len(bad):
        ids = ', '.join(str(int(i)) for i in bad[PID].head(5))
        qc_blocking(f"Counting-process reshape produced {len(bad)} row(s) with "
                    f"tstart >= tstop (subject id: {ids}) — the reshape is wrong")
    # Every subject must keep exactly one terminal event flag and contiguous intervals.
    for pid, g in cp.groupby(PID):
        if int(g['event'].sum()) > 1:
            qc_blocking(f"Counting-process reshape gave subject {pid} more than one event row")
            break
        if len(g) > 1 and not np.allclose(g['tstart'].values[1:], g['tstop'].values[:-1]):
            qc_blocking(f"Counting-process intervals for subject {pid} are not contiguous")
            break
        if len(g) > 1 and int(g['event'].values[-1]) == 0 and int(g['event'].sum()) == 1:
            qc_blocking(f"Counting-process reshape put the event of subject {pid} on a "
                        f"non-terminal interval")
            break

    n_ev_cp = int(cp['event'].sum())
    n_ev_raw = int(work[EVENT].sum())
    if n_ev_cp != n_ev_raw:
        qc_blocking(f"Counting-process reshape changed the number of events "
                    f"({n_ev_raw} -> {n_ev_cp})")
    return not _BLOCKING


def build_counting_process(work, terms):
    """One or two rows per subject:

        never exposed            -> [0, fu, event, exposed=0]
        exposed at time e (0<e<fu) -> [0, e, 0, 0] and [e, fu, event, 1]
        exposed at e >= fu       -> [0, fu, event, 0]   (no exposed person-time observed)
        exposed at e <= 0        -> [0, fu, event, 1]   (exposed from time zero)

    The first row of an exposed subject is exactly the immortal time the naive
    model mis-attributes to the exposed arm.
    """
    term_cols = [t for t, _ in terms]
    rows = []
    n_two_row = n_zero = n_after_fu = 0
    for _, r in work.iterrows():
        base = {c: r[c] for c in term_cols}
        base[PID] = int(r[PID])
        fu = float(r[TIME])
        ev = int(r[EVENT])
        e = r[EXPTIME]
        exposed_ever = int(r[EXPOSED])

        if exposed_ever == 1 and pd.notna(e) and 0 < float(e) < fu:
            n_two_row += 1
            rows.append(dict(base, tstart=0.0, tstop=float(e), event=0, exposed=0.0))
            rows.append(dict(base, tstart=float(e), tstop=fu, event=ev, exposed=1.0))
        elif exposed_ever == 1 and pd.notna(e) and float(e) <= 0:
            n_zero += 1
            rows.append(dict(base, tstart=0.0, tstop=fu, event=ev, exposed=1.0))
        else:
            if exposed_ever == 1:
                n_after_fu += 1
            rows.append(dict(base, tstart=0.0, tstop=fu, event=ev, exposed=0.0))

    cp = pd.DataFrame(rows)
    cp = cp.sort_values([PID, 'tstart']).reset_index(drop=True)
    check_counting_process(cp, work)
    meta = {
        'n_subjects': int(cp[PID].nunique()), 'n_rows': int(len(cp)),
        'n_two_row': n_two_row, 'n_exposed_at_zero': n_zero,
        'n_exposure_after_fu': n_after_fu,
        'exposed_person_time': float((cp.loc[cp['exposed'] == 1, 'tstop'] -
                                      cp.loc[cp['exposed'] == 1, 'tstart']).sum()),
        'unexposed_person_time': float((cp.loc[cp['exposed'] == 0, 'tstop'] -
                                        cp.loc[cp['exposed'] == 0, 'tstart']).sum()),
        'immortal_person_time': float((cp.loc[(cp['exposed'] == 0) & (cp['event'] == 0) &
                                              (cp[PID].isin(cp.loc[cp['exposed'] == 1, PID])),
                                              'tstop'] -
                                       cp.loc[(cp['exposed'] == 0) & (cp['event'] == 0) &
                                              (cp[PID].isin(cp.loc[cp['exposed'] == 1, PID])),
                                              'tstart']).sum()),
    }
    return cp, meta


def piecewise_cutpoints(work, config):
    """Cut points for the time-dependent HR. Explicit config wins; otherwise the
    tertiles of observed event times."""
    if config.get('tv_hr_cutpoints'):
        return sorted({float(c) for c in config['tv_hr_cutpoints']})
    ev_times = work.loc[work[EVENT] == 1, TIME]
    if len(ev_times) < 30:
        return []
    return sorted({float(x) for x in np.quantile(ev_times, [1 / 3, 2 / 3])})


def split_at_cutpoints(cp, cuts):
    """Split counting-process rows at the interval boundaries and tag the interval."""
    if not cuts:
        out = cp.copy()
        out['_interval'] = 0
        return out, [(0.0, float('inf'))]
    edges = [0.0] + list(cuts) + [float('inf')]
    intervals = list(zip(edges[:-1], edges[1:]))
    rows = []
    for _, r in cp.iterrows():
        for k, (lo, hi) in enumerate(intervals):
            a, b = max(r['tstart'], lo), min(r['tstop'], hi)
            if a >= b:
                continue
            new = r.to_dict()
            new['tstart'], new['tstop'] = float(a), float(b)
            new['event'] = int(r['event']) if b == r['tstop'] else 0
            new['_interval'] = k
            rows.append(new)
    out = pd.DataFrame(rows).sort_values([PID, 'tstart']).reset_index(drop=True)
    return out, intervals


def run_extended_cox(config, work, terms):
    term_cols = [t for t, _ in terms]
    cp, meta = build_counting_process(work, terms)
    if _BLOCKING:
        return None

    overall = fit_ctv(cp, term_cols, ['exposed'], context='Extended Cox (adjusted)')
    crude = fit_ctv(cp, [], ['exposed'], context='Extended Cox (crude)')
    if overall is None or crude is None:
        return None

    # ── Time-dependent HR: piecewise-constant exposure effect ────────────────
    cuts = piecewise_cutpoints(work, config)
    cp_split, intervals = split_at_cutpoints(cp, cuts)
    piecewise = []
    if len(intervals) > 1:
        targets = []
        for k in range(len(intervals)):
            col = f'exposed_p{k}'
            cp_split[col] = np.where(cp_split['_interval'] == k, cp_split['exposed'], 0.0)
            targets.append(col)
        pw = fit_ctv(cp_split, term_cols, targets, context='Extended Cox (piecewise HR)')
        if pw:
            for k, (lo, hi) in enumerate(intervals):
                res = pw.get(f'exposed_p{k}')
                sub = cp_split[cp_split['_interval'] == k]
                piecewise.append({
                    'interval': (lo, hi),
                    'events': int(sub['event'].sum()),
                    'result': res,
                })

    return {'cp': cp, 'meta': meta, 'overall': overall['exposed'],
            'crude': crude['exposed'], 'piecewise': piecewise, 'cutpoints': cuts}


def counting_process_sample(cp, work, n_exposed=3, n_unexposed=2):
    """Pick a few real patients and return their full expansion for manual checking."""
    exposed_ids = work.loc[work[EXPOSED] == 1, PID]
    # Prefer exposed subjects who died (2 rows + terminal event) — the interesting case.
    died = work.loc[(work[EXPOSED] == 1) & (work[EVENT] == 1), PID].tolist()
    alive = work.loc[(work[EXPOSED] == 1) & (work[EVENT] == 0), PID].tolist()
    picks = died[:max(1, n_exposed - 1)] + alive[:1]
    picks = picks[:n_exposed] or exposed_ids.head(n_exposed).tolist()
    unexp_died = work.loc[(work[EXPOSED] == 0) & (work[EVENT] == 1), PID].tolist()
    unexp_alive = work.loc[(work[EXPOSED] == 0) & (work[EVENT] == 0), PID].tolist()
    picks += (unexp_died[:1] + unexp_alive[:1])[:n_unexposed]
    sample = cp[cp[PID].isin(picks)].copy()
    info = work.set_index(PID)[[TIME, EVENT, EXPOSED, EXPTIME]]
    return sample, info


# ═══════════════════════════════════════════════════════════════════════════════
# Excel style constants (shared with the other pipeline sub-skills)
# ═══════════════════════════════════════════════════════════════════════════════

DARK_BLUE = "1F3864"
HEADER_BG = "2E5FA3"
SECTION_BG = "EBF2FA"
ALT_ROW = "F7FAFD"
SIG_COLOR = "C00000"
BORDER_COLOR = "BFBFBF"
FOOTNOTE_COLOR = "595959"
FLAG_BG = "FDEBD0"

THIN_BORDER = Border(
    left=Side(style='thin', color=BORDER_COLOR),
    right=Side(style='thin', color=BORDER_COLOR),
    top=Side(style='thin', color=BORDER_COLOR),
    bottom=Side(style='thin', color=BORDER_COLOR),
)

TITLE_FONT = Font(name='Arial', size=12, bold=True, color="FFFFFF")
SUBTITLE_FONT = Font(name='Arial', size=10, italic=True, color="FFFFFF")
HEADER_FONT = Font(name='Arial', size=10, bold=True, color="FFFFFF")
SECTION_FONT = Font(name='Arial', size=10, bold=True, color=DARK_BLUE)
DATA_FONT = Font(name='Arial', size=10)
SIG_FONT = Font(name='Arial', size=10, bold=True, color=SIG_COLOR)
FOOTNOTE_FONT = Font(name='Arial', size=8, italic=True, color=FOOTNOTE_COLOR)

TITLE_FILL = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type='solid')
HEADER_FILL = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type='solid')
SECTION_FILL = PatternFill(start_color=SECTION_BG, end_color=SECTION_BG, fill_type='solid')
ALT_FILL = PatternFill(start_color=ALT_ROW, end_color=ALT_ROW, fill_type='solid')
FLAG_FILL = PatternFill(start_color=FLAG_BG, end_color=FLAG_BG, fill_type='solid')

CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', indent=1, wrap_text=True)


def sheet_header(ws, title, subtitle, headers, widths):
    ncol = len(headers)
    for c, w in zip(range(1, ncol + 1), widths):
        ws.column_dimensions[get_column_letter(c)].width = w
    r = 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    cell = ws.cell(row=r, column=1, value=title)
    cell.font = TITLE_FONT
    cell.alignment = CENTER
    for c in range(1, ncol + 1):
        ws.cell(row=r, column=c).fill = TITLE_FILL
    r += 1
    if subtitle:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        cell = ws.cell(row=r, column=1, value=subtitle)
        cell.font = SUBTITLE_FONT
        cell.alignment = CENTER
        for c in range(1, ncol + 1):
            ws.cell(row=r, column=c).fill = TITLE_FILL
        r += 1
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER
    r += 1
    ws.freeze_panes = ws.cell(row=r, column=1)
    return r


def write_row(ws, r, values, alt=False, flag=False, bold_cols=(), sig_cols=()):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = SIG_FONT if c in sig_cols else DATA_FONT
        if c in bold_cols:
            cell.font = Font(name='Arial', size=10, bold=True)
        cell.border = THIN_BORDER
        cell.alignment = LEFT if c == 1 else CENTER
        if flag:
            cell.fill = FLAG_FILL
        elif alt:
            cell.fill = ALT_FILL
    return r + 1


def write_section(ws, r, ncol, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    cell = ws.cell(row=r, column=1, value=text)
    cell.font = SECTION_FONT
    cell.alignment = LEFT
    for c in range(1, ncol + 1):
        ws.cell(row=r, column=c).fill = SECTION_FILL
        ws.cell(row=r, column=c).border = THIN_BORDER
    return r + 1


def write_footnote(ws, r, ncol, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    cell = ws.cell(row=r, column=1, value=text)
    cell.font = FOOTNOTE_FONT
    cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    ws.row_dimensions[r].height = max(14, 12 * (1 + text.count('\n')))
    return r + 1


def fmt_p(p):
    if p is None or (isinstance(p, float) and math.isnan(p)):
        return "N/A"
    if p < 0.001:
        return "<0.001*"
    if p < 0.05:
        return f"{p:.3f}*"
    return f"{p:.3f}"


def fmt_hr(res):
    if not res:
        return "N/A"
    return f"{res['hr']:.2f} ({res['ci_low']:.2f}–{res['ci_high']:.2f})"


def fmt_num(x, nd=1):
    if x is None or (isinstance(x, float) and (math.isnan(x) or math.isinf(x))):
        return "N/A"
    return f"{x:.{nd}f}"


def significance_word(res):
    """Wording rule: a CI crossing 1 is 'no statistically
    significant difference', never 'no association'."""
    if not res:
        return "not estimable"
    if res['ci_low'] <= 1.0 <= res['ci_high']:
        return "no statistically significant difference"
    return "statistically significant" + (" (lower hazard)" if res['hr'] < 1 else " (higher hazard)")


# ═══════════════════════════════════════════════════════════════════════════════
# Excel — shared sheets
# ═══════════════════════════════════════════════════════════════════════════════

def write_estimand_sheet(wb, config, prep, terms, extra_rows=()):
    """Sheet 1 answers: what estimand is this, and on which population."""
    ws = wb.active
    ws.title = 'Estimand'
    headers = ['Item', 'Specification']
    r = sheet_header(ws, config['title'] or 'Immortal time bias — estimand',
                     config['subtitle'], headers, [42, 108])
    ncol = len(headers)
    exposed_value, exp_label, unexp_label = exposure_labels(config)
    unit = config['time_unit']
    mode = config['mode']

    if mode == 'landmark':
        estimand = ("Hazard ratio for the event comparing subjects who WERE ALREADY exposed at the "
                    "landmark time with those who were NOT, among subjects still alive and under "
                    "follow-up at that landmark, with follow-up re-started at the landmark and "
                    "the listed covariates held fixed. This is a conditional (landmark-population) "
                    "hazard ratio, not the effect of exposure in the whole enrolled cohort.")
        target_pop = ("Subjects event-free and still under observation at the landmark time. "
                      "Subjects with the event or censoring before the landmark are excluded by "
                      "design — see the Excluded_at_landmark sheet for who they were.")
    else:
        estimand = ("Hazard ratio for the event comparing exposed with unexposed person-time, "
                    "where exposure status changes at the moment exposure actually occurs "
                    "(counting-process / time-varying covariate), covariates held fixed. "
                    "Time before exposure is counted as UNEXPOSED person-time for the same subject.")
        target_pop = ("The whole analysis cohort; every subject contributes unexposed person-time "
                      "from time zero and exposed person-time only after exposure occurs.")

    r = write_section(ws, r, ncol, 'What is being estimated')
    r = write_row(ws, r, ['Analysis mode', mode], bold_cols=(1,))
    r = write_row(ws, r, ['Estimand', estimand], alt=True, bold_cols=(1,))
    r = write_row(ws, r, ['Target population', target_pop], bold_cols=(1,))
    r = write_row(ws, r, ['Effect measure', 'Cause-specific hazard ratio (Cox proportional hazards), '
                                            '95% Wald CI'], alt=True, bold_cols=(1,))
    r = write_row(ws, r, ['Why this analysis is needed',
                          'Exposure is defined by information that only becomes available after '
                          'time zero, so exposed subjects must survive event-free until exposure. '
                          'That immortal time makes exposure look protective if it is handed to a '
                          'Cox model as a baseline covariate.'], bold_cols=(1,))

    r += 1
    r = write_section(ws, r, ncol, 'Exposure definition')
    r = write_row(ws, r, ['Exposure variable',
                          f"{config['grouping']['column']} == {exposed_value!r} -> {exp_label}"],
                  bold_cols=(1,))
    r = write_row(ws, r, ['Reference arm', unexp_label], alt=True, bold_cols=(1,))
    r = write_row(ws, r, ['Exposure time variable', config['exposure_time_column']], bold_cols=(1,))
    r = write_row(ws, r, ['Median time to exposure (exposed only)',
                          f"{fmt_num(prep['median_exposure_time'])} {unit}"],
                  alt=True, bold_cols=(1,))
    r = write_row(ws, r, ['Immortal person-time in the exposed arm',
                          f"{fmt_num(prep['immortal_person_time'], 0)} {unit} "
                          f"(person-time before exposure that the naive model credits to the "
                          f"exposed arm)"], bold_cols=(1,))

    if mode == 'landmark':
        r += 1
        r = write_section(ws, r, ncol, 'Landmark times (pre-specified)')
        r = write_row(ws, r, ['Landmark times',
                              ', '.join(f"{t:g}" for t in config['landmark_times']) + f" ({unit})"],
                      bold_cols=(1,))
        r = write_row(ws, r, ['Rationale for these times',
                              config['landmark_rationale'] or
                              '[NOT SUPPLIED IN CONFIG — record the clinical rationale in the '
                              'decision log before submission]'],
                      alt=True, bold_cols=(1,), flag=not config['landmark_rationale'])
        r = write_row(ws, r, ['Reporting rule',
                              'Every landmark time in the config is reported. Selecting the most '
                              'favourable landmark after seeing the results is not permitted '
                              '(advanced_analyses.md section 7).'], bold_cols=(1,))
        r = write_row(ws, r, ['Late exposure handling',
                              'Subjects exposed AFTER a given landmark are classified as unexposed '
                              'at that landmark (status is judged as of the landmark).'],
                      alt=True, bold_cols=(1,))

    r += 1
    r = write_section(ws, r, ncol, 'Analysis population')
    r = write_row(ws, r, ['Rows in source data after filters', prep['n_raw']], bold_cols=(1,))
    r = write_row(ws, r, ['Dropped by listwise deletion', prep['n_listwise_dropped']],
                  alt=True, bold_cols=(1,))
    r = write_row(ws, r, ['Dropped for non-positive follow-up', prep['n_nonpositive_time']],
                  bold_cols=(1,))
    r = write_row(ws, r, ['Analysis cohort (n)', prep['n_analysis']], alt=True, bold_cols=(1,))
    r = write_row(ws, r, [f'Exposed ({exp_label})', prep['n_exposed']], bold_cols=(1,))
    r = write_row(ws, r, [f'Unexposed ({unexp_label})', prep['n_unexposed']],
                  alt=True, bold_cols=(1,))
    r = write_row(ws, r, ['Events', prep['n_events']], bold_cols=(1,))

    r += 1
    r = write_section(ws, r, ncol, 'Adjustment set')
    if terms:
        for i, (_, lab) in enumerate(terms):
            r = write_row(ws, r, ['Covariate', lab], alt=(i % 2 == 1))
    else:
        r = write_row(ws, r, ['Covariate', 'None — unadjusted models only'], flag=True)

    for label, value in extra_rows:
        r = write_row(ws, r, [label, value], bold_cols=(1,))

    r += 1
    r = write_footnote(ws, r, ncol,
                       'Wording: an interval containing 1.0 is reported as "no statistically '
                       'significant difference", which is not the same as evidence of no effect.')
    return ws


def write_naive_sheet(wb, config, naive, corrected_rows, prep):
    """Naive baseline-covariate Cox next to the corrected estimates — the gap is the bias."""
    ws = wb.create_sheet('Naive_comparison')
    headers = ['Model', 'Immortal time handled?', 'n', 'Events', 'HR (95% CI)', 'p', 'Interpretation']
    r = sheet_header(ws, 'Naive vs immortal-time-corrected estimates',
                     'The distance between the naive row and the corrected rows is the '
                     'immortal time bias', headers, [40, 20, 10, 10, 22, 12, 46])
    ncol = len(headers)

    r = write_section(ws, r, ncol, 'Naive model — exposure as a fixed baseline covariate (BIASED)')
    for label, res in (('Naive Cox, crude', naive['crude']),
                       ('Naive Cox, adjusted', naive['adjusted'])):
        r = write_row(ws, r, [label, 'No', res['n'] if res else 'N/A',
                              res['events'] if res else 'N/A', fmt_hr(res),
                              fmt_p(res['p']) if res else 'N/A', significance_word(res)],
                      flag=True)

    r += 1
    r = write_section(ws, r, ncol, 'Corrected models')
    for i, row in enumerate(corrected_rows):
        r = write_row(ws, r, row, alt=(i % 2 == 1))

    r += 1
    r = write_section(ws, r, ncol, 'Size of the bias')
    ref = naive['adjusted'] or naive['crude']
    for label, res in corrected_rows_estimates(corrected_rows):
        if ref and res:
            ratio = ref['hr'] / res['hr']
            r = write_row(ws, r, [f'Naive adjusted HR / {label}', '', '', '',
                                  f"{ratio:.2f}", '',
                                  f"log-HR difference {ref['coef'] - res['coef']:+.3f}"])
    r = write_row(ws, r, ['Immortal person-time credited to the exposed arm by the naive model',
                          '', '', '',
                          f"{fmt_num(prep['immortal_person_time'], 0)} {config['time_unit']}", '', ''])

    r += 1
    r = write_footnote(ws, r, ncol,
                       'The naive rows are printed for contrast only. They are not a valid estimate '
                       'of the exposure effect and must not be reported as the primary result.')
    return ws


_CORRECTED_ESTIMATES = []


def corrected_rows_estimates(_rows):
    return _CORRECTED_ESTIMATES


def write_method_sheet(wb, config, prep, terms, method_rows):
    ws = wb.create_sheet('Method')
    headers = ['Item', 'Detail']
    r = sheet_header(ws, 'Method and methodological decisions',
                     'Every default below is a decision that belongs in decision_log.md '
                     'with a permanent ID', headers, [40, 110])
    ncol = len(headers)
    import lifelines
    for i, (k, v) in enumerate(method_rows):
        if v is None:
            r = write_section(ws, r, ncol, k)
        else:
            r = write_row(ws, r, [k, v], alt=(i % 2 == 1), bold_cols=(1,))
    r += 1
    r = write_section(ws, r, ncol, 'Software')
    r = write_row(ws, r, ['Python', sys.version.split()[0]])
    r = write_row(ws, r, ['lifelines', lifelines.__version__], alt=True)
    r = write_row(ws, r, ['pandas / numpy', f"{pd.__version__} / {np.__version__}"])
    r += 1
    r = write_footnote(ws, r, ncol,
                       'Reporting rule: when both a landmark analysis and an extended Cox model are '
                       'run, both must be reported, together with an explanation of any difference '
                       'between them (advanced_analyses.md section 7).')
    return ws


# ═══════════════════════════════════════════════════════════════════════════════
# Excel — landmark mode
# ═══════════════════════════════════════════════════════════════════════════════

def write_landmark_excel(config, results, excluded_rows, naive, prep, terms):
    unit = config['time_unit']
    _, exp_label, unexp_label = exposure_labels(config)
    thr = config['thresholds']

    wb = Workbook()
    write_estimand_sheet(wb, config, prep, terms)

    # ── Landmark_results ────────────────────────────────────────────────────
    ws = wb.create_sheet('Landmark_results')
    headers = [f'Landmark ({unit})', 'n analyzed', 'n excluded', '% excluded',
               f'n {exp_label}', f'n {unexp_label}',
               f'Events {exp_label}', f'Events {unexp_label}',
               'Crude HR (95% CI)', 'Adjusted HR (95% CI)', 'p (adjusted)', 'Note']
    r = sheet_header(ws, 'Landmark analysis — all pre-specified landmark times',
                     config['subtitle'] or
                     'Every landmark time in the config is reported; none may be selected post hoc',
                     headers, [15, 12, 12, 11, 12, 14, 14, 16, 22, 22, 13, 40])
    ncol = len(headers)
    for i, res in enumerate(results):
        est = res['adjusted']
        notes = []
        if res['pct_excluded'] > thr['max_excluded_pct']:
            notes.append(f">{thr['max_excluded_pct']:g}% excluded")
        if min(res['events_exposed'], res['events_unexposed']) < thr['min_events_per_group']:
            notes.append(f"<{thr['min_events_per_group']} events in an arm")
        if res['n_late_exposure']:
            notes.append(f"{res['n_late_exposure']} exposed after this landmark, "
                         f"analyzed as unexposed")
        flag = bool(notes)
        r = write_row(ws, r, [
            f"{res['landmark']:g}", res['n_analyzed'], res['n_excluded'],
            fmt_num(res['pct_excluded']), res['n_exposed'], res['n_unexposed'],
            res['events_exposed'], res['events_unexposed'],
            fmt_hr(res['crude']), fmt_hr(est),
            fmt_p(est['p']) if est else 'N/A',
            '; '.join(notes) or significance_word(est),
        ], alt=(i % 2 == 1), flag=flag)
    r += 1
    r = write_footnote(ws, r, ncol,
                       'Follow-up is re-started at each landmark. Exposure status is judged as of '
                       'the landmark: subjects exposed later are counted as unexposed at that '
                       'landmark. Adjusted models include the covariates listed on the Estimand '
                       'sheet. An interval containing 1.0 means no statistically significant '
                       'difference at that landmark.')

    # ── Excluded_at_landmark ────────────────────────────────────────────────
    ws = wb.create_sheet('Excluded_at_landmark')
    headers = [f'Landmark ({unit})', 'n excluded', '% of cohort', 'Events among excluded',
               'Event rate among excluded (%)', 'Censored among excluded',
               f'Median follow-up of excluded ({unit})', 'Excluded who were ever exposed',
               'n retained', 'Events among retained', 'Event rate among retained (%)',
               'Rate ratio excluded / retained']
    r = sheet_header(ws, 'Who the landmark drops — and how sick they were',
                     'The excluded subjects are not a random subset of the cohort',
                     headers, [15, 12, 12, 18, 20, 18, 20, 20, 12, 18, 20, 18])
    ncol = len(headers)
    for i, e in enumerate(excluded_rows):
        rr = (e['excl_event_rate'] / e['kept_event_rate']
              if e['kept_event_rate'] else float('nan'))
        r = write_row(ws, r, [
            f"{e['landmark']:g}", e['n_excluded'], fmt_num(e['pct_excluded']),
            e['excl_events'], fmt_num(e['excl_event_rate']), e['excl_censored'],
            fmt_num(e['excl_median_fu']), e['excl_would_be_exposed'],
            e['n_kept'], e['kept_events'], fmt_num(e['kept_event_rate']),
            fmt_num(rr, 2),
        ], alt=(i % 2 == 1), flag=e['pct_excluded'] > config['thresholds']['max_excluded_pct'])
    r += 1
    r = write_footnote(ws, r, ncol,
                       'A rate ratio well above 1 means the landmark removed the highest-risk part '
                       'of the cohort; the landmark estimate then applies to a healthier population '
                       'than the one enrolled, and generalisability must be discussed.')

    # ── Naive_comparison ────────────────────────────────────────────────────
    corrected, _CORRECTED_ESTIMATES[:] = [], []
    for res in results:
        est = res['adjusted'] or res['crude']
        label = f"Landmark {res['landmark']:g} {unit}, adjusted"
        corrected.append([label, 'Yes (landmark)', res['n_analyzed'],
                          res['events_exposed'] + res['events_unexposed'],
                          fmt_hr(est), fmt_p(est['p']) if est else 'N/A',
                          significance_word(est)])
        _CORRECTED_ESTIMATES.append((label, est))
    write_naive_sheet(wb, config, naive, corrected, prep)

    # ── Method ──────────────────────────────────────────────────────────────
    method_rows = [
        ('Landmark analysis', None),
        ('Landmark times', ', '.join(f"{t:g}" for t in config['landmark_times']) + f' {unit}'),
        ('Rationale', config['landmark_rationale'] or '[NOT SUPPLIED — record before submission]'),
        ('Exclusion rule at each landmark',
         'Subjects with the event or with censoring at or before the landmark are excluded '
         '(follow-up time must be strictly greater than the landmark).'),
        ('Exposure classification',
         f"Exposed at the landmark if {config['exposure_time_column']} <= landmark; "
         f"exposure occurring later is classified as unexposed at that landmark."),
        ('Time origin', 'Follow-up time is re-started at the landmark (time - landmark).'),
        ('Model', 'Cox proportional hazards (lifelines CoxPHFitter; Efron tie handling, '
                  'Breslow baseline hazard), 95% Wald CI.'),
        ('Missing data', f"Listwise deletion of model variables "
                         f"({prep['n_listwise_dropped']} row(s) dropped)."),
        ('QC thresholds', f"exclusion warning at >{config['thresholds']['max_excluded_pct']:g}% of "
                          f"the cohort; event warning at <{config['thresholds']['min_events_per_group']} "
                          f"events per arm."),
        ('Naive comparator', None),
        ('Naive model', 'Cox with ever-exposed as a fixed baseline covariate over the full '
                        'follow-up. Reported only to quantify the immortal time bias.'),
    ]
    write_method_sheet(wb, config, prep, terms, method_rows)

    out = config['output_file']
    os.makedirs(os.path.dirname(os.path.abspath(out)) or '.', exist_ok=True)
    wb.save(out)
    print(f"  Excel saved: {out}")
    return out


# ═══════════════════════════════════════════════════════════════════════════════
# Excel — extended Cox mode
# ═══════════════════════════════════════════════════════════════════════════════

def write_extended_excel(config, ext, naive, prep, terms, work):
    unit = config['time_unit']
    _, exp_label, unexp_label = exposure_labels(config)
    meta = ext['meta']

    wb = Workbook()
    write_estimand_sheet(wb, config, prep, terms, extra_rows=[
        ('Counting-process rows', f"{meta['n_rows']} rows for {meta['n_subjects']} subjects"),
        ('Exposed person-time', f"{fmt_num(meta['exposed_person_time'], 0)} {unit}"),
        ('Unexposed person-time', f"{fmt_num(meta['unexposed_person_time'], 0)} {unit}"),
    ])

    # ── Extended_cox_results ────────────────────────────────────────────────
    ws = wb.create_sheet('Extended_cox_results')
    headers = ['Model', 'Subjects', 'Rows', 'Events', 'HR (95% CI)', 'p', 'Interpretation']
    r = sheet_header(ws, 'Extended Cox with time-varying exposure',
                     'Exposure switches on at the moment it occurs; time before that is '
                     'unexposed person-time for the same subject',
                     headers, [38, 12, 10, 10, 24, 12, 46])
    ncol = len(headers)
    for i, (label, res) in enumerate((('Time-varying exposure, crude', ext['crude']),
                                      ('Time-varying exposure, adjusted', ext['overall']))):
        r = write_row(ws, r, [label,
                              res['n_subjects'] if res else 'N/A',
                              res['n_rows'] if res else 'N/A',
                              res['events'] if res else 'N/A',
                              fmt_hr(res), fmt_p(res['p']) if res else 'N/A',
                              significance_word(res)], alt=(i % 2 == 1))

    if ext['piecewise']:
        r += 1
        r = write_section(ws, r, ncol, 'Time-dependent exposure effect (piecewise-constant HR)')
        for i, pw in enumerate(ext['piecewise']):
            lo, hi = pw['interval']
            hi_txt = '∞' if math.isinf(hi) else f"{hi:g}"
            res = pw['result']
            r = write_row(ws, r, [f"Interval {lo:g}–{hi_txt} {unit}", '', '',
                                  pw['events'], fmt_hr(res),
                                  fmt_p(res['p']) if res else 'N/A',
                                  significance_word(res)], alt=(i % 2 == 1))
        r = write_footnote(ws, r, ncol,
                           'Piecewise HRs come from one model in which the exposure indicator is '
                           'interacted with the follow-up interval; cut points are listed on the '
                           'Method sheet. They are a check on proportional hazards, not a set of '
                           'independent analyses to choose from.')
    r += 1
    r = write_footnote(ws, r, ncol,
                       'A confidence interval containing 1.0 is reported as no statistically '
                       'significant difference.')

    # ── Counting_process_sample ─────────────────────────────────────────────
    sample, info = counting_process_sample(ext['cp'], work)
    ws = wb.create_sheet('Counting_process_sample')
    headers = ['Subject id', 'Arm', f'Exposure time ({unit})', f'Follow-up ({unit})',
               'Event (subject)', 'tstart', 'tstop', 'event (row)', 'exposed (row)',
               'Why this row looks like this']
    r = sheet_header(ws, 'Counting-process expansion — sample subjects for manual checking',
                     'Every row is one at-risk interval; a subject with 2 rows switched exposure '
                     'during follow-up', headers, [11, 16, 18, 16, 14, 10, 10, 12, 12, 62])
    ncol = len(headers)
    prev_pid = None
    alt = False
    for _, row in sample.iterrows():
        pid = int(row[PID])
        if pid != prev_pid:
            alt = not alt
            prev_pid = pid
        s = info.loc[pid]
        arm = exp_label if s[EXPOSED] == 1 else unexp_label
        etime = s[EXPTIME]
        if s[EXPOSED] == 1 and pd.notna(etime) and row['exposed'] == 0:
            why = (f"Subject was exposed at {float(etime):g}, but this interval runs before that, "
                   f"so it is UNEXPOSED person-time. This is exactly the immortal time the naive "
                   f"model would have credited to the exposed arm.")
        elif row['exposed'] == 1:
            why = (f"From exposure at {float(etime):g} to end of follow-up: exposed person-time; "
                   f"the terminal event, if any, is attributed here.")
        else:
            why = ("Never exposed: one interval from time zero to end of follow-up, carrying the "
                   "subject's event indicator.")
        r = write_row(ws, r, [pid, arm,
                              fmt_num(float(etime), 1) if pd.notna(etime) else '—',
                              fmt_num(float(s[TIME]), 1), int(s[EVENT]),
                              fmt_num(float(row['tstart']), 1), fmt_num(float(row['tstop']), 1),
                              int(row['event']), int(row['exposed']), why], alt=alt)
    r += 1
    r = write_footnote(ws, r, ncol,
                       'Check by hand: intervals per subject must be contiguous (tstop of one row '
                       'equals tstart of the next), tstart < tstop everywhere, at most one row per '
                       'subject carries event = 1, and that row must be the last one.')

    # ── Naive_comparison ────────────────────────────────────────────────────
    corrected, _CORRECTED_ESTIMATES[:] = [], []
    for label, res in (('Extended Cox, crude', ext['crude']),
                       ('Extended Cox, adjusted', ext['overall'])):
        corrected.append([label, 'Yes (time-varying)',
                          res['n_subjects'] if res else 'N/A',
                          res['events'] if res else 'N/A', fmt_hr(res),
                          fmt_p(res['p']) if res else 'N/A', significance_word(res)])
        _CORRECTED_ESTIMATES.append((label, res))
    write_naive_sheet(wb, config, naive, corrected, prep)

    # ── Method ──────────────────────────────────────────────────────────────
    cuts = ext['cutpoints']
    method_rows = [
        ('Extended Cox (time-varying exposure)', None),
        ('Data structure', 'Counting process: one row per at-risk interval with tstart, tstop, '
                           'event, exposure. Never exposed -> 1 row; exposed during follow-up -> '
                           '2 rows split at the exposure time.'),
        ('Rows / subjects', f"{meta['n_rows']} rows for {meta['n_subjects']} subjects "
                            f"({meta['n_two_row']} subjects split into two rows)"),
        ('Exposure at time zero', f"{meta['n_exposed_at_zero']} subject(s) with exposure time <= 0 "
                                  f"were treated as exposed from time zero"),
        ('Exposure at or after end of follow-up',
         f"{meta['n_exposure_after_fu']} subject(s) whose exposure time was not before the end of "
         f"follow-up contribute unexposed person-time only"),
        ('Model', 'lifelines CoxTimeVaryingFitter (Newton-Raphson, Efron tie handling), '
                  '95% Wald CI.'),
        ('Time-dependent HR',
         ('Piecewise-constant exposure effect; cut points ' +
          (', '.join(f"{c:g}" for c in cuts) + f' {unit}' if cuts else 'not applied (too few events)')
          + '. Cut points are a pre-specified reporting device, not selected on the results.')),
        ('Missing data', f"Listwise deletion of model variables "
                         f"({prep['n_listwise_dropped']} row(s) dropped)."),
        ('Naive comparator', None),
        ('Naive model', 'Cox with ever-exposed as a fixed baseline covariate over the full '
                        'follow-up. Reported only to quantify the immortal time bias.'),
    ]
    write_method_sheet(wb, config, prep, terms, method_rows)

    out = config['output_file']
    os.makedirs(os.path.dirname(os.path.abspath(out)) or '.', exist_ok=True)
    wb.save(out)
    print(f"  Excel saved: {out}")
    return out


# ═══════════════════════════════════════════════════════════════════════════════
# Plots
# ═══════════════════════════════════════════════════════════════════════════════

PALETTE = ['#C0392B', '#2E5FA3']


def plot_landmark(config, results):
    """One KM panel per landmark time, follow-up re-started at the landmark."""
    from lifelines import KaplanMeierFitter

    _, exp_label, unexp_label = exposure_labels(config)
    unit = config['time_unit']
    n = len(results)
    ncols = min(3, n)
    nrows = int(math.ceil(n / ncols))
    fig, axes = plt.subplots(nrows, ncols, figsize=(5.0 * ncols, 4.2 * nrows), squeeze=False)

    for idx, res in enumerate(results):
        ax = axes[idx // ncols][idx % ncols]
        kept = res['kept']
        for val, label, color in ((1.0, exp_label, PALETTE[0]), (0.0, unexp_label, PALETTE[1])):
            sub = kept[kept[EXPOSED_LM] == val]
            if len(sub) == 0:
                continue
            kmf = KaplanMeierFitter()
            kmf.fit(sub['_lm_time'], sub[EVENT], label=f"{label} (n={len(sub)})")
            sf = kmf.survival_function_
            ax.step(sf.index, sf.iloc[:, 0], where='post', color=color, linewidth=1.8,
                    label=f"{label} (n={len(sub)})")
            ci = kmf.confidence_interval_survival_function_
            ax.fill_between(ci.index, ci.iloc[:, 0], ci.iloc[:, 1], step='post',
                            alpha=0.15, color=color)
        est = res['adjusted'] or res['crude']
        txt = (f"Adjusted HR {est['hr']:.2f} ({est['ci_low']:.2f}–{est['ci_high']:.2f})\n"
               f"p = {fmt_p(est['p'])}" if est else 'HR not estimable')
        ax.text(0.98, 0.04, txt, transform=ax.transAxes, fontsize=8, ha='right', va='bottom',
                bbox=dict(boxstyle='round,pad=0.3', facecolor='white', edgecolor='gray', alpha=0.85))
        ax.set_title(f"Landmark {res['landmark']:g} {unit}\n"
                     f"n = {res['n_analyzed']} ({res['pct_excluded']:.1f}% excluded)",
                     fontsize=10, fontweight='bold')
        ax.set_xlabel(f'Time since landmark ({unit})', fontsize=9)
        ax.set_ylabel('Survival probability', fontsize=9)
        ax.set_ylim(0, 1.02)
        ax.grid(axis='y', alpha=0.3, linestyle='--')
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.legend(fontsize=8, loc='lower left', framealpha=0.9)

    for j in range(n, nrows * ncols):
        axes[j // ncols][j % ncols].axis('off')

    if config['title']:
        fig.suptitle(config['title'], fontsize=13, fontweight='bold')
        fig.tight_layout(rect=(0, 0, 1, 0.96))
    else:
        fig.tight_layout()

    out = config['plot_file']
    os.makedirs(os.path.dirname(os.path.abspath(out)) or '.', exist_ok=True)
    fig.savefig(out, dpi=300, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    print(f"  Plot saved: {out}")
    return out


def plot_extended(config, ext, naive):
    """Time-dependent HR: piecewise-constant estimate with CI, against the naive HR."""
    unit = config['time_unit']
    fig, ax = plt.subplots(figsize=(8.0, 5.0))

    pw = ext['piecewise']
    if pw:
        max_t = max(p['interval'][1] for p in pw if not math.isinf(p['interval'][1])) \
            if any(not math.isinf(p['interval'][1]) for p in pw) else 1.0
        for p in pw:
            lo, hi = p['interval']
            hi = hi if not math.isinf(hi) else max_t * 1.35
            res = p['result']
            if not res:
                continue
            ax.hlines(res['hr'], lo, hi, color=PALETTE[1], linewidth=2.2, zorder=3)
            ax.fill_between([lo, hi], [res['ci_low']] * 2, [res['ci_high']] * 2,
                            color=PALETTE[1], alpha=0.16, zorder=1)
            ax.text((lo + hi) / 2, res['hr'], f" {res['hr']:.2f}", fontsize=8,
                    va='bottom', ha='center', color=PALETTE[1])
        xmax = max_t * 1.35
    else:
        xmax = 1.0

    overall = ext['overall']
    if overall:
        ax.axhline(overall['hr'], color='#1F3864', linestyle='-', linewidth=1.6,
                   label=f"Extended Cox, adjusted: HR {overall['hr']:.2f} "
                         f"({overall['ci_low']:.2f}–{overall['ci_high']:.2f})", zorder=4)
    nv = naive['adjusted'] or naive['crude']
    if nv:
        ax.axhline(nv['hr'], color=PALETTE[0], linestyle='--', linewidth=1.6,
                   label=f"Naive baseline-covariate Cox (biased): HR {nv['hr']:.2f}", zorder=4)
    ax.axhline(1.0, color='gray', linestyle=':', linewidth=1.2, label='No effect (HR = 1)', zorder=2)

    ax.set_xlim(0, xmax)
    ax.set_yscale('log')
    ax.set_xlabel(f'Time since time zero ({unit})', fontsize=11, fontweight='bold')
    ax.set_ylabel('Hazard ratio (log scale)', fontsize=11, fontweight='bold')
    ax.set_title(config['title'] or 'Time-varying exposure effect', fontsize=13, fontweight='bold')
    ax.grid(axis='y', alpha=0.3, linestyle='--')
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.legend(fontsize=8, loc='best', framealpha=0.9)
    fig.tight_layout()

    out = config['plot_file']
    os.makedirs(os.path.dirname(os.path.abspath(out)) or '.', exist_ok=True)
    fig.savefig(out, dpi=300, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    print(f"  Plot saved: {out}")
    return out


# ═══════════════════════════════════════════════════════════════════════════════
# Preview (--check-missing): never fits a model, never writes a file
# ═══════════════════════════════════════════════════════════════════════════════

def preview(config, df):
    cs = config['cox_settings']
    time_col, event_col = cs['time_column'], cs['event_column']
    grp_col = config['grouping']['column']
    exp_time_col = config['exposure_time_column']
    exposed_value, exp_label, unexp_label = exposure_labels(config)
    variables = covariate_variables(config)

    print("=" * 96)
    print(f"IMMORTAL TIME PREVIEW (--check-missing)  |  mode = {config['mode']}")
    print("=" * 96)
    print(f"Data file : {config['data_file']}")
    print(f"Rows after filters: {len(df)}")
    print()
    print(f"{'Column':<32} {'Role':<18} {'Present':<9} {'Missing n (%)':>18}")
    print("-" * 96)

    def line(col, role):
        if col not in df.columns:
            print(f"  {col:<30} {role:<18} {'NO':<9} {'COLUMN NOT FOUND':>18}")
            return False
        n_miss = df[col].isna().sum()
        pct = 100.0 * n_miss / len(df) if len(df) else 0.0
        print(f"  {col:<30} {role:<18} {'yes':<9} {n_miss:>8} ({pct:>5.1f}%)")
        return True

    ok = True
    ok &= line(time_col, 'follow-up time')
    ok &= line(event_col, 'event')
    ok &= line(grp_col, 'exposure group')
    ok &= line(exp_time_col, 'exposure time')
    for v in variables:
        line(v['column'], f"covariate/{v['type']}")

    print()
    if exp_time_col in df.columns and grp_col in df.columns:
        grp = pd.to_numeric(df[grp_col], errors='coerce') \
            if isinstance(exposed_value, (int, float)) and not isinstance(exposed_value, bool) \
            else df[grp_col]
        exposed = grp == exposed_value
        et = pd.to_numeric(df[exp_time_col], errors='coerce')
        n_exp = int(exposed.sum())
        n_missing_et = int(et[exposed].isna().sum())
        print(f"Exposed arm ({exp_label}): n = {n_exp}")
        print(f"Reference arm ({unexp_label}): n = {int((~exposed).sum())}")
        print(f"Exposed subjects with a missing exposure time: {n_missing_et}"
              f"{'   <-- would be BLOCKING in a full run' if n_missing_et else ''}")
        if n_exp and not et[exposed].isna().all():
            print(f"Exposure time among exposed: median {et[exposed].median():.1f}, "
                  f"range {et[exposed].min():.1f}–{et[exposed].max():.1f} {config['time_unit']}")

    if config['mode'] == 'landmark' and time_col in df.columns:
        t_all = pd.to_numeric(df[time_col], errors='coerce')
        e_all = pd.to_numeric(df[event_col], errors='coerce') if event_col in df.columns else None
        print()
        print("Landmark feasibility preview (before listwise deletion):")
        print(f"  {'Landmark':<12} {'retained':>10} {'excluded':>10} {'% excluded':>12} "
              f"{'events among excluded':>24}")
        for t in config['landmark_times']:
            keep = (t_all > t).sum()
            excl = int(len(df) - keep)
            ev = int(e_all[(t_all <= t)].sum()) if e_all is not None else -1
            print(f"  {t:<12g} {int(keep):>10} {excl:>10} "
                  f"{100.0 * excl / len(df):>11.1f}% {ev:>24}")

    print()
    print("[preview] No model was fitted and no Excel/PNG was written (--check-missing).")
    if not ok:
        print("[preview] At least one required column is missing; a full run would exit 3.")
    print("=" * 96)


# ═══════════════════════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════════════════════

def _apply_academic_style():
    """academic-style integration (Excel-only). Consumes --style=<journal> from
    sys.argv. No-op when the skill is not installed or the flag is absent."""
    from pathlib import Path as _P
    _as = _P(__file__).resolve().parents[2] / "academic-style"
    if not _as.exists():
        return None
    if str(_as) not in sys.path:
        sys.path.insert(0, str(_as))
    try:
        from core.biostats_integration import apply_style_from_argv
    except ImportError:
        return None
    try:
        return apply_style_from_argv(sys.argv)
    except Exception:
        return None


def main():
    _AS_STYLE = _apply_academic_style()
    if _AS_STYLE:
        print(f"  Applied {_AS_STYLE} journal style (Excel formatting hints)")

    if len(sys.argv) < 2:
        print("Usage: python immortal_time.py config.json [--check-missing] [--style=<journal>]")
        sys.exit(1)

    config = load_config(sys.argv[1])
    check_only = '--check-missing' in sys.argv
    df = load_data(config)

    if check_only:
        preview(config, df)
        sys.exit(0)

    print(f"Mode: {config['mode']}")
    print(f"Data loaded: n={len(df)}")

    work, terms, prep = prepare(config, df)
    print(f"Analysis cohort: n={prep['n_analysis']} "
          f"(exposed {prep['n_exposed']}, unexposed {prep['n_unexposed']}, "
          f"events {prep['n_events']})")

    naive = run_naive(work, terms)
    nv = naive['adjusted'] or naive['crude']
    if nv:
        print(f"  Naive (biased) adjusted HR: {nv['hr']:.3f} "
              f"({nv['ci_low']:.3f}–{nv['ci_high']:.3f}), p={nv['p']:.4g}")

    if config['mode'] == 'landmark':
        results, excluded_rows = run_landmark(config, work, terms)
        for res in results:
            est = res['adjusted'] or res['crude']
            if est:
                print(f"  Landmark {res['landmark']:g}: adjusted HR {est['hr']:.3f} "
                      f"({est['ci_low']:.3f}–{est['ci_high']:.3f}), p={est['p']:.4g}, "
                      f"n={res['n_analyzed']} ({res['pct_excluded']:.1f}% excluded)")
        if not _BLOCKING:
            write_landmark_excel(config, results, excluded_rows, naive, prep, terms)
            plot_landmark(config, results)
    else:
        ext = run_extended_cox(config, work, terms)
        if ext and not _BLOCKING:
            o = ext['overall']
            print(f"  Extended Cox adjusted HR: {o['hr']:.3f} "
                  f"({o['ci_low']:.3f}–{o['ci_high']:.3f}), p={o['p']:.4g}")
            write_extended_excel(config, ext, naive, prep, terms, work)
            plot_extended(config, ext, naive)

    print("\nDone!")
    n_block = len(_BLOCKING)
    n_warn = len(_QC_MESSAGES) - n_block
    print(f"  QC: {n_block} blocking, {n_warn} warning(s) written to stderr")
    if n_block:
        print("  Blocking findings present — output must not be delivered as is.")
        sys.exit(3)
    sys.exit(0)


if __name__ == '__main__':
    main()
