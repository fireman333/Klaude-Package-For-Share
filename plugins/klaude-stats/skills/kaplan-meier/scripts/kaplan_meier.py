#!/usr/bin/env python3
"""
Kaplan-Meier Survival Curve Generator
Generates publication-ready KM survival curves with log-rank test,
number at risk table, and formatted Excel summary.

Outputs: Excel (.xlsx) with survival summary + log-rank test, KM plot (.png).

Usage:
    python kaplan_meier.py config.json                  # Full analysis
    python kaplan_meier.py config.json --check-missing  # Missing data report only
"""

import sys
import json
import os
import operator
import pandas as pd
import numpy as np
# P-3 batch (2026-05-07 partitioned-flame audit): removed `import warnings;
# warnings.filterwarnings('ignore')` — global suppression hid lifelines / scipy /
# pandas convergence and data-import warnings (silent error pattern).
# 

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
from matplotlib.lines import Line2D
import matplotlib.gridspec as gridspec

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from lifelines import KaplanMeierFitter
from lifelines.statistics import logrank_test, multivariate_logrank_test


# ═══════════════════════════════════════════════════════════════════════════════
# Configuration
# ═══════════════════════════════════════════════════════════════════════════════

DEFAULT_COLORS = [
    '#E74C3C', '#3498DB', '#2ECC71', '#9B59B6', '#F39C12',
    '#1ABC9C', '#E67E22', '#34495E', '#E91E63', '#00BCD4'
]


def load_config(path):
    with open(path, 'r', encoding='utf-8') as f:
        config = json.load(f)
    config.setdefault('filters', [])
    config.setdefault('derived_variables', [])
    config.setdefault('title', 'Kaplan-Meier Survival Analysis')
    config.setdefault('subtitle', '')
    config.setdefault('abbreviations', '')
    config.setdefault('missing_data_handling', 'listwise')
    config.setdefault('time_scale', None)

    # Plot settings defaults
    ps = config.setdefault('plot_settings', {})
    ps.setdefault('confidence_interval', True)
    ps.setdefault('censor_marks', True)
    ps.setdefault('colors', None)
    ps.setdefault('x_max', None)
    ps.setdefault('x_tick_interval', None)
    ps.setdefault('y_min', 0.0)
    ps.setdefault('number_at_risk', True)
    ps.setdefault('risk_table_times', None)
    ps.setdefault('legend_position', 'best')
    ps.setdefault('figsize', [8, 6])
    ps.setdefault('dpi', 300)

    return config


# ═══════════════════════════════════════════════════════════════════════════════
# Data Loading & Preprocessing
# ═══════════════════════════════════════════════════════════════════════════════

# P-2 batch (2026-05-07 partitioned-flame audit): split equality vs numeric ops.
# Equality (==, !=) must compare on original dtype, otherwise string filters like
# `sex == "F"` silently coerce to NaN and drop every row.
EQ_OPS = {'==': operator.eq, '!=': operator.ne}
NUMERIC_OPS = {'>': operator.gt, '>=': operator.ge, '<': operator.lt, '<=': operator.le}
OPS = {**EQ_OPS, **NUMERIC_OPS}


def load_data(config):
    path = config['data_file']
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xls'):
        df = pd.read_excel(path)
    elif ext == '.csv':
        df = pd.read_csv(path)
    elif ext == '.tsv':
        df = pd.read_csv(path, sep='\t')
    elif ext == '.sas7bdat':
        df = pd.read_sas(path)
    elif ext == '.dta':
        df = pd.read_stata(path)
    else:
        df = pd.read_csv(path)

    # Apply filters
    for f in config['filters']:
        col, op, val = f['column'], f['operator'], f['value']
        if op in EQ_OPS:
            # P-2 batch: equality on original dtype so string filters work
            df = df[EQ_OPS[op](df[col], val)]
        elif op in NUMERIC_OPS:
            coerced = pd.to_numeric(df[col], errors='coerce')
            n_lost = int(coerced.isna().sum() - df[col].isna().sum())
            if n_lost > 0:
                print(f"[WARNING] Filter '{col} {op} {val}': {n_lost} non-numeric "
                      f"value(s) in column '{col}' coerced to NaN and excluded.",
                      file=sys.stderr)
            df = df[NUMERIC_OPS[op](coerced, val)]
        elif op == 'in':
            df = df[df[col].isin(val)]
        elif op == 'not_in':
            df = df[~df[col].isin(val)]

    # Compute derived variables
    for dv in config['derived_variables']:
        if dv['type'] == 'coalesce':
            result = pd.to_numeric(df[dv['columns'][0]], errors='coerce')
            for c in dv['columns'][1:]:
                result = result.combine_first(pd.to_numeric(df[c], errors='coerce'))
            df[dv['name']] = result
        elif dv['type'] == 'expression':
            local_ns = {}
            for dep in dv.get('dependencies', []):
                local_ns[dep] = pd.to_numeric(df[dep], errors='coerce')
            local_ns['np'] = np
            df[dv['name']] = eval(dv['expression'], {"__builtins__": {}}, local_ns)

    return df


def split_groups(df, config):
    """Split dataframe into groups defined in config. Supports 2+ groups.

    P-1 batch (2026-05-07): dtype-aware comparison + abort on empty group.
    Previously force-coerced to numeric, which silently wiped string-coded groups
    (e.g. "control"/"treatment", "M"/"F") to NaN and continued with empty groups.
    """
    gc = config['grouping']
    col = gc['column']
    g_vals = [g['value'] for g in gc['groups']]

    if all(isinstance(v, (int, float)) for v in g_vals):
        col_vals = pd.to_numeric(df[col], errors='coerce')
    else:
        col_vals = df[col]

    groups = []
    for g in gc['groups']:
        sub = df[col_vals == g['value']]
        if len(sub) == 0:
            sys.exit(
                f"[FATAL] split_groups: group '{g['label']}' (value={g['value']!r}) "
                f"produced empty subset. Check column '{col}' values vs configured "
                f"group value."
            )
        groups.append(sub)
    return groups


def prepare_data(df, config):
    """Convert time/event columns to numeric and handle missing data.

    P-6 audit fix (2026-05-08, km-A-2): event column encoding now validated.
    KM analysis assumes 1=event, 0=censor; reversed coding (`0=event`),
    string-coded (`Y/N`, `dead/alive`), or non-binary (`2/1`) silently
    inverts log-rank statistics if not flagged.
    """
    df_work = df.copy()
    time_col = config['time_column']
    event_col = config['event_column']

    # P-6 fix: detect non-numeric event values BEFORE coercion
    n_orig_event = int(df_work[event_col].notna().sum())
    df_work[time_col] = pd.to_numeric(df_work[time_col], errors='coerce')
    df_work[event_col] = pd.to_numeric(df_work[event_col], errors='coerce')

    n_lost = n_orig_event - int(df_work[event_col].notna().sum())
    if n_lost > 0:
        unique_orig = sorted([str(v) for v in df[event_col].dropna().unique()])[:5]
        print(f"[WARNING] Event column '{event_col}': {n_lost} non-numeric value(s) "
              f"coerced to NaN. Unique original values: {unique_orig}. KM expects "
              f"1=event, 0=censor; recode strings (e.g. dead/alive → 1/0) before run.",
              file=sys.stderr)

    # P-6 fix: validate non-NA unique values are subset of {0, 1}
    non_na_unique = sorted(df_work[event_col].dropna().unique().tolist())
    expected = {0, 1, 0.0, 1.0}
    extra = [v for v in non_na_unique if v not in expected]
    if extra:
        sys.exit(
            f"[FATAL] Event column '{event_col}' has unexpected values: {extra}. "
            f"KM requires strictly {{0, 1}} (1=event, 0=censor). "
            f"For multi-level status (competing event), use /competing-risk instead. "
            f"For reversed coding, recode the column before running."
        )
    if 1 not in non_na_unique and 1.0 not in non_na_unique:
        print(f"[WARNING] Event column '{event_col}' has no events (no 1 in data). "
              f"All observations censored. KM estimates will be vacuous.",
              file=sys.stderr)

    # Apply time scale conversion
    if config['time_scale'] is not None:
        df_work[time_col] = df_work[time_col] / config['time_scale']

    # Handle missing data
    if config['missing_data_handling'] == 'listwise':
        before_n = len(df_work)
        df_work = df_work.dropna(subset=[time_col, event_col])
        after_n = len(df_work)
        if before_n != after_n:
            print(f"  Listwise deletion: {before_n} → {after_n} rows ({before_n - after_n} removed)")

    # Exclude non-positive times
    df_work = df_work[df_work[time_col] > 0]

    return df_work


# ═══════════════════════════════════════════════════════════════════════════════
# Missing Data Check
# ═══════════════════════════════════════════════════════════════════════════════

def check_missing(df, group_dfs, config):
    """Print a comprehensive missing data report."""
    group_labels = [g['label'] for g in config['grouping']['groups']]
    time_col = config['time_column']
    event_col = config['event_column']

    lines = []
    lines.append("=" * 100)
    lines.append("MISSING DATA REPORT — Kaplan-Meier Survival Analysis")
    lines.append("=" * 100)

    header_parts = [f"{'Variable':<35}", f"{'Overall':>20}"]
    for label in group_labels:
        header_parts.append(f"{label:>20}")
    lines.append("  ".join(header_parts))
    lines.append("-" * 100)

    has_missing = False

    for col, col_label in [(time_col, 'Time'), (event_col, 'Event')]:
        display_label = f"{col_label} ({col})"
        row_parts = [f"  {display_label:<33}"]
        all_pops = [("Overall", df)] + list(zip(group_labels, group_dfs))
        for pop_name, pop_df in all_pops:
            if col in pop_df.columns:
                n_miss = pd.to_numeric(pop_df[col], errors='coerce').isna().sum()
                n_total = len(pop_df)
                pct = n_miss / n_total * 100 if n_total > 0 else 0
                marker = " ⚠" if n_miss > 0 else ""
                row_parts.append(f"{n_miss}/{n_total} ({pct:.1f}%){marker:>3}")
                if n_miss > 0:
                    has_missing = True
            else:
                row_parts.append(f"{'Column not found':>20}")
                has_missing = True
        lines.append("  ".join(row_parts))

    lines.append("-" * 100)

    # Effective sample size after listwise deletion
    before_n = len(df)
    valid = df.dropna(subset=[time_col, event_col])
    after_n = len(valid)
    lines.append(f"\nEffective sample size (after excluding missing time/event):")
    lines.append(f"  Overall: {after_n} / {before_n} ({after_n/before_n*100:.1f}%)" if before_n > 0 else "  Overall: 0")

    for label, gdf in zip(group_labels, group_dfs):
        g_before = len(gdf)
        g_valid = gdf.dropna(subset=[time_col, event_col])
        g_after = len(g_valid)
        pct = g_after / g_before * 100 if g_before > 0 else 0
        lines.append(f"  {label}: {g_after} / {g_before} ({pct:.1f}%)")

    lines.append("=" * 100)

    if not has_missing:
        lines.append("\n✓ No missing values detected in time or event columns.")
    else:
        lines.append("\n⚠ Missing values detected. Default handling: listwise deletion.")

    print("\n".join(lines))


# ═══════════════════════════════════════════════════════════════════════════════
# Kaplan-Meier Analysis
# ═══════════════════════════════════════════════════════════════════════════════

def fit_km(group_dfs, config):
    """Fit KM curves for each group and return fitters + summary data."""
    time_col = config['time_column']
    event_col = config['event_column']
    group_labels = [g['label'] for g in config['grouping']['groups']]
    time_points = config['time_points']

    fitters = []
    summaries = []

    for i, (gdf, label) in enumerate(zip(group_dfs, group_labels)):
        kmf = KaplanMeierFitter()
        T = gdf[time_col].values
        E = gdf[event_col].values

        n_total = len(T)
        n_events = int(E.sum())

        if n_total == 0:
            fitters.append(None)
            summaries.append({
                'label': label,
                'n': 0,
                'events': 0,
                'event_pct': 0.0,
                'median': 'N/A',
                'median_ci': ('N/A', 'N/A'),
                'survival_rates': {tp: ('N/A', 'N/A', 'N/A') for tp in time_points}
            })
            continue

        kmf.fit(T, event_observed=E, label=label)
        fitters.append(kmf)

        # Median survival with CI
        median_val = kmf.median_survival_time_
        if np.isinf(median_val) or np.isnan(median_val):
            median_str = 'NR'
            median_ci = ('NR', 'NR')
        else:
            median_str = f"{median_val:.2f}"
            # Get CI for median from the confidence interval of the survival function
            ci = kmf.confidence_interval_survival_function_
            try:
                # Find where survival crosses 0.5
                sf = kmf.survival_function_
                ci_lower_col = ci.columns[0]
                ci_upper_col = ci.columns[1]

                # Median CI bounds via inversion of survival CI (Brookmeyer-Crowley):
                # Lower bound of median CI = earliest time LOWER CI of S(t) reaches 0.5
                # Upper bound of median CI = earliest time UPPER CI of S(t) reaches 0.5
                # (Lower S_CI decays faster so it crosses 0.5 sooner; upper S_CI crosses later.)
                lower_ci_vals = ci[ci_lower_col]
                below_05_lower = lower_ci_vals[lower_ci_vals <= 0.5]
                if len(below_05_lower) > 0:
                    median_ci_lower = f"{below_05_lower.index[0]:.2f}"
                else:
                    median_ci_lower = 'NR'

                upper_ci_vals = ci[ci_upper_col]
                below_05_upper = upper_ci_vals[upper_ci_vals <= 0.5]
                if len(below_05_upper) > 0:
                    median_ci_upper = f"{below_05_upper.index[0]:.2f}"
                else:
                    median_ci_upper = 'NR'

                median_ci = (median_ci_lower, median_ci_upper)
            except Exception as e:
                # P-5 batch: surface the failure (was silent except: pass)
                print(f"[WARNING] Median CI estimation failed for group '{label}': "
                      f"{type(e).__name__}: {e}", file=sys.stderr)
                median_ci = ('N/A', 'N/A')

        # Survival rates at time points
        survival_rates = {}
        for tp in time_points:
            max_time = T.max()
            if tp > max_time:
                survival_rates[tp] = ('N/A', 'N/A', 'N/A')
            else:
                try:
                    # Get survival at time point
                    sf = kmf.survival_function_at_times(tp).values[0]
                    ci_at_t = kmf.confidence_interval_survival_function_
                    # Find the closest time in CI index
                    idx = ci_at_t.index
                    closest_idx = idx[idx <= tp]
                    if len(closest_idx) > 0:
                        ci_row = ci_at_t.loc[closest_idx[-1]]
                        ci_lo = ci_row.iloc[0]
                        ci_hi = ci_row.iloc[1]
                    else:
                        ci_lo = ci_hi = np.nan

                    sf_str = f"{sf*100:.1f}"
                    ci_lo_str = f"{ci_lo*100:.1f}" if not np.isnan(ci_lo) else 'N/A'
                    ci_hi_str = f"{ci_hi*100:.1f}" if not np.isnan(ci_hi) else 'N/A'
                    survival_rates[tp] = (sf_str, ci_lo_str, ci_hi_str)
                except Exception as e:
                    # P-5 batch: surface the failure (was silent except: pass)
                    print(f"[WARNING] Survival rate at t={tp} failed for group "
                          f"'{label}': {type(e).__name__}: {e}", file=sys.stderr)
                    survival_rates[tp] = ('N/A', 'N/A', 'N/A')

        summaries.append({
            'label': label,
            'n': n_total,
            'events': n_events,
            'event_pct': n_events / n_total * 100 if n_total > 0 else 0.0,
            'median': median_str,
            'median_ci': median_ci,
            'survival_rates': survival_rates
        })

    return fitters, summaries


def run_logrank(group_dfs, config):
    """Run log-rank test comparing survival between groups."""
    time_col = config['time_column']
    event_col = config['event_column']
    group_labels = [g['label'] for g in config['grouping']['groups']]

    # Filter out empty groups
    valid_groups = [(gdf, label) for gdf, label in zip(group_dfs, group_labels) if len(gdf) > 0]

    if len(valid_groups) < 2:
        return {'test_statistic': 'N/A', 'p_value': 'N/A', 'df': 'N/A',
                'comparison': 'Insufficient groups for comparison'}

    if len(valid_groups) == 2:
        gdf1, label1 = valid_groups[0]
        gdf2, label2 = valid_groups[1]
        result = logrank_test(
            gdf1[time_col], gdf2[time_col],
            event_observed_A=gdf1[event_col],
            event_observed_B=gdf2[event_col]
        )
        return {
            'test_statistic': f"{result.test_statistic:.4f}",
            'p_value': result.p_value,
            'df': 1,
            'comparison': f"{label1} vs {label2}"
        }
    else:
        # Multivariate log-rank for 3+ groups
        all_T = pd.concat([gdf[time_col] for gdf, _ in valid_groups])
        all_E = pd.concat([gdf[event_col] for gdf, _ in valid_groups])
        all_G = pd.concat([pd.Series([label] * len(gdf), index=gdf.index) for gdf, label in valid_groups])

        result = multivariate_logrank_test(all_T, all_G, all_E)
        return {
            'test_statistic': f"{result.test_statistic:.4f}",
            'p_value': result.p_value,
            'df': len(valid_groups) - 1,
            'comparison': ' vs '.join([label for _, label in valid_groups])
        }


def format_pvalue(p):
    """Format p-value for display."""
    if isinstance(p, str):
        return p
    if p < 0.001:
        return '<0.001'
    else:
        return f"{p:.3f}"


# ═══════════════════════════════════════════════════════════════════════════════
# Plot Generation
# ═══════════════════════════════════════════════════════════════════════════════

def generate_km_plot(fitters, summaries, logrank_result, config):
    """Generate publication-ready KM survival curve plot with number at risk table."""
    ps = config['plot_settings']
    time_unit = config['time_unit']
    group_labels = [g['label'] for g in config['grouping']['groups']]

    # Determine colors
    colors = ps['colors'] if ps['colors'] else DEFAULT_COLORS[:len(group_labels)]
    while len(colors) < len(group_labels):
        colors.append(DEFAULT_COLORS[len(colors) % len(DEFAULT_COLORS)])

    figsize = ps['figsize']
    show_risk_table = ps['number_at_risk']

    # Create figure with gridspec for KM plot + number at risk table
    if show_risk_table:
        fig = plt.figure(figsize=(figsize[0], figsize[1] + 0.4 * len(group_labels)))
        gs = gridspec.GridSpec(2, 1, height_ratios=[4, len(group_labels) * 0.6],
                               hspace=0.08)
        ax_km = fig.add_subplot(gs[0])
        ax_risk = fig.add_subplot(gs[1], sharex=ax_km)
    else:
        fig, ax_km = plt.subplots(figsize=figsize)
        ax_risk = None

    # Determine x-axis limits and ticks
    max_times = []
    for kmf in fitters:
        if kmf is not None:
            max_times.append(kmf.timeline.max())

    if not max_times:
        print("Error: No valid KM fitters to plot.")
        return

    data_max = max(max_times)
    x_max = ps['x_max'] if ps['x_max'] is not None else data_max
    x_tick_interval = ps['x_tick_interval']
    if x_tick_interval is None:
        # Auto-determine tick interval
        if x_max <= 5:
            x_tick_interval = 1
        elif x_max <= 12:
            x_tick_interval = 2
        elif x_max <= 30:
            x_tick_interval = 5
        elif x_max <= 120:
            x_tick_interval = 12
        else:
            x_tick_interval = max(1, int(x_max / 10))

    x_ticks = np.arange(0, x_max + x_tick_interval * 0.01, x_tick_interval)

    # Plot KM curves
    for i, (kmf, label, color) in enumerate(zip(fitters, group_labels, colors)):
        if kmf is None:
            continue

        # Plot survival function as step function
        sf = kmf.survival_function_
        ax_km.step(sf.index, sf.iloc[:, 0], where='post', color=color,
                   linewidth=1.8, label=label)

        # CI shading
        if ps['confidence_interval']:
            ci = kmf.confidence_interval_survival_function_
            ax_km.fill_between(ci.index,
                               ci.iloc[:, 0], ci.iloc[:, 1],
                               step='post', alpha=0.15, color=color)

        # Censor marks
        if ps['censor_marks']:
            event_table = kmf.event_table
            censor_times = event_table[event_table['censored'] > 0].index
            for ct in censor_times:
                if ct <= x_max:
                    sv = kmf.survival_function_at_times(ct).values[0]
                    ax_km.plot(ct, sv, '|', color=color, markersize=6,
                               markeredgewidth=1.2)

    # Plot formatting
    ax_km.set_xlim(0, x_max)
    ax_km.set_ylim(ps['y_min'], 1.02)
    ax_km.set_ylabel('Survival Probability', fontsize=12, fontweight='bold')
    ax_km.yaxis.set_major_formatter(mticker.PercentFormatter(xmax=1.0, decimals=0))
    ax_km.set_xticks(x_ticks)

    if not show_risk_table:
        ax_km.set_xlabel(f'Time ({time_unit})', fontsize=12, fontweight='bold')
    else:
        plt.setp(ax_km.get_xticklabels(), visible=False)

    ax_km.tick_params(axis='both', labelsize=10)
    ax_km.grid(axis='y', alpha=0.3, linestyle='--')
    ax_km.spines['top'].set_visible(False)
    ax_km.spines['right'].set_visible(False)

    # Add log-rank p-value text
    p_val = logrank_result['p_value']
    p_text = f"Log-rank p = {format_pvalue(p_val)}"
    ax_km.text(0.98, 0.02, p_text, transform=ax_km.transAxes,
               fontsize=10, ha='right', va='bottom',
               bbox=dict(boxstyle='round,pad=0.3', facecolor='white',
                          edgecolor='gray', alpha=0.8))

    # Legend
    ax_km.legend(loc=ps['legend_position'], fontsize=10, framealpha=0.9)

    # Title
    if config['title']:
        ax_km.set_title(config['title'], fontsize=14, fontweight='bold', pad=15)

    # Number at risk table
    if show_risk_table and ax_risk is not None:
        risk_times = ps['risk_table_times'] if ps['risk_table_times'] else x_ticks

        ax_risk.set_xlim(0, x_max)
        ax_risk.set_ylim(-0.5, len(group_labels) - 0.5)
        ax_risk.invert_yaxis()

        for i, (kmf, label, color) in enumerate(zip(fitters, group_labels, colors)):
            ax_risk.text(-0.02, i, label, transform=ax_risk.get_yaxis_transform(),
                         ha='right', va='center', fontsize=9, fontweight='bold',
                         color=color)

            if kmf is None:
                for t in risk_times:
                    ax_risk.text(t, i, '0', ha='center', va='center', fontsize=8)
                continue

            event_table = kmf.event_table
            for t in risk_times:
                # Number at risk at time t
                at_risk_rows = event_table[event_table.index <= t]
                if len(at_risk_rows) > 0:
                    last_row = at_risk_rows.iloc[-1]
                    n_at_risk = int(last_row['at_risk'] - last_row['observed'] - last_row['censored'])
                    if t == 0:
                        n_at_risk = int(event_table.iloc[0]['at_risk'])
                else:
                    n_at_risk = int(event_table.iloc[0]['at_risk'])

                n_at_risk = max(0, n_at_risk)
                ax_risk.text(t, i, str(n_at_risk), ha='center', va='center', fontsize=8)

        ax_risk.set_xlabel(f'Time ({time_unit})', fontsize=12, fontweight='bold')
        ax_risk.set_xticks(x_ticks)
        ax_risk.set_yticks(range(len(group_labels)))
        ax_risk.set_yticklabels([])
        ax_risk.tick_params(axis='y', length=0)
        ax_risk.tick_params(axis='x', labelsize=10)
        ax_risk.spines['top'].set_visible(False)
        ax_risk.spines['right'].set_visible(False)
        ax_risk.spines['left'].set_visible(False)

        # "No. at risk" label
        ax_risk.text(-0.02, -0.8, 'No. at risk', transform=ax_risk.get_yaxis_transform(),
                     ha='right', va='center', fontsize=9, fontweight='bold', fontstyle='italic')

    plt.tight_layout()

    # Save
    output_base = os.path.splitext(config['output_file'])[0]
    png_path = output_base + '.png'
    fig.savefig(png_path, dpi=ps['dpi'], bbox_inches='tight', facecolor='white')
    plt.close(fig)
    print(f"  KM plot saved: {png_path}")
    return png_path


# ═══════════════════════════════════════════════════════════════════════════════
# Excel Output
# ═══════════════════════════════════════════════════════════════════════════════

# Styling constants
DARK_BLUE = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
MEDIUM_BLUE = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
LIGHT_BLUE = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
LIGHT_GRAY = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
WHITE_FILL = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
WHITE_FONT = Font(color='FFFFFF', bold=True, size=12)
HEADER_FONT = Font(color='FFFFFF', bold=True, size=10)
DATA_FONT = Font(size=10)
BOLD_FONT = Font(size=10, bold=True)
THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)


def write_excel(summaries, logrank_result, config):
    """Write formatted Excel output with survival summary and log-rank test."""
    wb = Workbook()
    time_unit = config['time_unit']
    time_points = config['time_points']

    # ── Sheet 1: Survival Summary ──
    ws1 = wb.active
    ws1.title = "Survival Summary"

    # Calculate total N
    total_n = sum(s['n'] for s in summaries)

    # Row 1: Title
    row = 1
    title_text = config['title'] if config['title'] else 'Kaplan-Meier Survival Analysis'
    n_cols = 4 + len(time_points)  # Group, N, Events(%), Median(CI), + survival rates
    ws1.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
    cell = ws1.cell(row=row, column=1, value=title_text)
    cell.font = WHITE_FONT
    cell.fill = DARK_BLUE
    cell.alignment = CENTER_ALIGN
    ws1.row_dimensions[row].height = 30

    # Row 2: Subtitle
    row = 2
    subtitle = config['subtitle'] if config['subtitle'] else ''
    subtitle_full = f"{subtitle} (N = {total_n})" if subtitle else f"N = {total_n}"
    ws1.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
    cell = ws1.cell(row=row, column=1, value=subtitle_full)
    cell.font = Font(size=10, italic=True)
    cell.fill = LIGHT_BLUE
    cell.alignment = CENTER_ALIGN
    ws1.row_dimensions[row].height = 22

    # Row 3: Headers
    row = 3
    headers = ['Group', 'N', 'Events (%)', f'Median Survival ({time_unit})\n(95% CI)']
    for tp in time_points:
        headers.append(f'{tp}-{time_unit[0].upper()} Survival %\n(95% CI)')

    for j, header in enumerate(headers, 1):
        cell = ws1.cell(row=row, column=j, value=header)
        cell.font = HEADER_FONT
        cell.fill = MEDIUM_BLUE
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
    ws1.row_dimensions[row].height = 35

    # Data rows
    for idx, summary in enumerate(summaries):
        row = 4 + idx
        fill = LIGHT_GRAY if idx % 2 == 0 else WHITE_FILL

        # Group label
        cell = ws1.cell(row=row, column=1, value=summary['label'])
        cell.font = BOLD_FONT
        cell.fill = fill
        cell.alignment = LEFT_ALIGN
        cell.border = THIN_BORDER

        # N
        cell = ws1.cell(row=row, column=2, value=summary['n'])
        cell.font = DATA_FONT
        cell.fill = fill
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

        # Events (%)
        if summary['n'] > 0:
            events_text = f"{summary['events']} ({summary['event_pct']:.1f}%)"
        else:
            events_text = "0 (0.0%)"
        cell = ws1.cell(row=row, column=3, value=events_text)
        cell.font = DATA_FONT
        cell.fill = fill
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

        # Median (95% CI)
        if summary['median'] == 'NR':
            median_text = 'NR'
        elif summary['median'] == 'N/A':
            median_text = 'N/A'
        else:
            ci_lo, ci_hi = summary['median_ci']
            median_text = f"{summary['median']} ({ci_lo}–{ci_hi})"
        cell = ws1.cell(row=row, column=4, value=median_text)
        cell.font = DATA_FONT
        cell.fill = fill
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

        # Survival rates at time points
        for tp_idx, tp in enumerate(time_points):
            sr = summary['survival_rates'].get(tp, ('N/A', 'N/A', 'N/A'))
            if sr[0] == 'N/A':
                sr_text = 'N/A'
            else:
                sr_text = f"{sr[0]}% ({sr[1]}–{sr[2]})"
            cell = ws1.cell(row=row, column=5 + tp_idx, value=sr_text)
            cell.font = DATA_FONT
            cell.fill = fill
            cell.alignment = CENTER_ALIGN
            cell.border = THIN_BORDER

    ws1.row_dimensions[row].height = 22

    # Footnote rows
    footnote_row = row + 2

    # Log-rank p-value
    p_text = format_pvalue(logrank_result['p_value'])
    logrank_footnote = f"Log-rank test: χ² = {logrank_result['test_statistic']}, df = {logrank_result['df']}, p = {p_text}"
    ws1.merge_cells(start_row=footnote_row, start_column=1, end_row=footnote_row, end_column=n_cols)
    cell = ws1.cell(row=footnote_row, column=1, value=logrank_footnote)
    cell.font = Font(size=9, italic=True)
    cell.alignment = LEFT_ALIGN

    # Abbreviations
    if config['abbreviations']:
        footnote_row += 1
        ws1.merge_cells(start_row=footnote_row, start_column=1, end_row=footnote_row, end_column=n_cols)
        cell = ws1.cell(row=footnote_row, column=1, value=config['abbreviations'])
        cell.font = Font(size=9, italic=True)
        cell.alignment = LEFT_ALIGN

    # Method note
    footnote_row += 1
    method_note = "Survival estimated using Kaplan-Meier method. NR = not reached."
    ws1.merge_cells(start_row=footnote_row, start_column=1, end_row=footnote_row, end_column=n_cols)
    cell = ws1.cell(row=footnote_row, column=1, value=method_note)
    cell.font = Font(size=9, italic=True)
    cell.alignment = LEFT_ALIGN

    # Column widths
    ws1.column_dimensions['A'].width = 20
    ws1.column_dimensions['B'].width = 10
    ws1.column_dimensions['C'].width = 14
    ws1.column_dimensions['D'].width = 24
    for tp_idx in range(len(time_points)):
        col_letter = get_column_letter(5 + tp_idx)
        ws1.column_dimensions[col_letter].width = 20

    # ── Sheet 2: Log-Rank Test ──
    ws2 = wb.create_sheet("Log-Rank Test")

    # Row 1: Title
    ws2.merge_cells('A1:D1')
    cell = ws2.cell(row=1, column=1, value='Log-Rank Test Results')
    cell.font = WHITE_FONT
    cell.fill = DARK_BLUE
    cell.alignment = CENTER_ALIGN
    ws2.row_dimensions[1].height = 30

    # Row 2: Comparison
    ws2.merge_cells('A2:D2')
    cell = ws2.cell(row=2, column=1, value=f"Comparison: {logrank_result['comparison']}")
    cell.font = Font(size=10, italic=True)
    cell.fill = LIGHT_BLUE
    cell.alignment = CENTER_ALIGN
    ws2.row_dimensions[2].height = 22

    # Row 3: Headers
    for j, header in enumerate(['Metric', 'Value'], 1):
        cell = ws2.cell(row=3, column=j, value=header)
        cell.font = HEADER_FONT
        cell.fill = MEDIUM_BLUE
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

    # Data
    test_data = [
        ('Test Statistic (χ²)', logrank_result['test_statistic']),
        ('Degrees of Freedom', str(logrank_result['df'])),
        ('p-value', format_pvalue(logrank_result['p_value'])),
    ]
    for idx, (metric, value) in enumerate(test_data):
        r = 4 + idx
        fill = LIGHT_GRAY if idx % 2 == 0 else WHITE_FILL
        cell = ws2.cell(row=r, column=1, value=metric)
        cell.font = DATA_FONT
        cell.fill = fill
        cell.alignment = LEFT_ALIGN
        cell.border = THIN_BORDER
        cell = ws2.cell(row=r, column=2, value=value)
        cell.font = BOLD_FONT
        cell.fill = fill
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

    # Footnote
    footnote_row = 4 + len(test_data) + 1
    ws2.merge_cells(start_row=footnote_row, start_column=1, end_row=footnote_row, end_column=2)
    cell = ws2.cell(row=footnote_row, column=1,
                    value="Log-rank (Mantel-Cox) test for equality of survival distributions.")
    cell.font = Font(size=9, italic=True)
    cell.alignment = LEFT_ALIGN

    ws2.column_dimensions['A'].width = 25
    ws2.column_dimensions['B'].width = 20

    # Save
    output_path = config['output_file']
    wb.save(output_path)
    print(f"  Excel saved: {output_path}")


# ═══════════════════════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════════════════════

def _apply_academic_style():
    """academic-style integration (Wave 1). Consumes --style=<journal> from
    sys.argv before any other parsing. No-op when academic-style skill is not
    installed or --style is absent (byte-identical to legacy behavior)."""
    from pathlib import Path as _P
    _as = _P.home() / ".claude" / "skills" / "academic-style"
    if not _as.exists():
        return None
    if str(_as) not in sys.path:
        sys.path.insert(0, str(_as))
    try:
        from core.biostats_integration import apply_style_from_argv
    except ImportError:
        return None
    return apply_style_from_argv(sys.argv)


def main():
    _AS_STYLE = _apply_academic_style()
    if _AS_STYLE:
        print(f"  Applied {_AS_STYLE} journal style")

    if len(sys.argv) < 2:
        print("Usage: python kaplan_meier.py config.json [--check-missing] [--style=<journal>]")
        sys.exit(1)

    config_path = sys.argv[1]
    check_missing_mode = '--check-missing' in sys.argv

    config = load_config(config_path)

    print("Loading data...")
    df = load_data(config)
    print(f"  Data loaded: {len(df)} rows")

    # Prepare data (convert to numeric, apply time scale)
    df = prepare_data(df, config)
    print(f"  After preparation: {len(df)} rows")

    # Split into groups
    group_dfs = split_groups(df, config)
    group_labels = [g['label'] for g in config['grouping']['groups']]
    for gdf, label in zip(group_dfs, group_labels):
        print(f"  {label}: n = {len(gdf)}")

    # Prepare group data (time_scale / missing / non-positive already handled in prepare_data)
    prepared_groups = []
    for gdf in group_dfs:
        gdf_work = gdf.copy()
        time_col = config['time_column']
        event_col = config['event_column']
        gdf_work[time_col] = pd.to_numeric(gdf_work[time_col], errors='coerce')
        gdf_work[event_col] = pd.to_numeric(gdf_work[event_col], errors='coerce')
        gdf_work = gdf_work.dropna(subset=[time_col, event_col])
        gdf_work = gdf_work[gdf_work[time_col] > 0]
        prepared_groups.append(gdf_work)

    if check_missing_mode:
        # Use raw groups (before preparation) for missing report
        raw_groups = split_groups(load_data(config), config)
        check_missing(load_data(config), raw_groups, config)
        sys.exit(0)

    # Fit KM curves
    print("\nFitting Kaplan-Meier curves...")
    fitters, summaries = fit_km(prepared_groups, config)

    # Log-rank test
    print("Running log-rank test...")
    logrank_result = run_logrank(prepared_groups, config)
    print(f"  Log-rank p = {format_pvalue(logrank_result['p_value'])}")

    # Generate plot
    print("\nGenerating KM plot...")
    generate_km_plot(fitters, summaries, logrank_result, config)

    # Write Excel
    print("Writing Excel output...")
    write_excel(summaries, logrank_result, config)

    # Print summary
    print("\n" + "=" * 60)
    print("KAPLAN-MEIER ANALYSIS COMPLETE")
    print("=" * 60)
    for s in summaries:
        median_text = s['median']
        if median_text not in ('N/A', 'NR'):
            ci_lo, ci_hi = s['median_ci']
            median_text = f"{s['median']} ({ci_lo}–{ci_hi})"
        print(f"  {s['label']}: N={s['n']}, Events={s['events']} ({s['event_pct']:.1f}%), "
              f"Median={median_text}")
    print(f"\n  Log-rank p = {format_pvalue(logrank_result['p_value'])}")
    print("=" * 60)

    # ── QC Warnings ──────────────────────────────────────────────────────────
    # P-4 + P-5 batch (2026-05-07): track [QC_BLOCKING] emissions and exit
    # non-zero so pipeline detects failure; surface block-level exception.
    qc_blocking_count = 0

    def _qc_block(msg):
        nonlocal qc_blocking_count
        print(f"[QC_BLOCKING] {msg}", file=sys.stderr)
        qc_blocking_count += 1

    try:
        for s in summaries:
            if s['events'] == 0:
                _qc_block(f"Group '{s['label']}' has 0 events")
            if s['n'] < 30:
                print(f"[QC_WARNING] Group '{s['label']}' has n={s['n']} (< 30)", file=sys.stderr)
        # PH assumption check: KM and log-rank do not require fitting a Cox PH
        # model; PH diagnostics belong in Cox/survival-diagnostics workflows
        # (km-C-3 from audit; relevant when Cox is later layered on top).
    except Exception as e:
        # P-5 batch: surface QC exception (was silent except: pass)
        print(f"[QC_WARNING] QC block raised: {type(e).__name__}: {e}", file=sys.stderr)

    if qc_blocking_count > 0:
        sys.exit(2)


if __name__ == '__main__':
    main()
