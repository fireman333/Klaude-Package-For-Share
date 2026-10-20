#!/usr/bin/env python3
"""
Propensity Score Matching (PSM) Analysis
Performs greedy nearest-neighbor matching on logit(PS), then analyzes outcomes
(Cox HR or Logistic OR) on the matched cohort.

Outputs: Excel (outcome table + balance table) and Love plot (PNG).

Usage:
    python psm_matching.py config.json                  # Full analysis
    python psm_matching.py config.json --check-missing  # Missing data report only
"""

import hashlib
import sys
import json
import os
import operator
import pandas as pd
import numpy as np
# P-3 batch (2026-05-07 partitioned-flame audit): removed `import warnings;
# warnings.filterwarnings('ignore')` — global suppression hid statsmodels / scipy
# convergence, separation, and numerical warnings during PS estimation
# (silent error pattern).

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ═══════════════════════════════════════════════════════════════════════════════
# Configuration
# ═══════════════════════════════════════════════════════════════════════════════

def load_config(path):
    with open(path, 'r', encoding='utf-8') as f:
        config = json.load(f)
    config.setdefault('filters', [])
    config.setdefault('derived_variables', [])
    config.setdefault('title', 'Propensity Score Matched Analysis')
    config.setdefault('subtitle', '')
    config.setdefault('abbreviations', '')
    config.setdefault('missing_data_handling', 'listwise')

    # Matching defaults
    matching = config.setdefault('matching', {})
    caliper = matching.setdefault('caliper', {})
    caliper.setdefault('enabled', True)
    caliper.setdefault('width', 0.20)
    matching.setdefault('ratio', 1)

    return config


# ═══════════════════════════════════════════════════════════════════════════════
# Data Loading & Preprocessing
# ═══════════════════════════════════════════════════════════════════════════════

# P-2 batch (2026-05-07 partitioned-flame audit): split equality vs numeric ops.
# Equality (==, !=) must compare on original dtype, otherwise string filters like
# `sex == "F"` silently coerce to NaN and drop every row.
EQ_OPS = {'==': operator.eq, '!=': operator.ne}
NUMERIC_OPS = {'>': operator.gt, '>=': operator.ge, '<': operator.lt, '<=': operator.le}
OPS = {**EQ_OPS, **NUMERIC_OPS}


def load_data(config):
    path = config['data_file']
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xls'):
        df = pd.read_excel(path)
    elif ext == '.csv':
        df = pd.read_csv(path)
    elif ext == '.tsv':
        df = pd.read_csv(path, sep='\t')
    elif ext == '.sas7bdat':
        df = pd.read_sas(path)
    elif ext == '.dta':
        df = pd.read_stata(path)
    else:
        df = pd.read_csv(path)

    # Apply filters
    for f in config['filters']:
        col, op, val = f['column'], f['operator'], f['value']
        if op in EQ_OPS:
            # P-2 batch: equality on original dtype so string filters work
            df = df[EQ_OPS[op](df[col], val)]
        elif op in NUMERIC_OPS:
            coerced = pd.to_numeric(df[col], errors='coerce')
            n_lost = int(coerced.isna().sum() - df[col].isna().sum())
            if n_lost > 0:
                print(f"[WARNING] Filter '{col} {op} {val}': {n_lost} non-numeric "
                      f"value(s) in column '{col}' coerced to NaN and excluded.",
                      file=sys.stderr)
            df = df[NUMERIC_OPS[op](coerced, val)]
        elif op == 'in':
            df = df[df[col].isin(val)]
        elif op == 'not_in':
            df = df[~df[col].isin(val)]

    # Compute derived variables
    for dv in config['derived_variables']:
        if dv['type'] == 'coalesce':
            result = pd.to_numeric(df[dv['columns'][0]], errors='coerce')
            for c in dv['columns'][1:]:
                result = result.combine_first(pd.to_numeric(df[c], errors='coerce'))
            df[dv['name']] = result
        elif dv['type'] == 'expression':
            local_ns = {}
            for dep in dv.get('dependencies', []):
                local_ns[dep] = pd.to_numeric(df[dep], errors='coerce')
            local_ns['np'] = np
            df[dv['name']] = eval(dv['expression'], {"__builtins__": {}}, local_ns)

    return df


def split_groups(df, config):
    """P-1 batch (2026-05-07): dtype-aware comparison + abort on empty group.

    Previously force-coerced grouping column to numeric, which silently wiped
    string-coded treatment/control groups to NaN and continued with empty groups
    (PSM is especially sensitive — empty treatment group invalidates the entire
    matching procedure).
    """
    gc = config['grouping']
    col = gc['column']
    g1_val = gc['groups'][0]['value']
    g2_val = gc['groups'][1]['value']

    if isinstance(g1_val, (int, float)) and isinstance(g2_val, (int, float)):
        col_vals = pd.to_numeric(df[col], errors='coerce')
    else:
        col_vals = df[col]

    grp1 = df[col_vals == g1_val]
    grp2 = df[col_vals == g2_val]

    if len(grp1) == 0 or len(grp2) == 0:
        sys.exit(
            f"[FATAL] split_groups produced empty group(s): "
            f"{gc['groups'][0]['label']} n={len(grp1)}, "
            f"{gc['groups'][1]['label']} n={len(grp2)}. "
            f"Check column '{col}' values vs configured group values "
            f"({g1_val!r}, {g2_val!r})."
        )
    return grp1, grp2


def get_matching_columns(config):
    """Get all matching variable column names."""
    cols = []
    for section in config['matching_variables']:
        for var in section['variables']:
            cols.append(var['column'])
    return cols


def outcome_is_survival(outcome, config):
    """Decide per outcome, not per config.

    `analysis_type` is a single value describing the whole run, but a plan with a
    time-to-event primary and a binary secondary produces a genuinely mixed
    outcomes list (research-plan's generate_configs.gen_psm does exactly this,
    and it is a common shape). Reading analysis_type for every outcome then
    raised KeyError on the odd one out. The keys present on the outcome are the
    fact; analysis_type is only the fallback when neither marker is there.
    """
    if 'time_column' in outcome and 'event_column' in outcome:
        return True
    if 'outcome_column' in outcome:
        return False
    return config.get('analysis_type') == 'cox'


def get_outcome_columns(config):
    """Get all outcome-related column names."""
    cols = []
    for outcome in config['outcomes']:
        if outcome_is_survival(outcome, config):
            cols.append(outcome['event_column'])
            cols.append(outcome['time_column'])
        else:
            cols.append(outcome['outcome_column'])
    return cols


def prepare_analysis_df(df, config):
    """Prepare dataframe: convert to numeric and handle missing data for matching variables."""
    match_cols = get_matching_columns(config)
    grp_col = config['grouping']['column']
    all_cols = list(set(match_cols + [grp_col]))

    df_work = df.copy()
    for col in all_cols:
        if col in df_work.columns:
            df_work[col] = pd.to_numeric(df_work[col], errors='coerce')

    if config['missing_data_handling'] == 'listwise':
        existing_cols = [c for c in all_cols if c in df_work.columns]
        before_n = len(df_work)
        df_work = df_work.dropna(subset=existing_cols)
        after_n = len(df_work)
        if before_n != after_n:
            print(f"  Listwise deletion: {before_n} → {after_n} rows ({before_n - after_n} removed)")

    return df_work


# ═══════════════════════════════════════════════════════════════════════════════
# Missing Data Check
# ═══════════════════════════════════════════════════════════════════════════════

def check_missing(df, grp1, grp2, config):
    """Print a comprehensive missing data report."""
    g1_label = config['grouping']['groups'][0]['label']
    g2_label = config['grouping']['groups'][1]['label']

    lines = []
    lines.append("=" * 100)
    lines.append("MISSING DATA REPORT — Propensity Score Matching")
    lines.append("=" * 100)
    lines.append(f"{'Variable':<35} {'Overall':>20} {g1_label:>20} {g2_label:>20}")
    lines.append("-" * 100)

    has_missing = False

    # Check matching variables
    lines.append("\n  [Matching Variables]")
    for section in config['matching_variables']:
        lines.append(f"\n    {section['header']}")
        for var in section['variables']:
            col = var['column']
            label = var['label']
            row_parts = [f"    {label:<31}"]
            for pop_name, pop_df in [("Overall", df), (g1_label, grp1), (g2_label, grp2)]:
                if col in pop_df.columns:
                    n_miss = pd.to_numeric(pop_df[col], errors='coerce').isna().sum()
                    n_total = len(pop_df)
                    pct = n_miss / n_total * 100 if n_total > 0 else 0
                    marker = " ⚠" if n_miss > 0 else ""
                    row_parts.append(f"{n_miss}/{n_total} ({pct:.1f}%){marker:>3}")
                    if n_miss > 0:
                        has_missing = True
                else:
                    row_parts.append(f"{'Column not found':>20}")
                    has_missing = True
            lines.append("  ".join(row_parts))

    # Check outcome variables
    lines.append("\n  [Outcome Variables]")
    for outcome in config['outcomes']:
        if outcome_is_survival(outcome, config):
            for col_key, col_label in [('event_column', 'Event'), ('time_column', 'Time')]:
                col = outcome[col_key]
                label = f"{outcome['label']} ({col_label})"
                row_parts = [f"    {label:<31}"]
                for pop_name, pop_df in [("Overall", df), (g1_label, grp1), (g2_label, grp2)]:
                    if col in pop_df.columns:
                        n_miss = pd.to_numeric(pop_df[col], errors='coerce').isna().sum()
                        n_total = len(pop_df)
                        pct = n_miss / n_total * 100 if n_total > 0 else 0
                        marker = " ⚠" if n_miss > 0 else ""
                        row_parts.append(f"{n_miss}/{n_total} ({pct:.1f}%){marker:>3}")
                        if n_miss > 0:
                            has_missing = True
                    else:
                        row_parts.append(f"{'Column not found':>20}")
                        has_missing = True
                lines.append("  ".join(row_parts))
        else:
            col = outcome['outcome_column']
            label = outcome['label']
            row_parts = [f"    {label:<31}"]
            for pop_name, pop_df in [("Overall", df), (g1_label, grp1), (g2_label, grp2)]:
                if col in pop_df.columns:
                    n_miss = pd.to_numeric(pop_df[col], errors='coerce').isna().sum()
                    n_total = len(pop_df)
                    pct = n_miss / n_total * 100 if n_total > 0 else 0
                    marker = " ⚠" if n_miss > 0 else ""
                    row_parts.append(f"{n_miss}/{n_total} ({pct:.1f}%){marker:>3}")
                    if n_miss > 0:
                        has_missing = True
                else:
                    row_parts.append(f"{'Column not found':>20}")
                    has_missing = True
            lines.append("  ".join(row_parts))

    lines.append("-" * 100)

    # Effective sample size after listwise deletion on matching variables
    match_cols = get_matching_columns(config)
    grp_col = config['grouping']['column']
    all_cols = list(set(match_cols + [grp_col]))
    df_complete = df.copy()
    for col in all_cols:
        if col in df_complete.columns:
            df_complete[col] = pd.to_numeric(df_complete[col], errors='coerce')
    existing = [c for c in all_cols if c in df_complete.columns]
    df_complete = df_complete.dropna(subset=existing)

    g1_complete, g2_complete = split_groups(df_complete, config)
    lines.append(f"\nEffective sample size after listwise deletion (matching variables):")
    lines.append(f"  Overall: {len(df_complete)}  |  {config['grouping']['groups'][0]['label']}: {len(g1_complete)}  |  {config['grouping']['groups'][1]['label']}: {len(g2_complete)}")

    if not has_missing:
        lines.append("\n✓ All variables have zero missing values.")
    else:
        lines.append("\n⚠ Missing values detected. Default handling: listwise deletion on matching variables.")

    lines.append("=" * 100)
    return "\n".join(lines)


# ═══════════════════════════════════════════════════════════════════════════════
# Propensity Score Estimation
# ═══════════════════════════════════════════════════════════════════════════════

def estimate_propensity_scores(df_clean, config):
    """Estimate propensity scores using unregularized logistic regression (statsmodels)."""
    import statsmodels.api as sm

    match_cols = get_matching_columns(config)
    grp_col = config['grouping']['column']
    g1_val = config['grouping']['groups'][0]['value']

    df_ps = df_clean.copy()

    # Ensure numeric
    for col in match_cols:
        df_ps[col] = pd.to_numeric(df_ps[col], errors='coerce')

    # Treatment indicator: 1 = Group 1 (treatment), 0 = Group 2 (control)
    df_ps['_treatment'] = (pd.to_numeric(df_ps[grp_col], errors='coerce') == g1_val).astype(int)

    X = df_ps[match_cols].values.astype(float)
    y = df_ps['_treatment'].values.astype(float)

    # Fit unregularized logistic regression (statsmodels); standard for PS models
    X_sm = sm.add_constant(X, has_constant='add')
    ps_model = sm.Logit(y, X_sm).fit(disp=0, maxiter=1000)
    ps = ps_model.predict(X_sm)
    ps = np.clip(ps, 1e-6, 1 - 1e-6)
    logit_ps = np.log(ps / (1 - ps))

    df_ps['_ps'] = ps
    df_ps['_logit_ps'] = logit_ps

    # Log PS model coefficients (for reviewer inspection of PS model spec)
    coef_names = ['(Intercept)'] + list(match_cols)
    print(f"  PS model coefficients (unregularized logistic):", file=sys.stderr)
    for name, coef, se in zip(coef_names, ps_model.params, ps_model.bse):
        print(f"    {name}: beta={coef:.4f}, SE={se:.4f}", file=sys.stderr)

    print(f"  PS range: [{ps.min():.4f}, {ps.max():.4f}]")
    print(f"  Logit(PS) range: [{logit_ps.min():.4f}, {logit_ps.max():.4f}]")

    return df_ps


# ═══════════════════════════════════════════════════════════════════════════════
# Greedy Nearest-Neighbor Matching
# ═══════════════════════════════════════════════════════════════════════════════

def perform_matching(df_ps, config):
    """
    Greedy nearest-neighbor matching on logit(PS), without replacement.
    Returns matched dataframe with _stratum column and match info dict.
    """
    ratio = config['matching']['ratio']
    caliper_enabled = config['matching']['caliper']['enabled']
    # Caliper SD multiplier: prefer top-level 'caliper_sd_multiplier', fall back to
    # matching.caliper.width (legacy key). Default 0.2 (Austin 2011).
    caliper_width = config.get(
        'caliper_sd_multiplier',
        config['matching']['caliper'].get('width', 0.2)
    )

    treatment = df_ps[df_ps['_treatment'] == 1].copy()
    control = df_ps[df_ps['_treatment'] == 0].copy()

    logit_ps_std = df_ps['_logit_ps'].std()
    caliper_val = caliper_width * logit_ps_std if caliper_enabled else None

    if caliper_enabled:
        print(f"  Caliper: {caliper_width} × {logit_ps_std:.4f} = {caliper_val:.4f}")
        print(
            f"[PSM] caliper_sd_multiplier={caliper_width}, "
            f"SD(logit_ps)={logit_ps_std:.4f}, caliper={caliper_val:.4f}",
            file=sys.stderr,
        )
    else:
        print(f"  Caliper: disabled")
        print("[PSM] caliper disabled", file=sys.stderr)
    print(f"  Matching ratio: 1:{ratio}")
    print(f"  Treatment: n={len(treatment)}, Control: n={len(control)}")

    # Shuffle treatment subjects to avoid ordering bias
    treatment = treatment.sample(frac=1, random_state=42).reset_index(drop=False)
    treatment = treatment.rename(columns={'index': '_orig_idx'})

    control_logit = control['_logit_ps'].values
    control_indices = control.index.values
    used = np.zeros(len(control), dtype=bool)

    matched_treatment_indices = []
    matched_control_indices = []
    control_strata_list = []
    strata = []
    stratum_id = 0

    for _, t_row in treatment.iterrows():
        t_logit = t_row['_logit_ps']

        # Compute distances to all unused controls
        distances = np.abs(control_logit - t_logit)
        distances[used] = np.inf

        if caliper_val is not None:
            # Mask out controls beyond caliper
            beyond_caliper = distances > caliper_val
            distances[beyond_caliper] = np.inf

        # Find the nearest N controls
        n_to_match = min(ratio, np.sum(~np.isinf(distances)))
        if n_to_match == 0:
            continue

        nearest_idx = np.argpartition(distances, n_to_match)[:n_to_match]
        # Sort by distance to get the closest ones
        nearest_idx = nearest_idx[np.argsort(distances[nearest_idx])]

        # Mark as used
        used[nearest_idx] = True

        # Record matches
        matched_treatment_indices.append(t_row['_orig_idx'])
        for ci in nearest_idx:
            matched_control_indices.append(control_indices[ci])
            control_strata_list.append(stratum_id)

        strata.append(stratum_id)
        stratum_id += 1

    # Build matched dataframe
    treatment_matched = df_ps.loc[matched_treatment_indices].copy()
    treatment_matched['_stratum'] = strata

    control_matched = df_ps.loc[matched_control_indices].copy()
    control_matched['_stratum'] = control_strata_list

    matched_df = pd.concat([treatment_matched, control_matched], ignore_index=True)

    n_matched_treatment = len(matched_treatment_indices)
    n_matched_control = len(matched_control_indices)
    n_unmatched = len(treatment) - n_matched_treatment

    match_info = {
        'n_treatment_original': len(treatment),
        'n_control_original': len(control),
        'n_treatment_matched': n_matched_treatment,
        'n_control_matched': n_matched_control,
        'n_unmatched': n_unmatched,
        'n_strata': stratum_id,
        'caliper_enabled': caliper_enabled,
        'caliper_width': caliper_width,
        'caliper_value': caliper_val,
        'ratio': ratio,
    }

    print(f"  Matched: {n_matched_treatment} treatment × {n_matched_control} control ({n_unmatched} unmatched)")

    return matched_df, match_info


# ═══════════════════════════════════════════════════════════════════════════════
# SMD Calculation
# ═══════════════════════════════════════════════════════════════════════════════

def compute_smd(vals_t, vals_c, var_type):
    """Compute absolute standardized mean difference."""
    vals_t = pd.to_numeric(vals_t, errors='coerce').dropna()
    vals_c = pd.to_numeric(vals_c, errors='coerce').dropna()

    if len(vals_t) == 0 or len(vals_c) == 0:
        return np.nan

    if var_type == 'continuous':
        mean_t = vals_t.mean()
        mean_c = vals_c.mean()
        var_t = vals_t.var(ddof=1)
        var_c = vals_c.var(ddof=1)
        pooled_sd = np.sqrt((var_t + var_c) / 2)
        if pooled_sd == 0:
            return 0.0
        return abs(mean_t - mean_c) / pooled_sd
    elif var_type == 'binary':
        prop_t = vals_t.mean()
        prop_c = vals_c.mean()
        denom = np.sqrt((prop_t * (1 - prop_t) + prop_c * (1 - prop_c)) / 2)
        if denom == 0:
            return 0.0
        return abs(prop_t - prop_c) / denom
    else:
        return np.nan


def compute_balance_table(grp1_before, grp2_before, grp1_after, grp2_after, config):
    """Compute balance statistics (mean/proportion and SMD) before and after matching."""
    results = []

    for section in config['matching_variables']:
        for var in section['variables']:
            col = var['column']
            label = var['label']
            vtype = var['type']

            # Before matching
            t_before = pd.to_numeric(grp1_before[col], errors='coerce').dropna()
            c_before = pd.to_numeric(grp2_before[col], errors='coerce').dropna()

            # After matching
            t_after = pd.to_numeric(grp1_after[col], errors='coerce').dropna()
            c_after = pd.to_numeric(grp2_after[col], errors='coerce').dropna()

            if vtype == 'continuous':
                mean_t_before = t_before.mean() if len(t_before) > 0 else np.nan
                sd_t_before = t_before.std(ddof=1) if len(t_before) > 0 else np.nan
                mean_c_before = c_before.mean() if len(c_before) > 0 else np.nan
                sd_c_before = c_before.std(ddof=1) if len(c_before) > 0 else np.nan
                mean_t_after = t_after.mean() if len(t_after) > 0 else np.nan
                sd_t_after = t_after.std(ddof=1) if len(t_after) > 0 else np.nan
                mean_c_after = c_after.mean() if len(c_after) > 0 else np.nan
                sd_c_after = c_after.std(ddof=1) if len(c_after) > 0 else np.nan

                display_t_before = f"{mean_t_before:.2f} \u00b1 {sd_t_before:.2f}" if not np.isnan(mean_t_before) else "N/A"
                display_c_before = f"{mean_c_before:.2f} \u00b1 {sd_c_before:.2f}" if not np.isnan(mean_c_before) else "N/A"
                display_t_after = f"{mean_t_after:.2f} \u00b1 {sd_t_after:.2f}" if not np.isnan(mean_t_after) else "N/A"
                display_c_after = f"{mean_c_after:.2f} \u00b1 {sd_c_after:.2f}" if not np.isnan(mean_c_after) else "N/A"
            else:
                # Binary: show proportion as percentage
                prop_t_before = t_before.mean() * 100 if len(t_before) > 0 else np.nan
                prop_c_before = c_before.mean() * 100 if len(c_before) > 0 else np.nan
                prop_t_after = t_after.mean() * 100 if len(t_after) > 0 else np.nan
                prop_c_after = c_after.mean() * 100 if len(c_after) > 0 else np.nan

                n_t_before = int(t_before.sum()) if len(t_before) > 0 else 0
                n_c_before = int(c_before.sum()) if len(c_before) > 0 else 0
                n_t_after = int(t_after.sum()) if len(t_after) > 0 else 0
                n_c_after = int(c_after.sum()) if len(c_after) > 0 else 0

                display_t_before = f"{n_t_before} ({prop_t_before:.1f}%)" if not np.isnan(prop_t_before) else "N/A"
                display_c_before = f"{n_c_before} ({prop_c_before:.1f}%)" if not np.isnan(prop_c_before) else "N/A"
                display_t_after = f"{n_t_after} ({prop_t_after:.1f}%)" if not np.isnan(prop_t_after) else "N/A"
                display_c_after = f"{n_c_after} ({prop_c_after:.1f}%)" if not np.isnan(prop_c_after) else "N/A"

            smd_before = compute_smd(t_before, c_before, vtype)
            smd_after = compute_smd(t_after, c_after, vtype)

            results.append({
                'label': label,
                'column': col,
                'type': vtype,
                'section': section['header'],
                'display_t_before': display_t_before,
                'display_c_before': display_c_before,
                'smd_before': smd_before,
                'display_t_after': display_t_after,
                'display_c_after': display_c_after,
                'smd_after': smd_after,
            })

    return results


# ═══════════════════════════════════════════════════════════════════════════════
# Outcome Analysis (Post-Matching)
# ═══════════════════════════════════════════════════════════════════════════════

def analyze_outcome_cox(matched_df, outcome_def, config):
    """Analyze a single outcome using stratified Cox regression."""
    from lifelines import CoxPHFitter

    event_col = outcome_def['event_column']
    time_col = outcome_def['time_column']

    df = matched_df.copy()
    df[event_col] = pd.to_numeric(df[event_col], errors='coerce')
    df[time_col] = pd.to_numeric(df[time_col], errors='coerce')

    # Drop rows with missing outcome data
    df = df.dropna(subset=[event_col, time_col, '_treatment', '_stratum'])

    # Remove strata with only one group represented
    strata_check = df.groupby('_stratum')['_treatment'].nunique()
    valid_strata = strata_check[strata_check == 2].index
    df = df[df['_stratum'].isin(valid_strata)]

    n_treatment = int((df['_treatment'] == 1).sum())
    n_control = int((df['_treatment'] == 0).sum())
    events_treatment = int(df.loc[df['_treatment'] == 1, event_col].sum())
    events_control = int(df.loc[df['_treatment'] == 0, event_col].sum())

    total_events = events_treatment + events_control
    if len(df) < 10 or total_events < 5:
        return {
            'label': outcome_def['label'],
            'n_treatment': n_treatment,
            'n_control': n_control,
            'events_treatment': events_treatment,
            'events_control': events_control,
            'effect': None, 'ci_lower': None, 'ci_upper': None, 'p': None,
            'error': 'Insufficient data (n < 10 or events < 5)',
        }

    try:
        cph = CoxPHFitter()
        analysis_df = df[['_treatment', time_col, event_col, '_stratum']].copy()
        cph.fit(analysis_df, duration_col=time_col, event_col=event_col, strata=['_stratum'])

        hr = np.exp(cph.params_['_treatment'])
        ci = np.exp(cph.confidence_intervals_.loc['_treatment'].values)
        p_val = cph.summary.loc['_treatment', 'p']

        return {
            'label': outcome_def['label'],
            'n_treatment': n_treatment,
            'n_control': n_control,
            'events_treatment': events_treatment,
            'events_control': events_control,
            'effect': float(hr),
            'ci_lower': float(ci[0]),
            'ci_upper': float(ci[1]),
            'p': float(p_val),
            'error': None,
        }
    except Exception as e:
        return {
            'label': outcome_def['label'],
            'n_treatment': n_treatment,
            'n_control': n_control,
            'events_treatment': events_treatment,
            'events_control': events_control,
            'effect': None, 'ci_lower': None, 'ci_upper': None, 'p': None,
            'error': str(e),
        }


def analyze_outcome_logistic(matched_df, outcome_def, config):
    """Analyze a single outcome using conditional logistic regression."""
    from statsmodels.discrete.conditional_models import ConditionalLogit

    outcome_col = outcome_def['outcome_column']

    df = matched_df.copy()
    df[outcome_col] = pd.to_numeric(df[outcome_col], errors='coerce')
    df = df.dropna(subset=[outcome_col, '_treatment', '_stratum'])

    # Remove strata with only one group represented
    strata_check = df.groupby('_stratum')['_treatment'].nunique()
    valid_strata = strata_check[strata_check == 2].index
    df = df[df['_stratum'].isin(valid_strata)]

    # Remove strata where outcome is constant (all 0 or all 1)
    strata_outcome_var = df.groupby('_stratum')[outcome_col].nunique()
    informative_strata = strata_outcome_var[strata_outcome_var == 2].index
    df_informative = df[df['_stratum'].isin(informative_strata)]

    n_treatment = int((df['_treatment'] == 1).sum())
    n_control = int((df['_treatment'] == 0).sum())
    events_treatment = int(df.loc[df['_treatment'] == 1, outcome_col].sum())
    events_control = int(df.loc[df['_treatment'] == 0, outcome_col].sum())

    total_events = events_treatment + events_control
    if len(df) < 10 or total_events < 5:
        return {
            'label': outcome_def['label'],
            'n_treatment': n_treatment,
            'n_control': n_control,
            'events_treatment': events_treatment,
            'events_control': events_control,
            'effect': None, 'ci_lower': None, 'ci_upper': None, 'p': None,
            'error': 'Insufficient data (n < 10 or events < 5)',
        }

    try:
        model = ConditionalLogit(
            df_informative[outcome_col],
            df_informative[['_treatment']],
            groups=df_informative['_stratum']
        )
        result = model.fit(disp=0)

        or_val = np.exp(result.params['_treatment'])
        ci = np.exp(result.conf_int().loc['_treatment'].values)
        p_val = result.pvalues['_treatment']

        return {
            'label': outcome_def['label'],
            'n_treatment': n_treatment,
            'n_control': n_control,
            'events_treatment': events_treatment,
            'events_control': events_control,
            'effect': float(or_val),
            'ci_lower': float(ci[0]),
            'ci_upper': float(ci[1]),
            'p': float(p_val),
            'error': None,
        }
    except Exception as e:
        return {
            'label': outcome_def['label'],
            'n_treatment': n_treatment,
            'n_control': n_control,
            'events_treatment': events_treatment,
            'events_control': events_control,
            'effect': None, 'ci_lower': None, 'ci_upper': None, 'p': None,
            'error': str(e),
        }


def run_outcome_analyses(matched_df, config):
    """Run all outcome analyses."""
    results = []
    for outcome in config['outcomes']:
        print(f"  Outcome: {outcome['label']}...")
        if outcome_is_survival(outcome, config):
            res = analyze_outcome_cox(matched_df, outcome, config)
            res['effect_label'] = 'HR'
        else:
            res = analyze_outcome_logistic(matched_df, outcome, config)
            res['effect_label'] = 'OR'
        results.append(res)
    return results


# ═══════════════════════════════════════════════════════════════════════════════
# Formatting Helpers
# ═══════════════════════════════════════════════════════════════════════════════

def fmt_p(p):
    """Format p-value."""
    if p is None:
        return "N/A"
    if p < 0.001:
        return "<0.001*"
    elif p < 0.05:
        return f"{p:.3f}*"
    else:
        return f"{p:.3f}"


def fmt_effect(effect, ci_l, ci_u):
    """Format HR/OR with 95% CI using en-dash."""
    if effect is None:
        return "N/A"
    return f"{effect:.2f} ({ci_l:.2f}\u2013{ci_u:.2f})"


def fmt_smd(smd):
    """Format SMD value."""
    if smd is None or np.isnan(smd):
        return "N/A"
    return f"{smd:.3f}"


def fmt_n_events(n, events):
    """Format N (events)."""
    if n is None:
        return "N/A"
    return f"{n} ({events})"


# ═══════════════════════════════════════════════════════════════════════════════
# Excel Output
# ═══════════════════════════════════════════════════════════════════════════════

# Color constants (consistent with other skills)
DARK_BLUE = "1F3864"
HEADER_BG = "2E5FA3"
SECTION_BG = "EBF2FA"
ALT_ROW = "F7FAFD"
SIG_COLOR = "C00000"
BORDER_COLOR = "BFBFBF"
FOOTNOTE_COLOR = "595959"

THIN_BORDER = Border(
    left=Side(style='thin', color=BORDER_COLOR),
    right=Side(style='thin', color=BORDER_COLOR),
    top=Side(style='thin', color=BORDER_COLOR),
    bottom=Side(style='thin', color=BORDER_COLOR),
)


def write_excel(outcome_results, balance_results, match_info, config):
    """Write formatted Excel output with 2 sheets."""
    analysis_type = config['analysis_type']
    effect_label = "HR" if analysis_type == 'cox' else "OR"
    # A mixed outcomes list (survival primary + binary secondary) has no single
    # effect measure, so one column header cannot describe every row. Naming the
    # whole column "HR" when one row is an odds ratio puts a mislabelled number
    # in a manuscript table — the failure this pipeline exists to prevent.
    measures = {r.get('effect_label', effect_label) for r in outcome_results}
    mixed_measures = len(measures) > 1
    column_label = "Effect" if mixed_measures else effect_label
    g1_label = config['grouping']['groups'][0]['label']
    g2_label = config['grouping']['groups'][1]['label']
    output_path = config['output_file']
    abbreviations = config.get('abbreviations', '')

    # Common styles
    title_font = Font(name='Arial', size=12, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type='solid')
    subtitle_font = Font(name='Arial', size=10, italic=True, color="FFFFFF")

    header_font = Font(name='Arial', size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type='solid')

    section_font = Font(name='Arial', size=10, bold=True, color=DARK_BLUE)
    section_fill = PatternFill(start_color=SECTION_BG, end_color=SECTION_BG, fill_type='solid')

    data_font = Font(name='Arial', size=10)
    alt_fill = PatternFill(start_color=ALT_ROW, end_color=ALT_ROW, fill_type='solid')
    no_fill = PatternFill(fill_type=None)

    sig_font = Font(name='Arial', size=10, bold=True, color=SIG_COLOR)
    footnote_font = Font(name='Arial', size=8, italic=True, color=FOOTNOTE_COLOR)

    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', indent=1)

    wb = Workbook()

    # ── Sheet 1: Outcome Analysis ──────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Outcome Analysis"

    # Column widths
    total_cols_s1 = 5
    for c, w in zip(range(1, total_cols_s1 + 1), [32, 18, 18, 26, 12]):
        ws1.column_dimensions[get_column_letter(c)].width = w

    row_idx = 1

    # Title
    ws1.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=total_cols_s1)
    cell = ws1.cell(row=row_idx, column=1, value=config.get('title', 'Propensity Score Matched Analysis'))
    cell.font = title_font
    cell.fill = title_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    for c in range(2, total_cols_s1 + 1):
        ws1.cell(row=row_idx, column=c).fill = title_fill
    row_idx += 1

    # Subtitle
    caliper_str = f"caliper = {match_info['caliper_width']}×SD" if match_info['caliper_enabled'] else "no caliper"
    subtitle_text = config.get('subtitle', '')
    if subtitle_text:
        subtitle_text += " | "
    subtitle_text += (
        f"1:{match_info['ratio']} matching ({caliper_str}) | "
        f"{g1_label}: n={match_info['n_treatment_matched']}, "
        f"{g2_label}: n={match_info['n_control_matched']}"
    )
    ws1.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=total_cols_s1)
    cell = ws1.cell(row=row_idx, column=1, value=subtitle_text)
    cell.font = subtitle_font
    cell.fill = title_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    for c in range(2, total_cols_s1 + 1):
        ws1.cell(row=row_idx, column=c).fill = title_fill
    row_idx += 1

    # Column headers
    headers_s1 = ["Outcome", f"{g1_label}\nN (events)", f"{g2_label}\nN (events)",
                   f"{column_label} (95% CI)", "p-value"]
    for c, h in enumerate(headers_s1, 1):
        cell = ws1.cell(row=row_idx, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = THIN_BORDER
    row_idx += 1

    # Data rows
    for i, res in enumerate(outcome_results):
        fill = alt_fill if i % 2 == 1 else no_fill

        # Outcome label. When the measures differ, each row has to carry its own,
        # otherwise the column header is the only clue and it cannot be right for
        # every row.
        row_label = res['label']
        if mixed_measures:
            row_label = f"{row_label} ({res.get('effect_label', effect_label)})"
        cell = ws1.cell(row=row_idx, column=1, value=row_label)
        cell.font = data_font
        cell.fill = fill
        cell.alignment = left_align
        cell.border = THIN_BORDER

        # Treatment N (events)
        cell = ws1.cell(row=row_idx, column=2, value=fmt_n_events(res['n_treatment'], res['events_treatment']))
        cell.font = data_font
        cell.fill = fill
        cell.alignment = center
        cell.border = THIN_BORDER

        # Control N (events)
        cell = ws1.cell(row=row_idx, column=3, value=fmt_n_events(res['n_control'], res['events_control']))
        cell.font = data_font
        cell.fill = fill
        cell.alignment = center
        cell.border = THIN_BORDER

        # Effect (95% CI)
        cell = ws1.cell(row=row_idx, column=4, value=fmt_effect(res['effect'], res['ci_lower'], res['ci_upper']))
        cell.font = data_font
        cell.fill = fill
        cell.alignment = center
        cell.border = THIN_BORDER

        # p-value
        p_val = res['p']
        cell = ws1.cell(row=row_idx, column=5, value=fmt_p(p_val))
        cell.font = sig_font if (p_val is not None and p_val < 0.05) else data_font
        cell.fill = fill
        cell.alignment = center
        cell.border = THIN_BORDER

        row_idx += 1

    # Footnotes
    row_idx += 1
    if abbreviations:
        ws1.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=total_cols_s1)
        cell = ws1.cell(row=row_idx, column=1, value=abbreviations)
        cell.font = footnote_font
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        row_idx += 1

    # Statistical note
    if mixed_measures:
        note = (f"Time-to-event outcomes: stratified Cox proportional hazards regression "
                f"(stratified by matched pair), HR > 1 indicates higher risk in {g1_label} "
                f"group. Binary outcomes: conditional logistic regression (conditioned on "
                f"matched pair), OR > 1 indicates higher odds in {g1_label} group. The "
                f"measure for each outcome is given in the first column.")
    elif analysis_type == 'cox':
        note = f"Stratified Cox proportional hazards regression (stratified by matched pair). {effect_label} > 1 indicates higher risk in {g1_label} group."
    else:
        note = f"Conditional logistic regression (conditioned on matched pair). {effect_label} > 1 indicates higher odds in {g1_label} group."
    caliper_note = f"Caliper width = {match_info['caliper_width']} × SD(logit PS)." if match_info['caliper_enabled'] else "No caliper applied."
    note += f" Greedy nearest-neighbor matching on logit(PS), 1:{match_info['ratio']}, without replacement. {caliper_note}"
    note += " * p < 0.05."

    ws1.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=total_cols_s1)
    cell = ws1.cell(row=row_idx, column=1, value=note)
    cell.font = footnote_font
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # ── Sheet 2: Balance Assessment ────────────────────────────────────────
    ws2 = wb.create_sheet("Balance Assessment")

    # Column widths: Variable | Before(Treatment Mean, Control Mean, SMD) | After(Treatment Mean, Control Mean, SMD)
    total_cols_s2 = 7
    for c, w in zip(range(1, total_cols_s2 + 1), [32, 20, 20, 10, 20, 20, 10]):
        ws2.column_dimensions[get_column_letter(c)].width = w

    row_idx = 1

    # Title
    ws2.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=total_cols_s2)
    cell = ws2.cell(row=row_idx, column=1, value="Balance Assessment — Before and After Propensity Score Matching")
    cell.font = title_font
    cell.fill = title_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    for c in range(2, total_cols_s2 + 1):
        ws2.cell(row=row_idx, column=c).fill = title_fill
    row_idx += 1

    # Subtitle
    subtitle_s2 = (
        f"Before matching: {g1_label} n={match_info['n_treatment_original']}, {g2_label} n={match_info['n_control_original']} | "
        f"After matching: {g1_label} n={match_info['n_treatment_matched']}, {g2_label} n={match_info['n_control_matched']}"
    )
    ws2.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=total_cols_s2)
    cell = ws2.cell(row=row_idx, column=1, value=subtitle_s2)
    cell.font = subtitle_font
    cell.fill = title_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    for c in range(2, total_cols_s2 + 1):
        ws2.cell(row=row_idx, column=c).fill = title_fill
    row_idx += 1

    # Two-level headers
    # Level 1: merged headers
    ws2.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=4)
    cell = ws2.cell(row=row_idx, column=2, value="Before Matching")
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    for c in range(3, 5):
        ws2.cell(row=row_idx, column=c).fill = header_fill
        ws2.cell(row=row_idx, column=c).border = THIN_BORDER

    ws2.merge_cells(start_row=row_idx, start_column=5, end_row=row_idx, end_column=7)
    cell = ws2.cell(row=row_idx, column=5, value="After Matching")
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    for c in range(6, 8):
        ws2.cell(row=row_idx, column=c).fill = header_fill
        ws2.cell(row=row_idx, column=c).border = THIN_BORDER

    cell = ws2.cell(row=row_idx, column=1, value="")
    cell.fill = header_fill
    cell.border = THIN_BORDER
    row_idx += 1

    # Level 2: sub-headers
    sub_headers = ["Variable", g1_label, g2_label, "SMD", g1_label, g2_label, "SMD"]
    for c, h in enumerate(sub_headers, 1):
        cell = ws2.cell(row=row_idx, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = THIN_BORDER
    row_idx += 1

    # Data rows
    current_section = None
    data_row_idx = 0
    for res in balance_results:
        if res['section'] != current_section:
            current_section = res['section']
            # Section header row
            ws2.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=total_cols_s2)
            cell = ws2.cell(row=row_idx, column=1, value=current_section)
            cell.font = section_font
            cell.fill = section_fill
            cell.alignment = Alignment(horizontal='left', vertical='center')
            for c in range(2, total_cols_s2 + 1):
                ws2.cell(row=row_idx, column=c).fill = section_fill
                ws2.cell(row=row_idx, column=c).border = THIN_BORDER
            cell.border = THIN_BORDER
            row_idx += 1
            data_row_idx = 0

        fill = alt_fill if data_row_idx % 2 == 1 else no_fill

        # Variable label
        cell = ws2.cell(row=row_idx, column=1, value=res['label'])
        cell.font = data_font
        cell.fill = fill
        cell.alignment = left_align
        cell.border = THIN_BORDER

        # Before matching values
        cell = ws2.cell(row=row_idx, column=2, value=res['display_t_before'])
        cell.font = data_font
        cell.fill = fill
        cell.alignment = center
        cell.border = THIN_BORDER

        cell = ws2.cell(row=row_idx, column=3, value=res['display_c_before'])
        cell.font = data_font
        cell.fill = fill
        cell.alignment = center
        cell.border = THIN_BORDER

        # SMD before
        smd_before = res['smd_before']
        cell = ws2.cell(row=row_idx, column=4, value=fmt_smd(smd_before))
        cell.font = Font(name='Arial', size=10, bold=True, color=SIG_COLOR) if (not np.isnan(smd_before) and smd_before >= 0.1) else data_font
        cell.fill = fill
        cell.alignment = center
        cell.border = THIN_BORDER

        # After matching values
        cell = ws2.cell(row=row_idx, column=5, value=res['display_t_after'])
        cell.font = data_font
        cell.fill = fill
        cell.alignment = center
        cell.border = THIN_BORDER

        cell = ws2.cell(row=row_idx, column=6, value=res['display_c_after'])
        cell.font = data_font
        cell.fill = fill
        cell.alignment = center
        cell.border = THIN_BORDER

        # SMD after
        smd_after = res['smd_after']
        cell = ws2.cell(row=row_idx, column=7, value=fmt_smd(smd_after))
        cell.font = Font(name='Arial', size=10, bold=True, color=SIG_COLOR) if (not np.isnan(smd_after) and smd_after >= 0.1) else data_font
        cell.fill = fill
        cell.alignment = center
        cell.border = THIN_BORDER

        row_idx += 1
        data_row_idx += 1

    # Footnotes
    row_idx += 1
    if abbreviations:
        ws2.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=total_cols_s2)
        cell = ws2.cell(row=row_idx, column=1, value=abbreviations)
        cell.font = footnote_font
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        row_idx += 1

    note_s2 = (
        "Continuous variables: mean \u00b1 SD. Binary variables: n (%). "
        "SMD = standardized mean difference. SMD < 0.1 indicates good balance (shown in red bold if \u2265 0.1)."
    )
    ws2.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=total_cols_s2)
    cell = ws2.cell(row=row_idx, column=1, value=note_s2)
    cell.font = footnote_font
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # Save
    wb.save(output_path)
    print(f"Excel saved: {output_path}")

    # Also save CSV of outcome results
    csv_path = os.path.splitext(output_path)[0] + '.csv'
    rows = []
    for res in outcome_results:
        rows.append({
            'Outcome': res['label'],
            f'{g1_label} N': res['n_treatment'],
            f'{g1_label} Events': res['events_treatment'],
            f'{g2_label} N': res['n_control'],
            f'{g2_label} Events': res['events_control'],
            f"{res.get('effect_label', effect_label)}": res['effect'],
            '95% CI Lower': res['ci_lower'],
            '95% CI Upper': res['ci_upper'],
            f"{res.get('effect_label', effect_label)} (95% CI)": fmt_effect(res['effect'], res['ci_lower'], res['ci_upper']),
            'p-value': res['p'],
            'p-value (formatted)': fmt_p(res['p']),
        })
    pd.DataFrame(rows).to_csv(csv_path, index=False, encoding='utf-8-sig')
    print(f"CSV saved: {csv_path}")

    # Save balance CSV
    balance_csv_path = os.path.splitext(output_path)[0] + '_balance.csv'
    balance_rows = []
    for res in balance_results:
        balance_rows.append({
            'Section': res['section'],
            'Variable': res['label'],
            'Type': res['type'],
            f'{g1_label} Before': res['display_t_before'],
            f'{g2_label} Before': res['display_c_before'],
            'SMD Before': res['smd_before'],
            f'{g1_label} After': res['display_t_after'],
            f'{g2_label} After': res['display_c_after'],
            'SMD After': res['smd_after'],
        })
    pd.DataFrame(balance_rows).to_csv(balance_csv_path, index=False, encoding='utf-8-sig')
    print(f"Balance CSV saved: {balance_csv_path}")


# ═══════════════════════════════════════════════════════════════════════════════
# Love Plot
# ═══════════════════════════════════════════════════════════════════════════════

def write_love_plot(balance_results, config):
    """Generate a Love plot showing SMD before and after matching."""
    output_base = os.path.splitext(config['output_file'])[0]
    output_img = f"{output_base}_loveplot.png"

    n_vars = len(balance_results)
    if n_vars == 0:
        print("  No variables for Love plot.")
        return

    # Collect data for plotting (reverse order so first variable appears at top)
    labels = []
    smd_before = []
    smd_after = []
    sections = []
    section_breaks = []

    current_section = None
    for i, res in enumerate(reversed(balance_results)):
        if res['section'] != current_section:
            current_section = res['section']
            section_breaks.append(i)
            sections.append(current_section)
        labels.append(res['label'])
        smd_before.append(res['smd_before'] if not np.isnan(res['smd_before']) else 0)
        smd_after.append(res['smd_after'] if not np.isnan(res['smd_after']) else 0)

    y_positions = np.arange(n_vars)

    # Figure size
    fig_height = max(5, n_vars * 0.45 + 2)
    fig_width = 10

    fig, ax = plt.subplots(figsize=(fig_width, fig_height))

    # Draw connecting lines
    for i in range(n_vars):
        ax.plot([smd_before[i], smd_after[i]], [y_positions[i], y_positions[i]],
                color='#BFBFBF', linewidth=0.8, linestyle='--', zorder=1)

    # Plot SMD before (red circles)
    ax.scatter(smd_before, y_positions, color='#C00000', marker='o', s=60,
               zorder=3, label='Before Matching', edgecolors='white', linewidths=0.5)

    # Plot SMD after (blue squares)
    ax.scatter(smd_after, y_positions, color='#2E5FA3', marker='s', s=60,
               zorder=3, label='After Matching', edgecolors='white', linewidths=0.5)

    # Threshold line at 0.1
    ax.axvline(x=0.1, color='#555555', linestyle='--', linewidth=1, zorder=0)
    # Add threshold label
    max_smd = max(max(smd_before), max(smd_after), 0.15)
    ax.text(0.1, n_vars - 0.2, 'SMD = 0.1', fontsize=8, color='#555555',
            ha='left', va='bottom', fontstyle='italic')

    # Section labels
    for i, (break_idx, section_name) in enumerate(zip(section_breaks, sections)):
        y = y_positions[break_idx]
        # Add a subtle separator above each section (except the first)
        if i > 0:
            sep_y = (y_positions[break_idx] + y_positions[break_idx - 1]) / 2
            ax.axhline(y=sep_y, color='#EEEEEE', linewidth=0.5, zorder=0)

    # Configure axes
    ax.set_yticks(y_positions)
    ax.set_yticklabels(labels, fontsize=9)
    ax.set_xlabel('Absolute Standardized Mean Difference', fontsize=11, fontweight='bold')
    ax.set_xlim(left=0, right=max_smd * 1.15 + 0.02)
    ax.set_ylim(-0.8, n_vars - 0.2)

    # Remove top and right spines
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)

    # Grid
    ax.xaxis.grid(True, alpha=0.3, linestyle='-', linewidth=0.5)
    ax.set_axisbelow(True)

    # Legend
    ax.legend(loc='lower right', fontsize=9, framealpha=0.9, edgecolor='#BFBFBF')

    # Title
    title = config.get('title', 'Love Plot — SMD Before and After PSM')
    fig.suptitle(title, fontsize=13, fontweight='bold', color='#1F3864', y=0.98)

    plt.tight_layout(rect=[0, 0, 1, 0.96])
    plt.savefig(output_img, dpi=300, bbox_inches='tight', facecolor='white', edgecolor='none')
    plt.close()
    print(f"Love plot saved: {output_img}")


# ═══════════════════════════════════════════════════════════════════════════════
# Contrast sidecar
# ═══════════════════════════════════════════════════════════════════════════════

def _sidecar_analysis_frame(matched_df, outcome, config):
    """The rows an outcome model is fitted on, rebuilt from the matched cohort.

    This deliberately re-derives the analysis frame instead of reading the
    numbers out of `run_outcome_analyses()`. The sidecar's whole purpose is to
    be a SECOND, independent statement of which arm the estimate belongs to; a
    sidecar computed from the same dict that fills the output table would agree
    with the table unconditionally and could never catch anything.

    The drop rules mirror analyze_outcome_cox / analyze_outcome_logistic up to
    the point where each of those computes its per-arm n and events — i.e.
    before the informative-strata filter, which the logistic model applies to
    its fit but not to its reported counts.
    """
    survival = outcome_is_survival(outcome, config)
    value_col = outcome['event_column'] if survival else outcome['outcome_column']
    needed = [value_col, '_treatment', '_stratum']

    df = matched_df.copy()
    df[value_col] = pd.to_numeric(df[value_col], errors='coerce')
    if survival:
        time_col = outcome['time_column']
        df[time_col] = pd.to_numeric(df[time_col], errors='coerce')
        needed.append(time_col)
    df = df.dropna(subset=needed)

    strata_check = df.groupby('_stratum')['_treatment'].nunique()
    df = df[df['_stratum'].isin(strata_check[strata_check == 2].index)]
    return df, value_col


def _sha256(path):
    """Hex digest of a file, or None when it cannot be read."""
    h = hashlib.sha256()
    try:
        with open(path, 'rb') as f:
            for chunk in iter(lambda: f.read(1 << 20), b''):
                h.update(chunk)
    except OSError as exc:
        print(f"  [SKIP] could not hash {os.path.basename(path)}: "
              f"{type(exc).__name__}: {exc}", file=sys.stderr)
        return None
    return h.hexdigest()


def write_contrast_sidecar(matched_df, config, config_path=None):
    """Write <stem>_contrast.json — the producer's own statement of the contrast.

    Everything here comes from the dataframe the outcome models are fitted on,
    keyed on the `_treatment` indicator that enters the model as the exposure
    coefficient. `exposed_value` is read back out of the raw grouping column of
    the treated rows rather than copied from the config, so the file records
    what the model was actually fitted on rather than what the plan intended.

    Consumed by research-plan/scripts/precheck_configs.py, which blocks when
    `exposed_value` disagrees with `grouping.groups[0].value` or when the
    per-arm counts disagree with the ones printed in the output table.
    """
    output_path = config['output_file']
    sidecar_path = os.path.splitext(output_path)[0] + '_contrast.json'

    grouping = config['grouping']
    group_col = grouping['column']
    exposed_label = grouping['groups'][0]['label']
    control_label = grouping['groups'][1]['label']
    declared_value = grouping['groups'][0].get('value')

    treated = matched_df[matched_df['_treatment'] == 1]
    controls = matched_df[matched_df['_treatment'] == 0]

    # Read the exposure value back out of the data the model was fitted on. If
    # the treated rows do not all carry one value of the grouping column then
    # the indicator the model uses does not identify a group at all, and nothing
    # downstream can be trusted to be about the contrast it claims to be about.
    observed = pd.to_numeric(treated[group_col], errors='coerce').dropna().unique()
    if len(observed) == 0:
        observed = treated[group_col].dropna().unique()
    if len(observed) == 1:
        raw = observed[0]
        exposed_value = (float(raw) if isinstance(raw, (int, float, np.number))
                         else str(raw))
    else:
        exposed_value = None
        print(f"[QC_BLOCKING] the treated arm of the matched cohort carries "
              f"{len(observed)} distinct values of '{group_col}' "
              f"({list(observed)[:5]}); the model's exposure indicator does not "
              f"identify a single group", file=sys.stderr)

    if exposed_value is not None and declared_value is not None:
        same = (exposed_value == declared_value
                or str(exposed_value) == str(declared_value))
        if not same:
            print(f"[QC_BLOCKING] the matched cohort's treated arm is "
                  f"{group_col}={exposed_value!r} but grouping.groups[0].value is "
                  f"{declared_value!r}; the fitted contrast is not the declared one",
                  file=sys.stderr)

    payload = {
        # Binding fields: which table this describes, and proof it is that exact
        # table. Without them a sidecar is just a file in a directory — a
        # consumer looking for "the PSM sidecar" could validate one analysis's
        # table against another analysis's contrast and pass.
        'sidecar_version': 1,
        'table_basename': os.path.basename(output_path),
        'table_sha256': _sha256(output_path),
        'config_sha256': _sha256(config_path) if config_path else None,
        'exposed_label': exposed_label,
        'exposed_value': exposed_value,
        'control_label': control_label,
        'control_value': grouping['groups'][1].get('value'),
        'group_column': group_col,
        'source': 'matched cohort as fitted (_treatment indicator)',
        'matched_n': {'exposed': int(len(treated)), 'control': int(len(controls))},
        'outcomes': [],
    }

    processed = skipped = 0
    for outcome in config['outcomes']:
        try:
            frame, value_col = _sidecar_analysis_frame(matched_df, outcome, config)
        except Exception as exc:
            skipped += 1
            print(f"  [SKIP] contrast sidecar for outcome "
                  f"{outcome.get('label')!r}: {type(exc).__name__}: {exc}",
                  file=sys.stderr)
            continue
        processed += 1
        exposed_rows = frame[frame['_treatment'] == 1]
        control_rows = frame[frame['_treatment'] == 0]
        payload['outcomes'].append({
            'label': outcome.get('label'),
            'exposed': {'n': int(len(exposed_rows)),
                        'events': int(exposed_rows[value_col].sum())},
            'control': {'n': int(len(control_rows)),
                        'events': int(control_rows[value_col].sum())},
        })

    # The primary outcome is promoted to the top level: that is the row the
    # consumer checks the output table against, and it keeps the file readable.
    if payload['outcomes']:
        payload['exposed'] = payload['outcomes'][0]['exposed']
        payload['control'] = payload['outcomes'][0]['control']
        payload['primary_outcome'] = payload['outcomes'][0]['label']

    with open(sidecar_path, 'w', encoding='utf-8') as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)
    print(f"Contrast sidecar saved: {sidecar_path}")
    print(f"  contrast sidecar: processed {processed} / skipped {skipped} / "
          f"total {processed + skipped} outcome(s)")
    return sidecar_path


# ═══════════════════════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════════════════════

def _apply_academic_style():
    """academic-style integration (Wave 2). Consumes --style=<journal> from
    sys.argv. No-op when skill not installed or flag absent."""
    from pathlib import Path as _P
    _as = _P.home() / ".claude" / "skills" / "academic-style"
    if not _as.exists():
        return None
    if str(_as) not in sys.path:
        sys.path.insert(0, str(_as))
    try:
        from core.biostats_integration import apply_style_from_argv
    except ImportError:
        return None
    return apply_style_from_argv(sys.argv)


def main():
    _AS_STYLE = _apply_academic_style()
    if _AS_STYLE:
        print(f"  Applied {_AS_STYLE} journal style")

    if len(sys.argv) < 2:
        print("Usage: python psm_matching.py config.json [--check-missing] [--style=<journal>]")
        sys.exit(1)

    config_path = sys.argv[1]
    check_missing_only = '--check-missing' in sys.argv

    config = load_config(config_path)
    print("Loading data...")
    df = load_data(config)
    grp1, grp2 = split_groups(df, config)

    g1_label = config['grouping']['groups'][0]['label']
    g2_label = config['grouping']['groups'][1]['label']
    print(f"Overall: n={len(df)}")
    print(f"{g1_label}: n={len(grp1)}")
    print(f"{g2_label}: n={len(grp2)}")

    if check_missing_only:
        report = check_missing(df, grp1, grp2, config)
        print(report)
        sys.exit(0)

    # Prepare data (listwise deletion on matching variables)
    print("\nPreparing data...")
    df_clean = prepare_analysis_df(df, config)
    grp1_clean, grp2_clean = split_groups(df_clean, config)
    print(f"After preparation: {g1_label} n={len(grp1_clean)}, {g2_label} n={len(grp2_clean)}")

    # Step 1: Estimate propensity scores
    print("\nEstimating propensity scores...")
    df_ps = estimate_propensity_scores(df_clean, config)

    # Step 2: Perform matching
    print("\nPerforming matching...")
    matched_df, match_info = perform_matching(df_ps, config)

    # Step 3: Compute balance table
    print("\nComputing balance table...")
    matched_grp1 = matched_df[matched_df['_treatment'] == 1]
    matched_grp2 = matched_df[matched_df['_treatment'] == 0]
    balance_results = compute_balance_table(grp1_clean, grp2_clean, matched_grp1, matched_grp2, config)

    # Step 4: Analyze outcomes
    print("\nAnalyzing outcomes...")
    outcome_results = run_outcome_analyses(matched_df, config)

    # Step 5: Write outputs
    print("\nGenerating outputs...")
    write_excel(outcome_results, balance_results, match_info, config)
    write_love_plot(balance_results, config)
    # Written from `matched_df`, not from `outcome_results`, so it is a second
    # independent statement of which arm the estimate is about.
    write_contrast_sidecar(matched_df, config, config_path)

    # Summary
    print("\n" + "=" * 60)
    print("ANALYSIS COMPLETE")
    print("=" * 60)
    print(f"  Original: {g1_label} n={match_info['n_treatment_original']}, {g2_label} n={match_info['n_control_original']}")
    print(f"  Matched:  {g1_label} n={match_info['n_treatment_matched']}, {g2_label} n={match_info['n_control_matched']}")
    print(f"  Unmatched treatment subjects: {match_info['n_unmatched']}")
    if match_info['caliper_enabled']:
        print(f"  Caliper: {match_info['caliper_width']} × SD(logit PS) = {match_info['caliper_value']:.4f}")
    else:
        print(f"  Caliper: disabled")
    print(f"  Ratio: 1:{match_info['ratio']}")

    sig_count = sum(1 for r in outcome_results if r['p'] is not None and r['p'] < 0.05)
    total = len(outcome_results)
    print(f"  Significant outcomes: {sig_count}/{total} (p < 0.05)")

    failed = [r['label'] for r in outcome_results if r['error'] is not None]
    if failed:
        print(f"  Outcomes with errors: {', '.join(failed)}")

    # Balance summary
    max_smd_after = max((r['smd_after'] for r in balance_results if not np.isnan(r['smd_after'])), default=0)
    all_balanced = all(r['smd_after'] < 0.1 for r in balance_results if not np.isnan(r['smd_after']))
    if all_balanced:
        print(f"  Balance: All variables SMD < 0.1 after matching (max SMD = {max_smd_after:.3f})")
    else:
        imbalanced = [r['label'] for r in balance_results if not np.isnan(r['smd_after']) and r['smd_after'] >= 0.1]
        print(f"  Balance: {len(imbalanced)} variable(s) with SMD >= 0.1: {', '.join(imbalanced)}")

    print("=" * 60)

    # ── QC Warnings ──────────────────────────────────────────────────────────
    # Track blocking emissions and sys.exit(2) at the end so the pipeline sees a
    # non-zero exit when [QC_BLOCKING] is raised. Previously [QC_BLOCKING] was
    # emitted but exit was always 0, making the pipeline think the step succeeded
    # with warnings only.
    #
    # NOTE (intentional difference from table1): table1 runs its blocking check in
    # `pre_write_validation()` BEFORE write_excel/write_csv, so a blocking failure
    # leaves no artifact on disk. This script validates AFTER writing its outputs —
    # the QC checks depend on the matching result and fitted outcome models, so the
    # files already exist when we discover the problem. Consequence: on exit 2 the
    # output files DO exist; downstream consumers must key off the exit code, not
    # file presence. Kept as-is deliberately; do not reorder the flow for this.
    qc_blocking_count = 0

    def _qc_block(msg):
        nonlocal qc_blocking_count
        print(f"[QC_BLOCKING] {msg}", file=sys.stderr)
        qc_blocking_count += 1

    try:
        # Check matched pairs == 0
        if match_info['n_treatment_matched'] == 0:
            _qc_block("Matched pairs = 0 (no subjects matched)")

        # Check model convergence — if all outcomes failed, signal convergence issue
        all_outcome_errors = [r for r in outcome_results if r.get('error') is not None]
        if len(all_outcome_errors) == len(outcome_results) and len(outcome_results) > 0:
            _qc_block("Model did not converge for any outcome")

        # Check max SMD > 0.1 after matching
        if max_smd_after > 0.1:
            imbalanced_vars = [r['label'] for r in balance_results if not np.isnan(r['smd_after']) and r['smd_after'] >= 0.1]
            print(f"[QC_WARNING] Max SMD after matching = {max_smd_after:.3f} > 0.1; imbalanced: {', '.join(imbalanced_vars)}", file=sys.stderr)
    except Exception as e:
        # Surface the QC failure instead of silently swallowing it: a QC block that
        # itself crashes must not look like a clean run to the pipeline.
        print(f"[QC_WARNING] QC block raised: {type(e).__name__}: {e}", file=sys.stderr)

    if qc_blocking_count > 0:
        sys.exit(2)


if __name__ == '__main__':
    main()
