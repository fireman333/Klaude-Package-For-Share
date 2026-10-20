#!/usr/bin/env python3
r"""
STROBE domain report builder — the deterministic replacement for the reviewer-challenge
STROBE agent (domain 10).

Why this exists
---------------
`research-plan/scripts/strobe.py` already walks all 22/34 STROBE items and resolves each
one to a concrete artifact on disk, with a status of covered / partial / missing / n/a.
Sending an LLM agent to *read that checklist and report what it says* is a second pass
over a finished answer. What the agent cannot supply is the one thing STROBE actually
requires: a named author asserting that the human-written items were written. So the
report becomes script output plus a sign-off.

Contract
--------
  python strobe_signoff.py --checklist <Strobe.md|Strobe.xlsx> \
      [--rows-json <Strobe_rows.json>] [--signed-off-by "Name"] [--out report.json]

Input preference: `strobe.py` writes `<stem>_rows.json` next to the table, and that is
read first. Parsing the markdown table is the fallback for projects whose disk predates
that output — a brittle coupling (reorder a column in `write_md()` and the parse goes
wrong without erroring), so taking it prints a notice. It is never a silent downgrade.

The JSON also carries each row's `kind`, the only marker for the items no pipeline
can produce. The markdown does not carry it, and this script keeps **no local copy**
of that classification — a second copy drifts the moment strobe.py adds an item, and
it drifts silently in the direction that hides a missing artifact behind a sign-off.
So a report built from the fallback **cannot be a pass**: it reports what it can see
and blocks on the fact that it cannot tell author items from pipeline items. The
fallback exists so old projects do not crash, not so they can be signed off.

Provenance (schema `strobe_domain_report/1`)
--------------------------------------------
This report is aggregated as domain 10, so a stale `verdict: pass` at the same path
would be merged as this round's result. Three defences: the output is **deleted
before any work** and written **atomically**; it carries `schema` / `run_id` /
`checklist_sha256`; and it carries `plan_fingerprint`, delegated to
`reviewer_challenge.plan_fingerprint` so producer and consumer cannot disagree. The
plan is located from `--plan`, else the path strobe.py recorded in the rows JSON,
else `research_plan.json` beside the output dir or checklist. When it cannot be
established the report says so — `plan_fingerprint: null` plus
`plan_fingerprint_error` — which `provenance_problems()` blocks on. Unverifiable is
not clean, and the field is never simply left out.

The output is a **reviewer-challenge domain report**, written to the shape
`reviewer_challenge.py aggregate` validates: `domain` = `reporting_strobe` (the key in
its DOMAINS list), a top-level `verdict`, and `comments[]` each carrying
verdict / severity / comment / remedy / status / evidence. Feed it to:

  python reviewer_challenge.py aggregate --plan <plan> --reports <this> <the other 9> \
      --out <challenge.json>

It is **not** a guardian RoleReport — do not feed it to `runner.py aggregate`.

Getting `domain` wrong is not a cosmetic error: an unrecognised key fails schema
validation, the domain counts as unreported, and the manuscript gate blocks with
nothing to point at. That is precisely what `"D10_reporting_strobe"` did before
2026-08-12 — a value that read correctly to a human and matched nothing.

Exit 0  = rows were classified (`strobe_rows/1`), every actionable item is
          covered/partial, and the sign-off was supplied where author items are owed.
Exit 1  = actionable items are missing, author items are unsigned, OR the rows could
          not be classified at all (fallback input). Exit 0 is unreachable from the
          fallback by construction.
Exit 2  = the checklist could not be read at all (never a pass — a checklist that was
          not produced is not a checklist that passed).
"""

import argparse
import hashlib
import json
import os
import sys
import uuid
from datetime import datetime, timezone

# Bumped when the report shape changes in a way a consumer has to notice.
REPORT_SCHEMA = "strobe_domain_report/1"

STATUSES = ("covered", "partial", "missing", "n/a")

# There is deliberately NO local list of "items only an author can write".
#
# There used to be one, mirroring the rows marked `'human'` in strobe.py's ITEMS
# table, used whenever the structured rows were unavailable. It was wrong in both
# directions on its first outing — it excused 1a/3/4/11/19 (items the pipeline does
# resolve, so a real gap in them would have been waved through as "the author's
# prose") and it claimed 5/8/10 (setting, data sources, sample-size rationale), so
# items no pipeline can produce were reported as blocking defects.
#
# Fixing that list's contents did not fix the mechanism. A second copy of a
# classification drifts the moment strobe.py adds an item, and it fails silently in
# the direction that matters: a machine artifact misfiled as "author's prose"
# disappears behind a sign-off. So the copy is gone. `kind` has exactly one source
# (the `strobe_rows/1` JSON); when that source is absent this script cannot
# classify, and therefore must not be able to return a pass — see build_report().


def clear_stale_output(path):
    """Delete the report before any work. This one is read by
    `reviewer_challenge.py aggregate` as domain 10, so a stale `verdict: pass` left
    by a previous run would be aggregated as though this run had produced it."""
    if not path:
        return
    try:
        os.remove(path)
    except FileNotFoundError:
        pass
    except OSError as e:
        print(f"[GUARDIAN_BLOCKING] could not remove the previous report at {path} "
              f"({type(e).__name__}: {e}); refusing to run.", file=sys.stderr)
        sys.exit(2)


def write_json_atomic(path, payload):
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    tmp = f"{path}.{os.getpid()}.tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)
        f.write("\n")
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)


def plan_fingerprint_of(plan):
    """(fingerprint, error) via reviewer_challenge, so the two cannot disagree.

    Never computed locally: a second implementation would drift from the consumer's
    and produce mismatches that mean nothing — the same class of bug as the
    human-item mirror this file used to carry.
    """
    rc_path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
        "research-plan", "scripts", "reviewer_challenge.py")
    try:
        import importlib.util
        spec = importlib.util.spec_from_file_location("_rc_fp", rc_path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.plan_fingerprint(plan), None
    except Exception as e:
        return None, f"{type(e).__name__}: {e}"


def resolve_plan_path(explicit, rows_payload, checklist_path):
    """Where the plan for this checklist lives. Returns (path, how) or (None, why).

    Order: an explicit --plan; then the plan strobe.py recorded in the rows JSON
    (`artifacts.plan` — the plan it actually built the checklist from, so it is the
    right answer rather than a guess); then research_plan.json beside the recorded
    output dir or beside the checklist.
    """
    if explicit:
        if os.path.exists(explicit):
            return explicit, "--plan"
        return None, f"--plan {explicit} does not exist"

    art = (rows_payload or {}).get("artifacts") or {}
    recorded = art.get("plan")
    if recorded and os.path.exists(recorded):
        return recorded, "strobe_rows.json$.artifacts.plan"

    for base, how in ((art.get("output_dir"), "artifacts.output_dir"),
                      (os.path.dirname(os.path.abspath(checklist_path)), "checklist dir")):
        if not base:
            continue
        cand = os.path.join(base, "research_plan.json")
        if os.path.exists(cand):
            return cand, how
    return None, ("no plan found: pass --plan, or keep research_plan.json next to the "
                  "output dir (strobe.py records its path in strobe_rows.json)")


def sha256_of(path):
    try:
        h = hashlib.sha256()
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 16), b""):
                h.update(chunk)
        return h.hexdigest()[:16]
    except OSError:
        return None


def parse_rows_json(path):
    """Read strobe.py's `strobe_rows/1` output. Preferred over parsing the table."""
    with open(path, "r", encoding="utf-8") as f:
        payload = json.load(f)
    schema = payload.get("schema", "")
    if not str(schema).startswith("strobe_rows/"):
        raise ValueError(f"{path}: unexpected schema {schema!r}")
    rows = []
    for r in payload.get("rows") or []:
        status = r.get("status")
        if status not in STATUSES:
            raise ValueError(f"{path}: row {r.get('item')!r} has status {status!r}")
        rows.append({"item": str(r.get("item", "")), "section": r.get("section", ""),
                     "requirement": r.get("requirement", ""), "status": status,
                     "evidence": r.get("evidence", ""), "kind": r.get("kind")})
    return rows, payload


def is_human_row(row):
    """True only when strobe.py itself marked the row `human`.

    A row with no `kind` is not "probably machine" — it is unclassified, and the
    caller must refuse to produce a verdict rather than guess.
    """
    return row.get("kind") == "human"


def rows_are_classified(rows):
    """Every row carries strobe.py's own `kind`. False on any fallback input."""
    return bool(rows) and all(r.get("kind") is not None for r in rows)


def parse_md(path):
    """Parse the markdown checklist table written by strobe.py's write_md().

    Fallback only — positional, so a column reorder in write_md() would misread
    silently. Rows carry no `kind`, which is why a report built from this input can
    never be a pass.
    """
    rows = []
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            if not line.startswith("|"):
                continue
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if len(cells) < 5 or cells[0] in ("#", "---"):
                continue
            status = next((s for s in STATUSES if s in cells[3]), None)
            if status is None:
                continue
            rows.append({"item": cells[0], "section": cells[1], "requirement": cells[2],
                         "status": status, "evidence": cells[4]})
    return rows


def parse_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for r in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in r]
        if len(cells) < 5:
            continue
        status = next((s for s in STATUSES if s in cells[3]), None)
        if status is None:
            continue
        rows.append({"item": cells[0], "section": cells[1], "requirement": cells[2],
                     "status": status, "evidence": cells[4]})
    return rows


def build_report(rows, signed_off_by, checklist_path, source="rows_json"):
    """Emit a domain report in the shape `reviewer_challenge.py aggregate` accepts.

    The consumer's contract (reviewer_challenge.py validate_report) is:
      domain    one of DOMAINS — for this one, exactly "reporting_strobe"
      verdict   pass | concern | fatal, and it must be carried by the comments:
                `fatal` needs an open fatal comment, `pass` needs a non-empty
                `checked` and no open comments
      comments[] each with verdict / severity(major|minor) / comment / remedy /
                status(open|fixed|wontfix) / evidence — an empty evidence string
                silently demotes a fatal to concern, so every comment names its
                checklist row
    Guardian-side fields (counts, signed_off_by, input_source, checklist) ride
    along as extra keys; the validator ignores what it does not know.
    """
    counts = {s: sum(1 for r in rows if r["status"] == s) for s in STATUSES}

    if not rows_are_classified(rows):
        # No `kind` on the rows -> we cannot tell "the pipeline owes an artifact"
        # from "the author owes a paragraph". Every verdict below depends on that
        # split, so the honest output is a block naming what is missing without
        # pretending to categorise it. A sign-off cannot rescue this: signing for
        # the author items is meaningless when we cannot say which items those are.
        missing = [r for r in rows if r["status"] == "missing"]
        listed = ", ".join(r["item"] for r in missing) or "（無 missing 項目）"
        return {
            "domain": "reporting_strobe",
            "verdict": "fatal",
            "checked": (f"只讀到未分類的 checklist（{source}），共 {len(rows)} 項："
                        f"covered {counts['covered']}／partial {counts['partial']}／"
                        f"missing {counts['missing']}／n/a {counts['n/a']}。"
                        f"缺 `kind`，無法區分「流程該產而沒產」與「本來就該作者寫」。"),
            "comments": [{
                "id": "STROBE-UNCLASSIFIED",
                "check_id": "STROBE-UNCLASSIFIED",
                "verdict": "fatal",
                "severity": "major",
                "status": "open",
                "comment": (f"checklist 讀自 {source}，沒有 strobe.py 的 `kind` 欄位，"
                            f"無法判斷哪些項目只能由作者撰寫。這種輸入下本檢查"
                            f"**不產生通過**——分不出類別時的簽名等於簽一份不知道內容的清單。"
                            f"目前 missing 的項目（未分類）：{listed}。"),
                "remedy": ("重跑 research-plan 的 `strobe` step 產生 `<stem>_rows.json`"
                           "（`strobe_rows/1`，每列帶 `kind`），再跑一次本檢查。"),
                "evidence": f"{checklist_path}#unclassified_rows={len(rows)}",
            }],
            "produced_by": "biostats-guardian/scripts/strobe_signoff.py",
            "method": "deterministic (strobe.py checklist) + author sign-off",
            "input_source": source,
            "rows_classified": False,
            "checklist": checklist_path,
            "counts": counts,
            "n_items": len(rows),
            "signed_off_by": signed_off_by,
            "summary": (f"STROBE：{source} 無 `kind` 欄位，無法分類，本檢查不予通過"
                        f"（covered {counts['covered']}／partial {counts['partial']}／"
                        f"missing {counts['missing']}／n/a {counts['n/a']}，共 {len(rows)} 項）。"),
            "completed_at": datetime.now(timezone.utc).isoformat(),
        }

    human_rows = [r for r in rows if is_human_row(r)]
    actionable_missing = [r for r in rows
                          if r["status"] == "missing" and not is_human_row(r)]
    human_missing = [r for r in human_rows if r["status"] == "missing"]

    comments = []
    for r in actionable_missing:
        comments.append({
            "id": f"STROBE-{r['item']}",
            "check_id": "STROBE-ITEM",
            "verdict": "fatal",
            "severity": "major",
            "status": "open",
            "comment": (f"STROBE {r['item']}（{r['section']}）沒有任何產物對應："
                        f"{r['evidence']}"),
            "remedy": ("補跑對應的 pipeline step 產出該項的產物，或在 decision_log.md "
                       "開一列說明為何本研究不產生此項（wontfix 需附 decision_log_id）。"),
            "evidence": f"{checklist_path}#{r['item']}",
        })

    if human_missing and not signed_off_by:
        items = ", ".join(r["item"] for r in human_missing)
        comments.append({
            "id": "STROBE-SIGNOFF",
            "check_id": "STROBE-SIGNOFF",
            "verdict": "fatal",
            "severity": "major",
            "status": "open",
            "comment": (f"{len(human_missing)} 個 STROBE 項目只能由作者撰寫（{items}），"
                        f"但沒有作者簽名。腳本能證明產物齊備，不能證明文字寫了。"),
            "remedy": ("作者確認這些段落已寫入稿件後，用 "
                       "`strobe_signoff.py --signed-off-by \"姓名\"` 重跑。"),
            "evidence": f"{checklist_path}#human_items={items}",
        })

    verdict = "fatal" if comments else "pass"
    checked = (f"strobe.py 以 {len(rows)} 個 STROBE 項目逐項比對磁碟產物："
               f"covered {counts['covered']}／partial {counts['partial']}／"
               f"missing {counts['missing']}／n/a {counts['n/a']}；"
               f"其中 {len(human_rows)} 項標記為只能由作者撰寫。"
               f"來源：{checklist_path}（{source}）。"
               + (f"作者簽名：{signed_off_by}。" if signed_off_by else "尚未取得作者簽名。"))

    return {
        # Must match reviewer_challenge.py DOMAINS exactly — an id that only looks
        # right ("D10_reporting_strobe") fails schema validation, the domain counts
        # as unreported, and the manuscript gate blocks with nothing to point at.
        "domain": "reporting_strobe",
        "verdict": verdict,
        "checked": checked,
        "comments": comments,
        # ── guardian-side extras (ignored by the aggregate validator) ──
        "produced_by": "biostats-guardian/scripts/strobe_signoff.py",
        "method": "deterministic (strobe.py checklist) + author sign-off",
        "input_source": source,          # rows_json | markdown_fallback | xlsx_fallback
        "rows_classified": True,
        "checklist": checklist_path,
        "counts": counts,
        "n_items": len(rows),
        "signed_off_by": signed_off_by,
        "summary": (f"STROBE：covered {counts['covered']}／partial {counts['partial']}／"
                    f"missing {counts['missing']}／n/a {counts['n/a']}，共 {len(rows)} 項。"
                    + (f"作者簽名：{signed_off_by}。" if signed_off_by else "尚未取得作者簽名。")),
        "completed_at": datetime.now(timezone.utc).isoformat(),
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--checklist", required=True,
                    help="Strobe .md or .xlsx from strobe.py (fallback input)")
    ap.add_argument("--rows-json", default=None,
                    help="strobe.py's <stem>_rows.json; preferred input. "
                         "Default: derived from --checklist")
    ap.add_argument("--signed-off-by", default=None,
                    help="author asserting the human-written STROBE items are in the manuscript")
    ap.add_argument("--plan", default=None,
                    help="research_plan.json this checklist belongs to. Default: the "
                         "path strobe.py recorded in <stem>_rows.json, else "
                         "research_plan.json beside the output dir / checklist")
    ap.add_argument("--out", default=None, help="write the domain report JSON here")
    args = ap.parse_args()
    run_id = uuid.uuid4().hex[:12]
    clear_stale_output(args.out)

    # Preferred input: the structured rows. Parsing the table is the fallback and
    # says so out loud — a silent downgrade to a brittle parser is how a column
    # reorder becomes a wrong sign-off nobody notices.
    rows_json = args.rows_json or (os.path.splitext(args.checklist)[0] + "_rows.json")
    rows, source, read_from, why = None, None, None, "not found"
    rows_payload = None
    if os.path.exists(rows_json):
        try:
            rows, rows_payload = parse_rows_json(rows_json)
            source, read_from = "rows_json", rows_json
        except Exception as e:
            why = f"unreadable ({type(e).__name__}: {e})"

    if rows is None:
        if not os.path.exists(args.checklist):
            print(f"[GUARDIAN_BLOCKING] {rows_json} is {why} and the checklist "
                  f"{args.checklist} does not exist either. Run research-plan's `strobe` "
                  f"step first; a checklist that was never produced is not a checklist "
                  f"that passed.", file=sys.stderr)
            sys.exit(2)
        is_xlsx = args.checklist.endswith((".xlsx", ".xlsm"))
        print(f"[GUARDIAN_WARNING] {rows_json} is {why} — falling back to parsing "
              f"{'the Excel sheet' if is_xlsx else 'the markdown table'} "
              f"({args.checklist}). That parse is positional and would misread a column "
              f"reorder without erroring, and it carries no `kind`, so author-written "
              f"items cannot be told apart from pipeline artifacts. **This report will "
              f"therefore block rather than pass, whatever the sign-off says.** "
              f"Re-run strobe.py to get {os.path.basename(rows_json)}.",
              file=sys.stderr)
        rows = (parse_xlsx if is_xlsx else parse_md)(args.checklist)
        source, read_from = ("xlsx_fallback" if is_xlsx else "markdown_fallback"), args.checklist

    if not rows:
        print(f"[GUARDIAN_BLOCKING] no checklist rows parsed from {read_from}",
              file=sys.stderr)
        sys.exit(2)

    report = build_report(rows, args.signed_off_by, read_from, source)
    report["schema"] = REPORT_SCHEMA
    report["run_id"] = run_id
    report["checklist_sha256"] = sha256_of(read_from)

    # Bind the report to the plan it is about. This one is aggregated as domain 10,
    # so without it a stale `verdict: pass` from an earlier round is indistinguishable
    # from this round's result. `reviewer_challenge.provenance_problems()` blocks on a
    # mismatch and on null — so the honest failure (null + reason) is a block, and
    # omitting the field would be the only way to stay quiet. That is exactly why it
    # is never omitted.
    plan_path, how = resolve_plan_path(args.plan, rows_payload, read_from)
    fingerprint, fp_error = None, None
    if plan_path:
        try:
            with open(plan_path, "r", encoding="utf-8") as f:
                fingerprint, fp_error = plan_fingerprint_of(json.load(f))
        except Exception as e:
            fp_error = f"cannot read {plan_path}: {type(e).__name__}: {e}"
    else:
        fp_error = how
    report["plan"] = os.path.abspath(plan_path) if plan_path else None
    report["plan_source"] = how if plan_path else None
    report["plan_fingerprint"] = fingerprint
    report["plan_fingerprint_error"] = fp_error
    if fingerprint is None:
        print(f"[GUARDIAN_WARNING] plan fingerprint could not be computed ({fp_error}); "
              f"the report is written with plan_fingerprint=null, which the reviewer-"
              f"challenge aggregate treats as unverifiable and blocks on — it is not a "
              f"silent omission.", file=sys.stderr)
    if args.out:
        write_json_atomic(args.out, report)
    print(report["summary"])
    for c in report["comments"]:
        print(f"[GUARDIAN_BLOCKING] {c['check_id']} {c['comment']}", file=sys.stderr)
    sys.exit(1 if report["verdict"] == "fatal" else 0)


if __name__ == "__main__":
    main()
