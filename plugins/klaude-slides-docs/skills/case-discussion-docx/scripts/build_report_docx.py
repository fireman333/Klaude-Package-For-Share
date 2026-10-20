#!/usr/bin/env python3
"""
build_report_docx.py — render a structured JSON CONTENT block into a
書面報告 .docx with proper CJK fonts, inline citations,
and Vancouver-style References.

The JSON CONTENT is produced by Claude (main thread) after writing the
report draft. This script only handles
typography / layout — it does NOT generate prose.

Usage:
  python3 build_report_docx.py --content <path-to-content.json>

CONTENT schema (see SKILL.md):
{
  "title_main":       str   (required)
  "title_sub":        str   (optional)
  "reporter":         str   (required)
  "report_date":      str   (required)  # e.g. "2026/05/04"
  "rotation":         str   (optional)  # e.g. "神經外科實習"
  "cover_page":       bool  (default false)
  "font_body":        str   (default "標楷體") — body CJK font
  "font_ref":         str   (default "Times New Roman") — references font
  "citation_style":   str   (default "vancouver-bracket")
  "sections": [
    {"heading": str, "level": 1|2, "paragraphs": [str, ...]}
  ],
  "references":       [str, ...]   # full pre-formatted citation strings
  "output_path":      str
}
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print(
        json.dumps({
            "error": "python-docx not installed",
            "fix": "pip install python-docx",
        }),
        file=sys.stderr,
    )
    sys.exit(2)


def set_cjk_font(run, font_name: str = "標楷體", size_pt: int = 12) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)


def add_para(
    doc: Document,
    text: str,
    *,
    size: int = 12,
    bold: bool = False,
    font: str = "標楷體",
    align: str | None = None,
    space_before: int = 0,
    space_after: int = 6,
    line_spacing: float = 1.5,
    first_line_indent: float | None = None,
):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    run = p.add_run(text)
    set_cjk_font(run, font_name=font, size_pt=size)
    if bold:
        run.bold = True
    return p


def add_heading(doc, text: str, level: int = 1, font: str = "標楷體"):
    if level == 1:
        return add_para(doc, text, size=16, bold=True, font=font,
                        space_before=18, space_after=8)
    elif level == 2:
        return add_para(doc, text, size=14, bold=True, font=font,
                        space_before=12, space_after=6)
    else:
        return add_para(doc, text, size=12, bold=True, font=font,
                        space_before=8, space_after=4)


def add_body(doc, text: str, font: str = "標楷體"):
    return add_para(doc, text, size=12, font=font, space_after=8,
                    first_line_indent=0.85)


def add_ref(doc, text: str, font: str = "Times New Roman"):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.4
    pf.space_after = Pt(3)
    pf.left_indent = Cm(0.74)
    pf.first_line_indent = Cm(-0.74)
    run = p.add_run(text)
    set_cjk_font(run, font_name=font, size_pt=11)
    return p


def normalize_default_style(doc, font_body: str) -> None:
    style = doc.styles["Normal"]
    style.font.name = font_body
    style.font.size = Pt(12)
    el = style.element.get_or_add_rPr()
    rFonts = el.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        el.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), font_body)


def set_margins(doc) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)


def build_cover_page(doc, content, font_body: str) -> None:
    """Optional cover page — title block top-half, reporter info center."""
    add_para(doc, "", size=12, space_before=120)
    add_para(doc, content["title_main"], size=24, bold=True,
             align="center", font=font_body, space_after=12)
    if content.get("title_sub"):
        add_para(doc, content["title_sub"], size=14,
                 align="center", font=font_body, space_after=180)
    if content.get("rotation"):
        add_para(doc, content["rotation"], size=14,
                 align="center", font=font_body, space_after=24)
    add_para(doc, f"報告者：{content['reporter']}", size=14,
             align="center", font=font_body, space_after=8)
    add_para(doc, f"報告日期：{content['report_date']}", size=14,
             align="center", font=font_body, space_after=8)
    doc.add_page_break()


def build_inline_title(doc, content, font_body: str) -> None:
    """Inline (no cover page) — title block at top of doc."""
    add_para(doc, content["title_main"], size=18, bold=True,
             align="center", font=font_body, space_before=0, space_after=4)
    if content.get("title_sub"):
        add_para(doc, content["title_sub"], size=13,
                 align="center", font=font_body, space_after=12)
    info_bits = [f"報告者：{content['reporter']}",
                 f"報告日期：{content['report_date']}"]
    if content.get("rotation"):
        info_bits.append(content["rotation"])
    info_line = "　　".join(info_bits)
    add_para(doc, info_line, size=12, align="center",
             font=font_body, space_after=18)


def validate_citations(content: dict) -> list[str]:
    """Cross-check inline [N] usage vs References numbering."""
    warnings: list[str] = []
    inline_nums: set[int] = set()
    body = "\n".join(
        " ".join(s.get("paragraphs", []))
        for s in content.get("sections", [])
    )
    for match in re.finditer(r"\[(\d+(?:\s*[,\-]\s*\d+)*)\]", body):
        for piece in re.split(r"[,\-]", match.group(1)):
            piece = piece.strip()
            if piece.isdigit():
                inline_nums.add(int(piece))

    refs = content.get("references", [])
    ref_nums: set[int] = set()
    for r in refs:
        m = re.match(r"\s*(\d+)[\.\)]?\s", r)
        if m:
            ref_nums.add(int(m.group(1)))

    if inline_nums and ref_nums:
        only_inline = sorted(inline_nums - ref_nums)
        only_refs = sorted(ref_nums - inline_nums)
        if only_inline:
            nums = ", ".join(str(n) for n in only_inline)
            warnings.append(
                f"inline [{nums}] used but missing in References"
            )
        if only_refs:
            nums = ", ".join(str(n) for n in only_refs)
            warnings.append(
                f"References [{nums}] not cited in body"
            )
    elif inline_nums and not refs:
        warnings.append(
            f"{len(inline_nums)} inline citations but References list is empty"
        )

    return warnings


def main() -> int:
    ap = argparse.ArgumentParser(description="Build case discussion .docx")
    ap.add_argument("--content", required=True, help="Path to CONTENT JSON")
    args = ap.parse_args()

    json_path = Path(args.content).expanduser()
    if not json_path.is_file():
        print(json.dumps({"error": f"CONTENT JSON not found: {json_path}"}),
              file=sys.stderr)
        return 1

    with json_path.open("r", encoding="utf-8") as f:
        content = json.load(f)

    required = ["title_main", "reporter", "report_date", "sections", "output_path"]
    missing = [k for k in required if k not in content]
    if missing:
        print(json.dumps({"error": f"missing required keys: {missing}"}),
              file=sys.stderr)
        return 1

    font_body = content.get("font_body", "標楷體")
    font_ref = content.get("font_ref", "Times New Roman")
    output_path = Path(content["output_path"]).expanduser()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    normalize_default_style(doc, font_body)
    set_margins(doc)

    if content.get("cover_page"):
        build_cover_page(doc, content, font_body)
    else:
        build_inline_title(doc, content, font_body)

    for section in content["sections"]:
        heading = section.get("heading", "")
        level = int(section.get("level", 1))
        if heading:
            add_heading(doc, heading, level=level, font=font_body)
        for paragraph in section.get("paragraphs", []):
            if paragraph.strip():
                add_body(doc, paragraph, font=font_body)

    refs = content.get("references", [])
    if refs:
        add_heading(doc, "References", level=1, font=font_body)
        for r in refs:
            add_ref(doc, r, font=font_ref)

    doc.save(output_path)

    # Summary stats for stdout JSON
    total_chars = 0
    para_count = 0
    for s in content["sections"]:
        for p in s.get("paragraphs", []):
            total_chars += len(p)
            para_count += 1
    inline_cite_count = 0
    for s in content["sections"]:
        for p in s.get("paragraphs", []):
            inline_cite_count += len(re.findall(r"\[\d+(?:\s*[,\-]\s*\d+)*\]", p))

    warnings = validate_citations(content)

    summary = {
        "output_path": str(output_path),
        "file_size_bytes": output_path.stat().st_size,
        "section_count": len(content["sections"]),
        "paragraph_count": para_count,
        "body_char_count": total_chars,
        "inline_citation_count": inline_cite_count,
        "reference_count": len(refs),
        "citation_warnings": warnings,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if not warnings else 0  # warnings are not fatal


if __name__ == "__main__":
    sys.exit(main())
